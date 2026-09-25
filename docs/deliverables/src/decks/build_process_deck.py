"""Builds the BIBS business process deck (deliverable 41): As-Is, Envisioned, Gaps and Best Practice.

Content lives next to this script:

* ``process_deck/deck.yaml``: title, agenda, the holistic chapter and the closing summary;
* ``process_deck/areas/brdNN_*.yaml``: one file per business area (at a glance, As-Is swimlane,
  pain points, envisioned swimlane, before / after, gaps, best practice, speaker notes).

The swimlanes and enterprise maps are generated as Graphviz sources in ``figures/*.dot`` and
rendered to ``figures/*.png`` with the BDO colour tokens (tools/deliverables).

    python docs/deliverables/src/decks/build_process_deck.py              # deck
    python docs/deliverables/src/decks/build_process_deck.py --previews [dir]   # deck + slide PNGs
    python docs/deliverables/src/decks/build_process_deck.py --figures-only
"""

from __future__ import annotations

import argparse
import math
import tempfile
import sys
from pathlib import Path
from typing import Any

import yaml
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
sys.path.insert(0, str(ROOT / "tools" / "deliverables"))
sys.path.insert(0, str(HERE))

import brand  # noqa: E402
from bdoi_docx import render_dot  # noqa: E402
from bdoi_pptx import MARGIN, W, H, BdoiDeck, _rgb  # noqa: E402
import process_figures as pf  # noqa: E402

DATA = HERE / "process_deck"
FIG = HERE / "figures"
OUT = brand.OUT_DIR / "Decks" / brand.output_name(
    "Deck", "BRD-00", "Business Process AsIs Envisioned BestPractice", "1.0", "pptx")

CONTENT_W = W - 2 * MARGIN
TOP = Inches(1.4)

STATUS = {  # status chip: (label, fill, font)
    "BUILT": ("Built", brand.SUCCESS_BG, brand.SUCCESS),
    "IN BUILD": ("In build", brand.AMBER_BG, brand.AMBER),
    "DESIGNED": ("Designed", brand.BG_BLUE, brand.HEADER_BLUE),
    "PARTLY BUILT": ("Partly built", brand.AMBER_BG, brand.AMBER),
}
SEVERITY = {  # gap severity: (fill, font)
    "High": (brand.DANGER_BG, brand.DANGER),
    "Medium": (brand.AMBER_BG, brand.AMBER),
    "Low": (brand.DIRTY_WHITE, brand.MUTED),
    "Parked": ("EDE7F6", "4527A0"),
    "Answered": (brand.SUCCESS_BG, brand.SUCCESS),
}
PRACTICE = {  # best-practice support: (label, fill, font)
    "BIBS": ("Supported by BIBS", brand.SUCCESS_BG, brand.SUCCESS),
    "PARTLY": ("Partly in BIBS", brand.AMBER_BG, brand.AMBER),
    "ADOPT": ("BDOI to adopt", brand.BG_BLUE, brand.HEADER_BLUE),
}


def load(path: Path) -> Any:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def _shape(s, kind, x, y, w, h):
    return s.shapes.add_shape(kind, int(x), int(y), int(w), int(h))


def _textbox(s, x, y, w, h):
    return s.shapes.add_textbox(int(x), int(y), int(w), int(h))


def _table(s, rows, cols, x, y, w, h):
    return s.shapes.add_table(rows, cols, int(x), int(y), int(w), int(h))


# ====================================================================== deck with extra layouts

class ProcessDeck(BdoiDeck):
    """BdoiDeck plus the pictorial layouts of this deck. Uses the toolkit primitives only."""

    # -------------------------------------------------------------- helpers

    def _text(self, slide, x, y, w, h, *args, **kwargs):
        return super()._text(slide, int(x), int(y), int(w), int(h), *args, **kwargs)

    def _rect(self, slide, x, y, w, h, *args, **kwargs):
        return super()._rect(slide, int(x), int(y), int(w), int(h), *args, **kwargs)

    @property
    def slide(self):
        return self.prs.slides[len(self.prs.slides) - 1]

    def notes(self, text: str | list[str]) -> None:
        if isinstance(text, list):
            text = "\n".join(text)
        self.slide.notes_slide.notes_text_frame.text = text.strip()

    def page(self, title: str):
        s = self._slide()
        self._chrome(s, title)
        return s

    def para(self, s, x, y, w, h, runs, size=12, anchor=MSO_ANCHOR.TOP, align=None, space=4):
        """Text box with paragraphs; each paragraph is a list of (text, bold, colour) runs or a str."""
        box = _textbox(s, x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        first = True
        for p_runs in runs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(space)
            if align:
                p.alignment = align
            if isinstance(p_runs, str):
                p_runs = [(p_runs, False, brand.TEXT)]
            for text, bold, colour in p_runs:
                r = p.add_run()
                r.text = text
                r.font.name, r.font.size, r.font.bold = brand.FONT, Pt(size), bold
                r.font.color.rgb = _rgb(colour)
        return box

    def chip(self, s, x, y, text, fill, font, w=Inches(1.3), h=Inches(0.32), size=11):
        shape = _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        shape.line.color.rgb = _rgb(font)
        shape.line.width = Pt(0.75)
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name, r.font.size, r.font.bold = brand.FONT, Pt(size), True
        r.font.color.rgb = _rgb(font)
        return shape

    def circle(self, s, x, y, text, d=Inches(0.36), fill=brand.YELLOW, font=brand.NEAR_BLACK, size=12):
        shape = _shape(s, MSO_SHAPE.OVAL, x, y, d, d)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        shape.line.color.rgb = _rgb(brand.NEAR_BLACK)
        shape.line.width = Pt(0.75)
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(text)
        r.font.name, r.font.size, r.font.bold = brand.FONT, Pt(size), True
        r.font.color.rgb = _rgb(font)
        return shape

    def box(self, s, x, y, w, h, fill, line=None, radius=True):
        shape = _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
        if radius:
            shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        if line:
            shape.line.color.rgb = _rgb(line)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def arrow(self, s, x, y, w, h, fill=brand.HEADER_BLUE):
        shape = _shape(s, MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def chevrons(self, s, x, y, w, h, labels, fill=brand.BG_BLUE, font=brand.HEADER_BLUE, size=12):
        n = len(labels)
        step = w / n
        for i, label in enumerate(labels):
            shape = _shape(s, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
                                       int(x + i * step), y, int(step + Inches(0.12)), h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill if i % 2 == 0 else brand.WHITE)
            shape.line.color.rgb = _rgb(brand.HEADER_BLUE)
            shape.line.width = Pt(1)
            shape.shadow.inherit = False
            # the label sits in its own text box: a chevron's own text area is too narrow
            tip = Inches(0.2)
            self.para(s, x + i * step + (tip if i else Inches(0.04)), y, step - tip - Inches(0.04), h,
                      [[(label, True, font)]], size=size, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
                      space=0)

    def grid_table(self, s, x, y, w, headers, rows, widths, size=12, row_h=Inches(0.5), fills=None,
                   bold_cols=(), header_fill=brand.HEADER_BLUE):
        """Native table. ``fills`` maps (row, col) -> (fill, font) for heat-map cells."""
        n_rows, n_cols = len(rows) + 1, len(headers)
        shape = _table(s, n_rows, n_cols, x, y, w, row_h * n_rows)
        tbl = shape.table
        scale = w / sum(widths)
        for i, wd in enumerate(widths):
            tbl.columns[i].width = Emu(int(wd * scale))
        tbl.rows[0].height = Inches(0.4)
        for r in range(1, n_rows):
            tbl.rows[r].height = int(row_h)
        for r in range(n_rows):
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                value = headers[c] if r == 0 else rows[r - 1][c]
                cell.margin_left = cell.margin_right = Inches(0.07)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = str(value)
                run.font.name, run.font.size = brand.FONT, Pt(size)
                run.font.bold = r == 0 or c in bold_cols
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = _rgb(header_fill)
                    run.font.color.rgb = _rgb(brand.WHITE)
                elif fills and (r - 1, c) in fills:
                    fill, font = fills[(r - 1, c)]
                    cell.fill.fore_color.rgb = _rgb(fill)
                    run.font.color.rgb = _rgb(font)
                    run.font.bold = True
                    p.alignment = PP_ALIGN.CENTER
                else:
                    cell.fill.fore_color.rgb = _rgb(brand.BG_BLUE if r % 2 == 0 else brand.WHITE)
                    run.font.color.rgb = _rgb(brand.TEXT)
        return tbl

    # -------------------------------------------------------------- layouts

    def diagram(self, title: str, png: Path, caption: str, notes) -> None:
        from PIL import Image

        s = self.page(title)
        area_w, area_h = CONTENT_W, Inches(5.17)
        with Image.open(png) as im:
            iw, ih = im.size
        scale = min(area_w / iw, area_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(str(png), MARGIN + (area_w - w) // 2, Inches(1.36), width=w, height=h)
        if caption:
            self._text(s, MARGIN, Inches(1.36) + h + Inches(0.02), CONTENT_W, Inches(0.34), caption, size=12,
                       italic=True, colour=brand.MUTED, align=PP_ALIGN.CENTER)
        self.notes(notes)

    def agenda(self, title: str, tiles: list[dict]) -> None:
        s = self.page(title)
        cols = 4
        gap = Inches(0.18)
        tw = (CONTENT_W - gap * (cols - 1)) / cols
        th = Inches(1.18)
        for i, t in enumerate(tiles):
            x = MARGIN + (i % cols) * (tw + gap)
            y = TOP + (i // cols) * (th + gap)
            hl = t.get("highlight")
            self.box(s, x, y, tw, th, brand.HEADER_BLUE if hl else brand.WHITE, brand.HEADER_BLUE)
            colour = brand.WHITE if hl else brand.HEADER_BLUE
            self._text(s, x + Inches(0.1), y + Inches(0.06), Inches(0.6), Inches(0.5), t["num"], size=22,
                       bold=True, colour=brand.YELLOW if hl else brand.CTA_BLUE)
            self.para(s, x + Inches(0.62), y + Inches(0.08), tw - Inches(0.7), Inches(0.75),
                      [[(t["title"], True, colour)]], size=14)
            if t.get("brd"):
                self._text(s, x + Inches(0.12), y + th - Inches(0.42), Inches(1.3), Inches(0.35), t["brd"],
                           size=12, colour=brand.WHITE if hl else brand.MUTED)
            if t.get("status"):
                label, fill, font = STATUS[t["status"]]
                self.chip(s, x + tw - Inches(1.42), y + th - Inches(0.42), label, fill, font, w=Inches(1.3))

    def glance(self, a: dict) -> None:
        s = self.page(f"{a['num']}. {a['title']}: at a glance")
        colw = (CONTENT_W - Inches(0.4)) / 3
        y0, hh = TOP, Inches(3.95)
        # card 1: department and roles
        x = MARGIN
        self.box(s, x, y0, colw, hh, brand.WHITE, brand.BORDER)
        self.box(s, x, y0, colw, Inches(0.45), brand.HEADER_BLUE, radius=False)
        self._text(s, x + Inches(0.12), y0, colw, Inches(0.45), "Department and roles", size=14, bold=True,
                   colour=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
        runs = [[(a["department"], True, brand.NEAR_BLACK)]] + [[("•  ", False, brand.HEADER_BLUE),
                                                              (r, False, brand.TEXT)] for r in a["roles"]]
        self.para(s, x + Inches(0.1), y0 + Inches(0.55), colw - Inches(0.2), hh - Inches(0.6), runs, size=12,
                  space=3)
        # card 2: key figures
        x = MARGIN + colw + Inches(0.2)
        self.box(s, x, y0, colw, hh, brand.WHITE, brand.BORDER)
        self.box(s, x, y0, colw, Inches(0.45), brand.HEADER_BLUE, radius=False)
        self._text(s, x + Inches(0.12), y0, colw, Inches(0.45), "Volumes and NFRs (BRD)", size=14, bold=True,
                   colour=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
        figs = a.get("figures", [])
        fy = y0 + Inches(0.55)
        tile_h = Inches(0.8)
        for k, (value, label) in enumerate(figs[:4]):
            ty = fy + k * (tile_h + Inches(0.08))
            self.box(s, x + Inches(0.1), ty, colw - Inches(0.2), tile_h, brand.BG_BLUE)
            self._text(s, x + Inches(0.18), ty, Inches(1.55), tile_h, str(value), size=20, bold=True,
                       colour=brand.HEADER_BLUE, anchor=MSO_ANCHOR.MIDDLE)
            self.para(s, x + Inches(1.72), ty, colw - Inches(1.85), tile_h, [label], size=12,
                      anchor=MSO_ANCHOR.MIDDLE, space=0)
        # card 3: sources and status
        x = MARGIN + 2 * (colw + Inches(0.2))
        self.box(s, x, y0, colw, hh, brand.WHITE, brand.BORDER)
        self.box(s, x, y0, colw, Inches(0.45), brand.HEADER_BLUE, radius=False)
        self._text(s, x + Inches(0.12), y0, colw, Inches(0.45), "BRD and build status", size=14, bold=True,
                   colour=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
        label, fill, font = STATUS[a["status"]]
        self.chip(s, x + Inches(0.12), y0 + Inches(0.58), label, fill, font, w=Inches(1.5), h=Inches(0.36),
                  size=13)
        self.para(s, x + Inches(1.72), y0 + Inches(0.55), colw - Inches(1.8), Inches(0.5),
                  [[(a["status_text"], False, brand.TEXT)]], size=12, anchor=MSO_ANCHOR.MIDDLE)
        runs = [[("•  ", False, brand.HEADER_BLUE), (t, False, brand.TEXT)] for t in a["sources"]]
        self.para(s, x + Inches(0.1), y0 + Inches(1.1), colw - Inches(0.2), hh - Inches(1.2), runs, size=12,
                  space=3)
        # process in one line
        cy = y0 + hh + Inches(0.22)
        self._text(s, MARGIN, cy, CONTENT_W, Inches(0.3), "The process in one line", size=12, bold=True,
                   colour=brand.MUTED)
        self.chevrons(s, MARGIN, cy + Inches(0.32), CONTENT_W, Inches(0.72), a["chevrons"])
        self.notes(a.get("notes", {}).get("glance", "") + "\n\nSources: " + "; ".join(a["sources"]))

    def pain_legend(self, a: dict) -> None:
        s = self.page(f"{a['num']}. {a['title']}: pain points today")
        pains = a["pains"]
        cols = 2
        rows = math.ceil(len(pains) / cols)
        gap = Inches(0.3)
        cw = (CONTENT_W - gap) / cols
        rh = min(Inches(1.25), (Inches(5.35)) / rows)
        for i, p in enumerate(pains):
            c, r = i // rows, i % rows
            x, y = MARGIN + c * (cw + gap), TOP + r * rh
            self.circle(s, x, y + Inches(0.04), p["n"], d=Inches(0.42), size=14)
            self.para(s, x + Inches(0.52), y, cw - Inches(0.55), rh - Inches(0.05),
                      [[(p["title"], True, brand.NEAR_BLACK)],
                       [(p["text"], False, brand.TEXT)],
                       [(p["ref"], False, brand.CTA_BLUE)]], size=12, space=1)
        notes = [a.get("notes", {}).get("pains", "Pain points marked on the As-Is swimlane.")]
        notes += [f"{p['n']}. {p['title']}: {p['text']} ({p['ref']})" for p in pains]
        self.notes(notes)

    def before_after(self, a: dict) -> None:
        s = self.page(f"{a['num']}. {a['title']}: before and after")
        pains = {p["n"]: p for p in a["pains"]}
        rows = a["after"]
        head_y = TOP - Inches(0.02)
        xb, wb = MARGIN + Inches(0.5), Inches(4.0)
        xa, wa = xb + wb + Inches(0.45), Inches(4.9)
        xm, wm = xa + wa + Inches(0.12), CONTENT_W - (xa + wa + Inches(0.12) - MARGIN)
        for x, w, t in ((xb, wb, "Today (pain point)"), (xa, wa, "In BIBS"), (xm, wm, "Measure")):
            self._text(s, x, head_y, w, Inches(0.3), t, size=12, bold=True, colour=brand.HEADER_BLUE)
        rh = min(Inches(0.72), Inches(5.0) / len(rows))
        for i, r in enumerate(rows):
            y = TOP + Inches(0.32) + i * rh
            bh = rh - Inches(0.08)
            self.circle(s, MARGIN, y + (bh - Inches(0.36)) / 2, r["n"], size=12)
            self.box(s, xb, y, wb, bh, brand.DIRTY_WHITE, brand.BORDER)
            self.para(s, xb + Inches(0.05), y, wb - Inches(0.1), bh, [pains[r["n"]]["title"]], size=12,
                      anchor=MSO_ANCHOR.MIDDLE, space=0)
            self.arrow(s, xb + wb + Inches(0.08), y + bh / 2 - Inches(0.12), Inches(0.3), Inches(0.24))
            self.box(s, xa, y, wa, bh, brand.BG_BLUE, brand.CTA_BLUE)
            self.para(s, xa + Inches(0.05), y, wa - Inches(0.1), bh, [r["bibs"]], size=12,
                      anchor=MSO_ANCHOR.MIDDLE, space=0)
            self.para(s, xm, y, wm, bh, [[(r.get("measure") or "-", True, brand.SUCCESS)]], size=12,
                      anchor=MSO_ANCHOR.MIDDLE, space=0)
        notes = [a.get("notes", {}).get("after", "")]
        notes += [f"{r['n']}. {pains[r['n']]['title']} -> {r['bibs']}"
                  + (f" Measure: {r['measure']}." if r.get("measure") else "") for r in rows]
        self.notes(notes)

    def gaps(self, a: dict | None, title: str | None = None, rows_in: list | None = None,
             notes: str = "") -> None:
        rows_in = rows_in if rows_in is not None else a["gaps"]
        title = title or f"{a['num']}. {a['title']}: gaps in the envisioned process"
        s = self.page(title)
        headers = ["Level", "Gap", "Impact", "Owner", "Decision needed", "Ref"]
        rows, fills = [], {}
        for i, g in enumerate(rows_in):
            rows.append([g["sev"], g["gap"], g["impact"], g["owner"], g["decision"], g["ref"]])
            fills[(i, 0)] = SEVERITY[g["sev"]]
        size = 12 if len(rows) <= 5 else 11
        row_h = min(Inches(0.85), Inches(4.9) / max(1, len(rows)))
        self.grid_table(s, MARGIN, TOP, CONTENT_W, headers, rows, [1.0, 3.2, 3.0, 1.9, 3.3, 1.2], size=size,
                        row_h=row_h, fills=fills)
        n = [notes or (a.get("notes", {}).get("gaps", "") if a else "")]
        n += [f"{g['ref']} ({g['sev']}): {g['gap']}. Impact: {g['impact']}. Owner: {g['owner']}. "
              f"Decision: {g['decision']}" for g in rows_in]
        self.notes(n)

    def best_practice(self, a: dict) -> None:
        s = self.page(f"{a['num']}. {a['title']}: best practice")
        items = a["best"]
        cols = 3
        gap = Inches(0.18)
        cw = (CONTENT_W - gap * (cols - 1)) / cols
        rows = math.ceil(len(items) / cols)
        ch = (Inches(5.35) - gap * (rows - 1)) / rows
        for i, b in enumerate(items):
            x = MARGIN + (i % cols) * (cw + gap)
            y = TOP + (i // cols) * (ch + gap)
            self.box(s, x, y, cw, ch, brand.WHITE, brand.BORDER)
            self.box(s, x, y, cw, Inches(0.62), brand.HEADER_BLUE, radius=False)
            self.para(s, x + Inches(0.08), y, cw - Inches(0.16), Inches(0.62),
                      [[(b["practice"], True, brand.WHITE)]], size=13, anchor=MSO_ANCHOR.MIDDLE, space=0)
            label, fill, font = PRACTICE[b["status"]]
            self.chip(s, x + Inches(0.1), y + Inches(0.7), label, fill, font, w=Inches(1.75), h=Inches(0.28),
                      size=11)
            self.para(s, x + Inches(0.08), y + Inches(1.02), cw - Inches(0.16), ch - Inches(1.05),
                      [[("BIBS: ", True, brand.HEADER_BLUE), (b["bibs"], False, brand.TEXT)],
                       [("BDOI: ", True, brand.HEADER_BLUE), (b["bdoi"], False, brand.TEXT)]], size=12, space=4)
        notes = [a.get("notes", {}).get("best", "Practices that fit a Philippine insurance broker, and what "
                                                 "BIBS supports today or by design.")]
        notes += [f"{b['practice']} [{PRACTICE[b['status']][0]}]. BIBS: {b['bibs']} BDOI: {b['bdoi']}"
                  for b in items]
        self.notes(notes)

    def maturity(self, title: str, m: dict, notes: str) -> None:
        s = self.page(title)
        levels = m["levels"]
        label_w = Inches(2.3)
        x0 = MARGIN + label_w
        gw = CONTENT_W - label_w
        colw = gw / len(levels)
        # scale header
        for i, lv in enumerate(levels):
            x = x0 + i * colw
            self.box(s, x + Inches(0.03), TOP, colw - Inches(0.06), Inches(0.62), brand.BG_BLUE, radius=False)
            self.para(s, x + Inches(0.05), TOP, colw - Inches(0.1), Inches(0.62),
                      [[(f"{i + 1}  ", True, brand.HEADER_BLUE), (lv, False, brand.NEAR_BLACK)]], size=11,
                      anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=0)
        themes = m["themes"]
        rh = Inches(3.7) / len(themes)
        y0 = TOP + Inches(0.72)
        marks = [("now", brand.MUTED, "Today"), ("bibs", brand.CTA_BLUE, "BIBS envisioned"),
                 ("best", brand.YELLOW, "Best practice")]
        for k, t in enumerate(themes):
            y = y0 + k * rh
            if k % 2 == 0:
                self._rect(s, MARGIN, y, CONTENT_W, rh, brand.DIRTY_WHITE)
            self._text(s, MARGIN + Inches(0.05), y, label_w, rh, t["theme"], size=13, bold=True,
                       colour=brand.HEADER_BLUE, anchor=MSO_ANCHOR.MIDDLE)

            def cx(v):
                return int(x0 + (v - 0.5) * colw)

            d = Inches(0.3)
            cy = int(y + rh / 2)
            self._rect(s, cx(t["now"]), cy - Pt(1.5), cx(t["bibs"]) - cx(t["now"]), Pt(3), brand.FIELD_BLUE)
            for key, colour, _ in marks:
                v = t[key]
                off = {"now": 0, "bibs": 0, "best": 0}[key]
                self.circle(s, cx(v) - d // 2 + off, cy - d // 2, "", d=d, fill=colour, size=9)
        # legend and basis
        ly = y0 + len(themes) * rh + Inches(0.12)
        lx = MARGIN
        for _, colour, text in marks:
            self.circle(s, lx, ly, "", d=Inches(0.26), fill=colour)
            self._text(s, lx + Inches(0.3), ly - Inches(0.03), Inches(2.2), Inches(0.32), text, size=12,
                       colour=brand.TEXT)
            lx += Inches(2.3)
        self.para(s, MARGIN + Inches(7.0), ly - Inches(0.05), CONTENT_W - Inches(7.0), Inches(0.5),
                  [[(m["basis"], False, brand.MUTED)]], size=11)
        n = [notes] + [f"{t['theme']}: today {t['now']}, BIBS {t['bibs']}, best practice {t['best']}. {t['why']}"
                       for t in themes]
        self.notes(n)

    def heat_counts(self, title: str, rows: list[list], notes: str) -> None:
        """BRD x severity counts from the discrepancy register, as a heat map."""
        s = self.page(title)
        headers = ["Area", "BRD", "High", "Medium", "Low", "Open", "Answered"]
        fills = {}
        mx = max(r[2] for r in rows) or 1
        for i, r in enumerate(rows):
            for c, (base, font) in ((2, (brand.DANGER_BG, brand.DANGER)), (3, (brand.AMBER_BG, brand.AMBER)),
                                    (4, (brand.DIRTY_WHITE, brand.MUTED))):
                if r[c]:
                    fills[(i, c)] = (base, font)
            if r[2] >= 0.75 * mx:
                fills[(i, 2)] = (brand.DANGER, brand.WHITE)
        self.grid_table(s, MARGIN, TOP, Inches(8.2), headers, rows, [3.4, 1.2, 0.9, 0.9, 0.9, 0.9, 1.1],
                        size=12, row_h=Inches(0.36), fills=fills)
        self.para(s, MARGIN + Inches(8.45), TOP, CONTENT_W - Inches(8.45), Inches(5.3), notes.split("\n"),
                  size=12, space=6)
        self.notes(notes + "\nSource: docs/deliverables/src/registers/discrepancy_register.yaml. An item "
                           "that touches several BRDs is counted under each of them.")

    def roadmap(self, title: str, r: dict) -> None:
        s = self.page(title)
        cols = r["columns"]
        gap = Inches(0.2)
        cw = (CONTENT_W - gap * (len(cols) - 1)) / len(cols)
        ch = Inches(3.1)
        for i, c in enumerate(cols):
            x = MARGIN + i * (cw + gap)
            label, fill, font = STATUS[c["status"]]
            self.box(s, x, TOP, cw, ch, brand.WHITE, brand.BORDER)
            self.box(s, x, TOP, cw, Inches(0.5), fill, font, radius=False)
            self._text(s, x + Inches(0.1), TOP, cw, Inches(0.5), c["title"], size=15, bold=True, colour=font,
                       anchor=MSO_ANCHOR.MIDDLE)
            runs = [[("•  ", False, brand.HEADER_BLUE), (t, False, brand.TEXT)] for t in c["items"]]
            self.para(s, x + Inches(0.1), TOP + Inches(0.58), cw - Inches(0.2), ch - Inches(0.6), runs, size=12,
                      space=3)
            if i < len(cols) - 1:
                self.arrow(s, x + cw - Inches(0.02), TOP + Inches(0.1), gap + Inches(0.04), Inches(0.3))
        y = TOP + ch + Inches(0.18)
        self.box(s, MARGIN, y, CONTENT_W, Inches(5.45) - ch - Inches(0.18), brand.BG_BLUE, brand.HEADER_BLUE)
        self._text(s, MARGIN + Inches(0.12), y + Inches(0.05), CONTENT_W, Inches(0.35),
                   r["decisions_title"], size=14, bold=True, colour=brand.HEADER_BLUE)
        half = math.ceil(len(r["decisions"]) / 2)
        for k in range(2):
            items = r["decisions"][k * half:(k + 1) * half]
            runs = [[("•  ", False, brand.HEADER_BLUE), (t, False, brand.NEAR_BLACK)] for t in items]
            self.para(s, MARGIN + Inches(0.12) + k * CONTENT_W / 2, y + Inches(0.42), CONTENT_W / 2 - Inches(0.25),
                      Inches(1.7), runs, size=12, space=3)
        self.notes(r.get("notes", ""))


# ====================================================================== figures

def render(name: str, text: str) -> Path:
    path = pf.write(FIG / f"{name}.dot", text)
    return render_dot(path, dpi=200)


def area_figures(a: dict) -> tuple[Path, Path]:
    base = f"{a['id']}"
    asis = render(f"{base}_asis", pf.swimlane(f"{base}_asis", a["asis"], "asis",
                                              f"{a['brd']} {a['title']}: As-Is swimlane. {a['asis']['caption']}"))
    tobe = render(f"{base}_envisioned", pf.swimlane(f"{base}_envisioned", a["tobe"], "tobe",
                                                    f"{a['brd']} {a['title']}: envisioned swimlane in BIBS. "
                                                    f"{a['tobe']['caption']}"))
    return asis, tobe


def asis_map_spec(m: dict, areas: dict[str, dict]) -> dict:
    """Areas above and below, today's systems in the middle; one line per dependency."""
    nodes, edges = [], []
    w, h, gx = 136.0, 50.0, 145.4
    for key, y in (("areas_top", 26.0), ("areas_bottom", 274.0)):
        for i, aid in enumerate(m[key]):
            a = areas.get(aid)
            if a is None:  # chapter not written yet
                continue
            nodes.append({"id": aid, "x": 6 + i * gx, "y": y, "w": w, "h": h,
                          "text": f"{a['num']}  {a['title']}", "kind": "area", "size": 12,
                          "pain": [len(a["pains"])]})
    sw, sg = 116.0, 124.6
    for i, sy in enumerate(m["systems"]):
        nodes.append({"id": sy["id"], "x": 4 + i * sg, "y": 150.0, "w": sw, "h": 54, "text": sy["text"],
                      "kind": "legacy", "size": 12})
    top = set(m["areas_top"])
    for aid, targets in m["links"].items():
        for t in targets if aid in areas else []:
            ports = ("s", "n") if aid in top else ("n", "s")
            edges.append({"a": aid, "b": t, "style": "solid", "colour": "@MUTED", "arrow": False, "ports": ports,
                          "width": 0.9})
    bands = [{"x": 0, "y": 0, "w": pf.CANVAS_W, "h": 84, "fill": "@WHITE", "label": "Business areas 1-6",
              "label_w": 200},
             {"x": 0, "y": 124, "w": pf.CANVAS_W, "h": 90, "fill": "#EEF3F8",
              "label": "Today's systems and channels", "label_w": 260},
             {"x": 0, "y": 248, "w": pf.CANVAS_W, "h": 84, "fill": "@WHITE", "label": "Business areas 7-12",
              "label_w": 200}]
    return {"bands": bands, "nodes": nodes, "edges": edges,
            "legend": [["area", "Business area"], ["legacy", "Today's system or channel"],
                       ["pain", "Number of pain points in the area chapter"]]}


def envisioned_map_spec(m: dict) -> dict:
    nodes, edges = [], []
    ex_w, ex_h = 132.0, 52.0
    for i, x in enumerate(m["externals"]):
        nodes.append({"id": x["id"], "x": 0, "y": 6 + i * 65.0, "w": ex_w, "h": ex_h, "text": x["text"],
                      "sub": x.get("sub"), "kind": x.get("kind", "external"), "size": 12})
    nodes.append({"id": "chan", "x": 160, "y": 6, "w": 70, "h": 312, "text": m["channels"], "kind": "screen",
                  "size": 12})
    for xid, label in m["links"]:
        edges.append({"a": xid, "b": "chan", "style": "dashed" if label == "parked" else "solid",
                      "ports": ("e", "w")})
    x = 244.0
    for i, g in enumerate(m["groups"]):
        gid = f"g{i}"
        nodes.append({"id": gid, "x": x, "y": 26, "w": g["w"], "h": 38, "text": g["title"], "kind": "group",
                      "size": 11.5})
        nodes.append({"id": gid + "m", "x": x, "y": 68, "w": g["w"], "h": 148, "text": g["modules"],
                      "kind": "module", "size": 11})
        x += g["w"] + 4
    n = len(m["services"])
    sw = (pf.CANVAS_W - 244 - 4 * (n - 1)) / n
    for i, sv in enumerate(m["services"]):
        nodes.append({"id": f"s{i}", "x": 244 + i * (sw + 4), "y": 246, "w": sw, "h": 44, "text": sv,
                      "kind": "service", "size": 11})
    nodes.append({"id": "host", "x": 244, "y": 294, "w": pf.CANVAS_W - 244, "h": 24, "text": m["hosting"],
                  "kind": "stage", "size": 11})
    bands = [{"x": 238, "y": 0, "w": pf.CANVAS_W - 238, "h": 222, "fill": "@DIRTY_WHITE",
              "label": "BIBS modules by navigation group", "label_w": 300},
             {"x": 238, "y": 224, "w": pf.CANVAS_W - 238, "h": 98, "fill": "@BG_BLUE",
              "label": "Shared services and hosting", "label_w": 300}]
    return {"bands": bands, "nodes": nodes, "edges": edges,
            "legend": [["external", "External party"], ["group", "Navigation group"],
                       ["service", "Shared service"], ["parked", "Parked interface"]]}


def value_chain_spec(m: dict) -> dict:
    nodes, edges = [], []
    w, h, gx = 132.0, 104.0, 147.4
    stages = m["stages"]
    for i, st in enumerate(stages):
        row, k = divmod(i, 6)
        col = k if row == 0 else 5 - k
        status = st["status"].lower()
        kind = "vc_build" if "in build" in status else ("vc_designed" if status == "designed" else "vc_built")
        nodes.append({"id": f"v{i}", "x": 4 + col * gx, "y": 22 + row * 170, "w": w, "h": h, "text": st["text"],
                      "sub": [st["owner"], "BIBS: " + st["module"], f"({st['status']})"], "kind": kind,
                      "size": 13})
    for i in range(len(stages) - 1):
        edges.append({"a": f"v{i}", "b": f"v{i + 1}", "width": 2.0})
    return {"nodes": nodes, "edges": edges,
            "legend": [["vc_built", "Built (at least in part)"], ["vc_build", "In build"],
                       ["vc_designed", "Designed, not yet built"]]}


def holistic_figures(h: dict, areas: list[dict]) -> dict[str, Path]:
    by_key = {a["key"]: a for a in areas}
    specs = {"asis_map": asis_map_spec(h["asis_map"], by_key),
             "envisioned_map": envisioned_map_spec(h["envisioned_map"]),
             "value_chain": value_chain_spec(h["value_chain"])}
    return {key: render(f"holistic_{key}", pf.free_layout(f"holistic_{key}", spec, h[key]["caption"]))
            for key, spec in specs.items()}


# ====================================================================== build

def register_counts(areas: list[dict]) -> list[list]:
    reg = load(ROOT / "docs/deliverables/src/registers/discrepancy_register.yaml")
    rows = []
    for a in areas:
        key = "BRD-%02d" % int(a["brd"].split("-")[1])
        items = [i for i in reg["items"] if key in i["brds"]]
        rows.append([a["title"], a["brd"], sum(i["sev"] == "High" for i in items),
                     sum(i["sev"] == "Medium" for i in items), sum(i["sev"] == "Low" for i in items),
                     sum(i["status"] == "Open" for i in items), sum(i["status"] != "Open" for i in items)])
    return rows


def build(previews: str | None = None, figures_only: bool = False) -> Path:
    d = load(DATA / "deck.yaml")
    areas = [load(p) for p in sorted((DATA / "areas").glob("brd*.yaml"))]
    areas.sort(key=lambda a: a["num"])
    hol = holistic_figures(d["holistic"], areas)
    figs = {a["id"]: area_figures(a) for a in areas}
    if figures_only:
        return FIG

    deck = ProcessDeck(d["title"], version=d["version"], date=d["date"], subtitle=d["subtitle"])
    deck.title()
    deck.notes(d["notes"]["title"])
    tiles = [{"num": "0", "title": "Holistic view: the enterprise today, in BIBS and at best practice",
              "highlight": True}]
    tiles += [{"num": str(a["num"]), "title": a["title"], "brd": a["brd"], "status": a["status"]} for a in areas]
    tiles += [{"num": str(len(areas) + 1), "title": "Summary: maturity, decisions and roadmap", "highlight": True}]
    deck.agenda("Agenda", tiles)
    deck.notes(d["notes"]["agenda"])
    deck.content("How to read the process slides", d["reading"], size=16)
    deck.notes(d["notes"]["reading"])

    h = d["holistic"]
    deck.section("0. Holistic view", h["subtitle"])
    deck.notes(h["notes"])
    deck.diagram(h["asis_map"]["title"], hol["asis_map"], h["asis_map"]["caption"], h["asis_map"]["notes"])
    deck.diagram(h["envisioned_map"]["title"], hol["envisioned_map"], h["envisioned_map"]["caption"],
                 h["envisioned_map"]["notes"])
    deck.diagram(h["value_chain"]["title"], hol["value_chain"], h["value_chain"]["caption"],
                 h["value_chain"]["notes"])
    deck.maturity(h["maturity"]["title"], h["maturity"], h["maturity"]["notes"])
    deck.heat_counts(h["register"]["title"], register_counts(areas), h["register"]["text"])
    deck.gaps(None, title=h["decisions"]["title"], rows_in=h["decisions"]["rows"], notes=h["decisions"]["notes"])
    deck.roadmap(h["roadmap"]["title"], h["roadmap"])

    for a in areas:
        deck.section(f"{a['num']}. {a['title']}", f"{a['brd']} | {a['scope']}")
        deck.notes(a.get("notes", {}).get("section", f"{a['brd']} {a['title']}: {a['scope']}."))
        deck.glance(a)
        asis, tobe = figs[a["id"]]
        deck.diagram(f"{a['num']}. {a['title']}: As-Is process", asis, a["asis"]["caption"],
                     a["asis"]["notes"])
        deck.pain_legend(a)
        deck.diagram(f"{a['num']}. {a['title']}: envisioned in BIBS", tobe, a["tobe"]["caption"],
                     a["tobe"]["notes"])
        deck.before_after(a)
        deck.gaps(a)
        deck.best_practice(a)

    c = d["closing"]
    deck.section(f"{len(areas) + 1}. Summary", c["subtitle"])
    deck.notes(c["notes"])
    deck.maturity(c["maturity_title"], h["maturity"], c["maturity_notes"])
    deck.gaps(None, title=c["decisions_title"], rows_in=h["decisions"]["rows"], notes=c["decisions_notes"])
    deck.roadmap(c["roadmap_title"], h["roadmap"])
    deck.content(c["next_title"], c["next_steps"], size=16)
    deck.notes(c["next_notes"])

    out = deck.save(OUT)
    print(f"{out} ({deck._n} slides)")
    if previews:
        import render as rnd

        # previews are never written into the repository: they go to a scratch folder
        pdf = rnd.to_pdf(out, outdir=Path(previews))
        sheets = rnd.previews(pdf, dpi=60)
        print(f"previews: {sheets[0].parent if sheets else '-'}")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--previews", nargs="?", const=str(Path(tempfile.gettempdir()) / "bibs_process_deck"),
                    help="render slide PNGs and contact sheets into this folder (default: a temp folder)")
    ap.add_argument("--figures-only", action="store_true", help="only generate and render the figures")
    args = ap.parse_args()
    build(previews=args.previews, figures_only=args.figures_only)
    return 0


if __name__ == "__main__":
    sys.exit(main())
