"""Builds the BRD-13 Data Migration document set (client deliverable 29 and the migration templates of item 20).

Usage
  python docs/deliverables/src/migration/build_migration_pack.py --check      # checks only
  python docs/deliverables/src/migration/build_migration_pack.py              # build everything
  python docs/deliverables/src/migration/build_migration_pack.py --previews   # also page previews

Inputs (this folder)
  * dm_layouts.yaml   data objects, extract layouts, code map sets, data-quality rules, open decisions;
  * dm_cutover.yaml   cutover tasks, freeze windows, go / no-go checkpoints, rollback, communication,
                      hypercare and decommissioning;
  * DM_STRATEGY_AND_APPROACH.md, DM_CUTOVER_RUNBOOK.md, DM_RECONCILIATION_AND_SIGNOFF.md
                      Word sources in the bdoi_docx format; lines <!-- dm:<name> [opt=value] --> are
                      replaced by tables built from the YAML files (list in PLACEHOLDERS below).

Outputs (docs/deliverables/out/Drop-0_Setup_and_Data_Migration/Migration, from tools/deliverables/brand.py)
  * BIBS_Migration_BRD-13_Data_Requirements_Workbook_v<version>.xlsx
  * templates/<LAYOUT>_template.csv (header row = field names of the workbook), the Excel template
    P03_template.xlsx for the Renewal processing team, CONTROL_template.ctl.csv and README.md
  * BIBS_Migration_BRD-13_Cutover_Task_Plan_v<version>.xlsx
  * the three Word documents named in the "output" of their front matter.
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
from pathlib import Path
from typing import Any

import yaml

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402
from bdoi_docx import BdoiDocument, lint_source, load_source, meta_from, output_path, render_body  # noqa: E402
from bdoi_xlsx import BdoiWorkbook, Column  # noqa: E402

OUT = brand.out_dir("BRD-13", "Migration")  # Drop-0_Setup_and_Data_Migration/Migration (brand.BRD_DROP)
TEMPLATES = OUT / "templates"
DOCS = ["DM_STRATEGY_AND_APPROACH.md", "DM_RECONCILIATION_AND_SIGNOFF.md", "DM_CUTOVER_RUNBOOK.md"]
FIELD_KEYS = ["name", "description", "type", "length", "mandatory", "values", "format", "example", "target", "rule",
              "hint"]
TYPES = {"Text", "Code", "Date", "Timestamp", "Amount", "Integer", "Decimal", "Flag"}
CLASSES = ["MIGRATE", "CARRY_FORWARD", "CONDITIONAL", "ARCHIVE", "EXCLUDED"]
CLASS_LABEL = {"MIGRATE": "Migrate", "CARRY_FORWARD": "Carry forward", "CONDITIONAL": "Conditional",
               "ARCHIVE": "Archive", "EXCLUDED": "Excluded"}
SYSTEMS = ["EBIX", "QPS", "ISYS", "CMS", "Excel", "File shares"]
FILLER = ("seamless", "robust", "comprehensive", "leverage", "cutting-edge", "state-of-the-art", "best-in-class",
          "world-class")


# ------------------------------------------------------------------------------------------------ model

def load() -> tuple[dict[str, Any], dict[str, Any]]:
    lay = yaml.safe_load((HERE / "dm_layouts.yaml").read_text(encoding="utf-8"))
    cut = yaml.safe_load((HERE / "dm_cutover.yaml").read_text(encoding="utf-8"))
    for code, layout in lay["layouts"].items():
        layout["code"] = code
        layout["fields"] = parse_fields(layout["fields"], code)
    lay["control"]["fields"] = parse_fields(lay["control"]["fields"], "CONTROL")
    return lay, cut


def parse_fields(text: str, where: str) -> list[dict[str, str]]:
    fields = []
    for n, line in enumerate(text.strip().splitlines(), start=1):
        parts = [p.strip().replace("\\PIPE", "|") for p in line.replace("\\|", "\\PIPE").split(" | ")]
        if len(parts) != len(FIELD_KEYS):
            raise SystemExit(f"{where} field line {n}: {len(parts)} parts, expected {len(FIELD_KEYS)}: {line[:80]}")
        fields.append(dict(zip(FIELD_KEYS, parts)))
    return fields


def milestone_text(lay: dict[str, Any], ref: str) -> str:
    """'M4; final at T-2' -> 'M4 (T-14 weeks); final at T-2'."""
    names = {m[0]: m[2] for m in lay["milestones"]}
    return re.sub(r"\bM(\d)\b", lambda m: f"M{m.group(1)} ({names.get('M' + m.group(1), '?')})", str(ref))


def check(lay: dict[str, Any], cut: dict[str, Any]) -> list[str]:
    problems: list[str] = []
    objects = {o["code"]: o for o in lay["objects"]}
    for o in lay["objects"]:
        if o["decision"] not in CLASSES:
            problems.append(f"object {o['code']}: decision {o['decision']} not in {CLASSES}")
        for code in o["layouts"]:
            if code not in lay["layouts"]:
                problems.append(f"object {o['code']}: layout {code} does not exist")
        if o["decision"] in ("MIGRATE", "CARRY_FORWARD") and not o["layouts"] and o["code"] != "F05":
            problems.append(f"object {o['code']}: {o['decision']} without a layout")
    used = {c for o in lay["objects"] for c in o["layouts"]}
    for code, layout in lay["layouts"].items():
        if code not in used:
            problems.append(f"layout {code}: not referenced by any object")
        if layout["object"] not in objects:
            problems.append(f"layout {code}: object {layout['object']} unknown")
        names = [f["name"] for f in layout["fields"]]
        if len(names) != len(set(names)):
            problems.append(f"layout {code}: duplicate field names")
        for key in layout["key"]:
            if key not in names:
                problems.append(f"layout {code}: key {key} is not a field")
        for amount in layout["amounts"]:
            if amount not in names:
                problems.append(f"layout {code}: amount {amount} is not a field")
        for f in layout["fields"]:
            if f["type"] not in TYPES:
                problems.append(f"layout {code}.{f['name']}: type {f['type']}")
            if f["mandatory"] not in ("Y", "N", "C"):
                problems.append(f"layout {code}.{f['name']}: mandatory {f['mandatory']}")
            if not re.fullmatch(r"[a-z][a-z0-9_]*", f["name"]):
                problems.append(f"layout {code}.{f['name']}: field names are lower snake case")
            if f["type"] == "Amount" and f["name"] not in layout["amounts"] and code not in ("F01", "F03", "F04"):
                problems.append(f"layout {code}.{f['name']}: amount column without a control total")
    ids = [t[0] for t in cut["tasks"]]
    if len(ids) != len(set(ids)):
        problems.append("tasks: duplicate ids")
    for t in cut["tasks"]:
        if len(t) != 10:
            problems.append(f"task {t[0]}: {len(t)} items, expected 10")
            continue
        for pred in [p.strip() for p in str(t[7]).split(",") if p.strip()]:
            if pred not in ids:
                problems.append(f"task {t[0]}: predecessor {pred} unknown")
            elif ids.index(pred) > ids.index(t[0]):
                problems.append(f"task {t[0]}: predecessor {pred} comes later in the list")
        if t[5] not in {r[0] for r in cut["roles"]}:
            problems.append(f"task {t[0]}: owner {t[5]} not in roles")
    for r in cut["rollback"]:
        if r[3] not in {x[0] for x in cut["roles"]}:
            problems.append(f"rollback {r[0]}: owner {r[3]} not in roles")
    decision_ids = [d[0] for d in lay["decisions"]]
    for o in lay["objects"]:
        for q in re.findall(r"(DMQ\d\d|DCR-\d+)", str(o["questions"])):
            if q not in decision_ids:
                problems.append(f"object {o['code']}: question {q} is not in decisions")
    text = yaml.safe_dump(lay) + yaml.safe_dump(cut)
    for word in FILLER:
        if re.search(rf"\b{word}\b", text, re.I):
            problems.append(f"filler word: {word}")
    return problems


# ------------------------------------------------------------------------------------------------ workbook

SHORT = {"R01": "LOV and MIS values", "R02": "Branches", "R03": "Sales organisation", "R05": "Products and risk codes",
         "R06": "PACKAGE map", "R07": "Commission rates", "R08": "GL account map", "R11": "Receipt series",
         "C02": "Client addresses", "C03": "Payout accounts", "P01": "Policy headers", "P01S": "Policy shares",
         "P03": "RA already sent", "F01": "Invoice header", "F01S": "Invoice shares", "F01C": "Invoice components",
         "F02": "UPP", "F06": "PDCs pick-ups refunds", "G01": "GL trial balance", "G03": "True-up journals", "G03D": "True-up item detail", "H01": "Archive records",
         "H02": "Archive documents"}


def sheet_name(layout: dict[str, Any]) -> str:
    title = SHORT.get(layout["code"], layout["title"])
    title = title.replace("Open legacy invoices - ", "Invoices ").replace(" (", " ").replace(")", "")
    name = f"{layout['code']} {title}"
    while len(name) > 31:
        name = name.rsplit(" ", 1)[0]
    return name


def field_columns() -> list[Column]:
    return [
        Column("seq", "#", 5, "Column position in the file"),
        Column("name", "Field name", 22, "Column name in the header row of the file, exactly as written (lower case)"),
        Column("description", "Description", 40, "What the field holds"),
        Column("type", "Type", 11, "Text, Code (a coded value), Date, Timestamp, Amount (2 decimals), Integer, "
               "Decimal, Flag (Y or N)", values=sorted(TYPES)),
        Column("length", "Length", 8, "Maximum characters; for numbers the precision and decimals (19,2)"),
        Column("mandatory", "Mandatory", 10, "Y = always; N = optional; C = conditional (see the validation rule)",
               values=["Y", "N", "C"]),
        Column("values", "Allowed values / code list", 30, "Closed list of values, or the code map set that maps "
               "the legacy value to BIBS"),
        Column("format", "Format", 20, "How the value is written in the file"),
        Column("example", "Example", 20, "A made-up example value"),
        Column("target", "BIBS target field", 32, "Table and column (or service field) in BIBS; '(new)' = added by "
               "the migration build"),
        Column("rule", "Validation rule", 36, "Check BIBS applies in staging (data-quality rules catalogue)"),
        Column("hint", "Source system hint", 24, "Where the value is expected to come from in legacy; BDOI IT "
               "confirms"),
    ]


def register_columns() -> list[Column]:
    return [
        Column("code", "Object", 8, "Data object code"),
        Column("name", "Data object", 34, "What the object holds"),
        Column("category", "Category", 11, "Reference, Client, Policy, Open item, GL, History",
               values=["Reference", "Client", "Policy", "Open item", "GL", "History"]),
        Column("decision", "Decision (proposed)", 15, "Migrate / Carry forward / Conditional / Archive / Excluded; "
               "BDOI confirms at gate G1", values=CLASSES),
        Column("rationale", "Rationale", 44, "Why this class"),
        Column("source", "Source system", 14, "Legacy system(s) that hold the object"),
        Column("layouts", "Templates", 11, "Extract layouts (sheets of this workbook, CSV templates)"),
        Column("target", "BIBS target", 32, "Where the object lands in BIBS"),
        Column("owner", "Business owner (sign-off)", 24, "Signs gates G1, G2 and G6"),
        Column("steward", "Data steward", 20, "Resolves data-quality issues and prepares code maps"),
        Column("volume", "Volume estimate (BDOI to fill)", 14, "Row count per source system at the last extract"),
        Column("bound", "Planning bound", 18, "Known bound from the umbrella BRD or the design"),
        Column("history", "History depth", 22, "Which records and how far back"),
        Column("delta", "Delta frequency", 22, "Extracts after the first full one"),
        Column("format", "Extract format", 12, "CSV or XLSX, with a control file"),
        Column("controls", "Control totals required", 28, "Measures of the control file"),
        Column("due", "Due", 22, "Milestone of the first full extract and the final one"),
        Column("order", "Load order", 7, "Lower loads first; 0 = not loaded"),
        Column("questions", "Open decisions", 14, "Questions that affect the object"),
        Column("status", "Status", 12, "Progress of the object", values=["OPEN", "IN PROGRESS", "DONE", "N/A"],
               status=True),
    ]


def register_rows(lay: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for o in lay["objects"]:
        r = dict(o)
        r["layouts"] = ", ".join(o["layouts"]) or "-"
        r["due"] = milestone_text(lay, o["due"])
        r["volume"] = ""
        r["status"] = "N/A" if o["decision"] == "EXCLUDED" else "OPEN"
        rows.append(r)
    return rows


def control_rows(lay: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for code, layout in lay["layouts"].items():
        base = {"layout": code, "object": layout["object"]}
        rows.append({**base, "measure": "ROW_COUNT", "column": "-", "currency": "-", "rule": "Rows in the data "
                     "file, header excluded"})
        key = ", ".join(layout["key"])
        rows.append({**base, "measure": "HASH_TOTAL", "column": key, "currency": "-", "rule": layout["hash"]})
        for amount in layout["amounts"]:
            rule = "Sum per currency"
            if code == "F01C":
                rule = "Sum per currency and per component (filter component=<value>)"
            rows.append({**base, "measure": "AMOUNT_TOTAL", "column": amount, "currency": "Each currency",
                         "rule": rule})
        rows.append({**base, "measure": "SHA256", "column": "-", "currency": "-", "rule": "SHA-256 of the data "
                     "file as sent"})
    return rows


def build_workbook(lay: dict[str, Any]) -> Path:
    m = lay["meta"]
    wb = BdoiWorkbook("BDOI Data Requirements Workbook", doc_type="Data migration requirements", brd="BRD-13",
                      version=m["version"], date=m["date"],
                      subtitle="BRD-13 Data Migration - what BDOI extracts, in which form, and by when")
    wb.legend = [("OPEN", "Not yet supplied or decided"), ("IN PROGRESS", "Being prepared"),
                 ("DONE", "Supplied, checked and signed"), ("N/A", "Not applicable (excluded object)")]
    wb.cover_notes = [
        "Each layout sheet is the contract of one extract file; the CSV template of the same code in the "
        "templates folder has exactly these field names as its header row.",
        "Volumes, owners and stewards marked 'to name' or blank are filled by BDOI (DMQ01, DMQ28).",
        "Sources: BDOI_DM_BRD_SPEC.md and DATA_MIGRATION_DESIGN.md (sections 2, 5, 7, 8, 12, 15, 17, 27); dates on "
        "the BDOI programme timeline (go-live January 2028).",
    ]
    wb.sheet("Object Register", register_columns(), register_rows(lay),
             description="One row per data object with the proposed decision (BRID 1.1a); BDOI fills the volumes")
    contract = [
        ("File name", "<LAYOUT>_<SOURCE>_<yyyyMMdd>_<nn>.csv (or .xlsx), for example F01C_EBIX_20271231_01.csv; "
         "the date is the as-of date; nn is the sequence of that day"),
        ("One file per", "Layout, source system and extract (as-of date and sequence)"),
        ("CSV", "UTF-8 without BOM; comma separator; RFC 4180 quoting (double quotes around values that contain a "
         "comma, quote or line break); one header row with the field names of the layout, in any order; no "
         "total or blank rows"),
        ("XLSX", "First sheet only; header in row 1; no merged cells, formulas, hidden rows or colour-coded "
         "meaning"),
        ("Dates and times", "yyyy-MM-dd; timestamps yyyy-MM-dd HH:mm:ss in Philippine time"),
        ("Amounts", "Dot decimal, 2 decimals, no thousands separator, minus sign for negatives (return "
         "invoices), no currency symbol"),
        ("Currency", "ISO 4217 code (PHP, USD) in the currency column"),
        ("Codes", "Exactly as stored in legacy; BIBS maps them through the code maps; never translate codes in "
         "the extract"),
        ("Blank", "Empty value = no value; do not write NULL, N/A or 0 for an unknown value"),
        ("Control file", "<data file name>.ctl.csv in the CONTROL layout: row count, hash total of the key, "
         "amount totals per column and currency, SHA-256 of the data file (sheet Control File)"),
        ("Delivery", "Console upload or the SFTP drop agreed with BDOI IT; never e-mail; files land in the "
         "encrypted intake bucket (AWS ap-southeast-1)"),
        ("Personal data", "Production extracts are loaded only in production; every non-production environment "
         "receives masked data (masking at intake); staging and files are purged within 5 days of the batch "
         "sign-off"),
        ("Layout changes", "A layout is frozen at gate G2; a file whose header does not match the current layout "
         "version is rejected; changes go through the Data Migration Lead"),
        ("Rejections", "A file is rejected as a whole when its checksum, header, row count, amount totals or hash "
         "total do not match its control file; BIBS gives the reason and the difference"),
    ]
    wb.sheet("File Contract", [Column("item", "Item", 18, "Topic of the rule"),
                               Column("rule", "Rule", 110, "What every extract file must follow (R3 section 5.1)")],
             [{"item": a, "rule": b} for a, b in contract],
             description="Rules for every extract and control file (FR-DM-010)")
    for code, layout in lay["layouts"].items():
        rows = [{"seq": i, **f} for i, f in enumerate(layout["fields"], start=1)]
        desc = (f"Layout {code} of object {layout['object']}; key: {', '.join(layout['key'])}; template "
                f"templates/{code}_template.csv")
        wb.sheet(sheet_name(layout), field_columns(), rows, description=desc)
    rows = [{"seq": i, **f} for i, f in enumerate(lay["control"]["fields"], start=1)]
    wb.sheet("Control File", field_columns(), rows,
             description="Layout of the control file sent with every data file (one row per measure)")
    ctl_cols = [Column(f["name"], f["name"], w, f["description"]) for f, w in
                zip(lay["control"]["fields"], [8, 8, 9, 30, 19, 19, 16, 14, 18, 9, 18, 30])]
    wb.sheet("Control Example", ctl_cols, [list(r) for r in lay["control_example"]],
             description="Example control file of F01C (made-up values)")
    cols = [Column("layout", "Layout", 8, "Layout code"), Column("object", "Object", 8, "Data object"),
            Column("measure", "Measure", 14, "Control measure",
                   values=["ROW_COUNT", "HASH_TOTAL", "AMOUNT_TOTAL", "SHA256"]),
            Column("column", "Column", 26, "Column the measure applies to"),
            Column("currency", "Currency", 13, "Per currency or not"),
            Column("rule", "How it is computed", 44, "Computation rule"),
            Column("bdoi", "Value from BDOI (control file)", 18, "Filled from the control file"),
            Column("bibs", "Value in BIBS", 16, "Received / staged / loaded value (MIG-RECON-SUMMARY)"),
            Column("diff", "Difference", 12, "BDOI value minus BIBS value"),
            Column("status", "Status", 12, "MATCHED, BREAK or EXPLAINED", values=["MATCHED", "BREAK", "EXPLAINED"])]
    wb.sheet("Control Totals", cols, control_rows(lay),
             description="Control totals required per layout; also the tie-out template of the reconciliation (L1-L3)")
    cols = [Column("set", "Map set", 20, "Code map set"), Column("domain", "What is mapped", 46, "Legacy codes mapped"),
            Column("target", "BIBS master", 26, "Target master of the codes"),
            Column("owner", "Business owner (approves)", 26, "Approves the map versions (gate G2)"),
            Column("steward", "Data steward (prepares)", 24, "Prepares the entries"),
            Column("used", "Used by layouts", 26, "Layouts whose columns use the set"),
            Column("due", "Approved version due", 22, "Milestone")]
    wb.sheet("Code Map Sets", cols, [dict(zip(["set", "domain", "target", "owner", "steward", "used", "due"],
                                              [*r[:6], milestone_text(lay, r[6])])) for r in lay["map_sets"]],
             description="Code map sets (BRID 3.1): one per domain, versioned and approved")
    cols = [Column("set", "Map set", 20, "Code map set (sheet Code Map Sets)"),
            Column("source", "Source system", 12, "Legacy system of the code", values=["EBIX", "QPS", "ISYS", "CMS",
                                                                                        "EXCEL"]),
            Column("legacy_code", "Legacy code", 16, "Code as stored in legacy"),
            Column("legacy_desc", "Legacy description", 28, "Description in legacy"),
            Column("rows", "Rows using it", 11, "Filled by iorta TechNXT from profiling"),
            Column("action", "Action", 10, "MAP to a BIBS code; DEFAULT to the set default; REJECT the rows; CREATE "
                   "a new BIBS value", values=["MAP", "DEFAULT", "REJECT", "CREATE"]),
            Column("bibs_code", "BIBS code", 16, "Target code; mandatory for MAP and CREATE"),
            Column("bibs_desc", "BIBS description", 28, "Description of the BIBS value"),
            Column("remarks", "Remarks", 30, "Reason for DEFAULT, REJECT or CREATE"),
            Column("by", "Prepared by", 14, "Data steward"), Column("approved", "Approved by", 14, "Business owner")]
    example = [dict(zip(["set", "source", "legacy_code", "legacy_desc", "rows", "action", "bibs_code", "bibs_desc",
                         "remarks"], r)) for r in lay["map_example"]]
    wb.sheet("Code Map Template", cols, example,
             description="Template of a code map version (examples in the first rows; one file per set, "
                         "imported into the Migration Console as a DRAFT version)")
    cols = [Column("id", "Rule", 8, "Rule ID"), Column("layout", "Layout", 10, "Layout(s) checked"),
            Column("columns", "Columns", 20, "Columns checked"),
            Column("kind", "Kind", 15, "Kind of rule", values=["MANDATORY", "FORMAT", "LOOKUP", "UNIQUE", "REFERENTIAL",
                                                            "CROSS_FIELD", "BALANCE", "PLAUSIBILITY",
                                                            "DUPLICATE_CLIENT"]),
            Column("rule", "Rule", 50, "What is checked"), Column("severity", "Severity", 14, "ERROR = the row is not "
                                                                 "loaded; WARNING = loaded and reported"),
            Column("message", "Message shown in BIBS", 36, "Message text; codes are assigned at build"),
            Column("fixed_by", "Fixed by", 24, "Who corrects a failing row")]
    wb.sheet("DQ Rules", cols, [dict(zip(["id", "layout", "columns", "kind", "rule", "severity", "message",
                                          "fixed_by"], r)) for r in lay["dq_rules"]],
             description="Data-quality rules applied in staging (FR-DM-013); thresholds - master data 0.5 % "
                         "rejected or waived, financial objects 0")
    refs = lay.get("register_refs", {})
    cols = [Column("id", "ID", 9, "Question (DMQ##) or register ID (DCR-nnn)"),
            Column("register", "Register", 12, "Item of the BRD discrepancy and clarification register v1.2"),
            Column("topic", "Topic", 22, "Subject"),
            Column("question", "Decision needed", 60, "What BDOI decides"),
            Column("blocks", "Blocks", 28, "Objects, gates or build waves that wait for the answer"),
            Column("due", "Needed by", 20, "Milestone and date"), Column("owner", "BDOI owner", 22, "Who answers"),
            Column("status", "Status", 13, "OPEN, PARTIAL, ANSWERED, RECOMMENDED (answered with a recommendation that BDOI still confirms)",
                   values=["OPEN", "PARTIAL", "ANSWERED", "RECOMMENDED"],
                   status=True),
            Column("answer", "Answer", 40, "BDOI answer and date")]
    wb.sheet("Open Decisions", cols, [dict(zip(["id", "topic", "question", "blocks", "due", "owner", "status",
                                                "register", "answer"],
                                               [*r[:4], milestone_text(lay, r[4]), r[5], r[6], refs.get(r[0], "-"),
                                                r[7] if len(r) > 7 else ""]))
                                      for r in lay["decisions"]],
             description="Decisions that block the migration, with the date each is needed by")
    wb.sheet("Milestones", [Column("id", "Milestone", 10, "Milestone ID"), Column("what", "What is due", 60,
                                                                                 "Content of the milestone"),
                            Column("when", "When", 20, "Date or T-minus (T = go-live)")],
             [dict(zip(["id", "what", "when"], r)) for r in lay["milestones"]],
             description="Milestones used in the Due and Needed-by columns")
    return wb.save(OUT / brand.output_name("Migration", "BRD-13", "Data Requirements Workbook", m["version"], "xlsx"))


XLSX_TEMPLATES = {"P03": "RA_SENT"}  # layouts that business teams fill in Excel (DMQ38): layout -> first sheet name


def write_xlsx_template(layout: dict[str, Any], sheet: str, path: Path) -> Path:
    """Excel template of one layout: header in row 1 of the first sheet, lists on coded columns, text cells for
    dates, no merged cells or formulas (file contract), and a second sheet with the field instructions."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    fields = layout["fields"]
    ws.append([f["name"] for f in fields])
    for i, f in enumerate(fields, start=1):
        col = get_column_letter(i)
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.column_dimensions[col].width = max(14, len(f["name"]) + 3)
        if f["type"] in ("Date", "Timestamp", "Code", "Text"):
            for r in range(2, 2001):
                ws.cell(row=r, column=i).number_format = "@"
        if f["type"] == "Code" and re.fullmatch(r"[A-Z_]+(, [A-Z_]+)+", f["values"]):
            dv = DataValidation(type="list", formula1='"' + f["values"].replace(", ", ",") + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle=f["name"],
                                error=f"Allowed values: {f['values']}")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}2000")
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("Instructions")
    intro = [
        f"{layout['code']} - {layout['title']}",
        f"Fill the sheet {sheet}: one row per record under the header in row 1; do not rename, move or delete the "
        "header cells; no merged cells, formulas or colours with a meaning.",
        "Dates as yyyy-MM-dd text (for example 2027-11-16); amounts with a dot decimal and 2 decimals; codes exactly "
        "as in legacy.",
        f"Name the file {layout['code']}_EXCEL_<yyyyMMdd>_<nn>.xlsx (as-of date, sequence of the day) and send it with "
        "its control file (CONTROL_template.ctl.csv) through the Migration Console upload.",
        "Rejected rows come back in the rejection report (MIG-REJECTS) with correction columns: the maker corrects "
        "them and sends a resubmission file with the corrected rows only; the checker approves it in the console.",
        "",
    ]
    for line in intro:
        ins.append([line])
    header = ["Field", "Description", "Type", "Length", "Mandatory", "Allowed values", "Format", "Example", "Rule"]
    ins.append(header)
    for c in range(1, len(header) + 1):
        ins.cell(row=len(intro) + 1, column=c).font = Font(bold=True)
    for f in fields:
        ins.append([f["name"], f["description"], f["type"], f["length"], f["mandatory"], f["values"], f["format"],
                    f["example"], f["rule"]])
    for col, width in zip("ABCDEFGHI", [22, 48, 10, 8, 10, 28, 22, 22, 48]):
        ins.column_dimensions[col].width = width
    for row in ins.iter_rows(min_row=len(intro) + 2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ins.cell(row=1, column=1).font = Font(bold=True)
    wb.save(path)
    return path


def write_templates(lay: dict[str, Any]) -> list[Path]:
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    written = []
    for code, layout in lay["layouts"].items():
        path = TEMPLATES / f"{code}_template.csv"
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\r\n").writerow([f["name"] for f in layout["fields"]])
        path.write_bytes(buf.getvalue().encode("utf-8"))
        written.append(path)
    for code, sheet in XLSX_TEMPLATES.items():
        written.append(write_xlsx_template(lay["layouts"][code], sheet, TEMPLATES / f"{code}_template.xlsx"))
    path = TEMPLATES / "CONTROL_template.ctl.csv"
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerow([f["name"] for f in lay["control"]["fields"]])
    path.write_bytes(buf.getvalue().encode("utf-8"))
    written.append(path)
    objects = {o["code"]: o for o in lay["objects"]}
    lines = [
        "# BRD-13 Data Migration - extract templates",
        "",
        f"Version {lay['meta']['version']}, {lay['meta']['date']}. BDO Insurance and Reinsurance Brokers, Inc. "
        "(BDOI) - BIBS. Prepared by iorta TechNXT. Confidential - BDOI.",
        "",
        "One CSV template per extract layout. Each file holds only the header row: the field names exactly as in "
        "the layout sheet of the BDOI Data Requirements Workbook "
        f"(`../{brand.output_name('Migration', 'BRD-13', 'Data Requirements Workbook', lay['meta']['version'], 'xlsx')}`), "
        "which gives the type, length, mandatory flag, allowed values, format, example, BIBS target and "
        "validation rule of every field.",
        "",
        "How to use a template:",
        "",
        "1. Copy the template and name the copy `<LAYOUT>_<SOURCE>_<yyyyMMdd>_<nn>.csv`, for example "
        "`F01C_EBIX_20271231_01.csv` (as-of date, sequence of the day).",
        "2. Write one row per record under the header: UTF-8 without BOM, comma separator, RFC 4180 quoting, "
        "dates yyyy-MM-dd, amounts with a dot decimal and 2 decimals, codes exactly as stored in legacy.",
        "3. Produce the control file `<data file name>.ctl.csv` from `CONTROL_template.ctl.csv`: row count, hash "
        "total of the key, amount totals per column and currency, SHA-256 of the data file.",
        "4. Send both files through the Migration Console upload or the agreed SFTP drop. Never by e-mail.",
        "",
        "| Template | Layout | Object | Decision (proposed) | Source system | Fields | Key |",
        "|---|---|---|---|---|---|---|",
    ]
    for code, layout in lay["layouts"].items():
        o = objects[layout["object"]]
        lines.append(f"| `{code}_template.csv` | {layout['title']} | {o['code']} | {CLASS_LABEL[o['decision']]} | "
                     f"{o['source']} | {len(layout['fields'])} | {', '.join(layout['key'])} |")
    for code in XLSX_TEMPLATES:
        layout = lay["layouts"][code]
        o = objects[layout["object"]]
        lines.append(f"| `{code}_template.xlsx` | {layout['title']} (Excel template for the business team, "
                     f"with lists and an instructions sheet) | {o['code']} | {CLASS_LABEL[o['decision']]} | "
                     f"{o['source']} | {len(layout['fields'])} | {', '.join(layout['key'])} |")
    lines.append(f"| `CONTROL_template.ctl.csv` | Control file | all | - | - | {len(lay['control']['fields'])} | "
                 "data_file, measure, column_name, currency, filter |")
    lines += ["", "Objects without a template:", ""]
    for o in lay["objects"]:
        if not o["layouts"]:
            lines.append(f"- {o['code']} {o['name']}: {CLASS_LABEL[o['decision']]}. {o['rationale']}")
    lines.append("")
    readme = TEMPLATES / "README.md"
    readme.write_text("\n".join(lines), encoding="utf-8")
    written.append(readme)
    return written


# ------------------------------------------------------------------------------------------------ task plan

def build_task_plan(lay: dict[str, Any], cut: dict[str, Any]) -> Path:
    m = lay["meta"]
    wb = BdoiWorkbook("Cutover Task Plan", doc_type="Cutover runbook - task plan", brd="BRD-13",
                      version=m["version"], date=m["date"],
                      subtitle="BRD-13 Data Migration - production cutover from T-30 to hypercare exit")
    wb.legend = [("DONE", "Task completed and verified"), ("IN PROGRESS", "Task running"),
                 ("BLOCKED", "Task cannot start or finish"), ("N/A", "Not needed in this cutover")]
    wb.cover_notes = ["T = go-live: January 2028 (BDOI timeline), recommended Monday 3 January 2028 at the year-end "
                      "boundary (DMQ25; DMQ39 option A, awaiting Comptrollership confirmation). Days are calendar days; times are Philippine time; the sheet "
                      "Calendar gives the date of each relative day.",
                      "The same plan is loaded in the Migration Console (plan kind PRODUCTION); the console "
                      "records actual times and evidence (FR-DM-120)."]
    roles = {r[0]: r[1] for r in cut["roles"]}
    phases = {p[0]: p[1] for p in cut["phases"]}
    cols = [Column("id", "Task", 8, "Task ID"), Column("phase", "Phase", 14, "Phase of the cutover"),
            Column("day", "Day", 7, "Day relative to go-live (T)"), Column("time", "Start (PHT)", 8, "Planned start time"),
            Column("task", "Task", 58, "What is done"), Column("owner", "Owner", 8, "Accountable role (sheet Roles)"),
            Column("hours", "Hours", 6, "Planned elapsed hours", kind="number"),
            Column("pred", "Predecessors", 16, "Tasks that must be complete first"),
            Column("verification", "Verification / evidence", 40, "How completion is proven"),
            Column("checkpoint", "Checkpoint", 13, "Go / no-go checkpoint or milestone"),
            Column("assignee", "Assignee (name)", 14, "Named person"),
            Column("actual_start", "Actual start", 12, "Date and time"), Column("actual_end", "Actual end", 12,
                                                                                "Date and time"),
            Column("status", "Status", 12, "Task status", values=["NOT STARTED", "IN PROGRESS", "DONE", "BLOCKED",
                                                                   "N/A"], status=True)]
    rows = []
    for t in cut["tasks"]:
        rows.append({"id": t[0], "phase": f"{t[1]} {phases[t[1]]}", "day": t[2], "time": t[3], "task": t[4],
                     "owner": t[5], "hours": t[6], "pred": t[7], "verification": t[8], "checkpoint": t[9],
                     "status": "NOT STARTED"})
    wb.sheet("Tasks", cols, rows, description="Cutover tasks with owner, duration, predecessors and verification")
    wb.sheet("Roles", [Column("code", "Code", 8, "Role code used in the Owner column"),
                       Column("role", "Role", 70, "Role"), Column("org", "Organisation", 24, "Organisation"),
                       Column("name", "Name (to fill)", 26, "Named person and phone")],
             [dict(zip(["code", "role", "org"], r)) for r in cut["roles"]], description="Owner roles")
    wb.sheet("Phases", [Column("code", "Phase", 8, "Phase code"), Column("name", "Name", 40, "Phase"),
                        Column("when", "When", 26, "Days")],
             [dict(zip(["code", "name", "when"], p)) for p in cut["phases"]], description="Cutover phases")
    wb.sheet("Calendar", [Column("day", "Day", 10, "Day relative to go-live (T)"),
                          Column("date", "Date (T = 3-Jan-2028)", 20, "Calendar date for the recommended go-live"),
                          Column("note", "Note", 70, "Holiday or year-end note")],
             [dict(zip(["day", "date", "note"], c)) for c in cut["calendar"]],
             description="Relative days mapped to the calendar of the recommended go-live (DMQ25, DMQ39)")
    rows = []
    for c in cut["checkpoints"]:
        for i, crit in enumerate(c["criteria"], start=1):
            rows.append({"id": c["id"], "when": c["when"], "decision": c["decision"], "n": i, "criterion": crit})
    wb.sheet("Checkpoints", [Column("id", "Checkpoint", 14, "Checkpoint"), Column("when", "When", 10, "Day and time"),
                             Column("decision", "Decision", 30, "What is decided"), Column("n", "#", 4, "Criterion no."),
                             Column("criterion", "Criterion", 70, "Criterion that must be met"),
                             Column("value", "Measured value", 18, "Value at the decision"),
                             Column("met", "Met", 8, "Y / N", values=["Y", "N"])],
             rows, description="Go / no-go checkpoints and their criteria")
    wb.sheet("Freeze Windows", [Column("what", "Freeze", 20, "What is frozen"), Column("from", "From", 22, "Start"),
                                Column("to", "Until", 12, "End"), Column("scope", "Scope", 50, "What may not change"),
                                Column("exception", "Exception route", 46, "How an urgent change is handled")],
             [dict(zip(["what", "from", "to", "scope", "exception"], f)) for f in cut["freeze_windows"]],
             description="Freeze windows")
    wb.sheet("Rollback", [Column("id", "Step", 7, "Step"), Column("name", "Name", 16, "Step name"),
                          Column("action", "Action", 90, "What is done"),
                          Column("owner", "Owner", 8, "Role"),
                          Column("done", "Done at", 14, "Date and time")],
             [dict(zip(["id", "name", "action", "owner"], r)) for r in cut["rollback"]],
             description="Rollback procedure (only on NO-GO or before the point of no return)")
    wb.sheet("Communication", [Column("n", "#", 4, "Item"), Column("message", "Message", 50, "What is said"),
                               Column("audience", "Audience", 34, "Who receives it"),
                               Column("channel", "Channel", 26, "How"), Column("when", "When", 16, "Day"),
                               Column("owner", "Owner", 8, "Role"), Column("sent", "Sent on", 14, "Date")],
             [dict(zip(["n", "message", "audience", "channel", "when", "owner"], c)) for c in cut["communication"]],
             description="Communication plan")
    wb.sheet("Hypercare Roster", [Column("function", "Function", 26, "Support function"),
                                  Column("who", "Who", 44, "Team"), Column("hours", "Coverage", 36, "Hours"),
                                  Column("escalation", "Escalates to", 24, "Escalation"),
                                  Column("name", "Names and phone (to fill)", 30, "Named people")],
             [dict(zip(["function", "who", "hours", "escalation"], r)) for r in cut["hypercare_roster"]],
             description="Hypercare roster")
    wb.sheet("Hypercare Exit", [Column("n", "#", 4, "Criterion"), Column("criterion", "Exit criterion", 90, "Criterion"),
                                Column("met", "Met", 8, "Y / N", values=["Y", "N"]),
                                Column("evidence", "Evidence", 30, "Report or sign-off")],
             [{"n": i, "criterion": c} for i, c in enumerate(cut["hypercare_exit"], start=1)],
             description="Hypercare exit criteria")
    rows = []
    for system, items in cut["decommissioning"].items():
        for i, crit in enumerate(items, start=1):
            rows.append({"system": system, "n": i, "criterion": crit, "status": "OPEN"})
    wb.sheet("Decommissioning", [Column("system", "Legacy system", 20, "System or context"),
                                 Column("n", "#", 4, "Criterion"), Column("criterion", "Criterion", 80, "Criterion"),
                                 Column("evidence", "Evidence", 26, "Report, log or sign-off"),
                                 Column("signed", "Signed by / date", 20, "Signer"),
                                 Column("status", "Status", 11, "Status", values=["OPEN", "DONE", "N/A"], status=True)],
             rows, description="Decommissioning checklist per legacy system (BRID 12.1; DMQ27)")
    return wb.save(OUT / brand.output_name("Migration", "BRD-13", "Cutover Task Plan", m["version"], "xlsx"))


# ------------------------------------------------------------------------------------------------ Word placeholders

def md_table(headers: list[str], rows: list[list[Any]], opts: str) -> list[str]:
    def cell(v: Any) -> str:
        return str(v if v is not None else "").replace("|", "/").replace("\n", " ").strip() or "-"
    out = [f"<!-- table: {opts} -->", "| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(cell(v) for v in r) + " |" for r in rows]
    return out + [""]


def placeholders(lay: dict[str, Any], cut: dict[str, Any], name: str, opts: dict[str, str]) -> list[str]:
    objects = lay["objects"]
    roles = {r[0]: r[1] for r in cut["roles"]}
    if name == "register":
        rows = [[o["code"], o["name"], CLASS_LABEL[o["decision"]], o["rationale"]] for o in objects]
        return md_table(["Object", "Data object", "Decision", "Rationale"], rows,
                        'widths=1.3,5,2.2,8.1 caption="Data objects and proposed decisions (BDOI confirms at gate G1)"'
                        ' size=8 bold=first')
    if name == "register-detail":
        rows = [[o["code"], o["history"], o["delta"], o["format"], milestone_text(lay, o["due"])] for o in objects
                if o["decision"] != "EXCLUDED"]
        return md_table(["Object", "History depth", "Delta after the first full extract", "Format", "Due"], rows,
                        'widths=1.3,5.2,4.6,2.4,3.1 caption="Depth, delta frequency, format and due dates per object"'
                        ' size=8 bold=first')
    if name == "source-matrix":
        rows = []
        for o in objects:
            src = o["source"]
            marks = []
            for s in SYSTEMS:
                key = {"Excel": "Excel", "File shares": "File shares"}.get(s, s)
                marks.append("Y" if re.search(rf"\b{re.escape(key)}\b", src, re.I) else "")
            rows.append([o["code"], o["name"], *marks])
        return md_table(["Object", "Data object", *SYSTEMS], rows,
                        'widths=1.3,7,1.3,1.3,1.3,1.3,1.3,1.8 caption="Source systems per data object (Y = holds the '
                        'object; system of record per DMQ03)" size=8 bold=first')
    if name == "targets":
        rows = [[o["code"], o["target"], o["brd"], str(o["order"] or "-")] for o in objects
                if o["decision"] != "EXCLUDED"]
        return md_table(["Object", "BIBS target (service or table)", "BRD", "Load order"], rows,
                        'widths=1.3,10,3.6,1.7 caption="Target in BIBS and load order" size=8 bold=first')
    if name == "owners":
        rows = [[o["code"], o["owner"], o["steward"]] for o in objects if o["decision"] != "EXCLUDED"]
        return md_table(["Object", "Business owner (signs G1, G2, G6)", "Data steward"], rows,
                        'widths=1.3,8.3,7 caption="Business owner and data steward per object (proposal; DMQ01)" size=8'
                        ' bold=first')
    if name == "layouts":
        rows = []
        for code, layout in lay["layouts"].items():
            f = layout["fields"]
            rows.append([code, layout["title"], layout["object"], len(f), sum(1 for x in f if x["mandatory"] == "Y"),
                         ", ".join(layout["key"]), ", ".join(layout["amounts"]) or "-"])
        return md_table(["Layout", "Content", "Object", "Fields", "Mandatory", "Key", "Amount totals"], rows,
                        'widths=1.3,4.6,1.3,1.2,1.6,3.4,3.2 caption="Extract layouts (workbook sheets and CSV '
                        'templates)" size=8 bold=first')
    if name == "bdoi-inputs":
        rows = []
        for o in objects:
            if not o["layouts"]:
                continue
            rows.append([o["code"], o["name"], ", ".join(o["layouts"]),
                         o["source"], o["controls"], milestone_text(lay, o["due"])])
        return md_table(["Object", "Data object", "Templates", "Source", "Control totals", "Due"], rows,
                        'widths=1.2,5,1.8,2,4,2.6 caption="What BDOI provides per object" size=8 bold=first')
    if name == "decisions":
        refs = lay.get("register_refs", {})
        rows = [[d[0], refs.get(d[0], "-"), d[1], d[3], milestone_text(lay, d[4]), d[5], d[6]] for d in lay["decisions"]]
        return md_table(["ID", "Register", "Decision", "Blocks", "Needed by", "BDOI owner", "Status"], rows,
                        'widths=1.5,1.6,3,3.4,2.7,2.6,2.8 caption="Decisions, register items (v1.2) and the date '
                        'each is needed by" size=8 bold=first status=Status')
    if name == "milestones":
        return md_table(["Milestone", "What is due", "When"], [list(m) for m in lay["milestones"]],
                        'widths=1.8,10.8,4 caption="Milestones of BDOI inputs" bold=first')
    if name == "map-sets":
        rows = [[r[0], r[1], r[2], r[3]] for r in lay["map_sets"]]
        return md_table(["Map set", "What is mapped", "BIBS master", "Approves"], rows,
                        'widths=2.8,7,3.4,3.4 caption="Code map sets (BRID 3.1)" size=8 bold=first')
    if name == "dq-summary":
        kinds: dict[str, int] = {}
        for r in lay["dq_rules"]:
            kinds[r[3]] = kinds.get(r[3], 0) + 1
        rows = [[k, v, ", ".join(r[0] for r in lay["dq_rules"] if r[3] == k)] for k, v in kinds.items()]
        return md_table(["Kind", "Rules", "Rule IDs"], rows,
                        'widths=3.4,1.4,11.8 caption="Data-quality rules by kind (full catalogue in the workbook, sheet '
                        'DQ Rules)" size=8.5 bold=first')
    if name == "controls":
        rows = []
        for code, layout in lay["layouts"].items():
            amounts = ", ".join(layout["amounts"]) or "-"
            rows.append([code, layout["title"], layout["hash"], amounts])
        return md_table(["Layout", "Content", "Hash total (L3)", "Amount totals per currency (L2)"], rows,
                        'widths=1.3,5,5,5.3 caption="Control totals required per layout (row count and SHA-256 for '
                        'every layout)" size=8 bold=first')
    if name == "roles":
        return md_table(["Code", "Role", "Organisation"], [list(r) for r in cut["roles"]],
                        'widths=1.6,11,4 caption="Roles used in the task plan" size=8.5 bold=first')
    if name == "phases":
        rows = []
        for p in cut["phases"]:
            tasks = [t for t in cut["tasks"] if t[1] == p[0]]
            rows.append([p[0], p[1], p[2], len(tasks), f"{tasks[0][0]} to {tasks[-1][0]}" if tasks else "-"])
        return md_table(["Phase", "Name", "When", "Tasks", "IDs"], rows,
                        'widths=1.2,6,4,1.4,4 caption="Cutover phases" bold=first')
    if name == "tasks":
        phase = opts.get("phase")
        rows = [[t[0], t[2] + (f" {t[3]}" if t[3] else ""), t[4], t[5], t[6], t[7] or "-", t[8]]
                for t in cut["tasks"] if t[1] == phase]
        title = next(p[1] for p in cut["phases"] if p[0] == phase)
        return md_table(["ID", "When", "Task", "Owner", "Hours", "After", "Verification"], rows,
                        f'widths=1.4,1.7,8.6,1.2,1.1,2.4,6.4 caption="Phase {phase} - {title}" size=8 bold=first')
    if name == "freeze":
        return md_table(["Freeze", "From", "Until", "Scope", "Exception route"], [list(f) for f in cut["freeze_windows"]],
                        'widths=2.6,2.8,1.8,5.4,4 caption="Freeze windows" size=8.5 bold=first')
    if name == "checkpoints":
        out: list[str] = []
        for c in cut["checkpoints"]:
            rows = [[i, crit] for i, crit in enumerate(c["criteria"], start=1)]
            out += md_table(["#", "Criterion"], rows,
                            f'widths=0.8,15.8 caption="{c["id"]} at {c["when"]} - {c["decision"]}" size=8.5')
        return out
    if name == "rollback":
        return md_table(["Step", "Name", "Action", "Owner"], [list(r) for r in cut["rollback"]],
                        'widths=1.2,2.2,11.6,1.6 caption="Rollback procedure" size=8.5 bold=first')
    if name == "communication":
        return md_table(["#", "Message", "Audience", "Channel", "When", "Owner"], [list(c) for c in cut["communication"]],
                        'widths=0.6,5.4,3.8,3,2.2,1.2 caption="Communication plan" size=8 bold=first')
    if name == "roster":
        return md_table(["Function", "Who", "Coverage", "Escalates to"], [list(r) for r in cut["hypercare_roster"]],
                        'widths=3.2,5.4,4.8,3.2 caption="Hypercare roster (names in the task plan, sheet Hypercare '
                        'Roster)" size=8.5 bold=first')
    if name == "hypercare-exit":
        return md_table(["#", "Exit criterion"], [[i, c] for i, c in enumerate(cut["hypercare_exit"], start=1)],
                        'widths=0.8,15.8 caption="Hypercare exit criteria" size=9')
    if name == "decommissioning":
        out = []
        for system, items in cut["decommissioning"].items():
            out += md_table(["#", "Criterion"], [[i, c] for i, c in enumerate(items, start=1)],
                            f'widths=0.8,15.8 caption="Decommissioning checklist - {system}" size=8.5')
        return out
    if name == "calendar":
        return md_table(["Day", "Date (T = 3-Jan-2028)", "Note"], [list(c) for c in cut["calendar"]],
                        'widths=1.6,4,11 caption="Relative days on the calendar of the recommended go-live" size=8.5 '
                        'bold=first')
    if name == "role-legend":
        return md_table(["Code", "Role"], [[k, v] for k, v in roles.items()],
                        'widths=1.6,15 caption="Owner codes" size=8.5 bold=first')
    raise SystemExit(f"unknown placeholder dm:{name}")


PLACEHOLDER = re.compile(r"^\s*<!--\s*dm:([\w-]+)\s*(.*?)\s*-->\s*$")


def expand(lay: dict[str, Any], cut: dict[str, Any], lines: list[str]) -> list[str]:
    out: list[str] = []
    for line in lines:
        m = PLACEHOLDER.match(line)
        if not m:
            out.append(line)
            continue
        opts = dict(re.findall(r"(\w+)=(\S+)", m.group(2)))
        out += placeholders(lay, cut, m.group(1), opts)
    return out


def expanded_source(lay: dict[str, Any], cut: dict[str, Any], src: Path) -> Path:
    """Writes the expanded source next to the original (for lint_source) and returns its path."""
    front_text, body = src.read_text(encoding="utf-8").split("---", 2)[1:]
    lines = expand(lay, cut, body.splitlines())
    tmp = src.with_name(f".{src.stem}.expanded.md")
    tmp.write_text("---" + front_text + "---\n" + "\n".join(lines) + "\n", encoding="utf-8")
    return tmp


def build_doc(lay: dict[str, Any], cut: dict[str, Any], src: Path, pdf: bool, keep_pdf: bool):
    front, lines = load_source(src)
    doc = BdoiDocument(meta_from(front), h1_page_break=bool(front.get("h1_page_break", True)), base_dir=HERE)
    doc.cover()
    doc.front_matter()
    render_body(doc, expand(lay, cut, lines))
    return doc.publish(output_path(front, src), pdf=pdf, keep_pdf=keep_pdf)


# ------------------------------------------------------------------------------------------------ main

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="checks only")
    ap.add_argument("--previews", action="store_true", help="render page previews of every output")
    ap.add_argument("--no-pdf", action="store_true", help="skip the PDF pass of the Word documents")
    ap.add_argument("--only", help="build only this Word source (file name)")
    args = ap.parse_args(argv)
    lay, cut = load()
    problems = check(lay, cut)
    for name in DOCS:
        tmp = expanded_source(lay, cut, HERE / name)
        try:
            problems += [f"{name}: {p}" for p in lint_source(tmp)]
            text = tmp.read_text(encoding="utf-8")
            for word in FILLER:
                if re.search(rf"\b{word}\b", text, re.I):
                    problems.append(f"{name}: filler word {word}")
        finally:
            tmp.unlink(missing_ok=True)
    for p in problems:
        print(p)
    if problems:
        return 1
    print(f"check: {len(lay['objects'])} objects, {len(lay['layouts'])} layouts, "
          f"{sum(len(x['fields']) for x in lay['layouts'].values())} fields, {len(lay['dq_rules'])} rules, "
          f"{len(lay['decisions'])} decisions, {len(cut['tasks'])} cutover tasks - OK")
    if args.check:
        return 0
    import render

    outputs: list[Path] = []
    if not args.only:
        outputs.append(build_workbook(lay))
        templates = write_templates(lay)
        print(f"templates: {len(templates)} files in {TEMPLATES}")
        outputs.append(build_task_plan(lay, cut))
        for p in outputs:
            print(f"xlsx: {p}")
        if args.previews:
            for p in outputs:
                pdf = render.to_pdf(p)
                render.previews(pdf)
                pdf.unlink(missing_ok=True)
    for name in DOCS:
        if args.only and args.only != name:
            continue
        docx_path, pdf_path = build_doc(lay, cut, HERE / name, pdf=not args.no_pdf, keep_pdf=args.previews)
        print(f"docx: {docx_path}")
        if pdf_path and args.previews:
            print(f"pages: {render.page_count(pdf_path)}")
            render.previews(pdf_path)
            pdf_path.unlink(missing_ok=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
