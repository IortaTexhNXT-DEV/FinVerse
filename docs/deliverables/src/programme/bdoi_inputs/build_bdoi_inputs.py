"""The BDOI input workbooks of the programme (request of the BDOI product owner, 8 October 2026).

    python docs/deliverables/src/programme/bdoi_inputs/build_bdoi_inputs.py            # A, B and the Word guide
    python docs/deliverables/src/programme/bdoi_inputs/build_bdoi_inputs.py --check    # checks only
    python docs/deliverables/src/programme/bdoi_inputs/build_bdoi_inputs.py --a        # workbook A only

(A) BIBS_Inputs_BRD-00_BDOI_Requirements_and_Inputs_by_BRD_v1.0.xlsx: what the project needs from the BDOI business
    users, one tab per BRD (BRD-00 cross-cutting, BRD-01 to BRD-13, ReInsurance phase 2), with the Summary and the
    Instructions. The rows are consolidated and de-duplicated from
      * the proposed business rules and clarifications (CLR-*) and the open questions of every FRS (src/BRD-*/FRS_*.md),
      * the proposed rules, clarifications and open decisions of the BRD-13 set (pack/catalogue.yaml),
      * the discrepancy and clarification register (src/registers: items DCR-*, open questions and their status),
      * the earlier Drop 0 requirements list (drop0_requirements.yaml, items PM-Q*, UA-Q*, DM-Q*),
      * the Drop 0 configuration register (src/Drop-0_Closure/drop0.yaml, items CFG-*) and the platform document
        templates,
      * the programme: NFR targets to decide or confirm (requirements_quality/nfr_catalogue.yaml), the BDOI
        dependencies, sign-offs and test windows of the project plan (delivery_adoption/plan.yaml).
    A question, register item or earlier list item that a clarification already covers is merged into that row; the
    original IDs stay in the Source column so that the answers trace back.
(B) BIBS_Templates_BRD-00_Master_Data_and_Configuration_Upload_Templates_v1.0.xlsx with the short Word guide
    BIBS_Templates_BRD-00_Master_Data_Upload_Guide_v1.0.docx: every master, reference and configuration dataset BIBS
    needs before go-live, one guided template sheet per dataset in the column layout the platform loads (the Drop 0,
    BRD-03 and BRD-11 configuration templates, the Migration Console layouts of BRD-13 and the upload definitions of
    the platform), with the index in load order, the datasets configured on screens and the gaps.
Outputs: docs/deliverables/out/Programme/BDOI_Inputs/.
"""

from __future__ import annotations

import argparse
import datetime as dt
import importlib.util
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

import yaml

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402

SOURCE = HERE / "bdoi_inputs.yaml"
DROP0_LIST = HERE / "drop0_requirements.yaml"
GUIDE = HERE / "MASTER_DATA_GUIDE.md"
SRC = brand.SRC_DIR
KIND = "BDOI_Inputs"
STATUSES = ["Open", "Partly answered", "Answered"]
PRIORITIES = ["High", "Medium", "Low"]


# ============================================================================ loading helpers


def _module(path: Path, name: str):
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


@lru_cache(maxsize=1)
def cfg() -> dict[str, Any]:
    data = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    for key in ("screen_templates", "depends", "due"):
        data[key].update(data.get(f"{key}_more") or {})
    data["extra_rows"] = list(data.get("extra_rows") or []) + list(data.get("extra_rows_more") or [])
    return data


def excluded() -> dict[str, set[str]]:
    """Insurer-company items left out (parameters, accounting events and categories, configuration items)."""
    x = cfg().get("insurer_exclusions") or {}
    return {k: set(v or []) for k, v in x.items()}


def register_module():
    return _module(SRC / "registers" / "build_discrepancy_register.py", "build_discrepancy_register")


def dm_module():
    return _module(SRC / "BRD-13_Data_Migration" / "build_dm_pack.py", "build_dm_pack")


def drop_closure():
    import drop_closure as dc  # noqa: PLC0415

    return dc


def as_date(v: Any) -> dt.date | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def fmt(d: dt.date | None) -> str:
    return d.strftime("%d-%b-%Y") if d else ""


@lru_cache(maxsize=1)
def tabs() -> list[dict[str, Any]]:
    return [{**t, "due": as_date(t["due"])} for t in cfg()["tabs"]]


def tab(brd: str) -> dict[str, Any]:
    return next(t for t in tabs() if t["brd"] == brd)


# ============================================================================ client wording

# Development-status words and internal references are rewritten for the client (check_pack BUILD_STATUS).
WORDING = [
    (r"\b[Aa]s built\b", "as delivered"),
    (r"\b[Aa]s designed\b", "as specified"),
    (r"\bnot built\b", "not provided"),
    (r"\bwere built on\b", "were written on"),
    (r"\bwas built on\b", "was written on"),
    (r"\bbuilt\b(?!-in)", "provided"),
    (r"\bdesigned\b", "specified"),
    (r"\b[Rr]ebuild(s)?\b(?! from)", r"redo\1"),
    (r"\bbuilds\b", "prepares"),
    (r"\bbuild\b(?!-up)(?! Comparative)", "delivery"),
    (r"\bwork[- ]in[- ]progress\b", "in process"),
    (r"\b[Dd]efects?\b", "incidents"),
    (r"\bknown issues?\b", "open points"),
    (r"\bbugs?\b", "incidents"),
    (r"\bin development\b", "in preparation"),
    (r"\bunder development\b", "in preparation"),
    (r"\bto be built\b", "to be provided"),
    (r"\bnot (?:yet )?implemented\b", "not yet available"),
    (r"\bsprints?\b", "iterations"),
    (r"\bstub(?:s|bed)?\b", "placeholder"),
    (r"\bmocked\b", "simulated"),
    (r"\bseams?\b", "connection points"),
    (r"\badapters?\b", "connectors"),
    (r"\bidempotent\b", "repeatable without duplicates"),
    (r"\bsource code\b", "platform"),
    (r"\bthe code (reads|returns|checks)\b", r"BIBS \1"),
    (r"\b(?:in|from) the code\b(?! maps?\b| lists?\b| of\b)", "in BIBS"),
    # technical terms (client decision of 28-Sep-2026: business content only)
    (r"\bREST APIs?\b", "interfaces"),
    (r"\bAPIs\b", "interfaces"),
    (r"\bAPI\b", "interface"),
    (r"\bend-?points?\b", "addresses"),
    (r"\bJSONB?\b", "structured data"),
    (r"\bNoSQL\b", "document store"),
    (r"\bSQL\b", "query"),
    (r"\bdatabases\b", "data stores"),
    (r"\bdatabase\b", "data store"),
    (r"\bDBAs?\b", "data store administrators"),
    (r"\bschemas\b", "structures"),
    (r"\bschema\b", "structure"),
    (r"\bdata model\b", "data structure"),
    (r"\btable names?\b", "field names"),
    (r"\bpayloads?\b", "message contents"),
    (r"\bflyway\b", "release scripts"),
    (r"\bmigration scripts?\b", "release scripts"),
    (r"\bHTTPS?\b", "secure web"),
    (r"\bS?FTPS?\b", "secure file transfer"),
]
# Codes of the platform written in a business column: a code in brackets alone is dropped, any other is written
# in words.
CODE = re.compile(r"\b[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+\b")
LOWER_CODE = re.compile(r"\b[a-z]{2,}(?:_[a-z0-9]+)+\b")
CAMEL = re.compile(r"\b[A-Z][a-z]+(?:[A-Z][a-z0-9]+){2,}\b|\b[A-Z][a-z]+(?:[A-Z][a-z0-9]+)+(?:Service|Controller|"
                   r"Repository|Entity|Dto|Mapper|Job|Listener|Handler|Config|Page|Inbox|Client|Adapter|Port|"
                   r"Gateway)\b")
PATHS = re.compile(r"\b[\w.-]+\.(?:java|kt|tsx?|jsx?|py|sql|md|ya?ml|properties|xml|json|dot|sh)\b|"
                   r"(?:^|(?<=[\s(`'\"]))(?:\.{0,2}/)?(?:src|docs|tools|backend|frontend|deploy|main|test)/[\w./-]+|"
                   r"(?:^|(?<=[\s(]))(?:\.\.\.)?/[a-z][\w-]*(?:/[\w{}.:-]+)+|/api/\S*")


KEEP = ["Negative List Database System"]  # names of BDO systems (check_pack TECHNICAL_ALLOWED)


def wording(s: str) -> str:
    """The rewrites of WORDING, case-insensitive, keeping a leading capital."""
    keep = {f"\x00{i}\x00": k for i, k in enumerate(KEEP)}
    for token, phrase in keep.items():
        s = s.replace(phrase, token)
    for pat, repl in WORDING:
        def sub(m: re.Match, repl: str = repl) -> str:
            out = m.expand(repl)
            return out[:1].upper() + out[1:] if m.group(0)[:1].isupper() and out else out
        s = re.sub(pat, sub, s, flags=re.I)
    for token, phrase in keep.items():
        s = s.replace(token, phrase)
    return s


def clean(text: Any) -> str:
    """Plain text of a Markdown cell or a YAML value."""
    s = str(text or "")
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = s.replace("**", "").replace("`", "").replace("<br>", "; ").replace("\\|", "/")
    return re.sub(r"\s+", " ", s).strip()


def code_words(m: re.Match) -> str:
    return m.group(0).replace("_", " ").lower()


def business(text: Any) -> str:
    """Client wording of a text of the sources: no platform codes, paths or development-status words."""
    s = clean(text)
    if not s:
        return ""
    s = s.replace("BrokerVerse", "BIBS")
    s = PATHS.sub("", s)
    # (CODE) or (CODE, CODE) alone in brackets: dropped
    s = re.sub(r"\s*\((?:\s*(?:LOV |list |parameter |permission |role |status |event )?[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+"
               r"\s*[,;/]?)+\)", "", s)
    s = CODE.sub(lambda m: m.group(0) if m.group(0).startswith("BRID_") else code_words(m), s)
    s = LOWER_CODE.sub(code_words, s)
    s = CAMEL.sub(lambda m: re.sub(r"(?<!^)(?=[A-Z])", " ", m.group(0)).lower(), s)
    s = wording(s)
    s = re.sub(r"\(\s*\)", "", s)
    s = re.sub(r"\s+([,.;:])", r"\1", re.sub(r"\s{2,}", " ", s)).strip()
    return s


# ============================================================================ IDs

ID_RE = re.compile(r"\b(?:CLR-[A-Z]+-\d+|PR-DM-\d+|DCR-\d+|NFR-[A-Z]{3}-\d+|DEP-\d+|CFG-\d+|(?:PM|UA|DM)-Q\d+|"
                   r"(?:SANC|SP)-SQ\d+|(?:CRQ|IQ|XQ|DSQ|OQ|PQ|CQ|AQ|RQ|CLQ|EBQ|CSQ|UQ|DMQ|Q)\d+)\b")


def normalise_ids(text: str, brd: str = "") -> str:
    """SP SQ06 -> SP-SQ06; SQ06 of the BRD-10 / BRD-12 FRS -> SANC-SQ06 / SP-SQ06."""
    text = re.sub(r"\bSP[ ‑-]SQ(\d+)", r"SP-SQ\1", text)
    text = re.sub(r"\bSANC[ ‑-]SQ(\d+)", r"SANC-SQ\1", text)
    prefix = {"BRD-10": "SANC-", "BRD-12": "SP-"}.get(brd, "")
    if prefix:
        text = re.sub(r"(?<![-\w])SQ(\d+)", prefix + r"SQ\1", text)
    return text


def ids_in(text: str) -> list[str]:
    out: list[str] = []
    for m in ID_RE.finditer(text):
        if m.group(0) not in out:
            out.append(m.group(0))
    return out


def question_prefix(qid: str) -> str:
    m = re.match(r"((?:SANC|SP)-SQ|[A-Z]+)", qid)
    return m.group(1) if m else ""


def home_of(qid: str, default: str) -> str:
    return cfg()["question_home"].get(question_prefix(qid), default)


# ============================================================================ the rows


@dataclass
class Item:
    id: str
    kind: str  # clr, proposal, decision, drop0, question, dcr, nfr, cfg, doc, dep, test, signoff, ri
    tab: str
    need: str
    why: str = ""
    proposal: str = ""
    source: str = ""
    owner: str = ""
    due: dt.date | None = None
    priority: str = "Medium"
    status: str = ""  # Open, Partly answered, Answered ("" = from the attached items)
    answer: str = ""
    type: str = ""
    topic: str = ""
    text: str = ""  # every text of the item, for the references
    module: str = ""
    remarks: list[str] = field(default_factory=list)
    attached: list["Item"] = field(default_factory=list)
    order: int = 0
    brds: list[str] = field(default_factory=list)

    @property
    def refs(self) -> list[str]:
        return [i for i in ids_in(self.text) if i != self.id]


def md_section_rows(path: Path, heading: str, level: str = "#") -> list[list[str]]:
    """The rows of the pipe tables under a heading (to the next heading of the same or a higher level)."""
    lines = path.read_text(encoding="utf-8").splitlines()
    start = next((i for i, line in enumerate(lines) if re.match(rf"^{level} {re.escape(heading)}", line)), None)
    if start is None:
        return []
    depth = len(level)
    rows = []
    for line in lines[start + 1:]:
        m = re.match(r"^(#+) ", line)
        if m and len(m.group(1)) <= depth:
            break
        if line.startswith("|") and not re.match(r"\|\s*-", line):
            cells = [c.strip() for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]]
            rows.append(cells)
    return rows


def front_matter(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---"):
        return {}
    return yaml.safe_load(text.split("---", 2)[1]) or {}


def frs_files() -> dict[str, Path]:
    out = {}
    for t in tabs():
        if t["brd"] in ("RI", "BRD-13"):
            continue
        found = sorted(brand.src_dir(t["brd"]).glob("FRS_*.md"))
        if found:
            out[t["brd"]] = found[0]
    return out


def topic_refs(topic: str) -> tuple[str, str]:
    """('KYC at renewal', 'BRRN.028 and its annex; FR-RN-026') from 'KYC at renewal (BRRN.028 and its annex; FR-RN-026)'."""
    m = re.match(r"^(.*?)\s*\(([^()]*(?:\([^()]*\)[^()]*)*)\)\s*$", topic)
    return (m.group(1).strip(), m.group(2).strip()) if m else (topic.strip(), "")


def frs_items() -> tuple[list[Item], dict[str, dict[str, Any]]]:
    """The CLR rows of every FRS, and the open questions of the FRS keyed by question ID."""
    items: list[Item] = []
    questions: dict[str, dict[str, Any]] = {}
    for brd, path in frs_files().items():
        fm = front_matter(path)
        version = fm.get("version", "")
        doc = f"FRS {brd} v{version}"
        for cells in md_section_rows(path, "Proposed business rules and clarifications for confirmation"):
            if len(cells) < 5 or not cells[0].startswith("CLR-"):
                continue
            ref, topic, rule, reason, decision = [normalise_ids(c, brd) for c in cells[:5]]
            title, trefs = topic_refs(clean(topic))
            it = Item(id=clean(ref), kind="clr", tab=brd, need=f"{business(title)}: {business(decision)}",
                      why=business(reason), proposal=business(rule),
                      source=f"{doc}, Proposed business rules and clarifications" + (f"; {clean(trefs)}" if trefs else ""),
                      topic=business(title), text=" ".join(clean(c) for c in (topic, rule, reason, decision)))
            items.append(it)
        for cells in md_section_rows(path, "Open questions", "##"):
            if len(cells) < 4 or cells[0] in ("ID", "Ref"):
                continue
            raw = normalise_ids(clean(cells[0]), brd)
            qids = ids_in(raw) or [raw]
            status_raw = clean(cells[-1])
            for k, qid in enumerate(qids):
                q = {"id": qid, "brd": brd, "question": business(normalise_ids(cells[1], brd)),
                     "affects": clean(cells[2]), "status_raw": status_raw, "doc": doc,
                     "text": normalise_ids(" ".join(clean(c) for c in cells), brd), "same_row": qids if k == 0 else []}
                prev = questions.get(qid)
                if prev is None or (prev["brd"] != home_of(qid, prev["brd"]) and brd == home_of(qid, brd)):
                    questions[qid] = q
    return items, questions


def frs_status(raw: str) -> tuple[str, str]:
    """(status, answer) of a status cell of an FRS open-questions table: OPEN, PARTIAL (...), ANSWERED (...)."""
    m = re.match(r"^(OPEN|PARTIAL|ANSWERED|CLOSED|RECOMMENDED)\b\s*(?:\((.*)\))?", raw.strip(), re.I)
    if not m:
        return "Open", ""
    word, note = m.group(1).upper(), (m.group(2) or "").strip()
    return {"OPEN": "Open", "PARTIAL": "Partly answered", "RECOMMENDED": "Partly answered"}.get(word, "Answered"), note


ORDER = ["Open", "Partly answered", "Answered"]


def better(a: str, b: str) -> str:
    """The more advanced of two statuses."""
    return max(a or "Open", b or "Open", key=ORDER.index)


@lru_cache(maxsize=1)
def register_data() -> dict[str, Any]:
    return register_module().load()


def brd_of_short(short: str) -> str:
    for key, b in register_data()["brds"].items():
        if b["short"] == short:
            return key if key.startswith("BRD-") else "BRD-00"
    return "BRD-00"


def question_items(frs_questions: dict[str, dict[str, Any]]) -> list[Item]:
    """The open questions of the FRS and of the register, one item per question ID, with the latest status."""
    reg = {q["id"]: q for q in register_module().load_questions(register_data())}
    out: list[Item] = []
    for qid in list(frs_questions) + [q for q in reg if q not in frs_questions]:
        f, r = frs_questions.get(qid), reg.get(qid)
        default = f["brd"] if f else brd_of_short(r["brd"])
        home = home_of(qid, default)
        status, answer = "Open", ""
        if f:
            status, answer = frs_status(f["status_raw"])
        if r:
            reg_status = {"Partially answered": "Partly answered"}.get(r["status"], r["status"])
            if ORDER.index(reg_status) > ORDER.index(status):
                status, answer = reg_status, r["answer"]
            elif not answer and r["answer"] and status != "Open":
                answer = r["answer"]
        if r:
            need = f"{business(r['topic'])}: {business(r['question'])}"
        else:
            need = f["question"]
        sources = []
        if f:
            sources.append(f"{f['doc']}, Open questions {qid}" + (f" (affects {f['affects']})" if f["affects"] else ""))
        if r:
            sources.append(f"Discrepancy and clarification register v{register_data()['meta']['version']}, open "
                           f"question {qid}" + (f"; baseline {clean(r['source'])}" if r["source"] else ""))
        text = " ".join(x for x in ((f or {}).get("text", ""), (r or {}).get("question", ""),
                                     (r or {}).get("items", ""), (r or {}).get("same", "")) if x)
        it = Item(id=qid, kind="question", tab=home, need=need, source="; ".join(sources), status=status,
                  answer=business(answer), text=normalise_ids(text, home), topic=business((r or {}).get("topic", "")))
        it.group = (r or {}).get("under", "")  # type: ignore[attr-defined]
        out.append(it)
    return out


def dcr_items() -> list[Item]:
    reg = register_module()
    data = register_data()
    out = []
    for i in data["items"]:
        brds = [b for b in i["brds"] if b.startswith("BRD-")]
        home = brds[0] if len(set(brds)) == 1 else "BRD-00"
        need = business(i["question"])
        if need.rstrip(".").lower() in ("none", "-", ""):
            need = f"Note the proposed handling (no answer needed unless BDOI disagrees): {business(i['resolution'])}"
        elif len(need) < 40:
            need = f"{business(i['desc'])} {need}"
        it = Item(id=i["ref"], kind="dcr", tab=home, need=need,
                  why=business(f"{i['desc']} {i['impact']}"), proposal=business(i["resolution"]),
                  source=f"Discrepancy and clarification register v{data['meta']['version']}, {i['ref']} "
                         f"({i['type']}); {clean(i['loc'])}",
                  owner=business(reg.owner_of(i, data)), priority=i["sev"],
                  status="Answered" if i["status"] in ("Answered", "Closed") else "Open",
                  text=" ".join(str(i.get(k, "")) for k in ("loc", "desc", "question", "resolution")) + " " +
                  " ".join(i.get("related") or []), brds=sorted(set(i["brds"])))
        it.related = list(i.get("related") or [])  # type: ignore[attr-defined]
        if it.status == "Answered":
            it.answer = business(i.get("response") or i.get("bdoi_response") or "Answered in the register")
        out.append(it)
    return out


@lru_cache(maxsize=1)
def catalogue() -> dict[str, Any]:
    return yaml.safe_load((SRC / "BRD-13_Data_Migration" / "pack" / "catalogue.yaml").read_text(encoding="utf-8"))


def milestone_date(text: str) -> dt.date | None:
    m = re.search(r"\bM(\d)\b", str(text))
    if not m:
        return as_date(text)
    for code, _, date in catalogue()["milestones"]:
        if code == f"M{m.group(1)}":
            return as_date(date)
    return None


def brd13_items() -> list[Item]:
    """The proposed rules and clarifications (PR-DM, CLR-DM) and the open decisions of the BRD-13 set."""
    fm = front_matter(SRC / "BRD-13_Data_Migration" / "HANDBOOK_BRD13_DATA_MIGRATION.md")
    doc = f"Data Migration Handbook BRD-13 v{fm.get('version', '')}"
    out = []
    for p in catalogue()["proposals"]:
        pid, topic, rule, reason, requested, decisions, due = (list(p) + [""] * 7)[:7]
        out.append(Item(id=pid, kind="proposal", tab="BRD-13", need=f"{business(topic)}: {business(requested)}",
                        why=business(reason), proposal=business(rule),
                        source=f"{doc}, Proposed business rules and clarifications" +
                               (f"; decisions {clean(decisions)}" if decisions else ""),
                        due=milestone_date(due), topic=business(topic),
                        text=" ".join(str(x) for x in (topic, rule, reason, requested, decisions))))
    for d in catalogue()["decisions"]:
        did, topic, question, affects, milestone, owner, status = (list(d) + [""] * 7)[:7]
        answer = d[7] if len(d) > 7 else ""
        st = {"OPEN": "Open", "PARTIAL": "Partly answered", "RECOMMENDED": "Partly answered",
              "ANSWERED": "Answered", "CLOSED": "Answered"}.get(str(status).upper(), "Open")
        out.append(Item(id=did, kind="decision", tab="BRD-13", need=f"{business(topic)}: {business(question)}",
                        source=f"{doc}, Open decisions ({milestone}); affects {clean(affects)}",
                        owner=business(owner), due=milestone_date(milestone), status=st, answer=business(answer),
                        priority="High" if str(milestone).startswith("M1") else "Medium", topic=business(topic),
                        text=" ".join(str(x) for x in d)))
    return out


def drop0_list_items() -> list[Item]:
    data = yaml.safe_load(DROP0_LIST.read_text(encoding="utf-8"))
    out = []
    for brd, rows in data.items():
        for r in rows:
            due = as_date(r["needed_by"]) or as_date(cfg()["needed_by_text"].get(r["needed_by"]))
            need = business(r["need"]).rstrip("?").rstrip() + "."
            why = business(r["background"])
            if r.get("impact"):
                why = f"{why} If not provided: {business(r['impact'])}".strip()
            it = Item(id=r["ref"], kind="drop0", tab=brd, need=need, why=why, proposal=business(r["proposal"]),
                      source=f"Drop 0 requirements and questions list v1.0 (8-Oct-2026), {r['ref']}; "
                             f"{clean(r['sources'])}",
                      owner=business(r["owner"]), due=due, priority=r["priority"] or "Medium",
                      status={"Closed": "Answered"}.get(r["status"], r["status"] or "Open"),
                      type=cfg()["category_type"].get(r["category"], ""), topic=business(r["area"]),
                      text=" ".join(str(r.get(k, "")) for k in ("need", "background", "proposal", "sources")))
            out.append(it)
    return out


def nfr_items() -> list[Item]:
    data = yaml.safe_load((SRC / "programme" / "requirements_quality" / "nfr_catalogue.yaml").read_text("utf-8"))
    out = []
    for n in data["nfrs"]:
        if n["status"] not in ("DECISION", "PROPOSED"):
            continue
        decide = n["status"] == "DECISION"
        need = (f"{business(n['title'])}: decide the target; the BRDs give different values" if decide else
                f"{business(n['title'])}: confirm the proposed target or give the BDOI value")
        out.append(Item(id=n["id"], kind="nfr", tab="BRD-00", need=need, why=business(n["req"]),
                        proposal=business(n["target"]),
                        source=f"Non-functional requirements register v{data['meta']['version']}, {n['id']}; "
                               f"{clean(n.get('source', ''))}",
                        owner=business(n.get("acceptor") or ""), type="Decision" if decide else "Configuration value",
                        priority={"Must": "High", "Should": "Medium"}.get(n.get("prio"), "Low"),
                        status="Open", text=f"{n.get('source', '')} {n.get('frs', '')}"))
    return out


def module_tab(code: str) -> str:
    m = drop_closure().data()["modules"][code]
    return cfg()["module_tab"].get(code, m["brd"])


CONFIG_AREAS = {"System and security parameters", "Lists of values", "Document numbering", "Module rules and set-up",
                "Data migration reference data"}
DOCUMENT_ITEMS = {"CFG-082", "CFG-083", "CFG-098"}


def cfg_items() -> list[Item]:
    """The Drop 0 configuration register: one row per configuration item, with its template of workbook B."""
    dc = drop_closure()
    out = []
    for it, r in zip(dc.data()["items"], dc.register_rows()):
        if it["id"] in excluded().get("config_items", set()):
            continue
        target = b_target_text(it)
        if it["id"] in DOCUMENT_ITEMS:
            typ = "Sample document or template"
        elif it["route"] == "screen" or it["area"] in CONFIG_AREAS:
            typ = "Configuration value"
        else:
            typ = "Master data"
        due = dc.dues()[it["due"]]
        out.append(Item(id=it["id"], kind="cfg", tab=module_tab(it["module"]),
                        need=f"{business(it['name'])}: {business(it['provides'])}",
                        why=f"Needed for: {due['label']} ({due['note']}).",
                        proposal=target, source=f"Drop 0 configuration register {it['id']} (Drop 0 Configuration "
                                                f"Inputs v{dc.data()['meta']['version']}, area {it['area']})",
                        owner=business(r["owner"]), due=due["date"], type=typ,
                        priority="High" if it["due"] in ("D1", "D2", "D3") else "Medium", status="Open",
                        text="", module=dc.module(it["module"])["name"]))
    return out


def document_items() -> list[Item]:
    """One row per module: the BDOI layout and wording of the documents BIBS generates."""
    dc = drop_closure()
    mods = dc.data()["modules"]
    by_name = {m["name"]: code for code, m in mods.items()}
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in dc.document_rows():
        groups[r["module"]].append(r)
    out = []
    for n, (name, rows) in enumerate(sorted(groups.items()), start=1):
        code = by_name[name]
        titles = "; ".join(business(r["title"]) for r in rows)
        out.append(Item(id=f"DOC-{n:02d}", kind="doc", tab=module_tab(code), type="Sample document or template",
                        need=f"Documents of {name}: give the BDOI layout, logo placement and wording of the "
                             f"{len(rows)} documents BIBS generates ({titles}), as a filled-in sample of each.",
                        why="BIBS generates each document from its template; without the BDOI layout the documents "
                            "keep the proposed layout in SIT and UAT.",
                        proposal="The proposed layouts are reviewed on the Document Templates screen; BDOI marks the "
                                 "changes on a printout or sends its current form.",
                        source="Document templates of the platform (Drop 0 Configuration Inputs, Ref Document "
                               "templates); configuration register CFG-082",
                        owner=business(rows[0]["owner"]), due=dc.dues()["D5"]["date"], priority="Medium",
                        status="Open", remarks=["Review sheet Document templates of the Master Data and "
                                                "Configuration Upload Templates workbook."]))
    return out


def plan_data() -> dict[str, Any]:
    return yaml.safe_load((SRC / "programme" / "delivery_adoption" / "plan.yaml").read_text(encoding="utf-8"))


def dependency_items() -> list[Item]:
    c = cfg()
    out = []
    for d in plan_data()["dependencies"]:
        did, cat, what, owner, date, needed_for, impact = d
        if did in c["dependency_skip"]:
            continue
        out.append(Item(id=did, kind="dep", tab=c["dependency_tab"].get(did, "BRD-00"),
                        need=business(what), why=business(f"Needed for: {needed_for}. If late: {impact}."),
                        source=f"Project plan v1.0 (8-Oct-2026), dependency {did} ({cat})", owner=business(owner),
                        due=as_date(date), priority="High", status="Open",
                        type=c["dependency_type"].get(cat, "Clarification"), text=what))
    return out


def test_items() -> list[Item]:
    c = cfg()["test_data"]
    out = []
    for t in tabs():
        if t["brd"] in ("BRD-00", "RI"):
            continue
        fmtd = {"name": f"{t['brd']} {t['name']}", "drop": t["drop"]}
        out.append(Item(id=f"TD-{t['brd'][4:]}", kind="test", tab=t["brd"], type="Test data",
                        need=c["need"].format(**fmtd), why=c["why"].format(**fmtd), proposal=c["proposal"],
                        source=c["source"], owner=f"{t['owner']} (UAT testers of the business units)",
                        due=as_date(cfg()["type_due"]["Test data"]), priority="High", status="Open"))
    return out


def signoff_items() -> list[Item]:
    c = cfg()
    out = []
    for t in tabs():
        rows = list(c["signoff_rows"].get(t["brd"], []))
        if t["brd"] not in ("BRD-00", "RI"):
            rows += c["signoff_drop"].get(t["drop"], [])
        for n, (what, date, who) in enumerate(rows, start=1):
            f = {"brd": t["brd"], "name": t["name"], "owner": t["owner"]}
            out.append(Item(id=f"SO-{t['brd'][4:] if t['brd'].startswith('BRD') else t['brd']}-{n}", kind="signoff",
                            tab=t["brd"], type="Sign-off", need=what.format(**f) + ".",
                            why="The signed document fixes the scope for the next step of the plan; a late "
                                "sign-off moves the dates that follow it.",
                            source="Project plan v1.0 (8-Oct-2026), sign-offs and milestones", owner=who.format(**f),
                            due=as_date(date), priority="High", status="Open"))
    return out


def ri_items() -> list[Item]:
    out = []
    for r in cfg()["ri_rows"] + cfg().get("extra_rows", []):
        out.append(Item(id=r["ref"], kind="ri", tab=r.get("tab", "RI"), type=r["type"], need=r["need"], why=r["why"],
                        proposal=r["proposal"], source=r["source"], owner=r["owner"], due=as_date(r["due"]),
                        priority=r["priority"], status=r["status"],
                        remarks=[r["remarks"]] if r.get("remarks") else []))
    return out


# ============================================================================ owners


def owner_roles() -> list[str]:
    return [r[0] for r in cfg()["owner_roles"]]


def owner_role(text: str, tab_key: str) -> str:
    """The BDOI role of an owner text of the sources (first role named), one name per role."""
    first = re.split(r";| with |/", clean(text), maxsplit=1)[0]
    for pat, role in cfg()["owner_rules"]:
        if re.search(pat, first, re.I):
            return cfg()["tab_owner_role"][tab_key] if role == "TAB" else role
    for pat, role in cfg()["owner_rules"]:
        if role != "TAB" and re.search(pat, text, re.I):
            return role
    return cfg()["tab_owner_role"][tab_key]


# ============================================================================ consolidation


TYPE_REVIEW = HERE / "type_review.yaml"


@lru_cache(maxsize=1)
def type_review() -> dict[str, str]:
    """Type of each row after the row-by-row review: {row ID: type}."""
    if not TYPE_REVIEW.exists():
        return {}
    return {str(k): v for k, v in (yaml.safe_load(TYPE_REVIEW.read_text(encoding="utf-8")) or {}).items()}


def classify(it: Item) -> str:
    if it.id in type_review():
        return type_review()[it.id]
    if it.type:
        return it.type
    ask = it.need.split(": ", 1)[1] if it.topic and ": " in it.need else it.need
    for typ, pat, scope in cfg()["type_rules"]:
        if re.search(pat, ask if scope == "ask" else f"{it.topic} {it.need}", re.I):
            return typ
    return "Clarification"


@lru_cache(maxsize=1)
def consolidated() -> list[Item]:
    """Every row of workbook A, de-duplicated: a question, register item or earlier list item already covered by a
    row is attached to it (its ID stays in the Source)."""
    clr, frs_q = frs_items()
    dm = brd13_items()
    drop0 = drop0_list_items()
    questions = question_items(frs_q)
    dcrs = dcr_items()
    primaries: list[Item] = []
    by_id: dict[str, Item] = {}  # every ID -> the row that holds it

    def add(it: Item) -> None:
        primaries.append(it)
        by_id[it.id] = it

    def attach(it: Item, row: Item) -> None:
        row.attached.append(it)
        by_id[it.id] = row

    def holder(it: Item, candidates: list[Item]) -> Item | None:
        if not candidates:
            return None
        same = [c for c in candidates if c.tab == it.tab]
        return (same or candidates)[0]

    for it in clr:
        add(it)
    proposals = [x for x in dm if x.kind == "proposal"]
    decisions = [x for x in dm if x.kind == "decision"]
    for it in proposals:
        add(it)
    # BRD-13 decisions named by a proposal are merged into it
    for d in decisions:
        cands = [p for p in proposals if d.id in p.refs]
        h = holder(d, cands)
        attach(d, h) if h else add(d)
    # earlier Drop 0 list: merged into the clarification, proposal or decision it names
    for it in drop0:
        cands = [by_id[r] for r in it.refs if r in by_id and by_id[r].kind in ("clr", "proposal", "decision")]
        h = holder(it, cands)
        attach(it, h) if h else add(it)
    for it in nfr_items():
        add(it)
    # questions: merged into a row that names them, else into the head of their group, else rows
    referenced: dict[str, list[Item]] = defaultdict(list)
    for row in primaries:
        for r in [row.id] + row.refs + [r for a in row.attached for r in a.refs]:
            referenced[r].append(row)
    q_by_id = {q.id: q for q in questions}
    pending = []
    for q in questions:
        cands = [r for r in referenced.get(q.id, []) if r is not q]
        if q.id in by_id:  # a BRD-13 decision with the same ID: the decision row holds it
            attach(q, by_id[q.id])
            continue
        h = holder(q, cands)
        if h:
            attach(q, h)
        else:
            pending.append(q)
    for q in pending:
        head = getattr(q, "group", "")
        if head and head in q_by_id and head != q.id:
            continue
        add(q)
    for q in pending:
        head = getattr(q, "group", "")
        if head and head in q_by_id and head != q.id and q.id not in by_id:
            target = by_id.get(head)
            attach(q, target) if target else add(q)
    # register items: merged into the row that names them or holds one of their questions
    for row in primaries:
        for r in row.refs + [r for a in row.attached for r in a.refs]:
            referenced[r].append(row)
    for d in dcrs:
        if d.id in by_id:
            attach(d, by_id[d.id])
            continue
        cands = [r for r in referenced.get(d.id, []) if r is not d]
        cands += [by_id[q] for q in getattr(d, "related", []) if q in by_id]
        h = holder(d, cands)
        attach(d, h) if h else add(d)
    for it in cfg_items() + document_items() + dependency_items() + test_items() + signoff_items() + ri_items():
        add(it)
    # reinsurance points move to their tab
    ri = set(cfg()["ri_ids"])
    for row in primaries:
        if row.id in ri or any(a.id in ri for a in row.attached):
            if row.kind not in ("clr",):
                row.tab = "RI"
    finish(primaries)
    return primaries


def finish(rows: list[Item]) -> None:
    """Type, owner, due date, status, priority and the traces of every row."""
    for n, row in enumerate(rows):
        row.order = n
        t = tab(row.tab)
        row.type = classify(row)
        if not row.owner:
            owners = [a.owner for a in row.attached if a.owner]
            row.owner = owners[0] if owners else t["owner"]
            if row.type == "Integration detail" and "IT" not in row.owner:
                row.owner = f"BDOI IT with the {row.owner}"
        if not row.due:
            dues = [a.due for a in row.attached if a.due]
            row.due = min(dues) if dues else as_date(cfg()["type_due"].get(row.type)) or t["due"]
            if row.type == "Integration detail" and t["due"] and row.due < t["due"]:
                row.due = t["due"]
        # status: the row's own and that of the questions, decisions and items merged into it
        statuses = [row.status] if row.status else []
        statuses += [a.status for a in row.attached if a.status and a.kind in ("question", "decision", "dcr")]
        answers = [a.answer[:260] for a in [row] + [x for x in row.attached if x.kind in ("question", "decision", "dcr")
                                                   and x.id in row.refs + [row.id]]
                   if a.answer and a.status in ("Answered", "Partly answered")][:3]
        if not statuses:
            m = re.search(r"\bBDOI (?:decision|answer|decided|answered)\b[^.;]*", row.why, re.I)
            if m:
                statuses, answers = ["Partly answered"], [m.group(0)]
            else:
                statuses = ["Open"]
        if all(s == "Answered" for s in statuses):
            row.status = "Answered"
        elif any(s in ("Answered", "Partly answered") for s in statuses):
            row.status = "Partly answered"
        else:
            row.status = "Open"
        if answers and row.status != "Open":
            seen = []
            for a in answers:
                if a not in seen:
                    seen.append(a)
            row.answer = "; ".join(seen)
        row.owner = owner_role(row.owner, row.tab)
        prios = [row.priority] + [a.priority for a in row.attached if a.kind in ("dcr", "drop0", "decision")]
        row.priority = min((p for p in prios if p in PRIORITIES), key=PRIORITIES.index, default="Medium")
        also, seen_ids = [], {row.id}
        for a in row.attached:
            if a.id not in seen_ids:
                also.append(a)
                seen_ids.add(a.id)
        if also:
            label = {"question": "open question", "dcr": "register item", "drop0": "Drop 0 list item",
                     "decision": "open decision", "clr": "clarification"}
            row.source += "; also raised as " + ", ".join(f"{a.id} ({label.get(a.kind, a.kind)})" for a in also)
        other = sorted({b for a in [row] + row.attached for b in a.brds if b.startswith("BRD-") and b != row.tab})
        if other:
            row.remarks.append(f"Also affects {', '.join(other)}.")
        if row.answer:
            row.remarks.insert(0, f"Answer so far: {row.answer}.".replace("..", "."))
        if not row.why and row.kind == "question":
            affects = re.search(r"affects ([^)]*)\)", row.source)
            row.why = ("Open point of the FRS" + (f" ({affects.group(1)})" if affects else "") +
                       ": the requirement cannot be finalised or tested as BDOI intends until it is answered.")
        for a in row.attached:
            if a.kind == "drop0" and a.proposal and not row.proposal:
                row.proposal = a.proposal
            if a.why and not row.why:
                row.why = a.why
            if a.proposal and not row.proposal:
                row.proposal = a.proposal


# ============================================================================ workbook A

A_COLUMNS = [
    ("ref", "Ref", 13, "Original ID of the point: CLR clarification of an FRS, PR-DM proposed rule or DMQ decision of the "
                       "Data Migration set, open question (RQ06, AQ12 ...), register item DCR, Drop 0 list item "
                       "(PM-Q, UA-Q, DM-Q), configuration item CFG, NFR target, plan dependency DEP, documents DOC, "
                       "test data TD, sign-off SO, reinsurance RI"),
    ("module", "BRD / module", 18, "BRD and module the point belongs to"),
    ("type", "Type", 16, "What kind of input is asked for"),
    ("need", "What we need from BDOI", 60, "The decision, answer, value, data or document BDOI provides"),
    ("why", "Why / impact if not provided", 46, "Why it is needed and what happens without it"),
    ("proposal", "Our proposal (if any)", 46, "What the project team proposes; it applies as the working assumption "
                                               "until BDOI answers"),
    ("source", "Source (FRS section / CLR / BRD page)", 40, "Where the point is raised, with the IDs of the same point "
                                                             "in the other documents"),
    ("owner", "BDOI owner (role)", 26, "Role that answers or provides the input"),
    ("due", "Needed by", 12, "Date the input is needed by (project plan of 8-Oct-2026)"),
    ("priority", "Priority", 10, "High: blocks a sign-off, the set-up or a test window; Medium: needed before UAT; "
                                 "Low: can follow"),
    ("status", "Status", 14, "Open, Partly answered or Answered (from the latest sources)"),
    ("response", "BDOI response", 40, "The answer of BDOI (to fill in)"),
    ("remarks", "Remarks", 34, "Answer so far, other BRDs affected, template to fill in"),
]


def a_rows(t: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [r for r in consolidated() if r.tab == t["brd"]]
    rows.sort(key=lambda r: (r.due or dt.date(2099, 1, 1), PRIORITIES.index(r.priority), r.order))
    out = []
    for r in rows:
        module = f"{t['brd']} {t['name']}" if t["brd"] != "RI" else t["name"]
        if r.module:
            module = f"{t['brd']} {r.module}" if t["brd"] != "RI" else r.module
        out.append({"ref": r.id, "module": module, "type": r.type, "need": business(r.need), "why": business(r.why),
                    "proposal": business(r.proposal) or None, "source": business(r.source),
                    "owner": business(r.owner), "due": r.due, "priority": r.priority, "status": r.status,
                    "response": None, "remarks": business(" ".join(r.remarks)) or None})
    return out


def build_a() -> Path:
    from bdoi_xlsx import BdoiWorkbook, Column, _SheetInfo  # noqa: PLC0415
    from openpyxl.formatting.rule import CellIsRule  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415

    meta = cfg()["meta"]
    wb = BdoiWorkbook(meta["title_a"], doc_type="BDOI inputs workbook", brd="BRD-00 (all BRDs)",
                      version=meta["version"], date=meta["date"],
                      subtitle="What the project needs from the BDOI business users, BRD by BRD")
    wb.legend = [("Open", "Waiting for the BDOI answer"),
                 ("Partly answered", "BDOI has answered in part; the rest of the point is open"),
                 ("Answered", "BDOI has answered; the answer is applied and BDOI confirms it")]
    types = cfg()["types"]
    cols = []
    for key, head, width, desc in A_COLUMNS:
        values = {"type": types, "priority": PRIORITIES, "status": STATUSES, "owner": owner_roles()}.get(key)
        cols.append(Column(key, head, width, desc, values=values, status=key == "status",
                           kind="date" if key == "due" else "text"))
    owners_ws = wb.sheet("Owners", [
        Column("role", "BDOI owner (role)", 30, "Role named in the Owner column of the BRD tabs"),
        Column("who", "Who", 60, "Who holds the role"),
        Column("org", "Organisation", 18, "Organisation"),
        Column("answers", "Answers or provides", 50, "What the role answers or provides"),
        Column("items", "Items", 8, "Rows of this workbook owned by the role", kind="number"),
    ], [{"role": r, "who": w, "org": o, "answers": a,
         "items": sum(1 for x in consolidated() if x.owner == r)} for r, w, o, a in cfg()["owner_roles"]],
        description="The BDOI owner roles used in the Owner column (one name per role), from the programme RACI "
                    "and the BRD approval sheets")
    sheets = []
    for t in tabs():
        rows = a_rows(t)
        ws = wb.sheet(t["sheet"], cols, rows, description=f"{t['brd'] if t['brd'] != 'RI' else 'Phase 2'} "
                                                          f"{t['name']}: what BDOI provides; owner "
                                                          f"{cfg()['tab_owner_role'][t['brd']]}")
        rng = f"K5:K{5 + max(len(rows), 1) + 500}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Partly answered"'], stopIfTrue=True,
                                                      fill=PatternFill("solid", fgColor=brand.BG_BLUE,
                                                                       bgColor=brand.BG_BLUE),
                                                      font=Font(color=brand.HEADER_BLUE, bold=True)))
        sheets.append((t, ws, rows))
    instructions = wb.wb.create_sheet("Instructions")
    summary = wb.wb.create_sheet("Summary")
    _instructions(instructions, wb)
    _summary(summary, wb, sheets)
    for ws, desc in ((instructions, "How to read and answer this workbook"),
                     (summary, "Counts per BRD and type, status, owners and due dates")):
        wb._sheets.insert(0 if ws is instructions else 1, _SheetInfo(ws.title, desc, [], 0))
        wb._print_setup(ws, None)
    del Alignment
    wb._build_cover()
    wb._build_readme()
    order = [wb._cover, instructions, summary, owners_ws] + [s[1] for s in sheets] + [wb._readme, wb._lists]
    wb.wb._sheets = order
    wb.wb.active = 0
    props = wb.wb.properties
    props.title, props.creator, props.subject = wb.title, brand.VENDOR, brand.SYSTEM
    props.keywords = "BDOI inputs, BRD-00, BIBS, BDOI"
    path = brand.out_path("BRD-00", KIND, brand.output_name("Inputs", "BRD-00", meta["title_a"], meta["version"],
                                                              "xlsx"))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.wb.save(path)
    return path


def _cell(ws, row: int, col: int, value: Any, size: float = 10, bold: bool = False, colour: str = brand.TEXT,
          fill: str | None = None, wrap: bool = True, fmt_: str | None = None, border: bool = False):
    from bdoi_xlsx import BORDER  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415

    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=brand.FONT, size=size, bold=bold, color=colour)
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt_:
        c.number_format = fmt_
    if border:
        c.border = BORDER
    return c


TYPE_MEANING = {
    "Decision": "BDOI chooses between options or sets the scope; the choice changes what is delivered.",
    "Clarification": "The BRD text is unclear, incomplete or contradictory; BDOI explains what is meant.",
    "Business rule confirmation": "The project team proposes a rule; BDOI confirms it or gives its own.",
    "Configuration value": "A value BIBS is set up with: limits, days, rates, lists, numbering, parameters.",
    "Master data": "Records BIBS needs before go-live: organisation, users, products, insurers, accounts, clients.",
    "Sample document or template": "A letter, form, advice, slip or e-mail in the BDOI layout and wording.",
    "Report layout": "The columns, totals and filters of a report or an extract.",
    "Integration detail": "A file layout, interface, transport or contact of another system.",
    "Test data": "Testers, cases and business data for SIT and UAT.",
    "Sign-off": "A document or test result BDOI signs.",
}


def _instructions(ws, wb) -> None:
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    ws.sheet_view.showGridLines = False
    for i, w in enumerate([28, 70, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _cell(ws, 1, 1, "Instructions", 15, True, brand.HEADER_BLUE, wrap=False)
    _cell(ws, 2, 1, f"{wb.title}, version {wb.version}, {wb.date}", 9, colour=brand.MUTED, wrap=False)
    paras = [
        ("Purpose", "This workbook lists, BRD by BRD, everything the project needs from the BDOI business users to "
                    "complete the requirements, set BIBS up and test it: decisions, clarifications, confirmations of "
                    "the proposed rules, configuration values, master data, sample documents, report layouts, "
                    "integration details, test data and sign-offs. It answers the request of the BDOI product owner "
                    "of 8 October 2026 for the BRD-specific requirements from the BDOI users."),
        ("Where the rows come from", "The rows bring together the proposed business rules and clarifications (CLR) "
                                     "and the open questions of every FRS and of the Data Migration set, the "
                                     "discrepancy and clarification register, the earlier Drop 0 requirements list, "
                                     "the Drop 0 configuration register, the NFR targets and the dependencies, "
                                     "sign-offs and test windows of the project plan. A point raised in several "
                                     "documents is listed once; the Source column gives every ID under which it was "
                                     "raised, so an answer given against any of them is found here."),
        ("One tab per BRD", "BRD-00 holds the cross-cutting points (several BRDs, programme, NFR targets, "
                            "infrastructure and integrations). BRD-01 to BRD-13 hold the points of each BRD. The "
                            "last tab holds the few points of ReInsurance, which is phase 2."),
        ("How to answer", "Write the answer in BDOI response, set Status to Answered (or Partly answered) and send "
                          "the workbook back to the iorta TechNXT project team. Where a proposal is given, an answer "
                          "of 'Agree' is enough; otherwise give the decision, the value or attach the document and "
                          "name it in the response. Do not change the Ref, so that the answer is traced back to the "
                          "FRS, the register and the templates."),
        ("Needed by and priority", "Needed by is the date of the project plan of 8 October 2026 by which the input "
                                   "is needed: the sign-off of the FRS or sign-off set of the BRD for decisions and "
                                   "clarifications, the configuration milestones (D1 to D7) for values and master "
                                   "data, the integration specifications for interface details and the UAT "
                                   "preparation for test data. High blocks a sign-off, the set-up or a test window; "
                                   "Medium is needed before UAT; Low can follow."),
        ("Status", "The status is taken from the latest sources: Answered where BDOI has answered (the answer is in "
                   "Remarks and BDOI confirms it), Partly answered where part of the point is answered, Open "
                   "otherwise. Answered points stay in the list so that every answer is visible."),
        ("Master data and configuration", "The rows of type Master data and Configuration value name the template to "
                                          "fill in. The templates are in the companion workbook "
                                          "BIBS_Templates_BRD-00_Master_Data_and_Configuration_Upload_Templates_v1.0, "
                                          "with its guide; fill the data there, not in this workbook."),
        ("Proposals", "Until BDOI answers, the proposal of a row is the working assumption of the project team; "
                      "the screens, tests and set-up follow it, and a different answer is applied by the change "
                      "control of the programme."),
    ]
    row = 4
    for head, text in paras:
        _cell(ws, row, 1, head, 10, True, brand.HEADER_BLUE, fill=brand.BG_BLUE, border=True)
        c = _cell(ws, row, 2, text, border=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.row_dimensions[row].height = 15 * max(2, len(text) // 105 + 1)
        del c
        row += 1
    row += 1
    _cell(ws, row, 1, "Types of input", 12, True, brand.HEADER_BLUE, wrap=False)
    row += 1
    for c, h in enumerate(["Type", "What BDOI provides"], start=1):
        _cell(ws, row, c, h, 10, True, brand.WHITE, fill=brand.HEADER_BLUE, border=True)
    row += 1
    for typ in cfg()["types"]:
        _cell(ws, row, 1, typ, 10, True, border=True)
        _cell(ws, row, 2, TYPE_MEANING[typ], border=True)
        row += 1


def _summary(ws, wb, sheets) -> None:
    from openpyxl.styles import Font  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    ws.sheet_view.showGridLines = False
    types = cfg()["types"]
    heads = ["BRD", "Name", "Drop", "BDOI owner (role)", "Items"] + types + STATUSES + ["First needed by",
                                                                                       "Last needed by"]
    widths = [9, 26, 9, 30, 7] + [11] * len(types) + [9, 9, 9, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _cell(ws, 1, 1, "Summary", 15, True, brand.HEADER_BLUE, wrap=False)
    _cell(ws, 2, 1, "Counts per BRD, type and status follow the tabs as they are answered (formulas); owners and "
                    "dates as issued.", 9, colour=brand.MUTED, wrap=False)
    row = 4
    _cell(ws, row, 1, "Items per BRD and type", 12, True, brand.HEADER_BLUE, wrap=False)
    row += 1
    for c, h in enumerate(heads, start=1):
        _cell(ws, row, c, h, 9, True, brand.WHITE, fill=brand.HEADER_BLUE, border=True)
    ws.row_dimensions[row].height = 42
    first = row + 1
    for t, sheet, rows in sheets:
        row += 1
        q = f"'{sheet.title}'"
        last = 5 + max(len(rows), 1) + 500
        vals: list[Any] = [t["brd"] if t["brd"] != "RI" else "Phase 2", t["name"], t["drop"],
                           cfg()["tab_owner_role"][t["brd"]],
                           f"=COUNTA({q}!A5:A{last})"]
        vals += [f'=COUNTIF({q}!C5:C{last},"{x}")' for x in types]
        vals += [f'=COUNTIF({q}!K5:K{last},"{x}")' for x in STATUSES]
        vals += [f"=MIN({q}!I5:I{last})", f"=MAX({q}!I5:I{last})"]
        for c, v in enumerate(vals, start=1):
            cell = _cell(ws, row, c, v, 9, c == 1, border=True, fmt_="dd-mmm-yyyy" if c > len(vals) - 2 else None)
            if c == 1:
                cell.hyperlink = f"#{q}!A1"
                cell.font = Font(name=brand.FONT, size=9, bold=True, color=brand.CTA_BLUE, underline="single")
    row += 1
    _cell(ws, row, 1, "Total", 9, True, border=True, fill=brand.BG_BLUE)
    for c in range(2, len(heads) + 1):
        letter = get_column_letter(c)
        if c <= 4:
            _cell(ws, row, c, None, border=True, fill=brand.BG_BLUE)
        elif c > len(heads) - 2:
            fn = "MIN" if c == len(heads) - 1 else "MAX"
            _cell(ws, row, c, f"={fn}({letter}{first}:{letter}{row - 1})", 9, True, border=True, fill=brand.BG_BLUE,
                  fmt_="dd-mmm-yyyy")
        else:
            _cell(ws, row, c, f"=SUM({letter}{first}:{letter}{row - 1})", 9, True, border=True, fill=brand.BG_BLUE)
    ws.freeze_panes = ws.cell(row=first, column=2)
    # owners
    row += 2
    _cell(ws, row, 1, "Items per BDOI owner (role), as issued", 12, True, brand.HEADER_BLUE, wrap=False)
    row += 1
    for c, h in enumerate(["Owner (role)", None, None, None, "Items", "Open", "Partly answered", "Answered",
                           "First needed by", "BRDs"], start=1):
        if h:
            _cell(ws, row, c, h, 9, True, brand.WHITE, fill=brand.HEADER_BLUE, border=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    by_owner: dict[str, list[Item]] = defaultdict(list)
    for r in consolidated():
        by_owner[r.owner].append(r)
    ranked = sorted(by_owner.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    for owner, items in ranked:
        row += 1
        st = Counter(i.status for i in items)
        brds = sorted({i.tab for i in items})
        vals = [owner, None, None, None, len(items), st["Open"], st["Partly answered"], st["Answered"],
                min((i.due for i in items if i.due), default=None), ", ".join(brds)]
        for c, v in enumerate(vals, start=1):
            _cell(ws, row, c, v, 9, border=True, fmt_="dd-mmm-yyyy" if c == 9 else None)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.merge_cells(start_row=row, start_column=10, end_row=row, end_column=14)
    # due months
    row += 2
    _cell(ws, row, 1, "Items by month needed, as issued", 12, True, brand.HEADER_BLUE, wrap=False)
    row += 1
    for c, h in enumerate(["Month", None, "Items", "High", "Medium", "Low"], start=1):
        if h:
            _cell(ws, row, c, h, 9, True, brand.WHITE, fill=brand.HEADER_BLUE, border=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    by_month: dict[str, list[Item]] = defaultdict(list)
    for r in consolidated():
        if r.due:
            by_month[r.due.strftime("%Y-%m")].append(r)
    for month, items in sorted(by_month.items()):
        row += 1
        p = Counter(i.priority for i in items)
        label = dt.datetime.strptime(month, "%Y-%m").strftime("%b %Y")
        for c, v in enumerate([label, None, len(items), p["High"], p["Medium"], p["Low"]], start=1):
            _cell(ws, row, c, v, 9, border=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)


# ============================================================================ workbook B: the templates

ROUTE_UPLOAD = "Screen upload"
ROUTE_MIGRATION = "Migration Console"
ROUTE_SCREEN = "Configuration screen"
ROUTES = [ROUTE_UPLOAD, ROUTE_MIGRATION, ROUTE_SCREEN]
# Templates of the Drop 0, BRD-03 and BRD-11 sources whose rows the platform takes by file.
UPLOAD_TEMPLATES = {"D0-06", "UA-01"}
MIGRATION_TEMPLATES = {"PM-01"}  # delivered in the layout of R05; the set-up values on Products
LAYOUT_BRD = {"R01": "BRD-13", "R02": "BRD-05", "R03": "BRD-01", "R04": "BRD-03", "R04B": "BRD-03", "R05": "BRD-03",
              "R06": "BRD-03", "R07": "BRD-03", "R08": "BRD-05", "R11": "BRD-02", "C01": "BRD-01", "C02": "BRD-01",
              "C03": "BRD-05", "G01": "BRD-05"}
OWN_BRD = {"UP-01": "BRD-01", "UP-02": "BRD-05", "UP-03": "BRD-10", "MD-01": "BRD-05"}
REVIEW = {  # review sheets of the platform lists: key -> sheet name
    "LIST-LOV": "Lists of values",
    "LIST-PARAMETERS": "System parameters",
    "LIST-SLA": "SLA and turnaround",
    "LIST-NUMBERING": "Numbering prefixes",
    "LIST-NOTIFICATIONS": "Notification events",
    "LIST-DOCUMENTS": "Document templates",
    "LIST-EVENTS": "Accounting events",
}
SCREENS_SHEET = "Configured on screens"
CHARGES_SHEET = "PM-04 Charges checklist"
TAX_SHEET = "D0-10 Tax checklist"
COVERAGE_SHEET = "Coverage check"
GAPS_SHEET = "Gaps"
SLA = re.compile(r"SLA|_TAT_|_TAT$|TURNAROUND|REPLY_DAYS|FOLLOW_?UP_DAYS|ESCALATION|AGEING|AGING|REMINDER|LEAD_DAYS|"
                 r"NOTICE_DAYS|PAST_DUE|STALLED|ALERT_DAYS|GRACE|COMPARATIVE|ADVICE_DAYS|VALIDITY_DAYS|EXPIRING_DAYS|"
                 r"WORKING_DAYS")


def all_template_ids() -> list[str]:
    return [tid for g in cfg()["groups"] for tid in g["templates"]]


def group_of(tid: str) -> str:
    return next(g["name"] for g in cfg()["groups"] if tid in g["templates"])


@lru_cache(maxsize=1)
def dm_catalogue():
    return dm_module().Catalogue()


def _choices_of(expr: str, folder: Path) -> list[tuple[str, str]]:
    """Codes of .choices(...) / .codes(...) / .values(...) of a column definition."""
    import code_facts as cf  # noqa: PLC0415

    pairs = re.findall(r'new Choice\(\s*"([^"]+)"\s*,\s*"([^"]*)"\s*\)', expr)
    if pairs:
        return pairs
    m = re.fullmatch(r"\s*(\w+)\.class\s*", expr)
    if m:
        for p in cf.JAVA_ROOT.rglob(f"{m.group(1)}.java"):
            body = cf._strip_comments(p.read_text(encoding="utf-8"))
            e = re.search(r"enum\s+\w+\s*\{(.*?)(?:;|\})", body, re.S)
            if e:
                return [(c.strip(), c.strip().replace("_", " ").title()) for c in e.group(1).split(",") if c.strip()]
    lits = re.findall(r'"([^"]+)"', expr)
    if lits:
        return [(x, "") for x in lits]
    m = re.fullmatch(r"\s*(?:\w+\.)?(\w+)\s*", expr)
    if m:
        for p in list(folder.glob("*.java")) + list(cf.JAVA_ROOT.rglob("*.java")):
            body = cf._strip_comments(p.read_text(encoding="utf-8"))
            d = re.search(rf"\b{m.group(1)}\s*=\s*(List\.of\(.*?\));", body, re.S)
            if d:
                return re.findall(r'new Choice\(\s*"([^"]+)"\s*,\s*"([^"]*)"\s*\)', d.group(1))
    return []


def upload_columns(cls: str) -> list[dict[str, Any]]:
    """The columns of an upload definition of the platform, with their guide (when, list of values, codes, format,
    allowed values), read from its columns()."""
    import code_facts as cf  # noqa: PLC0415

    path = next(cf.JAVA_ROOT.rglob(f"{cls}.java"))
    text = cf._strip_comments(path.read_text(encoding="utf-8"))
    consts = {**cf._global_constants_simple(), **cf._package_constants(path.parent), **cf._java_constants(text)}
    body = cf._method_body(text, "columns") or ""
    m = re.search(r"List\.of\(", body)
    inner = body[m.end():cf.matching(body, m.end() - 1)]
    out = []
    for arg in cf.split_top(inner):
        base = cf._columns_of(arg, consts, path.parent)
        if not base:
            continue
        col = dict(base[0])
        for mm in re.finditer(r"\.(when|lov|codes|choices|values|format|allowed|master|example)\(", arg):
            start = mm.end() - 1
            val = arg[start + 1:cf.matching(arg, start)]
            key = mm.group(1)
            if key in ("codes", "choices", "values"):
                col["codes"] = _choices_of(val, path.parent)
            else:
                col[key] = cf.java_text(val, consts, cls)
        out.append(col)
    return out


def _upload_template(tid: str, spec: dict[str, Any], book) -> Any:
    import guided_xlsx as g  # noqa: PLC0415
    import sql_facts  # noqa: PLC0415

    cols = []
    lov_values = sql_facts.rows("lov_value", ("type_code", "code"))
    lov_types = {r["code"]: r["name"] for r in sql_facts.rows("lov_type")}
    for c in upload_columns(spec["handler"]):
        header = c["header"]
        mand = "Y" if c["required"] else (f"Cond.: {c['when']}" if c.get("when") else "N")
        fmt_ = {"DATE": "Date (dd-mmm-yyyy or yyyy-mm-dd)", "NUMBER": "Number", "YES_NO": "Y or N"}.get(c["type"], "Text")
        if c.get("format"):
            fmt_ = c["format"]
        check, allowed = "", c.get("allowed", "")
        codes = (spec.get("codes") or {}).get(header) or c.get("codes")
        if c["type"] == "YES_NO":
            check = "yn"
        elif c["type"] == "DATE":
            check = "date"
        elif c["type"] == "NUMBER":
            check = "number"
        elif codes:
            code = re.sub(r"[^A-Z0-9]+", "_", f"{tid}_{header}".upper()).strip("_")
            book.add_list(code, f"{header} ({tid})", [tuple(x) if isinstance(x, (list, tuple)) else (x, "")
                                                     for x in codes])
            check = f"list:{code}"
        elif c.get("lov"):
            values = [(v["code"], v["label"]) for v in lov_values if v["type_code"] == c["lov"]]
            if values:
                code = f"LOV_{c['lov']}"
                book.add_list(code, lov_types.get(c["lov"], c["lov"].replace("_", " ").title()), values,
                              "The values of the list as delivered with BIBS; a new value is added on Lists of Values "
                              "first.", strict=False)
                check = f"list:{code}"
        what = c.get("description", "")
        if c.get("master"):
            what += f" (a {c['master']} known to BIBS)"
        example = str(c.get("example") or "-")
        if check == "date" and re.fullmatch(r"\d{4}-\d{2}-\d{2}", example):
            example = dt.date.fromisoformat(example).strftime("%d-%b-%Y")
        cols.append(g.GuideColumn(header, mand, fmt_, business(what) or header, example, check, allowed=allowed))
    return g.Template(tid, spec["name"], f"{tid} {spec['sheet']}"[:31], spec["purpose"], spec["owner"],
                      due_text(tid), spec["load"], cols, source="Platform upload definition (column layout as the "
                                                                  "screen loads it)",
                      depends=list(cfg()["depends"].get(tid, [])), capacity=300)


def _screen_template(tid: str, spec: dict[str, Any], book) -> Any:
    import guided_xlsx as g  # noqa: PLC0415

    for code, lst in (spec.get("lists") or {}).items():
        book.add_list(code, lst["name"], [tuple(v) for v in lst["values"]], strict=lst.get("strict", True))
    cols = [g.GuideColumn(h, m, f, w, e or "-", c) for h, m, f, w, e, c in spec["columns"]]
    return g.Template(tid, spec["name"], f"{tid} {spec['sheet']}"[:31], spec["purpose"], spec["owner"], due_text(tid),
                      spec["load"], cols, source=spec.get("source", ""), depends=list(cfg()["depends"].get(tid, [])),
                      capacity=300)


def _layout_template(code: str, book) -> Any:
    m = dm_module()
    cat = dm_catalogue()
    layout = cat.layouts[code]
    t = m.layout_template(cat, layout)
    for c in t.columns:
        if c.check.startswith("list:"):
            raw = next(x for x in layout.columns if x["name"] == c.header).get("allowed_values")
            values = [v.strip() for v in m.client_text(raw).split(",")]
            book.add_list(c.check[5:], f"{layout.code} {c.header}", [(v, "") for v in values])
    for c in t.columns:
        if c.allowed_link == "Code maps!A1":
            c.allowed_link = ""
            c.allowed = c.allowed.replace("(sheet Code maps)", "(Code Maps screen of the Data Migration Console; "
                                                               "maps prepared in the Migration Workbook of BRD-13)")
    o = cat.by_code.get(layout.object, {})
    t.extra_rows, t.input_rows, t.extra_lists = ["Check when loaded"], [], {}
    t.confirm = ("Confirmed by (BDOI owner)", "Date")
    t.depends = list(cfg()["depends"].get(code, []))
    t.load = (f"Migration Console: the file {layout.template} (the header row as below) and its control file are "
              f"received on Data Migration › Data Migration › Extracts and loaded by a batch of object "
              f"{layout.object}; lands in: {o.get('target', '-')}")
    t.capacity = 200
    return t


@lru_cache(maxsize=1)
def drop_templates_b() -> dict[str, Any]:
    """The Drop 0, BRD-03 and BRD-11 templates of workbook B (copies), with the replaced templates' codes moved to the
    layouts that replace them."""
    import copy  # noqa: PLC0415

    replaced = {"PM-05": ("R04", "insurer_code"), "PM-07": ("R06", "bibs_package_version")}
    out = {}
    for t in drop_closure().guided():
        if t.id in cfg()["replaced"]:
            continue
        t = copy.deepcopy(t)
        for c in t.columns:
            kind, _, arg = c.check.partition(":")
            ref, _, _col = arg.partition("/")
            if kind in ("code", "ref") and ref in replaced:
                c.check = f"ref:{replaced[ref][0]}/{replaced[ref][1]}"
        t.depends = [replaced.get(d, (d,))[0] for d in t.depends]
        for d in cfg()["depends"].get(t.id, []):
            if d not in t.depends:
                t.depends.append(d)
        out[t.id] = t
    return out


def due_text(tid: str) -> str:
    code = cfg()["due"].get(tid)
    if code:
        d = drop_closure().dues()[code]
        return f"{drop_closure().fmt(d['date'])} – {d['label']}"
    return ""


def due_of(text: str) -> dt.date | None:
    """The date of a due text: '4-Jan-2027 – ...' or 'M3 (29 Jan 2027)'."""
    m = re.search(r"(\d{1,2})[- ]([A-Z][a-z]{2})[- ](\d{4})", text)
    if m:
        return dt.datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", "%d %b %Y").date()
    return milestone_date(text)


def route_of(tid: str) -> str:
    if tid in UPLOAD_TEMPLATES or tid.startswith("UP-"):
        return ROUTE_UPLOAD
    if tid in dm_catalogue().layouts or tid in MIGRATION_TEMPLATES:
        return ROUTE_MIGRATION
    return ROUTE_SCREEN


def brd_of_template(tid: str) -> str:
    if tid in LAYOUT_BRD:
        return LAYOUT_BRD[tid]
    if tid in OWN_BRD:
        return OWN_BRD[tid]
    if tid.startswith("PM-"):
        return "BRD-03"
    if tid.startswith("UA-"):
        return "BRD-11"
    dc = drop_closure()
    for it in dc.data()["items"]:
        if it.get("template") == tid:
            return module_tab(it["module"])
    return "BRD-05"


def where_of(tid: str, t: Any) -> str:
    """The screen or console where the rows are entered or loaded."""
    dc = drop_closure()
    for it in dc.data()["items"]:
        if it.get("template") == tid and it.get("screen"):
            return dc.menu()[it["screen"]]
    if route_of(tid) == ROUTE_MIGRATION and tid in dm_catalogue().layouts:
        return "Data Migration › Data Migration › Extracts and Batches"
    m = re.match(r"(?:Configuration screen|Screen upload): ([^,;(]+)", t.load)
    return m.group(1).strip() if m else business(t.load.split(";")[0])


def b_templates(book) -> list[Any]:
    """Every template of workbook B, in the order of the groups."""
    c = cfg()
    out = []
    drops = drop_templates_b()
    for tid in all_template_ids():
        if tid in drops:
            out.append(drops[tid])
        elif tid in c["uploads"]:
            out.append(_upload_template(tid, c["uploads"][tid], book))
        elif tid in c["screen_templates"]:
            out.append(_screen_template(tid, c["screen_templates"][tid], book))
        elif tid in dm_catalogue().layouts:
            out.append(_layout_template(tid, book))
        else:
            raise SystemExit(f"workbook B: unknown template {tid}")
    ids = {t.id for t in out}
    for t in out:
        t.depends = [d for d in t.depends if d in ids and d != t.id]
    return out


def b_target_text(item: dict[str, Any]) -> str:
    """Where a configuration item of the Drop 0 register is given in workbook B."""
    book = cfg()["meta"]["title_b"]
    tid = item.get("template") or ""
    obj = item.get("object") or ""
    if tid in cfg()["replaced"]:
        return f"Fill in the templates {', '.join(cfg()['replaced'][tid])} of the {book} workbook."
    if tid.startswith("LIST-"):
        sheets = [REVIEW[tid]] + ([REVIEW["LIST-SLA"]] if tid == "LIST-PARAMETERS" else [])
        return f"Give the values wanted on the sheet{'s' if len(sheets) > 1 else ''} {' and '.join(sheets)} of the {book} workbook."
    if tid and tid in all_template_ids():
        return f"Fill in the template {tid} of the {book} workbook."
    if item["route"] == "migration":
        layouts = [lc for lc in all_template_ids() if lc in dm_catalogue().layouts and
                   dm_catalogue().layouts[lc].object == obj]
        if obj == "R09":
            layouts = ["UP-02"]
        if layouts:
            return f"Fill in the template{'s' if len(layouts) > 1 else ''} {', '.join(layouts)} of the {book} workbook."
        return f"Loaded by the data migration (object {obj}, Migration Workbook of BRD-13)."
    if item["id"] in (cfg().get("covered_by") or {}):
        return f"Fill in the templates {', '.join(cfg()['covered_by'][item['id']])} of the {book} workbook."
    if item["route"] == "screen":
        return (f"Configured on the screen {drop_closure().menu()[item['screen']]} (sheet {SCREENS_SHEET} of the "
                f"{book} workbook).")
    return ""


# ============================================================================ workbook B: the writer

B_STATUSES = ["Not started", "In progress", "Ready for review", "Confirmed", "Not applicable"]


@lru_cache(maxsize=1)
def b_book() -> tuple[Any, list[Any], list[Any]]:
    """The guided book of workbook B with its lists and plans, the templates (group order) and the load order."""
    import guided_xlsx as g  # noqa: PLC0415

    dc = drop_closure()
    ci = dc.config_inputs()
    meta = cfg()["meta"]
    book = g.GuidedBook(meta["title_b"], meta["version"])
    insurer_codes = {r["code"] for r in dc.accounting_event_rows() if insurer_event(r)}
    for code, x in ci.resolve_lists(dc.combined()).items():
        values = [v for v in x["values"] if v[0] not in insurer_codes] if code == "ACCOUNTING_EVENT" else x["values"]
        book.add_list(code, x["name"], values, x["note"], x["strict"])
    templates = b_templates(book)
    by_id = {t.id: t for t in templates}
    for tid, cols in (cfg().get("list_overrides") or {}).items():
        for header, spec in cols.items():
            code = re.sub(r"[^A-Z0-9]+", "_", f"{tid}_{header}".upper()).strip("_")
            book.add_list(code, spec["name"], [tuple(v) for v in spec["values"]])
            col = next(c for c in by_id[tid].columns if c.header == header)
            col.check, col.allowed, col.allowed_link = f"list:{code}", "", ""
    for tid, notes in (cfg().get("template_notes") or {}).items():
        if tid in by_id:
            by_id[tid].notes = list(by_id[tid].notes) + list(notes)
    for t in templates:
        for c in t.columns:
            if c.example == "{baseCurrency}":
                c.example = "PHP"
    book.plan_all(templates)
    ci.finish_checks(book, templates)
    order = g.fill_in_order(templates)
    return book, templates, order


def _prepared_b(tid: str) -> str:
    book, _, _ = b_book()
    if tid in book.plans:
        return f"Template {tid} (sheet {book.plans[tid].sheet})"
    if tid in REVIEW:
        return "This sheet"
    return tid


def review_rows() -> dict[str, list[dict[str, Any]]]:
    dc = drop_closure()
    tl = dc._template_lists()
    lovs = []
    for r in dc.lov_rows():
        r = dict(r)
        if r["code"] in dc.migrated_lovs():
            r["route"] = ROUTE_MIGRATION
            r["prepared"] = f"Template R01 (legacy values mapped by code map {dc.migrated_lovs()[r['code']]}); new " \
                            f"values on Lists of Values"
        else:
            r["route"] = ROUTE_SCREEN
            r["prepared"] = _prepared_b(tl["lovs"].get(r["code"], "LIST-LOV"))
        lovs.append(r)
    params, sla, numbering = [], [], []
    for r in dc.parameter_rows():
        if r["key"] in excluded().get("parameters", set()):
            continue
        r = dict(r)
        r["route"] = ROUTE_SCREEN
        r["prepared"] = _prepared_b(r["template"])
        if r["key"].endswith("_PREFIX"):
            numbering.append(r)
        elif SLA.search(r["key"]) and r["group"] not in ("Security", "Data Migration"):
            sla.append(r)
        else:
            params.append(r)
    events = [dict(r, prepared="This sheet") for r in dc.event_rows()]
    docs = []
    for r in dc.document_rows():
        r = dict(r)
        r["prepared"] = _prepared_b(tl["documents"].get(r["code"], "LIST-DOCUMENTS"))
        docs.append(r)
    acc = [dict(r, prepared=_prepared_b("D0-08")) for r in dc.accounting_event_rows() if not insurer_event(r)]
    return {"LIST-LOV": lovs, "LIST-PARAMETERS": params, "LIST-SLA": sla, "LIST-NUMBERING": numbering,
            "LIST-NOTIFICATIONS": events, "LIST-DOCUMENTS": docs, "LIST-EVENTS": acc}


def insurer_event(r: dict[str, Any]) -> bool:
    x = excluded()
    return r["code"] in x.get("events", set()) or r["category"] in x.get("event_categories", set())


def review_columns() -> dict[str, tuple[str, list[Any]]]:
    from bdoi_xlsx import Column  # noqa: PLC0415

    def status_col():
        return Column("status", "Status", 14, "Status of the row", values=B_STATUSES)

    def comment_col():
        return Column("comments", "BDOI comments", 36, "Decision, value or comment of the owner")

    param_cols = [
        Column("key", "Parameter", 30, "Parameter on System Parameters"),
        Column("group", "Group", 16, "Group of the parameter"),
        Column("module", "Module", 20, "Module"),
        Column("owner", "Owner (BDOI unit)", 28, "Unit that decides the value"),
        Column("description", "What it controls", 60, "Meaning of the parameter"),
        Column("type", "Type", 12, "Kind of value"),
        Column("value", "Value proposed", 22, "Value proposed with the platform"),
        Column("range", "Allowed range", 12, "Minimum and maximum"),
        Column("prepared", "Given in", 30, "Where the value wanted is given"),
        Column("due", "Due", 22, "Due date"),
        Column("wanted", "Value wanted", 22, "Value BDOI wants"),
        Column("reason", "Reason", 30, "Policy reference or reason of a change"),
        Column("approver", "Approved by", 20, "BDOI approver of the value"),
        status_col(), comment_col(),
    ]
    return {
        "LIST-LOV": ("Every list of values of BIBS with its owner and its route: legacy values mapped through the "
                     "Migration Console (R01), the others entered on Lists of Values", [
            Column("code", "List", 26, "Code of the list on Lists of Values"),
            Column("name", "Name", 30, "Name of the list"),
            Column("module", "Module", 22, "Module that uses the list"),
            Column("owner", "Owner (BDOI unit)", 30, "Unit that owns the values"),
            Column("values", "Values proposed", 10, "Number of values proposed with the platform", kind="number"),
            Column("route", "Route", 16, "Migration Console or Configuration screen", values=[ROUTE_MIGRATION,
                                                                                            ROUTE_SCREEN]),
            Column("prepared", "Given in", 40, "Template where the values are given"),
            Column("due", "Due", 22, "Due date"),
            Column("keep", "Values to keep, add or retire", 40, "The owner's decision on the values"),
            status_col(), comment_col()]),
        "LIST-PARAMETERS": ("Every system parameter of BIBS (other than the SLA and turnaround parameters and the "
                            "numbering prefixes) with the value proposed and the value BDOI wants", param_cols),
        "LIST-SLA": ("The SLA, turnaround, reminder, escalation and ageing parameters: days and hours of the work "
                     "queues, insurer replies, reminders and alerts", param_cols),
        "LIST-NUMBERING": ("Prefixes of the numbered documents; receipt series come from R11, check numbers from "
                           "D0-12", [
            Column("key", "Parameter", 30, "Prefix parameter"),
            Column("module", "Module", 20, "Module of the document"),
            Column("description", "Numbered document", 60, "Document numbered with the prefix"),
            Column("value", "Prefix proposed", 16, "Prefix proposed with the platform"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that decides the prefix"),
            Column("wanted", "Prefix wanted", 16, "Prefix BDOI wants; the year and a running number follow it"),
            status_col(), comment_col()]),
        "LIST-NOTIFICATIONS": ("Every notification event with its default channels and the unit mailbox that also "
                               "receives it; users choose their own channels in Notification Settings", [
            Column("code", "Event", 28, "Notification event"),
            Column("name", "Name", 32, "Name shown in Notification Settings"),
            Column("module", "Module", 20, "Module"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that confirms the event"),
            Column("description", "Who is told and when", 50, "What the event tells the user"),
            Column("in_app", "In-app proposed", 10, "Y or N"),
            Column("email", "E-mail proposed", 10, "Y or N"),
            Column("in_app_wanted", "In-app wanted", 10, "Y or N", values=["Y", "N"]),
            Column("email_wanted", "E-mail wanted", 10, "Y or N", values=["Y", "N"]),
            Column("mailbox", "Unit mailbox (recipient)", 30, "Mailbox that also receives the e-mail (optional)"),
            status_col(), comment_col()]),
        "LIST-DOCUMENTS": ("Every generated document and letter with the unit that gives its BDOI wording and layout; "
                           "the layout is uploaded as a Word file on Document Templates", [
            Column("code", "Template", 26, "Template code on Document Templates"),
            Column("title", "Title", 34, "Title of the template"),
            Column("module", "Module", 20, "Module that generates the document"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that gives the wording and layout"),
            Column("fields", "Merge fields available", 50, "Fields BIBS fills in"),
            Column("prepared", "Given in", 30, "Where the text is given"),
            Column("due", "Due", 22, "Due date"),
            Column("layout", "BDOI layout attached", 12, "Y when the layout is sent", values=["Y", "N"]),
            status_col(), comment_col()]),
        "LIST-EVENTS": ("The accounting events that need a rule of template D0-08 (accounting rules mapping); owner "
                        "Head, Comptrollership", [
            Column("code", "Event type", 30, "Accounting event of the platform"),
            Column("name", "Name", 36, "Name of the event"),
            Column("category", "Category", 14, "Category"),
            Column("journal", "Journal type", 14, "Journal type of the entries"),
            Column("components", "Amount components", 50, "Amounts the rule lines can post"),
            Column("prepared", "Given in", 30, "Template of the rules"),
            status_col(), comment_col()]),
    }


def index_rows() -> list[dict[str, Any]]:
    book, templates, order = b_book()
    rows = []
    for i, t in enumerate(order, start=1):
        p = book.plans[t.id]
        brd = brd_of_template(t.id)
        rows.append({"Step": i, "ID": t.id, "Dataset": t.name, "Group": group_of(t.id), "BRD": brd,
                     "Drop": brand.drop_of(brd) if brd != "BRD-13" else "Drop 0",
                     "Provided by": t.owner, "Load route": route_label(t.id), "Entered or loaded on": where_of(t.id, t),
                     "Loads after": ", ".join(t.depends) or "-", "Due": t.due,
                     "Rows entered": p.rows_formula(), "Mandatory cells missing": p.missing_formula(),
                     "Status": "Not started", "Open": (f"{t.id} →", p.sheet)})
    return rows


def route_label(tid: str) -> str:
    """Route text of the index: the configuration-screen route with its note."""
    r = route_of(tid)
    return cfg()["route_screen"] if r == ROUTE_SCREEN else r


def screen_rows() -> list[dict[str, Any]]:
    dc = drop_closure()
    rows = []
    for it, r in zip(dc.data()["items"], dc.register_rows()):
        if (it["route"] != "screen" or it["id"] in excluded().get("config_items", set())
                or it["id"] in (cfg().get("covered_by") or {})):
            continue
        rows.append({"id": it["id"], "name": it["name"], "provides": r["provides"], "brd": module_tab(it["module"]),
                     "owner": r["owner"], "where": r["where"], "due": f"{r['due_date']} – {r['due']}",
                     "status": "Not started"})
    return rows


GAP_TEMPLATE = "No upload route: keyed on the screen from the template"
GAP_DATA = "Not held by BIBS: question to BDOI"


def gap_rows() -> list[dict[str, Any]]:
    book, templates, order = b_book()
    rows = []
    for t in order:
        if route_of(t.id) != ROUTE_SCREEN:
            continue
        p = book.plans[t.id]
        rows.append({"kind": GAP_TEMPLATE, "id": (t.id, f"{p.sheet}!A1"), "name": t.name,
                     "brd": brd_of_template(t.id), "owner": t.owner, "where": where_of(t.id, t),
                     "columns": len(t.columns), "mandatory": sum(1 for c in t.columns if c.mandatory.startswith("Y")),
                     "due": t.due, "how": cfg()["route_screen"]})
    return rows + data_gap_rows()


def data_gap_rows() -> list[dict[str, Any]]:
    """Data BDOI may need that BIBS does not hold (no field or no master): asked in the requirements workbook."""
    a = {r.id: r for r in consolidated()}
    rows = []
    seen = set()
    items = ([(x[0], x[1], x[5]) for x in cfg()["charges_checklist"] if x[4] == "No"]
             + [(x[0], x[1], x[5]) for x in cfg()["tax_checklist"] if x[4] == "No"]
             + [(x[0], x[3], x[4]) for x in cfg()["coverage"] if x[2] in ("No", "Partly") and x[4]])
    for name, why, ref in items:
        if name in seen:
            continue
        seen.add(name)
        q = a.get(ref)
        rows.append({"kind": GAP_DATA, "id": ref or "-", "name": name, "brd": q.tab if q else "-",
                     "owner": q.owner if q else "-", "where": "-", "columns": None, "mandatory": None,
                     "due": fmt(q.due) if q else "-",
                     "how": f"{why} Question {ref} of the BDOI Requirements and Inputs workbook." if ref else why})
    return rows


def checklist_rows(key: str) -> list[dict[str, Any]]:
    return [{"item": x[0], "rule": x[1], "where": x[2], "provides": x[3], "held": x[4], "question": x[5] or "-",
             "status": "Not started"} for x in cfg()[key]]


def coverage_rows() -> list[dict[str, Any]]:
    return [{"item": x[0], "where": x[1], "held": x[2], "note": x[3] or "-", "question": x[4] or "-"}
            for x in cfg()["coverage"]]


def prefill(book: Any, t: Any) -> None:
    """Writes the checklist rows of a template into its first input rows."""
    from openpyxl.styles import Font  # noqa: PLC0415

    rows = (cfg().get("prefill") or {}).get(t.id) or []
    if not rows:
        return
    p = book.plans[t.id]
    ws = book.wb[p.sheet[:31]]
    for i, values in enumerate(rows, start=1):
        for header, value in values.items():
            cell = ws[f"{p.letters[header]}{p.example_row + i}"]
            cell.value = value
            cell.font = Font(name=brand.FONT, size=10, color=brand.TEXT)


def build_b() -> Path:
    import guided_xlsx as g  # noqa: PLC0415
    from bdoi_xlsx import Column  # noqa: PLC0415

    meta = cfg()["meta"]
    book, templates, order = b_book()
    for i, t in enumerate(order, start=1):
        book.template_sheet(t, i, len(order))
        prefill(book, t)
    rcols = review_columns()
    rrows = review_rows()
    for key, sheet in REVIEW.items():
        desc, cols = rcols[key]
        book.table_sheet(sheet, f"Review list: {sheet.lower()}", desc, cols, rrows[key], B_STATUSES)
    book.table_sheet(SCREENS_SHEET, "Datasets configured on screens", "Set-up that BDOI decides and the BIBS "
                     "configuration team enters on the named screen in a working session; there is no template to "
                     "fill in, the decisions go into BDOI comments", [
        Column("id", "ID", 9, "Configuration item of the Drop 0 register"),
        Column("name", "Dataset", 30, "What is configured"),
        Column("provides", "What BDOI decides or provides", 50, "Content of the set-up"),
        Column("brd", "BRD", 9, "BRD"),
        Column("owner", "Owner (BDOI unit)", 26, "Unit that decides"),
        Column("where", "Configured on", 34, "Screen of BIBS"),
        Column("due", "Due", 26, "Due date"),
        Column("status", "Status", 14, "Status of the item", values=B_STATUSES),
        Column("comments", "BDOI comments", 36, "Decisions and values of the owner"),
    ], screen_rows(), B_STATUSES)
    book.table_sheet(GAPS_SHEET, "Gaps", "Datasets whose template rows the platform cannot take by file (keyed on "
                     "the screen from the filled-in template), and data BDOI may need that BIBS does not hold (asked "
                     "as a question in the BDOI Requirements and Inputs workbook)", [
        Column("kind", "Kind of gap", 22, "No upload route, or not held by BIBS", values=[GAP_TEMPLATE, GAP_DATA]),
        Column("id", "Template or question", 11, "Template of this workbook, or Ref of the question"),
        Column("name", "Dataset", 30, "Dataset"),
        Column("brd", "BRD", 9, "BRD"),
        Column("owner", "Provided by", 30, "Who fills in the template"),
        Column("where", "Keyed on", 34, "Screen where the rows are entered"),
        Column("columns", "Columns", 8, "Columns of the template", kind="number"),
        Column("mandatory", "Mandatory", 9, "Mandatory columns", kind="number"),
        Column("due", "Due", 24, "Due date of the template"),
        Column("how", "How the rows get into BIBS, or what is asked", 60, "Handling of the dataset"),
    ], gap_rows(), B_STATUSES)
    chk_cols = lambda what: [  # noqa: E731
        Column("item", what, 30, "Item"),
        Column("rule", "What BIBS does", 60, "The rule of the platform"),
        Column("where", "Where it is set up", 32, "Template, column or list"),
        Column("provides", "What BDOI provides or confirms", 44, "Input asked from BDOI"),
        Column("held", "Held by BIBS", 9, "Yes, or No (sheet Gaps)", values=["Yes", "No"]),
        Column("question", "Question", 10, "Ref of the question in the BDOI Requirements and Inputs workbook"),
        Column("confirm", "Rule confirmed", 14, "Agree, or Change (describe in BDOI comments)",
               values=["Agree", "Change"]),
        Column("comments", "BDOI comments", 36, "Decision, value or comment of the owner"),
        Column("status", "Status", 14, "Status of the item", values=B_STATUSES),
    ]
    book.table_sheet(CHARGES_SHEET, "PM-04 Charges checklist: premium charges billed to clients", "One row per "
                     "charge: how BIBS computes it, where its rate is given and what BDOI confirms (rates per line in "
                     "PM-04, LGT per insurer branch in R04B)", chk_cols("Charge"), checklist_rows("charges_checklist"),
                     B_STATUSES)
    book.table_sheet(TAX_SHEET, "D0-10 Tax checklist: BDOI's own taxes", "One row per kind of tax: what BIBS does, "
                     "where it is set up (D0-10 tax codes, TX-01 party tax profiles, TX-02 tax forms) and what BIBS "
                     "does not hold", chk_cols("Tax"), checklist_rows("tax_checklist"), B_STATUSES)
    book.table_sheet(COVERAGE_SHEET, "Coverage check", "The datasets the product owner asked about and where each "
                     "is in this workbook, or why BIBS does not hold it", [
        Column("item", "Dataset", 40, "Dataset asked about"),
        Column("where", "Where it is in this workbook", 46, "Template, list or sheet"),
        Column("held", "Held by BIBS", 10, "Yes, Partly or No", values=["Yes", "Partly", "No"]),
        Column("note", "Note", 60, "What BIBS keeps and what it does not"),
        Column("question", "Question", 10, "Ref of the question in the BDOI Requirements and Inputs workbook"),
    ], coverage_rows(), B_STATUSES)
    book.questions_sheet(order)
    idx = index_rows()
    columns = [("Step", 6), ("ID", 8), ("Dataset", 28), ("Group", 18), ("BRD", 8), ("Drop", 8), ("Provided by", 26),
               ("Load route", 30), ("Entered or loaded on", 28), ("Loads after", 13), ("Due", 20),
               ("Rows entered", 9), ("Mandatory cells missing", 10), ("Status", 14), ("Open", 9)]
    routes = Counter(r["Load route"] for r in idx)
    steps = [
        "The index below lists every master, reference and configuration dataset BIBS needs before go-live, one "
        "template sheet each, in the order the data is loaded: a dataset whose codes others use comes first (Loads "
        "after names the templates it depends on). Each template has the column layout BIBS loads.",
        f"Load route: {ROUTE_UPLOAD} ({routes[ROUTE_UPLOAD]} templates) - the filled-in sheet is uploaded on the named "
        f"screen, which checks every row; {ROUTE_MIGRATION} ({routes[ROUTE_MIGRATION]}) - the data is extracted from "
        f"the legacy systems by BDOI IT in this layout and loaded by the Data Migration Console; {ROUTE_SCREEN} "
        f"({routes[ROUTE_SCREEN]}) - upload on the screen being added; until then BDOI fills the template and it is "
        f"loaded on the screen with maker-checker (sheet {GAPS_SHEET}). Every record is authorised by a second user.",
        "On a template sheet, read the header block (purpose, who fills it in, how it is loaded, due, depends on) and "
        "the guide above each column (Mandatory, Format, Allowed values, What to enter; * marks a mandatory column). "
        "Overwrite or delete the grey example row and enter one row per record. Drop-downs offer the allowed values "
        "(sheet Reference lists) and the codes of the templates filled in before; a mandatory cell left empty turns "
        "red. For a large legacy dataset (clients, opening balances) BDOI IT sends the extract file in the layout "
        "instead of typing it.",
        "The review sheets (lists of values, system parameters, SLA and turnaround, numbering prefixes, notification "
        "events, document templates, accounting events) list what the platform delivers; write the value wanted or "
        f"the decision on each row you own. The sheet {SCREENS_SHEET} lists the set-up decided in working sessions "
        "and entered directly on the screens.",
        f"Set the Status of each template, follow Rows entered and Mandatory cells missing, raise questions on the "
        f"sheet {g.QUESTIONS}, and return the workbook to the {brand.VENDOR} project team by each due date. The owner "
        f"confirms each template (Confirmed by and Date on the sheet) before it is loaded; the guide "
        f"(BIBS_Templates_BRD-00_{meta['title_guide'].replace(' ', '_')}_v{meta['version']}) explains the "
        "validation and the sign-off of the master data.",
    ]
    identity = [("Client", brand.CLIENT), ("Scope", "Every master, reference and configuration dataset BIBS needs "
                                                     "before go-live, all BRDs"),
                ("Version and date", f"Version {meta['version']}, {meta['date']}"),
                ("Go-live", drop_closure().data()["meta"]["go_live_text"]),
                ("Owner", "Program Manager, Business Project Services, with the BDOI owner of each dataset"),
                ("Prepared by", brand.VENDOR), ("Classification", brand.CLASSIFICATION)]
    dc = drop_closure()

    def more(ws, row):
        row = book.small_table(ws, row, "Groups of datasets", ["Group", "Templates", "Upload", "Migration Console",
                                                               "Configuration screen"],
                               [[gr["name"], ", ".join(gr["templates"]),
                                 sum(1 for t in gr["templates"] if route_of(t) == ROUTE_UPLOAD),
                                 sum(1 for t in gr["templates"] if route_of(t) == ROUTE_MIGRATION),
                                 sum(1 for t in gr["templates"] if route_of(t) == ROUTE_SCREEN)]
                                for gr in cfg()["groups"]], spans=[3, 6, 2, 2, 2])
        row = book.small_table(ws, row, "Due milestones", ["Due", "Milestone", "Due date", "Before go-live",
                                                            "Why this date"],
                               [[c, x["label"], dc.fmt(x["date"]), f"T-{x['weeks']} weeks", x["note"]]
                                for c, x in dc.dues().items()], spans=[1, 4, 2, 2, 6])
        return book.small_table(ws, row, "Other sheets", ["Sheet", "What it holds"], [
            ((name, name), desc) for name, desc in
            [(REVIEW[k], rcols[k][0]) for k in REVIEW] +
            [(CHARGES_SHEET, "Premium charges billed to clients: one row per charge with the rule BIBS applies"),
             (TAX_SHEET, "BDOI's own taxes: one row per kind of tax, what BIBS does and what it does not hold"),
             (SCREENS_SHEET, "Set-up decided in working sessions and entered on the screens (no template)"),
             (GAPS_SHEET, "Datasets without an upload route, and data BIBS does not hold (questions)"),
             (COVERAGE_SHEET, "The datasets asked about and where each is, or why BIBS does not hold it"),
             (g.LISTS, "Every list of allowed values of the drop-downs, with code and label"),
             (g.QUESTIONS, "Questions and comments per template and column, with the answer")]
        ], spans=[3, 10])

    book.start_sheet(meta["title_b"], "Every master, reference and configuration dataset BIBS needs before go-live, "
                                      "with the templates in the order the data is loaded",
                     identity, steps, "Index of the datasets in load order (dependencies first)", columns, idx,
                     statuses=B_STATUSES, after=more)
    book.lists_sheet()
    tpl_sheets = []
    for t in order:
        tpl_sheets.append(book.plans[t.id].sheet)
        if t.id == "PM-04":
            tpl_sheets.append(CHARGES_SHEET)
        if t.id == "D0-10":
            tpl_sheets.append(TAX_SHEET)
    sheets = ([g.START] + tpl_sheets + list(REVIEW.values()) +
              [SCREENS_SHEET, GAPS_SHEET, COVERAGE_SHEET, g.LISTS, g.QUESTIONS])
    path = brand.out_path("BRD-00", KIND, brand.output_name("Templates", "BRD-00", meta["title_b"], meta["version"],
                                                              "xlsx"))
    return book.save(path, sheets, {"title": meta["title_b"],
                                    "keywords": "Master data templates, BRD-00, BIBS, BDOI"})


# ============================================================================ Word guide (```pack blocks)


def render(doc: Any, render: str, **_: Any) -> None:  # noqa: A002 - block key
    book, templates, order = b_book()
    if render == "groups":
        rows = []
        for gr in cfg()["groups"]:
            names = [f"{t} {next(x.name for x in templates if x.id == t)}" for t in gr["templates"]]
            rows.append([gr["name"], names])
        doc.table(["Group", "Templates"], rows, widths=[4.0, 13.6], caption="Datasets of the workbook by group",
                  size=8, keep_rows=False)
        return
    if render == "routes":
        routes = Counter(route_of(t.id) for t in templates)
        rows = [[ROUTE_UPLOAD, "The filled-in sheet (or a CSV or TXT file with the same header row) is uploaded on "
                               "the named screen. Every row is checked; the screen lists the rows refused with the "
                               "reason, and the accepted rows wait for a second user to authorise them.",
                 routes[ROUTE_UPLOAD]],
                [ROUTE_MIGRATION, "BDOI IT extracts the data from the legacy systems in the layout of the template, "
                                  "with its control file (row count and totals). The Data Migration Console checks "
                                  "the file, maps the legacy codes with the approved code maps and loads it in a "
                                  "batch that is reconciled and signed.", routes[ROUTE_MIGRATION]],
                [ROUTE_SCREEN, cfg()["route_screen"] + ". These datasets are listed on the sheet Gaps.",
                 routes[ROUTE_SCREEN]]]
        doc.table(["Route", "How the data gets into BIBS", "Templates"], rows, widths=[3.2, 12.2, 2.2],
                  caption="Load routes", size=8.5)
        return
    if render == "order":
        rows = [[r["Step"], r["ID"], r["Dataset"], r["Load route"], r["Loads after"], r["Due"]] for r in index_rows()]
        doc.table(["Step", "ID", "Dataset", "Route", "Loads after", "Due"], rows, widths=[0.9, 1.2, 5.6, 2.6, 2.5, 4.8],
                  caption="Load order of the datasets (dependencies first)", size=7.5, keep_rows=False)
        return
    if render == "gaps":
        rows = [[r["id"][0], r["name"], r["where"]] for r in gap_rows() if r["kind"] == GAP_TEMPLATE]
        doc.table(["Template", "Dataset", "Keyed on"], rows, widths=[1.6, 7.0, 9.0],
                  caption="Datasets without an upload route", size=8, keep_rows=False)
        return
    if render == "data_gaps":
        rows = [[r["id"], r["name"], r["how"]] for r in gap_rows() if r["kind"] == GAP_DATA]
        doc.table(["Question", "Data", "What BIBS keeps and what is asked"], rows, widths=[1.6, 5.0, 11.0],
                  caption="Data BIBS does not hold (questions to BDOI)", size=8, keep_rows=False)
        return
    if render == "charges":
        rows = [[x[0], x[1], x[2]] for x in cfg()["charges_checklist"]]
        doc.table(["Charge", "What BIBS does", "Where it is set up"], rows, widths=[3.6, 9.0, 5.0],
                  caption="Premium charges billed to clients", size=8, keep_rows=False)
        return
    if render == "screens":
        rows = [[r["id"], r["name"], r["where"]] for r in screen_rows()]
        doc.table(["ID", "Set-up", "Configured on"], rows, widths=[1.6, 7.0, 9.0],
                  caption="Set-up configured on screens in working sessions", size=8, keep_rows=False)
        return
    if render == "dues":
        dc = drop_closure()
        labels = {x["date"]: (c, x["label"]) for c, x in dc.dues().items()}
        for code, label, date in catalogue()["milestones"]:
            labels.setdefault(as_date(date), (code, label))
        per: dict[dt.date, list[str]] = defaultdict(list)
        for t in templates:
            d = due_of(t.due)
            if d:
                per[d].append(t.id)
        go = dc.go_live()
        rows = [[labels.get(d, ("-", ""))[0], labels.get(d, ("-", ""))[1], dc.fmt(d), f"T-{(go - d).days // 7} weeks",
                 ", ".join(ids)] for d, ids in sorted(per.items())]
        doc.table(["Due", "Milestone", "Date", "Before go-live", "Templates due"], rows,
                  widths=[1.0, 6.0, 2.2, 2.2, 6.2], caption="Due dates of the templates (go-live T = Monday 3 January "
                                                          "2028)", size=8, keep_rows=False)
        return
    raise ValueError(f"unknown render {render}")


def build_guide() -> Path:
    import bdoi_docx  # noqa: PLC0415

    meta = cfg()["meta"]
    out = brand.out_path("BRD-00", KIND, brand.output_name("Templates", "BRD-00", meta["title_guide"], meta["version"],
                                                             "docx"))
    results = bdoi_docx.build_markdown(GUIDE, out)
    docx, pdf = results[0]
    if pdf:
        Path(pdf).unlink(missing_ok=True)
    return docx


# ============================================================================ checks and main


def check() -> list[str]:
    problems: list[str] = []
    rows = consolidated()
    ids = Counter(r.id for r in rows)
    problems += [f"workbook A: row {i} twice" for i, n in ids.items() if n > 1]
    for r in rows:
        if r.tab not in {t["brd"] for t in tabs()}:
            problems.append(f"workbook A: {r.id} on an unknown tab {r.tab}")
        if r.type not in cfg()["types"]:
            problems.append(f"workbook A: {r.id} type {r.type}")
        if not r.due:
            problems.append(f"workbook A: {r.id} without a Needed by date")
        if not r.need:
            problems.append(f"workbook A: {r.id} without what is needed")
        if r.owner not in owner_roles():
            problems.append(f"workbook A: {r.id} owner {r.owner} is not a role of the Owners sheet")
    book, templates, order = b_book()
    known = set(all_template_ids())
    for tid in known:
        if tid not in book.plans:
            problems.append(f"workbook B: template {tid} not written")
    dc = drop_closure()
    for it in dc.data()["items"]:
        if it["route"] == "template" and it.get("template", "").startswith(("D0-", "PM-", "UA-")):
            if it["template"] not in known and it["template"] not in cfg()["replaced"]:
                problems.append(f"workbook B: configuration item {it['id']} names template {it['template']}, "
                                "not in the workbook")
    return problems


def scan(path: Path) -> list[str]:
    """Restricted words, development-status wording and technical terms in a written client file (check_pack)."""
    import check_pack as cp  # noqa: PLC0415

    allowed = cp.extract_field_names()
    hits: dict[str, set[str]] = defaultdict(set)
    for line in cp.office_lines(path):
        for m in cp.RESTRICTED.finditer(line):
            hits["restricted"].add(m.group(0))
        for label, pattern in cp.BUILD_STATUS + cp.TECHNICAL:
            for m in pattern.finditer(line):
                if label == "internal code" and m.group(0) in allowed:
                    continue
                hits[label].add(m.group(0).strip())
    return [f"{path.name}: {label}: {', '.join(sorted(w)[:12])}" for label, w in sorted(hits.items())]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build the BDOI input workbooks")
    ap.add_argument("--check", action="store_true", help="checks only")
    ap.add_argument("--a", action="store_true", help="workbook A only")
    ap.add_argument("--b", action="store_true", help="workbook B only")
    args = ap.parse_args(argv)
    problems = check()
    for p in problems:
        print(p)
    rows = consolidated()
    print(f"workbook A: {len(rows)} rows ({sum(len(r.attached) for r in rows)} merged); "
          + ", ".join(f"{t['brd']} {sum(1 for r in rows if r.tab == t['brd'])}" for t in tabs()))
    print("  types: " + ", ".join(f"{k} {v}" for k, v in Counter(r.type for r in rows).most_common()))
    print("  status: " + ", ".join(f"{k} {v}" for k, v in Counter(r.status for r in rows).most_common()))
    book, templates, order = b_book()
    print(f"workbook B: {len(templates)} templates; routes "
          + ", ".join(f"{k} {v}" for k, v in Counter(route_of(t.id) for t in templates).items())
          + f"; gaps {len(gap_rows())}; configured on screens {len(screen_rows())}")
    meta = cfg()["meta"]
    if args.check:
        import guided_xlsx  # noqa: PLC0415

        for path in (brand.out_path("BRD-00", KIND, brand.output_name("Inputs", "BRD-00", meta["title_a"],
                                                                       meta["version"], "xlsx")),
                     brand.out_path("BRD-00", KIND, brand.output_name("Templates", "BRD-00", meta["title_b"],
                                                                       meta["version"], "xlsx")),
                     brand.out_path("BRD-00", KIND, brand.output_name("Templates", "BRD-00", meta["title_guide"],
                                                                       meta["version"], "docx"))):
            if path.exists():
                found = scan(path)
                if path.suffix == ".xlsx" and "Templates" in path.name:
                    found += guided_xlsx.verify(path, [book.plans[t.id].sheet for t in templates])
                for p in found:
                    print(p)
                problems += found
        return 1 if problems else 0
    if problems:
        return 1
    built = []
    if not args.b:
        built.append(build_a())
    if not args.a:
        built.append(build_b())
        built.append(build_guide())
    for path in built:
        print(f"written: {path.relative_to(REPO)}")
        for p in scan(path):
            print(f"  {p}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
