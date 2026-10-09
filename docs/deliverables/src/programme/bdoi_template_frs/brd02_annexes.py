"""Appendices C to V, the summary for the Business Unit Review (section 1.4) and the change summary of the Operations
Cashiering FRS in BDOI's template, version 3.2. Used by build_brd02_bdoi_frs.py; the content is in brd02_document.yaml,
brd02_review.yaml and brd02_flows.yaml, and in the BIBS release set BRD-02 read at every build. The notice texts are
checked against the notification code of the platform at every build.
"""

from __future__ import annotations

import copy
import re
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
import yaml
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_brd02_bdoi_frs as m
import build_brd11_bdoi_frs as u11

REPO = m.REPO
OPS = m.OPS
DOC, REV, ADD, FLOWS = m.DOC, m.REV, m.ADD, m.FLOWS
JAVA = REPO / "backend" / "src" / "main" / "java" / "com" / "iortatechnxt" / "brokerverse"
PORTRAIT_W = 7.4
LANDSCAPE_W = 9.9
# The screens of Appendix E, in order (cashiering screens, then inquiry, reports and interfaces)
SCREENS = ["SCR-OP-07", "SCR-OP-08", "SCR-OP-09", "SCR-OP-10", "SCR-OP-11", "SCR-OP-12", "SCR-OP-13", "SCR-OP-14",
           "SCR-OP-15", "SCR-OP-16", "SCR-OP-18", "SCR-OP-20", "SCR-OP-21", "SCR-OP-22", "SCR-OP-02", "SCR-OP-03",
           "SCR-OP-05", "SCR-OP-63", "SCR-OP-64", "SCR-OP-61"]
# The platform files whose notices, alerts and e-mails belong to Cashiering (every one must be quoted in Appendix G)
NOTICE_SOURCES = ["cashiering/service/ReceiptActionService.java", "cashiering/service/DispositionService.java",
                  "cashiering/service/CashieringPaymentReversals.java",
                  "cashiering/service/CashieringDispositionRequests.java",
                  "cashiering/service/CashieringRefundValidationSource.java",
                  "cashiering/service/ReceiptSeriesService.java", "cashiering/service/PrebookedService.java",
                  "workflow/service/WorkflowService.java", "workflow/service/WorkAssignmentService.java",
                  "opsledger/service/HandoffService.java", "opsledger/service/DisbursementQueueService.java",
                  "opsledger/service/FlowInService.java", "messaging/service/JobFailureMailer.java"]
# Classes of the Cashiering package that only delegate a notice to a listed source
NOTICE_DELEGATES = {"cashiering/service/PaymentReversalSupport.java"}
# Places of the messages of the system (sign-off workbook of BRD-02, sheet Messages) that belong to Cashiering
MESSAGE_PLACES = re.compile(r"SCR-OP-(?:0[2357]|0[89]|1[0-8]|2[0-2]|6[134])\b|Every screen|Several screens|"
                            r"every upload type|Post-dated Checks$|Commission Payment Details$")
BRD_READS = {"Both": " (the BRD supports both readings)", "Our FRS": " (the BRD reads as the BIBS reference)",
             "BDOI's FRS": " (the BRD reads as BDOI's FRS)"}
WALK_STEPS = {"WT-A": [1, 2, 3, 4, 7], "WT-B": [8, 9, 10], "WT-D": [3, 4]}
PERSONAS = {"CASHIER": ("Cashier (Head Office)", "cashier"), "CASHIER_BR": ("Cashier (Cebu branch)", "cashbr"),
            "CASHIER_TL": ("Cashiering Team Leader / Head", "cashtl"), "PROCESSOR": ("Processing (New Business)", "proc"),
            "REMIT_PROCESSOR": ("Remittance Processor", "remit"), "REMIT_TL": ("Remittance Team Leader", "remittl"),
            "DISBURSEMENT": ("Disbursement", "disb"), "MKT_COLLECTION": ("Marketing Collection", "mktcoll"),
            "ADJUSTMENT": ("Adjustment Processor", "adjust"), "ADJUSTMENT_TL": ("Adjustment Team Leader", "adjtl"),
            "MKT_TL": ("Marketing Team Leader", "mkttl"), "AUDITOR": ("Auditor", "auditor"),
            "COMMREC_HANDLER": ("Commission Receivables Handler", "commrec"),
            "RECON_HANDLER": ("Production Reconciliation Handler", "recon")}


# ------------------------------------------------------------------------------------------- shared data
def screen_numbers() -> dict[str, str]:
    return {sid: f"E.{i}" for i, sid in enumerate(SCREENS, 1)}


def figure_numbers() -> dict[str, int]:
    order = ["summary", "C", "E", "J", "N"]
    figs = sorted(FLOWS, key=lambda f: order.index(f["where"]))
    return {f["id"]: i for i, f in enumerate(figs, 1)}


def screens() -> dict[str, dict]:
    out = {}
    for f in sorted((OPS / "pack" / "screens").glob("*.yaml")):
        for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
            out[s["id"]] = s
    return out


def signoff_sheet(name: str) -> list[tuple]:
    wb = openpyxl.load_workbook(m.SIGNOFF_XLSX, read_only=True)
    rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    return rows


def menu_rows(spec: dict) -> list[list]:
    seen = defaultdict(set)
    for r in signoff_sheet("Menu by persona")[3:]:
        if r[0]:
            seen[f"{r[1]} › {r[2]} › {r[3]}"].add(r[0])
    rows = []
    for screen in spec["menu_screens"]:
        if screen not in seen:
            raise SystemExit(f"menu entry {screen} not in the sign-off workbook")
        label = screen.split(" › ", 1)[1] if screen.startswith("Finance") else screen
        rows.append([label] + [p in seen[screen] for p in spec["menu_personas"]])
    return rows


def classify_items() -> dict[str, str]:
    oi = REV["open_items"]
    where: dict[str, str] = {}

    def put(ref, kind):
        if ref in where and where[ref] != kind:
            raise SystemExit(f"{ref} is classified twice ({where[ref]}, {kind})")
        where[ref] = kind
    for o in oi["open"]:
        for ref in re.findall(r"\b(?:OQ\d\d|IQ\d\d|CLR-OP-\d\d|CFG-\d{3})\b", o["refs"]):
            put(ref, "open")
    for ref in oi["covered_by_open"]:
        put(ref, "open")
    for ref in oi["inputs"]:
        put(ref, "input")
    for ref in oi["answered"]:
        put(ref, "answered")
    for a in oi["assumptions"]:
        put(a[0], "assumption")
    for d in oi["dependencies"]:
        put(d[0], "dependency")
    return where


def check_classification(where: dict[str, str]) -> int:
    needed = set(REV["open_items"]["scope_refs"])
    missing = sorted(needed - set(where))
    if missing:
        raise SystemExit(f"open items not classified: {missing}")
    clash = sorted(r for r, k in where.items() if k == "open" and r in REV["open_items"]["answered"])
    if clash:
        raise SystemExit(f"open questions answered in BDOI's FRS: {clash}")
    return len(needed)


def check_emails() -> int:
    items = REV["emails"]["items"]
    cat = yaml.safe_load((OPS / "pack" / "notifications.yaml").read_text(encoding="utf-8"))["notifications"]
    cash_nt = {n["id"] for n in cat if set(n.get("frs", [])) & set(m.SCOPE_FRS)}
    missing = sorted(cash_nt - {i["platform"] for i in items})
    if missing:
        raise SystemExit(f"Appendix G lacks the notices {missing}")
    code = {f: (JAVA / f).read_text(encoding="utf-8") for f in NOTICE_SOURCES}
    for f, text in code.items():
        if not any(frag in text for i in items for frag in i["code"]):
            raise SystemExit(f"no Appendix G notice quotes {f}")
    every = "\n".join(code.values())
    for i in items:
        for frag in i["code"]:
            if frag not in every:
                raise SystemExit(f"{i['id']}: text {frag!r} not found in the platform")
    for p in sorted((JAVA / "cashiering").rglob("*.java")):
        t = p.read_text(encoding="utf-8")
        if re.search(r"queueEmail\(|notifyUser\(|notifyPermission\(|alerts\.raise\(", t):
            rel = p.relative_to(JAVA).as_posix()
            if rel not in NOTICE_SOURCES and rel not in NOTICE_DELEGATES:
                raise SystemExit(f"notice source {rel} not quoted in Appendix G")
    return len(items)


def q_label(q_rows: list[dict], ref: str) -> str:
    for r in q_rows:
        if r["no"].endswith(f"({ref})"):
            return r["no"]
    return ref


def bdoi_messages(b) -> list[list[str]]:
    """The messages quoted in BDOI's FR table, with the item they belong to."""
    out, seen = [], set()
    t = b.tables[3]
    for row in t.rows[1:]:
        current = row.cells[0].text.strip()
        for el in row.cells[3]._tc.iter(qn("w:p")):
            text = u11.el_text(el).strip()
            idm = m.ID_RE.match(text)
            if idm:
                current = m.norm_id(idm.group(1))
            mm = re.search(r"[Mm]essage(?: prompt)?\s*:?\s*[“\"]?([^“”\"]{6,}?)[”\"]?\s*$", text)
            if not mm:
                mm = re.search(r"[Mm]essage(?: prompt)?\s+[“\"]([^”\"]+)[”\"]", text)
            if mm:
                msg = mm.group(1).strip().strip(".") + ("." if mm.group(1).strip().endswith(".") else "")
                if msg.lower() in seen or len(msg) < 6:
                    continue
                seen.add(msg.lower())
                out.append([current, msg])
    return out


# ------------------------------------------------------------------------------------------- the appendices
class Appendices:
    def __init__(self, b, figs, where):
        self.b = b
        self.figs = figs
        self.where = where
        self.els: list = []
        self.h1_ppr = m.find_par(b, "Appendix A", "Heading 1")._p.find(qn("w:pPr"))
        self.nums = b.fig_numbers
        self.scr = screens()
        self.scr_no = screen_numbers()
        self.cases = m.test_cases()
        self.trace = m.our_trace()
        self.our = m.our_frs()
        self.listed_refs: set[str] = set()
        self.listed_conflicts: set[str] = set()
        self.cmp_ids: set[str] = set()

    # ---------------------------------------------------------------- helpers
    def h1(self, text):
        p = OxmlElement("w:p")
        ppr = u11.strip_ids(copy.deepcopy(self.h1_ppr))
        for r in ppr.findall(qn("w:rPr")):
            ppr.remove(r)
        if ppr.find(qn("w:pageBreakBefore")) is None:
            u11.ppr_insert(ppr, OxmlElement("w:pageBreakBefore"))
        p.append(ppr)
        p.append(self.b.run(text))
        self.els.append(p)

    def h2(self, text):
        self.els.append(self.b.para(text, style="Heading2", no_num=True, keep_next=True))

    def p(self, text, **kw):
        self.els.append(self.b.para(text, **kw))

    def bullets(self, texts):
        n = self.b.new_num(self.b.BULLET_NUM)
        for t in texts:
            self.els.append(self.b.para(t, style="ListParagraph", num=(n, 0)))

    def table(self, header, rows, widths, size=8.5):
        self.els.append(self.b.table(header, rows, widths, size=size))
        self.els.append(self.b.para("", space_after=0))

    def kv(self, rows, widths=(1.6, 5.8), size=8.5):
        self.table(["Item", "Detail"], rows, list(widths), size=size)

    def figure(self, fid, max_h=8.4, max_w=7.0):
        spec = next(f for f in FLOWS if f["id"] == fid)
        self.els.append(self.b.picture(self.figs[fid], max_w=max_w, max_h=max_h))
        self.els.append(self.b.para(f"Figure {self.nums[fid]} – {spec['caption']}", italic=True, size=8.5, jc="center",
                                    space_after=40))
        self.els.append(self.b.para(f"How to read it: {spec['read']} Traces to: {spec['brd']}.", size=8.5,
                                    space_after=160))
        self.b.stats["figures_added"] += 1

    def items_of(self, frs) -> str:
        return ", ".join(sorted({x for fr in frs for x in self.b.fr_to_frum.get(fr, [])}, key=m.frs_key))

    def section(self, landscape: bool):
        self.els.append(u11.section_break(self.b, landscape=landscape))

    # ---------------------------------------------------------------- C
    def flows(self):
        self.h1("Appendix C – To-Be Process Flows")
        self.p(REV["annex_c"]["intro"])
        self.h2("C.1 Process flows")
        for f in FLOWS:
            if f["where"] == "C" and f["kind"] == "swimlane":
                self.figure(f["id"], max_h=8.6)
        self.h2("C.2 Status life-cycles")
        for f in FLOWS:
            if f["where"] == "C" and f["kind"] != "swimlane":
                self.figure(f["id"], max_h=6.0)

    # ---------------------------------------------------------------- D
    def files(self):
        d = DOC["annex_d"]
        self.h1("Appendix D – Payment and Upload File Layouts")
        self.p(d["intro"])
        for f in d["files"]:
            self.h2(f"{f['id']} {f['name']}")
            self.p("File name and format: " + f["convention"])
            self.table(f["header"], f["rows"], [0.55, 1.6, 2.15, 0.95, 2.15], size=8)
        self.b.stats["app_D_files"] = len(d["files"])

    # ---------------------------------------------------------------- E
    def screens(self):
        self.h1("Appendix E – Screen Specifications")
        self.p("This appendix specifies the screens of Operations Cashiering as the system shows them: the purpose of "
               "each screen, how the user reaches it, the FR items it serves, one screenshot, its fields with type, "
               "mandatory marker, format, list or source and validation, its actions and its rules. Mandatory: Y = "
               "mandatory, N = optional, Cond. = mandatory when the condition stated applies. The screenshots were "
               "taken on the SIT environment with SIT/UAT data. Where BDOI's FRS asks for a field, a filter or a step "
               "that the screen does not show today (for example the posting step, the bank account or the entry "
               "type), section 3 prevails and the point is in Appendix R.")
        for sid in SCREENS:
            s = self.scr[sid]
            k = self.scr_no[sid]
            self.h2(f"{k} {m.clean(s['title'])}")
            self.p("Purpose: " + m.clean(" ".join(str(s["purpose"]).split())))
            if s.get("entry"):
                self.p("Access: " + m.clean("; ".join(s["entry"])))
            items = self.items_of(s.get("frs", []))
            if items:
                self.p("FR items: " + items)
            shots = s.get("shots") or []
            if shots:
                n = int(sid.split("-")[-1])
                png = OPS / "screenshots" / f"scr-op-{n:02d}-01-{shots[0]['state']}.png"
                if not png.exists():
                    raise SystemExit(f"screenshot missing: {png}")
                self.els.append(self.b.picture(png, max_w=6.4, max_h=4.6))
                self.p(f"Screenshot {k}: {m.clean(shots[0].get('caption', s['title']))}", italic=True, size=8,
                       jc="center")
                self.b.stats["screenshots"] += 1
            rows = []
            for line in s.get("fields", []):
                c = [x.strip() for x in str(line).split(" | ")]
                c += ["-"] * (10 - len(c))
                label = c[1] if c[0] in ("-", "") else f"{c[0]}: {c[1]}"
                fmt = "; ".join(x for x in (c[3], c[5]) if x not in ("-", ""))
                val = "; ".join(x for x in (c[8], c[9]) if x not in ("-", ""))
                rows.append([m.clean(label), m.clean(c[2]), m.clean(c[4]), m.clean(fmt) or "-", m.clean(val) or "-"])
            if rows:
                self.table(["Field", "Type", "Mandatory", "Format / list / source", "Validation and message"],
                           rows, [1.5, 0.85, 0.9, 2.2, 1.95], size=7.5)
            acts = []
            for line in s.get("actions", []):
                c = [x.strip() for x in str(line).split(" | ")]
                c += ["-"] * (6 - len(c))
                acts.append([m.clean(c[0]), m.clean(c[1]), m.clean(c[2]), m.clean(c[3]),
                             m.clean(c[4]) if c[4] not in ("", "-") else "-"])
            if acts:
                self.table(["Action", "Who", "Available when", "What happens", "Resulting status"], acts,
                           [1.2, 1.3, 1.3, 2.6, 1.0], size=7.5)
            if s.get("rules"):
                self.p("Rules:", bold=True, keep_next=True)
                self.bullets([m.clean(r) for r in s["rules"]])
        self.b.stats["screens"] = len(SCREENS)
        self.h2(f"E.{len(SCREENS) + 1} Menu by persona")
        self.p("The Cashiering entries of the menu of each persona (the menu shows an entry only when one of the "
               "user's roles holds its permission).")
        self.figure("brd02_fig_menus", max_h=6.5)
        self.p("Screen standards: the layout, the formats of dates and amounts, the tables with row action menus and "
               "the common screen elements follow the BIBS programme screen standards; they are not repeated here.")

    # ---------------------------------------------------------------- F
    def messages(self):
        self.h1("Appendix F – Messages and Validations")
        self.p("F.1 lists the messages and confirmations of BDOI's FRS with the FR item that states them; F.2 lists the "
               "messages the system shows today on the Cashiering screens, word for word (the parts in angle brackets "
               "are filled in by the system). The wording of each message is confirmed by the Business Unit "
               "(requirements collection workbook, sheet Messages to confirm); where a message of the system differs "
               "from BDOI's, BDOI's wording prevails once confirmed.")
        self.h2("F.1 Messages of BDOI's FRS")
        rows = [[str(i), r[0], r[1]] for i, r in enumerate(bdoi_messages(self.b), 1)]
        self.table(["No.", "FR item", "Message"], rows, [0.4, 1.6, 5.4], size=8)
        self.b.stats["messages_bdoi"] = len(rows)
        self.h2("F.2 Messages of the system")
        rows = []
        for r in signoff_sheet("Messages")[4:]:
            if r[0] and r[1] and MESSAGE_PLACES.search(str(r[1])) and "SCR-OP-19" not in str(r[1]):
                rows.append([r[0].replace("MSG-OP-", "M-"), m.clean(r[1], keep_refs=True),
                             str(r[4]), str(r[5] or "-"), m.clean(r[6] or "-")])
        self.table(["No.", "Where it appears", "Message", "Type", "What the user does"], rows,
                   [0.55, 1.6, 2.6, 0.8, 1.85], size=7.5)
        self.b.stats["messages"] = len(rows)
        self.b.stats["messages_all"] = len(rows) + self.b.stats["messages_bdoi"]
        self.message_rows = rows

    # ---------------------------------------------------------------- G
    def emails(self):
        e = REV["emails"]
        self.h1("Appendix G – Notifications and E-mail Texts")
        self.p(e["intro"])
        rows = [[i["id"], i["name"], i["channel"], i["recipients"]] for i in e["items"]]
        self.table(["No.", "Notice", "Channel", "Recipients"], rows, [0.5, 1.9, 1.8, 3.2], size=7.5)
        for i in e["items"]:
            self.h2(f"{i['id']} {i['name']}")
            wording = ("Proposed by iorta TechNXT for BDOI approval; not given by the system today"
                       if i["platform"] == "Proposed" else
                       "Wording of the system today, proposed by iorta TechNXT for BDOI approval")
            self.kv([["Trigger", i["trigger"]], ["Recipients", i["recipients"]], ["Channel", i["channel"]],
                     ["Subject / title", i["subject"]], ["Text", i["body"]],
                     ["Attachments", e["common"]["attachments"]], ["Protection", e["common"]["protection"]],
                     ["Frequency", i["frequency"]],
                     ["Wording", f"{wording} (requirements collection workbook, sheet E-mail wordings, item {i['id']})"],
                     ["Traces to", f"{i['brd']}; {i['frs']}"]])
        self.b.stats["emails"] = len(e["items"])

    # ---------------------------------------------------------------- H
    def rules(self):
        h = DOC["annex_h"]
        self.h1("Appendix H – Business Rules, Roles and Configurable Parameters")
        self.p(h["intro"])
        self.h2("H.1 Personas")
        self.table(["Persona", "Responsibilities", "BRD"], h["personas"], [2.0, 4.0, 1.4], size=8)
        self.h2("H.2 Proposed role-to-action matrix")
        self.p("Y = the role holds the permission of the action. Proposal for BDOI to confirm with the access matrix "
               "of the User Access Maintenance FRS; the controls of H.3 apply whatever the roles grant.")
        self.table(["Action"] + h["matrix_roles"], h["matrix"], [2.1] + [0.76] * len(h["matrix_roles"]), size=7.5)
        self.h2("H.3 Controls")
        self.bullets(h["controls"])
        self.h2("H.4 Configurable parameters")
        self.table(["Parameter", "Value", "Meaning", "Maintained by"], h["parameters"], [1.8, 1.9, 2.3, 1.4], size=8)
        self.h2("H.5 Lists of values")
        self.table(["List", "Values (BDOI's FRS)"], h["lists"], [1.8, 5.6], size=8)
        self.h2("H.6 Scheduled runs")
        self.table(["Run", "When", "What it does"], h["jobs"], [2.0, 1.6, 3.8], size=8)
        self.h2("H.7 Statuses")
        self.table(["Record", "Statuses", "Source"], h["statuses"], [1.8, 3.9, 1.7], size=8)
        self.b.stats["app_H_parameters"] = len(h["parameters"])
        self.b.stats["app_H_lists"] = len(h["lists"])

    # ---------------------------------------------------------------- I
    def entries(self):
        i = DOC["annex_i"]
        self.h1("Appendix I – Accounting Entries")
        self.p(i["intro"])
        self.table(i["header"], i["rows"], [1.5, 1.5, 1.6, 1.8, 1.0], size=7.5)
        self.b.stats["app_I_entries"] = len(i["rows"])

    # ---------------------------------------------------------------- J
    def workflow(self):
        w = REV["workflow"]
        self.h1("Appendix J – Workflow and Approvals")
        self.p(w["intro"])
        self.figure("brd02_fig_approvals", max_h=6.0)
        self.table(w["header"], w["rows"], [1.3, 0.95, 1.35, 1.6, 1.25, 0.95], size=7.5)
        self.p(w["matrix_intro"], bold=True, keep_next=True)
        self.table(w["matrix_header"], w["matrix"], [1.8] + [0.8] * 7, size=7.5)
        self.p(w["escalation"])
        self.b.stats["workflow_rows"] = len(w["rows"])

    # ---------------------------------------------------------------- K
    def walkthroughs(self):
        data = yaml.safe_load((OPS / "pack" / "walkthroughs.yaml").read_text(encoding="utf-8"))["walkthroughs"]
        self.h1("Appendix K – Process Walkthroughs")
        self.p("The Cashiering steps of the end-to-end walkthroughs of the BIBS release set BRD-02, step by step: who "
               "acts, on which screen, what the user does, what the user sees and the result (the step numbers are "
               "those of the walkthroughs), then the walkthrough of this FRS for the Business Unit (WT-E). The "
               "screens are specified in Appendix E; the sign-in IDs are in the Business Unit Walkthrough Users "
               "workbook.")
        total = 0
        self.walk_rows = []
        for w in data:
            if w["id"] not in WALK_STEPS:
                continue
            self.h2(f"{w['id']} {m.clean(w['title'])}")
            self.p(m.clean(" ".join(str(w["summary"]).split())))
            rows = []
            for i, s in enumerate(w["steps"], 1):
                if i not in WALK_STEPS[w["id"]]:
                    continue
                rows.append(self.walk_row(f"{w['id']}.{i}", s[0], s[1], s[2], s[3], s[4]))
            total += len(rows)
            self.table(["Step", "Persona", "Screen", "What the user does", "What the user sees", "Result"], rows,
                       [0.6, 1.0, 1.2, 1.9, 1.75, 0.95], size=7.5)
        wt = REV["walkthrough"]
        self.h2(f"{wt['id']} {wt['title']}")
        self.p(wt["summary"])
        self.p(wt["data"], italic=True, size=9)
        rows = [self.walk_row(f"{wt['id']}.{i}", *s) for i, s in enumerate(wt["steps"], 1)]
        total += len(rows)
        self.table(["Step", "Persona", "Screen", "What the user does", "What the user sees", "Result"], rows,
                   [0.6, 1.0, 1.2, 1.9, 1.75, 0.95], size=7.5)
        self.b.stats["walk_steps"] = total

    def walk_row(self, step, persona, sid, does, sees, result):
        name, user = PERSONAS.get(persona, (m.clean(persona), "-"))
        screen = m.clean(self.scr[sid]["title"]) if sid in self.scr else "New Business screen"
        where = f"{screen} ({self.scr_no[sid]})" if sid in self.scr_no else screen
        self.walk_rows.append({"step": step, "persona": name, "user": user, "screen": where, "does": m.clean(does),
                               "sees": m.clean(sees), "result": m.clean(result)})
        return [step, name, where, m.clean(does), m.clean(sees), m.clean(result)]

    # ---------------------------------------------------------------- L
    def documents(self):
        d = REV["documents"]
        self.h1("Appendix L – Document Prints and Output Formats")
        self.p(d["intro"])
        self.p(d["file_name_rule"])
        for doc in d["items"]:
            self.h2(f"{doc['id']} {doc['name']}")
            fields = "\n".join(f"{n}. {f}" for n, f in enumerate(doc["fields"], 1))
            layout = ("Current layout of the system below; the Business Unit confirms BDOI's form or gives its layout: "
                      f"requirements collection workbook, item {doc['rc']}" if doc["shot"] else
                      f"Layout to be provided or confirmed by the Business Unit: requirements collection workbook, "
                      f"item {doc['rc']}")
            self.kv([["Purpose", doc["purpose"]], ["Recipients", doc["recipients"]], ["Format", doc["format"]],
                     ["Protection", doc["protection"]], ["Fields in order", fields], ["BDOI reference", doc["bdoi"]],
                     ["Layout", layout]], size=8)
            if doc["shot"]:
                png = OPS / "screenshots" / f"{doc['shot']}.png"
                if not png.exists():
                    raise SystemExit(f"image missing: {png}")
                self.els.append(self.b.picture(png, max_w=6.4, max_h=4.6))
                self.p(f"Image {doc['id']}: {doc['shot_caption']}", italic=True, size=8, jc="center")
        self.b.stats["documents"] = len(d["items"])

    # ---------------------------------------------------------------- M
    def reports(self):
        r = REV["reports"]
        self.h1("Appendix M – Reports and Schedules")
        self.p(r["intro"])
        self.table(r["header"], r["rows"], [0.95, 1.35, 1.05, 1.2, 1.35, 2.5, 0.7, 0.7], size=7)
        self.p(r["schedules_intro"], bold=True, keep_next=True)
        self.table(r["schedules_header"], r["schedules"], [1.9, 1.6, 2.2, 1.6, 1.4, 1.1], size=7.5)
        wanted = {f for f in self.b.bdoi_items if f.startswith("FRS.CSH.09.02.")}
        wanted |= {ns["id"] for ns in ADD["new_subitems"] if ns["row"] == "FRS.CSH.09.02"}
        missing = sorted(wanted - {row[0] for row in r["rows"]})
        if missing:
            raise SystemExit(f"Appendix M lacks the reports {missing}")
        self.b.stats["reports"] = len(r["rows"])
        self.b.stats["schedules"] = len(r["schedules"])

    # ---------------------------------------------------------------- N
    def integrations(self):
        w = REV["integrations"]
        self.h1("Appendix N – Integrations")
        self.p(w["intro"] + f" (Figure {self.nums['brd02_fig_context']}).")
        self.figure("brd02_fig_context", max_h=5.0)
        self.table(w["header"], w["rows"], [1.25, 0.55, 1.6, 0.85, 0.7, 0.85, 0.95, 0.75], size=7)
        self.b.stats["integrations"] = len(w["rows"])

    # ---------------------------------------------------------------- O
    def nfr(self):
        x = REV["nfr"]
        self.h1("Appendix O – Non-functional Requirements")
        self.p(x["intro"])
        self.table(x["header"], x["rows"], [0.8, 0.85, 1.7, 2.0, 1.05, 1.0], size=7.5)
        self.b.stats["nfr_rows"] = len(x["rows"])

    # ---------------------------------------------------------------- P
    def data_setup(self):
        y = REV["data_setup"]
        self.h1("Appendix P – Data Set-up and Migration at Go-live")
        self.p(y["intro"])
        self.table(y["header"], y["rows"], [1.5, 1.0, 1.1, 1.0, 0.95, 0.85, 1.0], size=7.5)
        self.p(y["not_loaded"])
        self.b.stats["data_rows"] = len(y["rows"])

    # ---------------------------------------------------------------- Q
    def traceability(self):
        self.h1("Appendix Q – Requirements Traceability")
        self.p("Q.1 traces every BRD ID of the Cashiering scope of the Operations BRD (v1.01 with the Cashiering annex "
               "of May 2026) to the FR items of this FRS, to the reference requirements of the BIBS FRS BRD-2 "
               "Operations v2.1 (FR-OP) and to the test conditions of the BIBS test plan BRD-2 (TC-OP-nnn.n; the cases "
               "of each condition are in the Test Cases and Traceability workbook). Q.2 lists every reference "
               "requirement with the FR items that incorporate it.")
        self.h2("Q.1 BRD requirement to FR items, BIBS reference and test conditions")
        rows = []
        for rid in self.b.ids:
            items = ", ".join(sorted(set(m.FRS_RE.findall(self.b.mapping[rid])), key=m.frs_key))
            tr = self.trace.get(rid, {"page": "-", "frs": []})
            refs = tr["frs"]
            if rid in m.REMOVED:
                items = items or "Removed from the BRD"
            rows.append([rid, tr["page"], items or "-", ", ".join(refs) or "-",
                         "; ".join(m.tc_text(r, self.cases) for r in refs) or "-"])
        self.table(["BRD ID", "BRD page", "FR items (this FRS)", "BIBS reference (FR-OP)", "Test conditions"],
                   rows, [1.0, 1.1, 4.0, 1.4, 2.4], size=7)
        self.b.stats["trace_rows"] = len(rows)
        self.h2("Q.2 BIBS reference requirement to FR items")
        rows = []
        for fr_id, items in sorted(self.b.fr_to_frum.items()):
            fr = self.our.get(fr_id)
            if fr is None:
                raise SystemExit(f"unknown BIBS reference {fr_id}")
            brd = ", ".join(x.split(" (")[0] for x in fr["brd"] if isinstance(x, str))
            rows.append([fr_id, m.clean(fr["title"]), brd, ", ".join(items), m.tc_text(fr_id, self.cases)])
            self.listed_refs.add(fr_id)
        self.table(["Reference", "Title", "BRD", "FR items (this FRS)", "Test conditions"], rows,
                   [0.85, 2.3, 1.5, 3.4, 1.85], size=7)

    # ---------------------------------------------------------------- R
    def observations(self):
        cmp_data = yaml.safe_load(m.CMP.read_text(encoding="utf-8"))
        self.h1("Appendix R – Observations and Points for BDOI Decision")
        self.p("Every point where BDOI's FRS and the BIBS reference FRS differ, the slips noticed in BDOI's text, the "
               "changes made to the Business Requirements Mapping and the other observations of this version. BDOI's "
               "text is kept in the body of the document; this appendix gives both readings, the BRD text, what "
               "version 3.2 proposes, the impact and who decides. Every point starts as Open.")
        header = ["No.", "Topic", "BDOI FRS text", "BRD text", "Proposed in v3.2", "Impact", "Decision by", "Status"]
        widths = [0.75, 1.0, 1.9, 1.5, 2.2, 0.95, 1.0, 0.6]
        no = 0
        section = "R.1"
        q_rows = self.b.q_rows

        def row(ref, topic, bdoi, brd, prop, impact, who):
            nonlocal no
            no += 1
            out = [f"R-{no:02d}" + (f" ({ref})" if ref else ""), topic, bdoi, brd, prop, impact, who, "Open"]
            q_rows.append(dict(zip(["no", "topic", "bdoi", "brd", "proposal", "impact", "decide_by", "status"], out),
                               section=section))
            return out
        decide = {"Between the documents": "Cashiering Process Owner; System Analyst (BDO ITG)",
                  "Within BDOI's FRS": "System Analyst (BDO ITG)"}
        self.h2("R.1 Conflicts between BDOI's FRS and the BIBS reference FRS")
        rows = []
        for c in cmp_data["conflicts"]:
            if c["kind"] != "Between the documents":
                continue
            prop = m.clean(f"{c['recommendation']} (BIBS reference FRS: {c['ours']})", keep_refs=True)
            brd = m.clean(c["brd_note"], keep_refs=True) + BRD_READS.get(c["brd"], "")
            rows.append(row(c["id"], m.clean(c["topic"], keep_refs=True), m.clean(c["bdoi"], keep_refs=True), brd,
                            prop, m.clean(c["decision"], keep_refs=True), decide[c["kind"]]))
            self.listed_conflicts.add(c["id"])
        self.table(header, rows, widths, size=7)
        section = "R.2"
        self.h2("R.2 Slips noticed in BDOI's FRS (proposed corrections; BDOI's text is not changed)")
        rows = []
        for c in cmp_data["conflicts"]:
            if c["kind"] == "Between the documents":
                continue
            rows.append(row(c["id"], m.clean(c["topic"], keep_refs=True), m.clean(c["bdoi"], keep_refs=True),
                            m.clean(c["brd_note"], keep_refs=True), m.clean(c["recommendation"], keep_refs=True),
                            m.clean(c["decision"], keep_refs=True), decide[c["kind"]]))
            self.listed_conflicts.add(c["id"])
        for o in DOC["slips"]:
            rows.append(row("", o["topic"], o["bdoi"], o["brd"], o["proposal"], o["impact"], o["decide_by"]))
        self.table(header, rows, widths, size=7)
        section = "R.3"
        self.h2("R.3 Changes to the Business Requirements Mapping")
        rows = []
        for ch in self.b.log:
            rows.append(row(ch["id"], "Mapping extended", f"Mapped to: {ch['before']}", ch["brd"],
                            f"{ch['after']}. {ch['why']}", "Coverage of the BRD requirement.",
                            "System Analyst (BDO ITG)"))
        rows.append(row("", "BRD IDs added", "Mapping without CSHID.009, CSHID.026 and CSHID.027",
                        "CSHID.009 removed by the Cashiering annex; CSHID.026 and CSHID.027 deleted in BRD v1.01",
                        "Three rows added with the BRD text and the reason of the removal, so that every BRD ID of the "
                        "Cashiering scope is accounted for.", "Coverage of every BRD ID.", "System Analyst (BDO ITG)"))
        self.table(header, rows, widths, size=7)
        section = "R.4"
        self.h2("R.4 Other observations and proposals of version 3.2")
        rows = [row("", o["topic"], o["bdoi"], o["brd"], o["proposal"], o["impact"], o["decide_by"])
                for o in DOC["observations"]]
        self.table(header, rows, widths, size=7)
        self.b.stats["observations"] = no
        self.b.stats["conflicts"] = sum(1 for c in cmp_data["conflicts"] if c["kind"] == "Between the documents")
        self.b.stats["slips"] = sum(1 for r in q_rows if r["section"] == "R.2")
        self.cmp_ids = {c["id"] for c in cmp_data["conflicts"]}

    # ---------------------------------------------------------------- S
    def open_items(self):
        o = REV["open_items"]
        self.h1("Appendix S – Assumptions, Dependencies and Open Questions")
        self.p(o["intro"])
        self.h2("S.1 Assumptions")
        self.table(["No.", "Our reference", "Assumption", "Related"],
                   [[f"S.1.{i}", r[0], r[1], r[2]] for i, r in enumerate(o["assumptions"], 1)],
                   [0.5, 0.9, 4.4, 1.6], size=8)
        self.h2("S.2 Dependencies")
        self.table(["No.", "Our reference", "Dependency", "Needed for"],
                   [[f"S.2.{i}", r[0], r[1], r[2]] for i, r in enumerate(o["dependencies"], 1)],
                   [0.5, 0.9, 4.4, 1.6], size=8)
        self.h2("S.3 Open questions not answered in BDOI's FRS")
        rows = []
        for i, q in enumerate(o["open"], 1):
            rows.append([f"S.3.{i}", q["refs"], q["q"], q["why"], q["fmt"], q["owner"],
                         q["by"].strftime("%d-%b-%Y") if hasattr(q["by"], "strftime") else str(q["by"])])
        self.table(o["open_header"], rows, [0.45, 0.85, 1.85, 1.45, 1.35, 0.8, 0.65], size=7)
        n_answered = sum(1 for k in self.where.values() if k == "answered")
        self.p(f"{n_answered} earlier questions, clarifications and assumptions are answered by BDOI's FRS and are not "
               f"repeated here; {sum(1 for k in self.where.values() if k == 'input')} items are values, data or "
               f"signatures the Business Unit provides (requirements collection workbook).", italic=True, size=9)
        self.b.stats["open_kept"] = len(o["open"])
        self.b.stats["answered"] = n_answered
        self.b.stats["assumptions"] = len(o["assumptions"])
        self.b.stats["dependencies"] = len(o["dependencies"])

    # ---------------------------------------------------------------- T, U, V
    def change_control(self):
        c = REV["change_control"]
        self.h1("Appendix T – Change Control after Sign-off")
        self.p(c["intro"])
        self.p("What signing freezes", bold=True, keep_next=True)
        self.bullets(c["freezes"])
        self.p(c["not_frozen"])
        self.p("How a change is raised and decided", bold=True, keep_next=True)
        self.table(c["steps_header"], c["steps"], [0.8, 3.0, 1.9, 1.7], size=8)
        self.p(c["candidates"])

    def checklist(self):
        c = REV["checklist"]
        self.h1("Appendix U – Business Unit Review Checklist")
        self.p(c["intro"])
        rows = [[str(i), r[0], r[1], r[2], r[3], "☐"] for i, r in enumerate(c["rows"], 1)]
        self.table(c["header"], rows, [0.4, 1.4, 1.5, 2.3, 1.3, 0.5], size=8)
        self.b.stats["checklist"] = len(rows)

    def glossary(self):
        self.h1("Appendix V – Glossary")
        self.p("Terms and abbreviations used in this document.")
        terms = sorted(DOC["glossary"] + REV["glossary_add"], key=lambda g: g[0].lower())
        self.table(["Term", "Meaning"], terms, [1.9, 5.5], size=9)
        self.b.stats["glossary"] = len(terms)

    # ---------------------------------------------------------------- order
    def write_all(self):
        self.flows()
        self.files()
        self.screens()
        self.messages()
        self.emails()
        self.rules()
        self.entries()
        self.workflow()
        self.walkthroughs()
        self.documents()
        self.section(landscape=False)
        self.reports()
        self.section(landscape=True)
        self.integrations()
        self.nfr()
        self.data_setup()
        self.section(landscape=False)
        self.traceability()
        self.observations()
        self.section(landscape=True)
        self.open_items()
        self.change_control()
        self.checklist()
        self.glossary()

    def place(self, anchor):
        self.b.insert_after(anchor, self.els)


# ------------------------------------------------------------------------------------------- 1.4 summary
def key_numbers(b) -> list[list[str]]:
    s = b.stats
    items = s["bdoi_items"] + s["new_subitems"] + s["new_item_subitems"]
    cases = m.test_cases()
    n_cases = sum(n for fr, (c, n) in cases.items() if fr in m.SCOPE_FRS)
    return [
        ["BRD IDs of the Cashiering scope in the mapping (removed and by reference included)", str(s["brd_ids"]),
         "Section 2; Appendix Q"],
        ["FR entries (BDOI's 18 and the new FRS.CSH.05.02 to FRS.CSH.10.04)", str(s["bdoi_entries"] + s["new_items"]),
         "Section 3"],
        ["FR items and sub-items (BDOI's and added)", str(items), "Section 3"],
        ["BDOI items elaborated / new sub-items and items", f"{s['elaborated']} / {s['new_subitems'] + s['new_item_subitems']}",
         "Section 3"],
        ["Acceptance criteria", str(s["ac"]), "Section 3"],
        ["Figures added (flows, life-cycles, menus, approvals, context)", str(len(FLOWS)), "1.4; Appendices C, E, J, N"],
        ["Screens", str(s["screens"]), "Appendix E"],
        ["Messages (BDOI's FRS / system)", f"{s['messages_bdoi']} / {s['messages']}", "Appendix F"],
        ["Notices, alerts and e-mails", str(s["emails"]), "Appendix G"],
        ["Documents and outputs", str(s["documents"]), "Appendix L"],
        ["Reports / scheduled runs", f"{s['reports']} / {s['schedules']}", "Appendix M"],
        ["Integrations", str(s["integrations"]), "Appendix N"],
        ["Test cases (BIBS test plan and one per acceptance criterion)", str(n_cases + s["ac"]),
         "Test Cases and Traceability workbook"],
        ["Open questions not answered in BDOI's FRS / answered and not repeated", f"{s['open_kept']} / {s['answered']}",
         "Appendix S"],
        ["Observations and points for decision", str(s["observations"]), "Appendix R"],
        ["Review checklist points", str(s["checklist"]), "Appendix U"],
    ]


def add_summary(b, figs):
    s = REV["summary"]
    model_h2 = m.find_par(b, "Document Preparation, Review and Approval", "Heading 2")._p
    model_body = m.find_par(b, "This Functional Requirements Specifications (FR) document outlines")._p

    def h2(text):
        h = u11.strip_ids(copy.deepcopy(model_h2))
        for el in list(h):
            if el.tag != qn("w:pPr"):
                h.remove(el)
        ppr = h.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is None:
            u11.ppr_insert(ppr, OxmlElement("w:pageBreakBefore"))
        h.append(b.run(text))
        return h

    def body(text):
        return b.clone_par(model_body, text)

    def sub(text):
        return b.para(text, bold=True, keep_next=True, size=10.5, color="014EA9", space_after=60)
    nums = b.fig_numbers
    spec = next(f for f in FLOWS if f["id"] == "brd02_fig_e2e")
    fig = [b.picture(figs["brd02_fig_e2e"], max_w=6.9, max_h=7.6),
           b.para(f"Figure {nums['brd02_fig_e2e']} – {spec['caption']}", italic=True, size=8.5, jc="center",
                  space_after=40),
           b.para(f"How to read it: {spec['read']} Traces to: {spec['brd']}.", size=8.5, space_after=160)]
    b.stats["figures_added"] += 1
    els = [h2("Summary for the Business Unit Review"), body(s["intro"])]
    els += [sub("Scope on a page"), b.table(s["scope_header"], s["scope"], [2.6, 2.3, 2.5], size=8), b.para("")]
    els += [sub("The process from end to end")] + fig
    els += [sub("Key numbers"), body(s["numbers_intro"]),
            b.table(["Item", "Number", "Where"], b.key_numbers, [3.9, 1.0, 2.5], size=8.5), b.para("")]
    rows = []
    for no, text, refs, who in s["decisions"]:
        where = "; ".join(q_label(b.q_rows, r.strip()) for r in refs.split(","))
        rows.append([str(no), text, where, who])
    els += [sub("Decisions the Business Unit must take"), body(s["decisions_intro"]),
            b.table(["No.", "Decision", "Appendix R", "Decided by"], rows, [0.4, 3.9, 1.3, 1.8], size=8), b.para("")]
    els += [sub("How to review this document"), body(s["review_intro"]),
            b.table(s["review_header"], s["review"], [1.8, 2.5, 2.5, 0.6], size=8), b.para("")]
    els += [body(s["workbooks_intro"]),
            b.table(["Workbook", "What it holds"], s["workbooks"], [2.2, 5.2], size=8.5), b.para("")]
    els.append(body(f"Before signing, the Business Unit confirms the {len(REV['checklist']['rows'])} points of the "
                    f"review checklist in Appendix U."))
    anchor = m.find_par(b, "Business Requirements Mapping", "Heading 1")._p
    for el in els:
        anchor.addprevious(el)


CHANGE_POINTS = [
    "Cover: drafting support line; Document Revision Log: row 3.2; Signoff Sheet: iorta TechNXT Project Team (drafting support) under Prepared by.",
    "Introduction: the drafting of this version, the scope of the BRD of 8-Oct-2026 and the new section 1.4 Summary for the Business Unit Review (scope on a page, the process from end to end, key numbers, the ten decisions, how to review).",
    "Business Requirements Mapping: {mapping} rows given FR items (the eight blank rows filled), {removed} BRD IDs removed from the BRD added as rows; List of Integrations completed.",
    "Functional Requirements: {elaborated} of BDOI's items elaborated, {new_sub} new sub-items continuing BDOI's numbering, {new_items} new entries FRS.CSH.05.02 to FRS.CSH.10.04 ({new_item_sub} sub-items); {statements} statements and {ac} numbered acceptance criteria; column Application Feature Model filled for every entry.",
    "Appendix B: references of version 3.2. New Appendices C To-Be Process Flows ({flows} figures); D Payment and Upload File Layouts ({files} files); E Screen Specifications ({screens} screens and the menu by persona); F Messages ({messages}); G Notifications and E-mail Texts ({emails}); H Business Rules, Roles and Parameters; I Accounting Entries ({entries} events); J Workflow and Approvals; K Process Walkthroughs ({walk} steps); L Document Prints ({documents}); M Reports and Schedules ({reports} reports, {schedules} runs); N Integrations ({integrations}); O Non-functional Requirements ({nfr}); P Data Set-up ({data}); Q Traceability ({trace} BRD IDs); R Observations ({obs}); S Open Questions ({open_kept} kept, {answered} answered by BDOI's FRS and not repeated); T Change Control; U Review Checklist ({checklist} points); V Glossary ({glossary} terms).",
    "BDOI's text of version 3.1 is unchanged: every paragraph, table row and FR ID of BDOI is still in place; the slips noticed are proposed in Appendix R only.",
]


def change_summary(b):
    s = b.stats
    values = dict(mapping=s["mapping_changes"], removed=s["removed_rows"], elaborated=s["elaborated"],
                  new_sub=s["new_subitems"], new_items=s["new_items"], new_item_sub=s["new_item_subitems"],
                  statements=s["requirements_added"], ac=s["ac"],
                  flows=sum(1 for f in FLOWS if f["where"] == "C"), files=s["app_D_files"], screens=s["screens"],
                  messages=s["messages_all"], emails=s["emails"], entries=s["app_I_entries"], walk=s["walk_steps"],
                  documents=s["documents"], reports=s["reports"], schedules=s["schedules"],
                  integrations=s["integrations"], nfr=s["nfr_rows"], data=s["data_rows"], trace=s["trace_rows"],
                  obs=s["observations"], open_kept=s["open_kept"], answered=s["answered"], checklist=s["checklist"],
                  glossary=s["glossary"])
    els = [b.para("Summary of changes in version 3.2", bold=True, size=14, color="014EA9", hl=False,
                  space_after=120),
           b.para("This review copy shows in light yellow every passage added in version 3.2: the full coverage of the "
                  "Cashiering scope of the BRD with acceptance criteria and the Business Unit review edition. BDOI's "
                  "text of version 3.1 is not shaded; the points for decision are in Appendix R and the open questions "
                  "in Appendix S.", size=9.5, hl=False)]
    n = b.new_num(b.BULLET_NUM)
    for pt in CHANGE_POINTS:
        els.append(b.para(pt.format(**values), style="ListParagraph", num=(n, 0), size=9.5, hl=False))
    els.append(b.para("The first decision BDOI must take is whether every AR and OR is posted by a second user before "
                      "its number is issued (Appendix R, C1); section 1.4 lists the ten decisions.", size=9.5, hl=False))
    els.append(b.para("Shading used in this copy: ", size=9.5, hl=False))
    els[-1].append(b.run("text added in version 3.2", size=9.5, hl=True))
    brk = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(copy.deepcopy(b.doc.sections[0]._sectPr))
    brk.append(ppr)
    els.append(brk)
    first = b.body[0]
    for e in els:
        first.addprevious(e)
