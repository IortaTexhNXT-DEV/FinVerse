"""Version 1.2 of the Product Maintenance FRS in BDOI's template: the Business Unit review edition.

Called by build_bdoi_frs_pm.py (bind() gives it the builder module) and, for the shared data, by
build_brd03_workbooks.py. It adds, in BDOI's template:

  * Introduction > Summary for the Business Unit Review (scope, end-to-end picture, key numbers, decisions, how to
    review, the checklist count);
  * the captions of every process flow, the process flows still missing, the status life-cycles and the integration
    context diagram (Graphviz, 240 dpi);
  * a menu-by-persona page in Annex G, section M.5 of Annex M, terms added to Annex N;
  * Annexes O to X (workflow and approvals, e-mail and notification texts, document prints, reports and schedules,
    integrations, non-functional requirements, data set-up at go-live, assumptions / dependencies / open questions
    not answered in BDOI's FRS, change control, review checklist).

The shared data (FRPM items and their acceptance criteria, the BRD mapping, the test cases re-keyed to BDOI's FR
numbers, the open items checked against BDOI's FRS) is computed here without the document, so the FRS and the
workbooks give the same numbers.
"""

from __future__ import annotations

import contextlib
import re
import shutil
import subprocess
import sys
import tempfile
from collections import OrderedDict
from pathlib import Path

import docx
import yaml
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import brd_common_final as final  # noqa: E402

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
SRC = REPO / "docs" / "deliverables" / "src"
PM = SRC / "BRD-03_Product_Maintenance"
NB = SRC / "BRD-01_New_Business"
V12 = yaml.safe_load((HERE / "brd03_v12_frs.yaml").read_text(encoding="utf-8"))
ITEMS = yaml.safe_load((HERE / "brd03_v12_items.yaml").read_text(encoding="utf-8"))
ADD = yaml.safe_load((HERE / "pm_frs_additions.yaml").read_text(encoding="utf-8"))
DOC = yaml.safe_load((HERE / "pm_document.yaml").read_text(encoding="utf-8"))
SOURCE = REPO / DOC["meta"]["source"]
CORE = None  # the builder module (build_bdoi_frs_pm), set by bind()
LANDSCAPE_W = 9.8
PORTRAIT_W = 7.4


def bind(core) -> None:
    global CORE
    CORE = core


# ===================================================================================================== shared data
_CACHE: dict = {}


def _source_doc():
    if "src" not in _CACHE:
        _CACHE["src"] = docx.Document(str(SOURCE))
    return _CACHE["src"]


def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def bdoi_items() -> "OrderedDict[str, dict]":
    """BDOI's items FRPM.001 to FRPM.021 with their sub-items (id -> title), read from the source FRS."""
    if "bdoi" in _CACHE:
        return _CACHE["bdoi"]
    out = OrderedDict()
    t = _source_doc().tables[3]
    for row in t.rows[1:]:
        rid = re.match(r"\s*(FRPM\.\d{3})", row.cells[0].text).group(1)
        subs = OrderedDict()
        for p in row.cells[3]._tc.iter(qn("w:p")):
            m = re.match(r"\s*(FRPM\.\d{3}\.\d{2})\s+(.*)", _text(p))
            if m and p.find(".//" + qn("w:b")) is not None and m.group(1) not in subs:
                subs[m.group(1)] = m.group(2).strip()
        out[rid] = {"component": row.cells[1].text.strip(), "capability": " ".join(row.cells[2].text.split()),
                    "subs": subs}
    _CACHE["bdoi"] = out
    return out


def frpm_items() -> list[dict]:
    """Every FR sub-item of the FRS (BDOI's and the added ones) with its item, title, origin, BIBS references and
    acceptance criteria, in document order (BDOI's items, the sub-items added inside them, the new items)."""
    if "items" in _CACHE:
        return _CACHE["items"]
    out = []
    bdoi = bdoi_items()
    added = {}
    for ns in ADD["new_subitems"]:
        added.setdefault(ns["row"], []).append(ns)
    for rid, item in bdoi.items():
        for sid, title in item["subs"].items():
            add = ADD["items"].get(sid, {})
            out.append({"id": sid, "item": rid, "capability": item["capability"], "title": title,
                        "origin": "BDOI FRS v1.0", "refs": add.get("refs", []), "ac": add.get("ac", [])})
        for ns in added.get(rid, []):
            out.append({"id": ns["id"], "item": rid, "capability": item["capability"], "title": ns["title"],
                        "origin": "Added in v1.1", "refs": ns["refs"], "ac": ns["ac"]})
    for it in ADD["new_items"]:
        for si in it["subitems"]:
            out.append({"id": si["id"], "item": it["id"], "capability": it["capability"], "title": si["title"],
                        "origin": "Added in v1.1 (new item)", "refs": si["refs"], "ac": si["ac"]})
    _CACHE["items"] = out
    return out


def fr_to_frpm() -> dict[str, list[str]]:
    """BIBS reference requirement (FR-PM, FR-NB) -> the FRPM sub-items that incorporate it (same order as the
    builder's note_refs)."""
    out: dict[str, list[str]] = {}

    def note(refs, sid):
        for r in refs:
            out.setdefault(r, [])
            if sid not in out[r]:
                out[r].append(sid)
    for sid, add in ADD["items"].items():
        note(add["refs"], sid)
    for ns in ADD["new_subitems"]:
        note(ns["refs"], ns["id"])
    for it in ADD["new_items"]:
        for si in it["subitems"]:
            note(si["refs"], si["id"])
    return out


def brd_mapping() -> "OrderedDict[str, dict]":
    """BRD ID -> {text, frpm: [FRPM ids as in the mapping], uam: bool}: BDOI's mapping with the changes of v1.1."""
    if "map" in _CACHE:
        return _CACHE["map"]
    out = OrderedDict()
    for row in _source_doc().tables[2].rows[1:]:
        rid = row.cells[0].text.replace("\n", "").strip()
        text = " ".join(row.cells[1].text.split())
        cell = " ".join(row.cells[2].text.split())
        out[rid] = {"text": text, "frpm": re.findall(r"FRPM\.\d{3}(?:\.\d{2})?", cell),
                    "uam": "User Access Maintenance" in cell}
    for rid, ch in DOC["mapping"].items():
        if "replace" in ch:
            out[rid]["frpm"] = list(ch["replace"])
        else:
            out[rid]["frpm"] += [f for f in ch["add"] if f not in out[rid]["frpm"]]
    _CACHE["map"] = out
    return out


def brd_of_item(sid: str) -> list[str]:
    """The BRD IDs whose mapping names the sub-item or its item."""
    item = sid[:8]
    return [rid for rid, m in brd_mapping().items() if sid in m["frpm"] or item in m["frpm"]]


def _plan(path: Path):
    sys.path.insert(0, str(SRC / "testplans"))
    import build_test_plan  # noqa: PLC0415
    return build_test_plan.load(path)


def _screens() -> dict[str, dict]:
    out = {}
    for f in sorted((PM / "pack" / "screens").glob("*.yaml")):
        for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
            out[s["id"]] = s
    return out


def walkthrough_steps() -> list[dict]:
    """The steps of the walkthroughs WT-A and WT-B with the FR-PM requirements of their screens."""
    if "wt" in _CACHE:
        return _CACHE["wt"]
    screens = _screens()
    out = []
    for w in yaml.safe_load((PM / "pack" / "walkthroughs.yaml").read_text(encoding="utf-8"))["walkthroughs"]:
        for i, s in enumerate(w["steps"], 1):
            out.append({"id": f"{w['id']}.{i}", "persona": s[0], "screen": s[1],
                        "title": screens.get(s[1], {}).get("title", s[1]),
                        "frs": screens.get(s[1], {}).get("frs", [])})
    _CACHE["wt"] = out
    return out


def _wt_of(refs) -> str:
    steps = [s["id"] for s in walkthrough_steps() if set(s["frs"]) & set(refs)]
    return ", ".join(steps[:6]) + (" ..." if len(steps) > 6 else "") if steps else "-"


NEGATIVE = {"Negative", "Boundary", "Security", "Error"}


def test_cases() -> list[dict]:
    """The test cases of the BDOI-format FRS: the cases of the BIBS test plan BRD-3 (and of BRD-1 for the
    non-package items) re-keyed to BDOI's FR numbers, plus one acceptance case per acceptance criterion of every FRPM
    sub-item."""
    if "tc" in _CACHE:
        return _CACHE["tc"]
    f2f = fr_to_frpm()
    out = []
    pm = _plan(PM / "brd03_cases.yaml")
    users = {k: v.get("user", "-") for k, v in pm.personas.items()}
    names = {k: v.get("name", k) for k, v in pm.personas.items()}
    for c in pm.cases:
        frpm = f2f.get(c.fr, [])
        out.append({"id": c.id, "title": c.title, "brd": ", ".join(c.brd), "frpm": ", ".join(frpm) or
                    "Refer to User Access Maintenance FRS", "frpm_list": frpm, "ref": c.fr,
                    "persona": names.get(c.persona, c.persona), "signin": users.get(c.persona, "-"),
                    "pre": c.pre, "steps": c.steps, "data": ", ".join(c.data), "expected": c.expected,
                    "type": "Negative" if c.negative else "Positive", "priority": c.priority, "screen": c.screen,
                    "wt": _wt_of([c.fr]), "source": "Test conditions of this FRS"})
    nb_frs = {r for r in f2f if r.startswith("FR-NB")}
    nb = _plan(NB / "brd01_cases.yaml")
    nusers = {k: v.get("user", "-") for k, v in nb.personas.items()}
    nnames = {k: v.get("name", k) for k, v in nb.personas.items()}
    for c in nb.cases:
        if c.fr not in nb_frs:
            continue
        frpm = f2f[c.fr]
        brd = sorted({b for s in frpm for b in brd_of_item(s)})
        out.append({"id": c.id, "title": c.title, "brd": ", ".join(brd), "frpm": ", ".join(frpm), "frpm_list": frpm,
                    "ref": c.fr, "persona": nnames.get(c.persona, c.persona), "signin": nusers.get(c.persona, "-"),
                    "pre": c.pre, "steps": c.steps, "data": ", ".join(c.data), "expected": c.expected,
                    "type": "Negative" if c.negative else "Positive", "priority": c.priority, "screen": c.screen,
                    "wt": "-", "source": "Test conditions of this FRS (non-package)"})
    screens = pm.screens
    for it in frpm_items():
        persona, screen = ITEMS["frpm_tests"][it["id"]]
        screen = screens.get(screen, screen)
        for i, ac in enumerate(it["ac"], 1):
            neg = bool(re.search(r"\b(cannot|refused|not|never|no)\b", ac, re.I))
            out.append({
                "id": f"TC-{it['id']}-AC{i}", "title": f"{it['title']}: acceptance criterion {i}",
                "brd": ", ".join(brd_of_item(it["id"])), "frpm": it["id"], "frpm_list": [it["id"]],
                "ref": ", ".join(it["refs"]), "persona": names.get(persona, persona),
                "signin": users.get(persona, "-"),
                "pre": "SIT/UAT data of the test plan; the user holds the role of the persona",
                "steps": (f"1. Sign in as {users.get(persona, '-')} ({names.get(persona, persona)}).\n"
                          f"2. Open {screen}.\n"
                          "3. Carry out the situation the criterion describes, with the test data.\n"
                          "4. Compare the result with the expected result."
                          if persona != "SYSTEM" else
                          "1. Prepare the record the criterion describes with the test data.\n"
                          "2. Let the daily run or event of the system take place (or ask the System "
                          "Administrator to start it).\n"
                          f"3. Open {screen} and compare the result with the expected result."),
                "data": "TD-PM-01 and the data of the scenario", "expected": ac,
                "type": "Negative" if neg else "Positive", "priority": "High", "screen": screen,
                "wt": _wt_of(it["refs"]), "source": "Acceptance of the BDOI-format FRS"})
    _CACHE["tc"] = out
    return out


def open_items() -> dict:
    """Our assumptions, dependencies, open questions and proposed rules checked against BDOI's FRS."""
    items = ITEMS["our_items"]
    answered = [i for i in items if i["status"] in ("Answered", "Answered earlier")]
    kept = [i for i in items if i["status"] in ("Open", "Partly")]
    return {"all": items, "answered": answered, "kept": kept,
            "partly": [i for i in items if i["status"] == "Partly"]}


# ================================================================================================= document helpers
@contextlib.contextmanager
def this_round(b):
    """The passages written inside are additions of version 1.2: shaded in the review copy."""
    old = b.hl
    b.hl = b.highlight
    try:
        yield
    finally:
        b.hl = old


def _render(b, dot_rel: str, dpi: int = 240) -> Path:
    if not hasattr(b, "tmp"):
        b.tmp = Path(tempfile.mkdtemp(prefix="pm12_"))
    png = b.tmp / (Path(dot_rel).stem + f"_{dpi}.png")
    if not png.exists():
        subprocess.run(["dot", "-Tpng", f"-Gdpi={dpi}", str(HERE / dot_rel), "-o", str(png)], check=True)
    return png


def _render_text(b, name: str, dot: str, dpi: int = 240) -> Path:
    if not hasattr(b, "tmp"):
        b.tmp = Path(tempfile.mkdtemp(prefix="pm12_"))
    src = b.tmp / f"{name}.dot"
    src.write_text(dot, encoding="utf-8")
    png = b.tmp / f"{name}.png"
    subprocess.run(["dot", "-Tpng", f"-Gdpi={dpi}", str(src), "-o", str(png)], check=True)
    return png


def cleanup(b) -> None:
    if hasattr(b, "tmp"):
        shutil.rmtree(b.tmp, ignore_errors=True)


def caption(b, text: str):
    return b.para(text, italic=True, size=9, jc="center", keep_next=False, space_after=160)


def figure_no(b) -> int:
    b.fig_no = getattr(b, "fig_no", 0) + 1
    return b.fig_no


def _level3_start(b, num_id: str) -> None:
    """BDOI's heading list has no start value at its third level (BDOI uses two levels); the third level starts at 1
    so that the sub-sections of 1.4 read 1.4.1, 1.4.2 ..."""
    n = b.numbering
    num = next(x for x in n.findall(qn("w:num")) if x.get(qn("w:numId")) == num_id)
    abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    ab = next(x for x in n.findall(qn("w:abstractNum")) if x.get(qn("w:abstractNumId")) == abstract)
    lvl = ab.findall(qn("w:lvl"))[2]
    if lvl.find(qn("w:start")) is None:
        st = OxmlElement("w:start")
        st.set(qn("w:val"), "1")
        lvl.insert(0, st)


def heading(b, text: str, level: int, numbered: bool = True):
    """A numbered heading of the Introduction in BDOI's format (a copy of BDOI's 'Document Preparation, Review and
    Approval' heading; level 3 one list level deeper)."""
    if not numbered:
        return b.para(text, style=f"Heading{level}", keep_next=True, no_num=True)
    model = CORE.find_par(b, "Document Preparation, Review and Approval", "Heading 2")._p
    h = CORE.clone_par(b, model, text)
    if level == 3:
        _level3_start(b, h.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val")))
        h.find(qn("w:pPr")).find(qn("w:pStyle")).set(qn("w:val"), "Heading3")
        h.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:ilvl")).set(qn("w:val"), "2")
        for sz in h.iter(qn("w:sz"), qn("w:szCs")):
            sz.set(qn("w:val"), "22")
    return h


def kv_table(b, rows, widths=(1.6, 5.8), size=8.5):
    return b.table(["Item", "Content"], rows, list(widths), size=size)


# ================================================================================ Introduction: summary for the BU
def summary(b, numbers: list[list[str]], checklist_count: int):
    s = V12["summary"]
    els = [heading(b, s["title"], 2), b.para(s["purpose"])]
    els.append(heading(b, "Scope on a page", 3))
    sc = s["scope"]
    n = max(len(sc["in"]), len(sc["out"]), len(sc["other"]))
    rows = []
    for i in range(n):
        a = f"{sc['in'][i][0]} ({sc['in'][i][1]})" if i < len(sc["in"]) else ""
        o = sc["out"][i] if i < len(sc["out"]) else ""
        h = f"{sc['other'][i][0]}: {sc['other'][i][1]}" if i < len(sc["other"]) else ""
        rows.append([a, o, h])
    els += [b.table(["In scope (BRD IDs)", "Out of scope", "Handled in another FRS"], rows, [3.1, 1.75, 1.75],
                    size=8), b.para("", space_after=0)]
    els.append(heading(b, "The process end to end", 3))
    els.append(b.para("Each column is a persona; the numbers give the order of the steps. Every process is drawn in "
                      "detail in section 4 (Process Flow), with the status life-cycles and the integration context."))
    els += [b.picture(_render(b, V12["e2e_figure"]["dot"]), max_w=6.6, max_h=7.6),
            caption(b, f"Figure 1 – {V12['e2e_figure']['caption']}")]
    els.append(heading(b, "Key numbers", 3))
    els += [b.table(["Item", "Number", "Where in this document"], numbers, [3.0, 1.5, 2.1], size=8.5),
            b.para("", space_after=0)]
    els.append(heading(b, "Decisions the Business Unit must take", 3))
    els.append(b.para("The ten decisions with the largest effect on the module. Each points to the row of the annex "
                      "where the options are given; every other point for decision is in Annex M and Annex V."))
    rows = [[str(i)] + d for i, d in enumerate(s["decisions"], 1)]
    els += [b.table(["No.", "Decision", "Why it matters", "Where", "Who decides"], rows, [0.35, 2.35, 1.6, 1.15, 1.15],
                    size=8), b.para("", space_after=0)]
    els.append(heading(b, "How to review", 3))
    els += [b.table(["Reviewer", "What to read", "About"], s["review"], [2.0, 3.8, 0.8], size=8.5),
            b.para("", space_after=0)]
    els.append(b.para(f"Before signing, the Business Unit confirms the {checklist_count} points of the Business Unit "
                      "Review Checklist (Annex X)."))
    anchor = CORE.find_par(b, "Business Requirements Mapping", "Heading 1")._p
    for e in els:
        anchor.addprevious(e)


def introduction(b):
    """The preparation sentence of v1.1 extended to v1.2, and the note on what v1.2 adds."""
    p = CORE.find_par(b, "This version 1.1 was elaborated")
    text = DOC["introduction"]["preparation"].replace("This version 1.1 was", "Versions 1.1 and 1.2 were")
    new = CORE.clone_par(b, p._p, text)
    p._p.addnext(new)
    new.addnext(CORE.clone_par(b, p._p, V12["intro_note"]))
    p._p.getparent().remove(p._p)


# ============================================================================================== process flow section
def process_flows(b):
    """Captions under every figure of the Process Flow section, then the flows, life-cycles and integration context
    added in this round, before the page break that precedes the Signoff Sheet."""
    start = CORE.find_par(b, "Process Flow", "Heading 1")._p
    end = CORE.find_par(b, "Signoff Sheet", "Heading 1")._p
    el, figs = start, []
    while el is not None and el is not end:
        if el.tag == qn("w:p") and el.find(".//" + qn("w:drawing")) is not None:
            figs.append(el)
        el = el.getnext()
    titles = []
    for fig in figs:
        prev = fig.getprevious()
        while prev is not None and not _text(prev).strip():
            prev = prev.getprevious()
        titles.append(_text(prev).strip())
    last = None
    for fig, title in zip(figs, titles):
        cap = caption(b, f"Figure {figure_no(b)} – {title}")
        fig.addnext(cap)
        last = cap
    head_model = CORE.find_par(b, "New Request Process Flow")._p
    els = []

    def head(text):
        h = CORE.strip_ids(__import__("copy").deepcopy(head_model))
        for r in h.findall(qn("w:r")):
            h.remove(r)
        CORE.ppr_insert(h.find(qn("w:pPr")), OxmlElement("w:pageBreakBefore"))
        h.append(b.run(text))
        return h
    for f in V12["flows"]:
        els += [head(f["title"]), b.picture(_render(b, f["dot"]), max_w=6.6, max_h=8.2),
                caption(b, f"Figure {figure_no(b)} – {f['caption']}")]
        b.stats["flows_v12"] += 1
    lc = V12["lifecycles"]
    els += [head(lc["title"]), b.para(lc["intro"])]
    for it in lc["items"]:
        els += [b.picture(_render(b, it["dot"]), max_w=6.6, max_h=4.4),
                caption(b, f"Figure {figure_no(b)} – {it['caption']}")]
        b.stats["lifecycles"] += 1
    ig = V12["integration_figure"]
    els += [head(ig["title"]), b.para("The systems and modules Product Maintenance exchanges information with, in "
                                      "business terms; each interface is described in Annex S."),
            b.picture(_render(b, ig["dot"]), max_w=6.6, max_h=6.0)]
    b.integration_fig = figure_no(b)
    els.append(caption(b, f"Figure {b.integration_fig} – {ig['caption']}"))
    b.stats["figures_total"] = b.fig_no
    CORE.Builder.insert_after(last, els)


# ===================================================================================================== Annex G page
def menu_page(a):
    b = a.b
    sys.path.insert(0, str(SRC / "signoff"))
    import signoff_pack  # noqa: PLC0415
    pack = signoff_pack._pack(PM / "pack" / "pack.yaml")
    menus = pack.menus()
    roles = [r for r in menus if any(x["own"] == "Yes" for x in menus[r])]
    cases = yaml.safe_load((PM / "brd03_cases.yaml").read_text(encoding="utf-8"))["personas"]
    screens = list(OrderedDict.fromkeys(x["screen"] for r in roles for x in menus[r] if x["own"] == "Yes"))
    has = {(r, x["screen"]) for r in roles for x in menus[r] if x["own"] == "Yes"}

    def esc(s):
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    head = "".join(f'<TD BGCOLOR="#014EA9"><FONT COLOR="white" POINT-SIZE="9">{esc(pack.persona_label(r))}'
                   f'<BR/>({esc(str(cases.get(r, {}).get("user", "-")))})</FONT></TD>' for r in roles)
    body = ""
    for i, sc in enumerate(screens):
        bg = "#F2F6FC" if i % 2 else "#FFFFFF"
        body += f'<TR><TD ALIGN="LEFT" BGCOLOR="{bg}">{esc(sc)}</TD>' + "".join(
            f'<TD BGCOLOR="{"#82B366" if (r, sc) in has else bg}"> </TD>' for r in roles) + "</TR>"
    dot = ('digraph menu { graph [pad=0.1, bgcolor="white"]; node [shape=plaintext, fontname="Arial", '
           'fontsize=10]; t [label=<<TABLE BORDER="0" CELLBORDER="1" CELLSPACING="0" CELLPADDING="5">'
           f'<TR><TD BGCOLOR="#014EA9"><FONT COLOR="white">Product Maintenance menu</FONT></TD>{head}</TR>'
           f'{body}</TABLE>>]; }}')
    png = _render_text(b, "menu_persona", dot)
    g = b.stats["screens"] + 1
    with this_round(b):
        a.h2(f"G.{g} Navigation: menu by persona")
        a.p("The Product Maintenance entries of the menu (Client & Policy › Product Maintenance) that each persona "
            "sees, with the sign-in ID of the persona in the BIBS UAT environment. A green cell means the persona "
            "sees the entry; the actions inside each screen follow the role (Annex J.5). The screen standards "
            "(layout, tables with row action menus, dates dd-MMM-yyyy, formatted amounts, messages) are those of "
            "the programme screen standards and are not repeated here.")
        a.els.append(b.picture(png, max_w=7.2, max_h=6.0))
        a.els.append(caption(b, f"Figure G.{g} – Product Maintenance menu by persona"))
        b.stats["menu_personas"] = len(roles)
        b.stats["menu_entries"] = len(screens)


# ================================================================================================ Annex M, section M.5
def observations_m5(a, row, header, widths):
    b = a.b
    with this_round(b):
        a.h2("M.5 Observations of version 1.2")
        rows = [row(o["ref"], o["topic"], o["bdoi"], o["brd"], o["proposes"], o["impact"], o["role"], "M.5")
                for o in V12["observations"]]
        a.table(header, rows, widths, size=7)
        b.stats["observations_v12"] = len(rows)


def glossary_add(a):
    b = a.b
    terms = [["BIBS", "BDO Insurance Brokering System, the system specified in this FRS"],
             ["CCM", "Centralized Communications Management, a BDO channel for sending e-mails"],
             ["ECM", "Enterprise content management: BDO's document store"],
             ["KPI", "Key performance indicator shown on the dashboard"],
             ["Maker-checker", "A record is made by one user and authorized by another (four eyes)"],
             ["NFR", "Non-functional requirement (Annex T)"],
             ["Service level", "Target hours of a workflow stage; amber from 80%, red once passed"],
             ["Sign-in ID", "The user ID of a reviewer in the BIBS UAT environment (walkthrough users workbook)"]]
    with this_round(b):
        a.h2("N.2 Terms added in version 1.2")
        a.table(["Term", "Meaning"], terms, [1.8, 5.4], size=9)
        b.stats["glossary_v12"] = len(terms)


# ========================================================================================================= Annexes O-X
def _section(a, landscape: bool):
    a.els.append(CORE.section_break(a.b, landscape=landscape))


def annex_o(a):
    b, w = a.b, V12["workflow"]
    a.h1("Annex O – Workflow and Approvals")
    a.p(w["intro"])
    a.h2("O.1 Who makes, checks and approves (picture)")
    roles = w["matrix_roles"]
    colour = {"M": "#DAE8FC", "A": "#F8CECC", "R": "#FFF2CC", "N": "#E1D5E7", "V": "#EEEEEE"}
    head = "".join(f'<TD BGCOLOR="#014EA9"><FONT COLOR="white" POINT-SIZE="9">{r}</FONT></TD>' for r in roles)
    body = ""
    for row in w["matrix"]:
        body += f'<TR><TD ALIGN="LEFT">{row[0]}</TD>' + "".join(
            f'<TD BGCOLOR="{colour.get(v, "#FFFFFF")}"><B>{v or " "}</B></TD>' for v in row[1:]) + "</TR>"
    legend = " ".join(f'<TD BGCOLOR="{c}">{k}</TD>' for k, c in colour.items())
    dot = ('digraph m { graph [pad=0.1, bgcolor="white"]; node [shape=plaintext, fontname="Arial", fontsize=10]; '
           't [label=<<TABLE BORDER="0" CELLBORDER="1" CELLSPACING="0" CELLPADDING="5">'
           f'<TR><TD BGCOLOR="#014EA9"><FONT COLOR="white">Process step</FONT></TD>{head}</TR>{body}</TABLE>>]; '
           'l [label=<<TABLE BORDER="0" CELLBORDER="1" CELLSPACING="0" CELLPADDING="4"><TR><TD>Legend</TD>'
           f'{legend}</TR><TR><TD></TD><TD>makes</TD><TD>approves</TD><TD>reviews</TD><TD>notified</TD>'
           '<TD>views</TD></TR></TABLE>>]; t -> l [style=invis]; }')
    a.els.append(b.picture(_render_text(b, "role_matrix", dot), max_w=9.4, max_h=4.6))
    a.els.append(caption(b, "Figure O.1 – Role-to-stage matrix of Product Maintenance"))
    a.p(w["matrix_intro"], size=9)
    a.h2("O.2 Stages, approvals and service levels per process")
    a.bullets(w["sla_rules"])
    b.stats["workflow_rows"] = 0
    for proc in w["processes"]:
        a.p(f"{proc['name']} (BRD: {proc['brd']})", bold=True, keep_next=True)
        a.table(w["columns"], proc["rows"], [1.55, 1.55, 1.5, 1.75, 1.2, 1.4, 0.95], size=7.5)
        b.stats["workflow_rows"] += len(proc["rows"])
    b.stats["workflow_processes"] = len(w["processes"])
    a.h2("O.3 Role-to-stage matrix (table)")
    a.table(["Process step"] + roles, w["matrix"], [2.3] + [0.75] * len(roles), size=8)
    a.h2("O.4 Status names: BDOI's FRS and the screens")
    a.p(w["status_map_intro"])
    a.table(["BDOI's status (FRPM.011.01, FRPM.016.01)", "Stage on the screens", "FRPM"], w["status_map"],
            [4.2, 3.6, 2.0], size=8.5)


def annex_p(a):
    b, n = a.b, V12["notifications"]
    a.h1("Annex P – E-mail and Notification Texts")
    a.p(n["intro"])
    a.h2("P.1 Overview")
    rows = [[x["id"], x["name"], x["channel"], x["recipients"], x["source"].split(" (")[0], x["frpm"]]
            for x in n["items"]]
    a.table(["No.", "Notice", "Channel", "Recipients", "Text", "FRPM"], rows, [0.5, 1.5, 1.05, 2.15, 0.9, 1.3],
            size=7.5)
    a.h2("P.2 The texts")
    for x in n["items"]:
        a.p(f"{x['id']} {x['name']}", bold=True, keep_next=True)
        body = x["body"].replace("\\n", "\n")
        rows = [["Trigger", x["trigger"]], ["Recipients (role)", x["recipients"]], ["Channel", x["channel"]],
                ["Subject / title", x["subject"]], ["Text", body], ["Attachments", x["attachments"]],
                ["Protection", x["protection"]], ["Frequency", x["frequency"]], ["Wording", x["source"]],
                ["Requirement", x["frpm"]]]
        a.els.append(kv_table(b, rows, (1.4, 6.0), size=8))
        a.els.append(b.para("", space_after=0))
    b.stats["notices"] = len(n["items"])
    b.stats["notices_bibs"] = sum(1 for x in n["items"] if x["source"].startswith("Sent by BIBS"))
    b.stats["notices_proposed"] = b.stats["notices"] - b.stats["notices_bibs"]


def _image(ref: str) -> Path | None:
    if not ref:
        return None
    brd, slug = ref.split(":")
    folder = PM if brd == "BRD-03" else NB
    p = folder / "screenshots" / f"{slug}.png"
    return p if p.exists() else None


def annex_q(a):
    b, d = a.b, V12["documents"]
    a.h1("Annex Q – Document Prints and Output Formats")
    a.p(d["intro"])
    a.h2("Q.1 Overview")
    rows = [[x["id"], x["name"], x["format"], x["protection"], x["layout"], x["frpm"]] for x in d["items"]]
    a.table(["No.", "Document", "Format", "Protection", "Layout", "FRPM"], rows, [0.5, 1.5, 1.7, 1.25, 1.55, 0.9],
            size=7.5)
    a.h2("Q.2 The documents")
    for x in d["items"]:
        a.p(f"{x['id']} {x['name']}", bold=True, keep_next=True)
        rows = [["Purpose", x["purpose"]], ["Recipients", x["recipients"]], ["Format", x["format"]],
                ["Protection", x["protection"]], ["Fields in order", x["fields"]], ["Layout", x["layout"]],
                ["Requirement", x["frpm"]]]
        a.els.append(kv_table(b, rows, (1.4, 6.0), size=8))
        img = _image(x["image"])
        if img:
            a.els.append(final.document_image(b, img))
            a.els.append(caption(b, f"Figure Q.{x['id'][3:]} – {x['name']}: current layout in the BIBS UAT environment"))
            b.stats["doc_images"] += 1
        else:
            a.els.append(b.para("", space_after=0))
    b.stats["documents"] = len(d["items"])


def annex_r(a):
    b, r = a.b, V12["reports"]
    a.h1("Annex R – Reports and Schedules")
    a.p(r["intro"])
    a.h2("R.1 Reports")
    a.table(["Report", "Frequency / time", "Recipients", "Filters", "Columns", "Sort", "Output", "FRPM"], r["items"],
            [1.35, 1.25, 1.05, 1.3, 2.4, 0.75, 0.95, 0.8], size=7.5)
    a.h2("R.2 Scheduled runs")
    a.p(r["schedules_intro"])
    a.table(["Run", "When", "What it does", "Recipients", "FRPM"], r["schedules"], [1.6, 1.3, 4.0, 1.5, 1.4],
            size=8)
    b.stats["reports"] = len(r["items"])
    b.stats["schedules"] = len(r["schedules"])


def annex_s(a):
    b, s = a.b, V12["integrations"]
    a.h1("Annex S – Integrations")
    a.p(s["intro"])
    a.p("The context diagram is Figure " + str(b.integration_fig) + " in section 4.", italic=True, size=9)
    a.table(s["columns"], s["items"], [0.45, 1.2, 0.8, 1.75, 1.15, 0.6, 0.85, 1.05, 0.9, 1.05], size=7.5)
    b.stats["integrations"] = len(s["items"])


def annex_t(a):
    b, n = a.b, V12["nfr"]
    a.h1("Annex T – Non-functional Requirements")
    a.p(n["intro"])
    a.table(n["columns"], n["items"], [0.45, 1.0, 2.3, 2.6, 1.5, 1.15, 0.75], size=8)
    b.stats["nfr"] = len(n["items"])


def annex_u(a):
    b, u = a.b, V12["data_setup"]
    a.h1("Annex U – Data Set-up and Migration at Go-live")
    a.p(u["intro"])
    a.table(u["columns"], u["items"], [0.45, 2.35, 1.4, 1.55, 1.0, 1.25, 0.85, 0.95], size=8)
    b.stats["data_setup"] = len(u["items"])


def annex_v(a):
    b = a.b
    oi = open_items()
    a.h1("Annex V – Assumptions, Dependencies and Open Questions")
    a.p("Only the items that BDOI's FRS does not answer are listed. Each assumption, dependency, open question and "
        "proposed rule of the earlier reviews was checked against BDOI's FRS: "
        f"{len(oi['answered'])} are answered by BDOI's FRS (or were closed earlier) and are not repeated here; they "
        "are listed with BDOI's clause in the traceability workbook (BIBS_RTM_BRD-03, sheet 'Answered in the BDOI "
        f"FRS'). {len(oi['partly'])} are answered in part: only the remaining part is asked below. The reference "
        "numbers of the earlier reviews (A-PM, D-PM, PQ, CLR-PM) are kept so that the earlier decisions can be traced; the "
        "last column gives the item of the requirements-collection workbook where the answer is recorded.")
    groups = [("V.1 Assumptions", "Assumption", "BDOI confirms each assumption with the sign-off of this FRS."),
              ("V.2 Dependencies", "Dependency", "Inputs other parties provide."),
              ("V.3 Open questions", "Open question", "Questions to answer before or at sign-off."),
              ("V.4 Proposed rules to confirm", "Proposed rule",
               "Rules proposed where the BRD needs a decision of BDOI.")]
    b.stats["open_kept"] = 0
    for title, kind, intro in groups:
        rows = []
        for i in oi["kept"]:
            if i["kind"] != kind:
                continue
            q = i.get("remaining") or i["text"]
            if i["status"] == "Partly":
                q = f"{q} (Answered in BDOI's FRS, {i['clause']}: {i['answer']})"
            rows.append([f"V-{b.stats['open_kept'] + len(rows) + 1:02d}", i["id"], q, i["format"], i["owner"],
                         i["needed"], i["rc"]])
        if not rows:
            continue
        a.h2(title)
        a.p(intro)
        a.table(["No.", "Ref.", "Item", "Format of the answer", "Owner", "Needed by", "Workbook item"], rows,
                [0.5, 0.85, 4.1, 1.6, 1.25, 0.8, 0.8], size=7.5)
        b.stats["open_kept"] += len(rows)
    b.stats["open_answered"] = len(oi["answered"])
    b.stats["open_partly"] = len(oi["partly"])
    b.v_text = " ".join(_text(e) for e in a.els)


def annex_w(a):
    b, c = a.b, V12["change_control"]
    a.h1("Annex W – Change Control after Sign-off")
    a.p(c["intro"])
    a.h2("W.1 What signing freezes")
    a.bullets(c["freezes"])
    a.h2("W.2 How a change is raised, assessed and approved")
    a.els.append(b.picture(_render(b, "figures/brd03_change_control.dot"), max_w=7.2, max_h=2.2))
    a.els.append(caption(b, "Figure W.1 – Change control after sign-off"))
    a.table(["Step", "What happens", "Who", "When"], c["steps"], [0.9, 3.9, 1.7, 0.9], size=8)
    a.p("Each change is assessed for its impact on scope, schedule, cost, the other BRDs and the tests; the owners of "
        "every BRD it touches approve it. The register: " + c["register"] + ".")


def annex_x(a):
    b, c = a.b, V12["checklist"]
    a.h1("Annex X – Business Unit Review Checklist")
    a.p(c["intro"])
    rows = [[str(i), r[0], r[1], r[2], r[3], "☐"] for i, r in enumerate(c["items"], 1)]
    a.table(["No.", "Point", "Where", "What to confirm", "Role", "Done"], rows, [0.4, 1.9, 1.6, 1.85, 1.3, 0.45],
            size=8)
    b.stats["checklist"] = len(rows)


def annexes(a):
    b = a.b
    with this_round(b):
        _section(a, landscape=False)      # ends the section of Annex N (portrait)
        annex_o(a)
        _section(a, landscape=True)       # O landscape
        annex_p(a)
        annex_q(a)
        _section(a, landscape=False)      # P, Q portrait
        annex_r(a)
        annex_s(a)
        annex_t(a)
        annex_u(a)
        annex_v(a)
        _section(a, landscape=True)       # R to V landscape
        annex_w(a)
        annex_x(a)                        # W, X in the last section (portrait)


# ===================================================================================================== self-checks
def platform_texts() -> list[str]:
    """The literal titles of the in-system notices and the e-mail subjects and fixed texts that the platform sends
    for Product Maintenance (package requests and catalogue)."""
    base = REPO / "backend" / "src" / "main" / "java" / "com" / "iortatechnxt" / "brokerverse"
    out = set()
    for pkg in ("productmaint", "catalog"):
        for f in (base / pkg).rglob("*.java"):
            s = f.read_text(encoding="utf-8")
            for m in re.finditer(r"new Notice\(\s*\"([^\"]+)\"", s):
                out.add(m.group(1))
            for m in re.finditer(r"new OutboundEmail\((?:[^;]*?)\"([^\"]+)\"", s, re.S):
                out.add(m.group(1).replace("\\n", ""))
            for m in re.finditer(r"\"(\\n\\nThe advisory is attached[^\"]*)\"", s):
                out.add(m.group(1).replace("\\n", ""))
    return sorted(out)


def report_names() -> list[str]:
    base = REPO / "backend" / "src" / "main" / "java" / "com" / "iortatechnxt" / "brokerverse" / "productmaint" / \
        "report"
    names = []
    for f in sorted(base.glob("*Report.java")):
        m = re.search(r"\"(Package [A-Z][\w ]+?(?:Report|History))\"", f.read_text(encoding="utf-8"))
        if m:
            names.append(m.group(1))
    return names


def self_checks(b, doc_text: str) -> list[str]:
    problems = []
    texts = " ".join(f"{x['subject']} {x['body']}" for x in V12["notifications"]["items"])
    for t in platform_texts():
        if t.strip() not in texts:
            problems.append(f"platform notice or e-mail not in Annex P: {t!r}")
    reps = " ".join(r[0] for r in V12["reports"]["items"])
    for n in report_names():
        if n not in reps:
            problems.append(f"report not in Annex R: {n}")
    oi = open_items()
    for i in oi["answered"]:
        if re.search(rf"\b{re.escape(i['id'])}\b", getattr(b, "v_text", "")):
            problems.append(f"answered item {i['id']} still asked in Annex V")
    if re.search(r"(?i)password:\s*(?!<password>)\S", doc_text):
        problems.append("a password value in the document")
    return problems
