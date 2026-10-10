"""Version 1.3 (Business Unit review edition) of the User Access Maintenance FRS in BDOI's template.

Used by build_brd11_bdoi_frs.py: the summary for the Business Unit Review (section 1.8), the process flows, status
life-cycles and integration context of Annex A, the menu by persona of Annex J, the observations of Annex Q.5, the
new Annexes S to AB, and the self-checks of this round. The content is in brd11_review.yaml and brd11_flows.yaml;
the e-mail texts are checked against the notification code of the platform at every build.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

import openpyxl
import yaml
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import brd_common_final as final  # noqa: E402

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
UA = REPO / "docs" / "deliverables" / "src" / "BRD-11_User_Access_Maintenance"
OUT_PROG = REPO / "docs" / "deliverables" / "out" / "Programme"
JAVA = REPO / "backend" / "src" / "main" / "java" / "com" / "iortatechnxt" / "brokerverse"
REV = yaml.safe_load((HERE / "brd11_review.yaml").read_text(encoding="utf-8"))
FLOWS = yaml.safe_load((HERE / "brd11_flows.yaml").read_text(encoding="utf-8"))["figures"]
INPUTS_XLSX = OUT_PROG / "BDOI_Inputs" / "BIBS_Inputs_BRD-00_BDOI_Requirements_and_Inputs_by_BRD_v1.2.xlsx"
NFR_XLSX = OUT_PROG / "Requirements" / "BIBS_NFR_BRD-00_Non_Functional_Requirements_Register_v1.0.xlsx"
# The platform files whose notices and e-mails belong to User Access Maintenance (every one must be quoted in Annex T)
NOTICE_SOURCES = ["nbadmin/service/AccessRequestNotifier.java", "nbadmin/service/AccessDecisionService.java",
                  "nbadmin/service/AccessScheduledChanges.java", "nbadmin/service/DirectRoleEditAlerts.java",
                  "nbadmin/service/SodRuleService.java", "nbadmin/service/SecurityParameterNotices.java",
                  "nbadmin/service/MfaResetNotices.java", "nbadmin/service/PasswordNoticeMailer.java",
                  "messaging/service/JobFailureMailer.java", "nbadmin/service/AccessImplementationService.java",
                  "identity/service/IdentityAlerts.java", "identity/service/BreakGlassAlerts.java"]
ANNEX_LETTERS = {"workflow": "S", "emails": "T", "documents": "U", "reports": "V", "integrations": "W",
                 "nfr": "X", "data": "Y", "open": "Z", "change": "AA", "checklist": "AB"}


# ------------------------------------------------------------------------------------------- shared data
def figure_numbers() -> dict[str, int]:
    """Figure numbers of the figures added in version 1.3, in the order they appear in the document."""
    order = ["summary", "A", "J", "S", "W"]
    figs = sorted(FLOWS, key=lambda f: order.index(f["where"]))
    return {f["id"]: i for i, f in enumerate(figs, 1)}


def q_label(q_rows: list[dict], ref: str) -> str:
    """'Q-01 (C01)' for a conflict or slip ID, as numbered in Annex Q."""
    for r in q_rows:
        if r["no"].endswith(f"({ref})"):
            return r["no"]
    return ref


def inputs_tab() -> list[dict]:
    """The BRD-11 rows of the programme workbook BDOI Requirements and Inputs by BRD v1.2 (same IDs)."""
    wb = openpyxl.load_workbook(INPUTS_XLSX, read_only=True)
    out = []
    for r in wb["BRD-11 User Access"].iter_rows(min_row=5, values_only=True):
        if r[0]:
            out.append({"id": r[0], "type": r[2], "need": r[3], "why": r[4], "proposal": r[5], "source": r[6],
                        "owner": r[7], "due": r[8].strftime("%d-%b-%Y") if hasattr(r[8], "strftime") else r[8],
                        "priority": r[9], "status": r[10]})
    wb.close()
    return out


def nfr_register() -> dict[str, str]:
    wb = openpyxl.load_workbook(NFR_XLSX, read_only=True)
    out = {r[0]: r[2] for r in wb["NFR register"].iter_rows(min_row=5, values_only=True) if r[0]}
    wb.close()
    return out


def classify_items() -> dict[str, str]:
    """Every reference of our sources classified once: open / input / answered / assumption / dependency."""
    oi = REV["open_items"]
    where: dict[str, str] = {}

    def put(ref, kind):
        if ref in where and where[ref] != kind:
            raise SystemExit(f"{ref} is classified twice ({where[ref]}, {kind})")
        where[ref] = kind
    for o in oi["open"]:
        for ref in re.findall(r"[A-Z]+(?:-[A-Z]+)?-?Q?\d+[\w-]*", o["refs"]):
            put(ref, "open")
    for ref in oi["covered_by_open"]:
        put(ref, "open")
    for ref in oi["inputs"]:
        put(ref, "input")
    for ref in oi["answered"]:
        put(ref, "answered")
    for a in oi["assumptions"]:
        put(a[0], "assumption")
    for d in oi["dependencies"]:
        put(d[0], "dependency")
    return where


def check_classification(where: dict[str, str]) -> int:
    """Every CLR-UA, UQ, IQ, A-UA, D-UA of the BIBS FRS and every row of the BRD-11 inputs tab is classified."""
    frs = (UA / "FRS_BRD11_USER_ACCESS_MAINTENANCE.md").read_text(encoding="utf-8")
    needed = set(re.findall(r"\bCLR-UA-\d\d\b", frs)) | set(re.findall(r"^\| (UQ\d\d|IQ0[45]|A-UA-\d\d|D-UA-\d\d) \|",
                                                                       frs, re.M))
    needed |= {r["id"] for r in inputs_tab()}
    missing = sorted(needed - set(where))
    if missing:
        raise SystemExit(f"open items not classified: {missing}")
    clash = sorted(r for r, k in where.items() if k == "open" and r in REV["open_items"]["answered"])
    if clash:
        raise SystemExit(f"open questions answered in BDOI's FRS: {clash}")
    return len(needed)


def check_emails() -> int:
    """Every notice of the BRD-11 notifications catalogue and every notice source of the platform is in Annex T,
    and the quoted texts are found in the platform code."""
    items = REV["emails"]["items"]
    ids = {i["id"] for i in items}
    cat = yaml.safe_load((UA / "pack" / "notifications.yaml").read_text(encoding="utf-8"))["notifications"]
    missing = sorted({n["id"] for n in cat} - ids)
    if missing:
        raise SystemExit(f"Annex T lacks the notices {missing}")
    code = "\n".join((JAVA / f).read_text(encoding="utf-8") for f in NOTICE_SOURCES)
    for f in NOTICE_SOURCES:
        text = (JAVA / f).read_text(encoding="utf-8")
        if not any(frag in text for i in items for frag in i["code"]):
            raise SystemExit(f"no Annex T notice quotes {f}")
    for i in items:
        for frag in i["code"]:
            if frag not in code:
                raise SystemExit(f"{i['id']}: text {frag!r} not found in the platform")
    # every class of the platform that queues an e-mail or a notice in the user-access packages is a known source
    for p in (list((JAVA / "nbadmin" / "service").glob("*.java")) + list((JAVA / "security").rglob("*.java"))
              + list((JAVA / "identity").rglob("*.java"))):
        t = p.read_text(encoding="utf-8")
        if re.search(r"queueEmail\(|notifyUser\(|notifyPermission\(|alerts\.raise\(", t):
            rel = p.relative_to(JAVA).as_posix()
            if rel not in NOTICE_SOURCES:
                raise SystemExit(f"notice source {rel} not quoted in Annex T")
    return len(items)


def check_reports(rows: list[list[str]]) -> int:
    docs = yaml.safe_load((UA / "pack" / "documents.yaml").read_text(encoding="utf-8"))["documents"]
    names = {r[0] for r in rows}
    missing = [d["name"] for d in docs if d["name"].replace(" (Excel)", "") not in names]
    if missing:
        raise SystemExit(f"Annex V lacks the reports {missing}")
    return len(rows)


# ------------------------------------------------------------------------------------------- figures
def figure_elements(b, figs: dict[str, Path], fid: str, nums: dict[str, int], max_h=8.4) -> list:
    spec = next(f for f in FLOWS if f["id"] == fid)
    els = [b.picture(figs[fid], max_w=6.9, max_h=max_h)]
    els.append(b.para(f"Figure {nums[fid]} - {spec['caption']}", italic=True, size=8.5, jc="center", space_after=40))
    els.append(b.para(f"How to read it: {spec['read']} Traces to: {spec['brd']}.", size=8.5, space_after=160))
    b.stats["figures_added"] += 1
    return els


# ------------------------------------------------------------------------------------------- 1.8 summary
def heading2(b, model_text="References"):
    for p in b.doc.paragraphs:
        if p.text.strip() == model_text and p.style.name == "Heading 2":
            return p._p
    raise SystemExit("Heading 2 'References' not found")


def add_summary(b, figs, nums, key_numbers: list[list[str]], q_rows: list[dict]):
    from build_brd11_bdoi_frs import find_par, strip_ids  # noqa: PLC0415
    s = REV["summary"]
    model_h2 = heading2(b)
    model_body = find_par(b, "The following documents are used as reference")._p

    def h2(text):
        h = strip_ids(copy.deepcopy(model_h2))
        for el in list(h):
            if el.tag != qn("w:pPr"):
                h.remove(el)
        ppr = h.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is None:
            from build_brd11_bdoi_frs import ppr_insert  # noqa: PLC0415
            ppr_insert(ppr, OxmlElement("w:pageBreakBefore"))
        h.append(b.run(text))
        return h

    def body(text):
        return b.clone_par(model_body, text)

    def sub(text):
        return b.para(text, bold=True, keep_next=True, size=10.5, color="014EA9", space_after=60)
    els = [h2("Summary for the Business Unit Review"), body(s["intro"])]
    els += [sub("Scope on a page"), b.table(s["scope_header"], s["scope"], [2.6, 2.3, 2.5], size=8), b.para("")]
    els += [sub("The process from end to end")] + figure_elements(b, figs, "brd11_fig_e2e", nums, max_h=7.6)
    els += [sub("Key numbers"), body(s["numbers_intro"]),
            b.table(["Item", "Number", "Where"], key_numbers, [3.6, 1.0, 2.8], size=8.5), b.para("")]
    rows = []
    for no, text, refs, who in s["decisions"]:
        where = "; ".join(q_label(q_rows, r.strip()) if re.fullmatch(r"[CS]\d\d", r.strip()) else r.strip()
                          for r in re.split(r",\s*(?=[CS]\d\d\b)|;\s*", refs))
        rows.append([str(no), text, where.replace("Z: ", "Annex Z: "), who])
    els += [sub("Decisions the Business Unit must take"), body(s["decisions_intro"]),
            b.table(["No.", "Decision", "Annex Q / Z", "Decided by"], rows, [0.4, 3.6, 1.6, 1.8], size=8),
            b.para("")]
    els += [sub("How to review this document"), body(s["review_intro"]),
            b.table(s["review_header"], s["review"], [2.0, 2.4, 2.4, 0.6], size=8), b.para("")]
    els += [body(s["workbooks_intro"]),
            b.table(["Workbook", "What it holds"], s["workbooks"], [2.2, 5.2], size=8.5), b.para("")]
    n_check = len(REV["checklist"]["rows"])
    els.append(body(f"Before signing, the Business Unit confirms the {n_check} points of the review checklist in "
                    f"Annex AB."))
    anchor = find_par(b, "Business Requirements Mapping", "Heading 1")._p
    for el in els:
        anchor.addprevious(el)


# ------------------------------------------------------------------------------------------- Annex A figures
def add_annex_a_figures(b, figs, nums):
    from build_brd11_bdoi_frs import find_par  # noqa: PLC0415
    head = find_par(b, "Annex A", "Heading 1")._p
    el = head.getnext()
    while el is not None and el.tag != qn("w:tbl"):
        el = el.getnext()
    if el is None:
        raise SystemExit("Annex A table not found")
    els = [b.para(""), b.para("Flows added in version 1.3", bold=True, size=11, color="014EA9", keep_next=True),
           b.para(REV["annex_a"]["intro"])]
    for f in FLOWS:
        if f["where"] == "A":
            els += figure_elements(b, figs, f["id"], nums, max_h=8.2)
    b.insert_after(el, els)


# ------------------------------------------------------------------------------------------- new annexes
class NewAnnexes:
    def __init__(self, a, figs, nums, nfr_rows, q_rows):
        self.a = a
        self.b = a.b
        self.figs = figs
        self.nums = nums
        self.nfr_rows = nfr_rows
        self.q_rows = q_rows

    def kv(self, rows, widths=(1.6, 5.8), size=8.5):
        self.a.table(["Item", "Detail"], rows, list(widths), size=size)

    def menus(self):
        """Annex J: menu by persona and the reference to the screen standards."""
        a = self.a
        k = self.b.stats["screens"] + 1
        a.h2(f"J.{k} Menu by persona")
        a.p("The User Access Maintenance entries of the menu of each persona (the sidebar shows an entry only when "
            "one of the user's active group profiles holds its permission). The figure is the one-page view for the "
            "review; the full sidebar of every persona of BIBS is in the BIBS sign-off workbook of BRD-11.")
        a.els += figure_elements(self.b, self.figs, "brd11_fig_menus", self.nums, max_h=6.0)
        a.p("Screen standards: the layout, formats of dates and amounts, tables with row action menus and the "
            "common screen elements follow the BIBS programme screen standards; they are not repeated here.")

    # S
    def workflow(self):
        w = REV["workflow"]
        a = self.a
        a.h1("Annex S – Workflow and Approvals")
        a.p(w["intro"])
        a.els += figure_elements(self.b, self.figs, "brd11_fig_approvals", self.nums, max_h=6.0)
        a.p("The process flows of Annex A draw each route step by step; the table below gives the stages, the "
            "controls and the time rules.")
        a.table(w["header"], w["rows"], [1.25, 1.0, 1.35, 1.6, 1.2, 1.0], size=7.5)
        a.p(w["matrix_intro"], bold=True, keep_next=True)
        a.table(w["matrix_header"], w["matrix"], [1.8] + [0.8] * 7, size=7.5)
        a.p(w["escalation"])
        self.b.stats["workflow_rows"] = len(w["rows"])

    # T
    def emails(self):
        e = REV["emails"]
        a = self.a
        a.h1("Annex T – E-mail and Notification Texts")
        a.p(e["intro"])
        rows = [[i["id"], i["name"], i["channel"], i["recipients"]] for i in e["items"]]
        a.table(["No.", "Notice", "Channel", "Recipients"], rows, [0.6, 2.0, 1.6, 3.2], size=7.5)
        for i in e["items"]:
            a.h2(f"{i['id']} {i['name']}")
            self.kv([["Trigger", i["trigger"]], ["Recipients", i["recipients"]], ["Channel", i["channel"]],
                     ["Subject / title", i["subject"]], ["Text", i["body"]],
                     ["Attachments", e["common"]["attachments"]], ["Protection", e["common"]["protection"]],
                     ["Frequency", i["frequency"]],
                     ["Wording", "Wording of the system today, proposed by iorta TechNXT for BDOI approval "
                                 "(requirements collection workbook, sheet E-mail wordings, item " + i["id"] + ")"],
                     ["Traces to", f"{i['brd']}; {i['frum']}"]])
        self.b.stats["emails"] = len(e["items"])

    # U
    def documents(self):
        from build_brd11_bdoi_frs import clean  # noqa: PLC0415
        d = REV["documents"]
        a = self.a
        docs = yaml.safe_load((UA / "pack" / "documents.yaml").read_text(encoding="utf-8"))["documents"]
        a.h1("Annex U – Document Prints and Output Formats")
        a.p(d["intro"])
        a.p(d["file_name_rule"])
        k = 0
        for doc in docs:
            k += 1
            extra = d["items"][doc["id"]]
            a.h2(f"U.{k} {doc['name']} ({doc['id']})")
            fields = "\n".join(f"{n}. {clean(f)}" for n, f in enumerate(doc["fields"], 1))
            self.kv([["Purpose", extra["purpose"]], ["Recipients", extra["recipients"]],
                     ["Format", clean(doc["format"])], ["Produced from", clean(doc["produced"])],
                     ["Protection", clean(doc["protected"])],
                     ["Fields in order", fields], ["BDOI reference", extra["bdoi"]],
                     ["Layout", "Current layout of the system below; the Business Unit confirms it or gives its "
                                f"layout: requirements collection workbook, item {extra['rc']}"]], size=8)
            png = UA / "screenshots" / f"{doc['shot']}.png"
            if not png.exists():
                raise SystemExit(f"image missing: {png}")
            a.els.append(final.document_image(self.b, png))
            a.p(f"Figure U.{k}: current layout of the {doc['name']}"
                + (f" ({clean(doc['shot_caption'])})" if doc.get("shot_caption") else ""), italic=True, size=8, jc="center")
        for doc in d["extra"]:
            k += 1
            a.h2(f"U.{k} {doc['name']} ({doc['id']})")
            fields = "\n".join(f"{n}. {f}" for n, f in enumerate(doc["fields"], 1))
            self.kv([["Purpose", doc["purpose"]], ["Recipients", doc["recipients"]], ["Format", doc["format"]],
                     ["Produced from", doc["produced"]], ["Protection", doc["protected"]],
                     ["Fields in order", fields], ["BDOI reference", doc["bdoi"]],
                     ["Layout", f"Layout to be provided or confirmed by the Business Unit: requirements collection "
                                f"workbook, item {doc['rc']}"]], size=8)
        self.b.stats["documents"] = k

    # V
    def reports(self):
        r = REV["reports"]
        a = self.a
        a.h1("Annex V – Reports and Schedules")
        a.p(r["intro"])
        a.table(r["header"], r["rows"], [0.95, 0.6, 1.0, 1.05, 1.75, 0.55, 0.65, 0.85], size=7)
        a.p(r["schedules_intro"], bold=True, keep_next=True)
        a.table(r["schedules_header"], r["schedules"], [1.15, 0.85, 1.9, 1.45, 1.05, 1.0], size=7.5)
        self.b.stats["reports"] = check_reports(r["rows"])
        self.b.stats["schedules"] = len(r["schedules"])

    # W
    def integrations(self):
        w = REV["integrations"]
        a = self.a
        a.h1("Annex W – Integrations")
        a.p(w["intro"] + f" (Figure {self.nums['brd11_fig_context']}).")
        a.table(w["header"], w["rows"], [0.95, 0.5, 1.25, 0.8, 0.8, 0.75, 1.15, 1.2], size=7)
        self.b.stats["integrations"] = len(w["rows"])

    # X
    def nfr(self):
        x = REV["nfr"]
        a = self.a
        a.h1("Annex X – Non-functional Requirements")
        a.p(x["intro"])
        a.table(x["header"], self.nfr_rows, [0.7, 0.85, 1.85, 2.0, 0.85, 1.15], size=7)
        self.b.stats["nfr_rows"] = len(self.nfr_rows)

    # Y
    def data_setup(self):
        y = REV["data_setup"]
        a = self.a
        a.h1("Annex Y – Data Set-up and Migration at Go-live")
        a.p(y["intro"])
        a.table(y["header"], y["rows"], [1.3, 1.1, 0.95, 1.0, 0.9, 1.25, 0.9], size=7.5)
        a.p(y["not_loaded"])
        self.b.stats["data_rows"] = len(y["rows"])

    # Z
    def open_items(self, where):
        o = REV["open_items"]
        a = self.a
        a.h1("Annex Z – Assumptions, Dependencies and Open Questions")
        a.p(o["intro"])
        a.h2("Z.1 Assumptions")
        a.table(["No.", "Reference", "Assumption", "Related"], [[f"Z.1.{i}", r[0], r[1], r[2]] for i, r in
                                                                    enumerate(o["assumptions"], 1)],
                [0.5, 0.9, 4.4, 1.6], size=8)
        a.h2("Z.2 Dependencies")
        a.table(["No.", "Reference", "Dependency", "Needed for"], [[f"Z.2.{i}", r[0], r[1], r[2]] for i, r in
                                                                      enumerate(o["dependencies"], 1)],
                [0.5, 0.9, 4.4, 1.6], size=8)
        a.h2("Z.3 Open questions not answered in BDOI's FRS")
        rows = []
        for i, q in enumerate(o["open"], 1):
            rows.append([f"Z.3.{i}", q["refs"], q["q"], q["why"], q["fmt"], q["owner"],
                         q["by"].strftime("%d-%b-%Y") if hasattr(q["by"], "strftime") else str(q["by"])])
        a.table(o["open_header"], rows, [0.45, 0.85, 1.85, 1.45, 1.35, 0.8, 0.65], size=7)
        n_answered = sum(1 for k in where.values() if k == "answered")
        a.p(f"{n_answered} earlier questions and assumptions are answered by BDOI's FRS and are not repeated here; "
            f"{sum(1 for k in where.values() if k == 'input')} items are values or data the Business Unit provides "
            f"(requirements collection workbook).", italic=True, size=9)
        self.b.stats["open_kept"] = len(o["open"])
        self.b.stats["answered"] = n_answered
        self.b.stats["assumptions"] = len(o["assumptions"])
        self.b.stats["dependencies"] = len(o["dependencies"])

    # AA
    def change_control(self):
        c = REV["change_control"]
        a = self.a
        a.h1("Annex AA – Change Control after Sign-off")
        a.p(c["intro"])
        a.p("What signing freezes", bold=True, keep_next=True)
        a.bullets(c["freezes"])
        a.p(c["not_frozen"])
        a.p("How a change is raised and decided", bold=True, keep_next=True)
        a.table(c["steps_header"], c["steps"], [0.8, 3.0, 1.9, 1.7], size=8)
        a.p(c["candidates"])

    # AB
    def checklist(self):
        c = REV["checklist"]
        a = self.a
        a.h1("Annex AB – Business Unit Review Checklist")
        a.p(c["intro"])
        rows = [[str(i), r[0], r[1], r[2], r[3], "☐"] for i, r in enumerate(c["rows"], 1)]
        a.table(c["header"], rows, [0.4, 1.4, 1.5, 2.3, 1.3, 0.5], size=8)
        self.b.stats["checklist"] = len(rows)


def nfr_rows(our_nfr_full: list[dict], frum_of, clean) -> list[list[str]]:
    reg = nfr_register()
    mp = REV["nfr"]["register"]
    over = REV["nfr"].get("approach", {})
    rows = []
    for n in our_nfr_full:
        num = int(n["id"].split("-")[-1])
        regs = mp.get(num, [])
        for r in regs:
            if r not in reg:
                raise SystemExit(f"{n['id']}: {r} not in the NFR register")
        frum = frum_of(n["frs"]) if n["frs"] else ""
        spec = frum or REV["nfr"]["specified_default"].get(num, "Technical Specification (BDOI IT)")
        rows.append([n["id"], clean(n["topic"]), clean(n["req"]), clean(over.get(num, n["approach"])), spec,
                     ", ".join(f"{r} {reg[r]}" for r in regs) or "-"])
    return rows
