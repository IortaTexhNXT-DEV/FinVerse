"""Product Maintenance FRS in BDOI's template, version 1.2 (Business Unit review edition).

    python docs/deliverables/src/programme/bdoi_template_frs/build_bdoi_frs_pm.py [--no-pdf]

Starts from a copy of BDOI's own FRS (docs/source-documents/FRS - BDOI BROKERSYS Product Maintenance v1.0 ...docx) and
keeps its template: styles, cover box, headers and footers, table formats, fonts and the table of contents (marked to
update when the document is opened). BDOI's text is kept word for word; the additions of version 1.1 come from:

  pm_frs_additions.yaml   requirements and acceptance criteria under BDOI's items, new sub-items, new items FRPM.022+
  pm_document.yaml        cover, revision log, introduction, mapping, process flows, sign-off, annex rows, Annex J, N
  pm_observations.yaml    Annex M: what v1.1 proposes for each conflict and slip, and the other observations
  figures/*.dot           the process flows added after BDOI's two figures
  brd03_v12.py            version 1.2: summary for the Business Unit, flows and life-cycles, Annexes O to X, M.5
  brd03_v12_frs.yaml      the content of version 1.2; brd03_v12_items.yaml the open items checked against BDOI's FRS

and from the BIBS release set BRD-3 (screens, notifications, walkthroughs, test cases, FRS) and the comparison data of
docs/deliverables/src/programme/comparisons (conflicts and slips).

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-03_Product_Maintenance/:
  BIBS_FRS-BDOI_BRD-03_Product_Maintenance_v1.2.docx                         clean copy
  BIBS_FRS-BDOI_BRD-03_Product_Maintenance_v1.2.pdf                          the clean copy as PDF (not committed)
  BIBS_FRS-BDOI_BRD-03_Product_Maintenance_v1.2_Changes_Highlighted.docx     the additions of version 1.2 shaded light
                                                                             yellow, with a change summary at the front

Self-checks (the build fails when one does not hold): every BRD ID of the Product Maintenance BRD is in the mapping and
maps to at least one FRPM item (or keeps BDOI's reference to the User Access Maintenance FRS); every FR-PM of the BIBS
FRS appears in Annex L; every conflict and slip of the comparison appears in Annex M; BDOI's original FRPM IDs and
texts are all still present; no internal code or restricted word in the added text; every notice, e-mail and
report of the platform for Product Maintenance is in Annex P and Annex R; no open question in Annex V that BDOI's FRS
answers; no password value.
"""

from __future__ import annotations

import argparse
import codecs
import contextlib
import copy
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter, OrderedDict
from pathlib import Path

import docx
import yaml
from docx.enum.text import WD_BREAK  # noqa: F401  (kept for readers extending the builder)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(HERE))
import brd03_v12 as v12  # noqa: E402

v12.bind(sys.modules[__name__])
SRC = REPO / "docs" / "deliverables" / "src"
PM = SRC / "BRD-03_Product_Maintenance"
NB = SRC / "BRD-01_New_Business"
CMP = SRC / "programme" / "comparisons" / "brd03_frs_comparison.yaml"
BRD_PDF = REPO / "docs" / "source-documents" / "Product Maintenance.pdf"

ADD = yaml.safe_load((HERE / "pm_frs_additions.yaml").read_text(encoding="utf-8"))
DOC = yaml.safe_load((HERE / "pm_document.yaml").read_text(encoding="utf-8"))
OBS = yaml.safe_load((HERE / "pm_observations.yaml").read_text(encoding="utf-8"))
OUT = REPO / "docs" / "deliverables" / "out" / DOC["meta"]["out_folder"]

HL_FILL = "FFF59D"          # light yellow of the review copy
HEAD_FILL = "D9E1F2"        # header fill of BDOI's annex tables
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"

# ----------------------------------------------------------------------------------------------------------- wording
# Internal codes of the BIBS sources replaced by business words in the generated annexes.
CODES = {
    "PRODUCT_MAINTAIN": "MBS", "MASTER_MAINTAIN": "master data maintainers", "PRODUCT_AUTHORIZE": "authorizers",
    "MASTER_AUTHORIZE": "authorizers", "INCENTIVE_CRITERIA_MAINTAIN": "incentive maintainers",
    "PKG_TSU_APPROVE": "TSU Head", "PKG_TSU_RECOMMEND": "TSU Team Lead", "PKG_MANCOM_SIGNOFF": "ManCom",
    "PKG_REQUEST_APPROVE": "Marketing approvers", "PKG_REQUEST": "requesters", "PKG_NEGOTIATE": "TSU",
    "PKG_ADVISORY": "TSU", "PRODUCT_VALIDATE": "validators", "WORK_ASSIGN": "team leaders",
    "QUOTE_MAINTAIN": "quotation users", "OPS_VIEW": "Operations", "ACCOUNT_PROCESS": "Processing",
    "PRODUCT_ARCHIVE_VIEW": "archive viewers", "PKG_QS_APPROVE": "Quotation Slip approvers",
    "MARKET_SEGMENT": "market segment list", "RETURN_REASON": "return reason list",
    "PKG_RESPONSE_OUTCOME": "response outcome list", "PKG_REQUEST_TYPE": "request type list",
    "PKG_REQUEST_REASON": "request reason list", "PACKAGE_EXPIRY_NOTICE_DAYS": "the expiry alert days",
    "PACKAGE_EXPIRY_REMINDER_DAYS": "the reminder days", "PKG_QUOTATION_SLIP": "package quotation slip",
    "PKG_RENEWAL_ADVISORY": "renewal advisory", "PKG_ADVISORY_GROUPS": "default advisory recipients",
    "PKG_ADVISORY_GROUP": "advisory recipient group list", "PACKAGE_RENEWAL_AUTODRAFT": "the automatic drafting setting",
    "PACKAGE_EXPIRY_MONITOR": "daily expiry check", "PKG_REQUEST_PREFIX": "request number prefix",
    "PKG_QS_REPLY_DAYS": "insurer reply days", "PKG_SLA_": "the service level settings",
    "DOCUMENT_TYPE": "document type list", "COVERAGE_KIND": "coverage kind list", "CLAUSE_KIND": "clause kind list",
    "INCENTIVE_TYPE": "incentive type list", "PACKAGE_EXPIRING": "Package expiring",
    "INCENTIVE_PRODUCT_INACTIVE": "Incentive criterion on an inactive product", "WORK_SLA_BREACH": "Work item overdue",
    "MKT_AO": "Marketing AO", "MKT_TL": "Marketing TL", "TSU_TL": "TSU Team Lead", "TSU_HEAD": "TSU Head",
    "NB_APPROVER": "New Business approver", "BUSINESS_ADMIN": "Business Administrator",
}
CODE_OK = {"FLEET_REPAIR",    # a clause code of the walkthrough data, shown on the screens as such
           "BIBS_CR_BRD", "BIBS_RTM_BRD", "BIBS_UAT_BRD"}  # the names of the workbooks of the review pack
CODE_RE = re.compile(r"\b[A-Z][A-Z0-9]*_[A-Z0-9][A-Z0-9_]*\b")
PHRASES = [
    # alert days of the BIBS release set, read with the 90 days of BDOI's FRPM.004.01 kept in this FRS
    (re.compile(r"the expiry monitor alerts at 60 days \(PACKAGE_EXPIRY_NOTICE_DAYS\) and again at 30 and 7 days"),
     "the expiry alerts start 90 days before the expiry date, with the configured reminder days"),
    (re.compile(r"the monitor alerts at 60, 30 and 7 days"), "the alerts follow FRPM.004.01 (from 90 days)"),
    (re.compile(r"\bSHA-256\b"), "Fingerprint"), (re.compile(r"\bregular expression\b"), "pattern"),
    (re.compile(r"Seed data:"), "SIT/UAT data:"), (re.compile(r"\bseed (client|data)\b"), r"SIT/UAT \1"),
    (re.compile(r"Exception [A-Z_]+ \(([^)]+)\)"), r"Alert '\1'"), (re.compile(r"Exception ([A-Z_]+): "), r"Alert: "),
    (re.compile(r"(?:every )?users? with ([A-Z][A-Z_]+) \(([^)]+)\)"), r"every \2 user"),
    (re.compile(r"\((?:permission |permissions )?([A-Z][A-Z0-9]*_[A-Z0-9_]+)(?:,? (?:or|and|/) [A-Z][A-Z0-9]*_[A-Z0-9_]+)*\)"), ""),
    (re.compile(r"\(([A-Z][A-Z0-9]*_[A-Z0-9_]+), ([^)]+)\)"), r"(\2)"),
    (re.compile(r"\bPHT\b"), ""), (re.compile(r"\bthe app\b"), "the system"), (re.compile(r"\bBIBS\b"), "the system"),
    (re.compile(r"\bin-app\b", re.I), "in-system"),
]
# Words a client document does not contain; kept encoded (rot13) so that this file does not contain them itself.
RESTRICTED = re.compile(r"\b(" + codecs.decode(
    "qrzb|qhzzl|snxr|fnzcyr qngn|cebgbglcr|fnaqobk|yberz vcfhz|gbqb|svkzr|pynhqr|naguebcvp|pungtcg|bcranv|pbcvybg|tvguho|trzvav|ntragf?|nffvfgnagf?|raqcbvagf?|qngnonfrf?|fpurznf?|cnlybnqf?|cyngsbez punatr|ohvyg|qrfvtarq|qrsrpgf?|fcevagf?", "rot13") + r")\b", re.I)
RESTRICTED_CS = re.compile(r"\b(AI|API|JSON|SQL|LLM|GPT)\b")
# Order of the children of a paragraph's and a run's properties (Word refuses some documents out of order).
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
             "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
RPR_AFTER_SHD = ["fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]


def ppr_insert(ppr, el):
    """Insert a child into a paragraph's properties at its place in the schema order."""
    name = el.tag.split("}")[1]
    pos = PPR_ORDER.index(name)
    for child in ppr:
        cname = child.tag.split("}")[1]
        if cname in PPR_ORDER and PPR_ORDER.index(cname) > pos:
            child.addprevious(el)
            return el
    ppr.append(el)
    return el


def rpr_shade(rpr):
    """Light-yellow shading of a run, placed before the run properties that follow it in the schema."""
    for s in rpr.findall(qn("w:shd")):
        rpr.remove(s)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), HL_FILL)
    for child in rpr:
        if child.tag.split("}")[1] in RPR_AFTER_SHD:
            child.addprevious(shd)
            return
    rpr.append(shd)


def clean(text) -> str:
    """Business wording of a text taken from the BIBS sources: no internal codes."""
    s = "" if text is None else str(text)
    for pat, rep in PHRASES:
        s = pat.sub(rep, s)
    for code, word in sorted(CODES.items(), key=lambda kv: -len(kv[0])):
        s = re.sub(r"\b" + re.escape(code) + (r"\b" if not code.endswith("_") else r"\w*"), word, s)
    s = re.sub(r"\b([A-Z][\w ]{1,40}?) \(\1\)", r"\1", s)
    s = re.sub(r"\s+([,.;:])", r"\1", re.sub(r"\(\s*\)", "", s))
    return re.sub(r"\s{2,}", " ", s).strip()


# ------------------------------------------------------------------------------------------------------ docx helpers
def strip_ids(el):
    for e in el.iter():
        for a in list(e.attrib):
            if a.startswith("{%s}" % W14) or a in (qn("w:rsidR"), qn("w:rsidRPr"), qn("w:rsidP"), qn("w:rsidRDefault"),
                                                   qn("w:rsidTr")):
                del e.attrib[a]
    for b in el.findall(".//" + qn("w:bookmarkStart")) + el.findall(".//" + qn("w:bookmarkEnd")):
        b.getparent().remove(b)
    return el


class Builder:
    def __init__(self, highlight: bool):
        self.highlight = highlight   # the review copy: the additions of version 1.2 are shaded
        self.hl = False              # shading of what is written now (v12.this_round switches it on)
        self.doc = docx.Document(str(REPO / DOC["meta"]["source"]))
        self.body = self.doc.element.body
        self.numbering = self.doc.part.numbering_part.element
        self.stats: Counter = Counter()
        self.log: list[dict] = []           # mapping changes for Annex M
        self.fr_to_frpm: dict[str, list[str]] = {}
        self.added_text: list[str] = []     # every text this builder writes (for the word checks)

    # ---- runs and paragraphs
    def run(self, text, bold=False, italic=False, size=None, font="Arial", color=None, hl=None):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), font)
        rpr.append(f)
        if bold:
            rpr.append(OxmlElement("w:b"))
            rpr.append(OxmlElement("w:bCs"))
        if italic:
            rpr.append(OxmlElement("w:i"))
        if color:
            c = OxmlElement("w:color")
            c.set(qn("w:val"), color)
            rpr.append(c)
        if size:
            for tag in ("w:sz", "w:szCs"):
                s = OxmlElement(tag)
                s.set(qn("w:val"), str(int(size * 2)))
                rpr.append(s)
        if (self.hl if hl is None else hl):
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), HL_FILL)
            rpr.append(shd)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = str(text)
        r.append(t)
        self.added_text.append(str(text))
        return r

    def para(self, text="", style=None, num=None, bold=False, italic=False, size=None, font="Arial", jc=None,
             keep_next=False, page_break_before=False, no_num=False, color=None, hl=None, space_after=None):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        if style:
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), style)
            ppr.append(ps)
        if keep_next:
            ppr.append(OxmlElement("w:keepNext"))
        if page_break_before:
            ppr.append(OxmlElement("w:pageBreakBefore"))
        if num or no_num:
            n = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), str(num[1] if num else 0))
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), str(num[0] if num else 0))
            n.append(il)
            n.append(ni)
            ppr.append(n)
        if space_after is not None:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:after"), str(space_after))
            ppr.append(sp)
        if jc:
            j = OxmlElement("w:jc")
            j.set(qn("w:val"), jc)
            ppr.append(j)
        mark = OxmlElement("w:rPr")          # font of the list number, as in BDOI's lists
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), font)
        mark.append(f)
        if size:
            for tag in ("w:sz", "w:szCs"):
                sz = OxmlElement(tag)
                sz.set(qn("w:val"), str(int(size * 2)))
                mark.append(sz)
        ppr.append(mark)
        p.append(ppr)
        if text != "":
            p.append(self.run(text, bold=bold, italic=italic, size=size, font=font, color=color, hl=hl))
        return p

    def new_num(self, like_num_id: int = 41) -> int:
        """A new list instance (numbering restarts at 1) on the abstract numbering of an existing list of BDOI's."""
        nums = self.numbering.findall(qn("w:num"))
        abstract = None
        for n in nums:
            if n.get(qn("w:numId")) == str(like_num_id):
                abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        n = OxmlElement("w:num")
        n.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        n.append(a)
        for lvl in range(3):
            o = OxmlElement("w:lvlOverride")
            o.set(qn("w:ilvl"), str(lvl))
            s = OxmlElement("w:startOverride")
            s.set(qn("w:val"), "1")
            o.append(s)
            n.append(o)
        mac = self.numbering.find(qn("w:numIdMacAtCleanup"))
        (mac.addprevious(n) if mac is not None else self.numbering.append(n))
        return new_id

    def items(self, entries, num_id, level=0):
        """Paragraphs of a requirement list: strings at this level, {text, sub} with a deeper list."""
        out = []
        for e in entries:
            if isinstance(e, dict):
                out.append(self.para(e["text"], style="ListParagraph", num=(num_id, level)))
                out += self.items(e.get("sub", []), num_id, level + 1)
            else:
                out.append(self.para(e, style="ListParagraph", num=(num_id, level)))
            if level == 0:
                self.stats["requirements_added"] += 1
        return out

    def acceptance(self, ac):
        out = [self.para("Acceptance criteria", bold=True, italic=True, keep_next=True)]
        n = self.new_num()
        for a in ac:
            out.append(self.para(a, style="ListParagraph", num=(n, 0)))
        self.stats["ac"] += len(ac)
        return out

    @staticmethod
    def insert_after(anchor, elements):
        for e in elements:
            anchor.addnext(e)
            anchor = e
        return anchor

    # ---- tables
    def table(self, header, rows, widths, size=8.5, head_fill=HEAD_FILL, shade_body=True):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl = t._tbl
        self.body.remove(tbl)
        tblpr = tbl.tblPr
        for old in tblpr.findall(qn("w:tblStyle")):
            tblpr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        tblpr.append(borders)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        total = sum(widths)
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), str(int(total * 1440)))
        tw.set(qn("w:type"), "dxa")
        for old in tblpr.findall(qn("w:tblW")):
            tblpr.remove(old)
        tblpr.append(tw)
        grid = tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        for ri, row in enumerate(t.rows):
            vals = header if ri == 0 else rows[ri - 1]
            if ri == 0:
                trpr = row._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:tblHeader"))
            for ci, cell in enumerate(row.cells):
                tc = cell._tc
                tcpr = tc.get_or_add_tcPr()
                tcw = OxmlElement("w:tcW")
                tcw.set(qn("w:w"), str(int(widths[ci] * 1440)))
                tcw.set(qn("w:type"), "dxa")
                for old in tcpr.findall(qn("w:tcW")):
                    tcpr.remove(old)
                tcpr.append(tcw)
                fill = head_fill if ri == 0 else (HL_FILL if (self.hl and shade_body) else None)
                if fill:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), fill)
                    tcpr.append(shd)
                for p in list(tc.findall(qn("w:p"))):
                    tc.remove(p)
                lines = str(vals[ci] if vals[ci] is not None else "").split("\n")
                for line in lines:
                    p = self.para(line, bold=(ri == 0), size=size, font="Calibri", hl=False, space_after=0)
                    tc.append(p)
        return tbl

    # ---- images
    def picture(self, path: Path, max_w=6.8, max_h=8.6):
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
        width = max_w
        if h / w * width > max_h:
            width = max_h * w / h
        p = self.doc.add_paragraph()
        p.add_run().add_picture(str(path), width=Inches(width))
        el = p._p
        self.body.remove(el)
        j = OxmlElement("w:jc")
        j.set(qn("w:val"), "center")
        ppr_insert(el.get_or_add_pPr(), j)
        ppr_insert(el.get_or_add_pPr(), OxmlElement("w:keepNext"))  # the caption stays with its figure
        return el


# ------------------------------------------------------------------------------------------------------- the sources
def our_frs() -> "OrderedDict[str, dict]":
    text = (PM / "FRS_BRD03_PRODUCT_MAINTENANCE.md").read_text(encoding="utf-8")
    out = OrderedDict()
    for block in re.findall(r"```fr\n(.*?)```", text, re.S):
        fr = yaml.safe_load(block)
        out[fr["id"]] = fr
    return out


def nb_frs() -> dict:
    text = (NB / "FRS_BRD01_NEW_BUSINESS.md").read_text(encoding="utf-8")
    out = {}
    for block in re.findall(r"```fr\n(.*?)```", text, re.S):
        fr = yaml.safe_load(block)
        out[fr["id"]] = fr
    return out


def brd_ids(our) -> list[str]:
    """The requirement IDs of the Product Maintenance BRD: read from the BRD itself (text pages 1-35; the scanned
    earlier version on pages 36-61 uses BRQID numbers), cross-checked with the BRD IDs traced in the BIBS FRS."""
    traced = sorted({re.match(r"(PMADD\d\d|BRPM\.\d{3})", b).group(1) for fr in our.values() for b in fr["brd"]})
    ids = set()
    if shutil.which("pdftotext") and BRD_PDF.exists():
        txt = subprocess.run(["pdftotext", "-l", "35", str(BRD_PDF), "-"], capture_output=True, text=True).stdout
        ids = set(re.findall(r"\b(PMADD0[1-9]|BRPM\.0\d\d)\b", txt))
    if ids and ids != set(traced):
        raise SystemExit(f"BRD IDs of the BRD and of the BIBS FRS differ: {sorted(ids ^ set(traced))}")
    key = lambda i: (0 if i.startswith("PMADD") else 1, i)  # noqa: E731
    return sorted(ids or traced, key=key)


def test_cases(path: Path, prefix: str) -> dict[str, tuple[list[str], int]]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    out = {}
    for fr_id, entry in (data.get("frs") or {}).items():
        n = fr_id.split("-")[-1]
        conds = [f"TC-{prefix}-{n}.{i}" for i in range(1, len(entry.get("conditions") or []) + 1)]
        out[fr_id] = (conds, len(entry.get("cases") or []))
    return out


def tc_text(fr_id, cases) -> str:
    conds, n = cases.get(fr_id, ([], 0))
    if not conds:
        return "-"
    return (f"{conds[0]} to {conds[-1].split('-')[-1]}" if len(conds) > 1 else conds[0]) + f" ({n} cases)"


# ------------------------------------------------------------------------------------------------------- the edits
def edit_cover(b: Builder):
    cell = b.doc.tables[0]._tbl
    for t in cell.iter(qn("w:t")):
        if t.text.strip() == DOC["meta"]["version_old"]:
            t.text = DOC["meta"]["version_new"]
            if b.highlight:
                rpr_shade(t.getparent().find(qn("w:rPr")))
            break
    else:
        raise SystemExit("cover: version not found")
    for p in cell.iter(qn("w:p")):
        if "Prepared by:" in "".join(t.text or "" for t in p.iter(qn("w:t"))):
            new = b.para(DOC["cover"]["drafting"], bold=True, size=14, color="014EA9", jc="center")
            p.addnext(new)
            return
    raise SystemExit("cover: 'Prepared by' not found")


def edit_revision_log(b: Builder):
    t = b.doc.tables[1]
    for i, log in ((2, DOC["revision_log"]), (3, v12.V12["revision"])):
        with (v12.this_round(b) if i == 3 else contextlib.nullcontext()):
            vals = [log[k] for k in ("date", "version", "description", "by")]
            for cell, v in zip(t.rows[i].cells, vals):
                p = cell.paragraphs[0]._p
                p.append(b.run(v, size=9))


def clone_par(b: Builder, p_el, text):
    """A new paragraph with the paragraph and run formatting of one of BDOI's paragraphs."""
    new = strip_ids(copy.deepcopy(p_el))
    runs = new.findall(qn("w:r"))
    rpr = copy.deepcopy(runs[0].find(qn("w:rPr"))) if runs and runs[0].find(qn("w:rPr")) is not None else None
    for el in list(new):
        if el.tag != qn("w:pPr"):
            new.remove(el)
    r = b.run(text)
    if rpr is not None:
        r.replace(r.find(qn("w:rPr")), rpr)
        if b.hl:
            rpr_shade(rpr)
    new.append(r)
    return new


def find_par(b: Builder, start: str, style: str | None = None):
    for p in b.doc.paragraphs:
        if p.text.strip().startswith(start) and (style is None or p.style.name == style):
            return p
    raise SystemExit(f"paragraph not found: {start}")


def edit_introduction(b: Builder):
    p = find_par(b, "This document once approved")
    find_par(b, "The System Analyst or equivalent")._p.addnext(clone_par(b, p._p, DOC["introduction"]["preparation"]))
    p = find_par(b, "The proposed solution will enable")
    p._p.addnext(clone_par(b, p._p, DOC["introduction"]["overview"]))


def edit_mapping(b: Builder, ids: list[str]):
    t = b.doc.tables[2]
    seen = {}
    for row in t.rows[1:]:
        rid = row.cells[0].text.replace("\n", "").strip()
        seen[rid] = row
    for rid, change in DOC["mapping"].items():
        row = seen[rid]
        cell = row.cells[2]._tc
        before = row.cells[2].text.replace("\n", " / ").strip()
        before = re.sub(r"\s+", " ", before).strip(" /")
        model = next(p for p in cell.findall(qn("w:p")) if p.find(qn("w:r")) is not None)
        if "replace" in change:
            for p in cell.findall(qn("w:p")):
                cell.remove(p)
            for f in change["replace"]:
                cell.append(clone_par(b, model, f))
            after = ", ".join(change["replace"])
        else:
            last = cell.findall(qn("w:p"))[-1]
            for f in change["add"]:
                new = clone_par(b, model, f)
                last.addnext(new)
                last = new
            after = before.replace(" / ", ", ") + " + " + ", ".join(change["add"])
        b.log.append({"id": rid, "before": before, "after": after, "why": change["why"],
                      "brd": re.sub(r"\s+", " ", row.cells[1].text).strip()})
    # self-check: every BRD ID present and mapped
    mapping = {}
    for row in t.rows[1:]:
        rid = row.cells[0].text.replace("\n", "").strip()
        mapping[rid] = re.sub(r"\s+", " ", row.cells[2].text).strip()
    missing = [i for i in ids if i not in mapping]
    unmapped = [i for i in ids if i in mapping and not re.search(r"FRPM\.\d{3}", mapping[i])
                and "User Access Maintenance" not in mapping[i]]
    if missing or unmapped:
        raise SystemExit(f"mapping check failed: missing {missing}, unmapped {unmapped}")
    b.stats["brd_ids"] = len(ids)
    b.stats["brd_uam"] = sum(1 for i in ids if "User Access Maintenance" in mapping[i])
    b.stats["mapping_changes"] = len(DOC["mapping"])
    return mapping


def subitem_blocks(cell_tc):
    """The sub-items of one cell of BDOI's Functional Requirements table: id -> list of child elements."""
    blocks = OrderedDict()
    current = None
    for el in cell_tc:
        if el.tag == qn("w:tcPr"):
            continue
        text = "".join(t.text or "" for t in el.iter(qn("w:t"))) if el.tag == qn("w:p") else ""
        m = re.match(r"\s*(FRPM\.\d{3}\.\d{2})\b", text)
        bold = el.tag == qn("w:p") and el.find(".//" + qn("w:b")) is not None
        if m and bold and el.find(qn("w:pPr") + "/" + qn("w:numPr")) is None:
            current = m.group(1)
            blocks[current] = [el]
        elif current:
            blocks[current].append(el)
    return blocks


def last_content(block):
    for el in reversed(block):
        if el.tag == qn("w:tbl") or "".join(t.text or "" for t in el.iter(qn("w:t"))).strip():
            return el
    return block[-1]


def block_num(block) -> int:
    ids = Counter()
    for el in block:
        n = el.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")) if el.tag == qn("w:p") else None
        if n is not None:
            ids[int(n.get(qn("w:val")))] += 1
    return ids.most_common(1)[0][0] if ids else 41


def note_refs(b: Builder, refs, frpm):
    for r in refs:
        b.fr_to_frpm.setdefault(r, [])
        if frpm not in b.fr_to_frpm[r]:
            b.fr_to_frpm[r].append(frpm)


def subitem_elements(b: Builder, sid, title, shall, ac, refs):
    out = [b.para(""), b.para(f"{sid} {title}", bold=True, keep_next=True)]
    n = b.new_num()
    out += b.items(shall, n)
    out += b.acceptance(ac)
    note_refs(b, refs, sid)
    return out


def edit_functional_requirements(b: Builder):
    t = b.doc.tables[3]
    rows = {}
    for row in t.rows[1:]:
        rid = re.match(r"\s*(FRPM\.\d{3})", row.cells[0].text).group(1)
        rows[rid] = row
    b.stats["bdoi_items"] = len(rows)
    all_blocks = {}
    for rid, row in rows.items():
        all_blocks.update(subitem_blocks(row.cells[3]._tc))
    b.stats["bdoi_subitems"] = len(all_blocks)
    for sid, add in ADD["items"].items():
        block = all_blocks[sid]
        anchor = last_content(block)
        num = block_num(block)
        els = b.items(add["shall"], num) + b.acceptance(add["ac"])
        b.insert_after(anchor, els)
        note_refs(b, add["refs"], sid)
        b.stats["elaborated"] += 1
    for ns in ADD["new_subitems"]:
        cell = rows[ns["row"]].cells[3]._tc
        prev = [s for s in all_blocks if s.startswith(ns["row"] + ".")]
        assert ns["id"] > max(prev), ns["id"]
        for el in subitem_elements(b, ns["id"], ns["title"], ns["shall"], ns["ac"], ns["refs"]):
            cell.append(el)
        b.stats["new_subitems"] += 1
    template = rows["FRPM.017"]._tr
    last = t.rows[-1]._tr
    for item in ADD["new_items"]:
        tr = strip_ids(copy.deepcopy(template))
        tcs = tr.findall(qn("w:tc"))
        for tc in tcs:
            for el in list(tc):
                if el.tag != qn("w:tcPr"):
                    tc.remove(el)
            shd = tc.find(qn("w:tcPr") + "/" + qn("w:shd"))
            if b.hl:
                if shd is None:
                    shd = OxmlElement("w:shd")
                    tc.find(qn("w:tcPr")).append(shd)
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), HL_FILL)
        tcs[0].append(b.para(item["id"], italic=True))
        tcs[1].append(b.para(item["component"]))
        tcs[2].append(b.para(item["capability"]))
        first = True
        for si in item["subitems"]:
            els = subitem_elements(b, si["id"], si["title"], si["shall"], si["ac"], si["refs"])
            for el in (els[1:] if first else els):
                tcs[3].append(el)
            first = False
        last.addnext(tr)
        last = tr
        b.stats["new_items"] += 1
        b.stats["new_item_subitems"] += len(item["subitems"])


def edit_process_flows(b: Builder):
    h = find_par(b, "New Request Process Flow")
    el = h._p
    nxt = el.getnext()
    while nxt is not None and nxt.find(".//" + qn("w:drawing")) is None:
        nxt = nxt.getnext()
    anchor = nxt
    work = Path(tempfile.mkdtemp(prefix="pmflows_"))
    for flow in DOC["process_flows"]:
        png = work / (Path(flow["dot"]).stem + ".png")
        subprocess.run(["dot", "-Tpng", "-Gdpi=200", str(HERE / flow["dot"]), "-o", str(png)], check=True)
        head = strip_ids(copy.deepcopy(el))
        for r in head.findall(qn("w:r")):
            head.remove(r)
        ppr_insert(head.find(qn("w:pPr")), OxmlElement("w:pageBreakBefore"))
        head.append(b.run(flow["title"]))
        anchor = b.insert_after(anchor, [head, b.picture(png, max_w=7.0, max_h=8.3)])
        b.stats["flows"] += 1
    brk = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    brk.append(r)
    anchor.addnext(brk)
    shutil.rmtree(work, ignore_errors=True)


def edit_signoff(b: Builder):
    cell = b.doc.tables[4].rows[0].cells[0]
    for p in cell.paragraphs:
        if p.text.strip() == "System Analyst":
            new = strip_ids(copy.deepcopy(p._p))
            for r in new.findall(qn("w:r")):
                new.remove(r)
            new.append(b.run(DOC["signoff"]["drafting"]))
            p._p.addnext(new)
            return
    raise SystemExit("sign-off: 'System Analyst' not found")


def edit_annex_rows(b: Builder):
    index = {"A": 5, "B": 7, "C": 8, "E": 9}
    for annex, rows in DOC["annex_rows"].items():
        t = b.doc.tables[index[annex]]
        template = t.rows[-1]._tr
        last = template
        for vals in rows:
            tr = strip_ids(copy.deepcopy(template))
            for tc, v in zip(tr.findall(qn("w:tc")), vals):
                paras = tc.findall(qn("w:p"))
                for p in paras[1:]:
                    tc.remove(p)
                p = paras[0]
                runs = p.findall(qn("w:r"))
                rpr = copy.deepcopy(runs[0].find(qn("w:rPr"))) if runs else None
                for r in runs:
                    p.remove(r)
                r = b.run(v, font="Calibri", size=11)
                if rpr is not None:
                    r.replace(r.find(qn("w:rPr")), rpr)
                    if b.hl:
                        rpr_shade(rpr)
                p.append(r)
                if b.hl:
                    tcpr = tc.find(qn("w:tcPr"))
                    for s in tcpr.findall(qn("w:shd")):
                        tcpr.remove(s)
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), HL_FILL)
                    # tcPr children order: shd comes after borders; appending before tcMar is tolerated by Word
                    mar = tcpr.find(qn("w:tcMar"))
                    (mar.addprevious(shd) if mar is not None else tcpr.append(shd))
            last.addnext(tr)
            last = tr
            b.stats["annex_rows"] += 1


# ------------------------------------------------------------------------------------------------- the new annexes
class Annexes:
    def __init__(self, b: Builder, mapping: dict, ids: list[str]):
        self.b = b
        self.mapping = mapping
        self.ids = ids
        self.els: list = []
        heading = find_par(b, "Annex A")
        self.h1_ppr = heading._p.find(qn("w:pPr"))

    def refs(self, text: str) -> str:
        """BIBS reference requirement IDs in a text replaced by the FRPM items of this FRS that incorporate them."""
        def one(m):
            items = self.b.fr_to_frpm.get(m.group(0))
            return items[0] if items else "User Access Maintenance FRS"
        return re.sub(r"FR-(?:PM|NB)-\d{3}", one, clean(text))

    def h1(self, text):
        p = OxmlElement("w:p")
        ppr = strip_ids(copy.deepcopy(self.h1_ppr))
        for r in ppr.findall(qn("w:rPr")):
            ppr.remove(r)
        ppr_insert(ppr, OxmlElement("w:pageBreakBefore"))
        p.append(ppr)
        p.append(self.b.run(text))
        self.els.append(p)

    def h2(self, text):
        self.els.append(self.b.para(text, style="Heading2", no_num=True, keep_next=True))

    def p(self, text, **kw):
        self.els.append(self.b.para(text, **kw))

    def bullets(self, texts):
        n = self.b.new_num()
        for t in texts:
            self.els.append(self.b.para(t, style="ListParagraph", num=(n, 0)))

    def table(self, header, rows, widths, size=8.5):
        self.els.append(self.b.table(header, rows, widths, size=size))
        self.els.append(self.b.para("", space_after=0))

    # Annex G – screens
    def screens(self):
        self.h1("Annex G – Screen Specifications")
        self.p("This annex specifies the Product Maintenance screens: the purpose of each screen, how the user reaches "
               "it, its fields with type, format, mandatory marker, list or source and validation, its actions with who "
               "may perform them and what happens, its rules, and one screenshot of the screen. Mandatory: Y = "
               "mandatory, N = optional, Cond. = mandatory when the condition stated applies. Dates are shown as "
               "dd-MMM-yyyy and amounts with thousands separators.")
        files = sorted((PM / "pack" / "screens").glob("*.yaml"))
        k = 0
        frs_titles = {}
        for f in files:
            for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
                k += 1
                frpm = sorted({x for fr in s.get("frs", []) for x in self.b.fr_to_frpm.get(fr, [])})
                frs_titles[s["id"]] = s["title"]
                self.h2(f"G.{k} {s['title']} ({s['id']})")
                self.p("Purpose: " + self.refs(" ".join(str(s["purpose"]).split())))
                if s.get("entry"):
                    self.p("Access: " + clean("; ".join(s["entry"])))
                if frpm:
                    self.p("Requirements: " + ", ".join(frpm))
                shots = s.get("shots") or []
                if shots:
                    n = int(s["id"].split("-")[-1])
                    png = PM / "screenshots" / f"scr-pm-{n:02d}-01-{shots[0]['state']}.png"
                    if png.exists():
                        self.els.append(self.b.picture(png, max_w=6.6, max_h=4.6))
                        self.p(f"Figure G.{k}: {clean(shots[0].get('caption', s['title']))}", italic=True, size=8,
                               jc="center")
                        self.b.stats["screenshots"] += 1
                rows = []
                for line in s.get("fields", []):
                    c = [x.strip() for x in str(line).split(" | ")]
                    c += ["-"] * (10 - len(c))
                    label = c[1] if c[0] in ("-", "") else f"{c[0]}: {c[1]}"
                    fmt = "; ".join(x for x in (c[3], c[5]) if x not in ("-", ""))
                    val = "; ".join(x for x in (c[8], c[9]) if x not in ("-", ""))
                    rows.append([clean(label), clean(c[2]), clean(c[4]), self.refs(fmt) or "-", self.refs(val) or "-"])
                if rows:
                    self.table(["Field", "Type", "Mandatory", "Format / list / source", "Validation and message"],
                               rows, [1.55, 0.8, 0.9, 1.9, 2.15], size=8)
                acts = []
                for line in s.get("actions", []):
                    c = [x.strip() for x in str(line).split(" | ")]
                    c += ["-"] * (6 - len(c))
                    acts.append([clean(c[0]), clean(c[1]), clean(c[2]), self.refs(c[3]),
                                 clean(c[4]) if c[4] not in ("", "-") else "-"])
                if acts:
                    self.table(["Action", "Who", "Available when", "What happens", "Resulting status"],
                               acts, [1.25, 1.35, 1.2, 2.5, 1.0], size=8)
                if s.get("rules"):
                    self.p("Rules:", bold=True, keep_next=True)
                    self.bullets([self.refs(r) for r in s["rules"]])
        self.b.stats["screens"] = k
        return frs_titles

    # Annex H – messages
    def messages(self, our, nb, screen_titles):
        self.h1("Annex H – Messages and Validations")
        self.p("The messages the user sees when a check fails or an action is not allowed, word for word as on the "
               "screens; <name> parts are filled in by the system. The fields and messages of each screen are also in "
               "Annex G.")
        rows, seen = [], set()
        for fr_id, fr in list(our.items()) + [(k, v) for k, v in nb.items() if k in self.b.fr_to_frpm]:
            frpm = ", ".join(self.b.fr_to_frpm.get(fr_id, [])) or "User Access Maintenance FRS"
            screen = clean(str(fr.get("screens", "")).split(";")[0])
            for v in fr.get("validations") or []:
                msg = clean(v[1])
                if msg in seen:
                    continue
                seen.add(msg)
                rows.append([screen, clean(v[0]), msg, frpm])
        for f in sorted((PM / "pack" / "screens").glob("*.yaml")):
            for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
                for line in s.get("fields", []):
                    c = [x.strip() for x in str(line).split(" | ")]
                    if len(c) >= 10 and c[9] not in ("-", ""):
                        for msg in (clean(m) for m in c[9].split(" / ")):
                            if msg and msg not in seen:
                                seen.add(msg)
                                rows.append([clean(s["title"]), clean(f"{c[1]}: {c[8]}"), msg,
                                             ", ".join(sorted({x for fr in s.get("frs", [])
                                                               for x in self.b.fr_to_frpm.get(fr, [])}))])
        rows = [[str(i)] + r for i, r in enumerate(rows, 1)]
        self.table(["No.", "Screen", "When", "Message shown", "Requirement"], rows, [0.4, 1.4, 1.75, 2.55, 1.1],
                   size=7.5)
        self.b.stats["messages"] = len(rows)

    # Annex I – notifications
    def notifications(self):
        self.h1("Annex I – Notifications and E-mails")
        self.p("The notifications the system sends: in-system notices (bell in the page header), alerts and e-mails. "
               "E-mails with documents are password protected and the password follows in a separate e-mail "
               "(FRPM.027.01).")
        data = yaml.safe_load((PM / "pack" / "notifications.yaml").read_text(encoding="utf-8"))["notifications"]
        override = {
            "NT-08": "The daily expiry check finds an active package at the expiry alert days before its expiry date "
                     "(90 days, FRPM.004.01) and at each reminder day; once per alert day",
            "NT-11": "The daily check finds package requests past the service level of their stage (FRPM.028.01)",
        }
        rows = []
        for n in data:
            frpm = ", ".join(sorted({x for fr in n.get("frs", []) for x in self.b.fr_to_frpm.get(fr, [])}))
            rows.append([n["id"], clean(n["channel"]), self.refs(override.get(n["id"], n["trigger"])),
                         clean(n["recipient"]), clean(n["template"]), frpm])
        self.table(["No.", "Channel", "Trigger", "Recipient", "Text", "Requirement"], rows,
                   [0.5, 0.7, 1.55, 1.25, 2.05, 1.15], size=7.5)
        self.b.stats["notifications"] = len(rows)

    # Annex J – rules and parameters
    def rules(self):
        j = DOC["annex_j"]
        self.h1("Annex J – Business Rules and Configurable Parameters")
        self.p("The business rules that apply across the Product Maintenance requirements, the values the System "
               "Administrator or the business maintains without a change to the system, and the proposed "
               "role-to-action matrix. Changes to parameters and lists are audited.")
        self.h2("J.1 Business rules")
        self.table(["No.", "Rule", "Requirement"], [[str(i), r[0], r[1]] for i, r in enumerate(j["rules"], 1)],
                   [0.45, 5.0, 1.75])
        self.h2("J.2 Configurable parameters")
        self.table(["Parameter", "Default", "Meaning", "Maintained by"], j["parameters"], [1.8, 1.2, 3.0, 1.2])
        self.h2("J.3 Lists of values")
        self.table(["List", "Values"], j["lists"], [1.8, 5.4])
        self.h2("J.4 Workflow stages and service levels")
        self.table(["Stage", "Owner", "Service level (hours, default)"], j["stages"], [2.6, 3.0, 1.6])
        self.h2("J.5 Proposed role-to-action matrix")
        self.p("Y = the role may perform the action. PMADD05 is specified in the User Access Maintenance FRS (sign-in, "
               "roles and their changes); the matrix below is the proposal for the Product Maintenance actions, for "
               "BDOI to confirm. The segregation of duties of FRPM.028.01 applies whatever the role grants.")
        self.table(["Action"] + j["matrix_roles"], j["matrix"], [2.25] + [0.55] * len(j["matrix_roles"]), size=7.5)
        self.h2("J.6 Masters maintained by the business")
        self.table(["Master", "Maintained by", "Authorized by", "Requirement"], j["masters"], [2.6, 1.6, 1.8, 1.2])

    # Annex K – walkthroughs
    def walkthroughs(self, screen_titles):
        cases = yaml.safe_load((PM / "brd03_cases.yaml").read_text(encoding="utf-8"))
        personas = {k: v["name"] for k, v in cases["personas"].items()}
        data = yaml.safe_load((PM / "pack" / "walkthroughs.yaml").read_text(encoding="utf-8"))["walkthroughs"]
        self.h1("Annex K – Process Walkthroughs")
        self.p("End-to-end walkthroughs of the Product Maintenance process, step by step (storyboard): who acts, on "
               "which screen, what the user does, what the user sees and the result. The screens are specified in "
               "Annex G.")
        total = 0
        for w in data:
            self.h2(f"{w['id']} {clean(w['title'])}")
            self.p(self.refs(" ".join(str(w["summary"]).split())))
            if w.get("data"):
                self.p(clean(w["data"]), italic=True, size=9)
            rows = []
            for i, s in enumerate(w["steps"], 1):
                rows.append([f"{w['id']}.{i}", personas.get(s[0], clean(s[0])),
                             f"{screen_titles.get(s[1], s[1])} ({s[1]})", self.refs(s[2]), self.refs(s[3]), clean(s[4])])
            total += len(rows)
            self.table(["Step", "Persona", "Screen", "What the user does", "What the user sees", "Result"], rows,
                       [0.65, 0.85, 1.1, 1.9, 1.7, 1.0], size=7.5)
            if w.get("downstream"):
                self.p("Afterwards:", bold=True, keep_next=True)
                self.bullets([self.refs(x) for x in w["downstream"]])
        self.b.stats["walk_steps"] = total

    # Annex L – traceability
    def traceability(self, our, nb):
        pm_cases = test_cases(PM / "brd03_cases.yaml", "PM")
        nb_cases = test_cases(NB / "brd01_cases.yaml", "NB")
        self.h1("Annex L – Requirements Traceability")
        self.p("L.1 traces every requirement of the Product Maintenance BRD to the functional requirements of this "
               "FRS (FRPM), to the reference requirements of the BIBS FRS BRD-3 Product Maintenance (FR-PM) and to the "
               "test conditions of the BIBS test plan BRD-3 (TC-PM-nnn.n; the cases of each condition are listed in "
               "the test plan workbook). L.2 lists every reference requirement with the FRPM items that incorporate it.")
        self.h2("L.1 BRD requirement to FRPM, BIBS reference and test cases")
        by_brd: dict[str, list[str]] = {}
        for fr_id, fr in our.items():
            for x in fr["brd"]:
                by_brd.setdefault(x.split(" ")[0], []).append(fr_id)
        rows = []
        for rid in self.ids:
            frpm = ", ".join(re.findall(r"FRPM\.\d{3}(?:\.\d{2})?", self.mapping[rid])) or \
                "Refer to User Access Maintenance FRS"
            refs = by_brd.get(rid, [])
            rows.append([rid, frpm, ", ".join(refs), "; ".join(tc_text(r, pm_cases) for r in refs)])
        self.table(["BRD ID", "FRPM (this FRS)", "BIBS reference (FR-PM)", "Test conditions (BIBS test plan BRD-3)"],
                   rows, [0.9, 3.4, 2.2, 3.4], size=7.5)
        self.h2("L.2 BIBS reference requirement to FRPM")
        rows = []
        for fr_id, fr in list(our.items()) + [(k, nb[k]) for k in sorted(self.b.fr_to_frpm) if k.startswith("FR-NB")]:
            frpm = ", ".join(self.b.fr_to_frpm.get(fr_id, [])) or "Refer to User Access Maintenance FRS"
            cases = pm_cases if fr_id.startswith("FR-PM") else nb_cases
            rows.append([fr_id, clean(fr["title"]), ", ".join(b.split(" ")[0] for b in fr["brd"]), frpm,
                         tc_text(fr_id, cases)])
        self.table(["Reference", "Title", "BRD", "FRPM (this FRS)", "Test conditions"], rows,
                   [0.85, 2.9, 1.5, 2.55, 2.1], size=7.5)
        self.listed_refs = {r[0] for r in rows}

    # Annex M – observations
    def observations(self):
        self.b.obs_rows = []
        cmp_data = yaml.safe_load(CMP.read_text(encoding="utf-8"))
        roles = OBS["roles"]
        self.h1("Annex M – Observations and Points for BDOI Decision")
        self.p("Every point where BDOI's FRS and the BIBS reference FRS differ, BDOI's open item, the changes made to "
               "the Business Requirements Mapping, the slips noticed in BDOI's text and the other observations of "
               "this version. BDOI's text is kept in the body of the document; this annex gives both readings, the "
               "BRD text, what version 1.1 proposes, the impact and who decides. Every point starts as Open.")
        header = ["No.", "Ref.", "Topic", "BDOI FRS text", "BRD text", "Proposed in v1.1", "Impact", "Decision by",
                  "Status"]
        widths = [0.45, 0.5, 0.95, 1.75, 1.45, 2.15, 1.05, 1.0, 0.5]
        no = 0

        def row(ref, topic, bdoi, brd, prop, impact, role_keys, section=None):
            nonlocal no
            no += 1
            who = "; ".join(roles[r] for r in role_keys)
            out = [f"M-{no:02d}", ref, topic, bdoi, brd, prop, impact, who, "Open"]
            self.b.obs_rows.append(out + [section or current[0]])
            return out
        current = ["M.1"]

        groups = [("M.1 Conflicts between BDOI's FRS and the BIBS reference FRS", "C"),
                  ("M.3 Slips noticed in BDOI's FRS (proposed corrections; the text is not changed)", "D")]
        conf = {c["id"]: c for c in cmp_data["conflicts"]}
        self.listed_conflicts = set()
        for title, kind in groups[:1]:
            self.h2(title)
            rows = []
            for cid, c in conf.items():
                if not cid.startswith(kind):
                    continue
                o = OBS["conflicts"][cid]
                brd = c["brd_note"] if c["brd_note"] not in ("-", "") else "Not stated in the BRD."
                rows.append(row(cid, c["topic"], c["bdoi"], brd, o["proposes"], o["impact"], o["role"]))
                self.listed_conflicts.add(cid)
            self.table(header, rows, widths, size=7)
        self.h2("M.2 BDOI's open item and changes to the Business Requirements Mapping")
        current[0] = "M.2"
        rows = []
        for o in OBS["others"][:1]:
            rows.append(row(o["ref"], o["topic"], o["bdoi"], o["brd"], o["proposes"], o["impact"], o["role"]))
        for ch in self.b.log:
            if ch["id"] == "BRPM.020":
                continue
            rows.append(row(ch["id"], "Mapping extended", f"Mapped to: {ch['before']}", ch["brd"],
                            f"Mapped to: {ch['after']}. {ch['why']}", "Coverage of the BRD requirement.",
                            ["SA", "PO"]))
        self.table(header, rows, widths, size=7)
        title, kind = groups[1]
        self.h2(title)
        current[0] = "M.3"
        rows = []
        for cid, c in conf.items():
            if cid.startswith(kind):
                o = OBS["conflicts"][cid]
                rows.append(row(cid, c["topic"], c["bdoi"], "Not stated in the BRD.", o["proposes"], o["impact"],
                                o["role"]))
                self.listed_conflicts.add(cid)
        self.table(header, rows, widths, size=7)
        self.h2("M.4 Scope, added requirements and added annex rows")
        current[0] = "M.4"
        rows = [row(o["ref"], o["topic"], o["bdoi"], o["brd"], o["proposes"], o["impact"], o["role"])
                for o in OBS["others"][1:]]
        self.table(header, rows, widths, size=7)
        v12.observations_m5(self, row, header, widths)
        self.b.stats["observations"] = no
        self.b.stats["conflicts"] = len([c for c in self.listed_conflicts if c.startswith("C")])
        self.b.stats["slips"] = len([c for c in self.listed_conflicts if c.startswith("D")])

    # Annex N – glossary
    def glossary(self):
        self.h1("Annex N – Glossary")
        self.p("Terms and abbreviations used in this document.")
        self.table(["Term", "Meaning"], sorted(DOC["glossary"], key=lambda g: g[0].lower()), [1.8, 5.4], size=9)
        self.b.stats["glossary"] = len(DOC["glossary"])


def section_break(b: Builder, landscape: bool):
    """A paragraph that ends a section; the section takes the page set-up of the document's last section."""
    sect = copy.deepcopy(b.doc.sections[-1]._sectPr)
    for t in sect.findall(qn("w:titlePg")):
        sect.remove(t)
    pg = sect.find(qn("w:pgSz"))
    w, h = int(pg.get(qn("w:w"))), int(pg.get(qn("w:h")))
    if landscape:
        pg.set(qn("w:w"), str(max(w, h)))
        pg.set(qn("w:h"), str(min(w, h)))
        pg.set(qn("w:orient"), "landscape")
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(sect)
    p.append(ppr)
    return p


def add_annexes(b: Builder, mapping, ids, our, nb):
    a = Annexes(b, mapping, ids)
    titles = a.screens()
    v12.menu_page(a)
    a.messages(our, nb, titles)
    a.notifications()
    a.rules()
    a.walkthroughs(titles)
    a.els.append(section_break(b, landscape=False))
    a.traceability(our, nb)
    a.observations()
    a.els.append(section_break(b, landscape=True))
    a.glossary()
    v12.glossary_add(a)
    v12.annexes(a)
    body_sect = b.body.find(qn("w:sectPr"))
    for el in a.els:
        body_sect.addprevious(el)
    return a


def change_summary(b: Builder):
    """The change summary page at the front of the review copy: the changes of version 1.2 (this round)."""
    s, oi = b.stats, v12.open_items()
    points = [
        "Cover and Document Revision Log: version 1.2, Business Unit review edition; BDOI's details, signatories and "
        "the entries of versions 1.0 and 1.1 unchanged.",
        f"Introduction: new section 1.4 'Summary for the Business Unit Review' (scope on a page, the process end to "
        f"end, key numbers, the ten decisions for the Business Unit, how to review, the {s['checklist']} points of "
        "the review checklist); a note on what version 1.2 adds.",
        f"Process Flow: a caption under every figure; {s['flows_v12']} process flows, {s['lifecycles']} status "
        "life-cycles and the integration context diagram added (Graphviz, by persona, with service levels).",
        f"Annex G: section G.23 'Navigation: menu by persona' ({s['menu_personas']} personas, {s['menu_entries']} "
        f"menu entries). Annex M: section M.5 with {s['observations_v12']} observations of version 1.2. Annex N: "
        f"{s['glossary_v12']} terms added.",
        f"New annexes: O Workflow and Approvals ({s['workflow_processes']} processes, {s['workflow_rows']} stages, "
        f"role-to-stage matrix, status names); P E-mail and Notification Texts ({s['notices']}: "
        f"{s['notices_bibs']} as sent by BIBS, {s['notices_proposed']} proposed for approval); Q Document Prints and "
        f"Output Formats ({s['documents']} documents, {s['doc_images']} current layouts); R Reports and Schedules "
        f"({s['reports']} reports, {s['schedules']} scheduled runs); S Integrations ({s['integrations']}); "
        f"T Non-functional Requirements ({s['nfr']}); U Data Set-up and Migration at Go-live ({s['data_setup']}); "
        f"V Assumptions, Dependencies and Open Questions ({s['open_kept']} items); W Change Control after "
        f"Sign-off; X Business Unit Review Checklist ({s['checklist']} points).",
        f"Open items: {len(oi['answered'])} assumptions, open questions and proposed rules of the BIBS reference FRS "
        f"are answered by BDOI's FRS (or were closed earlier) and are no longer asked; {len(oi['partly'])} are "
        "answered in part and only the remaining part is asked (Annex V).",
    ]
    els = [b.para("Version 1.2 – summary of changes for the Business Unit and BDO ITG review", bold=True, size=14,
                  color="014EA9", hl=False, space_after=120),
           b.para("This review copy is BDOI's Functional Requirements Specifications for Product Maintenance with the "
                  "changes of version 1.2 (Business Unit review edition). The passages added in version 1.2 are "
                  "shaded light yellow; BDOI's text and the content of version 1.1 are not shaded and are "
                  "unchanged. The clean copy has the same content without shading and without this page.",
                  size=9.5, hl=False)]
    n = b.new_num()
    for pt in points:
        els.append(b.para(pt, style="ListParagraph", num=(n, 0), size=9.5, hl=False))
    els.append(b.para("Where BDOI's FRS and the BIBS reference FRS differ, BDOI's text is kept and the point is listed "
                      "in Annex M with both readings and a recommendation.", size=9.5, hl=False))
    els.append(b.para("Shading used in this copy: ", size=9.5, hl=False))
    els[-1].append(b.run("text added in version 1.2", size=9.5, hl=True))
    # the summary is a section of its own, with the page set-up of the cover, so the cover keeps its first page
    brk = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(copy.deepcopy(b.doc.sections[0]._sectPr))
    brk.append(ppr)
    els.append(brk)
    first = b.body[0]
    for e in els:
        first.addprevious(e)


def set_update_fields(b: Builder):
    settings = b.doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    u = OxmlElement("w:updateFields")
    u.set(qn("w:val"), "true")
    later = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
             "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
             "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
             "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
    for child in settings:
        if child.tag.split("}")[1] in later:
            child.addprevious(u)
            return
    settings.append(u)


# ------------------------------------------------------------------------------------------------------- self-checks
def texts_of(cell_or_tbl) -> list[str]:
    out = []
    for p in cell_or_tbl.iter(qn("w:p")):
        t = "".join(x.text or "" for x in p.iter(qn("w:t"))).strip()
        if t:
            out.append(t)
    return out


def check_bdoi_text(new_doc):
    """BDOI's FRPM IDs and texts are all still present: in every table of the source, the paragraphs of each cell are
    found, in order, in the same cell of the new document (the mapping of BRPM.020, the empty revision log row and
    the cover version are the only planned changes)."""
    src = docx.Document(str(REPO / DOC["meta"]["source"]))
    ids_src = set(re.findall(r"FRPM\.\d{3}(?:\.\d{2})?", "\n".join(texts_of(src.element.body))))
    ids_new = set(re.findall(r"FRPM\.\d{3}(?:\.\d{2})?", "\n".join(texts_of(new_doc.element.body))))
    missing_ids = ids_src - ids_new
    problems = [f"FRPM ID missing: {sorted(missing_ids)}"] if missing_ids else []
    # BDOI's tables in the new document, in order: tables added in between (summary, annexes) are skipped
    new_tables, k = [], 0
    for ts in src.tables:
        head = [texts_of(c._tc) for c in ts.rows[0].cells]
        while k < len(new_doc.tables):
            tn = new_doc.tables[k]
            k += 1
            cand = [texts_of(c._tc) for c in tn.rows[0].cells] if len(tn.rows) >= len(ts.rows) else None
            if cand is not None and len(cand) == len(head) and all(
                    all(any(o.replace(DOC["meta"]["version_old"], DOC["meta"]["version_new"]) == n for n in c)
                        for o in h) for h, c in zip(head, cand)):
                new_tables.append(tn)
                break
        else:
            raise SystemExit("BDOI text check failed: a table of BDOI's FRS is missing")
    for ti, (ts, tn) in enumerate(zip(src.tables, new_tables)):
        for ri, row in enumerate(ts.rows):
            for ci, cell in enumerate(row.cells):
                old = texts_of(cell._tc)
                new = texts_of(tn.rows[ri].cells[ci]._tc)
                if ti == 0:
                    old = [x.replace(DOC["meta"]["version_old"], DOC["meta"]["version_new"]) for x in old]
                if ti == 2 and ci == 2 and "Under negotiation" in " ".join(old):
                    continue
                it = iter(new)
                if not all(any(o == n for n in it) for o in old):
                    problems.append(f"table {ti} row {ri} cell {ci}: BDOI text changed")
    body_src = [p.text.strip() for p in src.paragraphs if p.text.strip()]
    body_new = [p.text.strip() for p in new_doc.paragraphs if p.text.strip()]
    it = iter(body_new)
    lost = [o for o in body_src if not any(o == n for n in it)]
    if lost:
        problems.append(f"body paragraphs changed: {lost[:3]}")
    if problems:
        raise SystemExit("BDOI text check failed:\n  " + "\n  ".join(problems))
    return len(ids_src)


def check_words(b: Builder):
    hits = set()
    for t in b.added_text:
        for m in list(RESTRICTED.finditer(t)) + list(RESTRICTED_CS.finditer(t)):
            hits.add(m.group(0))
        for m in CODE_RE.finditer(t):
            if m.group(0) not in CODE_OK:
                hits.add(m.group(0))
    if hits:
        raise SystemExit(f"restricted words or internal codes in the added text: {sorted(hits)}")


# ------------------------------------------------------------------------------------------------------------ PDF
def to_pdf(src: Path, dest: Path, timeout=600):
    """The clean copy as PDF, with the table of contents updated (LibreOffice through UNO)."""
    import uno  # noqa: PLC0415
    from com.sun.star.beans import PropertyValue  # noqa: PLC0415

    def prop(name, value):
        p = PropertyValue()
        p.Name, p.Value = name, value
        return p

    profile = Path(tempfile.mkdtemp(prefix="pm_lo_"))
    pipe = f"pmfrs{int(time.time())}"
    proc = subprocess.Popen(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                             "--invisible", f"--accept=pipe,name={pipe};urp;"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        ctx = None
        resolver = uno.getComponentContext().ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", uno.getComponentContext())
        for _ in range(120):
            try:
                ctx = resolver.resolve(f"uno:pipe,name={pipe};urp;StarOffice.ComponentContext")
                break
            except Exception:  # noqa: BLE001
                time.sleep(1)
        if ctx is None:
            raise RuntimeError("LibreOffice did not start")
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        model = desktop.loadComponentFromURL(uno.systemPathToFileUrl(str(src)), "_blank", 0, (prop("Hidden", True),))
        idx = model.getDocumentIndexes()
        for _ in range(2):
            for i in range(idx.getCount()):
                idx.getByIndex(i).update()
            model.refresh()
        model.storeToURL(uno.systemPathToFileUrl(str(dest)), (prop("FilterName", "writer_pdf_Export"),))
        model.close(True)
        try:
            desktop.terminate()
        except Exception:  # noqa: BLE001
            pass
    finally:
        try:
            proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)


# ----------------------------------------------------------------------------------------------------------- main
def key_numbers(b: Builder, ids) -> list[list[str]]:
    s, oi = b.stats, v12.open_items()
    items = v12.frpm_items()
    n_items = len({i["item"] for i in items})
    n_bdoi = sum(1 for i in items if i["origin"].startswith("BDOI"))
    return [
        ["BRD requirement IDs (PMADD01 to PMADD08, BRPM.001 to BRPM.024)", str(len(ids)), "Section 2"],
        ["Functional requirement items and sub-items", f"{n_items} items; {len(items)} sub-items ({n_bdoi} of BDOI, "
         f"{len(items) - n_bdoi} added)", "Section 3"],
        ["Acceptance criteria", str(s["ac"]), "Section 3"],
        ["Process flows, life-cycles and other figures", f"{s['figures_total']} figures", "Section 1.4, section 4"],
        ["Screens", f"{s['screens']} (menu of {s['menu_personas']} personas)", "Annex G"],
        ["Messages and validations", str(s["messages"]), "Annex H"],
        ["E-mails and notices", f"{s['notices']} ({s['notices_bibs']} as sent by BIBS, {s['notices_proposed']} "
         "proposed)", "Annex P"],
        ["Documents printed, e-mailed or exported", str(s["documents"]), "Annex Q"],
        ["Reports and scheduled runs", f"{s['reports']} reports; {s['schedules']} runs (weekly Thursday 08:00)",
         "Annex R"],
        ["Integrations", str(s["integrations"]), "Annex S, Figure " + str(b.integration_fig)],
        ["Non-functional requirements", str(s["nfr"]), "Annex T"],
        ["Data set-up items at go-live", str(s["data_setup"]), "Annex U"],
        ["Open items for BDOI (not answered in BDOI's FRS)", f"{s['open_kept']} (another {len(oi['answered'])} "
         "answered by BDOI's FRS)", "Annex V"],
        ["Observations and points for decision", str(s["observations"]), "Annex M"],
        ["Test cases (traceability workbook)", str(len(v12.test_cases())), "BIBS_RTM_BRD-03 workbook"],
        ["Review checklist points", str(s["checklist"]), "Annex X"],
    ]


def build(highlight: bool, our, nb, ids) -> Builder:
    b = Builder(highlight)
    b.fig_no = 1  # Figure 1 is the end-to-end picture of the summary in the Introduction
    edit_cover(b)
    edit_revision_log(b)
    edit_introduction(b)
    with v12.this_round(b):
        v12.introduction(b)
    mapping = edit_mapping(b, ids)
    edit_functional_requirements(b)
    edit_process_flows(b)
    with v12.this_round(b):
        v12.process_flows(b)
    edit_signoff(b)
    edit_annex_rows(b)
    a = add_annexes(b, mapping, ids, our, nb)
    with v12.this_round(b):
        v12.summary(b, key_numbers(b, ids), b.stats["checklist"])
    missing_fr = [f for f in our if f not in a.listed_refs]
    if missing_fr:
        raise SystemExit(f"Annex L lacks {missing_fr}")
    cmp_ids = {c["id"] for c in yaml.safe_load(CMP.read_text(encoding="utf-8"))["conflicts"]}
    if cmp_ids - a.listed_conflicts:
        raise SystemExit(f"Annex M lacks {sorted(cmp_ids - a.listed_conflicts)}")
    unknown = [r for r in b.fr_to_frpm if r.startswith("FR-PM") and r not in our]
    if unknown:
        raise SystemExit(f"unknown BIBS references: {unknown}")
    if b.fr_to_frpm != v12.fr_to_frpm():
        raise SystemExit("the FRPM references of the document and of the shared data differ")
    if highlight:
        change_summary(b)
    set_update_fields(b)
    check_words(b)
    problems = v12.self_checks(b, " ".join(b.added_text))
    if problems:
        raise SystemExit("version 1.2 checks failed:\n  " + "\n  ".join(problems))
    return b


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true", help="skip the PDF")
    args = ap.parse_args(argv)
    our, nb = our_frs(), nb_frs()
    ids = brd_ids(our)
    OUT.mkdir(parents=True, exist_ok=True)
    results = {}
    for highlight, name in ((False, DOC["meta"]["clean"]), (True, DOC["meta"]["highlighted"])):
        b = build(highlight, our, nb, ids)
        path = OUT / name
        b.doc.save(str(path))
        v12.cleanup(b)
        n_ids = check_bdoi_text(docx.Document(str(path)))
        results[name] = (b, n_ids)
        print(f"wrote {path.relative_to(REPO)} ({path.stat().st_size // 1024} KB)")
    b, n_ids = results[DOC["meta"]["clean"]]
    if not args.no_pdf:
        pdf = OUT / DOC["meta"]["clean"].replace(".docx", ".pdf")
        to_pdf(OUT / DOC["meta"]["clean"], pdf)
        print(f"wrote {pdf.relative_to(REPO)} ({pdf.stat().st_size // 1024} KB)")
    s = b.stats
    print(f"BRD IDs mapped: {s['brd_ids']} ({s['brd_uam']} by BDOI's reference to the User Access Maintenance FRS); "
          f"mapping rows changed: {s['mapping_changes']}")
    print(f"FRPM: {s['bdoi_items']} items and {s['bdoi_subitems']} sub-items of BDOI kept ({n_ids} FRPM IDs checked); "
          f"{s['elaborated']} sub-items elaborated; {s['new_subitems']} new sub-items; {s['new_items']} new items "
          f"({s['new_item_subitems']} sub-items); {s['requirements_added']} requirements and {s['ac']} acceptance "
          f"criteria added")
    print(f"Annexes: rows added to A-E {s['annex_rows']}; G {s['screens']} screens ({s['screenshots']} screenshots); "
          f"H {s['messages']} messages; I {s['notifications']} notifications; K {s['walk_steps']} steps; "
          f"M {s['observations']} observations ({s['conflicts']} conflicts, {s['slips']} slips); N {s['glossary']} "
          f"terms; process flows added {s['flows']}")
    print(f"v1.2: figures {s['figures_total']} ({s['flows_v12']} flows, {s['lifecycles']} life-cycles added); G.23 "
          f"menu {s['menu_personas']} personas; M.5 {s['observations_v12']}; O {s['workflow_processes']} processes "
          f"{s['workflow_rows']} stages; P {s['notices']} ({s['notices_bibs']} BIBS, {s['notices_proposed']} "
          f"proposed); Q {s['documents']} ({s['doc_images']} images); R {s['reports']} reports, {s['schedules']} "
          f"runs; S {s['integrations']}; T {s['nfr']}; U {s['data_setup']}; V {s['open_kept']} kept, "
          f"{s['open_answered']} answered, {s['open_partly']} partly; X {s['checklist']}; "
          f"test cases {len(v12.test_cases())}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
