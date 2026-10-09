"""Operations Cashiering FRS in BDOI's template, version 3.2 (full BRD coverage and Business Unit review edition).

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd02_bdoi_frs.py [--no-pdf]

Starts from a copy of BDOI's own FRS (docs/source-documents/FRS - BDOI Core Replacement Operations Cashiering v3.1
(BDOI).docx, BDO ITG template, PRJ0012475) and keeps its template: styles, cover, headers and footers, table formats,
fonts, review comments and the table of contents (marked to update when the document is opened). BDOI's text is kept
word for word; the additions of version 3.2 come from:

  brd02_frs_additions.yaml   statements and acceptance criteria under BDOI's items, new sub-items, new items
  brd02_document.yaml        cover, revision log, introduction, mapping additions, Appendices D, H and I, slips and
                             observations of Appendix R, glossary
  brd02_review.yaml          the Business Unit review edition: summary (1.4), Appendices G and J to U, walkthrough WT-E
  brd02_flows.yaml           the figures (drawn by brd11_figures.render)
  ../comparisons/brd02_cashiering_comparison.yaml   the conflicts of the comparison workbook (Appendix R)

and from the BIBS release set BRD-02 (FRS v2.1, screens, screenshots, notifications, walkthroughs, test cases), its
sign-off workbook (messages, menus) and the Operations BRD (docs/source-documents/Operations_WS Addendum.pdf), read at
every build. The appendices are written by brd02_annexes.py.

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-02_Operations_Cashiering/:
  BIBS_FRS-BDOI_BRD-02_Operations_Cashiering_v3.2.docx                      clean copy
  BIBS_FRS-BDOI_BRD-02_Operations_Cashiering_v3.2.pdf                       the clean copy as PDF (not kept in git)
  BIBS_FRS-BDOI_BRD-02_Operations_Cashiering_v3.2_Changes_Highlighted.docx  every addition of version 3.2 shaded light
                                                                            yellow, change summary in front

Self-checks (the build fails when one does not hold): every BRD ID of the Cashiering scope of the BRD is in the
Business Requirements Mapping and maps to an FR item, or is marked removed or by reference; every FR-OP of the scope
is in Appendix Q; every conflict of the comparison is in Appendix R; BDOI's FR IDs, table texts and paragraphs are
all still present; every notice of the platform for Cashiering is in Appendix G and its text is found in the
platform; every report of FRS.CSH.09.02 is in Appendix M; every open item of our sources is classified once and none
that BDOI's FRS answers is kept in Appendix S; no internal code or restricted word in the added text.
"""

from __future__ import annotations

import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, OrderedDict
from pathlib import Path

import docx
import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import brd11_figures  # noqa: E402
import build_brd11_bdoi_frs as u11  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

REPO = HERE.parents[4]
SRC = REPO / "docs" / "deliverables" / "src"
OPS = SRC / "BRD-02_Operations"
CMP = SRC / "programme" / "comparisons" / "brd02_cashiering_comparison.yaml"
ADD = yaml.safe_load((HERE / "brd02_frs_additions.yaml").read_text(encoding="utf-8"))
DOC = yaml.safe_load((HERE / "brd02_document.yaml").read_text(encoding="utf-8"))
REV = yaml.safe_load((HERE / "brd02_review.yaml").read_text(encoding="utf-8"))
FLOWS = yaml.safe_load((HERE / "brd02_flows.yaml").read_text(encoding="utf-8"))["figures"]
OUT = REPO / "docs" / "deliverables" / "out" / DOC["meta"]["out_folder"]
SIGNOFF_XLSX = (REPO / "docs" / "deliverables" / "out" / "Drop-1_Transactional" / "BRD-02_Operations" /
                "03_BIBS_Signoff_BRD-02_Operations_v2.1.xlsx")

ID_RE = re.compile(r"^\s*(FRS\s*\.\s*(?:CSH|OPS)[\d .]*\d)")
FRS_RE = re.compile(r"FRS\.(?:CSH|OPS)\.\d{2,3}(?:\.\d{1,2})*")
BRD_ID_RE = re.compile(r"\b(?:BRQID|CSHID|MKTID|DBMID)\.\d{3}\b")
# BRD IDs of the Cashiering scope: the IDs of BDOI's mapping and the IDs the BRD removed
SCOPE_PREFIXES = ("BRQID.", "CSHID.")
BY_REFERENCE = set(DOC["by_reference"])
REMOVED = {r["id"] for r in DOC["removed"]}
# FR-OP of the BIBS FRS BRD-2 that the Cashiering scope traces to (each must be incorporated in an FR item)
SCOPE_FRS = ["FR-OP-%03d" % n for n in (1, 2, 3, 4, 5, 8, 9, 10, 11, 12, 13, 14, 15, 16, 18, 19, 20, 21, 22, 23, 24,
                                        25, 27, 28, 113, 120, 121, 130, 131, 132)]
CALIBRI = "Calibri"
# Values of BDOI's own FRS that look like codes (the certificate number of the sample AR)
BDOI_CODES = {"AC_125_022025_00482"}

# ------------------------------------------------------------------------------------------------ wording
CODES = {
    "AR_CLASS": "the list of AR classes", "OR_TYPE": "the list of OR types",
    "RECEIPT_CANCEL_REASON": "the list of cancellation reasons", "REINSTATEMENT_REASON": "the list of reinstatement reasons",
    "DISPOSITION_TYPE": "the list of disposition types", "CASH_APPROVE": "the approval permission",
    "CASH_RECEIPT": "the receipt permission", "CASH_UPLOAD": "the upload permission", "CASH_CANCEL": "the cancellation permission",
    "CASH_REINSTATE": "the reinstatement permission", "CASH_DISPOSITION": "the disposition permission",
    "CASH_DISPOSITION_APPROVE": "the disposition approval permission", "CASH_SERIES_MANAGE": "the receipt series permission",
    "CASH_PRINT": "the print permission", "CASH_APPLY": "the application permission", "OPS_VIEW": "the Operations view permission",
    "OPS_REPORT_VIEW": "the report view permission", "OPS_REPORT_EXPORT": "the report export permission",
    "FLOWIN_MANAGE": "the interface permission", "DISB_PROCESS": "the Disbursement permission",
    "CWT_APPLICATION_PERCENT": "the CWT share setting", "MIN_BALANCE_AUTO_MAX": "the minimal balance threshold",
    "PREBOOKED_AGEING": "pre-booked ageing", "RECEIPT_SERIES_LOW": "receipt series low", "OPS_EXTERNAL_LINK": "the list of links",
    "PAY_BILLS": "Bills Payment", "PAY_TRADE": "Trade", "PAY_CLPC": "CLPC", "PAY_PDC": "post-dated checks",
    "PAY_DIRECT_CREDIT": "Direct Credit", "FIXED_WIDTH": "fixed width", "AUTO": "automatic", "DELIMITED": "delimited",
    "CASHIER_TL": "Cashiering Team Leader", "CASHIER": "Cashier", "MKT_COLLECTION": "Marketing Collection",
    "OPS_AR_RECEIPT": "the AR entry", "OPS_PAYMENT_APPLY": "the application entry", "OPS_OR_ISSUE": "the OR entry",
    "OPS_RECEIPT_REINSTATE": "the reinstatement entry", "OPS_EXCESS_TO_OVERAGES": "the overages entry",
    "OPS_MINIMAL_BALANCE_REVERSAL": "the minimal balance entry", "OPS_UNAPPLIED_REFUND": "the refund entry",
    "OPS_UNAPPLIED_RECLASS": "the reclass entry", "OPS_AR_INSURANCE_RECEIPT": "the AR Insurance entry",
    "OPS_RECEIPT_ACTION": "the receipt action workflow", "OPS_DISPOSITION": "the disposition workflow",
    "CANCELLED_REFERENCE": "Cancelled reference", "UNAPPLIED_NO_MATCH": "No match", "MIN_BAL": "minimal balance",
    "COLLECTION_COMMISSION_PAYMENT": "commission payment details", "COMMISSION_PAYMENT": "commission payment details",
    "AR_INSURANCE": "AR Insurance", "NOT_APPLICABLE": "not applicable", "PARTIALLY_PAID": "partially paid",
    "WITH_OUTSTANDING_BALANCE": "with outstanding balance", "FOR_APPROVAL": "For Approval",
}
CODE_RE = re.compile(r"\b[A-Z][A-Z0-9]*_[A-Z0-9][A-Z0-9_]*\b")
PHRASES = u11.PHRASES + [(re.compile(r"\b(?:SCR-OP-\d\d|FR-OP-\d{3}|CLR-OP-\d\d|OQ\d\d|IQ\d\d)\b"), "")]


def clean(text, keep_refs=False) -> str:
    """Business wording of a text taken from the BIBS sources: no internal codes."""
    s = "" if text is None else str(text)
    for code, word in sorted(CODES.items(), key=lambda kv: -len(kv[0])):
        s = re.sub(r"\b" + re.escape(code) + r"\b", word, s)
    s = CODE_RE.sub(lambda m: m.group(0).replace("_", " ").lower(), s)
    for pat, rep in (u11.PHRASES if keep_refs else PHRASES):
        s = pat.sub(rep, s)
    for _ in range(3):
        s = re.sub(r"([,;])\s*[,;]", r"\1", s)
        s = re.sub(r"\(\s*[,;]\s*", "(", s)
        s = re.sub(r"\s*[,;]\s*\)", ")", s)
        s = re.sub(r"\(\s*\)", "", s)
    s = re.sub(r"\s+([,.;:])", r"\1", s)
    return re.sub(r"\s{2,}", " ", s).strip()


def norm_id(raw: str) -> str:
    return re.sub(r"\s+", "", raw).rstrip(".")


def frs_key(fid: str):
    parts = re.findall(r"\d+", fid.replace(" (2nd)", ""))
    head = 0 if ".OPS." in fid else 1
    return [head] + [int(x) for x in parts] + [1 if "(2nd)" in fid else 0]


# ------------------------------------------------------------------------------------------------ sources
def our_frs() -> "OrderedDict[str, dict]":
    text = (OPS / "FRS_BRD02_OPERATIONS.md").read_text(encoding="utf-8")
    out = OrderedDict()
    for block in re.findall(r"^```fr\n(.*?)^```", text, re.S | re.M):
        fr = yaml.safe_load(block)
        out[fr["id"]] = fr
    return out


def our_trace() -> dict[str, dict]:
    """BRD ID -> page and FR-OP of the traceability chapter of the BIBS FRS BRD-2."""
    text = (OPS / "FRS_BRD02_OPERATIONS.md").read_text(encoding="utf-8")
    part = text.split("# Traceability", 1)[1].split("# Navigation", 1)[0]
    out = {}
    for line in part.splitlines():
        c = [x.strip() for x in line.strip().strip("|").split("|")]
        if len(c) >= 5 and BRD_ID_RE.match(c[0]):
            rid = BRD_ID_RE.match(c[0]).group(0)
            out[rid] = {"page": c[1], "frs": re.findall(r"FR-OP-\d{3}", c[2]),
                        "removed": "removed" in c[0].lower() or "removed" in c[2].lower()}
    return out


def brd_ids() -> list[str]:
    """The BRD IDs of the Cashiering scope found in the Operations BRD (BRQID, CSHID) and the two IDs of BDOI's
    mapping held by reference."""
    work = Path(tempfile.mkdtemp(prefix="cshbrd_"))
    try:
        txt = work / "brd.txt"
        subprocess.run(["pdftotext", "-layout", str(REPO / DOC["meta"]["brd_pdf"]), str(txt)], check=True)
        text = txt.read_text(encoding="utf-8")
    finally:
        shutil.rmtree(work, ignore_errors=True)
    found = set(re.findall(r"\bCSHID\.\d{3}\b", text))
    found |= {"BRQID.%03d" % int(m) for m in re.findall(r"BRQID\.00\s+(\d)\s", text)}
    found |= {m for m in re.findall(r"\bBRQID\.\d{3}\b", text)}
    for rid in BY_REFERENCE:
        if rid not in text:
            raise SystemExit(f"{rid} not found in the BRD")
        found.add(rid)
    return sorted(found, key=lambda r: (0 if r.startswith("BRQ") else 1 if r.startswith("CSH") else 2, r))


def test_cases() -> dict[str, tuple[list[str], int]]:
    data = yaml.safe_load((OPS / "brd02_cases.yaml").read_text(encoding="utf-8"))
    out = {}
    for fr_id, entry in (data.get("frs") or {}).items():
        n = fr_id.split("-")[-1]
        conds = [f"TC-OP-{n}.{i}" for i in range(1, len(entry.get("conditions") or []) + 1)]
        out[fr_id] = (conds, len(entry.get("cases") or []))
    return out


def tc_text(fr_id, cases) -> str:
    conds, n = cases.get(fr_id, ([], 0))
    if not conds:
        return "-"
    return (f"{conds[0]} to {conds[-1].split('-')[-1]}" if len(conds) > 1 else conds[0]) + f" ({n} cases)"


def all_added_items() -> list[dict]:
    """Every FR item added or elaborated, with its refs: key, row, kind."""
    out = []
    for key, add in ADD["items"].items():
        out.append({"id": key, "refs": add["refs"], "kind": "elaborated", "ac": add["ac"], "shall": add["shall"]})
    for ns in ADD["new_subitems"]:
        out.append({"id": ns["id"], "refs": ns["refs"], "kind": "new sub-item", "row": ns["row"], "ac": ns["ac"],
                    "shall": ns["shall"], "title": ns["title"], "brd": ns.get("brd", [])})
    for item in ADD["new_items"]:
        for si in item["subitems"]:
            out.append({"id": si["id"], "refs": si["refs"], "kind": "new item", "row": item["id"], "ac": si["ac"],
                        "shall": si["shall"], "title": si["title"], "brd": item.get("brd", [])})
    return out


def fr_to_items() -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for it in all_added_items():
        for r in it["refs"]:
            out.setdefault(r, [])
            if it["id"] not in out[r]:
                out[r].append(it["id"])
    return {k: sorted(v, key=frs_key) for k, v in out.items()}


# ------------------------------------------------------------------------------------------------ the builder
class Builder(u11.Builder):
    """The BRD-11 builder of the BDOI-format documents, on BDOI's Cashiering FRS (fonts and lists of that document).
    Every addition of version 3.2 is shaded in the review copy."""

    BULLET_NUM = 19     # bullet list of BDOI's FR table
    DECIMAL_NUM = 6     # decimal list of BDOI's document

    def __init__(self, highlight: bool):  # noqa: D107  (the parent's init reads the BRD-11 source)
        self.highlight = highlight
        self.round13 = True
        self.doc = docx.Document(str(REPO / DOC["meta"]["source"]))
        self.body = self.doc.element.body
        self.tables = list(self.doc.tables)
        self.numbering = self.doc.part.numbering_part.element
        self.stats: Counter = Counter()
        self.log: list[dict] = []
        self.fr_to_frum = fr_to_items()
        self.added_text: list[str] = []
        self.q_rows: list[dict] = []

    def run(self, text, bold=False, italic=False, size=None, font=CALIBRI, color=None, hl=None):
        return super().run(text, bold=bold, italic=italic, size=size, font=font, color=color, hl=hl)

    def para(self, text="", style=None, num=None, bold=False, italic=False, size=None, font=CALIBRI, jc=None,
             keep_next=False, page_break_before=False, no_num=False, color=None, hl=None, space_after=None):
        return super().para(text, style=style, num=num, bold=bold, italic=italic, size=size, font=font, jc=jc,
                            keep_next=keep_next, page_break_before=page_break_before, no_num=no_num, color=color,
                            hl=hl, space_after=space_after)

    def new_num(self, like_num_id: int = BULLET_NUM) -> int:
        return super().new_num(like_num_id)

    def statements(self, entries, size=11):
        n = self.new_num(self.BULLET_NUM)
        out = []
        for e in entries:
            out.append(self.para(e, style="ListParagraph", num=(n, 0), size=size))
            self.stats["requirements_added"] += 1
        return out

    def acceptance(self, ac, size=11):
        out = [self.para("Acceptance criteria", bold=True, italic=True, keep_next=True, size=size)]
        n = self.new_num(self.DECIMAL_NUM)
        for a in ac:
            out.append(self.para(a, style="ListParagraph", num=(n, 0), size=size))
        self.stats["ac"] += len(ac)
        return out


def find_par(b: Builder, start: str, style: str | None = None):
    for p in b.doc.paragraphs:
        if p.text.strip().startswith(start) and (style is None or p.style.name == style):
            return p
    raise SystemExit(f"paragraph not found: {start}")


def plain_par(b: Builder, model_p, text):
    """A paragraph with the formatting of one of BDOI's paragraphs (first run's font kept)."""
    return b.clone_par(model_p, text)


# ------------------------------------------------------------------------------------------------ BDOI's sections
def edit_cover(b: Builder):
    cell = b.tables[0].rows[0].cells[0]
    for p in cell.paragraphs:
        if "Prepared by" in p.text:
            new = b.clone_par(p._p, DOC["cover"]["drafting"])
            p._p.addnext(new)
            return
    raise SystemExit("cover: 'Prepared by' not found")


def edit_revision_log(b: Builder):
    t = b.tables[1]
    empty = next((r for r in t.rows[1:] if not r.cells[0].text.strip()), None)
    if empty is None:
        raise SystemExit("revision log: no empty row")
    template = t.rows[2]._tr
    tr = b.clone_row(template, DOC["revision_row"])
    empty._tr.addprevious(tr)
    empty._tr.getparent().remove(empty._tr)


def edit_introduction(b: Builder):
    p = find_par(b, "The System Analyst or equivalent BDO ITG personnel")
    p._p.addnext(b.clone_par(p._p, DOC["introduction"]["preparation"]))
    obj = find_par(b, "Strengthen compliance and governance")
    model = find_par(b, "This Functional Requirements Specifications (FR) document outlines")._p
    obj._p.addnext(b.clone_par(model, DOC["introduction"]["overview"]))


def edit_references(b: Builder):
    p = find_par(b, "- Operations_WS Addendum")
    els = [b.clone_par(p._p, DOC["introduction"]["references_title"])]
    els += [b.clone_par(p._p, r) for r in DOC["introduction"]["references"]]
    els.append(b.clone_par(p._p, "- " + REV["meta"]["references_review"]))
    b.insert_after(p._p, els)
    return els[-1]


def edit_mapping(b: Builder, ids: list[str]) -> dict[str, str]:
    t = b.tables[2]
    rows = OrderedDict()
    for row in t.rows[1:]:
        rows[row.cells[0].text.replace("\n", "").strip()] = row
    for rid, change in DOC["mapping"].items():
        row = rows[rid]
        cell = row.cells[2]._tc
        before = re.sub(r"\s+", " ", row.cells[2].text).strip() or "(blank)"
        paras = cell.findall(qn("w:p"))
        model = next((p for p in paras if p.find(qn("w:r")) is not None), paras[0])
        text = "Added in v3.2: " + ", ".join(change["add"])
        if not row.cells[2].text.strip():
            for p in paras[1:]:
                cell.remove(p)
            paras[0].addprevious(b.clone_par(model, text))
            cell.remove(paras[0])
        else:
            paras[-1].addnext(b.clone_par(model, text))
        b.log.append({"id": rid, "before": before, "after": text, "why": change["why"],
                      "brd": re.sub(r"\s+", " ", row.cells[1].text).strip()[:300]})
    template = rows["CSHID.008"]._tr
    for r in DOC["removed"]:
        anchor = rows[r["after"]]._tr if r["after"] in rows else None
        tr = b.clone_row(template, [r["id"], r["text"], r["fr"]])
        anchor.addnext(tr)
        rows[r["id"]] = docx.table._Row(tr, t)
        b.stats["removed_rows"] += 1
    mapping = OrderedDict()
    for row in t.rows[1:]:
        mapping[row.cells[0].text.replace("\n", "").strip()] = re.sub(r"\s+", " ", row.cells[2].text).strip()
    missing = [i for i in ids if i not in mapping]
    unmapped = [i for i in ids if i in mapping and i not in REMOVED and not FRS_RE.search(mapping[i])]
    if missing or unmapped:
        raise SystemExit(f"mapping check failed: missing {missing}, unmapped {unmapped}")
    extra = [i for i in mapping if i not in ids]
    if extra:
        raise SystemExit(f"mapping rows not in the BRD: {extra}")
    b.stats["brd_ids"] = len(ids)
    b.stats["bdoi_rows"] = len(rows) - len(DOC["removed"])
    b.stats["mapping_changes"] = len(DOC["mapping"])
    # list of integrations
    p = find_par(b, "PMS to BIBS for PDC")
    els = [b.clone_par(p._p, DOC["integrations_title"])] + [b.clone_par(p._p, "- " + x) for x in DOC["integrations"]]
    b.insert_after(p._p, els)
    return mapping


# ------------------------------------------------------------------------------------------------ FR table
def cell_items(tc) -> list[tuple[str, int]]:
    """(item ID, index of the element in the cell) of every item paragraph of BDOI's cell; a repeated ID is marked."""
    out = []
    seen: Counter = Counter()
    for i, el in enumerate(tc):
        if el.tag != qn("w:p"):
            continue
        m = ID_RE.match(u11.el_text(el))
        if m:
            fid = norm_id(m.group(1))
            seen[fid] += 1
            out.append((fid + (" (2nd)" if seen[fid] > 1 else ""), i))
    return out


def block_end(tc, items, key, row_id):
    """The last element of item key (with its sub-items) in the cell; key == row ID: up to the next item."""
    children = list(tc)
    pos = next(i for fid, i in items if fid == key)
    base = key.replace(" (2nd)", "")
    nxt = len(children)
    for fid, i in items:
        if i <= pos:
            continue
        f = fid.replace(" (2nd)", "")
        if base == row_id or not f.startswith(base + "."):
            nxt = i
            break
    for el in reversed(children[pos:nxt]):
        if el.tag == qn("w:tbl") or u11.el_text(el).strip():
            return el
    return children[pos]


def edit_functional_requirements(b: Builder):
    t = b.tables[3]
    rows = OrderedDict()
    group_rows = {}
    for row in t.rows[1:]:
        cid = row.cells[0].text.strip()
        if FRS_RE.fullmatch(cid):
            rows[cid] = row
        elif cid:
            group_rows[cid] = row
    b.stats["bdoi_entries"] = len(rows)
    item_row = {}
    bdoi_items = []
    for rid, row in rows.items():
        for fid, _ in cell_items(row.cells[3]._tc):
            item_row[fid] = rid
            bdoi_items.append(fid)
    b.stats["bdoi_items"] = len(bdoi_items)
    b.bdoi_items = bdoi_items
    b.item_row = dict(item_row)
    row_refs: dict[str, list[str]] = {rid: [] for rid in rows}
    # elaboration of BDOI's items
    for key, add in ADD["items"].items():
        if key not in item_row:
            raise SystemExit(f"BDOI item {key} not found")
        rid = item_row[key]
        tc = rows[rid].cells[3]._tc
        last = block_end(tc, cell_items(tc), key, rid)
        b.insert_after(last, b.statements(add["shall"]) + b.acceptance(add["ac"]))
        row_refs[rid] += add["refs"]
        b.stats["elaborated"] += 1
    # new sub-items at the end of BDOI's entries
    for ns in ADD["new_subitems"]:
        rid = ns["row"]
        prefix = ns["id"].rsplit(".", 1)[0]
        prev = [f.replace(" (2nd)", "") for f in bdoi_items if item_row[f] == rid
                and f.replace(" (2nd)", "").rsplit(".", 1)[0] == prefix]
        if prev and frs_key(ns["id"]) <= max(frs_key(p) for p in prev):
            raise SystemExit(f"{ns['id']} does not continue BDOI's numbering")
        tc = rows[rid].cells[3]._tc
        els = [b.para("", size=11), b.para(f"{ns['id']} {ns['title']}", size=11, keep_next=True)]
        els += b.statements(ns["shall"]) + b.acceptance(ns["ac"])
        for el in els:
            tc.append(el)
        item_row[ns["id"]] = rid
        row_refs[rid] += ns["refs"]
        b.stats["new_subitems"] += 1
    # new entries
    template = rows["FRS.CSH.02.03"]._tr
    group_template = group_rows["Record Inquiry"]._tr
    placed = {rid: row._tr for rid, row in rows.items()}
    for item in ADD["new_items"]:
        anchor = placed[item["after"]]
        if item.get("group"):
            g = u11.strip_ids(copy.deepcopy(group_template))
            for tc in g.findall(qn("w:tc")):
                for el in list(tc):
                    if el.tag != qn("w:tcPr"):
                        tc.remove(el)
                tc.append(b.para(item["group"], bold=True, size=11))
                if b.hl:
                    u11.cell_shade(tc)
            anchor.addnext(g)
            anchor = g
        tr = u11.strip_ids(copy.deepcopy(template))
        tcs = tr.findall(qn("w:tc"))
        for tc in tcs:
            for el in list(tc):
                if el.tag != qn("w:tcPr"):
                    tc.remove(el)
            if b.hl:
                u11.cell_shade(tc)
        tcs[0].append(b.para(item["id"], size=11))
        tcs[1].append(b.para(item["component"], size=11))
        tcs[2].append(b.para(item["capability"], size=11))
        first = True
        for si in item["subitems"]:
            if not first:
                tcs[3].append(b.para("", size=11))
            first = False
            tcs[3].append(b.para(f"{si['id']} {si['title']}", size=11, keep_next=True))
            for el in b.statements(si["shall"]) + b.acceptance(si["ac"]):
                tcs[3].append(el)
            item_row[si["id"]] = item["id"]
            b.stats["new_item_subitems"] += 1
        row_refs[item["id"]] = [r for si in item["subitems"] for r in si["refs"]]
        anchor.addnext(tr)
        placed[item["id"]] = tr
        b.stats["new_items"] += 1
    # column Application Feature Model
    screen_no = b.screen_numbers
    for rid, tr in placed.items():
        tc = tr.findall(qn("w:tc"))[4]
        feature = DOC["feature"][rid]
        feature = re.sub(r"\{(SCR-OP-\d\d)\}", lambda m: f"Appendix {screen_no[m.group(1)]}", feature)
        refs = sorted(set(row_refs.get(rid, [])))
        paras = tc.findall(qn("w:p"))
        empty = len(paras) == 1 and not u11.el_text(paras[0]).strip()
        texts = [f"Feature: {feature}", "BIBS reference: " + (", ".join(refs) if refs else "-")]
        new = [b.para(x, size=10) for x in texts]
        if not paras:
            for el in new:
                tc.append(el)
        elif empty:
            paras[0].addprevious(new[0])
            new[0].addnext(new[1])
            tc.remove(paras[0])
        else:
            b.insert_after(paras[-1], new)
    b.item_row = item_row
    b.row_refs = row_refs


def edit_signoff(b: Builder):
    cell = b.tables[4].rows[0].cells[0]
    for p in cell.paragraphs:
        if p.text.strip() == "System Analyst":
            new = u11.strip_ids(copy.deepcopy(p._p))
            for r in new.findall(qn("w:r")):
                new.remove(r)
            new.append(b.run(DOC["signoff"]["drafting"]))
            p._p.addnext(new)
            return
    raise SystemExit("sign-off: 'System Analyst' not found")


# ------------------------------------------------------------------------------------------------ checks
def check_bdoi_text(path: Path) -> int:
    """BDOI's FR IDs and texts are all still present, in order, in the same table cells, and every body paragraph
    of the source is still there, in order."""
    src = docx.Document(str(REPO / DOC["meta"]["source"]))
    new = docx.Document(str(path))
    ids_src = Counter(FRS_RE.findall("\n".join(u11.texts_of(src.element.body))))
    ids_new = Counter(FRS_RE.findall("\n".join(u11.texts_of(new.element.body))))
    missing = [i for i, n in ids_src.items() if ids_new[i] < n]
    problems = [f"FR ID missing: {missing}"] if missing else []
    new_tables = list(new.tables)
    j = 0
    for ti, ts in enumerate(src.tables):
        key = u11.texts_of(ts._tbl)[:1]
        while j < len(new_tables) and u11.texts_of(new_tables[j]._tbl)[:1] != key:
            j += 1
        if j >= len(new_tables):
            problems.append(f"table {ti} of the source not found")
            break
        tn = new_tables[j]
        j += 1
        new_rows = list(tn.rows)
        k = 0
        for ri, row in enumerate(ts.rows):
            first = u11.texts_of(row.cells[0]._tc)
            while k < len(new_rows) and not u11.subsequence(first, u11.texts_of(new_rows[k].cells[0]._tc)):
                k += 1
            if k >= len(new_rows):
                problems.append(f"table {ti} row {ri}: not found")
                break
            for ci, cell in enumerate(row.cells):
                old = u11.texts_of(cell._tc)
                if ci < len(new_rows[k].cells) and not u11.subsequence(old, u11.texts_of(new_rows[k].cells[ci]._tc)):
                    problems.append(f"table {ti} row {ri} cell {ci}: BDOI text changed")
            k += 1
    body_src = [p.text.strip() for p in src.paragraphs if p.text.strip()]
    body_new = [p.text.strip() for p in new.paragraphs if p.text.strip()]
    it = iter(body_new)
    lost = [o for o in body_src if not any(o == n for n in it)]
    if lost:
        problems.append(f"body paragraphs changed: {lost[:3]}")
    if problems:
        raise SystemExit("BDOI text check failed:\n  " + "\n  ".join(problems))
    return sum(ids_src.values())


def check_words(b: Builder):
    hits = set()
    for t in b.added_text:
        for m in list(u11.RESTRICTED.finditer(t)) + list(u11.RESTRICTED_CS.finditer(t)):
            hits.add(m.group(0))
        for m in CODE_RE.finditer(t):
            if m.group(0) not in BDOI_CODES:
                hits.add(m.group(0))
        if re.search(r"\b[a-z]{2,}(?:_[a-z0-9]+)+\b", t):
            hits.add(re.search(r"\b[a-z]{2,}(?:_[a-z0-9]+)+\b", t).group(0))
    if hits:
        raise SystemExit(f"restricted words or internal codes in the added text: {sorted(hits)}")


# ------------------------------------------------------------------------------------------------ main
def build(highlight: bool, ids, figs) -> Builder:
    import brd02_annexes as annexes  # noqa: PLC0415
    b = Builder(highlight)
    b.ids = ids
    b.screen_numbers = annexes.screen_numbers()
    b.fig_numbers = annexes.figure_numbers()
    where = annexes.classify_items()
    b.stats["classified"] = annexes.check_classification(where)
    b.stats["emails_checked"] = annexes.check_emails()
    edit_cover(b)
    edit_revision_log(b)
    edit_introduction(b)
    mapping = edit_mapping(b, ids)
    b.mapping = mapping
    edit_functional_requirements(b)
    edit_signoff(b)
    last_ref = edit_references(b)
    a = annexes.Appendices(b, figs, where)
    a.write_all()
    a.place(last_ref)
    missing_fr = [f for f in SCOPE_FRS if f not in a.listed_refs]
    if missing_fr:
        raise SystemExit(f"Appendix Q lacks {missing_fr}")
    if a.cmp_ids - a.listed_conflicts:
        raise SystemExit(f"Appendix R lacks {sorted(a.cmp_ids - a.listed_conflicts)}")
    b.key_numbers = annexes.key_numbers(b)
    annexes.add_summary(b, figs)
    if highlight:
        annexes.change_summary(b)
    u11.set_update_fields(b)
    check_words(b)
    b.where = where
    return b


def build_review(figs_dir: Path | None = None) -> Builder:
    """The clean copy built in memory (not saved): the numbers, Appendix R rows and lists the workbooks use."""
    ids = brd_ids()
    work = Path(tempfile.mkdtemp(prefix="cshfig_")) if figs_dir is None else figs_dir
    try:
        figs = render_figures(work)
        b = build(False, ids, figs)
    finally:
        if figs_dir is None:
            shutil.rmtree(work, ignore_errors=True)
    b.ids = ids
    return b


def render_figures(work: Path) -> dict[str, Path]:
    import brd02_annexes as annexes  # noqa: PLC0415
    work.mkdir(parents=True, exist_ok=True)
    out = {}
    for spec in FLOWS:
        if spec["id"] == "brd02_fig_menus":
            spec = dict(spec, rows=annexes.menu_rows(spec))
        out[spec["id"]] = brd11_figures.render(spec, work)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true", help="skip the PDF")
    args = ap.parse_args(argv)
    ids = brd_ids()
    OUT.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="cshfig_"))
    results = {}
    try:
        figs = render_figures(work)
        for highlight, name in ((False, DOC["meta"]["clean"]), (True, DOC["meta"]["highlighted"])):
            b = build(highlight, ids, figs)
            path = OUT / name
            b.doc.save(str(path))
            n_ids = check_bdoi_text(path)
            results[name] = (b, n_ids)
            print(f"wrote {path.relative_to(REPO)} ({path.stat().st_size // 1024} KB)")
    finally:
        shutil.rmtree(work, ignore_errors=True)
    b, n_ids = results[DOC["meta"]["clean"]]
    if not args.no_pdf:
        pdf = OUT / DOC["meta"]["clean"].replace(".docx", ".pdf")
        u11.to_pdf(OUT / DOC["meta"]["clean"], pdf)
        print(f"wrote {pdf.relative_to(REPO)} ({pdf.stat().st_size // 1024} KB)")
    s = b.stats
    print(f"BRD IDs mapped: {s['brd_ids']} ({s['bdoi_rows']} BDOI rows kept, {s['removed_rows']} removed IDs added, "
          f"{s['mapping_changes']} rows given FR items)")
    print(f"FR: {s['bdoi_entries']} entries and {s['bdoi_items']} items of BDOI kept ({n_ids} FR IDs checked); "
          f"{s['elaborated']} items elaborated; {s['new_subitems']} new sub-items; {s['new_items']} new entries "
          f"({s['new_item_subitems']} sub-items); {s['requirements_added']} statements and {s['ac']} acceptance "
          f"criteria")
    for row in b.key_numbers:
        print("  ", " | ".join(str(x) for x in row))
    print("appendices:", {k: v for k, v in s.items() if k.startswith("app_")})
    return 0


if __name__ == "__main__":
    sys.exit(main())
