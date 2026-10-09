"""The review workbooks of BRD-02 Operations Cashiering that go with the FRS in BDOI's template v3.2.

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd02_workbooks.py

Builds the FRS v3.2 in memory first (build_brd02_bdoi_frs.build_review), so that every number of the workbooks is
the number of the document, then writes:

  BDOI_Template_FRS/BRD-02_Operations_Cashiering/
    BIBS_Inputs_BRD-02_Business_Unit_Requirements_Collection_v1.0.xlsx  what the Business Unit provides
    BIBS_FitGap_BRD-02_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx           every BRD requirement against the standard
    BIBS_CR_BRD-02_Change_Request_Register_v1.0.xlsx                    change control, candidates, decision log
    BIBS_RTM_BRD-02_Test_Cases_and_Traceability_v1.0.xlsx               traceability and test cases
    BIBS_UAT_BRD-02_Business_Unit_Walkthrough_Users_v1.0.xlsx           walkthrough users and script
  Quality/
    BIBS_Coverage_BRD-02_FRS_Coverage_in_BIBS_v1.0.xlsx                 internal: FRS coverage in the BIBS environment

Sources: brd02_workbooks.yaml and brd02_review.yaml (this folder), the BIBS FRS BRD-2 and its test cases, the
comparison of BDOI's Cashiering FRS, the programme workbooks (BDOI inputs v1.2, Feature List vs OOTB, conformance
register, UAT walkthrough users). No password is written: the password column says it is issued separately.

Self-checks: every BRD ID in scope has at least one test case and a fit-gap row; every FR item (BDOI's and added) has
at least one test case and a coverage row; every open question and input of Appendix S is in the requirements
collection workbook; every Appendix R observation has a decision row; the numbers agree with the FRS.
"""

from __future__ import annotations

import datetime as dt
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import docx
import openpyxl
import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import brd02_annexes as annexes  # noqa: E402
import build_brd02_bdoi_frs as frs  # noqa: E402
import build_brd11_bdoi_frs as u11  # noqa: E402
from brd11_xlsx import Column, ReviewWorkbook  # noqa: E402

REPO = frs.REPO
OPS = frs.OPS
OUT = REPO / "docs" / "deliverables" / "out"
WB = yaml.safe_load((HERE / "brd02_workbooks.yaml").read_text(encoding="utf-8"))
REV = frs.REV
DOC = frs.DOC
META = WB["meta"]
FOLDER = OUT / META["folder"]
QUALITY = OUT / META["quality"]
DATES = {k: (v if isinstance(v, dt.date) else dt.date.fromisoformat(str(v))) for k, v in META["dates"].items()}
INPUTS_XLSX = OUT / "Programme" / "BDOI_Inputs" / "BIBS_Inputs_BRD-00_BDOI_Requirements_and_Inputs_by_BRD_v1.2.xlsx"
STATUS = ["Open", "Received", "Confirmed"]
RESULT = ["Not run", "Pass", "Fail", "Blocked"]
FRS_RE = frs.FRS_RE
FROP_RE = re.compile(r"FR-OP-\d{3}")
OWNER = "Cashiering Process Owner"
VERSION = DOC["meta"]["version_new"]


def d(x) -> str:
    if isinstance(x, (dt.date, dt.datetime)):
        return x.strftime("%d-%b-%Y")
    try:
        return dt.date.fromisoformat(str(x)).strftime("%d-%b-%Y")
    except ValueError:
        return str(x or "")


def is_code(text: str) -> bool:
    return bool(re.fullmatch(r"[A-Z0-9_]+", text.strip()))


def expand(cid: str, items: list[dict]) -> list[str]:
    """The FR items a row of the comparison names: 'FRS.CSH.02.01.08.02-.12 (x)' -> the items and their sub-items."""
    cid = re.sub(r"\s*\(.*\)$", "", cid).strip()
    m = re.match(r"^(.*)\.(\d{2})-\.(\d{2})$", cid)
    heads = [f"{m.group(1)}.{n:02d}" for n in range(int(m.group(2)), int(m.group(3)) + 1)] if m else [cid]
    out = []
    for h in heads:
        out += [it["id"] for it in items if under(it, h)]
    return out


def under(it: dict, ref: str) -> bool:
    """The item is ref, a sub-item of ref, or an item of the FR entry ref (entry FRS.CSH.01.01 holds FRS.OPS.001)."""
    base = it["id"].replace(" (2nd)", "")
    return base == ref or base.startswith(ref + ".") or it["entry"] == ref


# ----------------------------------------------------------------------------------------------- the data
class Data:
    def __init__(self):
        self.b = frs.build_review()
        b = self.b
        self.a = b.annex
        self.ids = b.ids
        self.mapping = b.mapping
        self.trace = frs.our_trace()
        self.our = frs.our_frs()
        self.brd_text = self.read_brd_text()
        self.cases_yaml = yaml.safe_load((OPS / "brd02_cases.yaml").read_text(encoding="utf-8"))
        self.items = self.fr_items()
        self.item_ids = [it["id"] for it in self.items]
        self.item_by_id = {it["id"]: it for it in self.items}
        self.cmp = yaml.safe_load(frs.CMP.read_text(encoding="utf-8"))
        self.cmp_status = self.read_cmp()
        self.fr_items_map = self.fr_to_items()
        self.scr = annexes.screens()
        self.scr_no = annexes.screen_numbers()
        self.bibs_cases = self.make_bibs_cases()
        self.entry_frops: dict[str, set] = defaultdict(set)
        for it in self.items:
            self.entry_frops[it["entry"]] |= set(it["refs"]) | set(self.cmp_status.get(it["id"], ("", "", []))[2])
        for entry, refs in (getattr(self.b, "row_refs", {}) or {}).items():
            self.entry_frops[entry] |= set(refs)
        self.by_entry: set[str] = set()
        self.acc_cases = self.make_acceptance_cases()
        self.steps_by_fr = self.walkthrough_steps()
        self.inputs = self.read_inputs()
        self.conformance = self.read_conformance()

    # ---------------------------------------------------------------- sources
    def read_brd_text(self) -> dict[str, str]:
        out = {}
        for row in self.b.tables[2].rows[1:]:
            rid = row.cells[0].text.replace("\n", "").strip()
            out[rid] = re.sub(r"\s+", " ", row.cells[1].text).strip()
        return out

    def fr_items(self) -> list[dict]:
        src = docx.Document(str(REPO / DOC["meta"]["source"]))
        t = src.tables[3]
        add_items = frs.ADD["items"]
        out = []
        for row in t.rows[1:]:
            entry = row.cells[0].text.strip()
            if not FRS_RE.fullmatch(entry):
                continue
            tc = row.cells[3]._tc
            children = list(tc)
            for fid, i in frs.cell_items(tc):
                head = u11.el_text(children[i]).strip()
                title = re.sub(r"^\s*FRS\s*\.\s*(?:CSH|OPS)[\d .]*\d\.?\s*", "", head).strip() or head
                add = add_items.get(fid, {})
                out.append({"id": fid, "entry": entry, "title": title[:90], "origin": "BDOI FRS v3.1",
                            "refs": add.get("refs", []), "ac": add.get("ac", []),
                            "elaborated": fid in add_items})
        for ns in frs.ADD["new_subitems"]:
            out.append({"id": ns["id"], "entry": ns["row"], "title": ns["title"], "origin": "Added in v3.2",
                        "refs": ns["refs"], "ac": ns["ac"], "elaborated": False})
        for item in frs.ADD["new_items"]:
            for si in item["subitems"]:
                out.append({"id": si["id"], "entry": item["id"], "title": si["title"], "origin": "Added in v3.2",
                            "refs": si["refs"], "ac": si["ac"], "elaborated": False})
        s = self.b.stats
        if len(out) != s["bdoi_items"] + s["new_subitems"] + s["new_item_subitems"]:
            raise SystemExit(f"FR item count differs from the FRS ({len(out)})")
        if sum(len(it["ac"]) for it in out) != s["ac"]:
            raise SystemExit("acceptance criteria count differs from the FRS")
        return out

    def read_cmp(self) -> dict[str, tuple[str, str, list[str]]]:
        """Item -> (platform status, note, FR-OP) from the comparison of BDOI's FRS with the platform."""
        parts: dict[str, list] = defaultdict(list)
        for m in self.cmp["mapping"]:
            for i in expand(m["id"], self.items):
                parts[i].append((m["platform"], frs.clean(m.get("difference") or m.get("action_detail") or ""),
                                 FROP_RE.findall(str(m.get("ours", "")))))
        out = {}
        for i, lst in parts.items():
            st = {x[0] for x in lst}
            status = st.pop() if len(st) == 1 else "Partly"
            out[i] = (status, " ".join(dict.fromkeys(x[1] for x in lst if x[1])),
                      sorted({f for x in lst for f in x[2]}))
        return out

    def fr_to_items(self) -> dict[str, list[str]]:
        out: dict[str, list[str]] = defaultdict(list)
        for it in self.items:
            refs = set(it["refs"]) | set(self.cmp_status.get(it["id"], ("", "", []))[2])
            for r in refs:
                if it["id"] not in out[r]:
                    out[r].append(it["id"])
        return out

    def read_inputs(self) -> dict[str, dict]:
        wb = openpyxl.load_workbook(INPUTS_XLSX, read_only=True)
        out = {}
        for r in wb["BRD-02 Operations"].iter_rows(min_row=5, values_only=True):
            if r[0]:
                out[r[0]] = {"id": r[0], "type": r[2], "need": r[3], "owner": r[7], "due": d(r[8])}
        wb.close()
        return out

    def read_conformance(self) -> dict[str, list[tuple]]:
        p = OUT / "Programme" / "Quality" / "BIBS_Conformance_BRD-00_End-to-End_Conformance_Register_v1.0.xlsx"
        wb = openpyxl.load_workbook(p, read_only=True)
        out = defaultdict(list)
        for r in wb["FR conformance"].iter_rows(values_only=True):
            if r[0] == "BRD-02":
                out[r[1]].append((r[4], r[7]))
        wb.close()
        return out

    # ---------------------------------------------------------------- test cases
    def persona(self, code: str) -> tuple[str, str]:
        if code in annexes.PERSONAS:
            return annexes.PERSONAS[code]
        if code in WB["personas"]:
            return tuple(WB["personas"][code])
        p = self.cases_yaml["personas"].get(code, {"name": code, "user": "-"})
        return frs.clean(str(p.get("name"))), p.get("user", "-")

    def brd_of_frop(self, fr_id: str) -> str:
        return ", ".join(r for r, t in self.trace.items() if fr_id in t["frs"])

    def make_bibs_cases(self) -> list[dict]:
        cy = self.cases_yaml
        scen = {s["id"]: s for s in cy["scenarios"]}
        data_names = {x["id"]: x["name"] for x in cy["data"]}
        screens = cy["screens"]
        out = []
        for fr_id, e in cy["frs"].items():
            if fr_id not in frs.SCOPE_FRS:
                continue
            n = fr_id.split("-")[-1]
            seq: Counter = Counter()
            prio = self.our[fr_id].get("priority", "Must have")
            for c in e.get("cases") or []:
                ci = c["c"]
                seq[ci] += 1
                persona, user = self.persona(c.get("persona", e["persona"]))
                pre = c.get("pre") or scen.get(e.get("scenario"), {}).get("pre", "")
                exp = c["expected"]
                msgs = [x for x in (c.get("msg") or []) if not is_code(x)]
                if msgs:
                    exp += " Message: " + "; ".join(re.sub(r"[{}]", "", x) for x in msgs)
                scr = c.get("screen", e.get("screen"))
                out.append({
                    "id": f"TC-OP-{n}.{ci}-{seq[ci]:02d}", "title": frs.clean(c["title"]), "fr": fr_id,
                    "item": ", ".join(sorted(self.fr_items_map.get(fr_id, []), key=frs.frs_key)),
                    "brd": self.brd_of_frop(fr_id) or ", ".join(
                        x.split(" (")[0] for x in self.our[fr_id]["brd"] if isinstance(x, str)),
                    "persona": persona, "user": user,
                    "screen": frs.clean(screens.get(scr, scr) or "-"),
                    "pre": frs.clean(pre), "steps": "\n".join(f"{i}. {frs.clean(s)}" for i, s in enumerate(c["steps"], 1)),
                    "data": "; ".join(f"{x} {data_names.get(x, '')}" for x in (c.get("data") or e.get("data") or [])),
                    "expected": frs.clean(exp),
                    "type": "Negative" if c["type"] == "Negative" or c.get("neg") else "Positive",
                    "priority": "High" if str(prio).startswith("Must") else "Medium",
                    "source": "BIBS test plan BRD-2 v2.1"})
        return out

    def make_acceptance_cases(self) -> list[dict]:
        out = []
        for it in self.items:
            code, screen = WB["entry_persona"].get(it["entry"], ["CASHIER", "Finance › Cashiering"])
            persona, user = self.persona(code)
            for n, ac in enumerate(it["ac"], 1):
                neg = re.search(r"\b(refused|cannot|not offered|is not|are not|never|denied|rejected|blocked|"
                                r"stops|no longer)\b", ac, re.I)
                key = it["id"].replace("FRS.", "FRS-").replace(" (2nd)", "-2ND")
                out.append({
                    "id": f"TC-{key}-AC{n:02d}",
                    "title": f"Acceptance criterion {n} of {it['id']} {it['title']}"[:120], "fr": ", ".join(it["refs"]),
                    "item": it["id"], "brd": self.brd_of_item(it["id"]), "persona": persona, "user": user,
                    "screen": screen, "pre": "Users, receipt series, booked accounts and payment files of the SIT/UAT "
                                             "data (TD-OP-01 to TD-OP-08).",
                    "steps": f"1. Sign in as {user} ({persona}).\n2. Open {screen}.\n3. Bring about the situation of "
                             f"the acceptance criterion and observe the result.",
                    "data": "TD-OP-01 SIT/UAT users and roles; TD-OP-03 Cashiering storyline", "expected": ac,
                    "type": "Negative" if neg else "Positive", "priority": "High",
                    "source": f"Acceptance criterion of the FRS v{VERSION}"})
        if len(out) != self.b.stats["ac"]:
            raise SystemExit("acceptance criteria count differs from the FRS")
        return out

    def brd_of_item(self, item: str) -> str:
        it = self.item_by_id[item]
        base = item.replace(" (2nd)", "")
        out = []
        for r, mtext in self.mapping.items():
            for ref in FRS_RE.findall(mtext):
                if under(it, ref) or ref.startswith(base + "."):
                    out.append(r)
                    break
        return ", ".join(out)

    def cases_of_item(self, item: str, fallback: bool = True) -> list[str]:
        ids = [c["id"] for c in self.acc_cases if c["item"] == item]
        ids += [c["id"] for c in self.bibs_cases if item in c["item"].split(", ")]
        if not ids:   # a heading item: the cases of its sub-items
            base = item.replace(" (2nd)", "")
            for it in self.items:
                if it["id"].startswith(base + "."):
                    ids += self.cases_of_item(it["id"], fallback)
        if not ids and fallback:   # an item of BDOI's without its own reference: the cases of its FR entry
            entry = self.item_by_id[item]["entry"]
            frops = self.entry_frops.get(entry, set())
            ids += [c["id"] for c in self.bibs_cases if c["fr"] in frops]
            if ids:
                self.by_entry.add(item)
        return list(dict.fromkeys(ids))

    def walkthrough_steps(self) -> dict[str, list[str]]:
        out = defaultdict(list)
        for w in yaml.safe_load((OPS / "pack" / "walkthroughs.yaml").read_text(encoding="utf-8"))["walkthroughs"]:
            for i, s in enumerate(w["steps"], 1):
                if i in annexes.WALK_STEPS.get(w["id"], []):
                    for fr in self.scr.get(s[1], {}).get("frs", []):
                        out[fr].append(f"{w['id']}.{i}")
        wt = REV["walkthrough"]
        for i, s in enumerate(wt["steps"], 1):
            for fr in self.scr.get(s[1], {}).get("frs", []):
                out[fr].append(f"{wt['id']}.{i}")
        return out

    def screens_of(self, frops: list[str]) -> str:
        out = [f"{self.scr_no[sid]} {frs.clean(self.scr[sid]['title'])}" for sid in annexes.SCREENS
               if set(self.scr[sid].get("frs", [])) & set(frops)]
        return "; ".join(out) or "-"


def step_key(x: str):
    a, b = x.rsplit(".", 1)
    return (a, int(b))


# ----------------------------------------------------------------------------------------- 1 requirements collection
def rc_cols():
    return [
        {"key": "id", "header": "ID", "width": 12, "mandatory": True, "format": "Text", "what": "Reference of the item (kept from the BDOI inputs workbook where it exists)"},
        {"key": "item", "header": "Item", "width": 28, "mandatory": True, "what": "What the item is about"},
        {"key": "brd", "header": "BRD ID", "width": 14, "what": "BRD requirement the item serves"},
        {"key": "frs", "header": "FRS reference", "width": 18, "what": f"FR item or appendix of the FRS v{VERSION}"},
        {"key": "need", "header": "What is needed", "width": 46, "mandatory": True, "what": "What the Business Unit provides or decides"},
        {"key": "fmt", "header": "Format of the answer", "width": 26, "mandatory": True, "what": "How to answer"},
        {"key": "example", "header": "Example", "width": 24, "what": "An example of an answer"},
        {"key": "by", "header": "Provided by (role)", "width": 20, "mandatory": True, "what": "Role that answers"},
        {"key": "due", "header": "Needed by", "width": 12, "mandatory": True, "format": "Date dd-Mmm-yyyy", "what": "Date of the drop plan"},
        {"key": "resp", "header": "BU response", "width": 34, "what": "The answer of the Business Unit"},
        {"key": "status", "header": "Status", "width": 12, "mandatory": True, "values": STATUS, "what": "Open until answered; Received; Confirmed by the project"},
    ]


def frs_table_text(ref: str) -> str:
    text = (OPS / "FRS_BRD02_OPERATIONS.md").read_text(encoding="utf-8")
    m = re.search(r"^\| (?:[\w-]+ / )*" + re.escape(ref) + r"(?: / [\w-]+)* \| ([^|]+)\|", text, re.M)
    return frs.clean(m.group(1).strip()) if m else ""


def row(id_, item, brd, frsref, need, fmt, example, by, due) -> dict:
    return {"id": id_, "item": item, "brd": brd, "frs": frsref, "need": need, "fmt": fmt, "example": example,
            "by": by, "due": d(due), "resp": "", "status": "Open"}


def input_row(data: Data, ref: str, item: str, brd: str, frsref: str, fmt: str, example: str, default_due) -> dict:
    src = data.inputs.get(ref)
    need = frs.clean(src["need"]) if src else (frs_table_text(ref) or item)
    by = src["owner"] if src else OWNER
    due = src["due"] if src else d(default_due)
    return {"id": ref, "item": item, "brd": brd, "frs": frsref, "need": need, "fmt": fmt, "example": example,
            "by": by, "due": due, "resp": "", "status": "Open"}


def build_requirements(data: Data) -> tuple[Path, dict]:
    wb = ReviewWorkbook("Business Unit Requirements Collection", doc_type="Requirements collection", brd="BRD-02",
                        version=META["version"], subtitle="BRD-02 Operations Cashiering - what the Business Unit "
                                                          f"provides for the FRS v{VERSION}")
    wb.cover_notes = ["Fill in the column BU response and set the Status of each row; questions go to the sheet "
                      "Questions and comments. Items keep the IDs of the BDOI Requirements and Inputs by BRD v1.2 "
                      "workbook where they exist."]
    cols = rc_cols()
    sheets: "OrderedDict[str, list[dict]]" = OrderedDict()
    owner = {}
    dec, uat, sit, it_due = DATES["decision"], DATES["uat"], DATES["sit"], DATES["it"]
    # documents, print formats and report layouts
    rows = []
    for doc in REV["documents"]["items"]:
        rows.append(row(doc["rc"], doc["name"], re.search(r"(?:CSHID|BRQID)\.\d{3}", doc["purpose"]).group(0)
                        if re.search(r"(?:CSHID|BRQID)\.\d{3}", doc["purpose"]) else "CSHID.019",
                        f"Appendix L ({doc['id']})",
                        f"Confirm the layout of the {doc['name']} shown in Appendix L ({doc['id']}), or give BDOI's "
                        f"layout. Format: {doc['format']}",
                        "Confirm, or a sample (PDF or Excel) with the fields in order", "Confirmed as in Appendix L",
                        OWNER, dec))
    rows.append(row(f"RC-DOC-{len(rows) + 1:02d}", "Report file names", "CSHID.017", "FRS.CSH.09.01.02; Appendix M",
                    "Confirm BDOI's file names of the report exports, or give the pattern.",
                    "Confirm / other pattern", "Applied Premium Reports_10092026", OWNER, dec))
    for k, r in enumerate(REV["reports"]["rows"], 1):
        rows.append(row(f"RC-RPT-{k:02d}", r[1], "CSHID.017, 018, 023", f"{r[0]}; Appendix M",
                        f"Confirm the filters ({r[4]}), sort ({r[6]}) and timing ({r[2]}) of the report; the columns "
                        f"are BDOI's: {r[5]}.", "Confirm, or the change", "Confirmed", OWNER, dec))
    sheets["Documents and print formats"] = rows
    owner["Documents and print formats"] = OWNER
    # e-mails and notices
    rows = []
    for i in REV["emails"]["items"]:
        rows.append(row(i["id"], i["name"], i["brd"], f"Appendix G; {i['frs']}",
                        f"Approve the wording, the trigger and the recipients. Trigger: {i['trigger']} Recipients: "
                        f"{i['recipients']} Subject: {i['subject']}. Text: {i['body']}",
                        "Approve, or the new wording", "Approved", OWNER, dec))
    rows.append(row(f"RC-EML-{len(rows) + 1:02d}", "Sender and footer of the e-mails", "BRQID.005", "Appendix G",
                    "The sender address of each environment and, if required, the standard footer or disclaimer of "
                    "BDOI e-mails.", "Address per environment; footer text", "bibs-noreply@ (domain of BDOI)",
                    "BDOI IT", uat))
    sheets["E-mail wordings"] = rows
    owner["E-mail wordings"] = OWNER
    # messages
    rows = []
    for r in data.a.message_rows:
        rows.append(row(r[0], r[1][:60], "CSHID.001-025", "Appendix F.2",
                        f"Confirm the wording of the message, or give BDOI's wording: \"{r[2]}\" (type {r[3]}).",
                        "Confirm, or the new wording", "Confirmed", OWNER, uat))
    sheets["Messages to confirm"] = rows
    owner["Messages to confirm"] = OWNER
    # schedules
    rows = []
    for k, s in enumerate(REV["reports"]["schedules"], 1):
        rows.append(row(f"RC-SCH-{k:02d}", s[0], s[5].split(";")[0], "Appendix M",
                        f"Confirm the time ({s[1]}) and who is told: {s[3]}.", "Time and recipients, or Confirm",
                        s[1], s[4].split(";")[0], uat))
    for item, need, fmt, ex in [
            ("Posting of creation and cancellation records", "Time within which the poster acts on a record For "
             "Posting, the reminder and the escalation role (Appendix J).", "Working hours; role",
             "Same day; reminder at 15:00; Cashiering Team Head"),
            ("Approval of dispositions", "Time within which a disposition For Approval is approved, the reminder and "
             "the escalation role (FRS.CSH.06.02).", "Working days; role", "1 day; Cashiering Team Head"),
            ("Ageing of unapplied payments", "The ageing brackets of the Unapplied list and of the ageing report, and "
             "the age at which the unit head is told (FRS.CSH.06.01, 09.02.22).", "Brackets in days; age",
             "0-30, 31-60, 61-90, over 90; 60 days"),
            ("Cashier's day-end", "Cut-off time of the cashier's day for the receipts of the day and the day-end list "
             "(FRS.CSH.02.01.16).", "Time", "17:00")]:
        rows.append(row(f"RC-SCH-{len(rows) + 1:02d}", item, "CSHID.001, 024", "Appendix J; Appendix M", need, fmt,
                        ex, OWNER, dec))
    sheets["Schedules and timings"] = rows
    owner["Schedules and timings"] = OWNER
    # parameters
    rows = []
    for k, p in enumerate(DOC["annex_h"]["parameters"], 1):
        rows.append(row(f"RC-PAR-{k:02d}", p[0], "CSHID.001-025", "Appendix H",
                        f"Confirm the value of the parameter {p[0]}: {p[2]}. Value in this FRS: {p[1]}.", "Value",
                        str(p[1])[:40], p[3], uat))
    sheets["Parameters and thresholds"] = rows
    owner["Parameters and thresholds"] = "Cashiering Team Leader; Comptrollership"
    # lists
    rows = []
    for k, lst in enumerate(DOC["annex_h"]["lists"], 1):
        rows.append(row(f"RC-LOV-{k:02d}", lst[0], "CSHID.001-025", "Appendix H",
                        f"Give or confirm the values of the list {lst[0]} (in this FRS: {lst[1]}).",
                        "Values (code, name)", str(lst[1]).split(";")[0][:40], OWNER, sit))
    sheets["Lists of values"] = rows
    owner["Lists of values"] = OWNER
    # master data
    rows = [input_row(data, "CFG-050", "Receipt series in use (AR and OR)", "CSHID.006, 015", "FRS.CSH.02.03.06; Appendix P",
                      "Template R11 Receipt series in use of the Master Data and Configuration Upload Templates v1.2",
                      "Filled tab R11", DATES["freeze"]),
            input_row(data, "CFG-100", "Insurers and insurer branches", "CSHID.001, 007", "Appendix P",
                      "Template of BRD-03 (insurers)", "Filled template", sit)]
    for k, r in enumerate(REV["data_setup"]["rows"], 1):
        rows.append(row(f"RC-MD-{k:02d}", r[0], "CSHID.001-025", "Appendix P",
                        f"Provide {r[0]} from {r[1]}; checked by {r[4]}. {r[6]}", f"Template {r[2]}",
                        f"Filled template {r[2]}"[:40], r[3], DATES["freeze"] if "freeze" in str(r[5]).lower() else sit))
    sheets["Master and reference data"] = rows
    owner["Master and reference data"] = "Cashiering data steward"
    # approvers and users per role
    rows = []
    for k, p in enumerate(DOC["annex_h"]["personas"], 1):
        rows.append(row(f"RC-APR-{k:02d}", f"Users of the role {p[0]}"[:70], p[2], "Appendix H",
                        f"Name the users of the role {p[0]} with their branch or unit. The role: {p[1]}.",
                        "User ID, name, branch or unit", "a013000101, Juan dela Cruz, Head Office", OWNER, sit))
    rows.append(input_row(data, "OQ48", "Operations access matrix", "BRQID.001, 002", "FRS.CSH.01.02; Appendix H",
                          "The matrix of Appendix H confirmed or changed", "Confirmed", dec))
    for ref, item in [("TD-02", "UAT testers"), ("SO-02-1", "Signature of the FRS"),
                      ("SO-02-2", "Approval of the test plan"), ("SO-02-3", "UAT sign-off certificate")]:
        rows.append(input_row(data, ref, item, "BRQID.001", "Signoff Sheet; Appendix S", "Names, or signature", "-", dec))
    sheets["Approvers and users per role"] = rows
    owner["Approvers and users per role"] = OWNER
    # integration inputs and open decisions
    rows_i, rows_o = [], []
    for k, r in enumerate(REV["integrations"]["rows"], 1):
        rows_i.append(row(f"RC-INT-{k:02d}", r[0][:70], "BRQID.004, 005", "Appendix N",
                          f"{r[6]} (exchanged: {r[2]}; {r[3]}). Status of BDOI's input: {r[7]}.",
                          "Details, or Confirm", "Folder, time and layout", r[5].split(";")[0], it_due))
    rows_i.append(input_row(data, "CFG-044", "Payment file layouts of the collection channels", "CSHID.008",
                            "FRS.CSH.05.01.06; Appendix D", "Layout per channel (field, position, format)",
                            "Layout of the PMS file", it_due))
    for drow in REV["open_items"]["dependencies"]:
        rows_i.append(row(drow[0], "Dependency", "", f"Appendix S.2; {drow[2]}", drow[1], "Date when it is in place",
                          d(sit), "BDOI IT" if re.search(r"IT|MFT|relay|layout", drow[1]) else OWNER, sit))
    for k, q in enumerate(REV["open_items"]["open"], 1):
        first = q["refs"].split(",")[0].split(";")[0].strip()
        brd = re.search(r"(?:CSHID|BRQID|MKTID|DBMID)\.\d{3}", q["why"])
        r = row(first if re.match(r"(OQ|CLR|CFG|IQ)", first) else f"RC-DEC-{k:02d}",
                re.split(r"[:?(]", q["q"])[0][:70], brd.group(0) if brd else "CSHID.001-025",
                f"Appendix S.3.{k} ({q['refs']})", q["q"], q["fmt"], "", q["owner"], q["by"])
        (rows_i if "IT" in q["owner"].split() else rows_o).append(r)
    sheets["Integration inputs"] = rows_i
    owner["Integration inputs"] = "BDOI IT"
    sheets["Open decisions"] = rows_o
    owner["Open decisions"] = OWNER
    # self-check: every input and open item of Appendix S is here
    present = {r["id"] for rws in sheets.values() for r in rws}
    present_text = " ".join(r["frs"] for rws in sheets.values() for r in rws)
    missing = [ref for ref, kind in data.b.where.items() if kind == "input" and ref not in present]
    missing += [f"S.3.{k}" for k in range(1, len(REV["open_items"]["open"]) + 1)
                if f"Appendix S.3.{k} " not in present_text]
    if missing:
        raise SystemExit(f"requirements collection lacks {missing}")
    index = []
    for k, (name, rws) in enumerate(sheets.items(), 1):
        dues = sorted({r["due"] for r in rws if r["due"]}, key=lambda x: dt.datetime.strptime(x, "%d-%b-%Y"))
        index.append([k, name, f"{len(rws)} items", owner[name], dues[0] if dues else "", len(rws),
                      f"=COUNTIF('{name}'!K:K,\"Open\")", "Not started", name])
    index.append([len(index) + 1, "Questions and comments", "Questions of the Business Unit on any item", "Any reviewer",
                  "", 0, "", "Not started", "Questions and comments"])
    wb.index_sheet("Start here", "Start here – BRD-02 Operations Cashiering: what the Business Unit provides",
                   "Fill in the sheets in this order (decisions and lists first, then data). Each sheet is "
                   "self-contained: title block, column guide above the header, an example row to overwrite. "
                   "Rows open: the formula counts the rows still Open.",
                   ["No.", "Sheet", "Content", "Owner", "First due", "Rows", "Rows open", "Status", "Go to"],
                   index, [5, 30, 18, 30, 12, 8, 10, 14, 10], link_col=9, status_col=8,
                   status_values=["Not started", "In progress", "Completed"])
    for name, rws in sheets.items():
        ex = dict(rws[0], resp=rws[0]["example"] or "Confirmed", status="Received")
        wb.template_sheet(name, f"{name} – BRD-02 Operations Cashiering",
                          [("Purpose", f"Items of the kind '{name}' the FRS v{VERSION} needs from the Business Unit."),
                           ("Who fills it in", owner[name]),
                           ("How it is used", "The project applies each confirmed answer to the FRS, the configuration or "
                                              "the templates; data go through the upload templates v1.2."),
                           ("Due", "The date in the column Needed by (drop plan)."),
                           ("Depends on", "The decisions of the sheet Open decisions and of Appendix R where an item "
                                          "names one.")],
                          cols, rws, ex)
    q_cols = [
        {"key": "no", "header": "No.", "width": 6, "mandatory": True, "format": "Number", "what": "Running number"},
        {"key": "ref", "header": "Item or section", "width": 20, "mandatory": True, "what": "ID of the item, or FRS section / appendix"},
        {"key": "q", "header": "Question or comment", "width": 60, "mandatory": True, "what": "The question or comment"},
        {"key": "by", "header": "Raised by", "width": 20, "mandatory": True, "what": "Name and role"},
        {"key": "date", "header": "Date", "width": 12, "format": "Date dd-Mmm-yyyy", "what": "Date raised"},
        {"key": "answer", "header": "Answer of the project", "width": 50, "what": "Filled in by iorta TechNXT"},
        {"key": "status", "header": "Status", "width": 12, "values": STATUS, "what": "Open, Received, Confirmed"}]
    wb.template_sheet("Questions and comments", "Questions and comments", [
        ("Purpose", f"Any question or comment of the Business Unit on the items or on the FRS v{VERSION}."),
        ("Who fills it in", "Any reviewer; the project answers."), ("How it is used", "Answered within 3 working days."),
        ("Due", "Before the sign-off of the FRS."), ("Depends on", "-")], q_cols, [],
        {"no": 1, "ref": "G-03", "q": "Can the notice also go to the Cashiering Team Head?", "by": "Cashiering Process Owner",
         "date": "", "answer": "", "status": "Open"})
    path = wb.save(FOLDER / "BIBS_Inputs_BRD-02_Business_Unit_Requirements_Collection_v1.0.xlsx")
    return path, {"rows": sum(len(r) for r in sheets.values()), "sheets": len(sheets) + 1}


# ------------------------------------------------------------------------------------------------- 2 fit-gap
def env_status(data: Data, frops: list[str]) -> str:
    res = [x[1] for f in frops for x in data.conformance.get(f, [])]
    if not res:
        return "Planned for"
    if all(x in ("Conformant", "Partly conformant") for x in res):
        return "Available"
    return "Partly available"


def build_fitgap(data: Data) -> tuple[Path, dict]:
    feat = {}
    wbf = openpyxl.load_workbook(OUT / "Programme" / "Feature_List" /
                                 "BIBS_BDOI_FeatureList_vs_OOTB_and_Best_Practice_v1.0.xlsx", read_only=True)
    for r in wbf["Feature Comparison"].iter_rows(min_row=5, values_only=True):
        if r[0]:
            feat[r[0]] = {"feature": r[3], "ootb": r[8], "assess": r[13]}
    wbf.close()
    group = {}
    for g in WB["fitgap_groups"]:
        for f in g["features"]:
            if f not in feat:
                raise SystemExit(f"feature {f} not in the Feature List")
        for i in g["ids"]:
            group[i] = g
    rows = []
    for rid in data.ids:
        g = group.get(rid)
        if g is None:
            raise SystemExit(f"no fit-gap group for {rid}")
        frops = data.trace.get(rid, {"frs": []})["frs"]
        items = sorted(set(FRS_RE.findall(data.mapping.get(rid, ""))), key=frs.frs_key)
        rows.append({"id": rid, "text": data.brd_text.get(rid, ""), "group": g["name"], "item": ", ".join(items) or "-",
                     "fr": ", ".join(frops) or "-",
                     "features": "; ".join(f"{f} {feat[f]['feature']} (OOTB: {feat[f]['ootb']})" for f in g["features"]) or "-",
                     "ootb": g["ootb"], "best": g["best"], "fit": g["fit"], "rec": g["rec"], "impact": g["impact"],
                     "env": "-" if g["fit"] in ("Removed from the BRD",) else env_status(data, frops)})
    nf = WB["fitgap_nfr"]
    for r in REV["nfr"]["rows"]:
        frops = sorted({f for it in FRS_RE.findall(r[4]) for x in data.items if under(x, it)
                        for f in x["refs"] + data.cmp_status.get(x["id"], ("", "", []))[2]})
        rows.append({"id": r[0], "text": f"{r[1]}: {r[2]}", "group": "Non-functional (BRD Cashiering annex)",
                     "item": r[4], "fr": ", ".join(frops) or "-", "features": "-", "ootb": r[3],
                     "best": f"BIBS standards for all BRDs (NFR register {r[5]})", "fit": nf["fit"], "rec": nf["rec"],
                     "impact": "Appendix O", "env": env_status(data, frops) if frops else "-"})
    wb = ReviewWorkbook("BRD vs Out-of-the-Box vs Best Practice", doc_type="Fit-gap", brd="BRD-02",
                        version=META["version"], subtitle="BRD-02 Operations Cashiering")
    wb.legend = [("Fit", "Met by the standard system"), ("Configure", "Met by set-up or settings"),
                 ("Extend", "Met by an extension of the standard system"),
                 ("Recommend change", "A change to the BRD text is recommended"),
                 ("Removed from the BRD", "Withdrawn by the BRD; kept for completeness"),
                 ("Specified in another FRS", "Held by reference in BDOI's mapping")]
    counts = Counter(r["fit"] for r in rows)
    by_group = defaultdict(Counter)
    for r in rows:
        by_group[r["group"]][r["fit"]] += 1
    fv = WB["fit_values"]
    summ = [{"group": "All BRD requirements", "n": len(rows), **{f: counts[f] for f in fv}}]
    summ += [{"group": g, "n": sum(c.values()), **{f: c[f] for f in fv}} for g, c in by_group.items()]
    wb.sheet("Summary", [Column("group", "Group", 46, "Function of the BRD"), Column("n", "Requirements", 14, "Count")]
             + [Column(f, f, 14, "Count") for f in fv],
             summ, description=f"Counts by fit: {len(data.ids)} BRD IDs and {len(REV['nfr']['rows'])} non-functional rows")
    wb.sheet("Fit-Gap", [
        Column("id", "BRD ID", 14, "BRD requirement ID (CSH-NFR-nn for the non-functional rows)"),
        Column("text", "BRD text", 46, "Requirement as printed in the BRD"),
        Column("group", "Function", 24, "Function of the BRD"),
        Column("item", "FR item (BDOI format)", 24, f"FR items of the FRS v{VERSION}"),
        Column("fr", "BIBS reference FR", 16, "FR of the BIBS FRS BRD-2"),
        Column("features", "Feature list reference", 40, "Rows of the BIBS Feature List vs OOTB workbook"),
        Column("ootb", "What BIBS offers out of the box", 46, "Standard capability or configuration"),
        Column("best", "Insurance-broking best practice", 46, "Practice of the Philippine market and international broking"),
        Column("fit", "Fit", 18, "Fit class", values=fv, status=True),
        Column("rec", "Recommendation", 40, "What to do"),
        Column("impact", "Impact", 22, "Where it lands"),
        Column("env", "In the BIBS environment", 20, "Available / Partly available / Planned for")],
        rows, description="One row per BRD requirement")
    path = wb.save(FOLDER / "BIBS_FitGap_BRD-02_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx")
    if not set(data.ids) <= {r["id"] for r in rows}:
        raise SystemExit("fit-gap lacks BRD IDs")
    return path, {"rows": len(rows), **counts}


# ------------------------------------------------------------------------------------------------- 3 change requests
def build_cr(data: Data) -> tuple[Path, dict]:
    wb = ReviewWorkbook("Change Request Register", doc_type="Change control", brd="BRD-02", version=META["version"],
                        subtitle=f"BRD-02 Operations Cashiering - after the sign-off of the FRS v{VERSION}")
    steps = [{"step": s[0], "what": s[1], "who": s[2], "record": s[3]} for s in REV["change_control"]["steps"]]
    wb.sheet("How to use", [Column("step", "Step", 12, "Step of the process"), Column("what", "What happens", 60, ""),
                            Column("who", "Who", 36, ""), Column("record", "Record in this workbook", 36, "")],
             steps, description="Raise -> assess -> approve -> schedule -> test -> close (FRS Appendix T)")
    rows = []
    for k, c in enumerate(WB["cr_candidates"], 1):
        rows.append({"no": f"CR-CSH-{k:03d}", "date": "09-Oct-2026", "by": "iorta TechNXT Project Team", "brd": c[1],
                     "frs": c[0], "desc": f"{c[2]}: {c[3]}", "reason": c[4], "type": c[5], "scope": c[6],
                     "schedule": c[7], "cost": c[8], "other": c[9], "tests": c[10], "priority": c[11],
                     "status": "Candidate - for decision", "decision": "", "approved": "", "ddate": "", "drop": "",
                     "linked": ", ".join(data.cases_of_item(c[0])[:5])})
        if c[0] not in data.item_ids and not any(i.startswith(c[0] + ".") for i in data.item_ids):
            raise SystemExit(f"change candidate names an unknown FR item {c[0]}")
    wb.sheet("Register", [
        Column("no", "CR No.", 12, "Change request number"), Column("date", "Date raised", 12, "dd-Mmm-yyyy"),
        Column("by", "Raised by", 18, "Name and role"), Column("brd", "BRD ID", 14, "BRD requirement"),
        Column("frs", "FRS reference", 18, "FR item or appendix"), Column("desc", "Description", 46, "The change"),
        Column("reason", "Reason", 36, "Why"), Column("type", "Type", 18, "Kind of change", values=WB["cr_types"]),
        Column("scope", "Impact on scope", 22, ""), Column("schedule", "Impact on schedule", 14, ""),
        Column("cost", "Impact on cost", 12, ""), Column("other", "Impact on other BRDs", 18, ""),
        Column("tests", "Impact on tests", 18, ""), Column("priority", "Priority", 10, "", values=["High", "Medium", "Low"]),
        Column("status", "Status", 20, "", values=WB["cr_status"], status=True),
        Column("decision", "Decision", 24, "Decision of the approvers"), Column("approved", "Approved by", 18, ""),
        Column("ddate", "Decision date", 12, ""), Column("drop", "Target drop", 12, ""),
        Column("linked", "Linked test cases", 30, "Test cases of the Test Cases and Traceability workbook")],
        rows, description=f"Candidate changes known at v{VERSION} (status Candidate - for decision) and new requests")
    form_cols = [
        {"key": "field", "header": "Field", "width": 30, "mandatory": True, "what": "Part of the assessment"},
        {"key": "value", "header": "Assessment", "width": 70, "mandatory": True, "what": "Filled in by the project"},
        {"key": "by", "header": "Assessed by", "width": 20, "what": "Name"}]
    fields = ["CR No.", "Summary of the change", "BRD IDs and FR items affected", "Screens, messages, e-mails affected",
              "Other BRDs affected", "Effort (mandays)", "Schedule impact (drop)", "Cost",
              "Test cases to add or change", "Training and documents", "Risks", "Recommendation"]
    wb.template_sheet("Impact assessment form", "Impact assessment form", [
        ("Purpose", "The assessment of one change request before the decision."),
        ("Who fills it in", "iorta TechNXT Project Team with the System Analyst (BDO ITG)"),
        ("How it is used", "Copied for each change; its conclusion goes to the Register."),
        ("Due", "Within 5 working days of the request."), ("Depends on", "Register")],
        form_cols, [{"field": f, "value": "", "by": ""} for f in fields],
        {"field": "CR No.", "value": "CR-CSH-001", "by": "Project analyst"})
    dec = [{"no": q["no"], "section": q["section"], "topic": q["topic"], "proposal": q["proposal"],
            "by": q["decide_by"], "decision": "", "dby": "", "ddate": "", "cr": "", "status": "Open"}
           for q in data.b.q_rows]
    if len(dec) != data.b.stats["observations"]:
        raise SystemExit("decision log differs from Appendix R")
    wb.sheet("Decisions log", [
        Column("no", "Appendix R No.", 16, f"Row of Appendix R of the FRS v{VERSION}"), Column("section", "Section", 8, ""),
        Column("topic", "Topic", 26, ""), Column("proposal", "Proposed in the FRS", 60, ""),
        Column("by", "Decision by", 26, ""), Column("decision", "Decision", 30, "Adopt proposal / Keep BDOI text / Other"),
        Column("dby", "Decided by (name)", 18, ""), Column("ddate", "Date", 12, ""),
        Column("cr", "CR No. (if a change follows)", 14, ""),
        Column("status", "Status", 12, "", values=["Open", "Decided", "Withdrawn"], status=True)],
        dec, description="One row per observation of Appendix R, for the decision of the Business Unit")
    path = wb.save(FOLDER / "BIBS_CR_BRD-02_Change_Request_Register_v1.0.xlsx")
    return path, {"candidates": len(rows), "decisions": len(dec)}


# ------------------------------------------------------------------------------------------------- 4 RTM
def answered_rows(data: Data) -> list[dict]:
    out = []
    for ref, (clause, answer) in REV["open_items"]["answered"].items():
        q = (data.inputs.get(ref) or {}).get("need") or frs_table_text(ref)
        out.append({"ref": ref, "q": frs.clean(q) if q else "-", "clause": clause, "answer": answer,
                    "kept": "Appendix R decision" if re.search(r"Appendix R|\bC\d{1,2}\b", answer) else "-"})
    return out


def build_rtm(data: Data) -> tuple[Path, dict]:
    cases = data.acc_cases + data.bibs_cases
    case_ids_by_fr = defaultdict(list)
    for c in data.bibs_cases:
        case_ids_by_fr[c["fr"]].append(c["id"])
    trace_rows = []
    for rid in data.ids:
        items = sorted(set(FRS_RE.findall(data.mapping.get(rid, ""))), key=frs.frs_key)
        frops = data.trace.get(rid, {"frs": []})["frs"]
        ids = {i for f in frops for i in case_ids_by_fr[f]}
        ids |= {c["id"] for c in data.acc_cases for it in items if under(data.item_by_id[c["item"]], it)}
        steps = sorted({s for f in frops for s in data.steps_by_fr.get(f, [])}, key=step_key)
        note = "Removed from the BRD" if rid in frs.REMOVED else ("Specified in another FRS; held by reference"
                                                                    if rid in frs.BY_REFERENCE else "")
        trace_rows.append({"id": rid, "text": data.brd_text.get(rid, ""), "item": ", ".join(items) or note or "-",
                           "fr": ", ".join(frops) or "-", "n": len(ids), "cases": ", ".join(sorted(ids)) or note,
                           "screens": data.screens_of(frops), "steps": ", ".join(steps) or "-"})
    n_brd = len(trace_rows)
    for r in REV["nfr"]["rows"]:
        its = [x["id"] for x in data.items if any(under(x, i) for i in FRS_RE.findall(r[4]))]
        ids = sorted({c for i in its for c in data.cases_of_item(i)})
        cases_text = (", ".join(ids[:15]) + (" ..." if len(ids) > 15 else "")) if ids else \
            f"Non-functional test of the programme ({r[5]})"
        trace_rows.append({"id": r[0], "text": f"{r[1]}: {r[2]}", "item": r[4], "fr": "-", "n": len(ids),
                           "cases": cases_text, "screens": "-", "steps": "-"})
    no_case = [r["id"] for r in trace_rows if r["n"] == 0 and r["id"] not in frs.REMOVED and r["item"] != "-"]
    if no_case:
        raise SystemExit(f"BRD IDs without a test case: {no_case}")
    item_rows = []
    for it in data.items:
        ids = data.cases_of_item(it["id"])
        if not ids:
            raise SystemExit(f"{it['id']} has no test case")
        item_rows.append({"id": it["id"], "title": it["title"], "origin": it["origin"] +
                          (" (elaborated in v3.2)" if it["elaborated"] else ""), "ac": len(it["ac"]),
                          "n": len(ids), "cases": ", ".join(ids[:12]) + (" ..." if len(ids) > 12 else "")})
    expected = sum(n for fr, (c, n) in frs.test_cases().items() if fr in frs.SCOPE_FRS) + data.b.stats["ac"]
    if len(cases) != expected:
        raise SystemExit(f"test case count {len(cases)} differs from the FRS ({expected})")
    wb = ReviewWorkbook("Test Cases and Traceability", doc_type="Requirements traceability", brd="BRD-02",
                        version=META["version"], subtitle=f"BRD-02 Operations Cashiering - FRS in BDOI's template v{VERSION}")
    in_scope = [r for r in trace_rows[:n_brd] if r["id"] not in frs.REMOVED]
    ans = answered_rows(data)
    summary = [
        {"what": "BRD IDs of the Cashiering scope", "n": n_brd,
         "proof": f"{sum(1 for r in in_scope if r['n'])} of the {len(in_scope)} in force with at least one test case; "
                  f"{len(frs.REMOVED)} removed from the BRD (sheet Traceability)"},
        {"what": "Non-functional rows", "n": len(trace_rows) - n_brd,
         "proof": f"{sum(1 for r in trace_rows[n_brd:] if r['n'])} with test cases of their FR items; the others "
                  f"(availability, recovery) in the non-functional test of the programme"},
        {"what": "FR items (BDOI's and added)", "n": len(data.items),
         "proof": f"{sum(1 for r in item_rows if r['n'])} with at least one test case (sheet FR items)"},
        {"what": "Acceptance criteria", "n": data.b.stats["ac"], "proof": "One acceptance test case each"},
        {"what": "Test cases", "n": len(cases),
         "proof": f"{len(data.bibs_cases)} of the BIBS test plan BRD-2 re-keyed to the FR items; {len(data.acc_cases)} acceptance cases"},
        {"what": "Positive / negative", "n": f"{sum(1 for c in cases if c['type'] == 'Positive')} / "
                                             f"{sum(1 for c in cases if c['type'] == 'Negative')}", "proof": "-"},
        {"what": "Questions answered in the BDOI FRS", "n": len(ans), "proof": "Sheet Answered in the BDOI FRS"},
    ]
    wb.sheet("Summary", [Column("what", "Item", 40, ""), Column("n", "Count", 14, ""), Column("proof", "Proof", 80, "")],
             summary, description="Counts and the proof that every BRD ID and every FR item is covered")
    wb.sheet("Traceability", [
        Column("id", "BRD ID", 14, "BRD requirement ID"), Column("text", "BRD requirement", 46, ""),
        Column("item", "FR item (BDOI format)", 26, f"FR items of the FRS v{VERSION}"),
        Column("fr", "BIBS reference FR", 18, "FR of the BIBS FRS BRD-2"), Column("n", "Test cases", 10, "Count"),
        Column("cases", "Test case IDs", 60, "Sheet Test cases"), Column("screens", "Screen", 34, "Screens of Appendix E"),
        Column("steps", "Walkthrough steps", 26, "Steps of the walkthrough users workbook")],
        trace_rows, description="BRD ID -> FR item -> BIBS reference FR -> test cases -> screen -> walkthrough step")
    wb.sheet("FR items", [Column("id", "FR item", 22, ""), Column("title", "Title", 44, ""), Column("origin", "Origin", 22, ""),
                          Column("ac", "Acceptance criteria", 12, ""), Column("n", "Test cases", 10, ""),
                          Column("cases", "Test case IDs", 70, "")],
             item_rows, description=f"Every FR item of the FRS v{VERSION} with its test cases")
    case_cols = [
        Column("id", "Test case ID", 28, "TC-OP-nnn.c-ss: BIBS test plan; TC-FRS-...-ACnn: acceptance criterion"),
        Column("title", "Title", 36, ""), Column("brd", "BRD ID", 22, ""), Column("item", "FR item", 22, ""),
        Column("fr", "BIBS reference FR", 14, ""), Column("persona", "Persona", 18, ""),
        Column("user", "Sign-in ID", 12, "The password is issued separately by the iorta TechNXT Project Team"),
        Column("screen", "Screen", 22, ""), Column("pre", "Preconditions", 30, ""), Column("steps", "Steps", 50, ""),
        Column("data", "Test data", 22, ""), Column("expected", "Expected result", 50, ""),
        Column("type", "Type", 10, "", values=["Positive", "Negative"]),
        Column("priority", "Priority", 10, "", values=["High", "Medium", "Low"]),
        Column("result", "Result", 10, "", values=RESULT, status=True), Column("tester", "Tested by", 14, ""),
        Column("date", "Date", 12, ""), Column("obs", "Observation ref", 14, ""), Column("remarks", "Remarks", 24, ""),
        Column("source", "Source", 22, "")]
    for c in cases:
        c.update({"result": "Not run", "tester": "", "date": "", "obs": "", "remarks": ""})
    wb.sheet("Test cases", case_cols, cases, description=f"Test cases of the FRS v{VERSION} with the execution columns")
    wb.sheet("Answered in the BDOI FRS", [
        Column("ref", "Our reference", 14, "OQ, CLR-OP"), Column("q", "Our question", 60, ""),
        Column("clause", "BDOI FRS clause", 30, ""), Column("answer", "Answer in BDOI's FRS", 60, ""),
        Column("kept", "Still a decision?", 18, "Where BDOI's answer differs from the BRD, the decision stays in Appendix R")],
        ans, description="Open items dropped from the FRS because BDOI's FRS answers them")
    if len(ans) != data.b.stats["answered"]:
        raise SystemExit("answered questions differ from Appendix S")
    crit = [{"kind": "Entry", "no": i, "criterion": c} for i, c in enumerate(WB["uat_criteria"]["entry"], 1)]
    crit += [{"kind": "Exit", "no": i, "criterion": c} for i, c in enumerate(WB["uat_criteria"]["exit"], 1)]
    wb.sheet("UAT entry and exit", [Column("kind", "Entry / exit", 12, ""), Column("no", "No.", 6, ""),
                                    Column("criterion", "Criterion", 100, "")],
             crit, description="UAT entry and exit criteria of BRD-02 Operations Cashiering")
    path = wb.save(FOLDER / "BIBS_RTM_BRD-02_Test_Cases_and_Traceability_v1.0.xlsx")
    return path, {"trace": len(trace_rows), "cases": len(cases), "bibs": len(data.bibs_cases),
                  "acc": len(data.acc_cases), "answered": len(ans), "items": len(item_rows)}


# ------------------------------------------------------------------------------------------------- 5 UAT users
def build_uat(data: Data) -> tuple[Path, dict]:
    src = OUT / "Programme" / "UAT" / "BIBS_UAT_BRD-00_Walkthrough_Users_and_Sign-in_v1.0.xlsx"
    w = openpyxl.load_workbook(src, read_only=True)
    spec = next(f for f in frs.FLOWS if f["id"] == "brd02_fig_menus")
    menus = annexes.menu_rows(spec)
    personas = spec["menu_personas"]
    script = []
    for r in data.a.walk_rows:
        script.append({"wt": r["step"].rsplit(".", 1)[0], "step": r["step"], "persona": r["persona"], "signin": r["user"],
                       "screen": r["screen"], "action": r["does"], "expected": f"{r['sees']} {r['result']}".strip(),
                       "result": "Not run", "tester": "", "date": "", "remarks": ""})
    steps_of = defaultdict(list)
    for s in script:
        steps_of[s["signin"]].append(s["step"])
    users = []
    head = None
    for r in w["BRD-02 Operations"].iter_rows(values_only=True):
        if r and r[0] == "Persona":
            head = list(r)
            continue
        if head and r and r[0]:
            rw = dict(zip(head, r))
            sid = rw["Sign-in ID"]
            if sid not in WB["uat_signins"]:
                continue
            pname = WB["uat_menu_persona"][sid]
            if pname in personas:
                col = personas.index(pname) + 1
                menu = "; ".join(m[0] for m in menus if m[col])
            else:
                menu = "No Cashiering menu entry; the menus of the user's own team"
            users.append({"persona": pname, "name": rw["User name"], "signin": sid, "profiles": rw["Group profile"],
                          "unit": rw["Company / branch / unit"], "does": rw["What this user does in the walkthrough"],
                          "start": rw["Menu path to start"], "menus": menu,
                          "steps": ", ".join(steps_of.get(sid, [])) or "Not in a walkthrough step",
                          "password": META["password_text"]})
    w.close()
    found = {u["signin"] for u in users}
    if set(WB["uat_signins"]) - found:
        raise SystemExit(f"walkthrough users not found: {set(WB['uat_signins']) - found}")
    missing = {s["signin"] for s in script} - found - {"-"}
    if missing:
        raise SystemExit(f"script users without a row: {missing}")
    if len(script) != data.b.stats["walk_steps"]:
        raise SystemExit("walkthrough steps differ from Appendix K")
    wb = ReviewWorkbook("Business Unit Walkthrough Users", doc_type="UAT walkthrough", brd="BRD-02",
                        version=META["version"], subtitle="BRD-02 Operations Cashiering - sign-in IDs for the walkthrough")
    wb.cover_notes = [f"Environment: {META['environment']}. Passwords: {META['password_text']}; no password is "
                      f"written in this workbook."]
    wb.sheet("Users", [
        Column("persona", "Persona", 26, "Business role of the FRS"), Column("name", "User name", 22, "As shown on the screens"),
        Column("signin", "Sign-in ID", 12, "User ID to type on the Login page"),
        Column("profiles", "Group profile", 26, "Role held"), Column("unit", "Company / branch / unit", 32, ""),
        Column("does", "What the user does in the walkthrough", 40, ""), Column("start", "Menu path to start", 30, ""),
        Column("menus", "Menus the user should see", 50, "Cashiering entries of the user's menu (Figure of Appendix E)"),
        Column("steps", "Walkthrough steps", 24, "Steps of the sheet Walkthrough script"),
        Column("password", "Password", 28, "Never written here")],
        users, description=f"{META['environment']}. Password: {META['password_text']}")
    wb.sheet("Walkthrough script", [
        Column("wt", "Walkthrough", 10, ""), Column("step", "Step", 9, ""), Column("persona", "Persona", 22, ""),
        Column("signin", "Sign-in ID", 12, ""), Column("screen", "Screen", 28, "Screen of Appendix E"),
        Column("action", "Action", 50, ""), Column("expected", "Expected result", 56, ""),
        Column("result", "Result", 10, "", values=RESULT, status=True),
        Column("tester", "Done by", 14, ""), Column("date", "Date", 12, ""), Column("remarks", "Remarks", 24, "")],
        script, description="The Cashiering steps of the walkthroughs and the walkthrough WT-E for the Business Unit "
                            "(FRS Appendix K)")
    path = wb.save(FOLDER / "BIBS_UAT_BRD-02_Business_Unit_Walkthrough_Users_v1.0.xlsx")
    return path, {"users": len(users), "steps": len(script)}


# ------------------------------------------------------------------------------------------------- 6 coverage
def build_coverage(data: Data) -> tuple[Path, dict]:
    item_rows, ac_rows = [], []
    for it in data.items:
        cmp = data.cmp_status.get(it["id"])
        refs = list(dict.fromkeys(it["refs"] + (cmp[2] if cmp else [])))
        conf = [x for f in refs for x in data.conformance.get(f, [])]
        evidence = []
        if conf:
            evidence.append("Conformance register: " + "; ".join(
                f"{f} {dict(Counter(x[1] for x in data.conformance.get(f, [])))}" for f in refs if data.conformance.get(f)))
        steps = sorted({s for f in refs for s in data.steps_by_fr.get(f, [])}, key=step_key)
        if steps:
            evidence.append("Walkthrough " + ", ".join(steps[:8]))
        if it["id"] in WB["coverage_items"]:
            status, note = WB["coverage_items"][it["id"]]
        elif cmp:
            status = cmp[0] if cmp[0] in ("Yes", "Partly", "No") else "Partly"
            note = "Comparison of BDOI's FRS with the platform: " + (cmp[1] or "as specified.")
        elif not refs:
            status, note = "Partly", "Heading item: see its sub-items."
        elif conf and all(x[1] in ("Conformant", "Partly conformant") for x in conf):
            status, note = "Yes", "Shown working in the end-to-end run; criteria without an automated check are confirmed in UAT."
        else:
            status, note = "Partly", "Not reached by the end-to-end run; tested in UAT."
        cases = data.cases_of_item(it["id"])
        item_rows.append({"group": it["entry"], "id": it["id"], "title": it["title"], "origin": it["origin"],
                          "refs": ", ".join(refs) or "-", "status": status,
                          "evidence": f"{len(cases)} test cases" + ("; " + " | ".join(evidence) if evidence else ""),
                          "note": note})
        for n, ac in enumerate(it["ac"], 1):
            ast = status
            anote = "Available; confirmed with the acceptance test case in UAT." if status == "Yes" else note
            key = it["id"].replace("FRS.", "FRS-").replace(" (2nd)", "-2ND")
            ac_rows.append({"group": it["entry"], "id": it["id"], "no": n, "ac": ac, "status": ast,
                            "evidence": f"TC-{key}-AC{n:02d}", "note": anote})
    acc_ids = {c["id"] for c in data.acc_cases}
    if any(r["evidence"] not in acc_ids for r in ac_rows):
        raise SystemExit("coverage names an unknown acceptance case")

    def pct(rows):
        c = Counter(r["status"] for r in rows)
        n = len(rows)
        return c, (round(100 * c["Yes"] / n, 1) if n else 0), (round(100 * (c["Yes"] + c["Partly"]) / n, 1) if n else 0)
    summ = []
    groups = list(dict.fromkeys(r["group"] for r in item_rows))
    for g in ["All"] + groups:
        ir = [r for r in item_rows if g == "All" or r["group"] == g]
        ar = [r for r in ac_rows if g == "All" or r["group"] == g]
        ci, yi, pi = pct(ir)
        ca, ya, pa = pct(ar)
        summ.append({"group": "All FR items" if g == "All" else g, "items": len(ir), "iy": ci["Yes"], "ip": ci["Partly"],
                     "in": ci["No"], "ipct": yi, "ippct": pi, "acs": len(ar), "ay": ca["Yes"], "ap": ca["Partly"],
                     "an": ca["No"], "apct": ya, "appct": pa})
    wb = ReviewWorkbook("FRS Coverage in BIBS", doc_type="Internal coverage", brd="BRD-02", version=META["version"],
                        subtitle=f"BRD-02 Operations Cashiering - FRS in BDOI's template v{VERSION} (internal)",
                        classification="Internal - iorta TechNXT")
    wb.legend = [("Yes", "Available in the BIBS environment"), ("Partly", "Partly available (reason in the note)"),
                 ("No", "Planned for a later release or waiting for a decision (reason in the note)")]
    num = lambda k, h, w=10: Column(k, h, w, "", kind="number")  # noqa: E731
    wb.sheet("Summary", [Column("group", "FR group", 22, "FR entry"), num("items", "FR items"), num("iy", "Yes"),
                         num("ip", "Partly"), num("in", "No"), num("ipct", "% items available", 14),
                         num("ippct", "% items available or partly", 16), num("acs", "Acceptance criteria", 14),
                         num("ay", "Yes"), num("ap", "Partly"), num("an", "No"), num("apct", "% criteria available", 14),
                         num("appct", "% criteria available or partly", 16)],
             summ, description=f"Coverage of the FRS v{VERSION} in the BIBS environment, by item and by acceptance criterion")
    st = ["Yes", "Partly", "No"]
    wb.sheet("FR items", [Column("group", "FR group", 16, ""), Column("id", "FR item", 22, ""), Column("title", "Title", 40, ""),
                          Column("origin", "Origin", 14, ""), Column("refs", "BIBS reference FR", 22, ""),
                          Column("status", "Available", 12, "", values=st, status=True),
                          Column("evidence", "Evidence", 60, ""), Column("note", "Note", 60, "")],
             item_rows, description=f"Every FR item and sub-item of the FRS v{VERSION}")
    wb.sheet("Acceptance criteria", [Column("group", "FR group", 16, ""), Column("id", "FR item", 22, ""),
                                     Column("no", "AC", 6, ""), Column("ac", "Acceptance criterion", 60, ""),
                                     Column("status", "Available", 12, "", values=st, status=True),
                                     Column("evidence", "Test case", 30, ""), Column("note", "Note", 50, "")],
             ac_rows, description=f"Every acceptance criterion of the FRS v{VERSION}")
    path = wb.save(QUALITY / "BIBS_Coverage_BRD-02_FRS_Coverage_in_BIBS_v1.0.xlsx")
    return path, {"items": len(item_rows), "acs": len(ac_rows), "all": summ[0]}


def main() -> int:
    data = Data()
    FOLDER.mkdir(parents=True, exist_ok=True)
    for name, fn in (("requirements", build_requirements), ("fitgap", build_fitgap), ("cr", build_cr),
                     ("rtm", build_rtm), ("uat", build_uat), ("coverage", build_coverage)):
        path, stats = fn(data)
        print(f"wrote {path.relative_to(REPO)} ({path.stat().st_size // 1024} KB): {stats}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
