"""Product Maintenance document set v1.6 (BRD-03) in BDOI's template: FRS, BPS, UXD, RNO and the Interface
Specification v1.0.

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd03_v16.py [--no-pdf] [--only FRS,BPS,UXD,RNO,ICD]

Each document starts from BDOI's FRS template (template D051, templates/bdoi_d051_frs_template.docx) through
pack_common.PackDoc: cover, document details, headers and footers, styles, revision log and table of contents are
kept; the body is written from brd03_v16_content.py (FRS, ICD), brd03_v16_pages.py (process, storyboards, outputs)
and brd03_v16_screens.py (screens). The screenshots are in figures/brd03_v14_screens (flows/brd03/shots_v14.sh),
one full page each.

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-03_Product_Maintenance/ (PDFs not kept in git).
Self-checks: every screenshot named exists; every BRD row of the mapping points to an FRPM item; every screen of the
UX Design has a screenshot; no wording refused in client documents.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[4] / "tools" / "deliverables"))
import brd03_v16_content as C  # noqa: E402
import brd03_v16_pages as P  # noqa: E402
import brd03_v16_screens as S  # noqa: E402
import pack_common as pc  # noqa: E402
from build_brd11_interface_spec import _run  # noqa: E402
from colour_code import COLOURS, LABELS  # noqa: E402

OUT = pc.OUT_BDOI / "BRD-03_Product_Maintenance"
SHOTS = HERE / "figures" / "brd03_v14_screens"
FLOWS = HERE / "flows" / "brd03"
COL = {"k": None, "g": COLOURS["green"], "o": COLOURS["orange"], "b": COLOURS["blue"], "": None}
COLOUR_NAME = {"k": "Black", "g": "Green", "o": "Orange", "b": "Blue", "": "-"}
CREATED = {"FRS": "08/19/2026"}


def _colour(run, hexcol: str) -> None:
    rpr = run._r.get_or_add_rPr()
    for c in rpr.findall(qn("w:color")):
        rpr.remove(c)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), hexcol)
    rpr.append(c)


class Doc(pc.PackDoc):
    """A document of the Product Maintenance set."""

    def __init__(self, key: str, revisions):
        m = C.DOCS[key]
        version = {"ICD": C.ICD_VERSION, "UXD": C.UXD_VERSION}.get(key, C.VERSION)
        super().__init__(pc.DocMeta(file=m["file"], cover=m["cover"], header_left=m["header_left"],
                                    header_right=C.HEADER_RIGHT, doc_no=m["doc_no"], version=version,
                                    date_us=C.DATE_US, created_us=CREATED.get(key, C.DATE_US), revisions=revisions))
        self.key = key
        self.shots = 0

    def cpara(self, text: str, colour: str = "k") -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.64)
        p.paragraph_format.space_after = Pt(4)
        r = _run(p, text)
        if COL.get(colour):
            _colour(r, COL[colour])

    def ctable(self, headers, rows, widths, caption=None, colours=None, colour_cols=None):
        self.table(headers, rows, widths, caption=caption)
        tbl = self.doc.tables[-1]
        for i, code in enumerate(colours or []):
            if not COL.get(code):
                continue
            for c_i, cell in enumerate(tbl.rows[i + 1].cells):
                if colour_cols is None or c_i in colour_cols:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            _colour(r, COL[code])
        return tbl

    def screen(self, shot: str, caption: str) -> None:
        """One full-page screenshot on its own page, scaled to fit the page, never split."""
        png = SHOTS / f"{shot}.png"
        if not png.exists():
            raise SystemExit(f"missing screenshot {png.name}")
        w, h = Image.open(png).size
        width = min(16.0, 21.5 * w / h)
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        p.alignment = 1
        p.add_run().add_picture(str(png), width=Cm(width))
        self.caption(caption, kind="Figure")
        self.shots += 1

    def legend(self) -> None:
        self.para(C.LEGEND_NOTE)
        rows = [["Black", "In the BRD", "Asked by the BRD, BRD-00 or a programme NFR.", "In scope"],
                ["Green", "Not in the BRD: OOTB", "Offered by the iNXT BrokerVerse OOTB platform.", "No cost"],
                ["Orange", "Not in the BRD: change request", "Not offered by the OOTB platform.", "Only if approved"],
                ["Blue", "Not in the BRD: best practice", "Recommendation, separate table.", "Indicative effort"],
                ["Grey label", LABELS["integration"] + " / " + LABELS["migration"],
                 "Interfaces with other BDOI systems, or Data Migration (BRD-13).", "Not a change request"]]
        self.ctable(["Colour", "Meaning", "What it covers", "Cost"], rows, [1.2, 2.4, 4.2, 1.6],
                    caption="Colour code", colours=["k", "g", "o", "b", ""], colour_cols={0, 1})
        tbl = self.doc.tables[-1]
        for r in tbl.rows[5].cells[0].paragraphs[0].runs + tbl.rows[5].cells[1].paragraphs[0].runs:
            _colour(r, COLOURS["grey"])

    def personas(self, caption="Personas of Product Maintenance") -> None:
        self.table(["Persona", "Does", "Screens", "Walkthrough account"], C.PERSONAS, [1.8, 3.6, 2.6, 1.6],
                   caption=caption)


def _dot(src: str, name: str, tmp: Path) -> Path:
    dot = tmp / f"{name}.dot"
    dot.write_text(src, encoding="utf-8")
    png = tmp / f"{name}.png"
    subprocess.run(["dot", "-Tpng", "-Gdpi=170", str(dot), "-o", str(png)], check=True)
    return png


def _lifecycle(tmp: Path) -> Path:
    return _dot(
        'digraph g { rankdir=LR; graph [fontname="Arial", nodesep=0.25, ranksep=0.3]; node [shape=box, '
        'style="rounded", fontname="Arial", fontsize=10]; edge [fontname="Arial", fontsize=8];\n'
        'D [label="Draft request"]; M [label="For Marketing\\napproval"]; R [label="For TSU review"]; '
        'H [label="For TSU Head\\napproval"]; N [label="Insurer\\nnegotiation"]; T [label="Negotiated\\nterms review"];'
        ' K [label="Marketing review\\nof terms"]; Q [label="Requirements\\npreparation"]; '
        'C [label="For ManCom\\nsign-off"]; B [label="With MBS\\nfor set-up"]; V [label="Package version\\nfor '
        'validation"]; L [label="Released"]; X [label="Not proceeded /\\nVoided / Retired"];\n'
        'D -> M [label="Submit"]; D -> R [label="Submit (TSU)"]; M -> R [label="Approve"]; R -> H [label='
        '"Recommend"]; H -> N [label="Approve"]; N -> T [label="Terms final"]; T -> K [label="client-specific"]; '
        'T -> Q; K -> Q [label="Accept"]; K -> N [label="Changes"]; Q -> C [label="Submit"]; C -> B [label='
        '"Approved"]; C -> Q [label="Return"]; B -> V [label="Submit version"]; B -> Q [label="Return"]; '
        'V -> L [label="Validate"]; V -> B [label="Return"]; N -> X; H -> X; D -> X; }', "lifecycle", tmp)


def _pair(rows):
    return [list(r) for r in rows]


# ================================================================================================ FRS
def write_frs(tmp: Path) -> Doc:
    d = Doc("FRS", C.REVISIONS)
    i = C.FRS_INTRO
    d.h1("Introduction")
    d.h2("Document Purpose")
    d.para(i["purpose"])
    d.h2("Intended Audience")
    d.para(i["audience"])
    d.h2("Project Background")
    for t in i["background"]:
        d.para(t)
    d.h2("Scope")
    n = max(len(i["in_scope"]), len(i["out_scope"]))
    d.table(["In scope (BIBS)", "Out of scope or in another FRS"],
            [[(i["in_scope"] + [""] * n)[k], (i["out_scope"] + [""] * n)[k]] for k in range(n)], [4.6, 4.6],
            caption="Scope of Product Maintenance in BIBS")
    d.h2("Document Set")
    d.table(["Document", "Content"], C.DOC_SET, [3.2, 6.0], caption="Product Maintenance document set")
    d.h2("How to Read This Document")
    d.legend()
    d.h2("References")
    d.bullets(C.REFERENCES)

    d.h1("Business Requirements Mapping")
    d.para("Every BRD line of Product Maintenance maps to an item of section 3. The last column gives the change "
           "from version 1.3.")
    d.table(["BRD Req No.", "Business requirement", "FRPM items", "Where", "Change from v1.3"], C.MAPPING,
            [1.7, 2.8, 2.0, 1.3, 1.5], caption="Business requirements mapping")

    d.h1("Functional Requirements")
    d.para("Each item keeps BDOI's FRPM number. Status gives the change from version 1.3: Kept, Changed or New. The "
           "Example column holds the acceptance example; the full criteria are in the User Stories v1.0.")
    d.table(["FRPM", "Component", "Capability", "Status", "BRD"],
            [[f["id"], f["component"], f["capability"], f["status"], f["brd"]] for f in C.FRUM],
            [1.3, 1.8, 3.0, 1.1, 2.1], caption="Functional requirements at a glance")
    for f in C.FRUM:
        d.h2(f"{f['id']} {f['capability']}")
        d.para(f"Status: {f['status']}. Version 1.3: {f['was']}. BRD: {f['brd']}.", italic=True)
        rows, cols = [], []
        for rid, text, colour, label, rule, ex in f["reqs"]:
            lab = {"int": "  " + LABELS["integration"], "dm": "  " + LABELS["migration"]}.get(label, "")
            rows.append([rid, text + lab, rule, ex])
            cols.append(colour)
        d.ctable(["ID", "Requirement", "Rule", "Example"], rows, [1.8, 4.2, 1.9, 1.9], colours=cols,
                 colour_cols={1, 2})

    d.h1("Business Rules")
    d.ctable(["ID", "Rule", "Colour"], [[a, b, COLOUR_NAME[c]] for a, b, c in C.BUSINESS_RULES], [1.3, 6.8, 1.2],
             caption="Business rules", colours=[c for *_, c in C.BUSINESS_RULES], colour_cols={1, 2})
    d.h1("Validations and Messages")
    d.para("The main messages of each screen. The full list with the message codes of the services is in the "
           "Developer Specification, Annex A; the field checks of each screen are in the UX Design.")
    d.table(["Screen or step", "Condition", "Message"], C.MESSAGES, [2.2, 3.1, 3.9], caption="Validations and "
            "messages")
    d.h1("Personas and Permissions")
    d.personas()
    d.table(C.PERMISSION_HEADS, C.PERMISSIONS, [3.0] + [0.68] * 9, caption="What each persona may do (Y = yes); "
            "the segregation of duties of BR-PM-01 applies whatever the profile grants")
    d.h1("Parameters and Lists of Values")
    d.para("Settings are changed on System Parameters and applied after another user approves the change. Keys, "
           "types and ranges are in the Data Dictionary, sheet Parameters.")
    d.ctable(["Setting", "Default", "Range", "Area"], [r[:4] for r in C.PARAMETERS], [4.0, 2.4, 1.6, 1.2],
             caption="Product Maintenance settings", colours=[r[4] for r in C.PARAMETERS])
    d.table(["List", "Values"], C.LOVS, [3.0, 6.2], caption="Lists of values (codes in the Data Dictionary)")
    d.h1("Statuses")
    d.figure(_lifecycle(tmp), "Stages of a package request", width_cm=16.0)
    for name, rows in C.STATUSES.items():
        d.h2(name)
        d.table(["Status", "Meaning", "Next"], rows, [2.4, 3.6, 3.2], caption=f"Statuses: {name.lower()}")
    d.h2("BDOI's status names and the stages")
    d.table(["BDOI's status (FRPM.011.01, FRPM.016.01)", "Stage in BIBS"], C.BDOI_STATUS_MAP, [5.0, 4.2],
            caption="Status names to agree (Q-PM-09)")
    d.h1("Assumptions, Dependencies and Open Questions")
    d.table(["ID", "Assumption"], C.ASSUMPTIONS, [1.3, 7.9], caption="Assumptions")
    d.table(["ID", "Dependency"], C.DEPENDENCIES, [1.3, 7.9], caption="Dependencies")
    d.table(["ID", "Question", "Owner", "Answer by"], C.OPEN_QUESTIONS, [1.2, 5.0, 1.8, 1.2],
            caption="Open questions")
    d.h1("Change Control Summary")
    d.para(f"Change requests of the register v{C.WB_VERSION}. Effort in man-days (analysis and design, backend, "
           "frontend, testing, documentation and UAT support).")
    d.ctable(["CR", "Title", "Status", "FRS", "Man-days", "Note"],
             [[a, b, c, e, f"{sum(eff):g}", n] for a, b, c, e, eff, n in C.CRS], [1.3, 2.9, 0.9, 1.2, 0.9, 2.0],
             caption="Change requests", colours=["o" if c == "Open" else "" for _, _, c, *__ in C.CRS])
    d.ctable(["No.", "Not in the BRD, offered by the OOTB platform", "FRS"], C.OOTB_POINTS, [1.3, 5.6, 2.3],
             caption="Green points (no cost)", colours=["g"] * len(C.OOTB_POINTS), colour_cols={1})
    d.table(["No.", "Integration and data migration scope", "FRS"], C.INTEGRATION_POINTS, [1.3, 5.6, 2.3],
            caption="Grey label points (not change requests)")
    d.ctable(["No.", "Best-practice recommendation", "Priority", "Man-days"],
             [[a, b, c, str(e)] for a, b, c, e in C.BEST_PRACTICE], [1.3, 5.0, 1.4, 1.4],
             caption="Blue points (recommendations)", colours=["b"] * len(C.BEST_PRACTICE), colour_cols={1})
    d.h1("Changes from Version 1.3")
    d.table(["Item", "Change"], C.CHANGES_V13, [3.0, 6.2], caption="Changes from version 1.3")
    return d


# ================================================================================================ BPS
def write_bps(tmp: Path) -> Doc:
    d = Doc("BPS", C.REVISION_COMPANION)
    d.h1("Introduction")
    d.para(f"This document shows how Product Maintenance works end to end. The requirements are in the FRS "
           f"v{C.VERSION}; the screens in the UX Design v{C.VERSION}; the reports and notices in the Reports, "
           f"Notifications and Outputs v{C.VERSION}.")
    d.para("Each storyboard step is illustrated by a full-page screenshot taken on a test environment. The people, "
           "clients and insurers are fictitious.")
    d.personas("Personas")
    d.h1("End-to-end Journey")
    d.para("Fifteen steps from the package request to the end of the package.")
    d.figure(FLOWS / "pm_simple.png", "Product Maintenance: simplified process flow (swimlanes by actor)",
             width_cm=16.0)
    d.table(["Step", "Who", "What happens"], P.JOURNEY, [0.9, 2.0, 6.3], caption="The 15 steps")
    d.figure(FLOWS / "pm_e2e.png", "Product Maintenance: end-to-end flow by actor and system", width_cm=16.0)
    d.h1("To-be Processes")
    for p in P.PROCESSES:
        d.h2(f"{p['id']} {p['name']}")
        d.table(["Item", "Value"], [["Journey steps", p["steps"]], ["Trigger", p["trigger"]], ["Actors", p["actors"]],
                                    ["Outcome", p["outcome"]], ["Controls", p["controls"]]], [2.0, 7.2])
    d.h2("Stages of a package request")
    d.figure(_lifecycle(tmp), "Stages of a package request", width_cm=16.0)
    d.h1("Storyboards")
    for sb in P.STORYBOARDS:
        d.h2(f"{sb['id']} {sb['title']}")
        d.para(sb["summary"])
        d.table(["Step", "Who", "Screen", "Does", "Sees"],
                [[str(k + 1), w, s, a, r] for k, (w, s, a, r, _) in enumerate(sb["steps"])],
                [0.7, 1.9, 1.8, 2.6, 2.2], caption=f"{sb['id']} steps")
        for k, (w, s, a, r, shot) in enumerate(sb["steps"]):
            d.screen(shot, f"{sb['id']} step {k + 1}: {w}, {s}")
    return d


# ================================================================================================ UXD
def write_uxd(tmp: Path) -> Doc:
    d = Doc("UXD", C.REVISION_COMPANION + [("17 Oct 2026", C.UXD_VERSION, "Screenshots of the Ledger Accounts and Change History tabs of the Insurer Record (SCR-PM-18) and of the Master Data Change Log (SCR-PM-27); no change of the requirements.", C.AUTHOR)])
    d.h1("Introduction")
    d.para("This document describes every Product Maintenance screen of BIBS: purpose, users, access, rules and "
           "states, fields with their type, whether they are required, format and source, validation and message, "
           f"and the actions. The requirements and rules are in the FRS v{C.VERSION}. Each screen state is shown by "
           "one full-page screenshot on its own page, taken on a test environment with fictitious data.")
    d.para("Codes used: Y required, N optional, Cond. required in the case stated, - display only.")
    d.h1("Menu by Persona")
    d.para("The screens are under Client & Policy › Product Maintenance; the proposal requests of SCR-PM-28 are under "
           "New Business. A persona sees only its entries.")
    names = {"A": "AO", "P": "Mkt approver", "T": "TSU", "L": "TSU TL", "H": "TSU Head", "C": "ManCom", "M": "MBS",
             "B": "Bus. Admin", "U": "Auditor"}
    d.table(["Screen", "Purpose", "Personas"], [[a, b, ", ".join(names[x] for x in c.split())] for a, b, c in C.MENU],
            [2.6, 3.8, 2.8], caption="Product Maintenance menu by persona")
    d.h1("Screens")
    for s in S.SCREENS:
        d.h2(f"{s['id']} {s['name']}")
        d.table(["Item", "Value"], [["Purpose", s["purpose"]], ["Access", s["access"]],
                                    ["Requirements", s["frpm"]]], [2.2, 7.0])
        if s["rules"]:
            d.para("Rules and states:")
            d.bullets(s["rules"])
        d.table(["Field", "Type", "Req.", "Format, list or source", "Validation and message"],
                [list(r) for r in s["fields"]], [2.4, 1.3, 0.8, 2.4, 2.3], caption=f"{s['id']} fields")
        d.table(["Action", "Who", "Available when", "What happens", "Result"], [list(r) for r in s["actions"]],
                [1.6, 1.7, 1.8, 2.7, 1.4], caption=f"{s['id']} actions")
        for shot, cap in P.SHOTS[s["id"]]:
            d.screen(shot, f"{s['id']} {cap}")
    return d


# ================================================================================================ RNO
def write_rno(tmp: Path) -> Doc:
    d = Doc("RNO", C.REVISION_COMPANION)
    d.h1("Introduction")
    d.para("This document lists the reports, documents, exports, notices and e-mails and the scheduled jobs of "
           "Product Maintenance. Times are Manila time. Documents e-mailed outside BIBS are password protected and "
           "the password follows in a separate e-mail (BR-PM-11).")
    d.h1("Reports")
    d.table(["ID", "Report", "BRD", "Parameters", "Columns", "Output"], P.REPORTS, [1.2, 1.7, 1.1, 1.8, 2.4, 1.2],
            caption="Product Maintenance reports")
    d.para("Every report run is logged with the user, the parameters and the time. Reports open in the report runner "
           "of BIBS with their parameters.")
    d.screen("ux-23-dashboard", "Product Maintenance Dashboard (RPT-PM-01)")
    d.screen("ux-26-placement", "Placement Update Reports (RPT-PM-02)")
    d.h1("Documents")
    d.table(["ID", "Document", "Format and file name", "Protection", "FRPM"], P.DOCUMENTS, [0.9, 2.0, 3.0, 1.8, 1.5],
            caption="Documents produced")
    d.h1("Exports")
    d.table(["Output", "Format", "Where", "Content"], P.EXPORTS, [2.2, 2.0, 2.4, 2.6], caption="Exports")
    d.h1("Notices and E-mails")
    d.para("Notices appear under the bell and, where shown, by e-mail. Text in angle brackets is filled in. Alerts "
           "appear on the Alerts screen.")
    d.table(["ID", "Event", "To", "Title", "Text", "Channel"], P.NOTICES, [0.8, 1.7, 1.8, 2.0, 2.0, 0.9],
            caption="Notices and e-mails")
    d.h1("Scheduled Jobs")
    d.table(["Job", "When", "What it does", "Notices"], P.JOBS, [2.6, 1.8, 3.8, 1.0], caption="Jobs")
    d.para("Every run is recorded with its start, end, outcome and counts, and can be started at once from "
           "Operations › Jobs by a user with the job permission.")
    return d


# ================================================================================================ ICD
def write_icd(tmp: Path) -> Doc:
    d = Doc("ICD", [(C.DATE, C.ICD_VERSION, "First issue: product master transfer file and e-mail channel of "
                     "Product Maintenance.", C.AUTHOR)])
    d.h1("Introduction")
    d.para("This document specifies the interfaces of Product Maintenance with systems outside BIBS. Both are "
           f"Integration scope (FRS v{C.VERSION}, FRPM.029.02 and FRPM.027.01). The links with other BIBS modules "
           "(New Business, Operations, Renewal, User Access) are internal and not listed here.")
    d.table(["ID", "Interface", "Direction", "Means", "Frequency", "Content"], P.INTERFACES,
            [1.0, 1.7, 2.2, 1.6, 1.6, 1.6], caption="Interfaces")
    d.h1("IF-PM-01 Product Master Transfer")
    d.h2("Trigger and content")
    d.bullets(["A package version is released (change RELEASED), a package is retired (RETIRED) or a package "
               "expires without a newer version (EXPIRED): BIBS records a product master change.",
               "The job PRODUCT_MASTER_TRANSFER runs every 30 minutes; Send Now on Product Master Transfers runs it at "
               "once. It writes every pending change, oldest first, in one file.",
               "No file is written when no change is pending or no folder is set; the changes wait."])
    d.h2("File")
    d.table(["Item", "Value"], [["Name", "ProductMaster_yyyyMMddHHmmss.csv (Manila time of the run)"],
                                ["Format", "Text, UTF-8, comma separated, one record per line"],
                                ["Records", "One header (H), one detail (D) per change, one trailer (T)"],
                                ["Delivery", "Written to the folder of the receiving systems under the name .part, "
                                 "then renamed, so that a reader never sees half a file"],
                                ["Folder", "Set at deployment with the details "
                                 "given by BDOI IT (D-PM-01)"]], [2.2, 7.0], caption="File")
    d.table(["Record", "No.", "Field", "Format", "Example"], P.FILE_RECORDS, [1.6, 0.7, 2.6, 2.6, 1.9],
            caption="Records of the file")
    d.h2("Errors and recovery")
    d.bullets(["A file that cannot be written marks its changes Failed with the error and raises the alert "
               "PM_MASTER_TRANSFER_FAILED on the Alerts screen.",
               "Process Again on a failed change returns it to the queue and runs the transfer; a change that did not "
               "fail is refused (\"Only a failed product master change is processed again\").",
               "Each delivered file is recorded in the audit trail with the number of changes and the target.",
               "In test environments the file goes to a simulated receiving system whose files are listed on Product "
               "Master Transfers."])
    d.screen("ux-25-transfers", "Product Master Transfers (SCR-PM-25)")
    d.h1("IF-PM-02 E-mail Channel")
    d.bullets(["Quotation slips (NT-04), advisories (NT-09) and the password e-mails (NT-05) are queued by BIBS and "
               "sent through the mail relay of BDOI; the job MAIL_DISPATCH sends the queue every 2 minutes.",
               "Each e-mail is logged on its record with recipients, subject, attachment and status; a failed e-mail "
               "can be resent.",
               "Attachments are password-protected PDF or Excel files; the password follows in a separate e-mail.",
               "BDOI IT gives the relay, the sender mailbox and its credentials (D-PM-02)."])
    d.h1("Open Points")
    d.table(["ID", "Point", "Owner"], [["Q-PM-08", "Receiving systems, folder and timing of the product master file.",
                                        "Product Owner; BDO ITG"],
                                       ["D-PM-02", "Mail relay and sender mailbox.", "BDO ITG"]], [1.4, 5.8, 2.0],
            caption="Open points")
    return d


WRITERS = {"FRS": write_frs, "BPS": write_bps, "UXD": write_uxd, "RNO": write_rno, "ICD": write_icd}


def checks() -> None:
    ids = {r[0] for f in C.FRUM for r in f["reqs"]} | {f["id"] for f in C.FRUM}
    for row in C.MAPPING:
        refs = [t.strip() for t in row[2].replace(";", ",").split(",")]
        if not any(r in ids for r in refs):
            raise SystemExit(f"mapping row {row[0]} points to no FRPM item")
    missing = [s["id"] for s in S.SCREENS if s["id"] not in P.SHOTS]
    if missing:
        raise SystemExit(f"screens without a screenshot: {missing}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--only", default="FRS,BPS,UXD,RNO,ICD")
    args = ap.parse_args()
    checks()
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        for key in args.only.split(","):
            d = WRITERS[key](tmp)
            target = OUT / C.DOCS[key]["file"]
            n = pc.finish(d, target, tmp, pdf=not args.no_pdf)
            print(f"wrote {target.name} ({n} TOC lines, {d.shots} screenshots)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
