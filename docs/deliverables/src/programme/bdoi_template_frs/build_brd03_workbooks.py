"""The workbooks of the Business Unit review pack of BRD-03 Product Maintenance (with the FRS v1.2 in BDOI's template).

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd03_workbooks.py

Writes, in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-03_Product_Maintenance/:
  BIBS_Inputs_BRD-03_Business_Unit_Requirements_Collection_v1.0.xlsx   what the Business Unit provides (fill-in)
  BIBS_FitGap_BRD-03_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx           per BRD ID: OOTB, BIBS, best practice, fit
  BIBS_CR_BRD-03_Change_Request_Register_v1.0.xlsx                    change requests after sign-off, decisions log
  BIBS_RTM_BRD-03_Test_Cases_and_Traceability_v1.0.xlsx               BRD ID -> FRPM -> test cases; answered items
  BIBS_UAT_BRD-03_Business_Unit_Walkthrough_Users_v1.0.xlsx           users, sign-in IDs and the walkthrough script
and, internal, in docs/deliverables/out/Programme/Quality/:
  BIBS_Coverage_BRD-03_FRS_Coverage_in_BIBS_v1.0.xlsx                 availability of every FRPM item and criterion

The FRS is built in memory first (build_bdoi_frs_pm.build) so that the numbers of the FRS and of the workbooks are
the same: the FRPM items and acceptance criteria, the observations of Annex M, the open items of Annex V and the test
cases. Self-checks: every BRD ID is in the traceability sheet with at least one test case and in the fit-gap sheet;
every FRPM item has at least one test case and a coverage row; every workbook item that Annex V names exists in the
requirements-collection workbook; no password value is written.
"""

from __future__ import annotations

import datetime as dt
import re
import sys
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402
import build_bdoi_frs_pm as core  # noqa: E402
from bdoi_xlsx import BdoiWorkbook, Column, _SheetInfo  # noqa: E402
from guided_xlsx import GuidedBook, font as gfont, set_link, BORDER as GBORDER, HEADER_FILL as GHEAD  # noqa: E402

v12 = core.v12
OUT_DIR = REPO / "docs" / "deliverables" / "out" / "Programme" / "BDOI_Template_FRS" / "BRD-03_Product_Maintenance"
QUALITY = REPO / "docs" / "deliverables" / "out" / "Programme" / "Quality"
PROG = REPO / "docs" / "deliverables" / "out" / "Programme"
INPUTS = PROG / "BDOI_Inputs" / "BIBS_Inputs_BRD-00_BDOI_Requirements_and_Inputs_by_BRD_v1.2.xlsx"
FEATURES = PROG / "Feature_List" / "BIBS_BDOI_FeatureList_vs_OOTB_and_Best_Practice_v1.0.xlsx"
CONFORMANCE = PROG / "Quality" / "BIBS_Conformance_BRD-00_End-to-End_Conformance_Register_v1.0.xlsx"
CMP = REPO / "docs" / "deliverables" / "src" / "programme" / "comparisons" / "brd03_frs_comparison.yaml"
DATE = "9 October 2026"
PW_TEXT = "Issued separately by the iorta TechNXT Project Team"
ENVIRONMENT = "BIBS UAT environment (address issued with the passwords)"
ITEMS = v12.ITEMS
V12 = v12.V12


def fmt_date(v) -> str:
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%d-%b-%Y")
    return str(v or "")


# ================================================================================================ workbook helper
class Book(BdoiWorkbook):
    """BdoiWorkbook with the drop-down lists on a visible sheet 'Lists' (one column per list, its name on top) and
    sheets of free layout (forms, process tables) that keep their place in the cover index."""

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self._lists.title = "Lists"
        self._lists.sheet_state = "visible"
        self._queue: list[str] = []

    def sheet(self, name, columns, rows, description="", freeze_first_column=True):
        self._queue = [f"{name[:31]}: {c.header}" for c in columns if c.values]
        return super().sheet(name, columns, rows, description, freeze_first_column)

    def _list_ref(self, values):
        self._list_col += 1
        letter = get_column_letter(self._list_col)
        head = self._lists.cell(row=1, column=self._list_col, value=self._queue.pop(0) if self._queue else "List")
        head.font = Font(name=brand.FONT, bold=True, color=brand.WHITE)
        head.fill = PatternFill("solid", fgColor=brand.HEADER_BLUE)
        head.alignment = Alignment(wrap_text=True, vertical="top")
        self._lists.column_dimensions[letter].width = 26
        for i, v in enumerate(values, start=2):
            self._lists.cell(row=i, column=self._list_col, value=v)
        return f"'Lists'!${letter}$2:${letter}${len(values) + 1}"

    def free(self, name: str, description: str, rows: int = 0):
        ws = self.wb.create_sheet(name[:31])
        ws.sheet_view.showGridLines = False
        self._sheets.append(_SheetInfo(ws.title, description, [], rows))
        self._print_setup(ws, None)
        return ws

    def save(self, path):
        self._lists.freeze_panes = "A2"
        self._lists.row_dimensions[1].height = 45
        self._print_setup(self._lists, "1:1")
        return super().save(path)


THIN = Side(style="thin", color=brand.BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_table(ws, row: int, headers, rows, widths=None, title: str | None = None) -> int:
    if title:
        ws.cell(row=row, column=1, value=title).font = Font(name=brand.FONT, size=12, bold=True,
                                                           color=brand.HEADER_BLUE)
        row += 1
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name=brand.FONT, size=10, bold=True, color=brand.WHITE)
        cell.fill = PatternFill("solid", fgColor=brand.HEADER_BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
        if widths:
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in rows:
        row += 1
        for c, v in enumerate(r, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = Font(name=brand.FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BOX
    return row + 2


def title_block(ws, title: str, lines) -> int:
    ws["A1"] = title
    ws["A1"].font = Font(name=brand.FONT, size=14, bold=True, color=brand.HEADER_BLUE)
    row = 2
    for line in lines:
        ws.cell(row=row, column=1, value=line).font = Font(name=brand.FONT, size=10, color=brand.TEXT)
        row += 1
    return row + 1


# ================================================================================================== source data
def load_inputs() -> list[dict]:
    wb = openpyxl.load_workbook(INPUTS, read_only=True)
    rows = list(wb["BRD-03 Product Maintenance"].iter_rows(values_only=True))
    keys = ["ref", "brd", "type", "need", "why", "proposal", "source", "owner", "needed", "priority", "status",
            "response", "remarks"]
    return [dict(zip(keys, r)) for r in rows[4:] if r[0]]


def load_features() -> list[dict]:
    wb = openpyxl.load_workbook(FEATURES, read_only=True)
    rows = list(wb["Feature Comparison"].iter_rows(values_only=True))
    head = rows[3]
    return [dict(zip(head, r)) for r in rows[4:] if r[0]]


def load_conformance() -> dict[str, list[tuple]]:
    wb = openpyxl.load_workbook(CONFORMANCE, read_only=True)
    out: dict[str, list[tuple]] = {}
    for r in list(wb["FR conformance"].iter_rows(values_only=True))[4:]:
        if r[1]:
            out.setdefault(r[1], []).append((r[4], r[5], r[7], r[8]))
    return out


BRD_RE = re.compile(r"\b(?:BRPM\.\d{3}|PMADD\d\d)\b")


def brd_ids_in(*texts) -> list[str]:
    found = []
    for t in texts:
        for m in BRD_RE.findall(str(t or "")):
            if m not in found:
                found.append(m)
    return found


# ================================================================================ 1. requirements collection
KINDS = OrderedDict([
    ("DOC", ("Document templates and prints", "Layouts, wording and logo placement of the documents BIBS prints, "
             "e-mails or exports (Annex Q)", "Word, PDF or Excel sample with the wording; or 'Approve as in Annex Q'")),
    ("EML", ("E-mail wordings", "The wording of each e-mail and notice (Annex P): approve it or give the changed "
             "text", "'Approve as written' or the changed subject and text")),
    ("SCH", ("Schedules and timings", "Times and days of the scheduled runs and alerts (Annex R)",
             "Day and time (Philippine time); recipients")),
    ("PAR", ("Parameters and thresholds", "Service levels, alert days, limits and other values the system uses",
             "A value with its unit (hours, days, PHP, %)")),
    ("LOV", ("Lists of values", "The values of the drop-down lists of Product Maintenance",
             "The list of values, one per line; mark values to remove")),
    ("MDA", ("Master and reference data", "Master data to supply in the programme upload templates v1.2 (the tab "
             "is named; the template is not repeated here)", "Filled tab of the upload templates workbook v1.2")),
    ("APR", ("Approval matrix and users", "Who approves at each step, by name and role, and the users of each role",
             "Names, sign-in or e-mail, role, unit, limits")),
    ("INT", ("Integration inputs from BDOI IT", "What BDOI IT provides for each interface (Annex S)",
             "System, data, format, timing, contact person")),
    ("DEC", ("Open decisions", "Decisions and confirmations not answered in BDOI's FRS (Annex V and Annex M)",
             "The decision (option chosen) and any condition")),
])
TYPE_KIND = {"Sample document or template": "DOC", "Report layout": "DOC", "Configuration value": "PAR",
             "Master data": "MDA", "Integration detail": "INT", "Sign-off": "APR", "Test data": "APR",
             "Business rule confirmation": "DEC", "Decision": "DEC", "Clarification": "DEC"}
KIND_OVERRIDE = {"PM-Q54": "LOV", "PM-Q38": "SCH", "CLR-PM-22": "SCH", "PM-Q10": "APR", "CLR-PM-06": "APR",
                 "PM-Q09": "APR"}
STATUSES = ["Open", "Received", "Confirmed"]


def frs_ref_for(ref: str, brd: list[str]) -> str:
    kept = {i["rc"]: i for i in v12.open_items()["kept"]}
    if ref in kept or any(i["id"] == ref for i in v12.open_items()["kept"]):
        return "Annex V (" + ", ".join(i["id"] for i in v12.open_items()["kept"] if i["rc"] == ref or i["id"] == ref) \
            + ")"
    for u in V12["data_setup"]["items"]:
        if ref in u[7]:
            return f"Annex U ({u[0]})"
    for d in V12["documents"]["items"]:
        if ref in d["layout"]:
            return f"Annex Q ({d['id']})"
    m = v12.brd_mapping()
    frpm = [f for b in brd for f in m.get(b, {}).get("frpm", [])]
    return ", ".join(dict.fromkeys(frpm[:3])) or "Section 3"


def collection_items() -> "OrderedDict[str, list[dict]]":
    out: OrderedDict[str, list[dict]] = OrderedDict((k, []) for k in KINDS)
    answered_in = {k for k, v in ITEMS["inputs_bdoi"].items() if v[0] == "Answered"}
    for r in load_inputs():
        if r["ref"] in answered_in or r["status"] == "Answered":
            continue
        kind = KIND_OVERRIDE.get(r["ref"], TYPE_KIND.get(r["type"], "DEC"))
        brd = brd_ids_in(r["need"], r["source"]) or ["-"]
        need = r["need"]
        part = ITEMS["inputs_bdoi"].get(r["ref"])
        if part:
            need = f"{need} BDOI's FRS answers in part ({part[1]}: {part[2]}); still needed: {part[3]}"
        out[kind].append({"id": r["ref"], "item": r["type"], "brd": ", ".join(brd), "frs": frs_ref_for(r["ref"], brd),
                          "need": need, "format": KINDS[kind][2], "example": r["proposal"] or "-",
                          "by": r["owner"], "when": fmt_date(r["needed"]), "status": "Open",
                          "src": "Programme inputs workbook v1.2"})
    # e-mail wordings
    for x in V12["notifications"]["items"]:
        brd = sorted({b for f in re.findall(r"FRPM\.\d{3}\.\d{2}", x["frpm"]) for b in v12.brd_of_item(f)})
        out["EML"].append({"id": f"RC-EML-{x['id'][3:]}", "item": f"{x['id']} {x['name']}", "brd": ", ".join(brd),
                           "frs": f"Annex P ({x['id']}); {x['frpm']}",
                           "need": f"Confirm the {x['channel'].lower()} '{x['name']}': trigger, recipients and text "
                                   f"({x['source'].split(' (')[0]}).",
                           "format": KINDS["EML"][2], "example": f"Subject: {x['subject']}",
                           "by": "Head, TSU" if "ManCom" not in x["name"] else "ManCom secretariat",
                           "when": "30-Nov-2026", "status": "Open", "src": "FRS v1.2 Annex P"})
    for i, s in enumerate(V12["reports"]["schedules"], 1):
        brd = sorted({b for f in re.findall(r"FRPM\.\d{3}\.\d{2}", s[4]) for b in v12.brd_of_item(f)})
        out["SCH"].append({"id": f"RC-SCH-{i:02d}", "item": s[0], "brd": ", ".join(brd), "frs": f"Annex R.2; {s[4]}",
                           "need": f"Confirm the time ({s[1]}) and the recipients ({s[3]}) of the run: {s[2]}",
                           "format": KINDS["SCH"][2], "example": s[1], "by": "Head, TSU", "when": "30-Nov-2026",
                           "status": "Open", "src": "FRS v1.2 Annex R"})
    for i, st in enumerate(core.DOC["annex_j"]["stages"], 1):
        out["PAR"].append({"id": f"RC-PAR-{i:02d}", "item": f"Service level: {st[0]}", "brd": "BRPM.021",
                           "frs": "FRPM.028.01; Annex O", "need": f"Service level of the stage {st[0]} (owner "
                           f"{st[1]}); part of PM-Q17.", "format": "Hours", "example": f"{st[2]} hours (default)",
                           "by": "Head, TSU", "when": "16-Jul-2027", "status": "Open", "src": "FRS Annex J.4"})
    base = len(core.DOC["annex_j"]["stages"])
    for i, p in enumerate(core.DOC["annex_j"]["parameters"], base + 1):
        out["PAR"].append({"id": f"RC-PAR-{i:02d}", "item": p[0], "brd": "BRPM.021, BRPM.017", "frs": "Annex J.2",
                           "need": f"{p[2]}. Confirm or change the default.", "format": "A value with its unit",
                           "example": f"{p[1]} (default)", "by": p[3] if p[3] != "System Administrator" else
                           "Head, TSU", "when": "16-Jul-2027", "status": "Open", "src": "FRS Annex J.2"})
    for i, l in enumerate(core.DOC["annex_j"]["lists"], 1):
        out["LOV"].append({"id": f"RC-LOV-{i:02d}", "item": l[0], "brd": "BRPM.003, BRPM.004", "frs": "Annex J.3",
                           "need": f"Confirm or change the values of the list '{l[0]}'.", "format": KINDS["LOV"][2],
                           "example": l[1][:120], "by": "Head, MBS", "when": "16-Jul-2027", "status": "Open",
                           "src": "FRS Annex J.3 (template PM-09 of the upload templates v1.2)"})
    docs = [("RC-DOC-01", "Package Request Form layout", "DO-01", "BRPM.005, BRPM.008"),
            ("RC-DOC-02", "Comparative Table: client output layout", "DO-03", "PMADD03, BRPM.014"),
            ("RC-DOC-03", "Package slip layout", "DO-04", "BRPM.016"),
            ("RC-DOC-04", "ManCom approval record: confirm the standard layout", "DO-05", "BRPM.015"),
            ("RC-DOC-05", "Package advisory layout and wording", "DO-06", "BRPM.016"),
            ("RC-DOC-06", "Proposal Slip (non-package): confirm the layout", "DO-08", "BRPM.010"),
            ("RC-RPT-01", "Package reports: columns of the status, expiry and version history reports", "DO-10",
             "BRPM.017, BRPM.018")]
    for rid, item, do, brd in docs:
        out["DOC"].append({"id": rid, "item": item, "brd": brd, "frs": f"Annex Q ({do})",
                           "need": f"{item}: the BDOI layout with wording and logo placement, or approval of the "
                                   "current BIBS layout shown in Annex Q.", "format": KINDS["DOC"][2],
                           "example": "Approve as in Annex Q", "by": "Head, TSU", "when": "16-Jul-2027",
                           "status": "Open", "src": "FRS v1.2 Annex Q"})
    apr = [("RC-APR-01", "Marketing approvers per unit (TL, TH, UH)", "BRPM.008, BRPM.021", "FRPM.011.02"),
           ("RC-APR-02", "ManCom approver list, including the President", "BRPM.015", "FRPM.014.01"),
           ("RC-APR-03", "TSU Team Leads, TSU Head and the Quotation Slip co-officer", "BRPM.009, BRPM.012",
            "FRPM.011.03, FRPM.012.01"),
           ("RC-APR-04", "Validators of package configurations (post-setup validation)", "PMADD06", "FRPM.016.02"),
           ("RC-APR-05", "Approvers of rate exceptions", "BRPM.007", "FRPM.024.02"),
           ("RC-APR-06", "Approvers of deactivation requests and product matrix records", "BRPM.011, BRPM.021",
            "FRPM.003.03, FRPM.003.07"),
           ("RC-APR-07", "Business Unit reviewers of the walkthrough (one per role)", "All", "Annex X; walkthrough "
            "users workbook")]
    for rid, item, brd, frs in apr:
        out["APR"].append({"id": rid, "item": item, "brd": brd, "frs": f"{frs}; Annex O",
                           "need": f"{item}: names, role and unit of each person.", "format": KINDS["APR"][2],
                           "example": "Name, sign-in or e-mail, role, unit", "by": "Product Owner",
                           "when": "30-Nov-2026", "status": "Open", "src": "FRS v1.2 Annex O"})
    ints = [("RC-INT-01", "Microsoft 365 sending of the ManCom e-mails: mailbox and permission", "BRPM.015",
             "FRPM.014.01; Annex S (IN-08)"),
            ("RC-INT-02", "Document store (ECM): location and retention of the package documents", "BRPM.016, "
             "BRPM.024", "Annex S (IN-10)")]
    for rid, item, brd, frs in ints:
        out["INT"].append({"id": rid, "item": item, "brd": brd, "frs": frs, "need": item + ".",
                           "format": KINDS["INT"][2], "example": "-", "by": "BDO ITG", "when": "04-Jan-2027",
                           "status": "Open", "src": "FRS v1.2 Annex S"})
    present = {r["id"] for rows in out.values() for r in rows}
    for i in v12.open_items()["kept"]:
        if i["rc"] in present:
            continue
        out["DEC"].append({"id": i["id"], "item": i["kind"], "brd": "-", "frs": f"Annex V ({i['id']})",
                           "need": i.get("remaining") or i["text"], "format": i["format"], "example": "-",
                           "by": i["owner"], "when": i["needed"], "status": "Open", "src": "FRS v1.2 Annex V"})
        present.add(i["id"])
    return out


def build_collection(items) -> Path:
    path = OUT_DIR / "BIBS_Inputs_BRD-03_Business_Unit_Requirements_Collection_v1.0.xlsx"
    g = GuidedBook("Business Unit Requirements Collection – BRD-03 Product Maintenance", "1.0")
    status = g.list_for_values("Status", STATUSES)
    cols = [("ID *", 13, "Y", "Text", "Fixed", "The item number; keep it when you answer"),
            ("Item", 26, "N", "Text", "Fixed", "What the item is about"),
            ("BRD ID", 14, "N", "BRD IDs", "Fixed", "The BRD requirement(s) the item serves"),
            ("FRS reference", 20, "N", "FRPM item or annex", "Fixed", "Where the item is in the FRS v1.2"),
            ("What is needed", 46, "N", "Text", "Fixed", "What the Business Unit provides or decides"),
            ("Format of the answer", 24, "N", "Text", "Fixed", "How to give the answer"),
            ("Example", 26, "N", "Text", "Fixed", "An example or the proposal of the project team"),
            ("Provided by", 18, "N", "Role", "Fixed", "The role that answers"),
            ("Needed by", 13, "N", "dd-MMM-yyyy", "Fixed", "Date of the drop plan"),
            ("BU response *", 40, "Y", "Text", "Free entry", "The answer of the Business Unit"),
            ("Status *", 12, "Y", "List", ", ".join(STATUSES), "Open until answered; Received when answered; "
             "Confirmed after review")]
    index_rows = []
    order = ["MDA", "LOV", "PAR", "SCH", "APR", "DOC", "EML", "INT", "DEC"]  # codes others use first
    for step, kind in enumerate(order, 1):
        name, purpose, _ = KINDS[kind]
        rows = items[kind]
        ws = g.wb.create_sheet(name[:31])
        g._title(ws, f"{kind} {name}", purpose)
        block = [("Purpose", purpose), ("Who fills it in", "The Business Unit owner named in 'Provided by'"),
                 ("How it is used", "The answers are applied as configuration or recorded as decisions in the next "
                                    "version of the FRS (Annex V); master data goes into the upload templates v1.2"),
                 ("Due", "The date in 'Needed by' of each row"),
                 ("Depends on", "Decisions of the sheet 'Open decisions' where a row says so")]
        r = 3
        for k, v in block:
            ws.cell(row=r, column=1, value=k).font = gfont(9, True, brand.HEADER_BLUE)
            ws.cell(row=r, column=2, value=v).font = gfont(9)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            r += 1
        guide_top = r + 1
        labels = ["Mandatory", "Format", "Allowed values", "What to enter"]
        for gi, lab in enumerate(labels):
            cell = ws.cell(row=guide_top + gi, column=1, value=lab)
            cell.font = gfont(8, True, brand.HEADER_BLUE)
        header = guide_top + len(labels)
        for c, (h, w, mand, fmt, allowed, what) in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
            for gi, val in enumerate([mand, fmt, allowed, what]):
                cell = ws.cell(row=guide_top + gi, column=c, value=val if c > 1 or gi > 0 else val)
                if c == 1:
                    cell.value = f"{labels[gi]}: {val}"
                cell.font = gfont(8, False, brand.MUTED, italic=True)
                cell.fill = PatternFill("solid", fgColor=brand.DIRTY_WHITE)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = GBORDER
            cell = ws.cell(row=header, column=c, value=h)
            cell.font = gfont(9, True, brand.WHITE)
            cell.fill = GHEAD
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = GBORDER
        ex = header + 1
        example = ["Example – overwrite or delete", "Example item", "BRPM.021", "FRPM.028.01",
                   "The service level of the Marketing approval", "Hours", "24", "Head, Retail Marketing",
                   "30-Nov-2026", "24 hours", "Received"]
        for c, v in enumerate(example, start=1):
            cell = ws.cell(row=ex, column=c, value=v)
            cell.font = gfont(9, False, brand.MUTED, italic=True)
            cell.fill = PatternFill("solid", fgColor=brand.DIRTY_WHITE)
            cell.border = GBORDER
        first = ex + 1
        for i, it in enumerate(rows):
            vals = [it["id"], it["item"], it["brd"], it["frs"], it["need"], it["format"], it["example"], it["by"],
                    it["when"], None, it["status"]]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=first + i, column=c, value=v)
                cell.font = gfont(9, c == 1)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = GBORDER
        last = first + max(len(rows), 1) - 1 + 50
        g._list_validation(ws, f"K{first}:K{last}", status, "Status", "Open, Received or Confirmed")
        g._status_colours(ws, f"K{first}:K{last}", STATUSES)
        ws.freeze_panes = ws.cell(row=first, column=2)
        ws.auto_filter.ref = f"A{header}:K{first + max(len(rows), 1) - 1}"
        g._print_setup(ws, f"{header}:{header}")
        n = len(rows)
        index_rows.append({"Step": step, "Sheet": (name, f"'{ws.title[:31]}'!A1"), "Owner": ", ".join(
            sorted({str(x["by"]) for x in rows}))[:120], "Due": min((x["when"] for x in rows), default="-",
                                                               key=lambda d: dt.datetime.strptime(d, "%d-%b-%Y")
                                                               if re.match(r"\d\d-\w{3}-\d{4}$", str(d)) else
                                                               dt.datetime(2099, 1, 1)),
                           "Route": "Answer in the sheet; the project team applies it",
                           "Items": n,
                           "Answered": f"=COUNTIF('{ws.title}'!J{first}:J{last},\"?*\")",
                           "Status": "Not started", "Open →": ("Open →", f"'{ws.title}'!A1")})
    answered = [k for k, v in ITEMS["inputs_bdoi"].items() if v[0] == "Answered"]
    g.start_sheet(
        "Business Unit Requirements Collection", "BRD-03 Product Maintenance – what the Business Unit provides "
        "for the platform to work as specified in the FRS v1.2 (BDOI template)",
        [("Workbook", "BIBS_Inputs_BRD-03_Business_Unit_Requirements_Collection_v1.0"), ("Version", "1.0"),
         ("Date", DATE), ("FRS", "BIBS FRS-BDOI BRD-03 Product Maintenance v1.2"),
         ("Consistent with", "Programme inputs workbook BIBS_Inputs_BRD-00 v1.2 (sheet BRD-03): same IDs; "
          f"{len(answered)} of its items are answered by BDOI's FRS and are listed in the traceability workbook "
          "(sheet 'Answered in the BDOI FRS'); the upload templates v1.2 hold the master data templates"),
         ("Prepared by", brand.VENDOR)],
        ["Open the sheets in the order of the index: master data and lists first, decisions last.",
         "For each row, write the answer in 'BU response' in the format asked, and set Status to Received.",
         "Ask questions in the sheet 'Questions and comments'; the project team answers there.",
         "The project team reviews each answer and sets Status to Confirmed."],
        "Sheets in fill-in order", [("Step", 7), ("Sheet", 30), ("Owner", 34), ("Due", 13), ("Route", 30),
                                    ("Items", 8), ("Answered", 10), ("Status", 14), ("Open →", 10)],
        index_rows)
    _questions(g, [KINDS[k][0] for k in order])
    g.lists_sheet()
    order_names = ["Start here"] + [KINDS[k][0][:31] for k in order] + ["Questions and comments", "Reference lists"]
    return g.save(path, order_names, {"title": "Business Unit Requirements Collection BRD-03",
                                      "keywords": "Inputs, BRD-03, BIBS, BDOI"})


def _questions(g: GuidedBook, sheets):
    ws = g.wb.create_sheet("Questions and comments")
    g._title(ws, "Questions and comments", "Raise a question or a comment on a sheet or an item; the iorta TechNXT "
                                           "Project Team answers here.")
    cols = [("No.", 6), ("Sheet", 28), ("Item ID", 14), ("Question / comment", 50), ("Raised by", 20), ("Date", 13),
            ("Answer", 50), ("Answered by", 20), ("Status", 12)]
    for c, (h, w) in enumerate(cols, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = gfont(9, True, brand.WHITE)
        cell.fill = GHEAD
        cell.border = GBORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    for r in range(5, 65):
        ws.cell(row=r, column=1, value=r - 4)
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = GBORDER
    names = g.list_for_values("Sheet", list(sheets) + ["General"])
    g._list_validation(ws, "B5:B64", names, "Sheet", "The sheet the question is about")
    st = g.list_for_values("Question status", ["Open", "Answered", "Closed"])
    g._list_validation(ws, "I5:I64", st, "Status", "Open, Answered or Closed")
    ws.freeze_panes = "B5"
    g._print_setup(ws, "4:4")


# ======================================================================================================= 2. fit-gap
def build_fitgap() -> tuple[Path, list[dict]]:
    feats = load_features()
    m = v12.brd_mapping()
    rows = []
    for rid, mp in m.items():
        fg = ITEMS["fitgap"][rid]
        fl = [f for f in feats if rid in str(f.get("BDOI references") or "")]
        if fl:
            ootb = "\n".join(f"{f['ID']} {f['Feature']}: {f['OOTB platform']} – {f['OOTB note']}" for f in fl)
            best = "\n".join(f"{f['ID']}: {f['Best-practice expectation']}" for f in fl)
            refs = ", ".join(f["ID"] for f in fl)
        else:
            ex = ITEMS["fitgap_extra"][rid]
            ootb, best, refs = ex["ootb"], ex["best"], "-"
        ootb = re.sub(r"\b(?:[Oo]pen )?APIs?\b", "an open interface", ootb)
        rows.append({"id": rid, "text": mp["text"], "frpm": "Refer to User Access Maintenance FRS" if mp["uam"] and
                     not mp["frpm"] else ", ".join(mp["frpm"]), "ootb": ootb, "bibs": fg["bibs"], "best": best,
                     "fit": fg["fit"], "rec": fg["rec"], "impact": fg["impact"], "refs": refs})
    fits = ["Fit", "Configure", "Extend", "Recommend change"]
    wb = Book("BRD vs OOTB vs Best Practice", doc_type="Fit-gap workbook", brd="BRD-03", version="1.0", date=DATE,
              subtitle="BRD-03 Product Maintenance – per BRD requirement")
    wb.legend = [("Fit", "BIBS meets the requirement as specified"), ("Configure", "Met with values or layouts "
                 "BDOI provides"), ("Extend", "A platform extension is planned"), ("Recommend change", "Best "
                 "practice suggests a change to the BRD text")]
    cnt = Counter(r["fit"] for r in rows)
    summary = [{"fit": f, "n": cnt.get(f, 0), "ids": ", ".join(r["id"] for r in rows if r["fit"] == f)} for f in fits]
    summary.append({"fit": "Total", "n": len(rows), "ids": f"{len(rows)} BRD IDs (PMADD01 to PMADD08, BRPM.001 to "
                    "BRPM.024)"})
    wb.sheet("Summary", [Column("fit", "Fit", 20, "Fit class"), Column("n", "BRD IDs", 10, "Number of BRD IDs"),
                         Column("ids", "BRD IDs", 90, "The BRD IDs of the class")], summary,
             description="Counts of the BRD requirements by fit class")
    wb.sheet("Fit-Gap", [
        Column("id", "BRD ID", 11, "Requirement ID of the BRD"),
        Column("text", "BRD requirement", 40, "Text of the BRD (as in the Business Requirements Mapping)"),
        Column("frpm", "FRPM items", 18, "Functional requirements of the FRS v1.2"),
        Column("ootb", "What the platform offers out of the box", 45, "Standard capability of the base platform "
               "(programme Feature List)"),
        Column("bibs", "What BIBS offers for BDOI", 42, "Capability and configuration of BIBS for this BRD"),
        Column("best", "Insurance-broking best practice", 42, "Philippine market and international broking practice"),
        Column("fit", "Fit", 15, "Fit class", values=fits),
        Column("rec", "Recommendation", 36, "What the project team recommends"),
        Column("impact", "Impact", 30, "Effect on the business"),
        Column("refs", "Feature List rows", 16, "Rows of the programme Feature List workbook")], rows,
        description="One row per BRD requirement: BRD text, out of the box, BIBS, best practice, fit")
    path = wb.save(OUT_DIR / "BIBS_FitGap_BRD-03_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx")
    return path, rows


# ======================================================================================================== 3. CR register
def build_cr(b) -> tuple[Path, int, int]:
    tcs = v12.test_cases()
    rows = []
    seeds = [r for r in b.obs_rows if r[9] in ("M.1", "M.4", "M.5")]
    for i, r in enumerate(seeds, 1):
        no, ref, topic, bdoi, brd, prop = r[:6]
        ids = brd_ids_in(ref, brd, prop) or brd_ids_in(bdoi)
        frpm = re.findall(r"FRPM\.\d{3}(?:\.\d{2})?", f"{ref} {prop}")
        linked = [t["id"] for t in tcs if any(f in t["frpm"] for f in frpm)][:5]
        rows.append({"no": f"CR-PM-{i:03d}", "date": "09-Oct-2026", "by": "iorta TechNXT Project Team",
                     "brd": ", ".join(ids) or "-", "frs": f"Annex M {no}" + (f"; {', '.join(dict.fromkeys(frpm))}"
                                                                           if frpm else ""),
                     "desc": topic, "reason": prop, "type": "Proposal beyond the BRD text",
                     "scope": "To be assessed", "schedule": "To be assessed", "cost": "To be assessed",
                     "other": "To be assessed", "tests": "To be assessed", "priority": "Medium",
                     "status": "Candidate - for decision", "decision": None, "appr": None, "ddate": None,
                     "drop": "Drop 0", "linked": ", ".join(linked) or "-"})
    statuses = ["Candidate - for decision", "Raised", "Under assessment", "Approved", "Rejected", "Deferred",
                "Scheduled", "In test", "Closed"]
    wb = Book("Change Request Register", doc_type="Register", brd="BRD-03", version="1.0", date=DATE,
              subtitle="BRD-03 Product Maintenance – changes after the sign-off of the FRS v1.2")
    ws = wb.free("How to use", "The change process after sign-off: raise, assess, approve, schedule, test, close")
    r = title_block(ws, "How to use this register", [
        "Signing the FRS v1.2 freezes its requirements (FRS Annex W). A change after sign-off follows the steps "
        "below.", "The register is seeded with the candidate changes already known: the points where the FRS "
        "proposes more than the BRD text (Annex M). Each starts as 'Candidate - for decision'."])
    write_table(ws, r, ["Step", "What happens", "Who", "When"], V12["change_control"]["steps"], [14, 70, 40, 22])
    wb.sheet("Register", [
        Column("no", "CR No.", 12, "Change request number"), Column("date", "Date raised", 12, "dd-MMM-yyyy"),
        Column("by", "Raised by", 20, "Person or team"), Column("brd", "BRD ID", 14, "BRD requirement(s)"),
        Column("frs", "FRS reference", 22, "FRPM item or annex row"), Column("desc", "Description", 34, "The change"),
        Column("reason", "Reason", 44, "Why the change is needed"),
        Column("type", "Type", 16, "Type of change", values=["Proposal beyond the BRD text", "Scope change",
                                                             "Rule change", "Screen or field", "Document or e-mail",
                                                             "Report", "Integration", "Wording"]),
        Column("scope", "Impact on scope", 16, "Assessment"), Column("schedule", "Impact on schedule", 16,
                                                                     "Assessment"),
        Column("cost", "Impact on cost (mandays)", 14, "Assessment"),
        Column("other", "Impact on other BRDs", 16, "Assessment"), Column("tests", "Impact on tests", 16,
                                                                          "Assessment"),
        Column("priority", "Priority", 11, "Business priority", values=["High", "Medium", "Low"]),
        Column("status", "Status", 20, "Status of the change request", values=statuses, status=True),
        Column("decision", "Decision", 28, "Decision and conditions"), Column("appr", "Approved by", 20, "Names"),
        Column("ddate", "Decision date", 12, "dd-MMM-yyyy", kind="date"),
        Column("drop", "Target drop", 11, "Drop", values=["Drop 0", "Drop 1", "Drop 2", "Drop 3", "Later"]),
        Column("linked", "Linked test cases", 28, "Test cases of the traceability workbook")], rows,
        description="The change requests with their impact assessment and decision")
    ws = wb.free("Impact assessment form", "One form per change request: impact on scope, schedule, cost, other "
                                           "BRDs and tests")
    r = title_block(ws, "Impact assessment form", ["Copy this sheet for each change request, or fill in the "
                                                   "impact columns of the Register."])
    form = [["CR No.", ""], ["Change requested", ""], ["BRD ID and FRS reference", ""],
            ["Impact on scope (what is added, changed, removed)", ""], ["Impact on schedule (drop, milestones)", ""],
            ["Impact on cost (mandays of the project team)", ""], ["Impact on other BRDs (Annex S integrations)", ""],
            ["Impact on tests (cases to add or change)", ""], ["Impact on documents (FRS, e-mails, layouts)", ""],
            ["Risks of doing / of not doing the change", ""], ["Recommendation of the project team", ""],
            ["Assessed by and date", ""], ["Decision (approve, reject, defer) and conditions", ""],
            ["Approved by (owners of every BRD touched, Product Owner, Program Manager BPS) and date", ""]]
    write_table(ws, r, ["Field", "Entry"], form, [50, 90])
    obs = [{"no": o[0], "sec": o[9], "ref": o[1], "topic": o[2], "prop": o[5], "who": o[7], "decision": None,
            "by": None, "date": None, "status": "Open", "cr": None} for o in b.obs_rows]
    wb.sheet("Decisions log", [
        Column("no", "Observation", 12, "Number in Annex M of the FRS v1.2"), Column("sec", "Section", 8, "Annex M"),
        Column("ref", "Ref.", 14, "Conflict, slip or FRPM reference"), Column("topic", "Topic", 30, "Topic"),
        Column("prop", "Proposed in the FRS", 50, "What the FRS proposes"), Column("who", "Decision by", 24, "Roles"),
        Column("decision", "Decision of the Business Unit", 36, "Accept the proposal, keep BDOI's text, other"),
        Column("by", "Decided by", 18, "Name"), Column("date", "Date", 12, "dd-MMM-yyyy", kind="date"),
        Column("status", "Status", 12, "Status", values=["Open", "Decided", "Closed"], status=True),
        Column("cr", "CR No.", 12, "Change request raised, if any")], obs,
        description="One row per observation of Annex M, for the Business Unit to record its decision")
    path = wb.save(OUT_DIR / "BIBS_CR_BRD-03_Change_Request_Register_v1.0.xlsx")
    return path, len(rows), len(obs)


# ============================================================================================= 4. RTM and test cases
def build_rtm() -> tuple[Path, dict]:
    tcs = v12.test_cases()
    items = v12.frpm_items()
    m = v12.brd_mapping()
    trace = []
    by_item = {i["id"]: [t for t in tcs if i["id"] in t["frpm_list"]] for i in items}
    for rid, mp in m.items():
        targets = [i for i in items if i["id"] in mp["frpm"] or i["item"] in mp["frpm"]]
        if not targets:  # BDOI's reference to the User Access Maintenance FRS
            cases = [t for t in tcs if rid in t["brd"]]
            trace.append({"brd": rid, "text": mp["text"], "frpm": "Refer to User Access Maintenance FRS", "title": "-",
                          "ref": ", ".join(sorted({t["ref"] for t in cases})), "n": len(cases),
                          "tc": ", ".join(t["id"] for t in cases[:12]) + (" ..." if len(cases) > 12 else ""),
                          "screen": "; ".join(sorted({t["screen"] for t in cases}))[:200],
                          "wt": ", ".join(sorted({w for t in cases for w in t["wt"].split(", ") if w != "-"}))})
            continue
        for it in targets:
            cases = [t for t in by_item[it["id"]] if rid in t["brd"] or not t["brd"]] or by_item[it["id"]]
            trace.append({"brd": rid, "text": mp["text"], "frpm": it["id"], "title": it["title"],
                          "ref": ", ".join(it["refs"]), "n": len(cases),
                          "tc": ", ".join(t["id"] for t in cases[:12]) + (" ..." if len(cases) > 12 else ""),
                          "screen": "; ".join(sorted({t["screen"] for t in cases}))[:200],
                          "wt": ", ".join(sorted({w for t in cases for w in t["wt"].split(", ") if w not in
                                                  ("-", "...")}))[:200] or "-"})
    problems = [f"BRD ID {r} without a test case" for r in m if not any(t["brd"] == r and t["n"] for t in trace)]
    problems += [f"FRPM item {i['id']} without a test case" for i in items if not by_item[i["id"]]]
    if problems:
        raise SystemExit("RTM checks failed:\n  " + "\n  ".join(problems))
    oi = v12.open_items()
    answered = [{"ref": i["id"], "kind": i["kind"], "text": i["text"], "status": i["status"], "clause": i["clause"],
                 "answer": i["answer"], "remaining": "-", "src": "BIBS FRS BRD-3 v2.1"} for i in oi["answered"]]
    answered += [{"ref": i["id"], "kind": i["kind"], "text": i["text"], "status": "Partly answered",
                  "clause": i["clause"], "answer": i["answer"], "remaining": f"{i['remaining']} (FRS Annex V)",
                  "src": "BIBS FRS BRD-3 v2.1"} for i in oi["partly"]]
    inputs = {r["ref"]: r for r in load_inputs()}
    for k, v in ITEMS["inputs_bdoi"].items():
        if k in {a["ref"] for a in answered}:
            continue
        answered.append({"ref": k, "kind": "Programme input", "text": str(inputs.get(k, {}).get("need", ""))[:300],
                         "status": "Answered" if v[0] == "Answered" else "Partly answered", "clause": v[1],
                         "answer": v[2], "remaining": v[3] or "-", "src": "Programme inputs workbook v1.2"})
    for r in inputs.values():
        if r["status"] == "Answered" and r["ref"] not in {a["ref"] for a in answered}:
            answered.append({"ref": r["ref"], "kind": "Programme input", "text": str(r["need"])[:300],
                             "status": "Answered earlier", "clause": "BDOI answer recorded in the programme inputs "
                             "workbook v1.2", "answer": str(r["response"] or r["proposal"] or "-")[:300],
                             "remaining": "-", "src": "Programme inputs workbook v1.2"})
    results = ["Not run", "Passed", "Failed", "Blocked", "Not applicable"]
    wb = Book("Test Cases and Traceability", doc_type="Requirements traceability", brd="BRD-03", version="1.0",
              date=DATE, subtitle="BRD-03 Product Maintenance – FRS v1.2 in BDOI's template")
    src = Counter(t["source"] for t in tcs)
    summary = [
        {"k": "BRD requirement IDs", "v": len(m), "n": "Every BRD ID has at least one test case (sheet "
                                                         "Traceability, column Test cases)"},
        {"k": "FRPM sub-items (BDOI's and added)", "v": len(items), "n": "Every FRPM sub-item has at least one test "
                                                                       "case"},
        {"k": "Acceptance criteria of the FRPM sub-items", "v": sum(len(i["ac"]) for i in items),
         "n": "One acceptance test case per criterion (TC-FRPM...-ACn)"},
        {"k": "Test cases in total", "v": len(tcs), "n": "; ".join(f"{k}: {v}" for k, v in src.items())},
        {"k": "Positive / negative test cases", "v": f"{sum(t['type'] == 'Positive' for t in tcs)} / "
                                                    f"{sum(t['type'] == 'Negative' for t in tcs)}", "n": "-"},
        {"k": "Traceability rows (BRD ID x FRPM item)", "v": len(trace), "n": "-"},
        {"k": "Items answered in BDOI's FRS (dropped from the open questions)", "v": len(answered),
         "n": "Sheet 'Answered in the BDOI FRS'"}]
    for rid in m:
        n = sum(t["n"] for t in trace if t["brd"] == rid)
        summary.append({"k": f"  {rid}", "v": n, "n": "test cases (with repeats across FRPM items)"})
    wb.sheet("Summary", [Column("k", "Item", 50, "What is counted"), Column("v", "Number", 14, "Count"),
                         Column("n", "Note", 80, "Proof of coverage")], summary,
             description="Counts, and the proof that every BRD ID and every FRPM item is covered")
    wb.sheet("Traceability", [
        Column("brd", "BRD ID", 11, "Requirement ID of the BRD"), Column("text", "BRD requirement", 38, "BRD text"),
        Column("frpm", "BDOI FR item", 14, "FRPM sub-item of the FRS v1.2"),
        Column("title", "FR item title", 26, "Title of the sub-item"),
        Column("ref", "BIBS reference FR", 18, "FR of the BIBS FRS BRD-3 (or BRD-1)"),
        Column("n", "Test cases (count)", 10, "Number of test cases"),
        Column("tc", "Test cases", 46, "Test case IDs (first twelve)"), Column("screen", "Screen", 34, "Screens"),
        Column("wt", "Walkthrough steps", 22, "Steps of the walkthrough script")], trace,
        description="BRD ID -> BDOI FR item -> BIBS reference FR -> test cases -> screen -> walkthrough step")
    rows = [{**t, "result": "Not run", "tester": None, "date": None, "obs": None, "remarks": None} for t in tcs]
    wb.sheet("Test cases", [
        Column("id", "Test case ID", 20, "Test case"), Column("title", "Title", 34, "What is tested"),
        Column("brd", "BRD ID", 12, "BRD requirement(s)"), Column("frpm", "FR item (BDOI)", 16, "FRPM sub-item(s)"),
        Column("ref", "BIBS reference", 14, "FR of the BIBS FRS"), Column("persona", "Persona", 18, "Who tests"),
        Column("signin", "Sign-in ID", 10, "UAT sign-in ID (walkthrough users workbook)"),
        Column("pre", "Preconditions", 28, "Before the test"), Column("steps", "Steps", 46, "Steps"),
        Column("data", "Test data", 16, "Data sets of the test plan"),
        Column("expected", "Expected result", 40, "What the tester should see"),
        Column("type", "Type", 10, "Positive or negative", values=["Positive", "Negative"]),
        Column("priority", "Priority", 9, "Priority", values=["High", "Medium", "Low"]),
        Column("screen", "Screen", 26, "Screen"), Column("source", "Source", 22, "Where the case comes from"),
        Column("result", "Result", 12, "Execution result", values=results, status=True),
        Column("tester", "Tested by", 16, "Name"), Column("date", "Date", 12, "dd-MMM-yyyy", kind="date"),
        Column("obs", "Observation ref", 14, "Issue or observation number"),
        Column("remarks", "Remarks", 26, "Remarks")], rows,
        description="The test cases of the FRS v1.2 with their execution columns")
    wb.sheet("Answered in the BDOI FRS", [
        Column("ref", "Ref.", 12, "Reference in the BIBS FRS or the programme inputs workbook"),
        Column("kind", "Kind", 16, "Assumption, dependency, open question, proposed rule, programme input"),
        Column("text", "Item", 46, "The item as asked"),
        Column("status", "Status", 16, "Answered, partly answered or answered earlier",
               values=["Answered", "Partly answered", "Answered earlier"]),
        Column("clause", "BDOI clause", 26, "Clause of BDOI's FRS that answers it"),
        Column("answer", "Answer in BDOI's FRS", 46, "What BDOI's FRS says"),
        Column("remaining", "Still asked", 34, "The remaining part, asked in FRS Annex V"),
        Column("src", "Source of the item", 22, "Where the item was asked")], answered,
        description="The open items dropped from the FRS because BDOI's FRS answers them, with BDOI's clause")
    crit = [
        {"k": "Entry", "c": "The FRS v1.2 and this workbook are signed off; open decisions affecting the tested "
                            "steps are decided (FRS Annex V, Annex M)."},
        {"k": "Entry", "c": "The BIBS UAT environment is available with the SIT/UAT data and the users of the "
                            "walkthrough users workbook; passwords issued by the iorta TechNXT Project Team."},
        {"k": "Entry", "c": "Configuration of Annex U loaded for the lines tested (lists, parameters, products, "
                            "insurers, templates)."},
        {"k": "Entry", "c": "Business Unit testers named for each role (RC-APR-07, TD-03)."},
        {"k": "Exit", "c": "100% of the High priority test cases run; no open Severity 1 or 2 issue."},
        {"k": "Exit", "c": "At least 95% of all test cases passed; every failed case has an agreed resolution "
                           "date or change request."},
        {"k": "Exit", "c": "Every BRD ID has at least one passed test case."},
        {"k": "Exit", "c": "The walkthrough script of the walkthrough users workbook is completed by the Business "
                           "Unit with its results."},
        {"k": "Exit", "c": "The UAT sign-off certificate of BRD-03 is signed (SO-03-3)."}]
    wb.sheet("UAT entry and exit criteria", [Column("k", "Kind", 10, "Entry or exit", values=["Entry", "Exit"]),
                                             Column("c", "Criterion", 110, "The criterion")], crit,
             description="When UAT of Product Maintenance starts and when it is complete")
    path = wb.save(OUT_DIR / "BIBS_RTM_BRD-03_Test_Cases_and_Traceability_v1.0.xlsx")
    return path, {"trace": len(trace), "tcs": len(tcs), "answered": len(answered)}


# =============================================================================================== 5. UAT users
def build_uat() -> tuple[Path, int, int]:
    sys.path.insert(0, str(REPO / "docs" / "deliverables" / "src" / "programme" / "uat"))
    import build_uat_users as uat  # noqa: PLC0415
    model = uat.Model()
    b = next(x for x in model.brds if x["brd"] == "BRD-03")
    sys.path.insert(0, str(REPO / "docs" / "deliverables" / "src" / "signoff"))
    import signoff_pack  # noqa: PLC0415
    pack = signoff_pack._pack(REPO / "docs" / "deliverables" / "src" / "BRD-03_Product_Maintenance" / "pack" /
                              "pack.yaml")
    menus = pack.menus()
    users = []
    for r in uat.brd_rows(model, b):
        own = [x["screen"] for x in menus.get(r["code"], []) if x["own"] == "Yes"]
        users.append({"persona": r["persona"], "name": r["name"], "id": r["id"], "profile": r["profile"],
                      "unit": r["place"], "does": r["does"], "menu": r["menu"],
                      "menus": "Product Maintenance › " + ", ".join(own) if own else "-", "steps": r["steps"],
                      "scope": r["scope"], "env": ENVIRONMENT, "password": PW_TEXT, "remarks": r["remarks"]})
    script = [{"step": s["step"], "flow": s["flow"], "persona": s["persona"], "id": ", ".join(s["users"]) or "-",
               "action": s["action"], "expected": s["result"], "result": "Not run", "by": None, "date": None,
               "remarks": None} for s in b["script"]]
    results = ["Not run", "Passed", "Failed", "Blocked"]
    wb = Book("Business Unit Walkthrough Users", doc_type="UAT users", brd="BRD-03", version="1.0", date=DATE,
              subtitle="BRD-03 Product Maintenance – who signs in for the walkthrough")
    wb.cover_notes = [f"Environment: {ENVIRONMENT}.", f"Passwords: {PW_TEXT}; no password is written in this "
                      "workbook."]
    wb.sheet("Users", [
        Column("persona", "Persona", 20, "Business role in the FRS"), Column("name", "Name", 24, "User name"),
        Column("id", "Sign-in ID", 11, "User ID in the BIBS UAT environment"),
        Column("profile", "Role / group profile", 26, "Group profile of the user"),
        Column("unit", "Company / branch / unit", 30, "Where the user belongs"),
        Column("does", "What the user does in the walkthrough", 36, "Responsibilities"),
        Column("menu", "Menu path to start", 34, "Where the user starts"),
        Column("menus", "Product Maintenance menus the user should see", 40, "Entries of the menu"),
        Column("steps", "Walkthrough steps", 24, "Steps of the script"),
        Column("scope", "Limits or data scope", 24, "Approval limits and data scope"),
        Column("env", "Environment", 26, "Where to sign in"), Column("password", "Password", 26, "How the password "
                                                                     "is issued"),
        Column("remarks", "Remarks", 26, "Remarks")], users,
        description="The users of BRD-03 for the Business Unit walkthrough, with their sign-in IDs")
    wb.sheet("Walkthrough script", [
        Column("step", "Step", 9, "Step number"), Column("flow", "Walkthrough", 30, "Walkthrough"),
        Column("persona", "Persona", 20, "Who acts"), Column("id", "Sign-in ID", 11, "User ID"),
        Column("action", "Action", 50, "What the user does"), Column("expected", "Expected result", 40,
                                                                     "What the user sees"),
        Column("result", "Result", 11, "Result", values=results, status=True), Column("by", "Done by", 16, "Name"),
        Column("date", "Date", 12, "dd-MMM-yyyy", kind="date"), Column("remarks", "Remarks", 26, "Remarks")], script,
        description="The walkthrough of BRD-03 step by step, with the sign-in ID of each step")
    path = wb.save(OUT_DIR / "BIBS_UAT_BRD-03_Business_Unit_Walkthrough_Users_v1.0.xlsx")
    return path, len(users), len(script)


# ============================================================================================ 6. coverage (internal)
RANK = {"Yes": 2, "Partly": 1, "No": 0}


def build_coverage() -> tuple[Path, dict]:
    conf = load_conformance()
    comp = {m["id"]: m for m in yaml.safe_load(CMP.read_text(encoding="utf-8"))["mapping"]}
    tcs = v12.test_cases()
    items = v12.frpm_items()
    ac_rows, item_rows = [], []
    for it in items:
        refs = it["refs"]
        res = [x for r in refs for x in conf.get(r, [])]
        base, why = "Yes", "The reference requirements are shown working in the end-to-end conformance run"
        if any(x[2] == "Not conformant" for x in res):
            base, why = "Partly", "Part of a reference requirement is not conformant in the conformance register"
        elif any(x[2] == "Not testable yet" for x in res):
            base, why = "Partly", "Part of a reference requirement is not testable yet (conformance register)"
        elif not res:
            base, why = "Partly", "No conformance evidence for the reference requirements; to be shown in UAT"
        ev_conf = "; ".join(f"{r}: " + ", ".join(f"{k} {v}" for k, v in Counter(x[2] for x in conf.get(r, [])).items())
                            for r in refs if conf.get(r)) or "-"
        acs = []
        for i, ac in enumerate(it["ac"], 1):
            ov = (ITEMS["coverage"].get(it["id"]) or {}).get(i)
            avail, reason = (ov[0], ov[1]) if ov else (base, why)
            acs.append(avail)
            ac_rows.append({"item": it["id"], "group": it["item"], "n": i, "ac": ac, "avail": avail,
                            "tc": f"TC-{it['id']}-AC{i}", "conf": ev_conf, "note": reason})
        agg = min(acs, key=lambda a: RANK[a]) if acs else base
        if it["id"] in comp:
            p = comp[it["id"]]["platform"]
            avail = {"Yes": "Yes", "Partly": "Partly", "No": "No"}[p]
            if avail == "No" and any(a != "No" for a in acs):
                avail = "Partly"
            if RANK[agg] < RANK[avail]:
                avail = agg
            met = "BDOI's text is met" + (f": {comp[it['id']]['evidence']}" if comp[it["id"]].get("evidence") else "")
            note = (f"BDOI's text: {comp[it['id']]['action']} – {comp[it['id']]['action_detail']}" if p != "Yes" else
                    met) + ("" if agg == "Yes" else "; some acceptance criteria partly available")
        else:
            avail, note = agg, ("Added in v1.1 from the BIBS reference FRS; " + why) if agg == "Yes" else \
                "Added in v1.1; some acceptance criteria partly available"
        cases = [t["id"] for t in tcs if it["id"] in t["frpm_list"]]
        item_rows.append({"group": it["item"], "id": it["id"], "title": it["title"], "origin": it["origin"],
                          "avail": avail, "tcs": len(cases), "tc": ", ".join(cases[:8]) + (" ..." if len(cases) > 8
                                                                                        else ""),
                          "wt": v12._wt_of(refs), "conf": ev_conf, "note": note.replace("Platform change", "Planned "
                                                                                        "extension")})

    def pct(rows, key="avail"):
        n = len(rows)
        y = sum(r[key] == "Yes" for r in rows)
        p = sum(r[key] == "Partly" for r in rows)
        return n, y, p, (100.0 * y / n if n else 0), (100.0 * (y + 0.5 * p) / n if n else 0), \
            (100.0 * (y + p) / n if n else 0)
    groups = OrderedDict()
    for r in item_rows:
        groups.setdefault(r["group"], {"items": [], "acs": []})["items"].append(r)
    for r in ac_rows:
        groups[r["group"]]["acs"].append(r)
    summary = []
    for name, rows in [("Overall", (item_rows, ac_rows))] + [(g, (v["items"], v["acs"])) for g, v in groups.items()]:
        n, y, p, fy, fw, fa = pct(rows[0])
        an, ay, ap, afy, afw, afa = pct(rows[1])
        summary.append({"g": name, "n": n, "y": y, "p": p, "no": n - y - p, "fy": round(fy, 1), "fw": round(fw, 1),
                        "an": an, "ay": ay, "ap": ap, "ano": an - ay - ap, "afy": round(afy, 1),
                        "afw": round(afw, 1)})
    av = ["Yes", "Partly", "No"]
    wb = Book("FRS Coverage in BIBS", doc_type="Internal quality record", brd="BRD-03", version="1.0", date=DATE,
              subtitle="BRD-03 Product Maintenance – FRS v1.2 in BDOI's template", classification="Internal")
    wb.cover_notes = ["Internal record of the iorta TechNXT Project Team; shared with BDOI only if the product owner "
                      "decides.", "Sources: the end-to-end conformance register (Quality), the platform status of "
                      "the comparison workbook (BRD-03), the BIBS platform and the test cases of the traceability "
                      "workbook. 'Partly' gives the reason; nothing is assumed.",
                      "Weighted % counts a partly available item as one half."]
    wb.sheet("Summary", [
        Column("g", "FR group", 12, "Overall or FRPM item"), Column("n", "Sub-items", 9, "Number of sub-items"),
        Column("y", "Available", 9, "Available (Yes)"), Column("p", "Partly", 8, "Partly available"),
        Column("no", "Not yet", 8, "Not available yet"), Column("fy", "% available (items)", 11, "Yes / all",
                                                                kind="percent"),
        Column("fw", "% weighted (items)", 11, "(Yes + half of Partly) / all", kind="percent"),
        Column("an", "Criteria", 9, "Acceptance criteria"), Column("ay", "Available", 9, "Yes"),
        Column("ap", "Partly", 8, "Partly"), Column("ano", "Not yet", 8, "No"),
        Column("afy", "% available (criteria)", 11, "Yes / all", kind="percent"),
        Column("afw", "% weighted (criteria)", 11, "(Yes + half of Partly) / all", kind="percent")], summary,
        description="Percentage of the FRS available in the BIBS environment, by item and by acceptance criterion")
    wb.sheet("FR items", [
        Column("group", "FR group", 10, "FRPM item"), Column("id", "FR sub-item", 12, "FRPM sub-item"),
        Column("title", "Title", 30, "Title"), Column("origin", "Origin", 16, "BDOI's or added"),
        Column("avail", "Available in BIBS", 11, "Yes, Partly or No", values=av, status=True),
        Column("tcs", "Test cases", 8, "Number of test cases"), Column("tc", "Test cases (first eight)", 34, "IDs"),
        Column("wt", "Walkthrough steps", 22, "Steps"), Column("conf", "Conformance register", 34, "Results of the "
                                                               "reference requirements"),
        Column("note", "Note", 50, "Why partly or not yet")], item_rows,
        description="Availability of every FRPM sub-item (BDOI's and added) in the BIBS environment")
    wb.sheet("Acceptance criteria", [
        Column("item", "FR sub-item", 12, "FRPM sub-item"), Column("n", "No.", 5, "Criterion number"),
        Column("ac", "Acceptance criterion", 50, "Criterion of the FRS"),
        Column("avail", "Available in BIBS", 11, "Yes, Partly or No", values=av, status=True),
        Column("tc", "Test case", 20, "Acceptance test case"), Column("conf", "Conformance register", 34,
                                                                      "Evidence"),
        Column("note", "Note", 50, "Reason")], ac_rows,
        description="Availability of every acceptance criterion of the FRS v1.2")
    path = wb.save(QUALITY / "BIBS_Coverage_BRD-03_FRS_Coverage_in_BIBS_v1.0.xlsx")
    if {r["id"] for r in item_rows} != {i["id"] for i in items}:
        raise SystemExit("coverage rows do not match the FRPM items")
    return path, summary[0]


# ========================================================================================================== main
def main() -> int:
    our, nb = core.our_frs(), core.nb_frs()
    ids = core.brd_ids(our)
    b = core.build(False, our, nb, ids)
    v12.cleanup(b)
    items = collection_items()
    rc = {r["id"] for rows in items.values() for r in rows}
    missing = [i["rc"] for i in v12.open_items()["kept"] if i["rc"] not in rc]
    missing += [x for x in ("RC-DOC-01", "RC-DOC-02", "RC-DOC-03", "RC-DOC-04", "RC-DOC-05", "RC-DOC-06",
                            "RC-RPT-01", "RC-APR-02") if x not in rc]
    if missing:
        raise SystemExit(f"items named in the FRS but not in the collection workbook: {missing}")
    p1 = build_collection(items)
    p2, fit = build_fitgap()
    if {r["id"] for r in fit} != set(ids):
        raise SystemExit("fit-gap rows do not match the BRD IDs")
    p3, ncr, nobs = build_cr(b)
    if nobs != b.stats["observations"]:
        raise SystemExit("decisions log and Annex M differ")
    p4, rtm = build_rtm()
    p5, nusers, nsteps = build_uat()
    p6, cov = build_coverage()
    for p in (p1, p2, p3, p4, p5, p6):
        text = " ".join(str(c.value) for ws in openpyxl.load_workbook(p).worksheets for row in ws.iter_rows()
                        for c in row if c.value is not None)
        if re.search(r"(?i)password\s*[:=]\s*(?!<)\w", text):
            raise SystemExit(f"a password value in {p.name}")
        print(f"wrote {p.relative_to(REPO)} ({p.stat().st_size // 1024} KB)")
    print(f"collection: " + ", ".join(f"{k} {len(v)}" for k, v in items.items()) + f" (total "
          f"{sum(len(v) for v in items.values())})")
    print(f"fit-gap: {len(fit)} BRD IDs {dict(Counter(r['fit'] for r in fit))}; CR register: {ncr} candidates, "
          f"{nobs} decisions; RTM: {rtm}; UAT: {nusers} users, {nsteps} steps")
    print(f"coverage: items {cov['y']}/{cov['n']} available, {cov['p']} partly ({cov['fy']}%, weighted {cov['fw']}%); "
          f"criteria {cov['ay']}/{cov['an']} ({cov['afy']}%, weighted {cov['afw']}%)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
