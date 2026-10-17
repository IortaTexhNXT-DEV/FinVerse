"""Renewal document set v1.3 (BRD-06) in BDOI's template: FRS, BPS, UXD, RNO and the Interface Specification v1.0.

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd06_v13.py [--no-pdf] [--only FRS,BPS,UXD,RNO,ICD]

Each document starts from BDOI's FRS template (template D051, templates/bdoi_d051_frs_template.docx) through
pack_common.PackDoc: cover, document details, headers and footers, styles, revision log and table of contents are
kept; the body is written from brd06_v13_content.py (FRS), brd06_v13_pages.py (process, storyboards, outputs, interfaces)
and brd06_v13_screens.py (screens). The screenshots are in figures/brd06_v13_screens (flows/brd06/shots_v13.sh),
one full page each.

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-06_Renewal/ (PDFs not kept in git).
Self-checks: every screenshot named exists; every BRD row of the mapping points to an FRRN item; every screen of the
UX Design has a screenshot.
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[4] / "tools" / "deliverables"))
import brd06_v13_content as C  # noqa: E402
import brd06_v13_pages as P  # noqa: E402
import brd06_v13_screens as S  # noqa: E402
import pack_common as pc  # noqa: E402
from build_brd11_interface_spec import _run  # noqa: E402
from colour_code import COLOURS, LABELS  # noqa: E402

OUT = pc.OUT_BDOI / "BRD-06_Renewal"
SHOTS = HERE / "figures" / "brd06_v13_screens"
FLOWS = HERE / "flows" / "brd06"
COL = {"k": None, "g": COLOURS["green"], "o": COLOURS["orange"], "b": COLOURS["blue"], "": None}
COLOUR_NAME = {"k": "Black", "g": "Green", "o": "Orange", "b": "Blue", "": "-"}
CREATED = {"FRS": "06/11/2026"}


def _colour(run, hexcol: str) -> None:
    rpr = run._r.get_or_add_rPr()
    for c in rpr.findall(qn("w:color")):
        rpr.remove(c)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), hexcol)
    rpr.append(c)


class Doc(pc.PackDoc):
    """A document of the Renewal set."""

    def __init__(self, key: str, revisions):
        m = C.DOCS[key]
        version = C.ICD_VERSION if key == "ICD" else C.VERSION
        super().__init__(pc.DocMeta(file=m["file"], cover=m["cover"], header_left=m["header_left"],
                                    header_right=C.HEADER_RIGHT, doc_no=m["doc_no"], version=version,
                                    date_us=C.DATE_US, created_us=CREATED.get(key, C.DATE_US), revisions=revisions))
        self.key = key
        self.shots = 0

    def cpara(self, text: str, colour: str = "k") -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.64)
        p.paragraph_format.space_after = Pt(4)
        r = _run(p, text)
        if COL.get(colour):
            _colour(r, COL[colour])

    def ctable(self, headers, rows, widths, caption=None, colours=None, colour_cols=None):
        self.table(headers, rows, widths, caption=caption)
        tbl = self.doc.tables[-1]
        for i, code in enumerate(colours or []):
            if not COL.get(code):
                continue
            for c_i, cell in enumerate(tbl.rows[i + 1].cells):
                if colour_cols is None or c_i in colour_cols:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            _colour(r, COL[code])
        return tbl

    def screen(self, shot: str, caption: str) -> None:
        """One full-page screenshot on its own page, scaled to fit the page, never split."""
        png = SHOTS / f"{shot}.png"
        if not png.exists():
            raise SystemExit(f"missing screenshot {png.name}")
        w, h = Image.open(png).size
        width = min(16.0, 21.5 * w / h)
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        p.alignment = 1
        p.add_run().add_picture(str(png), width=Cm(width))
        self.caption(caption, kind="Figure")
        self.shots += 1

    def legend(self) -> None:
        self.para(C.LEGEND_NOTE)
        rows = [["Black", "In the BRD", "Asked by the BRD, BRD-00 or a programme NFR.", "In scope"],
                ["Green", "Not in the BRD: OOTB", "Offered by the iNXT BrokerVerse OOTB platform.", "No cost"],
                ["Orange", "Not in the BRD: change request", "Not offered by the OOTB platform.", "Only if approved"],
                ["Blue", "Not in the BRD: best practice", "Recommendation, separate table.", "Indicative effort"],
                ["Grey label", LABELS["integration"] + " / " + LABELS["migration"],
                 "Interfaces with other BDOI systems, or Data Migration (BRD-13).", "Not a change request"]]
        self.ctable(["Colour", "Meaning", "What it covers", "Cost"], rows, [1.2, 2.4, 4.2, 1.6],
                    caption="Colour code", colours=["k", "g", "o", "b", ""], colour_cols={0, 1})
        tbl = self.doc.tables[-1]
        for r in tbl.rows[5].cells[0].paragraphs[0].runs + tbl.rows[5].cells[1].paragraphs[0].runs:
            _colour(r, COLOURS["grey"])

    def personas(self, caption="Personas of Renewal") -> None:
        self.table(["Persona", "Does", "Screens", "Walkthrough account"], C.PERSONAS, [1.8, 3.6, 2.6, 1.6],
                   caption=caption)


def _dot(src: str, name: str, tmp: Path) -> Path:
    dot = tmp / f"{name}.dot"
    dot.write_text(src, encoding="utf-8")
    png = tmp / f"{name}.png"
    subprocess.run(["dot", "-Tpng", "-Gdpi=170", str(dot), "-o", str(png)], check=True)
    return png


def _lifecycle(tmp: Path) -> Path:
    return _dot(
        'digraph g { rankdir=LR; graph [fontname="Arial", nodesep=0.25, ranksep=0.3]; node [shape=box, '
        'style="rounded", fontname="Arial", fontsize=10]; edge [fontname="Arial", fontsize=8];\n'
        'E [label="Extracted"]; U [label="Unassigned\\nDisposition"]; T [label="Transfer\\nPending"]; '
        'D [label="For Disposition"]; R [label="Review in\\nProgress"]; N [label="For Proposal /\\nNew Business '
        'Path"]; P [label="For Processing /\\nIn Processing"]; W [label="With Insurer"]; A [label="RA Ready,\\n'
        'Generated, Sent"]; C [label="Accepted / For\\nPlacement and Booking"]; L [label="Letter Pending"]; '
        'Z [label="Renewed"]; X [label="Closed"];\n'
        'E -> U [label="Initiate"]; E -> P [label="matrix"]; E -> L [label="matrix"]; U -> D [label="Assign"]; '
        'U -> T [label="Transfer"]; T -> U [label="Accept / Decline"]; D -> R [label="Submit"]; '
        'R -> D [label="Return"]; R -> P [label="Post"]; R -> N [label="Post"]; R -> L [label="Post"]; '
        'P -> W [label="Batch"]; W -> A [label="Renew As Is"]; W -> D [label="Revise / Reject"]; '
        'A -> C [label="Acceptance"]; A -> X [label="Non-acceptance"]; C -> Z [label="Booking"]; N -> Z; N -> X; '
        'L -> X [label="Letter sent"]; }', "lifecycle", tmp)


def _pair(rows):
    return [list(r) for r in rows]


# ================================================================================================ FRS
def write_frs(tmp: Path) -> Doc:
    d = Doc("FRS", C.REVISIONS)
    i = C.FRS_INTRO
    d.h1("Introduction")
    d.h2("Document Purpose")
    d.para(i["purpose"])
    d.h2("Intended Audience")
    d.para(i["audience"])
    d.h2("Project Background")
    for t in i["background"]:
        d.para(t)
    d.h2("Scope")
    n = max(len(i["in_scope"]), len(i["out_scope"]))
    d.table(["In scope (BIBS)", "Out of scope or in another FRS"],
            [[(i["in_scope"] + [""] * n)[k], (i["out_scope"] + [""] * n)[k]] for k in range(n)], [4.6, 4.6],
            caption="Scope of Renewal in BIBS")
    d.h2("Document Set")
    d.table(["Document", "Content"], C.DOC_SET, [3.2, 6.0], caption="Renewal document set")
    d.h2("How to Read This Document")
    d.legend()
    d.h2("References")
    d.bullets(C.REFERENCES)

    d.h1("Business Requirements Mapping")
    d.para("Every BRRN requirement and every section of the main BRD maps to items of section 3. The 1,033 "
           "requirement lines of the main BRD are mapped one by one in the RTM workbook, sheet BRD Lines.")
    d.table(["BRD Req No.", "Business requirement", "FRRN items", "Where", "Change from v1.2"], C.MAPPING,
            [1.7, 3.0, 2.3, 0.9, 1.4], caption="Business requirements mapping")

    d.h1("Functional Requirements")
    d.para("Each item keeps BDOI's FRRN number. Status gives the change from version 1.2: Kept, Changed or New. The "
           "Example column holds the first acceptance criterion; all the criteria are in the User Stories v1.0 and "
           "the test cases of the RTM workbook.")
    d.table(["FRRN", "Component", "Capability", "Status", "BRD"],
            [[f["id"], f["component"], f["capability"], f["status"], f["brd"]] for f in C.FRUM],
            [1.3, 1.8, 3.0, 1.1, 2.1], caption="Functional requirements at a glance")
    for f in C.FRUM:
        d.h2(f"{f['id']} {f['capability']}")
        d.para(f"Status: {f['status']}. BRD: {f['brd']}.", italic=True)
        rows, cols = [], []
        for rid, text, colour, label, rule, ex in f["reqs"]:
            lab = {"int": "  " + LABELS["integration"], "dm": "  " + LABELS["migration"]}.get(label, "")
            rows.append([rid, text + lab, rule, ex])
            cols.append(colour)
        d.ctable(["ID", "Requirement", "Rule or setting", "Example"], rows, [1.8, 4.2, 1.9, 1.9], colours=cols,
                 colour_cols={1, 2})

    d.h1("Business Rules")
    d.ctable(["ID", "Rule", "Colour"], [[a, b, COLOUR_NAME[c]] for a, b, c in C.BUSINESS_RULES], [1.3, 6.8, 1.2],
             caption="Business rules", colours=[c for *_, c in C.BUSINESS_RULES], colour_cols={1, 2})
    d.h1("Validations and Messages")
    d.para("The main messages of the module. The full list with the message codes of the services is in the "
           "Developer Specification; the field checks of each screen are in the UX Design.")
    d.table(["Screen or step", "Condition", "Message"], C.MESSAGES, [2.2, 3.1, 3.9], caption="Validations and "
            "messages")
    d.h1("Personas and Permissions")
    d.personas()
    d.table(C.PERMISSION_HEADS, C.PERMISSIONS, [2.8] + [0.64] * 10, caption="What each persona may do (Y = yes); "
            "the rules of section 4 apply whatever the profile grants")
    d.h1("Parameters and Lists of Values")
    d.para("Settings are changed on System Parameters and applied after another user approves the change (BR-RN-01). "
           "Keys, types and ranges are in the Data Dictionary, sheet Parameters.")
    d.ctable(["Setting", "Default", "Range", "Area"], [r[:4] for r in C.PARAMETERS], [4.0, 2.4, 1.6, 1.2],
             caption="Renewal settings", colours=[r[4] for r in C.PARAMETERS])
    d.table(["List", "Values"], C.LOVS, [3.0, 6.2], caption="Lists of values (codes in the Data Dictionary)")
    d.h1("Statuses")
    d.figure(_lifecycle(tmp), "Stages of a renewal account", width_cm=16.0)
    for name, rows in C.STATUSES.items():
        d.h2(name)
        d.table(["Status", "Meaning", "Next"], rows, [2.4, 3.6, 3.2], caption=f"Statuses: {name.lower()}")
    d.h2("BDOI's status names and the stages")
    d.table(["BDOI's status", "Stage in BIBS"], C.BDOI_STATUS_MAP, [4.4, 4.8],
            caption="Status names shown on the lists by default (setting RNW_STATUS_NAMES; Q-RN-05)")
    d.h1("Assumptions, Dependencies and Open Questions")
    d.table(["ID", "Assumption"], C.ASSUMPTIONS, [1.3, 7.9], caption="Assumptions")
    d.table(["ID", "Dependency"], C.DEPENDENCIES, [1.3, 7.9], caption="Dependencies")
    d.table(["ID", "Question", "Owner", "Answer by"], C.OPEN_QUESTIONS, [1.2, 5.0, 1.8, 1.2],
            caption="Open questions")
    d.h1("Change Control Summary")
    d.para(f"Change requests of the register v{C.WB_VERSION}. Effort in man-days (analysis and design, backend, "
           "frontend, testing, documentation and UAT support).")
    d.ctable(["CR", "Title", "Status", "FRS", "Man-days", "Note"],
             [[a, b, c, e, f"{sum(eff):g}", n] for a, b, c, e, eff, n in C.CRS], [1.3, 2.9, 0.9, 1.2, 0.9, 2.0],
             caption="Change requests", colours=["o" if c == "Open" else "" for _, _, c, *__ in C.CRS])
    d.ctable(["No.", "Not in the BRD, offered by the OOTB platform", "FRS"], C.OOTB_POINTS, [1.3, 5.6, 2.3],
             caption="Green points (no cost)", colours=["g"] * len(C.OOTB_POINTS), colour_cols={1})
    d.table(["No.", "Integration and data migration scope", "FRS"], C.INTEGRATION_POINTS, [1.3, 5.6, 2.3],
            caption="Grey label points (not change requests)")
    d.ctable(["No.", "Best-practice recommendation", "Priority", "Man-days"],
             [[a, b, c, str(e)] for a, b, c, e in C.BEST_PRACTICE], [1.3, 5.0, 1.4, 1.4],
             caption="Blue points (recommendations)", colours=["b"] * len(C.BEST_PRACTICE), colour_cols={1})
    d.h1("Changes from Version 1.2")
    d.table(["Item", "Change"], C.CHANGES_V12, [3.0, 6.2], caption="Changes from version 1.2")
    return d


# ================================================================================================ BPS
def write_bps(tmp: Path) -> Doc:
    d = Doc("BPS", C.REVISION_COMPANION)
    d.h1("Introduction")
    d.para(f"This document shows how Renewal works end to end. The requirements are in the FRS v{C.VERSION}; the "
           f"screens in the UX Design v{C.VERSION}; the reports, letters and notices in the Reports, Notifications "
           f"and Outputs v{C.VERSION}.")
    d.para("Each storyboard step is illustrated by a full-page screenshot taken on a test environment. The people, "
           "clients and insurers are fictitious.")
    d.personas("Personas")
    d.h1("End-to-end Journey")
    d.para(f"{len(P.JOURNEY)} steps from the expiring policy to the booked renewal.")
    d.figure(FLOWS / "rn_simple.png", "Renewal: simplified process flow (swimlanes by actor)", width_cm=16.0)
    d.table(["Step", "Who", "What happens"], P.JOURNEY, [0.9, 2.0, 6.3], caption=f"The {len(P.JOURNEY)} steps")
    d.h1("To-be Processes")
    for p in P.PROCESSES:
        d.h2(f"{p['id']} {p['name']}")
        d.table(["Item", "Value"], [["Journey steps", p["steps"]], ["Trigger", p["trigger"]], ["Actors", p["actors"]],
                                    ["Outcome", p["outcome"]], ["Controls", p["controls"]]], [2.0, 7.2])
    d.h2("Stages of a renewal account")
    d.figure(_lifecycle(tmp), "Stages of a renewal account", width_cm=16.0)
    d.h1("Storyboards")
    for sb in P.STORYBOARDS:
        d.h2(f"{sb['id']} {sb['title']}")
        d.para(sb["summary"])
        d.para(f"Test data: {sb['data']}.", italic=True)
        d.table(["Step", "Who", "Screen", "Does", "Sees"],
                [[str(k + 1), w, s, a, r] for k, (w, s, a, r, _) in enumerate(sb["steps"])],
                [0.7, 1.9, 1.8, 2.6, 2.2], caption=f"{sb['id']} steps")
        for k, (w, s, a, r, shot) in enumerate(sb["steps"]):
            d.screen(shot, f"{sb['id']} step {k + 1}: {w}, {s}")
    return d


# ================================================================================================ UXD
def write_uxd(tmp: Path) -> Doc:
    d = Doc("UXD", C.REVISION_COMPANION)
    d.h1("Introduction")
    d.para("This document describes every Renewal screen of BIBS: purpose, users, access, rules and states, fields "
           "with their type, whether they are required, default, rule and message, and the actions. The requirements "
           f"and rules are in the FRS v{C.VERSION}. Each screen state is shown by one full-page screenshot on its own "
           "page, taken on a test environment with fictitious data.")
    d.para("Codes used: Y required, N optional, - display only. Dates show as dd-MMM-yyyy, times in Manila time, "
           "amounts with two decimals.")
    d.h1("Menu by Persona")
    d.para("The screens are under Client & Policy › Renewal. A persona sees only its entries.")
    d.table(["Screen", "Purpose", "Personas"],
            [[a, b, ", ".join(C.PERSONA_CODES[x] for x in c.split())] for a, b, c in C.MENU],
            [2.6, 3.6, 3.0], caption="Renewal menu by persona")
    d.h1("Screens")
    for s in S.SCREENS:
        d.h2(f"{s['id']} {s['name']}")
        d.table(["Item", "Value"], [["Purpose", s["purpose"]], ["Access", s["access"]],
                                    ["Requirements", s["frrn"]], ["Outcome", s["outcome"] or "-"]], [2.2, 7.0])
        if s["rules"]:
            d.para("Rules and states:")
            d.bullets(s["rules"])
        d.table(["Field", "Type", "Req.", "Default", "Rule", "Message"], s["fields"],
                [2.3, 1.1, 0.6, 1.2, 2.6, 1.6], caption=f"{s['id']} fields")
        d.table(["Action", "Who", "Available when", "What happens", "Result"], s["actions"],
                [1.6, 1.7, 1.8, 2.7, 1.4], caption=f"{s['id']} actions")
        for shot, cap, _ in s["shots"]:
            d.screen(shot, f"{s['id']}: {cap}")
    return d


# ================================================================================================ RNO
def write_rno(tmp: Path) -> Doc:
    d = Doc("RNO", C.REVISION_COMPANION)
    d.h1("Introduction")
    d.para("This document lists the reports, letters and documents, exports, notices and e-mails and the scheduled "
           "jobs of Renewal. Times are Manila time. Documents sent outside BIBS are password protected and the "
           "password follows separately (BR-RN-09).")
    d.h1("Reports")
    d.table(["ID", "Report", "BRD", "Parameters", "Content", "Output"], P.REPORTS, [1.2, 2.0, 1.2, 1.6, 2.2, 1.2],
            caption="Renewal reports (Report Centre › Renewal)")
    d.para("Every report applies the user's data scope, is logged with the user, parameters and time, and can be "
           "saved as a variant.")
    d.screen("ux-21-01-view", "Renewal Dashboard (FRRN.002.02)")
    d.screen("ux-23-01-view", "Processing Dashboard (FRRN.003)")
    d.h1("Letters and Documents")
    d.table(["ID", "Document", "Format and file name", "Protection", "FRRN"], P.DOCUMENTS, [0.9, 2.2, 3.2, 1.6, 1.3],
            caption="Letters and documents produced")
    d.h1("Exports")
    d.table(["Output", "Format", "Where", "Content"], P.EXPORTS, [2.0, 1.6, 2.8, 2.8], caption="Exports")
    d.h1("Notices and E-mails")
    d.para("Notices appear under the bell and, where shown, by e-mail; alerts on the Alerts screen. Text in angle "
           "brackets is filled in. Each user chooses the channels of the Renewal events on Notification Settings.")
    d.table(["ID", "Event", "To", "Text", "Channel"], P.NOTICES, [0.8, 2.4, 2.0, 3.0, 1.0],
            caption="Notices and e-mails")
    d.h1("Scheduled Jobs")
    d.table(["Job", "When", "What it does", "Notices"], P.JOBS, [2.6, 1.6, 4.0, 1.0], caption="Jobs")
    d.para("Every run is recorded with its start, end, outcome and counts, and can be started at once from "
           "Operations › Jobs by a user with the job permission.")
    return d


# ================================================================================================ ICD
def write_icd(tmp: Path) -> Doc:
    import code_facts as cf  # noqa: PLC0415
    d = Doc("ICD", [(C.DATE, C.ICD_VERSION, "First issue: CCM, MFT, insurer files, LAMD and BDOFC/SOLD reports and "
                     "CLPC billing files of Renewal.", C.AUTHOR)])
    d.h1("Introduction")
    d.para("This document specifies the interfaces of Renewal with systems outside BIBS. They are Integration scope "
           f"(FRS v{C.VERSION}, grey label). The links with other BIBS modules (New Business booking, Cashiering, "
           "Submitted Policies, Product Maintenance) are internal and not listed here.")
    d.table(["ID", "Interface", "Direction", "Means", "Frequency", "Content"], P.INTERFACES,
            [1.0, 1.8, 1.9, 1.3, 1.5, 1.8], caption="Interfaces")
    d.h1("IF-RN-01 CCM")
    d.bullets(["Each client document is queued Pending Transmission, protected with the password of its client "
               "(FRRN.024.01) and posted to CCM at <CCM address>/communications with the access key in the header "
               "X-Api-Key.",
               "CCM answers with its transaction reference; the message becomes Submitted to CCM. No reference, a "
               "refusal or an invalid recipient keeps the message Pending for the next attempt or fails it.",
               "The job RNW_CHANNEL_STATUS reads <CCM address>/communications/<reference> every 15 minutes and retries "
               "pending messages up to RNW_CHANNEL_MAX_ATTEMPTS (3).",
               "Settings: RNW_CCM_MODE (LIVE or SIMULATOR), RNW_CCM_ENDPOINT, RNW_CCM_CLIENT_ID, RNW_CCM_KEY_SETTING "
               "(name of the environment setting holding the key; the key is never stored in the database), "
               "RNW_CCM_TIMEOUT_SECONDS. RNW_DELIVERY_CHANNEL = EMAIL sends a protected e-mail instead."])
    d.table(["Field", "Content"], P.CCM_FIELDS, [3.2, 6.0], caption="CCM request")
    d.table(["CCM status", "Status in BIBS"], P.CCM_STATUS, [3.2, 6.0], caption="CCM delivery status")
    d.screen("ux-33-01-list", "Channel Monitor (SCR-RN-33)")
    d.h1("IF-RN-02 and IF-RN-03 MFT")
    d.bullets(["Outbound: a file for an insurer listed in RNW_MFT_ENROLLED_INSURERS (or whose placement channel is a "
               "file transfer) is put at <MFT address><RNW_MFT_OUTBOUND_FOLDER>/<insurer code>/<file name>; a put file is "
               "Delivered. Statuses: Pending Upload, Uploaded to MFT, Upload Failed (Processing Officer notified).",
               "Inbound: the job RNW_MFT_INBOX takes every 15 minutes the files of the folders placement-response, "
               "hold-cover-response and epolicy under the inbound folder, processes them as uploads by the user MFT, "
               "logs each as an inbound message (Received, Processed, Failed) and moves it to processed or refused.",
               "An e-policy summary file (.txt) waits for its ZIP file of the same name.",
               "Settings: RNW_MFT_MODE, RNW_MFT_ENDPOINT, RNW_MFT_KEY_SETTING, RNW_MFT_INBOUND_FOLDER, "
               "RNW_MFT_OUTBOUND_FOLDER, RNW_MFT_COMPANY_CODE."])
    d.h1("IF-RN-04 Insurer Files")
    d.table(["File", "Name", "Content"],
            [["Insurer request", "Excel per insurer batch", "28 columns of the BRD (Insurance Company to "
              "Packaged / Non-Packaged)"],
             ["Placement slip", "<Insurer>_<Line>_PlacementSlip_<ref>_MMDDYYYY.pdf", "One per insurer per account"],
             ["Placement file", "<Insurer>_<Line>_PlacementFile_MMDDYYYY.xlsx", "Packaged accounts of the insurer "
              "with each co-insurer's share"],
             ["Hold cover request", "Excel per insurer (CBG); PDF per account", ", ".join(P.HOLD_COVER_COLUMNS)],
             ["E-policy summary", ".txt, fields separated by |, tab or comma", "Sequence, reference, policy number, "
              "PDF file name, document tag; with a ZIP of the PDFs (at most 2,000 files, 200 MB)"]],
            [2.0, 3.4, 3.8], caption="Outbound and inbound insurer files")
    d.para("Inbound files are checked row by row; each row is matched on its reference, and every upload is listed "
           "with its file-level outcome and record-level results. The columns of each template are below.")
    for if_id, name, cls in P.INBOUND_HANDLERS:
        t = cf.bulk_templates([cls])
        if not t:
            raise SystemExit(f"no upload template {cls}")
        d.h2(f"{name} ({if_id})")
        d.table(["Column", "Required", "Type", "Content", "Example"],
                [[c["header"], "Y" if c["required"] else "N", c["type"].title(), c["description"] or "-",
                  c["example"] or "-"] for c in t[0].columns], [2.4, 0.8, 1.0, 3.4, 1.6])
    d.h1("IF-RN-05 CLPC Billing Files")
    d.bullets(["One CSV file per product line and kind (built-in for amortized accounts, non-built-in) with the date "
               "MMDDYYYY in its name, one per line and kind, from the job RNW_CLPC_BILLING or Generate Billing File.",
               "Each file is attached to its accounts, e-mailed to RNW_CLPC_BILLING_RECIPIENTS and written to the "
               "LMS directory (RNW_LMS_MODE LIVE, SIMULATOR or OFF); the upload status is Uploaded, Failed or Off."])
    d.table(["No.", "Column"], [[str(k), c] for k, c in enumerate(P.BILLING_COLUMNS, 1)], [1.0, 8.2],
            caption="Columns of the CLPC billing file (dates dd-MMM-yyyy)")
    d.h1("IF-RN-06 LAMD and BDOFC / BDOSOLD Reports")
    d.para("Uploaded on LAMD Reports by the LAMD user (layouts in IF-RN-04 above). PN is the matching key; serial and "
           "motor numbers are also matched for motor loans (RNW_LAMD_MOTOR_PRODUCTS).")
    d.h1("Open Points")
    d.table(["ID", "Point", "Owner"], [[a, b, "BDO ITG"] for a, b in C.DEPENDENCIES[:4]]
            + [["Q-RN-08", "CCM and MFT go-live dates; e-mail fall-back.", "BDO ITG"]], [1.4, 5.8, 2.0],
            caption="Open points")
    return d


WRITERS = {"FRS": write_frs, "BPS": write_bps, "UXD": write_uxd, "RNO": write_rno, "ICD": write_icd}


def checks() -> None:
    ids = {r[0] for f in C.FRUM for r in f["reqs"]} | {f["id"] for f in C.FRUM}
    for row in C.MAPPING:
        refs = [t.strip() for t in row[2].replace(";", ",").split(",")]
        if not any(r in ids for r in refs):
            raise SystemExit(f"mapping row {row[0]} points to no FRRN item")
    missing = [s["id"] for s in S.SCREENS if not s["shots"]]
    if missing:
        raise SystemExit(f"screens without a screenshot: {missing}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--only", default="FRS,BPS,UXD,RNO,ICD")
    args = ap.parse_args()
    checks()
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        for key in args.only.split(","):
            d = WRITERS[key](tmp)
            target = OUT / C.DOCS[key]["file"]
            n = pc.finish(d, target, tmp, pdf=not args.no_pdf)
            print(f"wrote {target.name} ({n} TOC lines, {d.shots} screenshots)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
