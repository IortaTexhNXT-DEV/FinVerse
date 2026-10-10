"""The review workbooks of BRD-06 Renewal that go with the FRS in BDOI's template v1.1.

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd06_workbooks.py

Builds the FRS v1.1 in memory first (build_brd06_bdoi_frs.build_review), so that every number of the workbooks is the
number of the document, then writes:

  BDOI_Template_FRS/BRD-06_Renewal/
    BIBS_Inputs_BRD-06_Business_Unit_Requirements_Collection_v1.0.xlsx  what the Business Unit provides
    BIBS_FitGap_BRD-06_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx           every BRD requirement against the standard
    BIBS_CR_BRD-06_Change_Request_Register_v1.0.xlsx                    change control, candidates, decision log
    BIBS_RTM_BRD-06_Test_Cases_and_Traceability_v1.0.xlsx               traceability and test cases
    BIBS_UAT_BRD-06_Business_Unit_Walkthrough_Users_v1.0.xlsx           walkthrough users and script
  Quality/
    BIBS_Coverage_BRD-06_FRS_Coverage_in_BIBS_v1.0.xlsx                 internal: FRS coverage in the BIBS environment

Sources: brd06_workbooks.yaml and brd06_review.yaml (this folder), the BIBS FRS BRD-06 and its test cases, the
comparison of BRD-06, the programme workbooks (BDOI inputs v1.2, Feature List vs OOTB, conformance register, UAT
walkthrough users). No password is written: the password column says it is issued separately.

Self-checks: every BRD ID has at least one test case and a fit-gap row; every FR item (BDOI's and added) has at least
one test case and a coverage row; every open question and input of Annex AN is in the requirements collection
workbook; every Annex AO observation has a decision row; the numbers agree with the FRS.
"""

from __future__ import annotations

import datetime as dt
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import docx
import openpyxl
import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import brd06_review_edition as review  # noqa: E402
import build_brd06_bdoi_frs as frs  # noqa: E402
from brd11_xlsx import Column, ReviewWorkbook  # noqa: E402

REPO = frs.REPO
RN = frs.RN
OUT = REPO / "docs" / "deliverables" / "out"
WB = yaml.safe_load((HERE / "brd06_workbooks.yaml").read_text(encoding="utf-8"))
REV = review.REV
DOC = frs.DOC
META = WB["meta"]
FOLDER = OUT / META["folder"]
QUALITY = OUT / META["quality"]
DATES = {k: (v if isinstance(v, dt.date) else dt.date.fromisoformat(str(v))) for k, v in META["dates"].items()}
STATUS = ["Open", "Received", "Confirmed"]
RESULT = ["Not run", "Pass", "Fail", "Blocked"]
FRRN_RE = frs.FRRN_RE
BRD = "BRD-06"
SUB = "BRD-06 Renewal"


def d(x) -> str:
    if isinstance(x, (dt.date, dt.datetime)):
        return x.strftime("%d-%b-%Y")
    return str(x or "")


def entry_of(item: str) -> str:
    m = re.match(r"FRRN\.(\d{2,3})", item)
    return f"FRRN.{int(m.group(1)):03d}"


# ----------------------------------------------------------------------------------------------- the data
class Data:
    def __init__(self):
        self.b = frs.build_review()
        b = self.b
        self.lines, self.trace, self.our, self.brrn = b.lines, b.trace, b.our, b.brrn
        self.mapping = b.mapping
        self.fr_to_frrn = b.fr_to_frum
        self.cases_yaml = yaml.safe_load((RN / "brd06_cases.yaml").read_text(encoding="utf-8"))
        self.personas = self.cases_yaml["personas"]
        self.items = self.frrn_items()
        self.brd_ids = list(self.brrn) + [f"BRD {r['id']}" for r in self.lines]
        self.brd_text = {**{k: v for k, v in self.brrn_texts().items()},
                         **{f"BRD {r['id']}": r["text"] for r in self.lines}}
        self.fr_of_brd = {**{k: v["frs"] for k, v in self.brrn.items()},
                          **{f"BRD {k}": v["frs"] for k, v in self.trace.items()}}
        self.bibs_cases = self.make_bibs_cases()
        self.acc_cases = self.make_acceptance_cases()
        self.steps_by_fr = self.walkthrough_steps()
        self.inputs = {r["id"]: r for r in review.inputs_tab()}
        self.conformance = self.read_conformance()
        self.cmp_platform = self.comparison_platform()

    @staticmethod
    def brrn_texts() -> dict[str, str]:
        src = docx.Document(str(REPO / DOC["meta"]["source"]))
        return {r.cells[0].text.strip(): " ".join(r.cells[1].text.split()) for r in src.tables[2].rows[1:]}

    def frrn_items(self) -> list[dict]:
        src = docx.Document(str(REPO / DOC["meta"]["source"]))
        t = src.tables[3]
        out = []
        seen: Counter = Counter()
        for row in t.rows[1:]:
            entry = row.cells[0].text.strip()
            for sid, block in frs.subitem_blocks(row.cells[3]._tc, seen).items():
                head = frs.el_text(block[0]).strip()
                title = re.sub(r"^\s*FRRN\.\d{2,3}(?:\.\d{2})*\s*", "", head).strip() or head
                add = frs.ADD["items"][sid]
                out.append({"id": sid, "entry": entry, "title": title[:90], "origin": "BDOI FRS v1.0",
                            "refs": add.get("refs") or [], "ac": add["ac"],
                            "statements": sum(1 for el in block[1:] if frs.el_text(el).strip())})
        for ns in frs.ADD["new_subitems"]:
            out.append({"id": ns["id"], "entry": ns["row"], "title": ns["title"], "origin": "Added in v1.1",
                        "refs": ns["refs"], "ac": ns["ac"], "statements": len(ns["shall"])})
        for item in frs.ADD["new_items"]:
            for si in item["subitems"]:
                out.append({"id": si["id"], "entry": item["id"], "title": si["title"], "origin": "Added in v1.1",
                            "refs": si["refs"], "ac": si["ac"], "statements": len(si["shall"])})
        s = self.b.stats
        if len(out) != s["bdoi_items"] + s["new_subitems"] + s["new_item_subitems"]:
            raise SystemExit("FR item count differs from the FRS")
        return out

    def persona_user(self, code: str) -> tuple[str, str]:
        p = self.personas.get(code, {"name": code, "user": "-"})
        return frs.clean(p["name"]), p.get("user", "-")

    def make_bibs_cases(self) -> list[dict]:
        scen = {s["id"]: s for s in self.cases_yaml["scenarios"]}
        data_names = {x["id"]: x["name"] for x in self.cases_yaml["data"]}
        screens = self.cases_yaml["screens"]
        brd_of = defaultdict(list)
        for rid, frs_ in self.fr_of_brd.items():
            for fr in frs_:
                brd_of[fr].append(rid)
        lmap = frs.line_items_map()
        out = []
        for fr_id, e in self.cases_yaml["frs"].items():
            n = fr_id.split("-")[-1]
            seq: Counter = Counter()
            prio = self.our[fr_id].get("priority", "Must have")
            for c in e.get("cases") or []:
                ci = c["c"]
                seq[ci] += 1
                persona, user = self.persona_user(c.get("persona", e["persona"]))
                pre = c.get("pre") or scen.get(e.get("scenario"), {}).get("pre", "")
                exp = c["expected"]
                if c.get("msg"):
                    exp += " Message: " + "; ".join(re.sub(r"[{}]", "", m) for m in c["msg"])
                ids = brd_of.get(fr_id, [])
                out.append({
                    "id": f"TC-RN-{n}.{ci}-{seq[ci]:02d}", "title": frs.clean(c["title"]), "fr": fr_id,
                    "frrn": ", ".join(lmap.get(fr_id, self.fr_to_frrn.get(fr_id, []))),
                    "brd": (", ".join(ids[:6]) + (f" and {len(ids) - 6} more" if len(ids) > 6 else "")) or "-",
                    "persona": persona, "user": user,
                    "screen": frs.clean(screens.get(c.get("screen", e.get("screen")), c.get("screen", e.get("screen")))
                                        ).replace(">", "›"),
                    "pre": frs.clean(pre), "steps": "\n".join(f"{i}. {frs.clean(s)}" for i, s in enumerate(c["steps"], 1)),
                    "data": "; ".join(f"{x} {data_names.get(x, '')}" for x in (c.get("data") or e.get("data") or [])),
                    "expected": frs.clean(exp),
                    "type": "Negative" if c["type"] == "Negative" else "Positive",
                    "priority": "High" if str(prio).startswith("Must") else "Medium",
                    "source": "BIBS test plan BRD-06 v2.0"})
        return out

    def make_acceptance_cases(self) -> list[dict]:
        out = []
        for it in self.items:
            code, screen = WB["entry_persona"][entry_of(it["id"])]
            persona, user = self.persona_user(code)
            for n, ac in enumerate(it["ac"], 1):
                neg = re.search(r"\b(refused|cannot|not offered|is not|never|denied|no user|without)\b", ac, re.I)
                cid = it["id"].replace("FRRN.", "FRRN-").replace(" (2nd)", "-2ND").replace(".", "-")
                out.append({
                    "id": f"TC-{cid}-AC{n:02d}", "title": f"Acceptance criterion {n} of {it['id']} {it['title']}"[:120],
                    "fr": ", ".join(it["refs"]) or "-", "frrn": it["id"], "brd": self.brd_of_item(it["id"]),
                    "persona": persona, "user": user, "screen": screen,
                    "pre": "Users, accounts and settings of the SIT/UAT data (TD-RN-01 and the data of the FR).",
                    "steps": f"1. Sign in as {user} ({persona}).\n2. Open {screen}.\n3. Bring about the situation of "
                             f"the acceptance criterion and observe the result.",
                    "data": "TD-RN-01 Test users per Renewal role", "expected": ac,
                    "type": "Negative" if neg else "Positive", "priority": "High",
                    "source": "Acceptance criterion of the FRS v1.1"})
        if len(out) != self.b.stats["ac"]:
            raise SystemExit("acceptance criteria count differs from the FRS")
        return out

    def brd_of_item(self, item: str) -> str:
        ids = [r for r, m in self.mapping.items() if item in FRRN_RE.findall(m)]
        return (", ".join(ids[:6]) + (f" and {len(ids) - 6} more" if len(ids) > 6 else "")) or "-"

    def cases_of_item(self, item: str) -> list[str]:
        ids = [c["id"] for c in self.acc_cases if c["frrn"] == item]
        ids += [c["id"] for c in self.bibs_cases if item in c["frrn"].split(", ")]
        return ids

    def walkthrough_steps(self) -> dict[str, list[str]]:
        scr_frs = {}
        for f in sorted((RN / "pack" / "screens").glob("*.yaml")):
            for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
                scr_frs[s["id"]] = s.get("frs", [])
        out = defaultdict(list)
        for w in yaml.safe_load((RN / "pack" / "walkthroughs.yaml").read_text(encoding="utf-8"))["walkthroughs"]:
            for i, s in enumerate(w["steps"], 1):
                for fr in scr_frs.get(s[1], []):
                    out[fr].append(f"{w['id']}.{i}")
        return out

    def read_conformance(self) -> dict[str, list[tuple]]:
        p = OUT / "Programme" / "Quality" / "BIBS_Conformance_BRD-00_End-to-End_Conformance_Register_v1.0.xlsx"
        wb = openpyxl.load_workbook(p, read_only=True)
        out = defaultdict(list)
        for r in wb["FR conformance"].iter_rows(values_only=True):
            if r[0] == BRD:
                out[r[1]].append((r[4], r[7]))
        wb.close()
        return out

    @staticmethod
    def comparison_platform() -> dict[str, tuple[str, str]]:
        cmp_data = yaml.safe_load(frs.CMP.read_text(encoding="utf-8"))
        out, seen = {}, Counter()
        for m in cmp_data["mapping"]:
            seen[m[0]] += 1
            key = m[0] + (" (2nd)" if seen[m[0]] > 1 else "")
            out[key] = (m[3], m[5])
        return out


# ----------------------------------------------------------------------------------------- 1 requirements collection
def rc_cols():
    return [
        {"key": "id", "header": "ID", "width": 12, "mandatory": True, "format": "Text", "what": "Reference of the item (kept from the BDOI inputs workbook where it exists)"},
        {"key": "item", "header": "Item", "width": 28, "mandatory": True, "what": "What the item is about"},
        {"key": "brd", "header": "BRD ID", "width": 14, "what": "BRD requirement the item serves"},
        {"key": "frs", "header": "FRS reference", "width": 18, "what": "FR item or annex of the FRS v1.1"},
        {"key": "need", "header": "What is needed", "width": 44, "mandatory": True, "what": "What the Business Unit provides or decides"},
        {"key": "fmt", "header": "Format of the answer", "width": 26, "mandatory": True, "what": "How to answer"},
        {"key": "example", "header": "Example", "width": 24, "what": "An example of an answer"},
        {"key": "by", "header": "Provided by (role)", "width": 18, "mandatory": True, "what": "Role that answers"},
        {"key": "due", "header": "Needed by", "width": 12, "mandatory": True, "format": "Date dd-Mmm-yyyy", "what": "Date of the drop plan"},
        {"key": "resp", "header": "BU response", "width": 34, "what": "The answer of the Business Unit"},
        {"key": "status", "header": "Status", "width": 12, "mandatory": True, "values": STATUS, "what": "Open until answered; Received; Confirmed by the project"},
    ]


def row(id_, item, brd, frsref, need, fmt, example, by, due) -> dict:
    return {"id": id_, "item": item, "brd": brd, "frs": frsref, "need": need, "fmt": fmt, "example": example,
            "by": by, "due": d(due), "resp": "", "status": "Open"}


def input_row(data: Data, ref, item, brd, frsref, fmt, example, default_due, need="", by="") -> dict:
    src = data.inputs.get(ref) or {}
    text = re.sub(r"\{\{(\w+)\}\}", lambda m: "[" + re.sub(r"(?<=[a-z])(?=[A-Z])", " ", m.group(1)).lower() + "]",
                  frs.clean(src.get("need") or need))
    return row(ref, item, brd, frsref, text, fmt, example,
               src.get("owner") or by or "Product Owner, Renewal", src.get("due") or default_due)


def build_requirements(data: Data):
    wb = ReviewWorkbook("Business Unit Requirements Collection", doc_type="Requirements collection", brd=BRD,
                        version=META["version"], subtitle="BRD-06 Renewal - what the Business Unit provides for the "
                                                          "FRS v1.1")
    wb.cover_notes = ["Fill in the column BU response and set the Status of each row; questions go to the sheet "
                      "Questions and comments. Items keep the IDs of the BDOI Requirements and Inputs by BRD v1.2 "
                      "workbook where they exist."]
    sheets: "OrderedDict[str, list[dict]]" = OrderedDict()
    owner = {}
    # documents and print formats
    rows = []
    docs = yaml.safe_load((RN / "pack" / "documents.yaml").read_text(encoding="utf-8"))["documents"]
    for doc in docs:
        e = REV["documents"]["items"][doc["id"]]
        rows.append(row(e["rc"], frs.clean(doc["name"]), "BRRN.010" if doc["id"] == "DO-01" else "BRD 1.003 / 3.009",
                        f"Annex AF ({doc['id']})", f"Confirm the layout of the {frs.clean(doc['name'])} shown in "
                        f"Annex AF, or give the layout. BDOI reference: {e['bdoi']}.",
                        "Confirm, or a sample (Excel or PDF) with the fields in order", "Confirmed as in Annex AF",
                        "Product Owner, Renewal", DATES["decision"]))
    for doc in REV["documents"]["extra"]:
        given = "given" in doc["bdoi"]
        rows.append(row(doc["rc"], doc["name"], "BRRN.040" if "lacement" in doc["name"] else "BRRN.010",
                        f"Annex AF ({doc['id']})",
                        ("Confirm BDOI's template of this FRS for: " if given else "Give the layout and wording of: ")
                        + doc["purpose"] + f". BDOI reference: {doc['bdoi']}.",
                        "Confirm" if given else "Template (Word or PDF) with the fields in order",
                        "Confirmed" if given else "Template attached", "Product Owner, Renewal", DATES["decision"]))
    for k, r in enumerate(REV["reports"]["rows"], 1):
        m = re.search(r"RC-RPT-\d\d", r[-1])
        if m:
            rows.append(row(m.group(0), r[0], "BRRN.019", "FRRN.044.01; Annex AG",
                            f"Give the columns, filters, sort and output of the report {r[0]}.",
                            "Columns in order; filters; Excel or PDF", "Account, status, premium; by month",
                            "Marketing and Processing Heads", DATES["decision"]))
    rows.append(input_row(data, "DOC-08", "Documents of Renewal", "BRRN.010", "Annex AF", "Templates", "-",
                          DATES["uat"]))
    rows.append(input_row(data, "CLR-RN-25", "Insurer extract (28 columns)", "BRD 3.009", "FRRN.025; Annex AF",
                          "Confirm, or the columns to keep on screen", "Confirmed", DATES["decision"]))
    sheets["Documents and print formats"] = rows
    owner["Documents and print formats"] = "Product Owner, Renewal"
    # e-mails
    rows = []
    for i in REV["emails"]["items"]:
        body = " ".join(str(i["body"]).replace("\\n", " ").split())
        rows.append(row(i["id"], i["name"], i["brd"], f"Annex AE; {i['frrn']}",
                        f"Approve the wording, the trigger and the recipients. Subject: {i['subject']}. Text: {body}",
                        "Approve, or the new wording", "Approved", "Product Owner, Renewal", DATES["decision"]))
    rows.append(row("RC-EML-01", "Sender and footer of the e-mails", "BRRN.010", "Annex AE",
                    "The sender address of each environment and the standard footer or disclaimer of BDOI e-mails.",
                    "Address per environment; footer text", "renewals@ (domain of BDOI)", "BDOI IT", DATES["uat"]))
    sheets["E-mail wordings"] = rows
    owner["E-mail wordings"] = "Product Owner, Renewal; Compliance"
    # schedules
    rows = []
    for k, s in enumerate(REV["reports"]["schedules"], 1):
        rows.append(row(f"RC-SCH-{k:02d}", s[0], s[5].split(";")[0], f"Annex AG; {s[5].split(';')[-1].strip()}",
                        f"Confirm when it runs ({s[1]}) and who is told: {s[3]}.", "Time and recipients, or Confirm",
                        s[1], "Product Owner, Renewal", DATES["uat"]))
    sheets["Schedules and timings"] = rows
    owner["Schedules and timings"] = "Product Owner, Renewal; BDOI IT"
    # parameters
    rows = []
    for k, p in enumerate(DOC["annex_ac"]["parameters"], 1):
        rows.append(row(f"RC-PAR-{k:02d}", p[0], p[2], "Annex AC.7", f"Confirm the value of {p[0].lower()}. As "
                        f"delivered: {p[1]}.", "Value", str(p[1]), p[3], DATES["uat"]))
    for ref, item in [("CLR-RN-27", "Exception ageing days and recipients"),
                      ("CLR-RN-33", "NRNS waiting days per segment")]:
        rows.append(input_row(data, ref, item, "BRRN.029 / 037", "Annex AC.7", "Days; roles", "5 working days",
                              DATES["decision"]))
    sheets["Parameters and thresholds"] = rows
    owner["Parameters and thresholds"] = "Product Owner, Renewal"
    # lists
    rows = []
    for k, lst in enumerate(DOC["annex_ac"]["lists"], 1):
        rows.append(row(f"RC-LOV-{k:02d}", lst[0], "BRD 5.003", "Annex AC.8; template R01",
                        f"Give or confirm the values of the list {lst[0]} (as delivered: {lst[1]}). Template: tab "
                        f"R01 Lists of values and MIS values of the Master Data and Configuration Upload Templates v1.2.",
                        "Values in template R01", "Unit Sold", "Business Administrator", DATES["sit"]))
    sheets["Lists of values"] = rows
    owner["Lists of values"] = "Business Administrator"
    # master and reference data
    rows = []
    for k, r in enumerate(REV["data_setup"]["rows"], 1):
        rows.append(row(f"RC-DAT-{k:02d}", r[0], "BRD 5.003", f"Annex AK; {r[5]}",
                        f"Provide {r[0].lower()} from {r[1]}; status: {r[6]}.",
                        f"Template {r[2]} of the Master Data and Configuration Upload Templates v1.2 (not copied here)",
                        f"Filled tab {r[2].split(';')[0]}", r[3], DATES["sit"]))
    rows.append(input_row(data, "CFG-124", "Renewal buckets and decision matrix", "BRRN.031, 034", "FRRN.040; Annex AC",
                          "Rules as Annex C", "-", DATES["uat"]))
    sheets["Master and reference data"] = rows
    owner["Master and reference data"] = "Product Owner, Renewal; Business Administrator"
    # approvers and users
    rows = []
    for k, p in enumerate(DOC["annex_ac"]["personas"][:8], 1):
        rows.append(row(f"RC-APR-{k:02d}", f"Users of the role {p[0]}", p[2], "Annex AC.1; Annex AD",
                        f"Name the users of the role {p[0]} with their unit, and the approvers of each stage of "
                        f"Annex AD where the role approves.", "User ID, name, unit", "a013000101, Juan dela Cruz, "
                        "Retail Marketing 1", "User Access Administrator", DATES["sit"]))
    for ref, item in [("TD-06", "UAT testers"), ("SO-06-1", "Signature of the FRS"), ("SO-06-2", "Approval of the test plan"),
                      ("SO-06-3", "UAT sign-off certificate"), ("CLR-RN-40", "Signatories")]:
        rows.append(input_row(data, ref, item, "-", "Signoff Sheet; Annex AO.4", "Names, or signature", "-",
                              DATES["decision"]))
    sheets["Approvers and users per role"] = rows
    owner["Approvers and users per role"] = "Product Owner, Renewal; PMO"
    # integration inputs and open decisions
    rows_i, rows_o = [], []
    for k, r in enumerate(REV["integrations"]["rows"], 1):
        if r[7].startswith("To be provided"):
            rows_i.append(row(f"RC-INT-{k:02d}", r[0], "BRRN.010 / 029 / 035 / 040", "Annex AH",
                              f"{r[6]}: {r[7]}" if r[6] != "-" else r[7], "Decision, date and contact", "-",
                              "BDOI IT", DATES["sit"]))
    for k, q in enumerate(REV["open_items"]["open"], 1):
        first = q["refs"].split(";")[0].strip()
        rid = first if first not in ("BDOFC", "CBG-HOME") else f"RC-DEC-{k:02d}"
        r = row(rid, re.split(r"[:?]", q["q"])[0][:70], "-", f"Annex AN.3.{k} ({q['refs']})", q["q"], q["fmt"], "",
                q["owner"], q["by"])
        (rows_i if "BDOI IT" in q["owner"] else rows_o).append(r)
    for dd in REV["open_items"]["dependencies"]:
        rows_i.append(row(dd[0], "Dependency", "-", f"Annex AN.2; {dd[2]}", dd[1], "Date when it is in place",
                          d(DATES["sit"]), "BDOI IT", DATES["sit"]))
    sheets["Integration inputs (BDOI IT)"] = rows_i
    owner["Integration inputs (BDOI IT)"] = "BDOI IT"
    sheets["Open decisions"] = rows_o
    owner["Open decisions"] = "Product Owner, Renewal"
    # self-check
    present = {r["id"] for rows in sheets.values() for r in rows}
    text = " ".join(r["frs"] for rows in sheets.values() for r in rows)
    missing = [ref for ref, kind in data.b.where.items() if kind == "input" and ref not in present]
    missing += [f"AN.3.{k}" for k in range(1, len(REV["open_items"]["open"]) + 1) if f"Annex AN.3.{k} " not in text]
    if missing:
        raise SystemExit(f"requirements collection lacks {missing}")
    index = []
    for k, (name, rows) in enumerate(sheets.items(), 1):
        dues = sorted({r["due"] for r in rows if r["due"]}, key=lambda x: dt.datetime.strptime(x, "%d-%b-%Y"))
        index.append([k, name, f"{len(rows)} items", owner[name], dues[0] if dues else "", len(rows),
                      f"=COUNTIF('{name}'!K:K,\"Open\")", "Not started", name])
    index.append([len(index) + 1, "Questions and comments", "Questions of the Business Unit on any item",
                  "Any reviewer", "", 0, "", "Not started", "Questions and comments"])
    wb.index_sheet("Start here", "Start here – BRD-06 Renewal: what the Business Unit provides",
                   "Fill in the sheets in this order (decisions and lists first, then data). Each sheet is "
                   "self-contained: title block, column guide above the header, an example row to overwrite. "
                   "Rows open: the formula counts the rows still Open.",
                   ["No.", "Sheet", "Content", "Owner", "First due", "Rows", "Rows open", "Status", "Go to"],
                   index, [5, 30, 18, 30, 12, 8, 10, 14, 10], link_col=9, status_col=8,
                   status_values=["Not started", "In progress", "Completed"])
    for name, rows in sheets.items():
        ex = dict(rows[0], resp=rows[0]["example"] or "Confirmed", status="Received")
        wb.template_sheet(name, f"{name} – BRD-06 Renewal",
                          [("Purpose", f"Items of the kind '{name}' the FRS v1.1 needs from the Business Unit."),
                           ("Who fills it in", owner[name]),
                           ("How it is used", "The project applies each confirmed answer to the FRS, the configuration or "
                                              "the templates; data go through the upload templates v1.2."),
                           ("Due", "The date in the column Needed by (drop plan)."),
                           ("Depends on", "The decisions of the sheet Open decisions where an item names one.")],
                          rc_cols(), rows, ex)
    q_cols = [
        {"key": "no", "header": "No.", "width": 6, "mandatory": True, "format": "Number", "what": "Running number"},
        {"key": "ref", "header": "Item or section", "width": 20, "mandatory": True, "what": "ID of the item, or FRS section / annex"},
        {"key": "q", "header": "Question or comment", "width": 60, "mandatory": True, "what": "The question or comment"},
        {"key": "by", "header": "Raised by", "width": 20, "mandatory": True, "what": "Name and role"},
        {"key": "date", "header": "Date", "width": 12, "format": "Date dd-Mmm-yyyy", "what": "Date raised"},
        {"key": "answer", "header": "Answer of the project", "width": 50, "what": "Filled in by iorta TechNXT"},
        {"key": "status", "header": "Status", "width": 12, "values": STATUS, "what": "Open, Received, Confirmed"}]
    wb.template_sheet("Questions and comments", "Questions and comments", [
        ("Purpose", "Any question or comment of the Business Unit on the items or on the FRS v1.1."),
        ("Who fills it in", "Any reviewer; the project answers."), ("How it is used", "Answered within 3 working days."),
        ("Due", "Before the sign-off of the FRS."), ("Depends on", "-")], q_cols, [],
        {"no": 1, "ref": "AE.22", "q": "Can the Renewal Advice e-mail also go to the Account Officer?",
         "by": "Product Owner", "date": "", "answer": "", "status": "Open"})
    path = wb.save(FOLDER / "BIBS_Inputs_BRD-06_Business_Unit_Requirements_Collection_v1.0.xlsx")
    return path, {"rows": sum(len(r) for r in sheets.values()), "sheets": len(sheets)}


# ------------------------------------------------------------------------------------------------- 2 fit-gap
def env_status(data: Data, frs_: list[str]) -> str:
    res = [x[1] for f in frs_ for x in data.conformance.get(f, [])]
    if not res:
        return "Planned for"
    if all(x == "Conformant" for x in res):
        return "Available"
    return "Partly available"


def build_fitgap(data: Data):
    feat = {}
    wbf = openpyxl.load_workbook(OUT / "Programme" / "Feature_List" /
                                 "BIBS_BDOI_FeatureList_vs_OOTB_and_Best_Practice_v1.0.xlsx", read_only=True)
    for r in wbf["Feature Comparison"].iter_rows(min_row=5, values_only=True):
        if r[0] and str(r[0]).startswith("RNW"):
            feat[r[0]] = {"feature": r[3], "ootb": r[8], "assess": r[13]}
    wbf.close()
    groups = WB["fitgap_groups"]

    def group_of(rid):
        key = rid.replace("BRD ", "")
        best = None
        for g in groups:
            for k in g["keys"]:
                if key == k or key.startswith(k + "."):
                    if best is None or len(k) > len(best[1]):
                        best = (g, k)
        if best is None:
            raise SystemExit(f"no fit-gap group for {rid}")
        return best[0]
    rows = []
    for rid in data.brd_ids:
        g = group_of(rid)
        fr_ = data.fr_of_brd[rid]
        out_scope = rid.startswith("BRD ") and data.trace[rid[4:]]["out"]
        rows.append({"id": rid, "text": data.brd_text[rid], "group": g["name"],
                     "frrn": ", ".join(FRRN_RE.findall(data.mapping[rid])), "fr": ", ".join(fr_) or "-",
                     "features": "; ".join(f"{f} {feat[f]['feature']} (OOTB: {feat[f]['ootb']})" for f in g["features"]),
                     "ootb": g["ootb"], "best": g["best"],
                     "fit": "Recommend change" if out_scope else g["fit"],
                     "rec": "Out of scope by BRD Addendum 1 (p.60); replaced as the FR items state." if out_scope
                     else g["rec"], "impact": g["impact"], "env": env_status(data, fr_)})
    for n in REV["nfr"]["rows"]:
        fit, rec = WB["fitgap_nfr"].get(n[0], ["Fit", "Keep; programme standard."])
        rows.append({"id": n[0], "text": n[2], "group": "Non-functional (BRD p.212-214)", "frrn": n[4], "fr": "-",
                     "features": n[5], "ootb": n[3], "best": "Programme register of non-functional requirements",
                     "fit": fit, "rec": rec, "impact": "Annex AJ", "env": "Available" if fit == "Fit" else "Planned for"})
    wb = ReviewWorkbook("BRD vs Out-of-the-Box vs Best Practice", doc_type="Fit-gap", brd=BRD,
                        version=META["version"], subtitle="BRD-06 Renewal")
    wb.legend = [("Fit", "Met by the standard system"), ("Configure", "Met by set-up or settings"),
                 ("Extend", "Met by an extension of the standard system"),
                 ("Recommend change", "A change to the BRD text is recommended")]
    counts = Counter(r["fit"] for r in rows)
    by_group = defaultdict(Counter)
    for r in rows:
        by_group[r["group"]][r["fit"]] += 1
    summ = [{"group": "All BRD requirements", "n": len(rows), **{f: counts[f] for f in WB["fit_values"]}}]
    summ += [{"group": g, "n": sum(c.values()), **{f: c[f] for f in WB["fit_values"]}} for g, c in by_group.items()]
    wb.sheet("Summary", [Column("group", "Group", 46, "Function of the BRD"), Column("n", "Requirements", 14, "Count"),
                         Column("Fit", "Fit", 10, "Count"), Column("Configure", "Configure", 12, "Count"),
                         Column("Extend", "Extend", 10, "Count"), Column("Recommend change", "Recommend change", 18, "Count")],
             summ, description=f"Counts by fit: {len(data.brd_ids)} BRD requirement IDs and "
                               f"{len(REV['nfr']['rows'])} non-functional rows")
    wb.sheet("Fit-Gap", [
        Column("id", "BRD ID", 18, "BRD requirement ID (RN-NFR-nn for the usage requirements)"),
        Column("text", "BRD text", 44, "Requirement as printed in the BRD"),
        Column("group", "Function", 24, "Function of the BRD"),
        Column("frrn", "FR item (BDOI format)", 22, "FR items of the FRS v1.1"),
        Column("fr", "BIBS reference FR", 16, "FR of the BIBS FRS BRD-06"),
        Column("features", "Feature list reference", 40, "Rows of the BIBS Feature List vs OOTB workbook"),
        Column("ootb", "What BIBS offers out of the box", 46, "Standard capability or configuration"),
        Column("best", "Insurance-broking best practice", 46, "Practice of the Philippine market and international broking"),
        Column("fit", "Fit", 16, "Fit class", values=WB["fit_values"], status=True),
        Column("rec", "Recommendation", 40, "What to do"),
        Column("impact", "Impact", 22, "Where it lands"),
        Column("env", "In the BIBS environment", 20, "Available / Partly available / Planned for")],
        rows, description="One row per BRD requirement")
    path = wb.save(FOLDER / "BIBS_FitGap_BRD-06_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx")
    if not set(data.brd_ids) <= {r["id"] for r in rows}:
        raise SystemExit("fit-gap lacks BRD IDs")
    return path, {"rows": len(rows), **counts}


# ------------------------------------------------------------------------------------------------- 3 change requests
def build_cr(data: Data):
    wb = ReviewWorkbook("Change Request Register", doc_type="Change control", brd=BRD, version=META["version"],
                        subtitle="BRD-06 Renewal - after the sign-off of the FRS v1.1")
    steps = [{"step": s[0], "what": s[1], "who": s[2], "record": s[3]} for s in REV["change_control"]["steps"]]
    wb.sheet("How to use", [Column("step", "Step", 12, "Step of the process"), Column("what", "What happens", 60, ""),
                            Column("who", "Who", 36, ""), Column("record", "Record in this workbook", 36, "")],
             steps, description="Raise -> assess -> approve -> schedule -> test -> close (FRS Annex AP)")
    rows = []
    for k, c in enumerate(WB["cr_candidates"], 1):
        rows.append({"no": f"CR-RN-{k:03d}", "date": "09-Oct-2026", "by": "iorta TechNXT Project Team", "brd": c[1],
                     "frs": c[0], "desc": f"{c[2]}: {c[3]}", "reason": c[4], "type": c[5], "scope": c[6],
                     "schedule": c[7], "cost": c[8], "other": c[9], "tests": c[10], "priority": c[11],
                     "status": "Candidate - for decision", "decision": "", "approved": "", "ddate": "", "drop": "",
                     "linked": ", ".join(data.cases_of_item(c[0])[:5])})
    wb.sheet("Register", [
        Column("no", "CR No.", 12, "Change request number"), Column("date", "Date raised", 12, "dd-Mmm-yyyy"),
        Column("by", "Raised by", 18, "Name and role"), Column("brd", "BRD ID", 16, "BRD requirement"),
        Column("frs", "FRS reference", 16, "FR item or annex"), Column("desc", "Description", 46, "The change"),
        Column("reason", "Reason", 36, "Why"), Column("type", "Type", 18, "Kind of change", values=WB["cr_types"]),
        Column("scope", "Impact on scope", 22, ""), Column("schedule", "Impact on schedule", 14, ""),
        Column("cost", "Impact on cost", 12, ""), Column("other", "Impact on other BRDs", 18, ""),
        Column("tests", "Impact on tests", 18, ""), Column("priority", "Priority", 10, "", values=["High", "Medium", "Low"]),
        Column("status", "Status", 20, "", values=WB["cr_status"], status=True),
        Column("decision", "Decision", 24, "Decision of the approvers"), Column("approved", "Approved by", 18, ""),
        Column("ddate", "Decision date", 12, ""), Column("drop", "Target drop", 12, ""),
        Column("linked", "Linked test cases", 30, "Test cases of the Test Cases and Traceability workbook")],
        rows, description="Candidate changes known at v1.1 (status Candidate - for decision) and new requests")
    form_cols = [
        {"key": "field", "header": "Field", "width": 30, "mandatory": True, "what": "Part of the assessment"},
        {"key": "value", "header": "Assessment", "width": 70, "mandatory": True, "what": "Filled in by the project"},
        {"key": "by", "header": "Assessed by", "width": 20, "what": "Name"}]
    fields = ["CR No.", "Summary of the change", "BRD IDs and FR items affected", "Screens, messages, e-mails affected",
              "Other BRDs affected (Annex AH)", "Effort (mandays)", "Schedule impact (drop)", "Cost",
              "Test cases to add or change", "Training and documents", "Risks", "Recommendation"]
    wb.template_sheet("Impact assessment form", "Impact assessment form", [
        ("Purpose", "The assessment of one change request before the decision."),
        ("Who fills it in", "iorta TechNXT Project Team with the System Analyst (BDO ITG)"),
        ("How it is used", "Copied for each change; its conclusion goes to the Register."),
        ("Due", "Within 5 working days of the request."), ("Depends on", "Register")],
        form_cols, [{"field": f, "value": "", "by": ""} for f in fields],
        {"field": "CR No.", "value": "CR-RN-001", "by": "Project analyst"})
    dec = [{"no": q["no"], "section": q["section"], "topic": q["topic"], "proposal": q["proposal"],
            "by": q["decide_by"], "decision": "", "dby": "", "ddate": "", "cr": "", "status": "Open"}
           for q in data.b.q_rows]
    if len(dec) != data.b.stats["observations"]:
        raise SystemExit("decision log differs from Annex AO")
    wb.sheet("Decisions log", [
        Column("no", "Annex AO No.", 14, "Row of Annex AO of the FRS v1.1"), Column("section", "Section", 8, ""),
        Column("topic", "Topic", 26, ""), Column("proposal", "Proposed in the FRS", 60, ""),
        Column("by", "Decision by", 26, ""), Column("decision", "Decision", 30, "Adopt proposal / Keep BDOI text / Other"),
        Column("dby", "Decided by (name)", 18, ""), Column("ddate", "Date", 12, ""),
        Column("cr", "CR No. (if a change follows)", 14, ""),
        Column("status", "Status", 12, "", values=["Open", "Decided", "Withdrawn"], status=True)],
        dec, description="One row per observation of Annex AO, for the decision of the Business Unit")
    path = wb.save(FOLDER / "BIBS_CR_BRD-06_Change_Request_Register_v1.0.xlsx")
    return path, {"candidates": len(rows), "decisions": len(dec)}


# ------------------------------------------------------------------------------------------------- 4 RTM
def frs_table_text(ref: str) -> str:
    text = (RN / "FRS_BRD06_RENEWAL.md").read_text(encoding="utf-8")
    m = re.search(r"^\| " + re.escape(ref) + r" \| ([^|]+)\|", text, re.M)
    return frs.clean(m.group(1).strip()) if m else ""


def answered_rows(data: Data) -> list[dict]:
    out = []
    for ref, (clause, answer) in REV["open_items"]["answered"].items():
        q = (data.inputs.get(ref) or {}).get("need") or frs_table_text(ref)
        if not q and ref.startswith("CLR-RN"):
            text = (RN / "FRS_BRD06_RENEWAL.md").read_text(encoding="utf-8")
            m = re.search(r"^\| " + re.escape(ref) + r" \| ([^|]+)\| ([^|]+)\|", text, re.M)
            q = f"{m.group(1).strip()}: {m.group(2).strip()}" if m else ""
        out.append({"ref": ref, "q": frs.clean(q) if q else "-", "clause": clause, "answer": answer,
                    "kept": "Annex AO decision" if re.search(r"Annex AO C\d\d", answer) else "-"})
    return out


def build_rtm(data: Data):
    cases = data.acc_cases + data.bibs_cases
    case_ids_by_fr = defaultdict(list)
    for c in data.bibs_cases:
        case_ids_by_fr[c["fr"]].append(c["id"])
    acc_by_item = defaultdict(list)
    for c in data.acc_cases:
        acc_by_item[c["frrn"]].append(c["id"])
    trace_rows = []
    for rid in data.brd_ids:
        frrn = FRRN_RE.findall(data.mapping[rid])
        fr_ = data.fr_of_brd[rid]
        ids = sorted({i for f in fr_ for i in case_ids_by_fr[f]} | {i for it in frrn for i in acc_by_item[it]})
        screens = "; ".join(dict.fromkeys(frs.clean(data.our[f].get("screens", "")) for f in fr_))
        steps = sorted({s for f in fr_ for s in data.steps_by_fr.get(f, [])},
                       key=lambda x: (x.split(".")[0], int(x.split(".")[1])))
        trace_rows.append({"id": rid, "text": data.brd_text[rid], "frrn": ", ".join(frrn), "fr": ", ".join(fr_) or "-",
                           "n": len(ids), "cases": ", ".join(ids[:15]) + (" ..." if len(ids) > 15 else ""),
                           "screens": screens or "-", "steps": ", ".join(steps) or "-"})
    no_case = [r["id"] for r in trace_rows if r["n"] == 0]
    if no_case:
        raise SystemExit(f"BRD IDs without a test case: {no_case[:10]}")
    item_rows = []
    for it in data.items:
        ids = data.cases_of_item(it["id"])
        if not ids:
            raise SystemExit(f"{it['id']} has no test case")
        item_rows.append({"id": it["id"], "title": it["title"], "origin": it["origin"], "ac": len(it["ac"]),
                          "n": len(ids), "cases": ", ".join(ids[:12]) + (" ..." if len(ids) > 12 else "")})
    wb = ReviewWorkbook("Test Cases and Traceability", doc_type="Requirements traceability", brd=BRD,
                        version=META["version"], subtitle="BRD-06 Renewal - FRS in BDOI's template v1.1")
    n_brrn = len(data.brrn)
    summary = [
        {"what": "BRD requirement IDs (42 BRRN and the lines of the main BRD)", "n": len(data.brd_ids),
         "proof": f"{sum(1 for r in trace_rows if r['n'])} with at least one test case (sheet Traceability); "
                  f"{n_brrn} BRRN, {len(data.lines)} lines"},
        {"what": "FR items (BDOI's and added)", "n": len(data.items),
         "proof": f"{sum(1 for r in item_rows if r['n'])} with at least one test case (sheet FR items)"},
        {"what": "Acceptance criteria", "n": data.b.stats["ac"], "proof": "One acceptance test case each"},
        {"what": "Test cases", "n": len(cases),
         "proof": f"{len(data.bibs_cases)} of the BIBS test plan BRD-06 re-keyed to the FR items; "
                  f"{len(data.acc_cases)} acceptance cases"},
        {"what": "Positive / negative", "n": f"{sum(1 for c in cases if c['type'] == 'Positive')} / "
                                             f"{sum(1 for c in cases if c['type'] == 'Negative')}", "proof": "-"},
        {"what": "Questions answered in the BDOI FRS", "n": len(REV["open_items"]["answered"]),
         "proof": "Sheet Answered in the BDOI FRS"},
    ]
    wb.sheet("Summary", [Column("what", "Item", 40, ""), Column("n", "Count", 14, ""), Column("proof", "Proof", 80, "")],
             summary, description="Counts and the proof that every BRD ID and every FR item is covered")
    wb.sheet("Traceability", [
        Column("id", "BRD ID", 18, "BRD requirement ID"), Column("text", "BRD requirement", 44, ""),
        Column("frrn", "FR item (BDOI format)", 26, "FR items of the FRS v1.1"),
        Column("fr", "BIBS reference FR", 16, "FR of the BIBS FRS BRD-06"), Column("n", "Test cases", 10, "Count"),
        Column("cases", "Test case IDs", 60, "Sheet Test cases"), Column("screens", "Screen", 30, "Screen of the FR"),
        Column("steps", "Walkthrough steps", 26, "Steps of the walkthrough users workbook")],
        trace_rows, description="BRD ID -> FR item -> BIBS reference FR -> test cases -> screen -> walkthrough step")
    wb.sheet("FR items", [Column("id", "FR item", 16, ""), Column("title", "Title", 44, ""), Column("origin", "Origin", 16, ""),
                          Column("ac", "Acceptance criteria", 12, ""), Column("n", "Test cases", 10, ""),
                          Column("cases", "Test case IDs", 70, "")],
             item_rows, description="Every FR item of the FRS v1.1 with its test cases")
    case_cols = [
        Column("id", "Test case ID", 26, "TC-RN-nnn.c-ss: BIBS test plan; TC-FRRN-...-ACnn: acceptance criterion"),
        Column("title", "Title", 36, ""), Column("brd", "BRD ID", 22, ""), Column("frrn", "FR item", 18, ""),
        Column("fr", "BIBS reference FR", 14, ""), Column("persona", "Persona", 18, ""),
        Column("user", "Sign-in ID", 12, "The password is issued separately by the iorta TechNXT Project Team"),
        Column("screen", "Screen", 22, ""), Column("pre", "Preconditions", 30, ""), Column("steps", "Steps", 50, ""),
        Column("data", "Test data", 22, ""), Column("expected", "Expected result", 50, ""),
        Column("type", "Type", 10, "", values=["Positive", "Negative"]),
        Column("priority", "Priority", 10, "", values=["High", "Medium", "Low"]),
        Column("result", "Result", 10, "", values=RESULT, status=True), Column("tester", "Tested by", 14, ""),
        Column("date", "Date", 12, ""), Column("obs", "Observation ref", 14, ""), Column("remarks", "Remarks", 24, ""),
        Column("source", "Source", 22, "")]
    for c in cases:
        c.update({"result": "Not run", "tester": "", "date": "", "obs": "", "remarks": ""})
    wb.sheet("Test cases", case_cols, cases, description="Test cases of the FRS v1.1 with the execution columns")
    ans = answered_rows(data)
    wb.sheet("Answered in the BDOI FRS", [
        Column("ref", "Our reference", 14, "RQ, CLR-RN, A-RN, D-RN, DMQ"), Column("q", "Our question or assumption", 60, ""),
        Column("clause", "BDOI FRS clause", 30, ""), Column("answer", "Answer in BDOI's FRS", 60, ""),
        Column("kept", "Still a decision?", 16, "Where BDOI's answer differs from the BRD, the decision stays in Annex AO")],
        ans, description="Open items dropped from the FRS because BDOI's FRS answers them")
    crit = [{"kind": "Entry", "no": i, "criterion": c} for i, c in enumerate(WB["uat_criteria"]["entry"], 1)]
    crit += [{"kind": "Exit", "no": i, "criterion": c} for i, c in enumerate(WB["uat_criteria"]["exit"], 1)]
    wb.sheet("UAT entry and exit", [Column("kind", "Entry / exit", 12, ""), Column("no", "No.", 6, ""),
                                    Column("criterion", "Criterion", 100, "")],
             crit, description="UAT entry and exit criteria of BRD-06")
    path = wb.save(FOLDER / "BIBS_RTM_BRD-06_Test_Cases_and_Traceability_v1.0.xlsx")
    return path, {"trace": len(trace_rows), "cases": len(cases), "answered": len(ans), "items": len(item_rows)}


# ------------------------------------------------------------------------------------------------- 5 UAT users
PERSONA_MENU = {"Marketing Team Leader": 1, "Marketing AO / Admin, Account Broker": 2, "Processing Team Leader": 3,
                "Processing Officer / Broker": 4, "Business Administrator": 5, "LAMD user": 6,
                "Contact Center user": 7}


def build_uat(data: Data):
    src = OUT / "Programme" / "UAT" / "BIBS_UAT_BRD-00_Walkthrough_Users_and_Sign-in_v1.0.xlsx"
    w = openpyxl.load_workbook(src, read_only=True)
    menus = next(f for f in review.FLOWS if f["id"] == "brd06_fig_menus")["rows"]
    users = []
    head = None
    for r in w[SUB].iter_rows(values_only=True):
        if r and r[0] == "Persona":
            head = list(r)
            continue
        if head and r and r[0]:
            rowd = dict(zip(head, r))
            col = PERSONA_MENU.get(rowd["Persona"])
            if col:
                menu = "; ".join("Client & Policy › Renewal › " + m[0] for m in menus if m[col])
            elif rowd["Persona"] == "System Administrator":
                menu = "Setup & Administration › Administration › Roles & Permissions (no Renewal entry)"
            else:
                menu = "Client & Policy › Renewal › Renewal Home, Expiry List (read only)"
            users.append({"persona": rowd["Persona"], "name": rowd["User name"], "signin": rowd["Sign-in ID"],
                          "profiles": rowd["Group profile"], "unit": rowd["Company / branch / unit"],
                          "does": rowd["What this user does in the walkthrough"], "start": rowd["Menu path to start"],
                          "menus": menu, "steps": rowd["Walkthrough steps"], "password": META["password_text"]})
    script = []
    for r in w["Walkthrough Script"].iter_rows(values_only=True):
        if r and r[0] == BRD:
            script.append({"wt": r[1], "step": r[2], "persona": r[3], "signin": r[4], "action": r[5],
                           "expected": r[6], "result": "Not run", "tester": "", "date": "", "remarks": ""})
    w.close()
    if not users or not script:
        raise SystemExit("walkthrough users or script not found")
    wb = ReviewWorkbook("Business Unit Walkthrough Users", doc_type="UAT walkthrough", brd=BRD,
                        version=META["version"], subtitle="BRD-06 Renewal - sign-in IDs for the walkthrough")
    wb.cover_notes = [f"Environment: {META['environment']}. Passwords: {META['password_text']}; no password is "
                      f"written in this workbook."]
    wb.sheet("Users", [
        Column("persona", "Persona", 24, "Business role of the FRS"), Column("name", "User name", 22, "As shown on the screens"),
        Column("signin", "Sign-in ID", 14, "User ID to type on the Login page"),
        Column("profiles", "Group profiles", 26, "Role held"), Column("unit", "Company / branch / unit", 30, ""),
        Column("does", "What the user does in the walkthrough", 40, ""), Column("start", "Menu path to start", 30, ""),
        Column("menus", "Menus the user should see", 50, "Renewal entries of the user's menu (FRS Annex AA)"),
        Column("steps", "Walkthrough steps", 24, "Steps of the sheet Walkthrough script"),
        Column("password", "Password", 26, "Never written here")],
        users, description=f"{META['environment']}. Password: {META['password_text']}")
    wb.sheet("Walkthrough script", [
        Column("wt", "Walkthrough", 30, ""), Column("step", "Step", 8, ""), Column("persona", "Persona", 22, ""),
        Column("signin", "Sign-in ID", 14, ""), Column("action", "Action", 56, ""),
        Column("expected", "Expected result", 56, ""), Column("result", "Result", 10, "", values=RESULT, status=True),
        Column("tester", "Done by", 14, ""), Column("date", "Date", 12, ""), Column("remarks", "Remarks", 24, "")],
        script, description="The walkthrough of BRD-06 for the Business Unit (FRS Annex AL)")
    path = wb.save(FOLDER / "BIBS_UAT_BRD-06_Business_Unit_Walkthrough_Users_v1.0.xlsx")
    return path, {"users": len(users), "steps": len(script)}


# ------------------------------------------------------------------------------------------------- 6 coverage
def available_evidence(item_id: str) -> str:
    """The evidence of an item available in full: the item itself, else the longest prefix it starts with."""
    avail = WB.get("coverage_available", {})
    if item_id in avail:
        return avail[item_id]
    base = re.sub(r"\s*\(.*\)$", "", item_id).strip()
    hits = [k for k in avail if "(" not in k and (base == k or base.startswith(k + "."))]
    return avail[max(hits, key=len)] if hits else ""


def build_coverage(data: Data):
    item_rows, ac_rows = [], []
    for it in data.items:
        refs = it["refs"]
        conf = [x for f in refs for x in data.conformance.get(f, [])]
        evidence = []
        if conf:
            evidence.append("Conformance register: " + "; ".join(
                f"{f} {dict(Counter(x[1] for x in data.conformance.get(f, [])))}" for f in refs))
        steps = sorted({s for f in refs for s in data.steps_by_fr.get(f, [])})
        if steps:
            evidence.append("Walkthrough " + ", ".join(steps[:8]))
        avail = available_evidence(it["id"])
        if avail:
            status, note = "Yes", avail
        elif it["id"] in data.cmp_platform:
            plat, cnote = data.cmp_platform[it["id"]]
            status = {"Yes": "Yes", "Partly": "Partly", "No": "No"}.get(plat, "Partly")
            note = "Comparison of BDOI's FRS with the system: " + frs.clean(cnote)
        elif conf and all(x[1] == "Conformant" for x in conf):
            status, note = "Yes", "Available; shown in the end-to-end run of the conformance register."
        elif conf:
            status, note = "Partly", "Partly available: some acceptance criteria are partly conformant in the end-to-end run."
        else:
            status, note = "Partly", "Not reached by the end-to-end run; confirmed in UAT."
        cases = data.cases_of_item(it["id"])
        item_rows.append({"group": entry_of(it["id"]), "id": it["id"], "title": it["title"], "origin": it["origin"],
                          "refs": ", ".join(refs) or "-", "status": status,
                          "evidence": f"{len(cases)} test cases; " + " | ".join(evidence), "note": note})
        for n, ac in enumerate(it["ac"], 1):
            ac_rows.append({"group": entry_of(it["id"]), "id": it["id"], "no": n, "ac": ac, "status": status,
                            "evidence": next(c["id"] for c in data.acc_cases
                                             if c["frrn"] == it["id"] and c["id"].endswith(f"AC{n:02d}")),
                            "note": {"Yes": "Available; confirmed with the acceptance test case in UAT.",
                                     "Partly": "Partly available: " + note,
                                     "No": "Planned for a later release: " + note}[status]})

    def pct(rows):
        c = Counter(r["status"] for r in rows)
        n = len(rows)
        return c, (round(100 * c["Yes"] / n, 1) if n else 0), (round(100 * (c["Yes"] + c["Partly"]) / n, 1) if n else 0)
    summ = []
    groups = list(dict.fromkeys(r["group"] for r in item_rows))
    for g in ["All"] + groups:
        ir = [r for r in item_rows if g == "All" or r["group"] == g]
        ar = [r for r in ac_rows if g == "All" or r["group"] == g]
        ci, yi, pi = pct(ir)
        ca, ya, pa = pct(ar)
        summ.append({"group": "All FR items" if g == "All" else g, "items": len(ir), "iy": ci["Yes"], "ip": ci["Partly"],
                     "in": ci["No"], "ipct": yi, "ippct": pi, "acs": len(ar), "ay": ca["Yes"], "ap": ca["Partly"],
                     "an": ca["No"], "apct": ya, "appct": pa})
    wb = ReviewWorkbook("FRS Coverage in BIBS", doc_type="Internal coverage", brd=BRD, version=META["version"],
                        subtitle="BRD-06 Renewal - FRS in BDOI's template v1.1 (internal)",
                        classification="Internal - iorta TechNXT")
    wb.legend = [("Yes", "Available in the BIBS environment"), ("Partly", "Partly available (reason in the note)"),
                 ("No", "Planned for a later release (reason in the note)")]
    num = lambda k, h, w=10: Column(k, h, w, "", kind="number")  # noqa: E731
    wb.sheet("Summary", [Column("group", "FR group", 22, "FR entry"), num("items", "FR items"), num("iy", "Yes"),
                         num("ip", "Partly"), num("in", "No"), num("ipct", "% items available", 14),
                         num("ippct", "% items available or partly", 16), num("acs", "Acceptance criteria", 14),
                         num("ay", "Yes"), num("ap", "Partly"), num("an", "No"), num("apct", "% criteria available", 14),
                         num("appct", "% criteria available or partly", 16)],
             summ, description="Coverage of the FRS v1.1 in the BIBS environment, by item count and by acceptance criterion")
    st = ["Yes", "Partly", "No"]
    wb.sheet("FR items", [Column("group", "FR group", 12, ""), Column("id", "FR item", 16, ""), Column("title", "Title", 40, ""),
                          Column("origin", "Origin", 14, ""), Column("refs", "BIBS reference FR", 22, ""),
                          Column("status", "Available", 12, "", values=st, status=True),
                          Column("evidence", "Evidence", 60, ""), Column("note", "Note", 60, "")],
             item_rows, description="Every FR item and sub-item of the FRS v1.1")
    wb.sheet("Acceptance criteria", [Column("group", "FR group", 12, ""), Column("id", "FR item", 16, ""),
                                     Column("no", "AC", 6, ""), Column("ac", "Acceptance criterion", 60, ""),
                                     Column("status", "Available", 12, "", values=st, status=True),
                                     Column("evidence", "Test case", 28, ""), Column("note", "Note", 50, "")],
             ac_rows, description="Every acceptance criterion of the FRS v1.1")
    path = wb.save(QUALITY / "BIBS_Coverage_BRD-06_FRS_Coverage_in_BIBS_v1.0.xlsx")
    return path, {"items": len(item_rows), "acs": len(ac_rows), "all": summ[0]}


def main() -> int:
    data = Data()
    FOLDER.mkdir(parents=True, exist_ok=True)
    for name, fn in (("requirements", build_requirements), ("fitgap", build_fitgap), ("cr", build_cr),
                     ("rtm", build_rtm), ("uat", build_uat), ("coverage", build_coverage)):
        path, stats = fn(data)
        print(f"wrote {path.relative_to(REPO)} ({path.stat().st_size // 1024} KB): {stats}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
