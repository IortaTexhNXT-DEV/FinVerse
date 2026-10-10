"""User Access Maintenance FRS in BDOI's template, version 1.3 (Business Unit review edition).

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd11_bdoi_frs.py [--no-pdf]

Starts from a copy of BDOI's own FRS (docs/source-documents/FRS - BDOI BROKERSYS User Access Maintenance v1.1
(BDOI).docx, BDO ITG template D051, FR document PRJ0012475) and keeps its template: styles, cover, headers and footers,
table formats, fonts and the table of contents (marked to update when the document is opened). BDOI's text is kept
word for word; the additions of version 1.2 come from:

  brd11_frs_additions.yaml   statements and acceptance criteria under BDOI's items, new sub-items, new items FRUM.009+
  brd11_document.yaml        cover, revision log, introduction, mapping additions, annex rows, Annexes M, N, R and the
                             observations of Annex Q.4
  ../comparisons/brd11_comparison.yaml   the conflicts and slips of the comparison workbook (Annex Q.1 and Q.2)

  brd11_review.yaml          version 1.3: summary for the Business Unit Review (1.8), Annex Q.5 and Annexes S to AB
  brd11_flows.yaml           version 1.3: the figures (process flows by persona, life-cycles, context, menus),
                             drawn by brd11_figures.py and placed by brd11_v13.py

and from the BIBS release set BRD-11 (FRS v2.1, screens, screenshots, notifications, walkthroughs, test cases) and the
User Access Maintenance BRD (docs/source-documents/User Access Maintenance.pdf), read at every build.

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-11_User_Access_Maintenance/:
  BIBS_FRS-BDOI_BRD-11_User_Access_Maintenance_v1.3.docx                      clean copy
  BIBS_FRS-BDOI_BRD-11_User_Access_Maintenance_v1.3.pdf                       the clean copy as PDF (not kept in git)
  BIBS_FRS-BDOI_BRD-11_User_Access_Maintenance_v1.3_Changes_Highlighted.docx  the additions of version 1.3 shaded
                                                                              light yellow, change summary in front

Self-checks (the build fails when one does not hold): every BRD ID of the BRD is in the Business Requirements Mapping
and maps to at least one FRUM item; every FR-UA of the BIBS FRS is in Annex P with its FRUM items; every conflict and
slip of the comparison is in Annex Q; BDOI's original FRUM IDs and texts are all still present; no internal code or
restricted word in the added text. Version 1.3 adds: every notice and e-mail of the platform for User Access
Maintenance is in Annex T and its text is found in the platform; every report is in Annex V; every open question of
our sources is classified once and none that BDOI's FRS answers is kept in Annex Z.
"""

from __future__ import annotations

import argparse
import codecs
import copy
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter, OrderedDict
from pathlib import Path

import docx
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import brd11_figures  # noqa: E402
import brd11_v13 as v13  # noqa: E402
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
SRC = REPO / "docs" / "deliverables" / "src"
UA = SRC / "BRD-11_User_Access_Maintenance"
CMP = SRC / "programme" / "comparisons" / "brd11_comparison.yaml"

ADD = yaml.safe_load((HERE / "brd11_frs_additions.yaml").read_text(encoding="utf-8"))
DOC = yaml.safe_load((HERE / "brd11_document.yaml").read_text(encoding="utf-8"))
DOC["meta"].update({k: v13.REV["meta"][k] for k in ("clean", "highlighted")})
DOC["meta"]["version_new"] = v13.REV["meta"]["version"]
DOC["revision_log"]["rows"].append(v13.REV["meta"]["revision_row"])
DOC["glossary"] = DOC["glossary"] + v13.REV["glossary_add"]
NEW_TERMS = {g[0] for g in v13.REV["glossary_add"]}
OUT = REPO / "docs" / "deliverables" / "out" / DOC["meta"]["out_folder"]

HL_FILL = "FFF59D"          # light yellow of the review copy
# Screens added or changed with the full-coverage build of October 2026 (shaded in the review copy of version 1.3)
SCREENS_V13 = {"SCR-UA-01", "SCR-UA-04", "SCR-UA-15", "SCR-UA-19", "SCR-UA-20", "SCR-UA-21", "SCR-UA-22"}
HEAD_FILL = "D9E1F2"        # header fill of the annex tables
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
PORTRAIT_W = 7.4            # usable width (inches) of the portrait pages of BDOI's main section
LANDSCAPE_W = 9.9

# ----------------------------------------------------------------------------------------------------------- wording
# Codes of the BIBS sources, in business words.
CODES = {
    "UAM_DIRECT_ROLE_EDIT": "the emergency role edit", "UAM_REQUESTOR": "Requestor", "UAM_APPROVER": "Approver",
    "UAM_SECOND_APPROVER": "Second Approver", "MKT_AO": "Marketing Account Officer", "NB_APPROVER": "New Business Approver",
    "BUSINESS_ADMIN": "Business Administrator", "INFOSEC_OFFICER": "Information Security Officer",
    "SESSION_TIMEOUT_MINUTES": "the session timeout", "PASSWORD_MIN_AGE_DAYS": "the minimum password age",
    "PASSWORD_HISTORY_COUNT": "the password history", "USER_ID_PATTERN": "the user ID format",
    "SESSION_IDLE_WARNING_MINUTES": "the inactivity warning time", "PASSWORD_MAX_AGE_DAYS": "the password expiry",
    "USER_ID_FORMAT_TEXT": "the user ID format in words", "UAM_WORKING_HOURS": "the working hours",
    "UAM_EXTERNAL_USERS": "the external users setting", "UAM_DORMANT_USERS": "the dormant-user run",
    "UAM_DORMANT_NOTICE_DAYS": "the dormancy notice", "UAM_ANY_APPROVER": "the any-approver setting",
    "SESSION_EXPIRY_WARNING_MINUTES": "the session-end warning time", "RENEWAL_ENQUIRY": "Renewal Enquiry",
    "LOGIN_MAX_FAILED_ATTEMPTS": "the lock-out attempts", "JOB_FAILURE_RECIPIENTS": "the failed-run recipients",
    "UAM_USER_LEVEL": "the user level list", "UAM_BUSINESS_UNIT": "the business unit group list",
    "UAM_ROLE_APPLY_ON_APPROVAL": "the apply-on-approval setting", "UAM_DORMANT_DAYS": "the dormancy period",
    "UAM_DEACTIVATION_REASON": "the deactivation reason list", "AUTH_MODE": "the sign-in mode",
    "PASSWORD_EXPIRY_NOTICE": "the password expiry notice", "UAM_EFFECTIVE_CHANGES": "the effective-changes run",
    "ACCESS_REQUEST": "the request permission", "ACCESS_APPROVE": "the approval permission",
    "USER_MANAGE": "the user management permission", "ROLE_MANAGE": "the role management permission",
    "AUDIT_VIEW": "the audit permission", "UAM_ENROLL": "the enrol permission", "UAM_MODIFY": "the modify permission",
    "UAM_DEACTIVATE": "the deactivate permission", "UAM_REACTIVATE": "the reactivate permission",
    "UAM_CORRECT": "the correction permission", "UAM_CANCEL": "the cancel permission", "UAM_VIEW": "the view permission",
    "UAM_GROUP_REQUEST": "the group-profile request permission", "UAM_REPORT_VIEW": "the report permission",
    "UAM_SECOND_APPROVE": "the second-approval permission", "UAM_SOD_MAINTAIN": "the rule maintenance permission",
    "UAM_SOD_AUTHORIZE": "the rule authorisation permission", "BULK_PROCESS": "the bulk upload permission",
    "SECURITY_PARAMETER_APPROVE": "the security setting approval permission",
    "SYSTEM_PARAMETER_MANAGE": "the parameter maintenance permission", "UAM_ACCESS_REQUEST": "access request",
    "PRIVILEGE_INCREASE": "privilege increase", "OUTSIDE_HOURS": "outside working hours",
    "PENDING_SECOND": "Pending Second Approval", "FOR_IMPLEMENTATION": "For Implementation",
    "UAM_PRIVILEGED_CHANGE": "Privileged access change", "UAM_SCHEDULED_APPLY_FAILED": "Scheduled change failed",
    "JOB_FAILURE": "Job failure", "CTL_AUDIT": "Audit Trail", "IDLE_TIMEOUT": "inactivity",
    "UAM_REQUEST_TO_APPROVE": "request to approve", "INFOSEC_OFFICER": "Information Security Officer", "UAM_ACCESS_CHANGED": "access changed",
}
CODE_RE = re.compile(r"\b[A-Z][A-Z0-9]*_[A-Z0-9][A-Z0-9_]*\b")
PHRASES = [
    (re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"), "a test e-mail address"),
    (re.compile(r"Seed data:\s*"), "SIT/UAT data: "), (re.compile(r"\bseed (data|batch|request|profile)\b", re.I), r"SIT/UAT \1"),
    (re.compile(r"\bThe seed\b"), "The SIT/UAT"), (re.compile(r"\bseed\b", re.I), "SIT/UAT"),
    (re.compile(r"\(([A-Z][A-Z0-9]*_[A-Z0-9_]+)(?:,? (?:or|and|/) [A-Z][A-Z0-9]*_[A-Z0-9_]+)*\)"), ""),
    (re.compile(r"\bBIBS\b"), "the system"), (re.compile(r"\bin-app\b", re.I), "in the system"),
    (re.compile(r"\bthe app\b"), "the system"), (re.compile(r"\bSYSADMIN\b"), "System Administrator"),
    (re.compile(r"\bAUDITOR\b"), "Auditor"), (re.compile(r"\bchapter (\d+)\b"), "this FRS"),
    (re.compile(r"\bthe code\b"), "the system"), (re.compile(r"\ba platform change\b"), "a change to the system"),
    (re.compile(r"\bplatform change\b"), "change to the system"),
]
# Words a client document does not contain; kept encoded (rot13) so that this file does not contain them itself.
RESTRICTED = re.compile(r"\b(" + codecs.decode(
    "qrzb|qhzzl|snxr|fnzcyr qngn|cebgbglcr|fnaqobk|yberz vcfhz|gbqb|svkzr|pynhqr|naguebcvp|pungtcg|bcranv|pbcvybg|"
    "tvguho|trzvav|ntragf?|nffvfgnagf?|raqcbvagf?|qngnonfrf?|fpurznf?|cnlybnqf?|cyngsbez punatr|ohvyg|ohvyqf?|"
    "qrfvtarq|qrsrpgf?|ohtf?|fcevagf?|wven|fghof?|nqncgref?|frnzf?|vqrzcbgrag|gbxraf?|wfba|fdy",
    "rot13") + r")\b", re.I)
RESTRICTED_CS = re.compile(r"\b(AI|API|APIs|JSON|SQL|LLM|GPT)\b")
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
             "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
RPR_AFTER_SHD = ["fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
FRUM_RE = re.compile(r"FRUM[.-]\d{3}(?:\.\d{2})*")


def clean(text) -> str:
    """Business wording of a text taken from the BIBS sources: no internal codes."""
    s = "" if text is None else str(text)
    for code, word in sorted(CODES.items(), key=lambda kv: -len(kv[0])):
        s = re.sub(r"\b" + re.escape(code) + r"\b", word, s)
    for pat, rep in PHRASES:
        s = pat.sub(rep, s)
    s = re.sub(r"\s+([,.;:])", r"\1", re.sub(r"\(\s*\)", "", s))
    return re.sub(r"\s{2,}", " ", s).strip()


def frum_key(fid: str):
    nums = [int(x) for x in re.findall(r"\d+", fid.replace(" (2nd)", ""))]
    return nums + [1 if "(2nd)" in fid else 0]


def ppr_insert(ppr, el):
    name = el.tag.split("}")[1]
    pos = PPR_ORDER.index(name)
    for child in ppr:
        cname = child.tag.split("}")[1]
        if cname in PPR_ORDER and PPR_ORDER.index(cname) > pos:
            child.addprevious(el)
            return el
    ppr.append(el)
    return el


def shd_el(fill=HL_FILL):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def rpr_shade(rpr):
    for s in rpr.findall(qn("w:shd")):
        rpr.remove(s)
    shd = shd_el()
    for child in rpr:
        if child.tag.split("}")[1] in RPR_AFTER_SHD:
            child.addprevious(shd)
            return
    rpr.append(shd)


def cell_shade(tc):
    tcpr = tc.get_or_add_tcPr()
    for s in tcpr.findall(qn("w:shd")):
        tcpr.remove(s)
    shd = shd_el()
    after = [tcpr.find(qn(t)) for t in ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                                         "w:hideMark")]
    after = [a for a in after if a is not None]
    (after[0].addprevious(shd) if after else tcpr.append(shd))


def strip_ids(el):
    for e in el.iter():
        for a in list(e.attrib):
            if a.startswith("{%s}" % W14) or a in (qn("w:rsidR"), qn("w:rsidRPr"), qn("w:rsidP"), qn("w:rsidRDefault"),
                                                   qn("w:rsidTr")):
                del e.attrib[a]
    for b in el.findall(".//" + qn("w:bookmarkStart")) + el.findall(".//" + qn("w:bookmarkEnd")):
        b.getparent().remove(b)
    return el


def el_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


# ------------------------------------------------------------------------------------------------------- the sources
def our_frs() -> "OrderedDict[str, dict]":
    text = (UA / "FRS_BRD11_USER_ACCESS_MAINTENANCE.md").read_text(encoding="utf-8")
    out = OrderedDict()
    for block in re.findall(r"^```fr\n(.*?)^```", text, re.S | re.M):
        fr = yaml.safe_load(block)
        out[fr["id"]] = fr
    return out


def our_trace() -> list[dict]:
    text = (UA / "FRS_BRD11_USER_ACCESS_MAINTENANCE.md").read_text(encoding="utf-8")
    part = text.split("<!-- TRACE:START -->", 1)[1].split("<!-- TRACE:END -->", 1)[0]
    rows = []
    for line in part.splitlines():
        c = [x.strip() for x in line.strip().strip("|").split("|")]
        if len(c) >= 6 and c[0].startswith("BRD ") and c[1].startswith("p."):
            rows.append({"id": c[0], "page": c[1], "req": c[2], "frs": re.findall(r"FR-UA-\d{3}", c[4]),
                         "tests": c[5]})
    return rows


def our_nfr() -> list[dict]:
    text = (UA / "FRS_BRD11_USER_ACCESS_MAINTENANCE.md").read_text(encoding="utf-8")
    part = text.split("<!-- NFR:START -->", 1)[1].split("<!-- NFR:END -->", 1)[0]
    rows = []
    for line in part.splitlines():
        c = [x.strip() for x in line.strip().strip("|").split("|")]
        if len(c) >= 5 and c[0].startswith("UAM-NFR-"):
            frs = ["FR-UA-" + n for n in re.findall(r"\b(\d{3})\b", c[4].replace("FR-UA-", ""))]
            rows.append({"id": c[0], "topic": c[1], "req": c[2], "approach": c[3], "frs": frs})
    return rows


def brd_lines() -> list[tuple[str, str]]:
    """Every requirement line of the BRD (ID as printed, text as in the BRD); a repeated ID is marked "(2nd)"."""
    work = Path(tempfile.mkdtemp(prefix="uambrd_"))
    try:
        txt = work / "brd.txt"
        subprocess.run(["pdftotext", "-layout", str(REPO / DOC["meta"]["brd_pdf"]), str(txt)], check=True)
        seen: Counter = Counter()
        out = []
        for line in txt.read_text(encoding="utf-8").splitlines():
            m = re.match(r"^(BRD \d[\d.]*\d)\s+(.*)$", line.rstrip())
            if not m:
                continue
            rid = m.group(1)
            parts = [p.strip() for p in re.split(r"\s{2,}", m.group(2)) if p.strip()]
            parts = [p for p in parts if p not in ("In scope", "In Scope", "Yes")]
            seen[rid] += 1
            if seen[rid] > 1:
                rid += " (2nd)"
            out.append((rid, " - ".join(parts)))
    finally:
        shutil.rmtree(work, ignore_errors=True)
    fix = {"BRD 4.002.2.9": "9. Assign Submit a User Access Group Profiles Request function to a specific user "
                            "profile/role"}
    return [(r, t or fix.get(r, "")) for r, t in out]


def test_cases() -> dict[str, tuple[list[str], int]]:
    data = yaml.safe_load((UA / "brd11_cases.yaml").read_text(encoding="utf-8"))
    out = {}
    for fr_id, entry in (data.get("frs") or {}).items():
        n = fr_id.split("-")[-1]
        conds = [f"TC-UA-{n}.{i}" for i in range(1, len(entry.get("conditions") or []) + 1)]
        out[fr_id] = (conds, len(entry.get("cases") or []))
    return out


def tc_text(fr_id, cases) -> str:
    conds, n = cases.get(fr_id, ([], 0))
    if not conds:
        return "-"
    return (f"{conds[0]} to {conds[-1].split('-')[-1]}" if len(conds) > 1 else conds[0]) + f" ({n} cases)"


def fr_to_frum_map() -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}

    def note(refs, fid):
        for r in refs:
            out.setdefault(r, [])
            if fid not in out[r]:
                out[r].append(fid)
    for fid, add in ADD["items"].items():
        note(add["refs"], fid)
    for ns in ADD["new_subitems"]:
        note(ns["refs"], ns["id"])
    for item in ADD["new_items"]:
        for si in item["subitems"]:
            note(si["refs"], si["id"])
    return {k: sorted(v, key=frum_key) for k, v in out.items()}


# ------------------------------------------------------------------------------------------------------- the builder
class Builder:
    def __init__(self, highlight: bool):
        self.highlight = highlight
        self.round13 = False        # True while the additions of version 1.3 are made (shaded in the review copy)
        self.doc = docx.Document(str(REPO / DOC["meta"]["source"]))
        self.body = self.doc.element.body
        self.tables = list(self.doc.tables)          # BDOI's tables, before any insertion
        self.numbering = self.doc.part.numbering_part.element
        self.stats: Counter = Counter()
        self.log: list[dict] = []
        self.fr_to_frum = fr_to_frum_map()
        self.added_text: list[str] = []
        self.q_rows: list[dict] = []

    @property
    def hl(self) -> bool:
        return self.highlight and self.round13

    def run(self, text, bold=False, italic=False, size=None, font="Arial", color=None, hl=None):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), font)
        rpr.append(f)
        if bold:
            rpr.append(OxmlElement("w:b"))
            rpr.append(OxmlElement("w:bCs"))
        if italic:
            rpr.append(OxmlElement("w:i"))
        if color:
            c = OxmlElement("w:color")
            c.set(qn("w:val"), color)
            rpr.append(c)
        if size:
            for tag in ("w:sz", "w:szCs"):
                s = OxmlElement(tag)
                s.set(qn("w:val"), str(int(size * 2)))
                rpr.append(s)
        if self.hl if hl is None else hl:
            rpr.append(shd_el())
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = str(text)
        r.append(t)
        self.added_text.append(str(text))
        return r

    def para(self, text="", style=None, num=None, bold=False, italic=False, size=None, font="Arial", jc=None,
             keep_next=False, page_break_before=False, no_num=False, color=None, hl=None, space_after=None):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        if style:
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), style)
            ppr.append(ps)
        if keep_next:
            ppr.append(OxmlElement("w:keepNext"))
        if page_break_before:
            ppr.append(OxmlElement("w:pageBreakBefore"))
        if num or no_num:
            n = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), str(num[1] if num else 0))
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), str(num[0] if num else 0))
            n.append(il)
            n.append(ni)
            ppr.append(n)
        if space_after is not None:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:after"), str(space_after))
            ppr.append(sp)
        if jc:
            j = OxmlElement("w:jc")
            j.set(qn("w:val"), jc)
            ppr.append(j)
        mark = OxmlElement("w:rPr")
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            f.set(qn(a), font)
        mark.append(f)
        if size:
            for tag in ("w:sz", "w:szCs"):
                sz = OxmlElement(tag)
                sz.set(qn("w:val"), str(int(size * 2)))
                mark.append(sz)
        ppr.append(mark)
        p.append(ppr)
        if text != "":
            p.append(self.run(text, bold=bold, italic=italic, size=size, font=font, color=color, hl=hl))
        return p

    def new_num(self, like_num_id: int = 59) -> int:
        """A new list instance, numbered from 1, on the abstract numbering of one of BDOI's lists."""
        nums = self.numbering.findall(qn("w:num"))
        abstract = None
        for n in nums:
            if n.get(qn("w:numId")) == str(like_num_id):
                abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        n = OxmlElement("w:num")
        n.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        n.append(a)
        for lvl in range(3):
            o = OxmlElement("w:lvlOverride")
            o.set(qn("w:ilvl"), str(lvl))
            s = OxmlElement("w:startOverride")
            s.set(qn("w:val"), "1")
            o.append(s)
            n.append(o)
        mac = self.numbering.find(qn("w:numIdMacAtCleanup"))
        (mac.addprevious(n) if mac is not None else self.numbering.append(n))
        return new_id

    def items(self, entries, num_id, level=0):
        out = []
        for e in entries:
            out.append(self.para(e, style="ListParagraph", num=(num_id, level)))
            self.stats["requirements_added"] += 1
        return out

    def acceptance(self, ac):
        out = [self.para("Acceptance criteria", bold=True, italic=True, keep_next=True)]
        n = self.new_num()
        for a in ac:
            out.append(self.para(a, style="ListParagraph", num=(n, 0)))
        self.stats["ac"] += len(ac)
        return out

    @staticmethod
    def insert_after(anchor, elements):
        for e in elements:
            anchor.addnext(e)
            anchor = e
        return anchor

    def table(self, header, rows, widths, size=8.5, head_fill=HEAD_FILL, shade_body=True, shade_rows=None):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl = t._tbl
        self.body.remove(tbl)
        tblpr = tbl.tblPr
        for old in tblpr.findall(qn("w:tblStyle")):
            tblpr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        tblpr.append(borders)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblpr.append(lay)
        for old in tblpr.findall(qn("w:tblW")):
            tblpr.remove(old)
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), str(int(sum(widths) * 1440)))
        tw.set(qn("w:type"), "dxa")
        tblpr.append(tw)
        grid = tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        for ri, row in enumerate(t.rows):
            vals = header if ri == 0 else rows[ri - 1]
            if ri == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for ci, cell in enumerate(row.cells):
                tc = cell._tc
                tcpr = tc.get_or_add_tcPr()
                for old in tcpr.findall(qn("w:tcW")):
                    tcpr.remove(old)
                tcw = OxmlElement("w:tcW")
                tcw.set(qn("w:w"), str(int(widths[ci] * 1440)))
                tcw.set(qn("w:type"), "dxa")
                tcpr.append(tcw)
                shaded = (self.hl and shade_body) if shade_rows is None else (self.highlight and ri in shade_rows)
                fill = head_fill if ri == 0 else (HL_FILL if shaded else None)
                if fill:
                    tcpr.append(shd_el(fill))
                for p in list(tc.findall(qn("w:p"))):
                    tc.remove(p)
                for line in str(vals[ci] if vals[ci] is not None else "").split("\n"):
                    tc.append(self.para(line, bold=(ri == 0), size=size, font="Calibri", hl=False, space_after=0))
        return tbl

    def picture(self, path: Path, max_w=6.8, max_h=8.6):
        from PIL import Image  # noqa: PLC0415
        with Image.open(path) as im:
            w, h = im.size
        width = max_w
        if h / w * width > max_h:
            width = max_h * w / h
        p = self.doc.add_paragraph()
        p.add_run().add_picture(str(path), width=Inches(width))
        el = p._p
        self.body.remove(el)
        j = OxmlElement("w:jc")
        j.set(qn("w:val"), "center")
        ppr_insert(el.get_or_add_pPr(), j)
        return el

    def clone_par(self, p_el, text):
        """A new paragraph with the paragraph and run formatting of one of BDOI's paragraphs."""
        new = strip_ids(copy.deepcopy(p_el))
        runs = new.findall(qn("w:r"))
        rpr = copy.deepcopy(runs[0].find(qn("w:rPr"))) if runs and runs[0].find(qn("w:rPr")) is not None else None
        for el in list(new):
            if el.tag != qn("w:pPr"):
                new.remove(el)
        r = self.run(text)
        if rpr is not None:
            for b in rpr.findall(qn("w:b")) + rpr.findall(qn("w:bCs")):
                rpr.remove(b)
            r.replace(r.find(qn("w:rPr")), rpr)
            if self.hl:
                rpr_shade(rpr)
        new.append(r)
        return new

    def clone_row(self, template_tr, values, font_size=None):
        """A copy of one of BDOI's table rows with new cell texts (first paragraph and run formatting kept)."""
        tr = strip_ids(copy.deepcopy(template_tr))
        for h in tr.findall(qn("w:trPr") + "/" + qn("w:trHeight")):
            h.getparent().remove(h)
        for tc, v in zip(tr.findall(qn("w:tc")), values):
            paras = tc.findall(qn("w:p"))
            for p in paras[1:]:
                tc.remove(p)
            p = paras[0]
            runs = p.findall(qn("w:r"))
            rpr = copy.deepcopy(runs[0].find(qn("w:rPr"))) if runs and runs[0].find(qn("w:rPr")) is not None else None
            for el in list(p):
                if el.tag != qn("w:pPr"):
                    p.remove(el)
            num = p.find(qn("w:pPr") + "/" + qn("w:numPr")) if p.find(qn("w:pPr")) is not None else None
            if num is not None:
                num.getparent().remove(num)
            lines = str(v).split("\n")
            for i, line in enumerate(lines):
                target = p if i == 0 else strip_ids(copy.deepcopy(p))
                if i:
                    for el in list(target):
                        if el.tag != qn("w:pPr"):
                            target.remove(el)
                    tc.append(target)
                r = self.run(line, size=font_size)
                if rpr is not None:
                    new_rpr = copy.deepcopy(rpr)
                    for b in new_rpr.findall(qn("w:b")) + new_rpr.findall(qn("w:bCs")):
                        new_rpr.remove(b)
                    r.replace(r.find(qn("w:rPr")), new_rpr)
                    if self.hl:
                        rpr_shade(new_rpr)
                target.append(r)
            if self.hl:
                cell_shade(tc)
        return tr


def find_par(b: Builder, start: str, style: str | None = None):
    for p in b.doc.paragraphs:
        if p.text.strip().startswith(start) and (style is None or p.style.name == style):
            return p
    raise SystemExit(f"paragraph not found: {start}")


# ------------------------------------------------------------------------------------------------------- the edits
def edit_cover(b: Builder):
    t = b.tables[1]
    rev = next(r for r in t.rows if r.cells[0].text.strip().startswith("Revision No"))
    ts = list(rev.cells[1]._tc.iter(qn("w:t")))
    if "".join(t.text or "" for t in ts).strip() != DOC["meta"]["version_old"]:
        raise SystemExit("cover: revision not found")
    ts[0].text = DOC["meta"]["version_new"]
    for extra in ts[1:]:
        extra.text = ""
    if b.highlight:
        rpr = ts[0].getparent().find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ts[0].getparent().insert(0, rpr)
        rpr_shade(rpr)
    prep = next(r for r in t.rows if r.cells[0].text.strip().startswith("Prepared by"))
    tr = b.clone_row(prep._tr, [DOC["cover"]["drafting_label"], DOC["cover"]["drafting_value"]])
    prep._tr.addnext(tr)


def add_revision_log(b: Builder):
    toc = find_par(b, "Table of Contents", "Heading 1")._p
    head = strip_ids(copy.deepcopy(toc))
    for el in list(head):
        if el.tag != qn("w:pPr"):
            head.remove(el)
    head.append(b.run(DOC["revision_log"]["title"]))
    tbl = b.table(DOC["revision_log"]["header"], DOC["revision_log"]["rows"], [1.5, 0.8, 3.4, 1.7], size=9,
                  shade_rows={len(DOC["revision_log"]["rows"])})
    brk = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    brk.append(r)
    for el in (head, tbl, b.para("", hl=False), brk):
        toc.addprevious(el)


def edit_introduction(b: Builder):
    p = find_par(b, "The System Analyst, Business System Analyst or equivalent BDO ITG personnel")
    p._p.addnext(b.clone_par(p._p, DOC["introduction"]["preparation"]))
    p = find_par(b, "The following documents are used as reference")
    ref12 = b.clone_par(p._p, DOC["introduction"]["references"])
    p._p.addnext(ref12)
    b.round13 = True
    ref12.addnext(b.clone_par(p._p, v13.REV["meta"]["references"]))
    b.round13 = False


def bdoi_row_key(rid: str) -> str:
    parts = rid.replace("BRD ", "").replace(" (2nd)", "").split(".")
    return ".".join(parts[:3] if parts[:2] in (["3", "002"], ["3", "003"]) else parts[:2])


def edit_mapping(b: Builder, lines: list[tuple[str, str]], trace: dict) -> dict:
    t = b.tables[3]
    rows = OrderedDict()
    for row in t.rows[1:]:
        rows[row.cells[0].text.replace("\n", "").strip()] = row
    # FR IDs added to BDOI's rows
    for rid, change in DOC["mapping"].items():
        row = rows[rid]
        cell = row.cells[2]._tc
        before = re.sub(r"\s+", " ", row.cells[2].text).strip()
        model = next(p for p in cell.findall(qn("w:p")) if p.find(qn("w:r")) is not None)
        last = cell.findall(qn("w:p"))[-1]
        last.addnext(b.clone_par(model, "Added in v1.2: " + ", ".join(change["add"])))
        b.log.append({"id": rid, "before": before, "after": before + "; added in v1.2: " + ", ".join(change["add"]),
                      "why": change["why"], "brd": re.sub(r"\s+", " ", row.cells[1].text).strip()[:300]})
    # every BRD line under the row of its function
    by_key = {bdoi_row_key(rid): rid for rid in rows}
    groups: dict[str, list] = {}
    for rid, text in lines:
        if rid in rows:
            continue
        groups.setdefault(by_key[bdoi_row_key(rid)], []).append((rid, text))
    template = t.rows[1]._tr
    for gid, items in groups.items():
        anchor = rows[gid]._tr
        for rid, text in items:
            frum = sorted({x for fr in trace[rid]["frs"] for x in b.fr_to_frum.get(fr, [])}, key=frum_key)
            tr = b.clone_row(template, [rid, text, ", ".join(frum)])
            anchor.addnext(tr)
            anchor = tr
            b.stats["leaf_rows"] += 1
    mapping = {}
    for row in t.rows[1:]:
        mapping[row.cells[0].text.replace("\n", "").strip()] = re.sub(r"\s+", " ", row.cells[2].text).strip()
    ids = [rid for rid, _ in lines]
    missing = [i for i in ids if i not in mapping]
    unmapped = [i for i in ids if i in mapping and not FRUM_RE.search(mapping[i])]
    if missing or unmapped:
        raise SystemExit(f"mapping check failed: missing {missing}, unmapped {unmapped}")
    b.stats["brd_ids"] = len(ids)
    b.stats["bdoi_rows"] = len(rows)
    b.stats["mapping_changes"] = len(DOC["mapping"])
    return mapping


def subitem_blocks(cell_tc):
    blocks = OrderedDict()
    current = None
    seen: Counter = Counter()
    for el in cell_tc:
        if el.tag == qn("w:tcPr"):
            continue
        text = el_text(el) if el.tag == qn("w:p") else ""
        m = re.match(r"\s*(FRUM[.-]\d{3}(?:\.\d{2})+)", text)
        bold = el.tag == qn("w:p") and el.find(".//" + qn("w:b")) is not None
        if m and bold and el.find(qn("w:pPr") + "/" + qn("w:numPr")) is None:
            seen[m.group(1)] += 1
            current = m.group(1) + (" (2nd)" if seen[m.group(1)] > 1 else "")
            blocks[current] = [el]
        elif current:
            blocks[current].append(el)
    return blocks


def last_content(block):
    for el in reversed(block):
        if el.tag == qn("w:tbl") or el_text(el).strip():
            return el
    return block[-1]


def block_num(block, default) -> int:
    ids = Counter()
    for el in block:
        n = el.find(qn("w:pPr") + "/" + qn("w:numPr") + "/" + qn("w:numId")) if el.tag == qn("w:p") else None
        if n is not None:
            ids[int(n.get(qn("w:val")))] += 1
    return ids.most_common(1)[0][0] if ids else default


def subitem_elements(b: Builder, sid, title, shall, ac):
    out = [b.para(""), b.para(f"{sid} {title}", bold=True, keep_next=True)]
    out += b.items(shall, b.new_num())
    out += b.acceptance(ac)
    return out


def edit_functional_requirements(b: Builder):
    t = b.tables[4]
    rows = OrderedDict()
    for row in t.rows[1:]:
        rows[row.cells[0].text.strip()] = row
    b.stats["bdoi_entries"] = len(rows)
    all_blocks = OrderedDict()
    for row in rows.values():
        all_blocks.update(subitem_blocks(row.cells[3]._tc))
    b.stats["bdoi_items"] = len(all_blocks)
    b.bdoi_item_ids = list(all_blocks)
    for sid, add in ADD["items"].items():
        block = all_blocks[sid]
        num = block_num(block, 59)
        els = b.items(add["shall"], num) + b.acceptance(add["ac"])
        b.insert_after(last_content(block), els)
        b.stats["elaborated"] += 1
    for ns in ADD["new_subitems"]:
        cell = rows[ns["row"]].cells[3]._tc
        prev = [s for s in all_blocks if re.sub(r"FRUM-", "FRUM.", s).startswith(ns["id"][:8] + ".")]
        if prev and frum_key(ns["id"]) <= max(frum_key(p) for p in prev):
            raise SystemExit(f"{ns['id']} does not continue BDOI's numbering")
        for el in subitem_elements(b, ns["id"], ns["title"], ns["shall"], ns["ac"]):
            cell.append(el)
        b.stats["new_subitems"] += 1
    template = rows["FRUM.007"]._tr
    last = t.rows[-1]._tr
    for item in ADD["new_items"]:
        tr = strip_ids(copy.deepcopy(template))
        tcs = tr.findall(qn("w:tc"))
        for tc in tcs:
            for el in list(tc):
                if el.tag != qn("w:tcPr"):
                    tc.remove(el)
            if b.hl:
                cell_shade(tc)
        tcs[0].append(b.para(item["id"]))
        tcs[1].append(b.para(item["component"]))
        tcs[2].append(b.para(item["capability"]))
        tcs[4].append(b.para(""))
        for s in item.get("scope") or []:
            label, rest = s.split(":", 1)
            p = b.para("")
            p.append(b.run(label + ": ", bold=True))
            p.append(b.run(rest.strip()))
            tcs[3].append(p)
            tcs[3].append(b.para(""))
        first = not item.get("scope")
        for si in item["subitems"]:
            els = subitem_elements(b, si["id"], si["title"], si["shall"], si["ac"])
            for el in (els[1:] if first else els):
                tcs[3].append(el)
            first = False
            b.stats["new_item_subitems"] += 1
        last.addnext(tr)
        last = tr
        b.stats["new_items"] += 1


def edit_signoff(b: Builder):
    cell = b.tables[12].rows[2].cells[0]
    for p in cell.paragraphs:
        if p.text.strip().endswith("Business System Analyst"):
            new = strip_ids(copy.deepcopy(p._p))
            for r in new.findall(qn("w:r")):
                new.remove(r)
            new.append(b.run(DOC["signoff"]["drafting"]))
            p._p.addnext(new)
            return
    raise SystemExit("sign-off: 'Business System Analyst' not found")


def edit_annex_rows(b: Builder):
    index = {"E": 8, "F": 9, "G": 10, "H": 11}
    b.annex_log = []
    for annex, rows in DOC["annex_rows"].items():
        t = b.tables[index[annex]]
        last = t.rows[-1]._tr
        for vals in rows:
            tr = b.clone_row(last, vals)
            last.addnext(tr)
            last = tr
            b.stats["annex_rows"] += 1
        b.annex_log.append((annex, rows))


# ------------------------------------------------------------------------------------------------- the new annexes
class Annexes:
    def __init__(self, b: Builder, mapping: dict, lines, trace: dict, our: dict):
        self.b = b
        self.mapping = mapping
        self.lines = lines
        self.trace = trace
        self.our = our
        self.els: list = []
        self.h1_ppr = find_par(b, "Annex A", "Heading 1")._p.find(qn("w:pPr"))
        self.screen_titles: dict[str, str] = {}

    def refs(self, text: str) -> str:
        def one(m):
            items = self.b.fr_to_frum.get(m.group(0))
            return ", ".join(items[:2]) if items else m.group(0)
        return re.sub(r"FR-UA-\d{3}", one, clean(text))

    def frum_of(self, frs) -> str:
        return ", ".join(sorted({x for fr in frs for x in self.b.fr_to_frum.get(fr, [])}, key=frum_key))

    def h1(self, text):
        p = OxmlElement("w:p")
        ppr = strip_ids(copy.deepcopy(self.h1_ppr))
        for r in ppr.findall(qn("w:rPr")):
            ppr.remove(r)
        if ppr.find(qn("w:pageBreakBefore")) is None:
            ppr_insert(ppr, OxmlElement("w:pageBreakBefore"))
        p.append(ppr)
        p.append(self.b.run(text))
        self.els.append(p)

    def h2(self, text):
        self.els.append(self.b.para(text, style="Heading2", no_num=True, keep_next=True))

    def p(self, text, **kw):
        self.els.append(self.b.para(text, **kw))

    def bullets(self, texts):
        n = self.b.new_num()
        for t in texts:
            self.els.append(self.b.para(t, style="ListParagraph", num=(n, 0)))

    def table(self, header, rows, widths, size=8.5, changed=None):
        """A table; `changed` names the first cells of the rows corrected in this round (shaded in the review copy)."""
        shade = None if not changed else {i for i, r in enumerate(rows, 1) if r[0] in changed}
        self.els.append(self.b.table(header, rows, widths, size=size, shade_rows=shade))
        self.els.append(self.b.para("", space_after=0))

    # Annex J - screens
    def screens(self):
        self.h1("Annex J – Screen Specifications")
        self.p("This annex specifies the User Access Maintenance screens: the purpose of each screen, how the user "
               "reaches it, the requirements it serves, one screenshot, its fields with type, mandatory marker, format, "
               "list or source and validation, its actions, and its rules. Mandatory: Y = mandatory, N = optional, "
               "Cond. = mandatory when the condition stated applies. The screenshots were taken on the BIBS test "
               "environment with SIT/UAT data on 10-Oct-2026.")
        k = 0
        for f in sorted((UA / "pack" / "screens").glob("*.yaml")):
            for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
                k += 1
                self.screen_titles[s["id"]] = clean(s["title"])
                self.b.round13 = s["id"] in SCREENS_V13
                self.h2(f"J.{k} {clean(s['title'])} ({s['id']})")
                self.p("Purpose: " + self.refs(" ".join(str(s["purpose"]).split())))
                if s.get("entry"):
                    self.p("Access: " + clean("; ".join(s["entry"])))
                frum = self.frum_of(s.get("frs", []))
                if frum:
                    self.p("Requirements: " + frum)
                shots = s.get("shots") or []
                n = int(s["id"].split("-")[-1])
                pick = 1 if s["id"] == "SCR-UA-07" and len(shots) > 1 else 0
                if shots:
                    png = UA / "screenshots" / f"scr-ua-{n:02d}-{pick + 1:02d}-{shots[pick]['state']}.png"
                    if not png.exists():
                        raise SystemExit(f"screenshot missing: {png}")
                    self.els.append(self.b.picture(png, max_w=6.4, max_h=4.8 if s["id"] != "SCR-UA-07" else 7.5))
                    self.p(f"Figure J.{k}: {clean(shots[pick].get('caption', s['title']))}", italic=True, size=8,
                           jc="center")
                    self.b.stats["screenshots"] += 1
                rows = []
                for line in s.get("fields", []):
                    c = [x.strip() for x in str(line).split(" | ")]
                    c += ["-"] * (10 - len(c))
                    if "[Group profile check boxes]" in c[1]:
                        rows.append(["Group Profiles: profile picker", "Picker", "Cond.: enrolment (at least one)",
                                     "Active group profiles grouped by business area, in groups that open and close "
                                     "with their counts; search by profile, area or what the profile does; each "
                                     "profile with a one-line description, a details button (key permissions, kind, "
                                     "level) and the marker Privileged for High and Admin profiles; the chosen "
                                     "profiles listed apart, marked Current, Added or Removed for a modification and "
                                     "removable one by one",
                                     "Active profiles only; a user cannot change his own profiles; two chosen "
                                     "profiles of an active separation-of-duties rule are flagged with a warning in "
                                     "the picker and refused on Submit; Select at least one role; One user may not "
                                     "hold both <profile> and <profile> (separation-of-duties rule <rule>)"])
                        continue
                    label = c[1] if c[0] in ("-", "") else f"{c[0]}: {c[1]}"
                    fmt = "; ".join(x for x in (c[3], c[5]) if x not in ("-", ""))
                    val = "; ".join(x for x in (c[8], c[9]) if x not in ("-", ""))
                    rows.append([clean(label), clean(c[2]), clean(c[4]), self.refs(fmt) or "-", self.refs(val) or "-"])
                if rows:
                    self.table(["Field", "Type", "Mandatory", "Format / list / source", "Validation and message"],
                               rows, [1.5, 0.85, 0.9, 2.2, 1.95], size=7.5)
                acts = []
                for line in s.get("actions", []):
                    c = [x.strip() for x in str(line).split(" | ")]
                    c += ["-"] * (6 - len(c))
                    acts.append([clean(c[0]), clean(c[1]), clean(c[2]), self.refs(c[3]),
                                 clean(c[4]) if c[4] not in ("", "-") else "-"])
                if acts:
                    self.table(["Action", "Who", "Available when", "What happens", "Resulting status"], acts,
                               [1.2, 1.4, 1.2, 2.6, 1.0], size=7.5)
                if s.get("rules"):
                    self.p("Rules:", bold=True, keep_next=True)
                    self.bullets([self.refs(r) for r in s["rules"]])
                self.b.round13 = False
        self.b.stats["screens"] = k

    # Annex K - messages
    def messages(self):
        self.h1("Annex K – Messages and Validations")
        self.p("The messages the user sees when a check fails or an action is not allowed, word for word as on the "
               "screens; <name> parts are filled in by the system. The messages of each screen are also in Annex J.")
        rows, seen = [], set()
        for fr_id, fr in self.our.items():
            frum = self.frum_of([fr_id]) or "-"
            screen = clean(str(fr.get("screens", "")).split(";")[0])
            for v in fr.get("validations") or []:
                msg = clean(v[1])
                if not msg or msg in seen or msg == "-":
                    continue
                seen.add(msg)
                rows.append([screen, clean(v[0]), msg, frum])
        for f in sorted((UA / "pack" / "screens").glob("*.yaml")):
            for s in yaml.safe_load(f.read_text(encoding="utf-8"))["screens"]:
                for line in s.get("fields", []):
                    c = [x.strip() for x in str(line).split(" | ")]
                    if len(c) >= 10 and c[9] not in ("-", ""):
                        for msg in (clean(m) for m in re.split(r";\s+(?=[A-Z])", c[9])):
                            if msg and msg not in seen:
                                seen.add(msg)
                                rows.append([clean(s["title"]), clean(f"{c[1]}: {c[8]}"), msg,
                                             self.frum_of(s.get("frs", [])) or "-"])
        rows = [[str(i)] + r for i, r in enumerate(rows, 1)]
        self.table(["No.", "Screen", "When", "Message shown", "Requirement"], rows, [0.4, 1.5, 1.9, 2.4, 1.2],
                   size=7.5)
        self.b.stats["messages"] = len(rows)

    # Annex L - notifications
    def notifications(self):
        self.h1("Annex L – Notifications and E-mails")
        self.p("The notifications the system sends: notices in the system (bell in the page header), e-mails and "
               "alerts. The temporary password of a new user is never e-mailed.")
        b13 = self.b.round13
        self.b.round13 = True
        self.p("Corrected in version 1.3: the notices of access requests are given in the system; the e-mails are "
               "the access-change, dormant-account, password and failed-run messages (Annex Q.5). The full text of "
               "every notice and e-mail, with its subject, recipients and rules, is in Annex T, which prevails.")
        self.b.round13 = b13
        data = yaml.safe_load((UA / "pack" / "notifications.yaml").read_text(encoding="utf-8"))["notifications"]
        texts = {i["id"]: i for i in v13.REV["emails"]["items"]}
        rows, changed = [], set()
        for k, n in enumerate(data, 1):
            t = texts[n["id"]]
            channel, text = clean(n["channel"]), clean(n["template"])
            if channel != t["channel"] and "e-mail by preference" in n["channel"]:
                channel = t["channel"]
                changed.add(k)
            if n["id"] == "NT-18":
                text = f"Title: {t['subject']}. Text: {t['body']}"
                changed.add(k)
            rows.append([n["id"], channel, self.refs(n["trigger"]), clean(n["recipient"]), text,
                         self.frum_of(n.get("frs", [])) or "-"])
        self.els.append(self.b.table(["No.", "Channel", "Trigger", "Recipient", "Text", "Requirement"], rows,
                                     [0.5, 0.9, 1.7, 1.4, 2.0, 0.9], size=7.5, shade_rows=changed))
        self.els.append(self.b.para("", space_after=0))
        self.b.stats["notifications"] = len(rows)

    # Annex M - rules and parameters
    def rules(self):
        m = DOC["annex_m"]
        self.h1("Annex M – Business Rules and Parameters")
        self.p("The rules that apply across the User Access Maintenance requirements and the values the System "
               "Administrator maintains without a change to the system. A change of a security setting applies only "
               "after the approval of Information Security (FRUM.012.02).")
        self.h2("M.1 Personas")
        self.table(["Persona", "Responsibilities", "BRD"], m["personas"], [1.6, 4.6, 1.2])
        self.h2("M.2 Permissions of the access functions")
        self.p("Each access function has its own permission, so that it can be granted to any group profile "
               "(BRD 4.002.2; FRUM.005.06).")
        self.table(["Function", "Action class", "BRD"], m["permissions"], [4.2, 1.4, 1.8])
        self.h2("M.3 Proposed role-to-action matrix")
        self.p("Y = the persona's group profile holds the permission. Proposal for BDOI to confirm with the User Access "
               "Matrix of Annex C; the segregation rules of M.6 apply whatever the roles grant.")
        self.table(["Function"] + m["matrix_roles"], m["matrix"], [2.2] + [0.65] * len(m["matrix_roles"]), size=7.5)
        self.h2("M.4 Password, lock-out and second factor")
        self.table(["Rule", "Value", "Applies to"], m["passwords"], [1.6, 3.9, 1.9],
                   changed={"Password history", "Password reset"})
        self.h2("M.5 Session policy")
        self.table(["Rule", "Value", "Note"], m["sessions"], [1.8, 3.0, 2.6], changed={"Inactivity sign-out"})
        self.h2("M.6 Segregation and separation-of-duties rules")
        self.bullets(m["controls"])
        self.h2("M.7 Configurable parameters")
        self.table(["Parameter", "Default", "Meaning", "Value decided by"], m["parameters"], [1.7, 1.9, 2.6, 1.2],
                   size=8, changed={"Password history"})
        self.h2("M.8 Lists of values")
        self.table(["List", "Values"], m["lists"], [1.8, 5.6])
        self.h2("M.9 Request statuses and transitions")
        self.table(["Status", "Meaning", "Who acts"], m["statuses"], [1.6, 3.6, 2.2])
        self.table(["From", "Action", "To", "Who", "Remarks"], m["transitions"], [1.6, 1.1, 1.9, 1.9, 0.9], size=8)
        b13 = self.b.round13
        self.b.round13 = True
        self.h2("M.10 Settings for the points BDOI decides")
        self.p("The points where BDOI's FRS and the BRD or the earlier BIBS behaviour differ are offered both ways in "
               "BIBS, each chosen by a setting that the System Administrator changes with the approval of "
               "Information Security, without a release. The default is BDOI's FRS unless stated. BDOI's decision "
               "only confirms the value; the points stay in Annex Q as Configurable - BDOI to confirm the setting.")
        self.table(["Setting", "Values", "Default", "Settles (Annex Q)", "Decided by"], m["decision_settings"],
                   [1.5, 3.3, 1.1, 0.8, 1.2], size=8)
        self.b.round13 = b13
        self.b.stats["decision_settings"] = len(m["decision_settings"])

    # Annex N - sign-in and identity integration
    def identity(self):
        n = DOC["annex_n"]
        self.h1("Annex N – Sign-in and Identity Integration")
        self.p("The business view of the sign-in through the Enterprise SSO platform (EIAM), the provisioning of user "
               "accounts from UIDM-ISC and the break-glass access. The technical settings are in the Authentication "
               "and Identity Integration Requirements reviewed by BDOI IT.")
        self.h2("N.1 Sign-in modes")
        self.table(["Mode", "How the user signs in", "Who"], n["modes"], [1.6, 4.3, 1.5])
        self.h2("N.2 Single sign-on with EIAM")
        self.bullets(n["steps"])
        self.p("Rules:", bold=True, keep_next=True)
        self.bullets(n["rules"])
        self.h2("N.3 Provisioning from UIDM-ISC")
        b13 = self.b.round13
        self.b.round13 = True
        self.p("BIBS offers the three options as settings (Annex M.10); as delivered option (a) applies. BDOI confirms "
               "the settings (Annex Q, conflict C01).")
        self.table(["Option", "Requests and approvals", "Exchange with UIDM-ISC"], n["provisioning"], [2.0, 3.6, 1.8])
        self.b.round13 = b13
        self.bullets(n["provisioning_rules"][:1])
        self.b.round13 = True
        self.bullets(n["provisioning_rules"][1:4])
        self.b.round13 = b13
        self.bullets(n["provisioning_rules"][4:])
        self.h2("N.4 Break-glass administrators")
        self.table(["Rule", "Value"], n["break_glass"], [1.4, 6.0])
        self.h2("N.5 What BDOI IT provides")
        self.bullets(n["bdoi_it"])

    # Annex O - walkthroughs
    def walkthroughs(self):
        cases = yaml.safe_load((UA / "brd11_cases.yaml").read_text(encoding="utf-8"))
        personas = {k: clean(v["name"]) for k, v in cases["personas"].items()}
        data = yaml.safe_load((UA / "pack" / "walkthroughs.yaml").read_text(encoding="utf-8"))["walkthroughs"]
        self.h1("Annex O – Process Walkthroughs")
        self.p("End-to-end walkthroughs of User Access Maintenance, step by step: who acts, on which screen, what the "
               "user does, what the user sees and the result. The screens are specified in Annex J.")
        total = 0
        for w in data:
            self.h2(f"{w['id']} {clean(w['title'])}")
            self.p(self.refs(" ".join(str(w["summary"]).split())))
            if w.get("data"):
                self.p(clean(w["data"]), italic=True, size=9)
            rows = []
            for i, s in enumerate(w["steps"], 1):
                rows.append([f"{w['id']}.{i}", personas.get(s[0], clean(s[0])),
                             f"{self.screen_titles.get(s[1], s[1])} ({s[1]})", self.refs(s[2]), self.refs(s[3]),
                             clean(s[4])])
            total += len(rows)
            self.table(["Step", "Persona", "Screen", "What the user does", "What the user sees", "Result"], rows,
                       [0.6, 0.95, 1.15, 1.95, 1.75, 1.0], size=7.5)
            if w.get("downstream"):
                self.p("Afterwards:", bold=True, keep_next=True)
                self.bullets([self.refs(x) for x in w["downstream"]])
        self.b.stats["walk_steps"] = total

    # Annex P - traceability
    def traceability(self, nfr):
        cases = test_cases()
        self.h1("Annex P – Requirements Traceability")
        self.p("P.1 traces every requirement line of the User Access Maintenance BRD, and the non-functional rows that "
               "carry a function, to the functional requirements of this FRS (FRUM), to the reference requirements of "
               "the BIBS FRS BRD-11 User Access Maintenance (FR-UA) and to the test conditions of the BIBS test plan "
               "BRD-11 (TC-UA-nnn.n; the cases of each condition are in the test plan). The BRD gives no IDs to its "
               "non-functional rows; they carry the IDs UAM-NFR-nn of the BIBS FRS. P.2 lists every reference "
               "requirement with the FRUM items that incorporate it.")
        self.h2("P.1 BRD requirement to FRUM, BIBS reference and test cases")
        rows = []
        for rid, _ in self.lines:
            frum = ", ".join(FRUM_RE.findall(self.mapping[rid]))
            refs = self.trace[rid]["frs"]
            rows.append([rid, frum, ", ".join(refs), "; ".join(tc_text(r, cases) for r in refs)])
        for n in nfr:
            if not n["frs"]:
                continue
            rows.append([f"{n['id']} ({clean(n['topic'])})", self.frum_of(n["frs"]), ", ".join(n["frs"]),
                         "; ".join(tc_text(r, cases) for r in n["frs"])])
        self.table(["BRD ID", "FRUM (this FRS)", "BIBS reference (FR-UA)", "Test conditions (BIBS test plan BRD-11)"],
                   rows, [1.7, 3.2, 1.6, 3.4], size=7.5)
        self.b.stats["trace_rows"] = len(rows)
        self.h2("P.2 BIBS reference requirement to FRUM")
        rows = []
        self.listed_refs = set()
        for fr_id, fr in self.our.items():
            frum = self.frum_of([fr_id])
            if not frum:
                raise SystemExit(f"{fr_id} is not incorporated in any FRUM item")
            brd = ", ".join(x.split(" (")[0] for x in fr["brd"] if isinstance(x, str))
            rows.append([fr_id, clean(fr["title"]), clean(brd), frum, tc_text(fr_id, cases)])
            self.listed_refs.add(fr_id)
        self.table(["Reference", "Title", "BRD", "FRUM (this FRS)", "Test conditions"], rows,
                   [0.85, 2.6, 2.4, 2.2, 1.85], size=7.5)

    # Annex Q - observations
    def observations(self):
        cmp_data = yaml.safe_load(CMP.read_text(encoding="utf-8"))
        self.h1("Annex Q – Observations and Points for BDOI Decision")
        self.p("Every point where BDOI's FRS and the BIBS reference FRS differ, the slips noticed in BDOI's text, the "
               "changes made to the Business Requirements Mapping, the rows added to BDOI's annexes and the other "
               "observations of this version. BDOI's text is kept in the body of the document; this annex gives both "
               "readings, the BRD text, what version 1.2 proposes, the impact and who decides. Every point starts as "
               "Open; a point that BIBS offers both ways, chosen by a system setting, reads Configurable - BDOI to "
               "confirm the setting.")
        header = ["No.", "Topic", "BDOI FRS text", "BRD text", "Proposed in v1.2", "Impact", "Decision by", "Status"]
        widths = [0.75, 1.0, 1.9, 1.5, 2.2, 0.95, 1.0, 0.6]
        no = 0

        def row(ref, topic, bdoi, brd, prop, impact, who, status="Open"):
            nonlocal no
            no += 1
            out = [f"Q-{no:02d}" + (f" ({ref})" if ref else ""), topic, bdoi, brd, prop, impact, who, status]
            self.b.q_rows.append(dict(zip(["no", "topic", "bdoi", "brd", "proposal", "impact", "decide_by", "status"],
                                          out), section=section))
            return out
        section = "Q.1"
        self.listed_conflicts = set()
        self.h2("Q.1 Conflicts between BDOI's FRS and the BIBS reference FRS")
        rows = []
        for c in cmp_data["conflicts"]:
            if c["kind"] != "Between the documents":
                continue
            prop = f"{clean(c['recommendation'])} (BIBS reference FRS: {clean(c['ours'])})"
            if c.get("setting"):
                prop += f" Available in BIBS: {clean(c['setting'])}"
            rows.append(row(c["id"], clean(c["topic"]), clean(c["bdoi"]), clean(c["brd"]), prop, clean(c["impact"]),
                            c["decide_by"], status=c.get("status", "Open")))
            self.listed_conflicts.add(c["id"])
        self.table(header, rows, widths, size=7)
        section = "Q.2"
        self.h2("Q.2 Slips noticed in BDOI's FRS (proposed corrections; BDOI's text is not changed)")
        rows = []
        for c in cmp_data["conflicts"]:
            if c["kind"] == "Between the documents":
                continue
            rows.append(row(c["id"], clean(c["topic"]), clean(c["bdoi"]), clean(c["brd"]), clean(c["recommendation"]),
                            clean(c["impact"]), c["decide_by"]))
            self.listed_conflicts.add(c["id"])
        self.table(header, rows, widths, size=7)
        section = "Q.3"
        self.h2("Q.3 Changes to the Business Requirements Mapping")
        rows = []
        for ch in self.b.log:
            rows.append(row(ch["id"], "Mapping extended", f"Mapped to: {ch['before']}", ch["brd"],
                            f"Mapped to: {ch['after']}. {ch['why']}", "Coverage of the BRD requirement.",
                            "System Analyst (BDO ITG); Process Owner (Business Administration)"))
        rows.append(row("", "BRD lines added", f"{self.b.stats['bdoi_rows']} rows, one per BRD function",
                        f"{self.b.stats['brd_ids']} requirement lines",
                        f"{self.b.stats['leaf_rows']} BRD lines added under the row of their function, each with the "
                        f"BRD text and the FRUM items that cover it; BDOI's rows are kept as they are.",
                        "Coverage of every BRD ID.", "System Analyst (BDO ITG)"))
        self.table(header, rows, widths, size=7)
        section = "Q.4"
        self.h2("Q.4 Rows added to BDOI's annexes and other observations")
        rows = []
        for annex, added in self.b.annex_log:
            fields = ", ".join(r[1] for r in added)
            rows.append(row(f"Annex {annex}", f"Fields added to Annex {annex}", f"Annex {annex} as issued",
                            "-", f"Rows added: {fields}. {DOC['annex_why'][annex]}", f"Annex {annex}.",
                            "System Analyst (BDO ITG); Process Owner (Business Administration)"))
        for o in DOC["observations"]:
            rows.append(row("", o["topic"], o["bdoi"], o["brd"], o["proposal"], o["impact"], o["decide_by"]))
        self.table(header, rows, widths, size=7)
        section = "Q.5"
        self.b.round13 = True
        self.h2("Q.5 Observations of version 1.3")
        rows = [row("", o["topic"], o["bdoi"], o["brd"], o["proposal"], o["impact"], o["decide_by"])
                for o in v13.REV["observations"]]
        self.table(header, rows, widths, size=7)
        self.b.round13 = False
        self.b.stats["observations_v13"] = len(rows)
        self.b.stats["observations"] = no
        self.b.stats["conflicts"] = sum(1 for c in self.listed_conflicts if c.startswith("C"))
        self.b.stats["slips"] = sum(1 for c in self.listed_conflicts if c.startswith("S"))
        self.cmp_ids = {c["id"] for c in cmp_data["conflicts"]}

    # Annex R - glossary
    def glossary(self):
        self.h1("Annex R – Glossary")
        self.p("Terms and abbreviations used in this document.")
        terms = sorted(DOC["glossary"], key=lambda g: g[0].lower())
        self.els.append(self.b.table(["Term", "Meaning"], terms, [1.9, 5.5], size=9,
                                     shade_rows={i for i, g in enumerate(terms, 1) if g[0] in NEW_TERMS}))
        self.els.append(self.b.para("", space_after=0))
        self.b.stats["glossary"] = len(DOC["glossary"])


def section_break(b: Builder, landscape: bool):
    """A paragraph that ends a section, with the page set-up of the document's last section."""
    sect = copy.deepcopy(b.doc.sections[-1]._sectPr)
    for t in sect.findall(qn("w:titlePg")):
        sect.remove(t)
    pg = sect.find(qn("w:pgSz"))
    w, h = int(pg.get(qn("w:w"))), int(pg.get(qn("w:h")))
    if landscape:
        pg.set(qn("w:w"), str(max(w, h)))
        pg.set(qn("w:h"), str(min(w, h)))
        pg.set(qn("w:orient"), "landscape")
    else:
        pg.set(qn("w:w"), str(min(w, h)))
        pg.set(qn("w:h"), str(max(w, h)))
        if qn("w:orient") in pg.attrib:
            del pg.attrib[qn("w:orient")]
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(sect)
    p.append(ppr)
    return p


def add_annexes(b: Builder, mapping, lines, trace, our, nfr, figs, nums, where):
    a = Annexes(b, mapping, lines, trace, our)
    a.screens()
    b.round13 = True
    new = v13.NewAnnexes(a, figs, nums, v13.nfr_rows(nfr, a.frum_of, clean), b.q_rows)
    new.menus()
    b.round13 = False
    a.messages()
    a.notifications()
    a.rules()
    a.identity()
    a.walkthroughs()
    a.els.append(section_break(b, landscape=False))
    a.traceability(nfr)
    a.observations()
    a.els.append(section_break(b, landscape=True))
    a.glossary()
    b.round13 = True
    new.workflow()
    new.emails()
    new.documents()
    new.reports()
    new.integrations()
    new.nfr()
    new.data_setup()
    new.open_items(where)
    new.change_control()
    new.checklist()
    b.round13 = False
    # before the page break that precedes the Signoff Sheet
    sign = find_par(b, "Signoff Sheet", "Heading 1")._p
    anchor = sign
    prev = anchor.getprevious()
    while prev is not None and prev.tag == qn("w:p") and not el_text(prev).strip() \
            and prev.find(".//" + qn("w:sectPr")) is None:
        anchor = prev
        prev = anchor.getprevious()
    for el in a.els:
        anchor.addprevious(el)
    return a


CHANGE_POINTS = [
    "Cover: revision {version}; Document Revision Log: row {version} (Business Unit review edition).",
    "Introduction: new section 1.8 Summary for the Business Unit Review - scope on a page, the process from end to end, key numbers, the {decisions} decisions the Business Unit must take and how to review; references of version {version}.",
    "Annex A: {flows} figures added - one process flow by persona for every process of the BRD, the status life-cycles of the request, the group-profile request, the user account and the separation-of-duties rule, and the integration context.",
    "Annex J: menu by persona (figure) and the reference to the programme screen standards. Annex L: the channel of the request notices corrected. Annex Q.5: {obs13} observations of this round. Annex R: {terms} terms added.",
    "New annexes: S Workflow and Approvals ({workflow_rows} stages and the role-to-stage matrix); T E-mail and Notification Texts ({emails} notices and e-mails, word for word); U Document Prints and Output Formats ({documents}); V Reports and Schedules ({reports} reports, {schedules} runs); W Integrations ({integrations}); X Non-functional Requirements ({nfr_rows} rows); Y Data Set-up and Migration at Go-live ({data_rows}); Z Assumptions, Dependencies and Open Questions ({open_kept} open questions kept, {answered} answered by BDOI's FRS and not repeated); AA Change Control after Sign-off; AB Business Unit Review Checklist ({checklist} points).",
    "Aligned with BIBS on 10 Oct 2026: Annex J screens added (Identity Synchronisation, Session Timed Out, Audit Trail) and updated (Login, inactivity warning in BDOI's wording, Users with the directory details, Separation of Duties with the conflicting permission combinations), every screenshot taken again; Annex M.10 with the {decision_settings} settings for the points BDOI decides; Annex N (provisioning options as settings, simulator for SIT and UAT), Annex T (notices of refused identity events and break-glass sign-ins), Annexes V and W updated.",
    "BDOI's text of version 1.1 and the additions of version 1.2 are unchanged, except the corrections shaded in this copy (Annex L channel, password history of 10, password reset, inactivity sign-out, provisioning options).",
]


def change_summary(b: Builder):
    s = b.stats
    values = dict(version=v13.REV["meta"]["version"], decisions=len(v13.REV["summary"]["decisions"]),
                  flows=sum(1 for f in v13.FLOWS if f["where"] == "A"), obs13=s["observations_v13"],
                  terms=len(v13.REV["glossary_add"]))
    for k in ("workflow_rows", "emails", "documents", "reports", "schedules", "integrations", "nfr_rows",
              "data_rows", "open_kept", "answered", "checklist", "decision_settings"):
        values[k] = s[k]
    els = [b.para(f"Summary of changes in version {values['version']}", bold=True, size=14, color="014EA9", hl=False,
                  space_after=120),
           b.para("This review copy shows in light yellow every passage added in version 1.3, the Business Unit "
                  "review edition. BDOI's text of version 1.1 and the additions of version 1.2 are not shaded; the "
                  "points for decision are in Annex Q and the open questions in Annex Z.", size=9.5, hl=False)]
    n = b.new_num()
    for pt in CHANGE_POINTS:
        els.append(b.para(pt.format(**values), style="ListParagraph", num=(n, 0), size=9.5, hl=False))
    els.append(b.para("The first decision BDOI must take is still where access requests are raised and approved "
                      "(Annex Q, C01): BIBS offers each option as a setting (Annex M.10); section 1.8 lists the ten "
                      "decisions.", size=9.5, hl=False))
    els.append(b.para("Shading used in this copy: ", size=9.5, hl=False))
    els[-1].append(b.run("text added or corrected in version 1.3", size=9.5, hl=True))
    brk = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ppr.append(copy.deepcopy(b.doc.sections[0]._sectPr))
    brk.append(ppr)
    els.append(brk)
    first = b.body[0]
    for e in els:
        first.addprevious(e)


def bibs_case_count() -> int:
    data = yaml.safe_load((UA / "brd11_cases.yaml").read_text(encoding="utf-8"))
    return sum(len(e.get("cases") or []) for e in data["frs"].values())


def key_numbers(b: Builder, nfr) -> list[list[str]]:
    s = b.stats
    items = s["bdoi_items"] + s["new_subitems"] + s["new_item_subitems"]
    return [
        ["BRD requirement IDs mapped (each to at least one FR item)", str(s["brd_ids"]), "Section 2; Annex P"],
        ["Non-functional rows of the BRD", str(len(nfr)), "Annex X"],
        ["FR entries (BDOI's 8 and FRUM.009 to FRUM.013)", str(s["bdoi_entries"] + s["new_items"]), "Section 3"],
        ["FR items and sub-items (BDOI's and added)", str(items), "Section 3"],
        ["Acceptance criteria", str(s["ac"]), "Section 3"],
        ["Screens", str(s["screens"]), "Annex J"],
        ["Messages shown to the user", str(s["messages"]), "Annex K"],
        ["Notices and e-mails", str(s["emails"]), "Annex T"],
        ["Documents and outputs", str(s["documents"]), "Annex U"],
        ["Reports / scheduled runs", f"{s['reports']} / {s['schedules']}", "Annex V"],
        ["Integrations", str(s["integrations"]), "Annex W"],
        ["Figures added in version 1.3", str(len(v13.FLOWS)), "1.8; Annexes A, J, S"],
        ["Test cases (BIBS test plan and one per acceptance criterion)", str(bibs_case_count() + s["ac"]),
         "Test Cases and Traceability workbook"],
        ["Open questions not answered in BDOI's FRS / answered and not repeated", f"{s['open_kept']} / {s['answered']}",
         "Annex Z"],
        ["Observations and points for decision", str(s["observations"]), "Annex Q"],
        ["Review checklist points", str(s["checklist"]), "Annex AB"],
    ]


def set_update_fields(b: Builder):
    settings = b.doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    u = OxmlElement("w:updateFields")
    u.set(qn("w:val"), "true")
    later = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
             "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
             "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
             "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
    for child in settings:
        if child.tag.split("}")[1] in later:
            child.addprevious(u)
            return
    settings.append(u)


# ------------------------------------------------------------------------------------------------------- self-checks
def texts_of(el) -> list[str]:
    out = []
    for p in el.iter(qn("w:p")):
        t = el_text(p).strip()
        if t:
            out.append(t)
    return out


def subsequence(old: list[str], new: list[str]) -> bool:
    it = iter(new)
    return all(any(o == n for n in it) for o in old)


def check_bdoi_text(path: Path):
    """BDOI's FRUM IDs and texts are all still present, in order, in the same table cells (rows of the source keep
    their position; added rows are skipped), and every body paragraph of the source is still there, in order."""
    src = docx.Document(str(REPO / DOC["meta"]["source"]))
    new = docx.Document(str(path))
    ids_src = Counter(FRUM_RE.findall("\n".join(texts_of(src.element.body))))
    ids_new = Counter(FRUM_RE.findall("\n".join(texts_of(new.element.body))))
    missing = [i for i, n in ids_src.items() if ids_new[i] < n]
    problems = [f"FRUM ID missing: {missing}"] if missing else []
    # match the tables of the source with those of the new document by their first cell text, in order
    new_tables = list(new.tables)
    j = 0
    for ti, ts in enumerate(src.tables):
        key = texts_of(ts._tbl)[:1]
        while j < len(new_tables) and texts_of(new_tables[j]._tbl)[:1] != key:
            j += 1
        if j >= len(new_tables):
            problems.append(f"table {ti} of the source not found")
            break
        tn = new_tables[j]
        j += 1
        new_rows = [r for r in tn.rows]
        k = 0
        for ri, row in enumerate(ts.rows):
            first = texts_of(row.cells[0]._tc)
            while k < len(new_rows) and not subsequence(first, texts_of(new_rows[k].cells[0]._tc)):
                k += 1
            if k >= len(new_rows):
                problems.append(f"table {ti} row {ri}: not found")
                break
            for ci, cell in enumerate(row.cells):
                old = texts_of(cell._tc)
                if ti == 1:
                    old = [x.replace(DOC["meta"]["version_old"], DOC["meta"]["version_new"])
                           if x == DOC["meta"]["version_old"] else x for x in old]
                it = iter(texts_of(new_rows[k].cells[ci]._tc))
                if not all(any(o == n for n in it) for o in old):
                    problems.append(f"table {ti} row {ri} cell {ci}: BDOI text changed")
            k += 1
    body_src = [p.text.strip() for p in src.paragraphs if p.text.strip()]
    body_new = [p.text.strip() for p in new.paragraphs if p.text.strip()]
    it = iter(body_new)
    lost = [o for o in body_src if not any(o == n for n in it)]
    if lost:
        problems.append(f"body paragraphs changed: {lost[:3]}")
    if problems:
        raise SystemExit("BDOI text check failed:\n  " + "\n  ".join(problems))
    return sum(ids_src.values())


def check_words(b: Builder):
    hits = set()
    for t in b.added_text:
        for m in list(RESTRICTED.finditer(t)) + list(RESTRICTED_CS.finditer(t)):
            hits.add(m.group(0))
        for m in CODE_RE.finditer(t):
            hits.add(m.group(0))
    if hits:
        raise SystemExit(f"restricted words or internal codes in the added text: {sorted(hits)}")


# ------------------------------------------------------------------------------------------------------------ PDF
def to_pdf(src: Path, dest: Path):
    """The clean copy as PDF, with the table of contents updated (LibreOffice)."""
    import uno  # noqa: PLC0415
    from com.sun.star.beans import PropertyValue  # noqa: PLC0415

    def prop(name, value):
        p = PropertyValue()
        p.Name, p.Value = name, value
        return p

    profile = Path(tempfile.mkdtemp(prefix="uam_lo_"))
    pipe = f"uamfrs{int(time.time())}"
    proc = subprocess.Popen(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                             "--invisible", f"--accept=pipe,name={pipe};urp;"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        ctx = None
        resolver = uno.getComponentContext().ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", uno.getComponentContext())
        for _ in range(120):
            try:
                ctx = resolver.resolve(f"uno:pipe,name={pipe};urp;StarOffice.ComponentContext")
                break
            except Exception:  # noqa: BLE001
                time.sleep(1)
        if ctx is None:
            raise RuntimeError("LibreOffice did not start")
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        model = desktop.loadComponentFromURL(uno.systemPathToFileUrl(str(src)), "_blank", 0, (prop("Hidden", True),))
        idx = model.getDocumentIndexes()
        for _ in range(2):
            for i in range(idx.getCount()):
                idx.getByIndex(i).update()
            model.refresh()
        model.storeToURL(uno.systemPathToFileUrl(str(dest)), (prop("FilterName", "writer_pdf_Export"),))
        model.close(True)
        try:
            desktop.terminate()
        except Exception:  # noqa: BLE001
            pass
    finally:
        try:
            proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)


# ----------------------------------------------------------------------------------------------------------- main
def build(highlight: bool, our, lines, trace, nfr, figs: dict) -> Builder:
    b = Builder(highlight)
    unknown = [r for r in b.fr_to_frum if r not in our]
    if unknown:
        raise SystemExit(f"unknown BIBS references: {unknown}")
    nums = v13.figure_numbers()
    where = v13.classify_items()
    b.stats["classified"] = v13.check_classification(where)
    b.stats["emails_checked"] = v13.check_emails()
    edit_cover(b)
    edit_introduction(b)
    mapping = edit_mapping(b, lines, trace)
    edit_functional_requirements(b)
    edit_signoff(b)
    edit_annex_rows(b)
    a = add_annexes(b, mapping, lines, trace, our, nfr, figs, nums, where)
    missing_fr = [f for f in our if f not in a.listed_refs]
    if missing_fr:
        raise SystemExit(f"Annex P lacks {missing_fr}")
    if a.cmp_ids - a.listed_conflicts:
        raise SystemExit(f"Annex Q lacks {sorted(a.cmp_ids - a.listed_conflicts)}")
    b.round13 = True
    add_annex_flows = v13.add_annex_a_figures
    add_annex_flows(b, figs, nums)
    b.key_numbers = key_numbers(b, nfr)
    v13.add_summary(b, figs, nums, b.key_numbers, b.q_rows)
    b.round13 = False
    add_revision_log(b)
    if highlight:
        change_summary(b)
    set_update_fields(b)
    check_words(b)
    b.mapping = mapping
    b.where = where
    return b


def build_review(figs_dir: Path | None = None) -> Builder:
    """The clean copy built in memory (not saved): the numbers, Annex Q rows and lists the workbooks use."""
    our = our_frs()
    lines = brd_lines()
    trace = {t["id"]: t for t in our_trace()}
    work = Path(tempfile.mkdtemp(prefix="uamfig_")) if figs_dir is None else figs_dir
    try:
        figs = brd11_figures.render_all(work)
        b = build(False, our, lines, trace, our_nfr(), figs)
    finally:
        if figs_dir is None:
            shutil.rmtree(work, ignore_errors=True)
    b.lines, b.trace, b.our = lines, trace, our
    return b


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true", help="skip the PDF")
    args = ap.parse_args(argv)
    our = our_frs()
    lines = brd_lines()
    trace = {t["id"]: t for t in our_trace()}
    if [r for r, _ in lines] != list(trace):
        raise SystemExit("the BRD lines of the BRD and of the BIBS FRS traceability differ")
    nfr = our_nfr()
    OUT.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="uamfig_"))
    results = {}
    try:
        figs = brd11_figures.render_all(work)
        for highlight, name in ((False, DOC["meta"]["clean"]), (True, DOC["meta"]["highlighted"])):
            b = build(highlight, our, lines, trace, nfr, figs)
            path = OUT / name
            b.doc.save(str(path))
            n_ids = check_bdoi_text(path)
            results[name] = (b, n_ids)
            print(f"wrote {path.relative_to(REPO)} ({path.stat().st_size // 1024} KB)")
    finally:
        shutil.rmtree(work, ignore_errors=True)
    b, n_ids = results[DOC["meta"]["clean"]]
    if not args.no_pdf:
        pdf = OUT / DOC["meta"]["clean"].replace(".docx", ".pdf")
        to_pdf(OUT / DOC["meta"]["clean"], pdf)
        print(f"wrote {pdf.relative_to(REPO)} ({pdf.stat().st_size // 1024} KB)")
    s = b.stats
    print(f"BRD IDs mapped: {s['brd_ids']} ({s['bdoi_rows']} BDOI rows kept, {s['leaf_rows']} BRD lines added, "
          f"{s['mapping_changes']} BDOI rows given FRUM items)")
    print(f"FRUM: {s['bdoi_entries']} entries and {s['bdoi_items']} items of BDOI kept ({n_ids} FRUM IDs checked); "
          f"{s['elaborated']} items elaborated; {s['new_subitems']} new sub-items; {s['new_items']} new entries "
          f"({s['new_item_subitems']} sub-items); {s['requirements_added']} requirements and {s['ac']} acceptance "
          f"criteria")
    print(f"Annexes: J {s['screens']} screens; K {s['messages']} messages; L {s['notifications']} notifications; "
          f"O {s['walk_steps']} steps; P {s['trace_rows']} rows; Q {s['observations']} observations "
          f"({s['observations_v13']} of v1.3); R {s['glossary']} terms")
    print(f"v1.3: {s['figures_added']} figures; S {s['workflow_rows']} stages; T {s['emails']} notices "
          f"({s['emails_checked']} checked against the platform); U {s['documents']} documents; V {s['reports']} "
          f"reports and {s['schedules']} runs; W {s['integrations']} interfaces; X {s['nfr_rows']} rows; "
          f"Y {s['data_rows']} rows; Z {s['assumptions']} assumptions, {s['dependencies']} dependencies, "
          f"{s['open_kept']} open questions kept, {s['answered']} answered in BDOI's FRS; AB {s['checklist']} points; "
          f"{s['classified']} references classified")
    for row in b.key_numbers:
        print("  ", " | ".join(row))
    return 0


if __name__ == "__main__":
    sys.exit(main())
