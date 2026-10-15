"""User Access interface specification for BDOI IT (BRD-11): SCIM 2.0 provisioning from UIDM-ISC and OpenID Connect
sign-in with Microsoft Entra ID, in BDOI's template.

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd11_interface_spec.py [--no-pdf]

Starts from the User Access Maintenance FRS v1.4 in BDOI's template (BDO ITG template D051): keeps its cover, page set-up,
headers and footers, styles, revision-log table and table of contents, and replaces the body. The content is in
brd11_icd_content.py; the context figure in figures/brd11_icd_context.dot.

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-11_User_Access_Maintenance/:
  BIBS_ICD-BDOI_BRD-11_User_Access_Interface_Specification_v1.0.docx   the specification (and its PDF, not in git)
  BIBS_ICD-BDOI_BRD-11_SCIM_OpenAPI_v1.0.yaml                          OpenAPI 3.0.3 description of the SCIM service
  BIBS_ICD-BDOI_BRD-11_Interface_Test_Pack_v1.0.xlsx                   interface test cases for BDOI IT and UIDM-ISC

Self-checks: the OpenAPI description parses and has every operation of the specification; every test case points to
a sample that exists; the document text holds none of the development wording refused for client documents.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import yaml
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brd_common_final as final  # noqa: E402
import brd11_icd_content as C  # noqa: E402
from brd11_xlsx import ReviewWorkbook  # noqa: E402
from bdoi_xlsx import Column  # noqa: E402

OUT = REPO / "docs" / "deliverables" / "out" / "Programme" / "BDOI_Template_FRS" / "BRD-11_User_Access_Maintenance"
TEMPLATE = OUT / "BIBS_FRS-BDOI_BRD-11_User_Access_Maintenance_v1.4.docx"
FIGURE = HERE / "figures" / "brd11_icd_context.dot"
FONT = "Arial"
HEAD_FILL = "D9E2F3"
CODE_FILL = "F2F2F2"
TABLE_TW = 8640  # twips, as the tables of the template


# ================================================================================================ Word helpers
class Spec:
    def __init__(self, template: Path):
        self.doc = docx.Document(str(template))
        self.figures = 0
        self.tables = 0
        self._strip_body()

    # -------------------------------------------------------------------------------------------- template
    def _strip_body(self) -> None:
        body = self.doc.element.body
        els = list(body)
        rev = next(i for i, e in enumerate(els) if e.tag == qn("w:tbl") and "".join(t.text or "" for t in e.iter(qn("w:t"))).startswith("Date"))
        toc_head = next(i for i, e in enumerate(els) if e.tag == qn("w:p") and "".join(e.itertext()).strip()
                        .startswith("Table of Contents") and i > rev)
        # the TOC field: from the next paragraph to the one whose field ends
        end, depth = toc_head + 1, 0
        for j in range(toc_head + 1, len(els)):
            for fc in els[j].iter(qn("w:fldChar")):
                t = fc.get(qn("w:fldCharType"))
                depth += 1 if t == "begin" else -1 if t == "end" else 0
            if depth == 0:
                end = j
                break
        keep = set(range(0, rev + 1)) | set(range(toc_head, end + 1))
        for i, e in enumerate(els):
            if i not in keep and e.tag != qn("w:sectPr"):
                body.remove(e)
        self.rev_table = els[rev]
        # the TOC heading: no page break of its own (the revision log is short)
        self.toc_heading = els[toc_head]

    def cover(self) -> None:
        body = self.doc.element.body
        els = list(body)
        first = els[0]
        texts = [t for t in first.iter(qn("w:t"))]
        for t in texts:
            if "Brokering System (BIBS)" in (t.text or ""):
                t.text = (" Brokering System (BIBS)\u00a0-\u00a0Interface Specification: User Access (SCIM 2.0 with "
                          "UIDM-ISC, OpenID Connect with Microsoft Entra ID)\u00a0")
            elif t.text == "John Barbachano":
                t.text = C.AUTHOR
        details = next(e for e in els if e.tag == qn("w:tbl")
                       and "Template Reference" in "".join(t.text or "" for t in e.iter(qn("w:t"))))
        values = {"FR Document No.": C.DOC_NO, "Prepared by:": C.AUTHOR, "Drafting support:": C.AUTHOR,
                  "Creation Date:": "10/15/2026", "Revision No.": C.VERSION, "Last Updated by:": C.AUTHOR,
                  "Last Updated Date:": "10/15/2026"}
        for tr in details.iter(qn("w:tr")):
            tcs = tr.findall(qn("w:tc"))
            label = "".join(t.text or "" for t in tcs[0].iter(qn("w:t"))).strip()
            for key, value in values.items():
                if label.replace(" ", "") == key.replace(" ", ""):
                    _set_tc(tcs[-1], value)
            if label.startswith("FR"):
                _set_tc(tcs[0], "Document No.", bold=True)
        # headers of every section
        for s in self.doc.sections:
            for part in (s.header, s.first_page_header, s.even_page_header):
                for p in part.paragraphs:
                    for r in p.runs:
                        if r.text == "Functional Requirements Specification":
                            r.text = C.HEADER_LEFT
                        elif "User Access Maintenance" in r.text:
                            r.text = C.HEADER_RIGHT

    def revision_log(self, rows: list[tuple[str, str, str, str]]) -> None:
        trs = self.rev_table.findall(qn("w:tr"))
        proto = trs[1]
        for tr in trs[1:]:
            self.rev_table.remove(tr)
        for row in rows:
            tr = copy.deepcopy(proto)
            for tc, value in zip(tr.findall(qn("w:tc")), row):
                _set_tc(tc, value)
            self.rev_table.append(tr)

    # -------------------------------------------------------------------------------------------- content
    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def heading(self, text: str, level: int = 1, new_page: bool = False) -> None:
        p = self.doc.add_paragraph(text, style=f"Heading {level}")
        if new_page:
            p.paragraph_format.page_break_before = True

    def para(self, text: str, bold_lead: str | None = None, italic: bool = False) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.64)
        p.paragraph_format.space_after = Pt(6)
        if bold_lead:
            _run(p, bold_lead + " ", bold=True)
        _run(p, text, italic=italic)

    def bullets(self, items: list[str]) -> None:
        for text in items:
            p = self.doc.add_paragraph(style="List Paragraph")
            pf = p.paragraph_format
            pf.left_indent = Cm(1.2)
            pf.first_line_indent = Cm(-0.5)
            pf.space_after = Pt(3)
            _run(p, "\u2022\t" + text)
            tabs = pf.tab_stops
            tabs.add_tab_stop(Cm(1.2))

    def caption(self, text: str, kind: str = "Table") -> None:
        if kind == "Table":
            self.tables += 1
            n = self.tables
        else:
            self.figures += 1
            n = self.figures
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.64)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        _run(p, f"{kind} {n}: {text}", bold=True, size=9)

    def table(self, headers: list[str], rows: list, widths: list[float], caption: str | None = None,
              mono: set[int] | None = None) -> None:
        if caption:
            self.caption(caption)
        mono = mono or set()
        total = sum(widths)
        tw = [int(TABLE_TW * w / total) for w in widths]
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        tbl = t._tbl
        tblPr = tbl.tblPr
        for tag in ("w:tblStyle",):
            for old in tblPr.findall(qn(tag)):
                tblPr.remove(old)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tblPr.append(jc)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        tblPr.append(borders)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        w = tblPr.find(qn("w:tblW"))
        if w is None:
            w = OxmlElement("w:tblW")
            tblPr.append(w)
        w.set(qn("w:w"), str(TABLE_TW))
        w.set(qn("w:type"), "dxa")
        grid = tbl.tblGrid
        for gc, width in zip(grid.findall(qn("w:gridCol")), tw):
            gc.set(qn("w:w"), str(width))
        for r_i, values in enumerate([headers] + [list(r) for r in rows]):
            row = t.rows[r_i]
            if r_i == 0:
                trPr = row._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader")
                trPr.append(h)
            cant = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant)
            for c_i, value in enumerate(values):
                cell = row.cells[c_i]
                tcPr = cell._tc.get_or_add_tcPr()
                tcw = tcPr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW")
                    tcPr.append(tcw)
                tcw.set(qn("w:w"), str(tw[c_i]))
                tcw.set(qn("w:type"), "dxa")
                if r_i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), HEAD_FILL)
                    tcPr.append(shd)
                lines = str(value).split("\n")
                p = cell.paragraphs[0]
                for k, line in enumerate(lines):
                    if k:
                        p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    _run(p, line, bold=r_i == 0, size=8 if c_i in mono and r_i else 9,
                         font="Courier New" if c_i in mono and r_i else FONT)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, title: str, obj, headers: str | None = None) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.64)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        _run(p, title, bold=True, size=9)
        text = obj if isinstance(obj, str) else json.dumps(obj, indent=2, ensure_ascii=False)
        if headers:
            text = headers + ("\n\n" + text if text else "")
        t = self.doc.add_table(rows=1, cols=1)
        tblPr = t._tbl.tblPr
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tblPr.append(jc)
        w = tblPr.find(qn("w:tblW"))
        if w is None:
            w = OxmlElement("w:tblW")
            tblPr.append(w)
        w.set(qn("w:w"), str(TABLE_TW))
        w.set(qn("w:type"), "dxa")
        cell = t.rows[0].cells[0]
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), CODE_FILL)
        tcPr.append(shd)
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        for line in text.split("\n"):
            q = cell.add_paragraph()
            q.paragraph_format.space_before = Pt(0)
            q.paragraph_format.space_after = Pt(0)
            _run(q, line if line else " ", size=7.5, font="Courier New")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, png: Path, caption: str, width_cm: float = 15.0) -> None:
        p = self.doc.add_paragraph()
        p.alignment = 1
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(png), width=Cm(width_cm))
        self.caption(caption, kind="Figure")

    def save(self, path: Path) -> None:
        settings = self.doc.settings.element
        for old in settings.findall(qn("w:updateFields")):
            settings.remove(old)
        u = OxmlElement("w:updateFields")
        u.set(qn("w:val"), "true")
        settings.append(u)
        # drop the parts of the template body that are no longer referenced (pictures of the FRS annexes)
        part = self.doc.part
        xml = part.element.xml
        for rid, rel in list(part.rels.items()):
            if rel.reltype.endswith(("/image", "/oleObject", "/package", "/chart", "/diagramData")) \
                    and f'"{rid}"' not in xml:
                del part.rels[rid]
        self.doc.core_properties.author = "iorta TechNXT"
        self.doc.core_properties.title = C.TITLE
        self.doc.save(str(path))


def _run(p, text: str, bold: bool = False, italic: bool = False, size: float = 10, font: str = FONT):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = font
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), font)
    return r


def _set_tc(tc, text: str, bold: bool | None = None) -> None:
    """Sets the text of a table cell, keeping the format of its first run."""
    ps = tc.findall(qn("w:p"))
    for p in ps[1:]:
        tc.remove(p)
    p = ps[0]
    runs = p.findall(qn("w:r"))
    proto = copy.deepcopy(runs[0]) if runs else None
    for r in runs:
        p.remove(r)
    r = proto if proto is not None else OxmlElement("w:r")
    for child in list(r):
        if child.tag != qn("w:rPr"):
            r.remove(child)
    if bold is not None:
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        if rpr.find(qn("w:b")) is None and bold:
            rpr.append(OxmlElement("w:b"))
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    p.append(r)


# ================================================================================================ document
def render_figure() -> Path:
    png = FIGURE.with_suffix(".png")
    if shutil.which("dot"):
        subprocess.run(["dot", "-Tpng", "-Gdpi=200", str(FIGURE), "-o", str(png)], check=True)
    return png


def write_docx(out_path: Path) -> list[str]:
    s = Spec(TEMPLATE)
    s.cover()
    s.revision_log([
        (C.ISSUE_DATE, C.VERSION, "First issue for BDOI IT review: SCIM 2.0 provisioning of Users and Groups from "
         "UIDM-ISC through the API gateway, and OpenID Connect sign-in with Microsoft Entra ID, as configured in "
         "BIBS. Companion files: OpenAPI description and interface test pack.", C.AUTHOR),
    ])
    s.page_break()
    heads: list[str] = []

    def h1(text, new_page=True):
        heads.append(text)
        s.heading(text, 1, new_page=new_page)

    def h2(text):
        heads.append(text)
        s.heading(text, 2)

    # 1 -------------------------------------------------------------------------------------------- introduction
    h1("Introduction", new_page=False)
    h2("Purpose")
    s.para("This document specifies the two user access interfaces of the BDO Insure Brokering System (BIBS). "
           "UIDM-ISC provisions BIBS users and group-profile membership through a SCIM 2.0 service. Users sign in "
           "to BIBS with Microsoft Entra ID through OpenID Connect.")
    s.para("BDOI IT and the UIDM-ISC team use it to configure the connector, the API gateway and the Entra ID "
           "app registration, and to run the interface tests. It describes the service as it is configured in "
           "BIBS for BDOI.")
    h2("Scope")
    s.para("In scope:", bold_lead=None)
    s.bullets([
        "Users: joiner, mover, status change, leaver and rehire, and the user extract.",
        "Groups: the group-profile extract and the change of the members of a group profile.",
        "Security of the calls: gateway token, scope, size and rate limits.",
        "Event logging in BIBS and the daily reconciliation with the UIDM-ISC directory.",
        "Single sign-on with Microsoft Entra ID (OpenID Connect, authorisation code with PKCE).",
    ])
    s.para("Out of scope:")
    s.bullets([
        "The request and approval steps inside UIDM-ISC, and the account steps inside Entra ID.",
        "The creation and change of group profiles and their permissions. They stay in BIBS under the "
        "group-profile requests of the User Access Maintenance FRS.",
        "Bulk files. Every change travels as a single call.",
    ])
    h2("Intended Audience")
    s.para("BDOI IT (integration, identity and network teams), the UIDM-ISC team, BDOI Information Security, the "
           "BIBS System Administrators and the iorta TechNXT project team.")
    h2("References")
    s.table(["Reference", "Document"], [
        ("FRS", "BIBS FRS-BDOI BRD-11 User Access Maintenance v1.4 (FRUM.001 to FRUM.003, FRUM.013; Annex N)"),
        ("OpenAPI", f"{C.OPENAPI.rsplit('.', 1)[0]} (OpenAPI 3.0.3 description issued with this document)"),
        ("Test pack", f"{C.XLSX.rsplit('.', 1)[0]} (issued with this document)"),
        ("RFC 7643, RFC 7644", "SCIM 2.0 core schema and protocol"),
        ("RFC 6749, RFC 7519, RFC 7636", "OAuth 2.0, JSON Web Token, PKCE"),
        ("OpenID Connect Core 1.0", "Sign-in protocol of Entra ID"),
    ], [4, 11])
    h2("Conventions")
    s.bullets([
        "Paths of the SCIM service are written relative to the base address of Section 4, for example Users or "
        "Groups/{group-id}.",
        "{user-id} is a BIBS user ID or a Windows ID. {group-id} is a group profile code.",
        "Examples use fictitious people and request numbers.",
        "MUST and SHOULD carry their usual meaning for the UIDM-ISC connector.",
    ])

    # 2 -------------------------------------------------------------------------------------------- actors
    h1("Actors and Systems")
    s.table(["Actor or system", "Who or what", "Role in the interfaces"], C.ACTORS, [3.2, 3.6, 8.2],
            caption="Actors and systems")
    png = render_figure()
    if png.exists():
        s.figure(png, "Context of the user access interfaces")

    # 3 -------------------------------------------------------------------------------------------- journey
    h1("End-to-end Journey")
    s.para("The journey below runs from the access request to the periodic certification. Steps 5, 8, 11, 12 and "
           "15 use the SCIM service. Steps 10 and 13 use OpenID Connect.")
    s.table(["No.", "Actor", "System", "What happens", "Interface"], C.JOURNEY, [0.9, 2.8, 2.6, 6.6, 2.4],
            caption="End-to-end journey of a user's access")

    # 4 -------------------------------------------------------------------------------------------- connectivity
    h1("Connectivity")
    s.para("UIDM-ISC never calls BIBS directly. It calls the BDOI API gateway, which checks the client, issues the "
           "token and forwards the call to BIBS over TLS. BIBS accepts the calls of the interface only with a "
           "gateway token.")
    s.table(["Environment", "SCIM base address (at the gateway)", "BIBS address (sign-in)", "Status"],
            C.ENVIRONMENTS, [2.2, 6.2, 4.2, 2.4], caption="Addresses per environment", mono={1, 2})
    s.bullets([
        f"The service path inside BIBS is {C.BASE_PATH}. The public path at the gateway is set by BDOI IT "
        "(open point OP-01).",
        "TLS 1.2 or higher on both legs. Certificates are issued by a CA trusted by both sides.",
        "Content type: application/scim+json (application/json is also accepted). Character set UTF-8.",
        "Only HTTPS port 443 is opened between the gateway and BIBS. BDOI IT gives the network path and "
        "allow-lists (OP-09).",
    ])

    # 5 -------------------------------------------------------------------------------------------- security
    h1("Security")
    h2("Gateway Token")
    s.para("Every call carries a bearer token issued by the gateway to the UIDM-ISC client with the OAuth 2.0 "
           "client credentials grant. BIBS checks it as follows.")
    s.table(["Item", "Rule"], C.SECURITY, [3.4, 11.6], caption="Security rules of the SCIM service")
    s.code("Token claims (decoded, example)", {
        "iss": "https://{gateway-issuer}", "aud": "bibs-scim", "sub": "uidm-isc", "scope": "identity.provision",
        "iat": 1792040400, "nbf": 1792040400, "exp": 1792044000})
    h2("Limits")
    s.bullets([
        "Body: 256 KB at most. A larger body is refused with status 413 before the token is read.",
        "Rate: 300 requests per minute per client address. Above it, status 429 with Retry-After: 60.",
        "Group profiles per user in one call: 50. Operations in one PatchOp: 100.",
    ])

    # 6 -------------------------------------------------------------------------------------------- operations
    h1("Resources and Operations")
    h2("Overview")
    s.table(["Method", "Path", "Purpose", "Success", "BIBS event"],
            [(m, p, purpose, ok, ev) for _, m, p, purpose, ok, ev in C.OPERATIONS], [1.6, 3.4, 6.4, 1.4, 2.6],
            caption="Operations of the SCIM service", mono={1})
    s.para("Every answer with a body has the content type application/scim+json. Filters, sorting, paging, bulk "
           "and the discovery resources of SCIM are not offered: UIDM-ISC reads the full extracts.")
    for key, method, rpath, purpose, ok, _ in C.OPERATIONS:
        d = C.OPERATION_DETAIL[key]
        h2(f"{LABELS[key]} ({method}, {rpath})")
        s.para(purpose + ".")
        s.bullets(d["notes"])
        hdr = "Authorization: Bearer <token>\nContent-Type: application/scim+json"
        if key == "membership":
            hdr += "\nX-UIDM-Request-No: UIDM-2026-004811"
        if d.get("request") is not None:
            s.code("Request", d["request"], headers=hdr)
        for title, extra in d.get("extra", []):
            s.code(f"Request: {title.lower()}", extra)
        status, answer = d["answer"]
        s.code(f"Answer {status}", answer if answer is not None else "(no body)")
        s.para("Refusals: " + "; ".join(d["refusals"]) + ". Section 9 gives the full list.")
    h2("Error Answer")
    s.para("Every refusal has the SCIM error format.")
    s.code("Answer 400 (example)", C.ERR_SOD)
    s.code("Answer 409 (example)", C.ERR_TAKEN)

    # 7 -------------------------------------------------------------------------------------------- mapping
    h1("Attribute Mapping")
    h2("Core User Attributes")
    s.table(["SCIM attribute", "BIBS field", "Mandatory", "Format", "Validation"], C.CORE_MAP,
            [3.0, 2.4, 2.3, 2.5, 4.8], caption="Core User attributes", mono={0})
    h2("Enterprise Extension")
    s.para(f"Schema {C.ENTERPRISE}.")
    s.table(["SCIM attribute", "BIBS field", "Mandatory", "Format", "Validation"], C.ENTERPRISE_MAP,
            [3.0, 2.4, 2.3, 2.5, 4.8], caption="Enterprise extension attributes", mono={0})
    h2("BIBS Extension")
    s.para(f"Schema {C.BIBS}. It carries the BDOI attributes that SCIM has no place for.")
    s.table(["Attribute", "BIBS field", "Mandatory", "Format", "Validation"], C.BIBS_MAP,
            [3.0, 2.4, 2.3, 2.5, 4.8], caption="BIBS extension attributes", mono={0})
    h2("Status Values")
    s.table(["AD status", "Can sign in", "Effect in BIBS"], C.STATUS_VALUES, [3, 2.5, 9.5],
            caption="AD status values")
    s.para("Any attribute not listed in this section is refused (status 400, invalidValue, naming the attribute). "
           "Text values are trimmed; an empty text counts as absent.")

    # 8 -------------------------------------------------------------------------------------------- rules
    h1("Processing Rules")
    h2("UIDM Request Number")
    s.bullets([
        "On a User call, send the request number in the BIBS extension attribute uidmRequestNo.",
        "On a Groups patch, send it in the header X-UIDM-Request-No.",
        "Format: up to 40 letters, digits and . _ / -. Any other value is refused.",
        "BIBS writes it as the reference of every line of the access change log. Without it BIBS uses "
        "IDE-<event no.>, which BDOI cannot trace to a request. BDOI's rule is to always send it.",
    ])
    h2("Repeated Calls")
    s.para("A call sent again, for example after a time-out, is safe. BIBS compares it with the latest event of "
           "the same user and kind. When the source, the UIDM request number, the type and the content are the "
           "same, and the earlier event was applied or had nothing to change, BIBS returns the earlier answer and "
           "records nothing new.")
    s.para("A call that differs, or that follows a refusal or a failure, is processed as new. So a corrected "
           "request is applied, and a request refused earlier is checked again.")
    h2("Group Profiles on a User")
    s.bullets([
        "groups absent on POST or PUT: the group profiles do not change.",
        "groups present: they become the user's group profiles at once, with no approval step in BIBS. UIDM-ISC "
        "holds the approval.",
        "An empty list removes every group profile of the user.",
        "A leaver keeps the group profiles for the record; a rehire gets back the group profiles held before.",
    ])
    h2("All or Nothing")
    s.para("Group profiles are checked before anything is applied: the profile exists, it is active, the "
           "profiles of the user break no separation-of-duties rule and combine no conflicting permissions.")
    s.bullets([
        "On a joiner or mover with groups: when a check fails, neither the user change nor the group change is "
        "applied.",
        "On a Groups patch: every user whose group profiles would change is checked first. When one is refused, "
        "nobody changes.",
        "The refused event is kept in the identity event log and the System Administrators are told.",
    ])
    h2("Order of Events")
    s.para("UIDM-ISC SHOULD send the events of one user one at a time and wait for the answer. Events of "
           "different users may run in parallel.")

    # 9 -------------------------------------------------------------------------------------------- errors
    h1("Error Catalogue")
    s.table(["Status", "scimType", "When", "Detail", "UIDM-ISC action"], C.ERRORS, [1.5, 2.4, 4.6, 3.2, 3.3],
            caption="Errors of the SCIM service")
    s.para("A refused event is recorded with its number. The detail starts with Event <no.>, which the BIBS "
           "System Administrator finds in Identity Synchronisation. Status 401, 403, 413 and 429 are answered "
           "before the request is read and leave no event.")

    # 10 ------------------------------------------------------------------------------------------- logging
    h1("Event Logging and Reconciliation")
    h2("Event Log")
    s.table(["Record", "Content", "Where to see it"], [
        ("Identity event", "Time, source UIDM-ISC, type (joiner, mover, status change, leaver, rehire, "
         "membership), Windows ID, UIDM request number, outcome (applied, no change, refused, failed) and "
         "message", "Identity Synchronisation (System Administrator)"),
        ("Access change log", "Each attribute and group profile changed, old and new value, actor UIDM-ISC, "
         "reference UIDM request number", "User record and the User Access Report"),
        ("Audit trail", "Event applied or refused, user created, deactivated or reactivated", "Audit Trail"),
        ("Notices", "Your access changed (to the user); Identity event to review (to the System "
         "Administrators, in BIBS and by e-mail)", "Notification bell and e-mail"),
    ], [3.0, 8.0, 4.0], caption="What BIBS records")
    s.para("A refused or failed event can be processed again or dismissed with a remark by a System "
           "Administrator. The remark is audited.")
    h2("Daily Reconciliation")
    s.bullets([
        "The job UAM_IDENTITY_RECONCILIATION runs daily at 01:40 Philippine time. The parameter "
        "UAM_RECONCILIATION_INTERVAL_DAYS sets how often it compares (1 = daily).",
        "It reads every BIBS user with a Windows ID and the matching account in the UIDM-ISC directory.",
        "Mismatches: NOT_IN_DIRECTORY (unknown to the directory), STATUS (active in one, not in the other), "
        "GROUP_PROFILES (different group profiles, when the directory names them).",
        "BIBS changes nothing. The holders of the user-management permission are told (notice "
        "UAM_RECONCILIATION_MISMATCH).",
        "The correction comes from UIDM-ISC as a normal event. The System Administrator then marks the mismatch "
        "resolved with a remark.",
        "A System Administrator may also start the run at once (Run Now).",
    ])

    # 11 ------------------------------------------------------------------------------------------- OIDC
    h1("Single Sign-on with Microsoft Entra ID")
    h2("Flow")
    s.bullets([
        "The user opens BIBS and chooses single sign-on. BIBS sends the browser to Entra ID with the "
        "authorisation code flow, PKCE (S256), a state and a nonce. The request is valid 10 minutes.",
        "Entra ID signs the user in with BDO's policies (password, multi-factor, conditional access) and "
        "returns a code to the BIBS redirect URI.",
        "BIBS exchanges the code for the ID token at the token address of Entra ID, as a confidential client.",
        "BIBS checks the ID token: RS256 signature from the Entra ID key set, issuer, audience equal to the "
        "client ID, validity with 2 minutes of clock skew, and the nonce.",
        "BIBS links the identity to an existing BIBS user through the user name claim. It never creates a user. "
        "The user must be active and not locked.",
        "BIBS opens the session through a one-time ticket valid 2 minutes. The sign-in is audited and logged.",
    ])
    h2("What BDOI IT Provides")
    s.table(["Item", "Value or rule", "Owner"], C.OIDC_ITEMS, [3.2, 8.4, 3.4],
            caption="Entra ID app registration items")
    s.table(["Environment", "BIBS address", "Redirect URI path (after the address)"],
            [(e, b, C.REDIRECT_PATH) for e, _, b, _ in C.ENVIRONMENTS], [2.4, 5.6, 7.0],
            caption="Redirect URIs to register", mono={1, 2})
    h2("Claims Used")
    s.table(["Claim", "Meaning", "Use in BIBS"], C.CLAIMS, [3.2, 6.4, 5.4], caption="ID token claims")
    s.para("Access in BIBS comes only from the group profiles UIDM-ISC gives through the SCIM service. Entra ID "
           "groups and app roles are not read for authorisation. BIBS offers an optional group check at sign-in; "
           "it stays off for BDOI.")
    h2("Leavers and Break-glass Accounts")
    s.bullets([
        "A leaver or a disabled user is deactivated in BIBS by the SCIM call. A later sign-in is refused even "
        "when Entra ID still issues a token.",
        "A user unknown to BIBS is refused. No account is created at sign-in.",
        "Named break-glass System Administrators keep a local password with a second factor, for use when Entra "
        "ID is unavailable (OP-10).",
    ])

    # 12 ------------------------------------------------------------------------------------------- NFR
    h1("Non-functional Requirements")
    s.table(["Area", "Requirement"], C.NFR, [3.0, 12.0], caption="Non-functional requirements")

    # 13 ------------------------------------------------------------------------------------------- tests
    h1("Test Pack Summary")
    groups: dict[str, int] = {}
    for tc in C.TEST_CASES:
        groups[tc["group"]] = groups.get(tc["group"], 0) + 1
    s.para(f"The interface test pack holds {len(C.TEST_CASES)} test cases. Each gives the preconditions, the "
           "request (method, path and sample body), the expected status and answer, and the expected record in "
           "BIBS (identity event, access change log, audit trail and notice). Result and comment columns have "
           "drop-down lists.")
    ids: dict[str, list[str]] = {}
    for tc in C.TEST_CASES:
        ids.setdefault(tc["group"], []).append(tc["id"])
    s.table(["Group", "Cases", "Test case IDs"],
            [(g, str(n), f"{ids[g][0]} to {ids[g][-1]}" if n > 1 else ids[g][0]) for g, n in groups.items()],
            [4, 2, 9], caption="Test cases by group")
    s.para("Entry: the SIT addresses, the gateway client and the Entra ID registration are in place (OP-01 to "
           "OP-04). Exit: every case passed or its failure accepted by BDOI IT.")

    # 14 ------------------------------------------------------------------------------------------- open points
    h1("Open Points for BDOI IT")
    s.table(["No.", "Point", "Owner", "Needed by"], C.OPEN_POINTS, [1.4, 8.2, 3.0, 2.4],
            caption="Open points")

    # 15 ------------------------------------------------------------------------------------------- glossary
    h1("Glossary")
    s.table(["Term", "Meaning"], C.GLOSSARY, [4, 11], caption="Glossary")

    s.save(out_path)
    return heads


LABELS = {"create": "Joiner", "list": "User Extract", "read": "One User", "replace": "Mover",
          "patch": "Status Change and Rehire", "delete": "Leaver", "groups": "Group Profile Extract",
          "group": "One Group Profile", "membership": "Members of a Group Profile"}


# ================================================================================================ OpenAPI
def openapi() -> dict:
    ref = lambda name: {"$ref": f"#/components/schemas/{name}"}  # noqa: E731
    err = lambda desc: {"description": desc, "content": {"application/scim+json": {  # noqa: E731
        "schema": ref("Error")}}}

    def body(schema, obj, extra=None):
        ex = {"default": {"value": obj}}
        for i, (title, o) in enumerate(extra or [], start=1):
            ex[f"example{i}"] = {"summary": title, "value": o}
        return {"required": True, "content": {"application/scim+json": {"schema": ref(schema), "examples": ex}}}

    def ok(desc, schema, obj):
        return {"description": desc, "content": {"application/scim+json": {"schema": ref(schema), "example": obj}}}

    common = {
        "401": {"$ref": "#/components/responses/Unauthorized"},
        "403": {"$ref": "#/components/responses/Forbidden"},
        "413": {"$ref": "#/components/responses/TooLarge"},
        "429": {"$ref": "#/components/responses/TooMany"},
        "500": {"$ref": "#/components/responses/Failed"},
    }
    uid = {"name": "id", "in": "path", "required": True, "description": "BIBS user ID or Windows ID",
           "schema": {"type": "string", "pattern": "^[A-Za-z0-9._\\\\@-]{1,50}$"}}
    gid = {"name": "id", "in": "path", "required": True, "description": "Group profile code",
           "schema": {"type": "string", "pattern": "^[A-Za-z0-9._\\\\@-]{1,50}$"}}
    d = C.OPERATION_DETAIL
    paths = {
        "/Users": {
            "get": {"operationId": "listUsers", "tags": ["Users"], "summary": "User extract",
                    "description": " ".join(d["list"]["notes"]),
                    "responses": {"200": ok("Every user with a Windows ID", "UserList", d["list"]["answer"][1]),
                                  **common}},
            "post": {"operationId": "createUser", "tags": ["Users"], "summary": "Joiner",
                     "description": " ".join(d["create"]["notes"]),
                     "requestBody": body("User", C.JOINER),
                     "responses": {"201": ok("User created", "User", d["create"]["answer"][1]),
                                   "400": err("Invalid attribute or refused group profile (invalidValue)"),
                                   "404": err("A group profile of groups does not exist"),
                                   "409": err("User ID or Windows ID taken (uniqueness)"),
                                   "415": err("Unsupported content type"), **common}},
        },
        "/Users/{id}": {
            "parameters": [uid],
            "get": {"operationId": "getUser", "tags": ["Users"], "summary": "One user",
                    "responses": {"200": ok("The user", "User", d["read"]["answer"][1]),
                                  "404": err("No user with that ID"), **common}},
            "put": {"operationId": "replaceUser", "tags": ["Users"], "summary": "Mover",
                    "description": " ".join(d["replace"]["notes"]),
                    "requestBody": body("User", C.MOVER),
                    "responses": {"200": ok("User updated", "User", d["replace"]["answer"][1]),
                                  "400": err("Invalid attribute, unknown user or refused group profile"),
                                  "404": err("A group profile of groups does not exist"),
                                  "409": err("Windows ID taken (uniqueness)"), **common}},
            "patch": {"operationId": "patchUser", "tags": ["Users"], "summary": "Status change or rehire",
                      "description": " ".join(d["patch"]["notes"]),
                      "requestBody": body("PatchOp", C.STATUS_PATCH, d["patch"]["extra"]),
                      "responses": {"200": ok("User with the new status", "User", d["patch"]["answer"][1]),
                                    "400": err("No status operation, unknown AD status or unknown user"),
                                    **common}},
            "delete": {"operationId": "deleteUser", "tags": ["Users"], "summary": "Leaver",
                       "description": " ".join(d["delete"]["notes"]),
                       "responses": {"204": {"description": "User deactivated, sessions ended"},
                                     "400": err("Unknown user"), **common}},
        },
        "/Groups": {
            "get": {"operationId": "listGroups", "tags": ["Groups"], "summary": "Group profile extract",
                    "description": " ".join(d["groups"]["notes"]),
                    "responses": {"200": ok("Every group profile", "GroupList", d["groups"]["answer"][1]),
                                  **common}},
        },
        "/Groups/{id}": {
            "parameters": [gid],
            "get": {"operationId": "getGroup", "tags": ["Groups"], "summary": "One group profile",
                    "responses": {"200": ok("The group profile", "Group", C.GROUP_ANSWER),
                                  "404": err("No group profile with that code"), **common}},
            "patch": {"operationId": "patchGroupMembers", "tags": ["Groups"], "summary": "Change the members",
                      "description": " ".join(d["membership"]["notes"]),
                      "parameters": [{"name": "X-UIDM-Request-No", "in": "header", "required": False,
                                      "description": "UIDM request number of the change",
                                      "schema": {"type": "string", "pattern": "^[A-Za-z0-9._/-]{1,40}$"}}],
                      "requestBody": body("PatchOp", C.GROUP_ADD, d["membership"]["extra"]),
                      "responses": {"200": ok("The group profile after the change", "Group", C.GROUP_ANSWER),
                                    "400": err("invalidPath, invalidSyntax or invalidValue (refused membership, "
                                               "unknown member)"),
                                    "404": err("No group profile with that code"), **common}},
        },
    }
    text = lambda n: {"type": "string", "maxLength": n}  # noqa: E731
    schemas = {
        "User": {
            "type": "object", "additionalProperties": False, "required": ["schemas"],
            "description": "SCIM User with the enterprise and BIBS extensions. Attributes outside this list are "
                           "refused.",
            "properties": {
                "schemas": {"type": "array", "items": {"type": "string"}},
                "id": {"type": "string", "readOnly": True, "description": "BIBS user ID"},
                "externalId": {"type": "string", "pattern": "^[A-Za-z0-9._\\\\@-]{1,50}$",
                               "description": "Windows ID"},
                "userName": {"type": "string", "pattern": "^[A-Za-z0-9._-]{1,50}$", "description": "BIBS user ID"},
                "name": {"type": "object", "properties": {"givenName": text(60), "familyName": text(60)}},
                "displayName": text(120),
                "emails": {"type": "array", "items": {"type": "object", "properties": {
                    "value": {"type": "string", "pattern": "^[^@\\s]{1,64}@[^@\\s]{1,120}$"},
                    "primary": {"type": "boolean"}}}},
                "active": {"type": "boolean"},
                "groups": {"type": "array", "maxItems": 50, "items": ref("Reference"),
                           "description": "Absent: unchanged. Present: replaces the group profiles."},
                "meta": {"type": "object"},
                C.ENTERPRISE: ref("EnterpriseUser"),
                C.BIBS: ref("BibsUser"),
            },
        },
        "EnterpriseUser": {"type": "object", "properties": {
            "department": text(120), "division": text(120),
            "manager": {"type": "object", "properties": {"displayName": text(120)}}}},
        "BibsUser": {"type": "object", "properties": {
            "windowsId": {"type": "string", "pattern": "^[A-Za-z0-9._\\\\@-]{1,50}$"},
            "adGroup": text(120),
            "adStatus": {"type": "string", "enum": [v[0] for v in C.STATUS_VALUES]},
            "teamLeaderName": text(120), "teamHeadName": text(120), "sectionHeadName": text(120),
            "unitHeadName": text(120), "unitSegment": text(120), "location": text(120),
            "uidmRequestNo": {"type": "string", "pattern": "^[A-Za-z0-9._/-]{1,40}$"},
            "groupProfiles": {"type": "array", "items": {"type": "string"},
                              "description": "Returned on every answer; read on input only with the setting "
                                             "UAM_PROVISIONING_ROLES"}}},
        "Reference": {"type": "object", "required": ["value"], "properties": {
            "value": {"type": "string", "pattern": "^[A-Za-z0-9_.-]{1,40}$"}}},
        "Group": {"type": "object", "properties": {
            "schemas": {"type": "array", "items": {"type": "string"}},
            "id": {"type": "string", "description": "Group profile code"},
            "displayName": {"type": "string"}, "active": {"type": "boolean"},
            "members": {"type": "array", "items": {"type": "object", "properties": {"value": {"type": "string"}}}},
            "permissions": {"type": "array", "items": {"type": "string"}}}},
        "PatchOp": {"type": "object", "required": ["schemas", "Operations"], "properties": {
            "schemas": {"type": "array", "items": {"type": "string", "enum": [C.PATCHOP]}},
            "Operations": {"type": "array", "minItems": 1, "maxItems": 100, "items": {
                "type": "object", "required": ["op"], "properties": {
                    "op": {"type": "string", "description": "add, remove or replace (any case)"},
                    "path": {"type": "string"}, "value": {}}}}}},
        "UserList": {"type": "object", "properties": {
            "schemas": {"type": "array", "items": {"type": "string"}}, "totalResults": {"type": "integer"},
            "itemsPerPage": {"type": "integer"}, "startIndex": {"type": "integer"},
            "Resources": {"type": "array", "items": ref("User")}}},
        "GroupList": {"type": "object", "properties": {
            "schemas": {"type": "array", "items": {"type": "string"}}, "totalResults": {"type": "integer"},
            "itemsPerPage": {"type": "integer"}, "startIndex": {"type": "integer"},
            "Resources": {"type": "array", "items": ref("Group")}}},
        "Error": {"type": "object", "required": ["schemas", "status", "detail"], "properties": {
            "schemas": {"type": "array", "items": {"type": "string", "enum": [C.ERROR]}},
            "status": {"type": "string", "description": "HTTP status as text"},
            "scimType": {"type": "string", "enum": ["invalidValue", "invalidSyntax", "invalidPath", "uniqueness"]},
            "detail": {"type": "string"}}, "example": C.ERR_SOD},
    }
    responses = {
        "Unauthorized": {"description": "No token, or a token with a wrong issuer, audience, signature or validity",
                         "headers": {"WWW-Authenticate": {"schema": {"type": "string"}}}},
        "Forbidden": {"description": "Token without the scope identity.provision"},
        "TooLarge": err("Body larger than 256 KB"),
        "TooMany": {"description": "More than 300 requests in one minute from one client address",
                    "headers": {"Retry-After": {"schema": {"type": "integer", "example": 60}}},
                    "content": {"application/scim+json": {"schema": ref("Error")}}},
        "Failed": err("Unexpected failure; the detail carries a reference only"),
    }
    return {
        "openapi": "3.0.3",
        "info": {
            "title": "BIBS User Provisioning (SCIM 2.0)",
            "version": C.VERSION,
            "description": (
                "SCIM 2.0 service of the BDO Insure Brokering System (BIBS) for UIDM-ISC, reached through the BDOI "
                "API gateway. Joiner (POST /Users), mover (PUT), status change and rehire (PATCH), leaver "
                "(DELETE), user and group profile extracts (GET), and the members of a group profile "
                "(PATCH /Groups/{id}). Companion of the User Access Interface Specification v" + C.VERSION + ". "
                "Path inside BIBS: /" + C.BASE_PATH + ". Sign-in of the users is by OpenID Connect with "
                "Microsoft Entra ID; the redirect URI to register is <BIBS address>/" + C.REDIRECT_PATH + "."),
            "contact": {"name": "iorta TechNXT Project Team"},
        },
        "servers": [{"url": "https://{gateway}/" + C.BASE_PATH, "description": "BDOI API gateway (per environment)",
                     "variables": {"gateway": {"default": "gateway-host",
                                               "description": "Gateway host of SIT, UAT or production"}}}],
        "security": [{"gatewayToken": ["identity.provision"]}],
        "tags": [{"name": "Users"}, {"name": "Groups"}],
        "paths": paths,
        "components": {
            "securitySchemes": {"gatewayToken": {
                "type": "oauth2",
                "description": "Client credentials token of the BDOI API gateway: signed JWT (RS256) with the "
                               "configured issuer and audience and the scope identity.provision; 60 seconds of "
                               "clock skew.",
                "flows": {"clientCredentials": {"tokenUrl": "https://{gateway}/oauth/token",
                                                "scopes": {"identity.provision": "Provision BIBS users and group "
                                                                                 "membership"}}}}},
            "schemas": schemas,
            "responses": responses,
        },
    }


def write_openapi(path: Path) -> int:
    spec = openapi()
    text = "# BIBS User Provisioning (SCIM 2.0) - OpenAPI 3.0.3. Issued by iorta TechNXT with the User Access\n" \
           "# Interface Specification v" + C.VERSION + " for BDOI IT and the UIDM-ISC team.\n" + \
           yaml.safe_dump(spec, sort_keys=False, allow_unicode=True, width=110)
    path.write_text(text, encoding="utf-8")
    return validate_openapi(path)


def validate_openapi(path: Path) -> int:
    spec = yaml.safe_load(path.read_text(encoding="utf-8"))
    try:
        from openapi_spec_validator import validate  # noqa: PLC0415

        validate(spec)
    except ImportError:
        assert spec["openapi"] == "3.0.3" and spec["info"]["title"] and spec["paths"], "OpenAPI structure"

        def refs(node):
            if isinstance(node, dict):
                for k, v in node.items():
                    if k == "$ref":
                        yield v
                    else:
                        yield from refs(v)
            elif isinstance(node, list):
                for v in node:
                    yield from refs(v)
        for r in refs(spec):
            parts = r.lstrip("#/").split("/")
            node = spec
            for p in parts:
                node = node[p]
    ops = [(m.upper(), p) for p, item in spec["paths"].items() for m in item if m != "parameters"]
    want = {(m, "/" + p.replace("{user-id}", "{id}").replace("{group-id}", "{id}")) for _, m, p, *_ in C.OPERATIONS}
    if set(ops) != want:
        raise SystemExit(f"OpenAPI operations differ from the specification: {sorted(set(ops) ^ want)}")
    return len(ops)


# ================================================================================================ test pack
def write_xlsx(path: Path) -> None:
    wb = ReviewWorkbook("User Access Interface Test Pack", doc_type="Interface test pack", brd="BRD-11",
                        version=C.VERSION, date="15 October 2026",
                        subtitle="SCIM 2.0 with UIDM-ISC and OpenID Connect with Microsoft Entra ID")
    wb.cover_notes = [
        "Run the cases in SIT in the order given; later cases use the users created by earlier ones.",
        "Request bodies are on the sheet Samples. Paths are relative to the SCIM base address of the environment.",
        "Record the result, the date, the tester and the UIDM request number used. Attach evidence in the comments.",
    ]
    missing = [tc["id"] for tc in C.TEST_CASES if re.match(r"S-\d\d", tc["body"])
               and tc["body"][:4] not in C.SAMPLES]
    if missing:
        raise SystemExit(f"test cases with an unknown sample: {missing}")
    cols = [
        Column("id", "Test ID", 9, "Test case identifier"),
        Column("group", "Scenario group", 14, "Joiner, mover, status change, leaver, rehire, membership, repeat, "
               "refusal, limits, security, extract, sign-in, reconciliation"),
        Column("scenario", "Scenario", 30, "What the case proves"),
        Column("pre", "Preconditions", 28, "State needed before the run"),
        Column("method", "Method", 9, "Call method, or Browser / Job"),
        Column("path", "Path", 22, "Path relative to the SCIM base address, or the screen"),
        Column("body", "Body sample", 18, "Sample of the sheet Samples, with the change stated"),
        Column("status", "Expected status", 10, "Expected answer status"),
        Column("answer", "Expected answer", 34, "Expected answer content"),
        Column("event", "Expected identity event", 28, "Expected record in the identity event log"),
        Column("acl", "Expected access change log", 30, "Expected lines of the access change log"),
        Column("audit", "Expected audit and notice", 28, "Expected audit trail entry and notices"),
        Column("uidm", "UIDM request no. used", 16, "Request number sent in the run"),
        Column("tester", "Tester", 14, "Name of the tester"),
        Column("date", "Date run", 12, "Date of the run"),
        Column("result", "Result", 12, "Outcome of the run", values=["NOT RUN", "PASS", "FAIL", "BLOCKED", "N/A"],
               status=True),
        Column("severity", "Severity if failed", 13, "Severity of a failure",
               values=["Critical", "High", "Medium", "Low"]),
        Column("owner", "Fix owner", 13, "Who acts on a failure",
               values=["BDOI IT", "UIDM-ISC team", "iorta TechNXT", "Gateway team"]),
        Column("comments", "Comments", 34, "Evidence and remarks"),
    ]
    rows = [dict(tc, result="NOT RUN") for tc in C.TEST_CASES]
    wb.sheet("Test Cases", cols, rows, description="Interface test cases for BDOI IT and the UIDM-ISC team")
    sample_cols = [Column("ref", "Sample", 9, "Sample reference"), Column("title", "Content", 28, "What it is"),
                   Column("body", "Request body", 100, "Body to send (SCIM JSON)")]
    wb.sheet("Samples", sample_cols,
             [{"ref": k, "title": t, "body": json.dumps(o, indent=2, ensure_ascii=False)}
              for k, (t, o) in C.SAMPLES.items()],
             description="Request bodies of the test cases (the examples of the specification)")
    err_cols = [Column("status", "Status", 9, "Answer status"), Column("type", "scimType", 14, "SCIM error type"),
                Column("when", "When", 50, "Cause"), Column("detail", "Detail", 36, "Text of the detail"),
                Column("action", "UIDM-ISC action", 40, "What UIDM-ISC does")]
    wb.sheet("Error Catalogue", err_cols, [list(e) for e in C.ERRORS],
             description="Errors of the SCIM service (Section 9 of the specification)")
    wb.save(path)


# ================================================================================================ checks and main
def _check_pack():
    import check_pack  # noqa: PLC0415

    return check_pack


def check_text(path: Path) -> None:
    d = docx.Document(str(path))
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                texts += [p.text for p in c.paragraphs]
    cp = _check_pack()
    hits = sorted({m.group(0) for t in texts for _, pat in cp.BUILD_STATUS for m in pat.finditer(t)}
                  | {m.group(0) for t in texts for m in cp.RESTRICTED.finditer(t)})
    if hits:
        raise SystemExit(f"{path.name}: wording refused in client documents: {hits}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / C.DOCX
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp) / C.DOCX
        heads = write_docx(work)
        n = final.populate_toc(work)
        shutil.copy(work, target)
    check_text(target)
    print(f"wrote {target.relative_to(REPO)} ({len(heads)} headings, {n} TOC lines)")
    ops = write_openapi(OUT / C.OPENAPI)
    print(f"wrote {C.OPENAPI} ({ops} operations, valid)")
    write_xlsx(OUT / C.XLSX)
    print(f"wrote {C.XLSX} ({len(C.TEST_CASES)} test cases)")
    if not args.no_pdf:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(target)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        print(f"wrote {target.with_suffix('.pdf').name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
