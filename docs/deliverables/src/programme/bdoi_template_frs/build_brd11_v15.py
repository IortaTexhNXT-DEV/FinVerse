"""User Access Maintenance document set v1.5 (BRD-11) in BDOI's template.

    python docs/deliverables/src/programme/bdoi_template_frs/build_brd11_v15.py [--no-pdf] [--only FRS,BPS,UXD,RNO]

Starts each document from BDOI's FRS template (BDO ITG template D051, templates/bdoi_d051_frs_template.docx): cover,
document details table, headers and footers, styles, revision log and table of contents are kept; the body is
written from brd11_v15_content.py (FRS) and brd11_v15_pages.py (process, storyboards, screens, outputs). The
screenshots are in figures/brd11_v15_screens (flows/brd11/shots_v15.sh), one full page each.

Outputs in docs/deliverables/out/Programme/BDOI_Template_FRS/BRD-11_User_Access_Maintenance/:
  BIBS_FRS-BDOI_BRD-11_User_Access_Maintenance_v1.5.docx
  BIBS_BPS-BDOI_BRD-11_User_Access_Process_and_Storyboards_v1.5.docx
  BIBS_UXD-BDOI_BRD-11_User_Access_UX_Design_v1.5.docx
  BIBS_RNO-BDOI_BRD-11_User_Access_Reports_Notifications_Outputs_v1.5.docx
  and their PDFs (not kept in git).

Self-checks: every screenshot named exists; every BRD group of the mapping points to FRUM items of section 3; no
wording refused in client documents.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brd_common_final as final  # noqa: E402
import brd11_v15_content as C  # noqa: E402
import brd11_v15_pages as P  # noqa: E402
import build_brd11_interface_spec as icd  # noqa: E402
from colour_code import COLOURS, LABELS  # noqa: E402

OUT = icd.OUT
TEMPLATE = HERE / "templates" / "bdoi_d051_frs_template.docx"
SHOTS = HERE / "figures" / "brd11_v15_screens"
JOURNEY = HERE / "figures" / "brd11_v15_journey.dot"
COL = {"k": None, "g": COLOURS["green"], "o": COLOURS["orange"], "b": COLOURS["blue"], "": None}
COLOUR_NAME = {"k": "Black", "g": "Green", "o": "Orange", "b": "Blue", "": "-"}


class Doc(icd.Spec):
    """A document of the set: BDOI's template with a new body."""

    def __init__(self, key: str):
        super().__init__(TEMPLATE)
        self.key = key
        self.meta = C.DOCS[key]
        self.shots = 0

    def set_cover(self) -> None:
        m = self.meta
        body = self.doc.element.body
        for t in list(body)[0].iter(qn("w:t")):
            if "Brokering System (BIBS)" in (t.text or ""):
                t.text = f" Brokering System (BIBS) - {m['cover']} "
            elif t.text in ("John Barbachano", icd.C.AUTHOR):
                t.text = C.AUTHOR
        details = next(e for e in body if e.tag == qn("w:tbl")
                       and "Template Reference" in "".join(t.text or "" for t in e.iter(qn("w:t"))))
        values = {"Document No.": m["doc_no"], "FR Document No.": m["doc_no"], "Prepared by:": C.AUTHOR,
                  "Drafting support:": C.AUTHOR, "Revision No.": C.VERSION, "Last Updated by:": C.AUTHOR,
                  "Last Updated Date:": C.DATE_US, "Creation Date:": "08/26/2026" if self.key == "FRS" else C.DATE_US}
        for tr in details.iter(qn("w:tr")):
            tcs = tr.findall(qn("w:tc"))
            label = "".join(t.text or "" for t in tcs[0].iter(qn("w:t"))).strip()
            for k, v in values.items():
                if label.replace(" ", "") == k.replace(" ", ""):
                    icd._set_tc(tcs[-1], v)
        for s in self.doc.sections:
            for part in (s.header, s.first_page_header, s.even_page_header):
                for p in part.paragraphs:
                    for r in p.runs:
                        if r.text in (icd.C.HEADER_LEFT, "Functional Requirements Specification"):
                            r.text = m["header_left"]
                        elif "User Access" in r.text and "BIBS" in r.text:
                            r.text = C.HEADER_RIGHT

    # ------------------------------------------------------------------ text helpers
    def h1(self, text: str, new_page: bool = True) -> None:
        self.heading(text, 1, new_page=new_page)

    def h2(self, text: str) -> None:
        self.heading(text, 2)

    def cpara(self, text: str, colour: str = "k", label: str = "", lead: str | None = None) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.64)
        p.paragraph_format.space_after = Pt(4)
        if lead:
            icd._run(p, lead + " ", bold=True)
        r = icd._run(p, text)
        if COL.get(colour):
            _colour(r, COL[colour])
        if label:
            g = icd._run(p, "  " + LABELS["integration" if label == "int" else "migration"], bold=True, size=8)
            _colour(g, COLOURS["grey"])

    def ctable(self, headers, rows, widths, caption=None, colours=None, colour_cols=None):
        """Table whose rows can be coloured (colours: one code per row; colour_cols: which columns)."""
        self.table(headers, rows, widths, caption=caption)
        tbl = self.doc.tables[-1]
        if colours:
            for i, code in enumerate(colours):
                if not COL.get(code):
                    continue
                for c_i, cell in enumerate(tbl.rows[i + 1].cells):
                    if colour_cols is None or c_i in colour_cols:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                _colour(r, COL[code])
        return tbl

    def screen(self, shot: str, caption: str) -> None:
        """One full-page screenshot on its own page, scaled to fit the page, never split."""
        png = SHOTS / f"{shot}.png"
        if not png.exists():
            raise SystemExit(f"missing screenshot {png.name}")
        w, h = Image.open(png).size
        max_w, max_h = 16.0, 21.5
        width = min(max_w, max_h * w / h)
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        p.alignment = 1
        p.add_run().add_picture(str(png), width=Cm(width))
        self.caption(caption, kind="Figure")
        self.shots += 1

    def legend(self) -> None:
        self.para(C.LEGEND_NOTE)
        rows = [["Black", "In the BRD", "Asked by the BRD, BRD-00 or a programme NFR.", "In scope"],
                ["Green", "Not in the BRD: OOTB", "Offered by the iNXT BrokerVerse OOTB platform.", "No cost"],
                ["Orange", "Not in the BRD: change request", "Not offered by the OOTB platform.", "Only if approved"],
                ["Blue", "Not in the BRD: best practice", "Recommendation, separate table.", "Indicative effort"],
                ["Grey label", LABELS["integration"] + " / " + LABELS["migration"],
                 "UIDM-ISC, Microsoft Entra ID, or Data Migration (BRD-13).", "Not a change request"]]
        self.ctable(["Colour", "Meaning", "What it covers", "Cost"], rows, [1.2, 2.4, 4.2, 1.6],
                    caption="Colour code", colours=["k", "g", "o", "b", ""], colour_cols={0, 1})
        tbl = self.doc.tables[-1]
        for r in tbl.rows[5].cells[0].paragraphs[0].runs + tbl.rows[5].cells[1].paragraphs[0].runs:
            _colour(r, COLOURS["grey"])


def _colour(run, hexcol: str) -> None:
    rpr = run._r.get_or_add_rPr()
    for c in rpr.findall(qn("w:color")):
        rpr.remove(c)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), hexcol)
    rpr.append(c)


def _render(dot: Path, name: str, tmp: Path) -> Path:
    png = tmp / f"{name}.png"
    subprocess.run(["dot", "-Tpng", "-Gdpi=170", str(dot), "-o", str(png)], check=True)
    return png


def _lifecycle(tmp: Path) -> Path:
    dot = tmp / "lifecycle.dot"
    dot.write_text(
        'digraph g { rankdir=LR; graph [fontname="Arial", nodesep=0.3]; node [shape=box, style="rounded", '
        'fontname="Arial", fontsize=10]; edge [fontname="Arial", fontsize=8];\n'
        'Draft -> Pending [label="Submit"]; Pending -> "Pending Second" [label="High/Admin or\\nout of hours"];\n'
        'Pending -> Returned [label="Return"]; Returned -> Pending [label="Resubmit"];\n'
        'Pending -> "For Implementation" [label="Approve"]; "Pending Second" -> "For Implementation" '
        '[label="Second approval"];\n "For Implementation" -> Implemented [label="Implement"];\n'
        'Pending -> Approved [label="Approve\\n(apply at approval)", style=dashed];\n'
        'Pending -> Rejected [label="Reject"]; "Pending Second" -> Rejected; Pending -> Cancelled [label="Withdraw"];'
        ' Returned -> Cancelled; Draft -> Cancelled; }', encoding="utf-8")
    return _render(dot, "lifecycle", tmp)


# ================================================================================================ FRS
def write_frs(path: Path, tmp: Path) -> Doc:
    d = Doc("FRS")
    d.set_cover()
    d.revision_log(C.REVISIONS)
    d.page_break()
    i = C.FRS_INTRO
    d.h1("Introduction")
    d.h2("Document Purpose")
    d.para(i["purpose"])
    d.h2("Intended Audience")
    d.para(i["audience"])
    d.h2("Project Background")
    for t in i["background"]:
        d.para(t)
    d.h2("Scope")
    d.table(["In scope (BIBS)", "Out of scope"],
            [[a, b] for a, b in zip(i["in_scope"], i["out_scope"] + [""] * (len(i["in_scope"]) - len(i["out_scope"])))],
            [4.6, 4.6], caption="Scope of User Access Maintenance in BIBS")
    d.h2("Document Set")
    d.table(["Document", "Content"], C.DOC_SET, [3.2, 6.0], caption="User Access Maintenance document set")
    d.h2("How to Read This Document")
    d.legend()
    d.h2("References")
    d.bullets(C.REFERENCES)

    d.h1("Business Requirements Mapping")
    d.para("Every BRD line of User Access Maintenance maps to an item of section 3. The column Where says where the "
           "function is performed in the agreed model; the last column gives the change from version 1.4.")
    d.table(["BRD Req No.", "Business requirement", "FRUM items", "Where", "Change from v1.4"], C.MAPPING,
            [1.6, 2.6, 1.8, 1.7, 1.6], caption="Business requirements mapping")

    d.h1("Functional Requirements")
    d.para("Each item keeps BDOI's FRUM number. Status gives the change from version 1.4: Kept, Changed, Withdrawn "
           "or New. A withdrawn item stays listed so that the numbering and the traceability hold.")
    d.table(["FRUM", "Component", "Capability", "Status", "BRD"],
            [[f["id"], f["component"], f["capability"], f["status"], f["brd"]] for f in C.FRUM],
            [1.3, 1.8, 3.0, 1.1, 2.1], caption="Functional requirements at a glance")
    for f in C.FRUM:
        d.h2(f"{f['id']} {f['capability']}")
        d.para(f"Status: {f['status']}. Version 1.4: {f['was']}. BRD: {f['brd']}.", italic=True)
        rows, cols = [], []
        for rid, text, colour, label, rule, ex in f["reqs"]:
            req = text + ("  " + LABELS["integration"] if label == "int" else "")
            rows.append([rid, req, rule, ex])
            cols.append(colour)
        d.ctable(["ID", "Requirement", "Rule", "Example"], rows, [1.8, 3.5, 2.5, 2.0], colours=cols,
                 colour_cols={1, 2})

    d.h1("Business Rules")
    d.ctable(["ID", "Rule", "Colour"], [[a, b, COLOUR_NAME[c]] for a, b, c in C.BUSINESS_RULES], [1.3, 6.8, 1.2],
             caption="Business rules", colours=[c for *_, c in C.BUSINESS_RULES], colour_cols={1, 2})
    d.h1("Validations and Messages")
    d.table(["Screen or step", "Condition", "Message"], C.MESSAGES, [2.0, 3.3, 3.9], caption="Validations and "
            "messages")
    d.h1("Personas and Permissions")
    d.table(["Persona", "Does", "Screens", "Walkthrough account"], C.PERSONAS, [1.8, 3.6, 2.6, 1.6],
            caption="Personas of User Access")
    d.para(f"{C.SUPPORT_ROLE[0]}: {C.SUPPORT_ROLE[1]}")
    d.table(["Function", "Admin", "Maker", "Checker", "InfoSec", "Auditor"], C.PERMISSIONS,
            [4.0, 1.0, 1.0, 1.0, 1.0, 1.0], caption="What each persona may do (Y = yes)")
    d.h1("Parameters and Lists of Values")
    d.ctable(["Setting", "Default", "Range", "Area"], [r[:4] for r in C.PARAMETERS], [3.6, 2.4, 1.8, 1.6],
             caption="UAM settings (each change approved by another user)", colours=[r[4] for r in C.PARAMETERS])
    d.table(["List", "Values"], C.LOVS, [2.6, 6.6], caption="Lists of values")
    d.h1("Statuses")
    d.figure(_lifecycle(tmp), "Statuses of a group-profile change", width_cm=16.0)
    for name, rows in C.STATUSES.items():
        d.h2(name)
        d.table(["Status", "Meaning", "Next"], rows, [2.0, 4.0, 3.2], caption=f"Statuses: {name.lower()}")
    d.h1("Assumptions, Dependencies and Open Questions")
    d.table(["ID", "Assumption"], C.ASSUMPTIONS, [1.3, 7.9], caption="Assumptions")
    d.table(["ID", "Dependency"], C.DEPENDENCIES, [1.3, 7.9], caption="Dependencies")
    d.table(["ID", "Question", "Owner", "Answer by"], C.OPEN_QUESTIONS, [1.2, 5.0, 1.8, 1.2],
            caption="Open questions")
    d.h1("Change Control Summary")
    d.para("Change requests of the register v1.2. Effort in man-days (analysis and design, backend, frontend, "
           "testing, documentation and UAT support).")
    d.ctable(["CR", "Title", "Status", "FRS", "Man-days", "Note"],
             [[a, b, c, e, f"{sum(eff):g}", n] for a, b, c, e, eff, n in C.CRS], [1.3, 2.3, 1.1, 1.2, 0.9, 2.6],
             caption="Change requests", colours=["o" if c == "Open" else "" for _, _, c, *__ in C.CRS])
    d.ctable(["No.", "Not in the BRD, offered by the OOTB platform", "FRS"], C.OOTB_POINTS, [1.3, 5.6, 2.3],
             caption="Green points (no cost)", colours=["g"] * len(C.OOTB_POINTS), colour_cols={1})
    d.table(["No.", "Integration and data migration scope", "FRS"], C.INTEGRATION_POINTS, [1.3, 5.6, 2.3],
            caption="Grey label points (not change requests)")
    d.ctable(["No.", "Best-practice recommendation", "Priority", "Man-days"],
             [[a, b, c, str(e)] for a, b, c, e in C.BEST_PRACTICE], [1.3, 5.0, 1.4, 1.4],
             caption="Blue points (recommendations)", colours=["b"] * len(C.BEST_PRACTICE), colour_cols={1})
    d.h1("Changes from Version 1.4")
    ch = C.CHANGES_V14
    d.table(["Removed", "Added"], _pair(ch["removed_screens"], ch["added_screens"]), [4.6, 4.6],
            caption="Screens")
    d.table(["Removed", "Added"], _pair(ch["removed_rules"], ch["added_rules"]), [4.6, 4.6], caption="Rules")
    d.table(["Removed", "Added"], _pair(ch["removed_notices"], ch["added_notices"]), [4.6, 4.6],
            caption="Notices")
    d.save(path)
    return d


def _pair(a, b):
    n = max(len(a), len(b))
    return [[(a + [""] * n)[k], (b + [""] * n)[k]] for k in range(n)]


# ================================================================================================ BPS
def write_bps(path: Path, tmp: Path) -> Doc:
    d = Doc("BPS")
    d.set_cover()
    d.revision_log(C.REVISION_COMPANION)
    d.page_break()
    d.h1("Introduction")
    d.para("This document shows how User Access works end to end once UIDM-ISC, Microsoft Entra ID and BIBS are "
           "connected. The requirements are in the FRS v1.5; the screens in the UX Design v1.5; the interfaces in "
           "the Interface Specification v1.0.")
    d.para("Each storyboard step is illustrated by a full-page screenshot taken on a test environment. The people "
           "and records are fictitious. Events from UIDM-ISC were sent through the provisioning test tool of that "
           "environment, which feeds the same processing as the SCIM interface.")
    d.table(["Persona", "Does", "Screens", "Walkthrough account"], C.PERSONAS, [1.8, 3.6, 2.6, 1.6],
            caption="Personas")
    d.h1("End-to-end Journey")
    d.para("Fifteen steps from the access request in UIDM-ISC to the evidence for the auditor.")
    d.figure(_render(JOURNEY, "journey", tmp), "User access journey across UIDM-ISC, Microsoft Entra ID and BIBS",
             width_cm=13.5)
    d.table(["Step", "System", "What happens"], P.JOURNEY, [0.9, 1.6, 6.7], caption="The 15 steps")
    d.h1("To-be Processes")
    for p in P.PROCESSES:
        d.h2(f"{p['id']} {p['name']}")
        d.table(["Item", "Value"], [["Journey steps", p["steps"]], ["Trigger", p["trigger"]], ["Actors", p["actors"]],
                                    ["Outcome", p["outcome"]], ["Controls", p["controls"]]], [2.0, 7.2])
    d.h2("Statuses of a group-profile change")
    d.figure(_lifecycle(tmp), "Statuses of a group-profile change", width_cm=16.0)
    d.h1("Storyboards")
    for sb in P.STORYBOARDS:
        d.h2(f"{sb['id']} {sb['title']}")
        d.para(sb["summary"])
        d.table(["Step", "Who", "Screen", "Does", "Sees"],
                [[str(k + 1), w, s, a, r] for k, (w, s, a, r, _) in enumerate(sb["steps"])],
                [0.7, 1.7, 1.9, 2.6, 2.6], caption=f"{sb['id']} steps")
        for k, (w, s, a, r, shot) in enumerate(sb["steps"]):
            d.screen(shot, f"{sb['id']} step {k + 1}: {w}, {s}")
    d.save(path)
    return d


# ================================================================================================ UXD
def write_uxd(path: Path, tmp: Path) -> Doc:
    d = Doc("UXD")
    d.set_cover()
    d.revision_log(C.REVISION_COMPANION)
    d.page_break()
    d.h1("Introduction")
    d.para("This document describes every User Access screen of BIBS: purpose, fields, validations, actions, states "
           "and empty or error states. Each screen state is shown by one full-page screenshot on its own page, taken "
           "on a test environment with fictitious data. The rules behind the screens are in the FRS v1.5.")
    d.h1("Menu by Persona")
    d.para("All screens are under Setup & Administration › User Access. A persona sees only its entries.")
    d.table(["Screen", "Purpose", "Personas"],
            [[a, b, c.replace("A", "Administrator").replace("M", "Maker").replace("C", "Checker")
              .replace("I", "InfoSec").replace("U", "Auditor").replace(" ", ", ")] for a, b, c in C.MENU],
            [2.2, 4.2, 2.8], caption="User Access menu by persona")
    d.h1("Screens")
    for s in P.SCREENS:
        d.h2(f"{s['id']} {s['name']}")
        d.table(["Item", "Value"], [["Menu", s["menu"]], ["Personas", s["personas"]], ["Purpose", s["purpose"]],
                                    ["Actions", ", ".join(s["actions"])], ["States", s["states"]],
                                    ["Empty or error state", s["empty"]]], [2.2, 7.0])
        d.table(["Field", "Description"], s["fields"], [2.6, 6.6], caption=f"{s['id']} fields")
        if s["validations"] != ["-"]:
            d.table(["Validation message"], [[v] for v in s["validations"]], [9.2], caption=f"{s['id']} validations")
        for shot, cap in s["shots"]:
            d.screen(shot, f"{s['id']} {cap}")
    d.save(path)
    return d


# ================================================================================================ RNO
def write_rno(path: Path, tmp: Path) -> Doc:
    d = Doc("RNO")
    d.set_cover()
    d.revision_log(C.REVISION_COMPANION)
    d.page_break()
    d.h1("Introduction")
    d.para("This document lists the reports, the certification extract, the exports, the notices and e-mails and "
           "the scheduled jobs of User Access in BIBS. Only the current set is listed; those withdrawn with version "
           "1.5 are named at the end.")
    d.h1("Reports")
    d.table(["ID", "Report", "BRD", "Parameters", "Fields"], P.REPORTS, [1.2, 1.8, 1.3, 2.0, 3.0],
            caption="User access reports")
    d.para("Each report is exported to PDF (landscape), Excel and CSV. The file is named after the report and the "
           "date: User Access Report_10162026.")
    d.screen("ux-26-reports", "Reports & Certification")
    d.h1("Certification Extract")
    d.cpara("The extract lists every account with each group profile and each permission, one row per permission. "
            "It is taken in CSV or XLSX from Reports & Certification. Each extract is logged with who took it, when, "
            "the number of rows and the SHA-256 of the file, and Information Security is told.", "g")
    d.table(["Field"], [[f] for f in P.EXTRACT_FIELDS], [9.2], caption="Fields of the certification extract")
    d.screen("wt5-02-extract", "Taking the certification extract")
    d.h1("Exports")
    d.table(["Output", "Format", "Where", "Content"], P.EXPORTS, [2.0, 1.8, 2.8, 2.6], caption="Exports")
    d.h1("Notices and E-mails")
    d.para("Notices appear under the bell and, where shown, by e-mail. Text in angle brackets is filled in. Notices "
           "of user access requests (to approve, returned, decided) are sent by UIDM-ISC.")
    d.table(["ID", "Event", "To", "Title", "Text", "Channel"], P.NOTICES, [0.8, 1.7, 1.5, 2.0, 2.4, 1.2],
            caption="Notices and e-mails")
    d.h1("Scheduled Jobs")
    d.table(["Job", "When", "What it does", "Notices"], P.JOBS, [2.4, 1.8, 4.0, 1.0], caption="Jobs")
    d.h1("Withdrawn with Version 1.5")
    d.table(["v1.4 notice", "Subject", "Why"], P.REMOVED_NOTICES, [2.0, 3.4, 3.8], caption="Notices withdrawn")
    d.save(path)
    return d


WRITERS = {"FRS": write_frs, "BPS": write_bps, "UXD": write_uxd, "RNO": write_rno}


def check_mapping() -> None:
    ids = {f["id"] for f in C.FRUM}
    for row in C.MAPPING:
        refs = [t.strip(" ;,") for t in row[2].replace(";", ",").split(",")]
        if not any(r.split()[0] in ids or r.split()[0].rsplit(".", 1)[0] in ids for r in refs if r):
            raise SystemExit(f"mapping row {row[0]} points to no FRUM item")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--only", default="FRS,BPS,UXD,RNO")
    args = ap.parse_args()
    check_mapping()
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        for key in args.only.split(","):
            name = C.DOCS[key]["file"]
            work = tmp / name
            d = WRITERS[key](work, tmp)
            n = final.populate_toc(work)
            target = OUT / name
            shutil.copy(work, target)
            icd.check_text(target)
            print(f"wrote {name} ({n} TOC lines, {d.shots} screenshots)")
            if not args.no_pdf:
                subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(target)],
                               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return 0


if __name__ == "__main__":
    sys.exit(main())
