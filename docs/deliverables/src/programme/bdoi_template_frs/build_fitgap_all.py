"""Fit-Gap workbooks of the ten BRDs without a BDOI-template FRS pack, and the summary of all fourteen BRDs.

    python docs/deliverables/src/programme/bdoi_template_frs/build_fitgap_all.py

Writes in docs/deliverables/out/Programme/FitGap/:
  BIBS_FitGap_BRD-xx_<Name>_BRD_vs_OOTB_vs_Best_Practice_v1.0.xlsx   for BRD-00, 01, 04, 05, 07, 08, 09, 10, 12, 13
  BIBS_FitGap_BRD-00_Summary_All_BRDs_v1.0.xlsx                      counts and fitment of all fourteen BRDs

Method (same as the Fit-Gap workbooks of BRD-02, 03, 06 and 11):
  rows      every BRD requirement line of the programme RTM (Matrix sheet) for the BRD;
  OOTB      the iNXT BrokerVerse OOTB platform (BDOI-OOTB, branch brokerverse-platform) through the Feature List vs
            OOTB workbook, which assesses each OOTB feature against the BRD IDs and FRs it cites;
  classes   Fit (OOTB as delivered), Configure (OOTB through settings, masters or templates), Extend (new or changed
            functions delivered by BIBS), Recommend change (BRD wording or practice to restate); grey rows
            (Integration scope, Data Migration scope) and excluded rows (Removed from the BRD, Specified in another
            FRS) are outside the fitment; best-practice rows (blue) are recommendations only.
  fitment   (Fit + Configure) / in-scope rows, in-scope = Fit + Configure + Extend + Recommend change.

BRD-02, 03, 06 and 11 are classified with the same method into sheets "Programme method BRD-xx" of the summary;
their in-pack Fit-Gap workbooks are not changed and keep the pack's delivery view.
"""

from __future__ import annotations

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[4] / "tools" / "deliverables"))
import pack_common as pc  # noqa: E402
from brd11_xlsx import Column, ReviewWorkbook  # noqa: E402

PROG = pc.OUT_PROGRAMME
OUT = PROG / "FitGap"
RTM = PROG / "Requirements" / "BIBS_RTM_BRD-00_Requirements_Traceability_Matrix_v1.2.xlsx"
FEATURES = PROG / "Feature_List" / "BIBS_BDOI_FeatureList_vs_OOTB_and_Best_Practice_v1.0.xlsx"
V = "1.0"
DATE = "16 October 2026"

NEW = ["BRD-00", "BRD-01", "BRD-04", "BRD-05", "BRD-07", "BRD-08", "BRD-09", "BRD-10", "BRD-12", "BRD-13"]
PACK = ["BRD-02", "BRD-03", "BRD-06", "BRD-11"]   # programme method computed into the summary only
EXISTING = {"BRD-02": "BRD-02_Operations_Cashiering", "BRD-03": "BRD-03_Product_Maintenance",
            "BRD-06": "BRD-06_Renewal", "BRD-11": "BRD-11_User_Access_Maintenance"}
NAMES = {"BRD-00": "Core Replacement", "BRD-01": "New Business", "BRD-02": "Operations Cashiering",
         "BRD-03": "Product Maintenance", "BRD-04": "Collections", "BRD-05": "Accounting Disbursement ACSL",
         "BRD-06": "Renewal", "BRD-07": "Claims", "BRD-08": "Employee Benefits",
         "BRD-09": "Customer Servicing Facility", "BRD-10": "Sanction Screening", "BRD-11": "User Access Maintenance",
         "BRD-12": "Submitted Policies", "BRD-13": "Data Migration"}

IN_SCOPE = ["Fit", "Configure", "Extend", "Recommend change"]
GREY = ["Integration scope", "Data Migration scope"]
EXCLUDED = ["Removed from the BRD", "Specified in another FRS"]
BLUE = "Best practice"
CLASSES = IN_SCOPE + GREY + EXCLUDED + [BLUE]
COLOUR = {"Fit": "Black (BRD)", "Configure": "Black (BRD)", "Extend": "Black (BRD)",
          "Recommend change": "Black (BRD)", "Integration scope": "Grey (Integration)",
          "Data Migration scope": "Grey (Data Migration)", "Removed from the BRD": "-",
          "Specified in another FRS": "-", BLUE: "Blue (best practice)"}
FONT_COLOUR = {"Grey (Integration)": "616161", "Grey (Data Migration)": "616161", "Blue (best practice)": "1565C0",
               "-": "9E9E9E"}

INTEGRATION = re.compile(
    r"\b(interface|MFT|API|CCM|LMS|ESB|SAP|Finacle|core banking|host system|Entra|EIAM|UIDM|SSO|single sign-on|"
    r"LDAP|Active Directory|SMS gateway|e-?mail gateway|SFTP|file transfer|web service|external system|feed from|"
    r"feed to|eFPS|eBIRForms|World-Check|Dow Jones|ALeA|LAMD|BDOFC|BDOSOLD|SOLD|OBPCS|CLPC|payment channels?|"
    r"bank (channel|file|system|portal|statement|upload|feed)s?|insurer (system|portal|file|folder)s?|"
    r"BIR\W+(\w+\W+){0,3}(submi|upload|transmi|e-?fil|portal|eFPS)\w*|"
    r"post-dated checks? (from|of) PMS|product master transfer)\b", re.I)
# External interfaces of the four Interface Specifications (ICD) of BRD-02, 03, 06 and 11; internal flows between
# BIBS modules (IF-CSH-05 to 08) are not integration. BRD-11 has no IF- row; its identity systems are named above.
ICD_EXTERNAL = {"IF-CSH-01": "OBPCS Bills Payment", "IF-CSH-02": "PMS post-dated checks",
                "IF-CSH-03": "Bank channel files", "IF-CSH-04": "Insurer commission schedule",
                "IF-PM-01": "Product master transfer", "IF-PM-02": "E-mail channel to insurers",
                "IF-RN-01": "CCM", "IF-RN-02": "MFT outbound", "IF-RN-03": "MFT inbound",
                "IF-RN-04": "Insurer files", "IF-RN-05": "CLPC billing to LMS", "IF-RN-06": "LAMD and BDOFC / BDOSOLD"}
ICD_REF = re.compile(r"IF-(CSH|PM|RN|UA)-\d\d")
MIGRATION = re.compile(r"\bmigrat", re.I)
CONFIG = re.compile(r"\b(parameter|set-?up|setting|configur|maintain|maintenance|master|template|table|rate|limit|"
                    r"matrix|threshold|code list|lookup|schedule|calendar|approval level|authority)\w*", re.I)
PRACTICE = re.compile(r"\b(manual(ly)?|by hand|spreadsheet|fax|photocop\w*|wet signature|hard ?cop(y|ies)|"
                      r"print(ed)? and sign\w*)\b", re.I)
ID_TOKEN = re.compile(r"[A-Z]{2,5}[-.]?[A-Z]{0,3}[-.]?\d{2,3}(?:\.\d+)?")
DM_IN_BIBS = {"4", "5", "6", "7", "8", "9", "10"}   # BRD-13 sections processed in BIBS after the load

# Best-practice points per BRD: Feature List IDs (Best-Practice Opportunities sheet) and points of this workbook.
BEST = {
    "BRD-00": ["INT-022", "INT-027", "NFR-023", "CMP-037"],
    "BRD-01": ["CRM-027", "CRM-028"],
    "BRD-04": ["CLX-026", "BIL-027"],
    "BRD-05": ["DIS-026", "CMP-037"],
    "BRD-07": ["CLM-036", "CLM-037"],
    "BRD-08": ["EB-033", "EB-034", "EB-035"],
    "BRD-09": ["CSF-024", "CSF-025"],
    "BRD-10": ["CMP-024", "CMP-037"],
    "BRD-02": [],
    "BRD-03": [],
    "BRD-06": [],
    "BRD-11": [],
    "BRD-12": [("BP-SP-01", "Expiry watch on client-provided policies: the account officer is told before a "
                "submitted policy expires so that the client is asked for the renewed policy in time.",
                "Lenders and brokers track the insurance cover of collateral continuously."),
               ("BP-SP-02", "Automated check of the submitted policy against the lender's cover requirements "
                "(sum insured, mortgagee clause, perils) with a list of exceptions for follow-up.",
                "Collateral insurance is checked against the loan terms, not only recorded.")],
    "BRD-13": [("BP-DM-01", "Parallel-run comparison of the key reports (aging, production, GL trial balance) "
                "between the legacy system and BIBS before the cut-over sign-off.",
                "A parallel run with report comparison is the usual cut-over gate."),
               ("BP-DM-02", "Data quality scorecard per object (completeness, validity, duplicates) published to "
                "the data owners at each mock load.", "Data owners sign off on measured quality, not on samples.")],
}


def load_features():
    wb = openpyxl.load_workbook(FEATURES, read_only=True)
    feats = [r for r in wb["Feature Comparison"].iter_rows(min_row=5, values_only=True) if r[0]]
    idx = defaultdict(list)
    for f in feats:
        for tok in set(ID_TOKEN.findall(str(f[6] or ""))):
            idx[tok].append(f)
    return {f[0]: f for f in feats}, idx


def load_matrix():
    ws = openpyxl.load_workbook(RTM, read_only=True)["Matrix"]
    return [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]


def _join(xs, n=4):
    xs = list(dict.fromkeys(x for x in xs if x))
    return ", ".join(xs[:n]) + (f" and {len(xs) - n} more" if len(xs) > n else "")


def classify(brd, r, idx):
    """Returns (class, OOTB features, note) for one RTM row."""
    rid, text, frs, cov, screens = str(r[2]), str(r[4] or ""), str(r[5] or ""), r[16], r[14]
    if cov == "REMOVED" or (cov == "OUT" and not text.strip()) or (cov == "OUT" and "removed" in text.lower()):
        return "Removed from the BRD", "", "Removed from the BRD by BDOI; no delivery."
    if cov in ("OUT", "REPLACED") or (cov == "CLARIFICATION" and re.search(r"covered by (BRD|Re)", text)):
        where = re.search(r"covered by ([^;]+)", text)
        return "Specified in another FRS", "", (f"Specified in {where.group(1).strip()}." if where
                                                 else "Specified in another FRS of the programme.")
    keys = [rid] + re.findall(r"FR-[A-Z]+-\d+", frs)
    feats = list({f[0]: f for k in keys for f in idx.get(k, [])}.values())
    fids = _join(f[0] for f in feats)
    built = f" BIBS screens: {screens}." if screens else ""
    if brd == "BRD-13":
        sec = rid.replace("BRID_MIG", "").strip().split(".")[0]
        if sec not in DM_IN_BIBS:
            return "Data Migration scope", fids, "Extraction, load, reconciliation or sign-off of the migration; " \
                "delivered by the Data Migration scope." + built
    elif MIGRATION.search(text):
        return "Data Migration scope", fids, "One-time or recurring load of legacy data; Data Migration scope."
    icd = [m.group(0) for m in ICD_REF.finditer(f"{text} {frs} {r[20] or ''}") if m.group(0) in ICD_EXTERNAL]
    if icd:
        return "Integration scope", fids, f"Interface {icd[0]} ({ICD_EXTERNAL[icd[0]]}) of the Interface " \
            "Specification; Integration scope." + built
    if INTEGRATION.search(text) or INTEGRATION.search(str(r[6] or "")) or (feats and all(f[0].startswith("INT-") for f in feats)):
        return "Integration scope", fids, "Depends on an interface with another BDO or third-party system; " \
            "Integration scope." + built
    if cov in ("CLARIFICATION", "FR ONLY") or PRACTICE.search(text):
        return "Recommend change", fids, ("The BRD line is open-ended or describes a manual practice; we recommend "
                                          "restating it as a system rule before sign-off." + built)
    ranks = {"Fit": 0, "OOTB extra": 0, "Partial fit": 1, "Gap": 2, "BDOI-specific": 2}
    worst = max((ranks.get(f[13], 1) for f in feats), default=None)
    notes = "; ".join(dict.fromkeys(str(f[10]) for f in feats if f[10]))[:400]
    cfg = bool(CONFIG.search(text) or CONFIG.search(notes))
    if worst == 0:
        cls = "Configure" if CONFIG.search(text) else "Fit"
    elif worst == 1:
        cls = "Configure" if cfg else "Extend"
    elif worst == 2:
        cls = "Extend"
    else:
        cls = "Configure" if CONFIG.search(text) else "Extend"
        notes = "No OOTB feature cites this line; classified from what BIBS delivers."
    return cls, fids, (notes + built).strip()


def best_rows(brd, fmap):
    rows = []
    for item in BEST[brd]:
        if isinstance(item, str):
            f = fmap[item]
            rows.append({"id": f[0], "fr": "-", "req": f"{f[3]}: {f[4]}", "colour": COLOUR[BLUE], "fit": BLUE,
                         "ootb": f"{f[8]}", "note": f[12]})
        else:
            rows.append({"id": item[0], "fr": "-", "req": item[1], "colour": COLOUR[BLUE], "fit": BLUE,
                         "ootb": "Not available", "note": item[2]})
    return rows


def counts(classes):
    c = Counter(classes)
    scope = sum(c[k] for k in IN_SCOPE)
    pct = round(100 * (c["Fit"] + c["Configure"]) / scope, 1) if scope else None
    return c, scope, pct


def _colour_rows(ws, n, col):
    for r in range(6, 6 + n):
        hexcol = FONT_COLOUR.get(ws.cell(row=r, column=col).value)
        if hexcol:
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(name=cell.font.name, size=cell.font.size, bold=cell.font.bold, color=hexcol)


def file_of(brd):
    return OUT / f"BIBS_FitGap_{brd}_{NAMES[brd].replace(' ', '_')}_BRD_vs_OOTB_vs_Best_Practice_v{V}.xlsx"


def brd_rows(brd, matrix, idx):
    rows = []
    for r in (x for x in matrix if x[0] == brd):
        cls, fids, note = classify(brd, r, idx)
        rows.append({"id": str(r[2]), "fr": r[5] or "-", "req": str(r[4] or "(line removed from the BRD)"),
                     "colour": COLOUR[cls], "fit": cls, "ootb": fids or "-", "note": note})
    return rows


COLS = [Column("id", "BRD requirement", 18), Column("fr", "FR", 22), Column("req", "Requirement", 56),
        Column("colour", "Colour", 18), Column("fit", "Fit", 20, values=CLASSES),
        Column("ootb", "OOTB feature", 22), Column("note", "OOTB note and BIBS delivery", 60)]


def build_brd(brd, matrix, fmap, idx):
    rows = brd_rows(brd, matrix, idx)
    c, scope, pct = counts(x["fit"] for x in rows)
    bp = best_rows(brd, fmap)
    wb = ReviewWorkbook("BRD vs Out-of-the-Box vs Best Practice", doc_type="Fit-gap", brd=brd, version=V, date=DATE,
                        subtitle=f"{NAMES[brd]}; each BRD requirement against the iNXT BrokerVerse OOTB platform")
    cols = COLS
    allrows = rows + bp
    ws = wb.sheet("Fit-Gap", cols, allrows,
                  description=f"Each requirement line of the {brd} {NAMES[brd]} BRD (programme RTM v1.2) against the "
                  "OOTB platform, with the best-practice points in blue at the end.")
    _colour_rows(ws, len(allrows), 4)
    summ = [{"k": k, "n": c[k], "g": g} for k, g in
            [(k, "In scope") for k in IN_SCOPE] + [(k, "Grey, outside the fitment") for k in GREY]
            + [(k, "Excluded") for k in EXCLUDED]]
    summ.append({"k": BLUE, "n": len(bp), "g": "Recommendation only"})
    summ.append({"k": "Fitment % (Fit + Configure) / in scope", "n": pct if pct is not None else "n/a", "g": "-"})
    wb.sheet("Summary", [Column("k", "Class", 40), Column("n", "Rows", 12), Column("g", "Group", 28)], summ,
             description=f"Counts of the Fit-Gap sheet; {scope} rows in scope.")
    wb.sheet("Best Practice", [Column("id", "No.", 12), Column("req", "Recommendation", 60),
                               Column("ootb", "OOTB platform", 16), Column("note", "Why", 50)], bp,
             description="Blue points: in no BRD, FRS or OOTB; recommendation only, not in any change request total.")
    path = pc.save_workbook(wb, file_of(brd))
    return path, c, scope, pct, len(bp)


def pack_method(brd, matrix, fmap, idx):
    """BRD-02, 03, 06, 11 classified with the programme method; the in-pack Fit-Gap workbook is not changed. The
    blue rows are the best-practice points of the in-pack workbook."""
    rows = brd_rows(brd, matrix, idx)
    path = sorted((pc.OUT_BDOI / EXISTING[brd]).glob("*FitGap*.xlsx"))[-1]
    ws = openpyxl.load_workbook(path, read_only=True)["Best Practice"]
    bp = [{"id": r[0], "fr": "-", "req": r[1], "colour": COLOUR[BLUE], "fit": BLUE, "ootb": "Not available",
           "note": f"Best-practice point of {path.name}."} for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    c, scope, pct = counts(x["fit"] for x in rows)
    return (f"Sheet Programme method {brd} of this workbook", c, scope, pct, len(bp)), rows + bp


def build_summary(results, pack_rows):
    wb = ReviewWorkbook("Fit-Gap Summary of All BRDs", doc_type="Fit-gap", brd="BRD-00", version=V, date=DATE,
                        subtitle="Counts and fitment of the fourteen BRDs against the iNXT BrokerVerse OOTB platform")
    rows, tot = [], Counter()
    for brd in sorted(results):
        path, c, scope, pct, bp = results[brd]
        tot.update(c)
        tot["bp"] += bp
        rows.append({"brd": brd, "name": NAMES[brd], "src": Path(path).name,
                     **{k: c[k] for k in CLASSES if k != BLUE}, "bp": bp, "scope": scope,
                     "pct": pct if pct is not None else "n/a"})
    scope = sum(tot[k] for k in IN_SCOPE)
    rows.append({"brd": "Total", "name": "All BRDs", "src": "-", **{k: tot[k] for k in CLASSES if k != BLUE},
                 "bp": tot["bp"], "scope": scope, "pct": round(100 * (tot["Fit"] + tot["Configure"]) / scope, 1)})
    num = lambda k, h: Column(k, h, 11, kind="number")  # noqa: E731
    wb.sheet("Summary", [Column("brd", "BRD", 10), Column("name", "Module", 26)]
             + [num(k, k) for k in IN_SCOPE] + [num("scope", "In scope")]
             + [num(k, k) for k in GREY + EXCLUDED] + [num("bp", "Best practice (blue)"),
             Column("pct", "Fitment %", 11, kind="percent"), Column("src", "Fit-Gap workbook", 60)], rows,
             description="Fitment % = (Fit + Configure) / in scope; grey, excluded and blue rows are outside it.")
    method = [
        ("Rows", "Every BRD requirement line of the programme RTM v1.2, for all 14 BRDs."),
        ("OOTB", "The iNXT BrokerVerse OOTB platform (BDOI-OOTB, branch brokerverse-platform), through the Feature "
         "List vs OOTB and Best Practice workbook v1.0, which assesses each OOTB feature against the BRD IDs and FRs."),
        ("Fit", "Met by the OOTB platform as delivered."),
        ("Configure", "Met by the OOTB platform through settings, masters, templates or rules, without new code."),
        ("Extend", "New or changed functions delivered in BIBS."),
        ("Recommend change", "BRD line open-ended or describing a manual practice; to be restated before sign-off."),
        ("Integration scope", "Grey: maps to an external interface of the Interface Specifications of BRD-02, 03, 06 "
         "or 11, or names an external system (bank, insurer system, BIR filing, LMS, CCM, MFT, UIDM, Entra, EIAM, "
         "SSO, LAMD, BDOFC, SOLD, payment channel); outside the fitment and outside change requests."),
        ("Data Migration scope", "Grey: extraction, load and reconciliation of legacy data (BRD-13)."),
        ("Removed from the BRD", "Excluded: removed by BDOI (or withdrawn change requests)."),
        ("Specified in another FRS", "Excluded: specified and counted in another BRD's FRS (for example BRD-00 "
         "lines covered by a function BRD, or Phase 2 Reinsurance)."),
        ("Best practice", "Blue: in no BRD, FRS or OOTB; recommendation only, outside the fitment."),
        ("Same method", "All 14 rows of the Summary use the same programme method; the in-pack Fit-Gap workbooks of "
         "BRD-02, 03, 06 and 11 keep the pack's delivery view."),
        ("BRD-00", "The umbrella BRD repeats the function BRDs; its fitment is not added to theirs in a business "
         "reading, but it is kept in the total as the RTM keeps it."),
    ]
    wb.sheet("Method", [Column("k", "Term", 26), Column("t", "Meaning", 100)],
             [{"k": a, "t": b} for a, b in method], description="How the classes and the fitment are set.")
    for brd, rows in pack_rows.items():
        ws = wb.sheet(f"Programme method {brd}", COLS, rows,
                      description=f"{brd} {NAMES[brd]} classified with the programme method, for the Summary.")
        _colour_rows(ws, len(rows), 4)
    return pc.save_workbook(wb, OUT / f"BIBS_FitGap_BRD-00_Summary_All_BRDs_v{V}.xlsx")


def main() -> int:
    fmap, idx = load_features()
    matrix = load_matrix()
    results = {}
    for brd in NEW:
        results[brd] = build_brd(brd, matrix, fmap, idx)
    pack_rows = {}
    for brd in PACK:
        results[brd], pack_rows[brd] = pack_method(brd, matrix, fmap, idx)
    s = build_summary(results, pack_rows)
    for brd in sorted(results):
        _, c, scope, pct, bp = results[brd]
        print(brd, {k: c[k] for k in CLASSES if c[k]}, "in scope", scope, "fitment", pct, "bp", bp)
    print("wrote", s.name)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
