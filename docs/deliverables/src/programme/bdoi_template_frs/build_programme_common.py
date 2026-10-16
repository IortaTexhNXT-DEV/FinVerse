"""Programme documents referred to by every BRD pack, in docs/deliverables/out/Programme/Common/:

    python docs/deliverables/src/programme/bdoi_template_frs/build_programme_common.py [--no-pdf] [--only NFR,GLO,LOG,REG]

  BIBS_NFR_BRD-00_Non_Functional_and_Security_Requirements_v2.0.docx   (BDOI template D051; replaces the NFR
                                                                        document and register v1.0 of Requirements/)
  BIBS_Glossary_BRD-00_Programme_Glossary_v1.0.xlsx
  BIBS_Log_BRD-00_Decision_and_Observation_Log_v1.0.xlsx
  BIBS_Register_BRD-00_Document_Register_v1.0.xlsx                      (every current document of BRD-11, BRD-03,
                                                                        BRD-06 and BRD-02, read from out/)

The NFR rows come from requirements_quality/nfr_catalogue.yaml; the rest from programme_common.py. Run the Document
Register last, after the BRD packs are written.
"""

from __future__ import annotations

import argparse
import re
import sys
import tempfile
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[4] / "tools" / "deliverables"))
import pack_common as pc  # noqa: E402
import programme_common as G  # noqa: E402

CATALOGUE = HERE.parent / "requirements_quality" / "nfr_catalogue.yaml"
NAME_RE = re.compile(r"BIBS_(?P<type>[A-Za-z-]+)_(?P<brd>BRD-\d\d)_(?P<name>.+?)_v(?P<ver>\d+(?:\.\d+)*)"
                     r"(?P<extra>_Changes_Highlighted)?\.(?P<ext>docx|xlsx|yaml)$")


def write_nfr(tmp: Path, pdf: bool) -> tuple[Path, int]:
    cat = yaml.safe_load(CATALOGUE.read_text(encoding="utf-8"))
    meta = pc.DocMeta(file=G.FILES["NFR"], cover="Non-Functional and Security Requirements",
                      header_left="Non-Functional and Security Requirements", header_right=G.HEADER_RIGHT,
                      doc_no="BIBS-NFR-BRD-00", version=G.NFR_VERSION, date_us=G.DATE_US, created_us="10/08/2026",
                      revisions=[("08 Oct 2026", "1.0", "First issue of the catalogue and its register.",
                                  pc.PREPARED),
                                 (G.DATE, G.NFR_VERSION, "One document for every BRD pack, with the security "
                                  "requirements, the ASR checklist summary and the measured User Access results. "
                                  "Replaces the NFR document and register v1.0. NFR-SEC-02 and NFR-SEC-03 follow the "
                                  "agreed access model.", pc.PREPARED)])
    d = pc.PackDoc(meta)
    d.h1("Introduction")
    d.para(G.NFR_INTRO)
    d.para(G.NFR_STATUS_NOTE)
    d.table(["ID", "Change in version 2.0"], [list(x) for x in G.NFR_CHANGES], [1.8, 7.4],
            caption="Requirements changed")
    d.h1("Sizing Basis")
    d.kv(G.SIZING, caption="Sizing")
    d.h1("Requirements")
    nfrs = cat["nfrs"]
    for code, name in cat["categories"].items():
        rows = [[n["id"], f"{n['title']}. {n['req']}", n["target"], f"{n['test']}; {n['stage']}", n["status"]]
                for n in nfrs if n["cat"] == code]
        if not rows:
            continue
        d.h2(name)
        d.table(["ID", "Requirement", "Target", "Verified by, when", "Status"], rows, [1.4, 2.4, 2.8, 1.8, 1.0],
                caption=name)
    d.h1("Measured Results: User Access")
    d.para(G.MEASURED_NOTE)
    d.table(["Scenario", "Load", "Requests", "Errors", "p95", "p99", "NFR", "Met"], [list(x) for x in G.MEASURED],
            [1.6, 1.4, 0.9, 0.7, 1.5, 1.3, 1.4, 0.6], caption="Load test of User Access, 16 Oct 2026")
    d.h1("Security")
    d.h2("BDO Application Security Requirements")
    d.para(f"Self-assessment of BIBS against the {G.ASR['source']}. BIBS is an internal back-office application "
           "with no card data, no customer-facing service and no artificial-intelligence component, so many items do not apply.")
    d.table(["Sheet", "Met", "Partly", "Not met", "Not applicable", "Total"],
            [[str(x) for x in r] for r in G.ASR["summary"]], [2.0, 1.2, 1.2, 1.2, 1.6, 1.2],
            caption="ASR checklist summary")
    d.table(["Item", "Requirement", "Status", "Action"], [list(x) for x in G.ASR["open"]], [1.5, 3.2, 1.0, 3.5],
            caption="ASR items not fully met")
    d.h2("OWASP Self-Assessment")
    d.table(["Area", "Result"], [list(x) for x in G.OWASP], [4.0, 5.2],
            caption="OWASP Top 10 2021 and API Top 10 2023")
    d.table(["Open item", "Action"], [list(x) for x in G.SECURITY_OPEN], [5.0, 4.2], caption="Open items")
    out = pc.OUT_COMMON / meta.file
    return out, pc.finish(d, out, tmp, pdf=pdf)


def write_glossary() -> tuple[Path, int]:
    wb = pc.workbook("Programme Glossary", "Glossary", "BRD-00", G.VERSION, G.DATE,
                     "Terms and abbreviations used in every BRD pack")
    cats = sorted({c for _, _, c in G.GLOSSARY})
    wb.sheet("Glossary", [pc.Column("t", "Term", 28), pc.Column("m", "Meaning", 80),
                          pc.Column("c", "Area", 14, values=cats)],
             [{"t": t, "m": m, "c": c} for t, m, c in sorted(G.GLOSSARY, key=lambda x: x[0].lower())],
             description="One line per term; the packs use these meanings.")
    return pc.save_workbook(wb, pc.OUT_COMMON / G.FILES["GLO"]), len(G.GLOSSARY)


def write_log() -> tuple[Path, int]:
    wb = pc.workbook("Decision and Observation Log", "Decision log", "BRD-00", G.VERSION, G.DATE,
                     "Decisions taken and observations raised across the BRD packs")
    keys = ["id", "date", "type", "brd", "subject", "text", "by", "impact", "status", "ref"]
    rows = [dict(zip(keys, r)) for r in G.LOG]
    wb.sheet("Log", [pc.Column("id", "ID", 8), pc.Column("date", "Date", 12), pc.Column("type", "Type", 12,
                     values=["Decision", "Observation"]), pc.Column("brd", "BRD", 14),
                     pc.Column("subject", "Subject", 28), pc.Column("text", "Decision or observation", 56),
                     pc.Column("by", "Made or raised by", 20), pc.Column("impact", "Impact or action", 40),
                     pc.Column("status", "Status", 18), pc.Column("ref", "Reference", 22)], rows,
             description="D-nnn: decision; O-nnn: observation. A pack refers to these IDs.")
    return pc.save_workbook(wb, pc.OUT_COMMON / G.FILES["LOG"]), len(rows)


def _docx_date(p: Path) -> str:
    import docx  # noqa: PLC0415
    d = docx.Document(str(p))
    for t in d.tables[:6]:
        if t.rows and t.rows[0].cells[0].text.strip().startswith("Date") and len(t.rows) > 1:
            return t.rows[-1].cells[0].text.strip()
    return "-"


def _xlsx_date(p: Path) -> str:
    import openpyxl  # noqa: PLC0415
    wb = openpyxl.load_workbook(p, read_only=True)
    for ws in wb.worksheets[:2]:
        for row in ws.iter_rows(max_row=40, values_only=True):
            for v in row:
                if isinstance(v, str):
                    m = re.search(r"\b(\d{1,2}[ -](?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w*[ -]20\d\d)\b", v)
                    if m:
                        return m.group(1)
    return "-"


def _norm(date: str) -> str:
    from datetime import datetime  # noqa: PLC0415
    for fmt in ("%d %B %Y", "%d %b %Y", "%d-%b-%Y", "%d-%B-%Y"):
        try:
            return datetime.strptime(date, fmt).strftime("%d %b %Y")
        except ValueError:
            pass
    return date


TYPE_NAMES = {"FRS-BDOI": "FRS", "BPS-BDOI": "BPS", "UXD-BDOI": "UXD", "RNO-BDOI": "RNO", "ICD-BDOI": "ICD",
              "DSP-BDOI": "DSP", "DDI-BDOI": "DDI", "USR-BDOI": "USR"}


def write_register() -> tuple[Path, int]:
    rows = []
    for brd, folder, name, owner in G.REGISTER_BRDS:
        present = set()
        for p in sorted((pc.OUT_BDOI / folder).glob("BIBS_*")):
            m = NAME_RE.match(p.name)
            if not m:
                continue
            kind = TYPE_NAMES.get(m["type"], m["type"])
            if kind == "Inputs" and "How_to_Read" in m["name"]:
                kind = "README"
            present.add(kind)
            title = m["name"].replace("_", " ") + (" (changes highlighted)" if m["extra"] else "")
            date = _docx_date(p) if p.suffix == ".docx" else _xlsx_date(p) if p.suffix == ".xlsx" else "-"
            rows.append({"brd": brd, "kind": kind, "title": title, "file": p.stem, "fmt": p.suffix[1:].upper(), "ver": m["ver"], "date": date,
                         "owner": owner, "prep": pc.PREPARED, "status": "Issued for BDOI review"})
        for r in rows:
            r["date"] = _norm(r["date"])
            if r["date"] == "-" and r["brd"] == brd and r["ver"] != "-":
                r["date"] = next((x["date"] for x in rows if x["brd"] == brd and x["kind"] == r["kind"]
                                  and x["ver"] == r["ver"] and x["date"] != "-"), "-")
        for kind, title in G.STANDARD_SET:
            if kind not in present and not (kind == "ICD" and brd != "BRD-11"):
                rows.append({"brd": brd, "kind": kind, "title": title, "file": "-", "fmt": "-", "ver": "-", "date": "-",
                             "owner": owner, "prep": pc.PREPARED, "status": "To be issued"})
    for key, fname in G.FILES.items():
        if key == "REG":
            continue
        m = NAME_RE.match(fname)
        rows.append({"brd": "Programme", "kind": m["type"], "title": m["name"].replace("_", " "), "file": fname.rsplit(".", 1)[0],
                     "fmt": fname.rsplit(".", 1)[1].upper(), "ver": m["ver"], "date": G.DATE, "owner": "BIBS Product Owner", "prep": pc.PREPARED,
                     "status": "Issued for BDOI review"})
    for p in sorted((pc.OUT_PROGRAMME / "FitGap").glob("BIBS_FitGap_*.xlsx")):
        m = re.match(r"BIBS_FitGap_(BRD-\d\d)_(.+)_v(\d+\.\d+)\.xlsx$", p.name)
        if m:
            rows.append({"brd": m[1], "kind": "FitGap", "title": m[2].replace("_", " "), "file": p.stem,
                         "fmt": "XLSX", "ver": m[3], "date": _norm(_xlsx_date(p)), "owner": "BIBS Product Owner",
                         "prep": pc.PREPARED, "status": "Issued for BDOI review"})
    wb = pc.workbook("Document Register", "Document register", "BRD-00", G.VERSION, G.DATE,
                     "Current documents of BRD-11, BRD-03, BRD-06 and BRD-02 and the programme documents")
    status = ["Issued for BDOI review", "Signed off", "To be issued", "Superseded"]
    wb.sheet("Register", [pc.Column("brd", "BRD", 10), pc.Column("kind", "Type", 10), pc.Column("title", "Document",
                          40), pc.Column("file", "File name", 60), pc.Column("fmt", "Format", 8),
                          pc.Column("ver", "Version", 9),
                          pc.Column("date", "Date", 13), pc.Column("owner", "Owner (BDOI)", 30),
                          pc.Column("prep", "Prepared by", 26), pc.Column("status", "Status", 20, values=status,
                                                                          status=True)], rows,
             description="One row per current document; standard documents not yet issued are listed as To be "
             "issued. Older versions are not kept.")
    return pc.save_workbook(wb, pc.OUT_COMMON / G.FILES["REG"]), len(rows)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--only", default="NFR,GLO,LOG,REG")
    args = ap.parse_args()
    for old in G.SUPERSEDED:
        pc.retire(pc.OUT_PROGRAMME / old)
    with tempfile.TemporaryDirectory() as t:
        for key in args.only.split(","):
            if key == "NFR":
                path, n = write_nfr(Path(t), not args.no_pdf)
                print(f"wrote {path.name} ({n} TOC lines)")
            else:
                path, n = {"GLO": write_glossary, "LOG": write_log, "REG": write_register}[key]()
                print(f"wrote {path.name} ({n} rows)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
