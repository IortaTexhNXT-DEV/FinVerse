"""Comparison of BDOI's Operations Cashiering FRS with the BIBS FRS BRD-2 Operations (cashiering scope), and the plan to
bring our content into BDOI's format.

    python docs/deliverables/src/programme/comparisons/build_brd02_cashiering_comparison.py            # the workbook
    python docs/deliverables/src/programme/comparisons/build_brd02_cashiering_comparison.py --check    # checks only

Sources:
  * docs/source-documents/FRS - BDOI Core Replacement Operations Cashiering v3.1 (BDOI).docx – BDOI's FRS (BDO ITG,
    revision 1.1 of 09-Oct-2026), read for its counts (FR items, numbered sub-items, mapped BRD IDs, images, review
    comments) and to check that every BDOI reference quoted in the data exists in it;
  * brd02_cashiering_comparison.yaml (this folder) – summary and recommendation, structure, pros and cons, BRD coverage,
    requirement mapping, conflicts and the incorporation plan;
  * brd02_cashiering_comparison_lists.yaml (this folder) – only in ours, only in BDOI's FRS, file and field comparison;
  * our side, counted from the BRD-2 sources (docs/deliverables/src/BRD-02_Operations): the cashiering FRs with their
    rules, validations, acceptance criteria and fields, the test cases of those FRs, the cashiering screens of the
    sign-off pack, and the messages of the cashiering screens in the issued FRS v2.1.

Output: docs/deliverables/out/Programme/Comparisons/BIBS_Comparison_BRD-02_BDOI_FRS_vs_BIBS_FRS_Operations_Cashiering_v1.0.xlsx

Every text of the workbook is checked against the wording rules of tools/deliverables/check_pack.py before it is
written (restricted words, development-status wording, technical terms).
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402
import check_pack  # noqa: E402
from bdoi_xlsx import BdoiWorkbook, Column  # noqa: E402

DATA = HERE / "brd02_cashiering_comparison.yaml"
LISTS = HERE / "brd02_cashiering_comparison_lists.yaml"
BDOI_DOCX = REPO / "docs" / "source-documents" / "FRS - BDOI Core Replacement Operations Cashiering v3.1 (BDOI).docx"
OPS_SRC = brand.SRC_DIR / "BRD-02_Operations"
OPS_FRS_OUT = brand.out_dir("BRD-02", "FRS") / "02_BIBS_FRS_BRD-02_Operations_v2.1.docx"
OUT_FOLDER = brand.OUT_DIR / brand.DROPS["Programme"]["folder"] / "Comparisons"

# The cashiering FRs of our FRS (FR-OP-017 and FR-OP-026 are removed from scope) and the cashiering screens.
CASH_FRS = [f"FR-OP-{n:03d}" for n in list(range(8, 29)) + [132] if n not in (17, 26)]
CASH_SCREENS = re.compile(r"SCR-OP-(?:0[7-9]|1\d|2[0-2])\b")

MATCH = ["Same", "Partial", "Different", "Not in ours"]
PLATFORM = ["Yes", "Partly", "No"]
ACTIONS = ["Adopt BDOI wording", "Keep ours", "Needs BDOI decision", "Platform change"]
COVERAGE = ["Full", "Partial", "By reference", "Not covered", "Removed"]
RECOMMEND = ["Include", "Include as annex", "Keep as supporting document"]
BDOI_ACTIONS = ["Add to our FRS", "Platform change needed", "Ask BDOI"]
SIDES = ["BDOI's FRS", "Our FRS", "Both", "Neither", "BRD silent", "Not stated", "-"]
KINDS = ["Between the documents", "Within BDOI's FRS"]
PRIORITIES = ["High", "Medium", "Low"]
DECISIONS = ["Open", "Decided", "Deferred"]
# Colours of the values: each value takes the colour of a status of the pack's palette (brand.STATUS_COLOURS).
COLOURS = {"SAME": "FIT", "DIFFERENT": "CHANGE", "NOT IN OURS": "GAP", "YES": "FIT", "PARTLY": "PARTIAL", "NO": "GAP",
           "FULL": "FIT", "BY REFERENCE": "N/A", "NOT COVERED": "GAP", "REMOVED": "OUT", "INCLUDE": "FIT",
           "INCLUDE AS ANNEX": "CONFIGURE", "KEEP AS SUPPORTING DOCUMENT": "N/A", "ADD TO OUR FRS": "CONFIGURE",
           "PLATFORM CHANGE NEEDED": "CHANGE", "ASK BDOI": "OPEN", "ADOPT BDOI WORDING": "CONFIGURE",
           "KEEP OURS": "FIT", "NEEDS BDOI DECISION": "OPEN", "PLATFORM CHANGE": "CHANGE", "HIGH": "GAP",
           "MEDIUM": "PARTIAL", "LOW": "N/A", "DECIDED": "DONE", "DEFERRED": "ON HOLD"}
BDOI_REF = re.compile(r"FRS\.CSH\.\d\d\.\d\d(?:\s?\.\s?\d+)*")
FR_REF = re.compile(r"FR-OP-(\d{3})")


def load(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


# ------------------------------------------------------------------------------------------------------- BDOI's FRS

def bdoi_facts() -> dict:
    """Counts and reference IDs of BDOI's FRS, read from the Word file."""
    import docx  # noqa: PLC0415

    document = docx.Document(str(BDOI_DOCX))
    tables = document.tables
    fr_table = next(t for t in tables if t.rows[0].cells[0].text.strip() == "Functional Requirement ID")
    map_table = next(t for t in tables if t.rows[0].cells[0].text.strip() == "BRD Req No.")
    text = "\n".join(c.text for row in fr_table.rows for c in row.cells)
    refs = [re.sub(r"\s", "", r) for r in BDOI_REF.findall(text)]
    unique = list(dict.fromkeys(refs))
    items = [r for r in unique if r.count(".") == 3]
    subs = [r for r in unique if r.count(".") > 3]
    brd_ids, mapped = [], []
    for row in map_table.rows[1:]:
        brd_id, frs = row.cells[0].text.strip(), row.cells[2].text.strip()
        brd_ids.append(brd_id)
        if BDOI_REF.search(frs):
            mapped.append(brd_id)
    with zipfile.ZipFile(BDOI_DOCX) as z:
        names = z.namelist()
        comments = len(re.findall(r"<w:comment ", z.read("word/comments.xml").decode("utf-8"))) \
            if "word/comments.xml" in names else 0
        images = len([n for n in names if re.search(r"(^|/)media/image\d+\.\w+$", n)])
        footers = " ".join(z.read(n).decode("utf-8") for n in names if re.match(r"word/footer\d*\.xml$", n))
    pages = re.search(r"of\s*</w:t>.*?<w:t[^>]*>\s*(\d+)\s*<", footers, re.S)
    return {"items": items, "subs": subs, "brd_ids": brd_ids, "mapped": mapped, "comments": comments,
            "images": images, "pages": pages.group(1) if pages else "", "text": text,
            "all_text": "\n".join(p.text for p in document.paragraphs) + "\n" + text}


# ------------------------------------------------------------------------------------------------------- our FRS

def our_facts() -> dict:
    """Counts of the cashiering scope of our FRS: FRs, rules, validations, acceptance criteria, fields, test cases,
    screens and messages."""
    frs_text = (OPS_SRC / "FRS_BRD02_OPERATIONS.md").read_text(encoding="utf-8")
    blocks = [yaml.safe_load(b) for b in re.findall(r"```fr\n(.*?)```", frs_text, re.S)]
    by_id = {b["id"]: b for b in blocks}
    cash = [by_id[f] for f in CASH_FRS if f in by_id]
    cases = load(OPS_SRC / "brd02_cases.yaml")["frs"]
    tests = sum(len((cases.get(f) or {}).get("cases") or []) for f in CASH_FRS)
    screens = [s for s in load(OPS_SRC / "pack" / "screens" / "02_cashiering.yaml")["screens"]
               if CASH_SCREENS.match(s["id"])]
    removed = set(re.findall(r"^### (FR-OP-\d{3}) ", frs_text, re.M))
    return {"all_frs": set(by_id) | removed, "frs": len(cash), "total_frs": len(blocks),
            "rules": sum(len(b.get("rules") or []) for b in cash),
            "validations": sum(len(b.get("validations") or []) for b in cash),
            "acceptance": sum(len(b.get("acceptance") or []) for b in cash),
            "fields": sum(len(b.get("fields") or []) for b in cash),
            "tests": tests, "screens": len(screens), "messages": cash_messages()}


def cash_messages() -> int:
    """Messages of the cashiering screens in the messages catalogue of the issued FRS v2.1 (0 when not built)."""
    if not OPS_FRS_OUT.exists():
        return 0
    import docx  # noqa: PLC0415

    document = docx.Document(str(OPS_FRS_OUT))
    body = document.element.body
    ids: set[str] = set()
    in_catalogue = in_cash = False
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            style = child.xpath("string(.//w:pStyle/@w:val)")
            text = "".join(child.itertext())
            if style.startswith("Heading1"):
                in_catalogue = "Messages catalogue" in text
            elif style.startswith("Heading2") and in_catalogue:
                in_cash = bool(CASH_SCREENS.search(text) or "Post-dated Checks" in text
                               or "Commission Payment Details" in text)
        elif tag == "tbl" and in_catalogue and in_cash:
            ids.update(re.findall(r"MSG-OP-\d+", "".join(child.itertext())))
    return len(ids)


# ------------------------------------------------------------------------------------------------------- checks

def strings(node) -> list[str]:
    if isinstance(node, dict):
        return [s for v in node.values() for s in strings(v)]
    if isinstance(node, list):
        return [s for v in node for s in strings(v)]
    return [str(node)] if node is not None else []


def wording_errors(texts: list[str]) -> list[str]:
    """The check_pack rules on the texts of the workbook: restricted words, development-status wording, technical
    terms."""
    out = []
    for t in texts:
        for m in check_pack.RESTRICTED.finditer(t):
            out.append(f"restricted word “{m.group(0)}” in: {t[:90]}")
        for label, pattern in check_pack.BUILD_STATUS:
            for m in pattern.finditer(t):
                out.append(f"{label} “{m.group(0)}” in: {t[:90]}")
        for label, words in check_pack.technical_hits([t]).items():
            out.append(f"technical term ({label}) {sorted(words)} in: {t[:90]}")
    return out


def check(data: dict, lists: dict, bdoi: dict, ours: dict) -> list[str]:
    errors: list[str] = []
    known = set(bdoi["items"]) | set(bdoi["subs"])
    rows = data["mapping"] + lists["only_bdoi"] + [{"id": c["bdoi"]} for c in data["conflicts"]]
    for r in rows:
        for ref in BDOI_REF.findall(str(r.get("id", "")) + " " + str(r.get("ref", ""))):
            ref = re.sub(r"\s", "", ref)
            base = ".".join(ref.split(".")[:4])
            if base not in known:
                errors.append(f"BDOI reference {ref} is not in BDOI's FRS")
    for fr in sorted({f"FR-OP-{n}" for s in strings([data, lists]) for n in FR_REF.findall(s)}):
        if fr not in ours["all_frs"]:
            errors.append(f"{fr} is not an FR of our FRS")
    conflict_ids = {c["id"] for c in data["conflicts"]}
    for s in strings([data["mapping"], lists, data["decisions_order"], data["plan"]]):
        for cid in re.findall(r"\bC(\d{1,2})\b", s):
            if f"C{cid}" not in conflict_ids:
                errors.append(f"conflict C{cid} referenced but not defined")
    cov_ids = [c["id"] for c in data["coverage"]]
    if len(cov_ids) != len(set(cov_ids)):
        errors.append("duplicate BRD ID in the coverage")
    for b in bdoi["brd_ids"]:
        if b not in cov_ids:
            errors.append(f"BRD ID {b} of BDOI's mapping is missing in the coverage")
    for sheet, key, allowed in (("mapping", "match", MATCH), ("mapping", "platform", PLATFORM),
                                ("mapping", "action", ACTIONS), ("coverage", "bdoi_status", COVERAGE),
                                ("coverage", "ours_status", COVERAGE), ("conflicts", "kind", KINDS),
                                ("conflicts", "priority", PRIORITIES), ("conflicts", "brd", SIDES),
                                ("conflicts", "platform", SIDES)):
        for r in data[sheet]:
            if r[key] not in allowed:
                errors.append(f"{sheet}: {key} “{r[key]}” not in {allowed}")
    for sheet, key, allowed in (("only_ours", "recommendation", RECOMMEND), ("only_bdoi", "action", BDOI_ACTIONS),
                                ("files", "match", MATCH)):
        for r in lists[sheet]:
            if r[key] not in allowed:
                errors.append(f"{sheet}: {key} “{r[key]}” not in {allowed}")
    errors += wording_errors(strings([data, lists]))
    return errors


# ------------------------------------------------------------------------------------------------------- sheets

def count_line(counter: Counter, order: list[str]) -> str:
    return "; ".join(f"{k} {counter[k]}" for k in order if counter[k])


def summary_rows(data: dict, lists: dict, bdoi: dict, ours: dict) -> list[dict]:
    s = data["summary"]
    rows: list[dict] = []

    def add(section, item, b="", o="", note=""):
        rows.append({"section": section, "item": item, "bdoi": b, "ours": o, "note": note})

    for p in s["purpose"]:
        add("What each document is", p["item"], p["bdoi"], p["ours"], p["note"])
    cov = data["coverage"]
    in_scope = [c for c in cov if c["ours_status"] != "Removed"]
    blank = [c["id"] for c in cov if c["bdoi_fr"] == "(blank)"]
    add("Counts", "BRD IDs of the cashiering scope",
        f"{len(bdoi['brd_ids'])} listed in the mapping; {len(bdoi['mapped'])} with FR IDs; {len(blank)} blank "
        f"({', '.join(blank)}); 2 sent to “Com rec”",
        f"{len(cov)} listed (BRQID.001-007, CSHID.001-027, MKTID.013, DBMID.001), {len(in_scope)} in scope; "
        f"every in-scope ID traced to FRs and test cases",
        "Removed by BRD v1.01 and the May 2026 annex: CSHID.009, 026, 027 (sheet 4).")
    add("Counts", "BRD coverage (in-scope IDs)",
        count_line(Counter(c["bdoi_status"] for c in in_scope), COVERAGE),
        count_line(Counter(c["ours_status"] for c in in_scope), COVERAGE),
        "Full, Partial, By reference (another FRS), Not covered.")
    add("Counts", "Functional requirements",
        f"{len(bdoi['items'])} FR items with {len(bdoi['subs'])} numbered sub-items; {bdoi['images']} images; "
        f"{bdoi['comments']} review comments; {bdoi['pages']} pages",
        f"{ours['frs']} cashiering FRs (of {ours['total_frs']} in the FRS) with {ours['rules']} rules, "
        f"{ours['validations']} validations, {ours['acceptance']} acceptance criteria and {ours['fields']} field "
        f"rows; {ours['screens']} screen specifications; {ours['messages']} messages; {ours['tests']} test cases",
        "Our counts cover FR-OP-008 to 028 and 132 (FR-OP-017 and 026 removed from scope).")
    mapping = data["mapping"]
    add("Counts", "Requirement mapping (sheet 5)",
        f"{len(mapping)} BDOI items and sub-item groups compared",
        count_line(Counter(m["match"] for m in mapping), MATCH),
        "BIBS today: " + count_line(Counter(m["platform"] for m in mapping), PLATFORM)
        + ". Proposed action: " + count_line(Counter(m["action"] for m in mapping), ACTIONS) + ".")
    ob, oo, cf = lists["only_bdoi"], lists["only_ours"], data["conflicts"]
    add("Counts", "Only in one document (sheets 6 and 7)",
        f"{len(ob)} items only in BDOI's or different in ours: "
        + count_line(Counter(r["action"] for r in ob), BDOI_ACTIONS),
        f"{len(oo)} items only in ours: " + count_line(Counter(r["recommendation"] for r in oo), RECOMMEND))
    kinds = Counter(c["kind"] for c in cf)
    add("Counts", "Conflicts (sheet 8)",
        f"{kinds[KINDS[1]]} within BDOI's FRS",
        f"{kinds[KINDS[0]]} between the two documents",
        f"{len(cf)} in all; priority " + count_line(Counter(c["priority"] for c in cf), PRIORITIES) + ".")
    files = lists["files"]
    add("Counts", "Files, forms and screen fields (sheet 9)",
        f"{len(files)} fields of BDOI's sample files, forms and screens",
        count_line(Counter(f["match"] for f in files), MATCH),
        "Bills Payment, CLPC, Trade, PDC, Direct Credit and Commission Schedule files; AR and OR forms; "
        "five screens.")
    pc = s["pros_cons"]
    add("Headline pros and cons", "Top 5 pros", pc["bdoi_pros"], pc["ours_pros"])
    add("Headline pros and cons", "Top 5 cons", pc["bdoi_cons"], pc["ours_cons"])
    for r in s["recommendation"]:
        add("Recommendation", r["item"], note=r["note"])
    effort = sum(float(p["effort"] or 0) for p in data["plan"])
    add("Recommendation", "Effort and order",
        note=f"About {effort:.0f} person-days of the project team in the order of sheet 10, plus one BDOI ITG review "
             f"cycle of five working days; the decision session comes first.")
    by_id = {c["id"]: c for c in cf}
    for cid in data["decisions_order"]:
        c = by_id[cid]
        add("Decisions needed from BDOI", f"{cid} {c['topic']}", note=f"{c['decision']} Our recommendation: "
            f"{c['recommendation']}")
    return rows


def build(data: dict, lists: dict, bdoi: dict, ours: dict) -> Path:
    for value, status in COLOURS.items():
        brand.STATUS_COLOURS.setdefault(value, brand.STATUS_COLOURS[status])
    meta = data["meta"]
    wb = BdoiWorkbook(meta["title"], doc_type=meta["doc_type"], brd=meta["brd"], version=meta["version"],
                      date=meta["date"], subtitle=meta["subtitle"])
    wb.legend = [("Same", "The two documents say the same"), ("Partial", "Partly the same; see the difference"),
                 ("Different", "The two documents say different things"),
                 ("Not in ours", "In BDOI's FRS only (or, on sheet 9, not in the BIBS layout)"),
                 ("Yes", "BIBS does it today"), ("Partly", "BIBS does part of it today"), ("No", "BIBS does not do it today"),
                 ("Needs BDOI decision", "A decision of BDOI is needed (sheet 8)"),
                 ("Platform change", "BIBS changes once the merged FRS is agreed")]
    wb.cover_notes = [f"Compares BDOI's file “{meta['bdoi_file']}” with the {meta['ours_name']}. Quotes of BDOI's FRS "
                      "are short and word for word. “BIBS today” states what the platform does now; it is not a "
                      "commitment until the merged FRS is signed.",
                      "C1 to C29 refer to the conflicts of sheet 8; FR-OP-nnn and SCR-OP-nn to our FRS; "
                      "FRS.CSH.nn.nn to BDOI's FRS."]

    wb.sheet("1 Summary", [
        Column("section", "Section", 20, "Part of the summary"),
        Column("item", "Item", 30, "What is described or counted"),
        Column("bdoi", "BDOI's FRS", 52, "BDOI's FRS (BDO ITG, revision 1.1 of 09-Oct-2026)"),
        Column("ours", "Our FRS", 52, "BIBS FRS BRD-2 Operations v2.1 (08-Oct-2026), cashiering scope"),
        Column("note", "Comment / recommendation", 72, "Comment, recommendation or trade-off"),
    ], summary_rows(data, lists, bdoi, ours),
        description="What each document is for, the counts, the headline pros and cons, the recommendation and the "
                    "decisions needed")

    wb.sheet("2 Structure", [
        Column("bdoi", "BDOI section", 34, "Section of BDOI's FRS"),
        Column("ours", "Our chapter(s)", 38, "Chapter(s) or FRs of our FRS"),
        Column("bdoi_content", "What BDOI's contains", 46, "Content of BDOI's section"),
        Column("ours_content", "What ours contains", 46, "Content of our chapter(s)"),
        Column("depth", "Depth", 16, "Which is deeper"),
        Column("assessment", "Assessment", 50, "What to do with the section in the merged document"),
    ], data["structure"], description="BDOI's sections against our chapters, section by section")

    pc_rows = []
    for g in data["pros_cons"]:
        pc_rows.append({"group": g["group"], "doc": "BDOI's FRS", "pros": g["bdoi"]["pros"], "cons": g["bdoi"]["cons"]})
        pc_rows.append({"group": g["group"], "doc": "Our FRS", "pros": g["ours"]["pros"], "cons": g["ours"]["cons"]})
    wb.sheet("3 Pros and cons", [
        Column("group", "Group", 24, "Aspect compared"),
        Column("doc", "Document", 13, "BDOI's FRS or our FRS", values=["BDOI's FRS", "Our FRS"]),
        Column("pros", "Pros", 72, "Real strengths of the document for this aspect"),
        Column("cons", "Cons", 72, "Real weaknesses of the document for this aspect"),
    ], pc_rows, description="Strengths and weaknesses of each document, grouped by aspect")

    wb.sheet("4 BRD coverage", [
        Column("id", "BRD ID", 11, "Requirement ID of the Operations BRD v1.01 and its Cashiering annex"),
        Column("page", "BRD page", 12, "Page of the file Operations_WS Addendum (annex p.n = May 2026 annexes)"),
        Column("req", "Business requirement (short)", 40, "The requirement in short"),
        Column("bdoi_fr", "BDOI FR IDs", 18, "FR IDs in BDOI's mapping table"),
        Column("bdoi_status", "BDOI coverage", 13, "How far BDOI's FRS meets the requirement", values=COVERAGE,
               status=True),
        Column("bdoi_note", "BDOI note", 42, "What BDOI's FRS says or leaves out"),
        Column("ours_fr", "Our FRs", 18, "FRs of our FRS"),
        Column("ours_status", "Our coverage", 13, "How far our FRS meets the requirement", values=COVERAGE,
               status=True),
        Column("note", "Note", 40, "Comment; C-numbers refer to sheet 8"),
    ], data["coverage"], description="Every BRD ID of the cashiering scope: BDOI's FR IDs, our FR IDs and the coverage "
                                     "in each")

    wb.sheet("5 Requirement mapping", [
        Column("id", "BDOI ID", 20, "FR item or sub-items of BDOI's FRS"),
        Column("capability", "BDOI capability", 22, "Capability or topic"),
        Column("quote", "BDOI text (short quote)", 46, "Short quote of BDOI's text"),
        Column("ours", "Our FR IDs", 18, "Matching FRs and screens of our FRS"),
        Column("match", "Match", 12, "Same, partial, different or not in ours", values=MATCH, status=True),
        Column("difference", "Difference in plain words", 54, "What differs"),
        Column("platform", "BIBS today", 10, "Does BIBS do it today", values=PLATFORM, status=True),
        Column("action", "Proposed action", 18, "What to do", values=ACTIONS, status=True),
        Column("action_detail", "Action detail", 40, "Detail of the action; C-numbers refer to sheet 8"),
    ], data["mapping"], description="Every BDOI FR item with its sub-items against our FRs, with the platform status "
                                    "and the proposed action")

    wb.sheet("6 Only in ours", [
        Column("ref", "Our reference", 24, "FR, screen or chapter of our FRS"),
        Column("kind", "Kind", 18, "Kind of content"),
        Column("content", "Content", 62, "What it states"),
        Column("brd", "BRD", 16, "BRD requirement it serves"),
        Column("where", "Where it goes in BDOI's template", 38, "Section, FR item or annex"),
        Column("recommendation", "Recommendation", 18, "Include, include as annex or keep as supporting document",
               values=RECOMMEND, status=True),
    ], lists["only_ours"], description="Content of our FRS (cashiering scope) that BDOI's FRS does not have, and where "
                                       "it goes")

    wb.sheet("7 Only in BDOI FRS", [
        Column("ref", "BDOI reference", 26, "Item, sub-item, appendix or comment of BDOI's FRS"),
        Column("item", "Item", 62, "What BDOI's FRS has that ours lacks or states differently"),
        Column("impact", "Impact", 40, "Effect on our FRS or on BIBS"),
        Column("action", "Action", 18, "Add to our FRS, platform change needed or ask BDOI", values=BDOI_ACTIONS,
               status=True),
        Column("detail", "Detail", 34, "Detail; C-numbers refer to sheet 8"),
    ], lists["only_bdoi"], description="Items of BDOI's FRS missing or different in ours, including Appendix A and B")

    wb.sheet("8 Conflicts", [
        Column("id", "ID", 6, "Conflict number"),
        Column("kind", "Kind", 16, "Between the documents or within BDOI's FRS", values=KINDS),
        Column("priority", "Priority", 9, "Effect on the merge and on BIBS", values=PRIORITIES, status=True),
        Column("topic", "Topic", 24, "What the conflict is about"),
        Column("bdoi", "BDOI's FRS says", 46, "BDOI's text (short quotes)"),
        Column("ours", "Our FRS says", 42, "Our text"),
        Column("brd", "Matches the BRD", 12, "Which side matches the BRD", values=SIDES),
        Column("brd_note", "BRD text", 38, "What the BRD says"),
        Column("platform", "Matches BIBS", 12, "Which side matches BIBS today", values=SIDES),
        Column("decision", "Decision needed from BDOI", 38, "The question for BDOI"),
        Column("recommendation", "Our recommendation", 44, "What the project team recommends"),
        Column("status", "Decision status", 11, "Recorded at the decision session", values=DECISIONS, status=True),
        Column("bdoi_decision", "BDOI decision", 30, "Recorded at the decision session"),
    ], [{**c, "status": "Open", "bdoi_decision": ""} for c in data["conflicts"]],
        description="Where the two documents say different things, and points to correct within BDOI's FRS")

    wb.sheet("9 File and field", [
        Column("file", "File, form or screen", 24, "BDOI sample file (Appendix A), form or screen"),
        Column("num", "Pos.", 6, "Position or row in BDOI's file or form"),
        Column("field", "Field", 24, "Field"),
        Column("bdoi", "BDOI: position, format, sample", 40, "What BDOI's sample or text shows"),
        Column("bdoi_mandatory", "Mandatory (BDOI)", 12, "Mandatory in BDOI's FRS (its files do not state it)"),
        Column("platform", "BIBS layout or field", 38, "Upload template, form or screen field of BIBS today"),
        Column("platform_mandatory", "Mandatory (BIBS)", 14, "Mandatory in BIBS"),
        Column("ours", "Our FRS", 26, "What our FRS states"),
        Column("match", "Match", 12, "Same, partial, different or not in ours (not in the BIBS layout)",
               values=MATCH, status=True),
        Column("difference", "Difference", 46, "What differs; C-numbers refer to sheet 8"),
    ], lists["files"], description="BDOI's sample AR, Billspayment, CLPC and Trade files, the other channel files and "
                                   "BDOI's screen fields against the BIBS layouts and our FRS",
        freeze_first_column=False)

    wb.sheet("10 Incorporation plan", [
        Column("part", "Part", 16, "A ID scheme, B section by section, C order of work, D ITG review, E platform "
                                   "changes"),
        Column("bdoi_section", "BDOI section", 30, "Section of BDOI's FRS (or new annex)"),
        Column("insert", "Content to insert from ours", 64, "What goes in"),
        Column("how", "How (form in BDOI's FRS)", 26, "Form it takes in BDOI's document"),
        Column("source", "Source in ours", 24, "Where it comes from"),
        Column("effort", "Effort (person-days)", 11, "Indicative effort of the project team", kind="number"),
        Column("itg", "BDOI ITG review", 32, "What BDOI ITG reviews or decides"),
    ], data["plan"], description="How our content is placed into BDOI's FRS format: ID scheme, sections, order of "
                                 "work, review and platform changes")
    return wb.save(OUT_FOLDER / meta["filename"])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="checks only")
    args = ap.parse_args()
    data, lists = load(DATA), load(LISTS)
    bdoi, ours = bdoi_facts(), our_facts()
    errors = check(data, lists, bdoi, ours)
    for e in errors:
        print("error:", e)
    if errors:
        return 1
    if args.check:
        print("check: 0 errors")
        return 0
    path = build(data, lists, bdoi, ours)
    late = wording_errors(check_pack.office_lines(path))
    for e in late:
        print("error:", e)
    print(f"wrote {path.relative_to(REPO)} ({path.stat().st_size // 1024} KB)")
    return 1 if late else 0


if __name__ == "__main__":
    sys.exit(main())
