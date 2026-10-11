"""Builds the delivery and adoption documents of the BIBS programme (BRD-00).

Outputs (docs/deliverables/out/Programme/Delivery and .../Adoption):
  Delivery  BIBS_Delivery_BRD-00_Project_Plan_v1.0.xlsx            timeline (Gantt), milestones, sign-offs, test
                                                                   windows, cut-over and hypercare, BDOI dependencies
            BIBS_Delivery_BRD-00_Project_Plan_Summary_v1.0.docx
            BIBS_Delivery_BRD-00_Delivery_Methodology_v1.0.docx
            BIBS_Delivery_BRD-00_RACI_Matrix_v1.0.xlsx
            BIBS_Delivery_BRD-00_Risk_Register_RAID_Log_v1.0.xlsx
  Adoption  BIBS_Adoption_BRD-00_Knowledge_Transfer_Plan_v1.0.docx
            BIBS_Adoption_BRD-00_Change_Management_and_Training_Framework_v1.0.docx
            BIBS_Adoption_BRD-00_User_Manual_Framework_v1.0.docx
            BIBS_Adoption_BRD-00_QRG_<persona>_v1.0.docx (one per persona of guides.yaml)

Inputs (this folder): plan.yaml, raid.yaml, raci.yaml, guides.yaml and the Word sources *.md. A line
`<!-- da:<name> -->` of a Word source is replaced by the Markdown table that render_<name>() builds from the YAML
files; figures/*.dot are rendered by the Word builder; figures/pp_gantt.png is drawn here from plan.yaml.

Usage
  python docs/deliverables/src/programme/delivery_adoption/build_delivery_adoption.py --check
  python docs/deliverables/src/programme/delivery_adoption/build_delivery_adoption.py [--no-pdf] [--only NAME]
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path
from typing import Any

import yaml
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402
from bdoi_docx import BdoiDocument, load_source, meta_from, render_body  # noqa: E402
from bdoi_xlsx import HEADER_ROW, BdoiWorkbook, Column, _SheetInfo  # noqa: E402

VERSION = "1.0"
DATE = "08 October 2026"
OUT_DELIVERY = brand.OUT_DIR / brand.DROPS["Programme"]["folder"] / "Delivery"
OUT_ADOPTION = brand.OUT_DIR / brand.DROPS["Programme"]["folder"] / "Adoption"
PLACEHOLDER = re.compile(r"^<!--\s*da:(\w+)\s*(.*?)-->\s*$")
THIN = Side(style="thin", color=brand.BORDER)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
KIND_COLOURS = {"req": "004EA8", "dev": "2E9BD6", "test": "D99400", "mig": "2E7D32", "ready": "B5446E",
                "live": "004EA8"}
RATING_COLOURS = {"High": (brand.DANGER_BG, brand.DANGER), "Medium": (brand.AMBER_BG, brand.AMBER),
                  "Low": (brand.SUCCESS_BG, brand.SUCCESS)}
RACI_COLOURS = {"A": (brand.YELLOW, brand.NEAR_BLACK), "A/R": (brand.YELLOW, brand.NEAR_BLACK), "R": (brand.HEADER_BLUE, brand.WHITE),
                "C": (brand.BG_BLUE, brand.HEADER_BLUE), "I": (brand.DIRTY_WHITE, brand.MUTED)}


def load(name: str) -> dict[str, Any]:
    return yaml.safe_load((HERE / name).read_text(encoding="utf-8"))


PLAN, RAID, RACI, GUIDES = load("plan.yaml"), load("raid.yaml"), load("raci.yaml"), load("guides.yaml")


def d(value: Any) -> str:
    """dd-MMM-yyyy of a date (or the text as it is)."""
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%d-%b-%Y")
    return str(value)


def cell(value: Any) -> str:
    return d(value).replace("|", "/").replace("\n", " ").strip()


def table(headers: list[str], rows: list[list[Any]], widths: str, caption: str, size: str = "8.5",
          status: str = "") -> list[str]:
    opts = f'widths={widths} caption="{caption}" size={size}' + (f" status={status}" if status else "")
    out = [f"<!-- table: {opts} -->", "| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(cell(c) for c in r) + " |" for r in rows]
    return out + [""]


# ------------------------------------------------------------------------------------------------ RAID helpers

def score(r: dict[str, Any]) -> int:
    return int(r["prob"]) * int(r["impact"])


def rating(r: dict[str, Any]) -> str:
    s = score(r)
    return "High" if s >= 15 else "Medium" if s >= 8 else "Low"


def raci_map(spec: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for part in spec.split():
        letter, roles = part.split(":")
        for role in roles.split(","):
            out[role] = letter if role not in out else out[role] + "/" + letter
    return out


# ------------------------------------------------------------------------------------------------ checks

def check() -> list[str]:
    problems: list[str] = []
    for row in PLAN["timeline"]:
        if row[4] > row[5]:
            problems.append(f"{row[0]}: start after end")
        if row[8] not in PLAN["kinds"]:
            problems.append(f"{row[0]}: kind {row[8]}")
    dates = [m[1] for m in PLAN["milestones"]]
    if dates != sorted(dates):
        problems.append("milestones not in date order")
    roles = set(RACI["order"])
    if roles != set(RACI["roles"]):
        problems.append("RACI order and roles differ")
    for area, act, _, spec in RACI["activities"]:
        m = raci_map(spec)
        unknown = set(m) - roles
        if unknown:
            problems.append(f"RACI {act}: unknown roles {sorted(unknown)}")
        if sum(1 for v in m.values() if "A" in v.split("/")) != 1:
            problems.append(f"RACI {act}: needs exactly one A")
        if not any("R" in v.split("/") for v in m.values()):
            problems.append(f"RACI {act}: no R")
    for r in RAID["risks"]:
        if r["category"] not in RAID["categories"] or r["status"] not in RAID["statuses"]:
            problems.append(f"{r['id']}: category or status")
    for p in GUIDES["personas"]:
        for t in p["tasks"]:
            for s in t["steps"]:
                if len(s) != 3:
                    problems.append(f"{p['key']} {t['title']}: step needs 3 cells")
    return problems


# ------------------------------------------------------------------------------------------------ Gantt figure

def gantt_png(path: Path) -> Path:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt

    rows = PLAN["timeline"]
    fig, ax = plt.subplots(figsize=(11.5, 9.2), dpi=200)
    for i, r in enumerate(rows):
        start, end = r[4], r[5]
        width = max((end - start).days, 4)
        ax.barh(i, width, left=mdates.date2num(start), height=0.62, color="#" + KIND_COLOURS[r[8]],
                edgecolor="white", linewidth=1.5)
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels([f"{r[0]}  {short(r[2])}" for r in rows], fontsize=6.6, fontname="DejaVu Sans")
    ax.invert_yaxis()
    ax.xaxis_date()
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b\n%y"))
    ax.tick_params(axis="x", labelsize=6.5)
    lo, hi = dt.date(2026, 9, 1), dt.date(2028, 5, 31)
    ax.set_xlim(mdates.date2num(lo), mdates.date2num(hi))
    ax.grid(axis="x", color="#E3E3E3", linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    for when, label in ((PLAN["as_of"], "Today 08-Oct-2026"), (PLAN["go_live"], "Go-live 03-Jan-2028")):
        x = mdates.date2num(when)
        ax.axvline(x, color="#C62828" if "Go" in label else "#656565", linewidth=1.1, linestyle="--")
        ax.text(x + 4, -1.1, label, fontsize=6.8, ha="left", color="#333333")
    handles = [plt.Rectangle((0, 0), 1, 1, color="#" + KIND_COLOURS[k]) for k in PLAN["kinds"] if k != "live"]
    labels = [v for k, v in PLAN["kinds"].items() if k != "live"]
    labels[0] = labels[0] + "; cut-over and go-live"
    ax.legend(handles, labels, loc="lower left", fontsize=6.6, frameon=False, ncol=3, bbox_to_anchor=(0, -0.13))
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def short(text: str, limit: int = 62) -> str:
    text = re.sub(r"\s*\(.*?\)", "", text)
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + " ..."


# ------------------------------------------------------------------------------------------------ workbooks

class Workbook(BdoiWorkbook):
    """BdoiWorkbook with free-form sheets registered in the cover index."""

    def custom(self, name: str, description: str):
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws["A1"] = name
        ws["A1"].font = Font(name=brand.FONT, size=14, bold=True, color=brand.HEADER_BLUE)
        ws["A2"] = description
        ws["A2"].font = Font(name=brand.FONT, size=9, color=brand.MUTED)
        ws.row_dimensions[1].height = 24
        self._print_setup(ws, f"{HEADER_ROW}:{HEADER_ROW}")
        self._sheets.append(_SheetInfo(name, description, [], None))
        return ws


def fills(ws, first_col: int, last_col: int, n: int, colours: dict[str, tuple[str, str]]) -> None:
    rng = f"{get_column_letter(first_col)}{HEADER_ROW + 1}:{get_column_letter(last_col)}{HEADER_ROW + max(n, 1) + 200}"
    for value, (bg, fg) in colours.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], stopIfTrue=True,
                                                      fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                                                      font=Font(color=fg, bold=True)))


def project_plan_xlsx() -> Path:
    wb = Workbook("BIBS Project Plan", doc_type="Project plan", brd="BRD-00", version=VERSION, date=DATE,
                  subtitle="Drops, milestones, sign-offs, test windows, trial migrations, cut-over and hypercare to "
                           "the go-live of January 2028")
    wb.legend = [("DONE", "Milestone reached"), ("OPEN", "Milestone ahead")]
    wb.cover_notes = ["Dates are those of the BDOI drop plan and programme timeline (26-Sep-2026), the Drop 0 closure "
                      "summary v2.0 and the Data Migration Handbook; dates marked proposed are iorta TechNXT "
                      "proposals listed as open points for BDOI. Status as of 08-Oct-2026."]
    timeline_sheet(wb)
    wb.sheet("Milestones", [
        Column("id", "ID", 9, "Milestone identifier"),
        Column("date", "Date", 13, "Date of the milestone", kind="date"),
        Column("ms", "Milestone", 62, "What is reached on the date"),
        Column("drop", "Drop", 14, "BDOI drop"),
        Column("gate", "Gate / sign-off", 14, "Kind of milestone"),
        Column("owner", "Owner", 30, "Who reaches or signs the milestone"),
        Column("status", "Status", 11, "Done or Open", values=["Done", "Open"], status=True),
        Column("src", "Source", 30, "Document that gives the date"),
    ], [dict(zip(("id", "date", "ms", "drop", "gate", "owner", "status", "src"), m)) for m in PLAN["milestones"]],
        description="Milestones of the programme in date order, from the Drop 0 sets to the close of the true-ups")
    wb.sheet("Sign-offs", [
        Column("set", "Document set", 34, "Set or decision that is signed"),
        Column("brds", "BRDs", 24, "BRDs covered"),
        Column("drop", "Drop", 13, "BDOI drop"),
        Column("ver", "Version", 22, "Version signed"),
        Column("due", "Due", 13, "Sign-off due date", kind="date"),
        Column("who", "Signatories", 55, "Who signs (BRD approval sheets and closure summary)"),
        Column("status", "Status", 11, "Open or Done", values=["Open", "Done"], status=True),
    ], [dict(zip(("set", "brds", "drop", "ver", "due", "who", "status"), s)) for s in PLAN["signoffs"]],
        description="Sign-offs of the document sets, test results and go-live decisions")
    wb.sheet("Test windows", [
        Column("w", "Window", 28, "Test or migration window"),
        Column("drop", "Drop", 13, "BDOI drop"),
        Column("start", "Start", 13, "First day", kind="date"),
        Column("end", "End", 13, "Last day", kind="date"),
        Column("env", "Environment", 12, "Environment used"),
        Column("entry", "Entry criteria", 55, "What must be true to start"),
        Column("exit", "Exit criteria", 55, "What must be true to finish"),
    ], [dict(zip(("w", "drop", "start", "end", "env", "entry", "exit"), t)) for t in PLAN["test_windows"]],
        description="SIT and UAT windows per drop, trial migrations, dress rehearsal and the performance test")
    wb.sheet("Cut-over and hypercare", [
        Column("ph", "Phase", 8, "Phase of the cut-over plan of the Data Migration Handbook"),
        Column("name", "Name", 28, "Phase name"),
        Column("days", "Days", 14, "Days relative to go-live T"),
        Column("dates", "Dates", 20, "Calendar dates for T = Monday 3-Jan-2028"),
        Column("what", "Main activities", 70, "What happens in the phase"),
        Column("owner", "Owner", 30, "Who leads the phase"),
    ], [dict(zip(("ph", "name", "days", "dates", "what", "owner"), c)) for c in PLAN["cutover"]],
        description="Cut-over phases A-H from readiness to the final true-up")
    wb.sheet("BDOI dependencies", [
        Column("id", "ID", 9, "Dependency identifier"),
        Column("type", "Type", 12, "Decision, data, environment, integration or resource",
               values=["Decision", "Data", "Environment", "Integration", "Resource"]),
        Column("what", "Dependency", 60, "What BDOI provides or decides"),
        Column("by", "Provided by", 30, "BDOI or BDO owner"),
        Column("need", "Needed by", 13, "Date needed", kind="date"),
        Column("for", "Needed for", 34, "Activity that waits for it"),
        Column("late", "Effect if late", 40, "What happens when it is late"),
    ], [dict(zip(("id", "type", "what", "by", "need", "for", "late"), x)) for x in PLAN["dependencies"]],
        description="Decisions, data, environments, integrations and people the plan needs from BDOI")
    path = OUT_DELIVERY / f"BIBS_Delivery_BRD-00_Project_Plan_v{VERSION}.xlsx"
    return wb.save(path)


def months() -> list[dt.date]:
    y, m = map(int, PLAN["months"]["start"].split("-"))
    ey, em = map(int, PLAN["months"]["end"].split("-"))
    out = []
    while (y, m) <= (ey, em):
        out.append(dt.date(y, m, 1))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def timeline_sheet(wb: Workbook) -> None:
    ws = wb.custom("Timeline", "Gantt view by month from September 2026 to May 2028; bars coloured by kind of work "
                               "(legend below the chart); go-live month outlined in red")
    heads = [("ID", 7), ("Stream", 13), ("Activity", 52), ("Drop", 13), ("Start", 11), ("End", 11), ("Owner", 28),
             ("Gate or output", 30)]
    mlist = months()
    for c, (h, w) in enumerate(heads, start=1):
        head(ws, c, h, w)
    for k, mo in enumerate(mlist, start=len(heads) + 1):
        head(ws, k, mo.strftime("%b %y"), 5.2, rotate=True)
    ws.row_dimensions[HEADER_ROW].height = 42
    golive_col = None
    for i, r in enumerate(PLAN["timeline"], start=HEADER_ROW + 1):
        vals = [r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7]]
        for c, v in enumerate(vals, start=1):
            x = ws.cell(row=i, column=c, value=v)
            x.font = Font(name=brand.FONT, size=9)
            x.alignment = Alignment(wrap_text=True, vertical="top")
            x.border = BOX
            if c in (5, 6):
                x.number_format = "dd-mmm-yyyy"
        for k, mo in enumerate(mlist, start=len(heads) + 1):
            nxt = dt.date(mo.year + (mo.month == 12), 1 if mo.month == 12 else mo.month + 1, 1)
            x = ws.cell(row=i, column=k)
            x.border = BOX
            if r[4] < nxt and r[5] >= mo:
                x.fill = PatternFill("solid", fgColor=KIND_COLOURS[r[8]])
            if mo.year == 2028 and mo.month == 1:
                golive_col = k
    last = HEADER_ROW + len(PLAN["timeline"])
    if golive_col:
        red = Side(style="medium", color=brand.DANGER)
        for row in range(HEADER_ROW, last + 1):
            ws.cell(row=row, column=golive_col).border = Border(left=red, right=red, top=THIN, bottom=THIN)
    row = last + 2
    ws.cell(row=row, column=2, value="Legend").font = Font(name=brand.FONT, size=10, bold=True,
                                                            color=brand.HEADER_BLUE)
    for j, (k, label) in enumerate(PLAN["kinds"].items(), start=1):
        ws.cell(row=row + j, column=2).fill = PatternFill("solid", fgColor=KIND_COLOURS[k])
        ws.cell(row=row + j, column=3, value=label).font = Font(name=brand.FONT, size=9)
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=4)
    ws.auto_filter.ref = f"A{HEADER_ROW}:H{last}"
    wb._sheets[-1].rows = len(PLAN["timeline"])


def head(ws, col: int, text: str, width: float, rotate: bool = False) -> None:
    x = ws.cell(row=HEADER_ROW, column=col, value=text)
    x.font = Font(name=brand.FONT, size=9 if rotate else 10, bold=True, color=brand.WHITE)
    x.fill = PatternFill("solid", fgColor=brand.HEADER_BLUE)
    x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if rotate else None,
                            text_rotation=90 if rotate else 0)
    x.border = BOX
    ws.column_dimensions[get_column_letter(col)].width = width


def raci_xlsx() -> Path:
    wb = Workbook("BIBS RACI Matrix", doc_type="RACI matrix", brd="BRD-00", version=VERSION, date=DATE,
                  subtitle="Who is responsible, accountable, consulted and informed for each activity of the programme")
    wb.legend = []
    wb.cover_notes = ["R = responsible (does the work); A = accountable (approves; exactly one per activity); C = "
                      "consulted before; I = informed after. BDOI roles first, then iorta TechNXT roles."]
    order = RACI["order"]
    cols = [Column("area", "Area", 13, "Phase or area of the programme"),
            Column("act", "Activity", 46, "Activity"),
            Column("out", "Deliverable or evidence", 26, "What shows the activity is done")]
    cols += [Column(r, r, 6.2, f"{RACI['roles'][r][0]} ({RACI['roles'][r][1]})", values=["R", "A", "C", "I", "A/R"])
             for r in order]
    rows = []
    for area, act, out, spec in RACI["activities"]:
        m = raci_map(spec)
        row = {"area": area, "act": act, "out": out}
        row.update({r: m.get(r) for r in order})
        rows.append(row)
    ws = wb.sheet("RACI", cols, rows, description="One row per activity; role codes are explained on the sheet Roles")
    fills(ws, 4, 3 + len(order), len(rows), RACI_COLOURS)
    for c in range(4, 4 + len(order)):
        for r in range(HEADER_ROW + 1, HEADER_ROW + len(rows) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=HEADER_ROW, column=c).alignment = Alignment(text_rotation=90, horizontal="center")
    ws.row_dimensions[HEADER_ROW].height = 40
    wb.sheet("Roles", [
        Column("code", "Code", 8, "Column code on the sheet RACI"),
        Column("name", "Role", 34, "Role"),
        Column("org", "Organisation", 18, "BDOI, BDO Unibank or iorta TechNXT"),
        Column("desc", "What the role does in the programme", 70, "Scope of the role"),
        Column("a", "Accountable for", 9, "Number of activities where the role is accountable", kind="number"),
        Column("r", "Responsible for", 9, "Number of activities where the role is responsible", kind="number"),
    ], [{"code": c, "name": RACI["roles"][c][0], "org": RACI["roles"][c][1], "desc": RACI["roles"][c][2],
         "a": sum(1 for x in rows if "A" in str(x.get(c) or "").split("/")),
         "r": sum(1 for x in rows if "R" in str(x.get(c) or "").split("/"))} for c in order],
        description="The roles of the matrix with their organisation")
    return wb.save(OUT_DELIVERY / f"BIBS_Delivery_BRD-00_RACI_Matrix_v{VERSION}.xlsx")


def raid_xlsx() -> Path:
    wb = Workbook("BIBS Risk Register (RAID log)", doc_type="Risk register", brd="BRD-00", version=VERSION, date=DATE,
                  subtitle="Programme risks, assumptions, issues and dependencies with owner, mitigation and status")
    wb.legend = [("OPEN", "Not yet mitigated or resolved"), ("IN PROGRESS", "Mitigation or action under way"),
                 ("CLOSED", "Closed")]
    wb.cover_notes = ["Probability and impact 1 (very low) to 5 (very high); score = probability x impact; rating High "
                      "from 15, Medium 8 to 14, Low up to 7. High items go to the steering committee; the log is "
                      "reviewed at the weekly programme meeting. Status as of 08-Oct-2026."]
    risks = RAID["risks"]
    summary = []
    for cat in RAID["categories"]:
        rs = [r for r in risks if r["category"] == cat and r["status"] != "Closed"]
        if rs:
            summary.append({"cat": cat, "h": sum(rating(r) == "High" for r in rs),
                            "m": sum(rating(r) == "Medium" for r in rs), "l": sum(rating(r) == "Low" for r in rs),
                            "n": len(rs), "top": ", ".join(r["id"] for r in sorted(rs, key=score, reverse=True)[:3])})
    wb.sheet("Summary", [
        Column("cat", "Category", 16, "Risk category"),
        Column("h", "High", 8, "Open High risks", kind="number"),
        Column("m", "Medium", 9, "Open Medium risks", kind="number"),
        Column("l", "Low", 8, "Open Low risks", kind="number"),
        Column("n", "Open total", 10, "Open risks of the category", kind="number"),
        Column("top", "Highest scores", 30, "Risks with the highest score"),
    ], summary, description=f"Open risks by category and rating; {len(RAID['assumptions'])} assumptions, "
                            f"{len(RAID['issues'])} issues, {len(PLAN['dependencies'])} dependencies")
    ws = wb.sheet("Risks", [
        Column("id", "ID", 7, "Risk identifier"),
        Column("category", "Category", 12, "Category", values=RAID["categories"]),
        Column("title", "Risk", 30, "Short name of the risk"),
        Column("cause", "Cause", 46, "What causes the risk"),
        Column("effect", "Effect", 36, "What happens if it occurs"),
        Column("prob", "P", 5, "Probability 1-5", kind="number"),
        Column("impact", "I", 5, "Impact 1-5", kind="number"),
        Column("score", "Score", 7, "Probability x impact", kind="number"),
        Column("rating", "Rating", 9, "High, Medium or Low", values=["High", "Medium", "Low"]),
        Column("owner", "Owner", 24, "Owner of the mitigation"),
        Column("mitigation", "Mitigation", 50, "Actions that reduce probability or impact"),
        Column("contingency", "Contingency", 36, "What we do if it occurs"),
        Column("due", "Review / trigger date", 12, "Date of the next decision or trigger", kind="date"),
        Column("status", "Status", 11, "Status", values=RAID["statuses"], status=True),
        Column("source", "Source", 26, "Where the risk comes from"),
    ], [dict(r, score=score(r), rating=rating(r)) for r in risks],
        description="Programme risks with probability, impact, owner, mitigation and status")
    fills(ws, 9, 9, len(risks), RATING_COLOURS)
    wb.sheet("Assumptions", [
        Column("id", "ID", 7, "Assumption identifier"),
        Column("a", "Assumption", 60, "What the plan assumes"),
        Column("basis", "Basis", 30, "BDOI answer or document behind it"),
        Column("eff", "Effect if false", 34, "What changes if the assumption is wrong"),
        Column("owner", "Owner", 24, "Who validates it"),
        Column("by", "Validate by", 12, "Date", kind="date"),
        Column("status", "Status", 11, "Status", values=RAID["statuses"], status=True),
    ], [dict(zip(("id", "a", "basis", "eff", "owner", "by", "status"), a)) for a in RAID["assumptions"]],
        description="Planning assumptions to validate with BDOI")
    ws = wb.sheet("Issues", [
        Column("id", "ID", 7, "Issue identifier"),
        Column("issue", "Issue", 52, "What has happened"),
        Column("eff", "Effect", 32, "Effect on the programme"),
        Column("owner", "Owner", 24, "Owner of the action"),
        Column("prio", "Priority", 10, "High, Medium or Low", values=["High", "Medium", "Low"]),
        Column("action", "Action", 46, "Action to resolve it"),
        Column("due", "Due", 12, "Date", kind="date"),
        Column("status", "Status", 11, "Status", values=RAID["statuses"], status=True),
    ], [dict(zip(("id", "issue", "eff", "owner", "prio", "action", "due", "status"), i)) for i in RAID["issues"]],
        description="Current issues of the programme")
    fills(ws, 5, 5, len(RAID["issues"]), RATING_COLOURS)
    wb.sheet("Dependencies", [
        Column("id", "ID", 9, "Dependency identifier (same as the Project Plan)"),
        Column("type", "Type", 12, "Decision, data, environment, integration or resource"),
        Column("what", "Dependency", 56, "What BDOI provides or decides"),
        Column("by", "Provided by", 28, "Owner"),
        Column("need", "Needed by", 12, "Date", kind="date"),
        Column("for", "Needed for", 32, "Activity that waits for it"),
        Column("late", "Effect if late", 36, "Effect"),
        Column("state", "State", 11, "On track, At risk, Late or Delivered",
               values=["On track", "At risk", "Late", "Delivered"]),
    ], [dict(zip(("id", "type", "what", "by", "need", "for", "late"), x), state="On track")
        for x in PLAN["dependencies"]], description="Dependencies on BDOI (same list as the Project Plan)")
    return wb.save(OUT_DELIVERY / f"BIBS_Delivery_BRD-00_Risk_Register_RAID_Log_v{VERSION}.xlsx")


# ------------------------------------------------------------------------------------------------ Word tables

def render_milestones(_: str) -> list[str]:
    rows = [[m[0], m[1], m[2], m[3], m[5], m[6]] for m in PLAN["milestones"]]
    return table(["ID", "Date", "Milestone", "Drop", "Owner", "Status"], rows, "1.2,1.9,7,2,3.6,1.2",
                 "Programme milestones", size="8", status="Status")


def render_timeline(_: str) -> list[str]:
    rows = [[r[0], r[1], r[2], r[4], r[5], r[6]] for r in PLAN["timeline"]]
    return table(["ID", "Stream", "Activity", "Start", "End", "Owner"], rows, "1.1,2.3,6.6,1.9,1.9,3.4",
                 "Activities of the timeline", size="8")


def render_signoffs(_: str) -> list[str]:
    rows = [[s[0], s[1], s[3], s[4], s[5]] for s in PLAN["signoffs"]]
    return table(["Document set or decision", "BRDs", "Version", "Due", "Signatories"], rows, "3.8,2.6,2.6,1.9,6",
                 "Sign-offs", size="8")


def render_test_windows(_: str) -> list[str]:
    rows = [[t[0], t[2], t[3], t[4], t[5], t[6]] for t in PLAN["test_windows"]]
    return table(["Window", "Start", "End", "Env.", "Entry criteria", "Exit criteria"], rows,
                 "2.6,1.8,1.8,1.4,4.6,4.6", "SIT, UAT, trial migrations and the performance test", size="7.5")


def render_cutover(_: str) -> list[str]:
    rows = [[c[0], c[1], c[2], c[3], c[4], c[5]] for c in PLAN["cutover"]]
    return table(["Phase", "Name", "Days", "Dates", "Main activities", "Owner"], rows, "1.3,2.5,1.8,2.2,6,2.8",
                 "Cut-over and hypercare phases", size="8")


def render_dependencies(_: str) -> list[str]:
    rows = [[x[0], x[1], x[2], x[3], x[4], x[6]] for x in PLAN["dependencies"]]
    return table(["ID", "Type", "Dependency", "Provided by", "Needed by", "Effect if late"], rows,
                 "1.4,2.1,5.3,2.8,1.9,3.5", "Dependencies on BDOI", size="7.5")


def render_risks(arg: str) -> list[str]:
    rs = sorted(RAID["risks"], key=score, reverse=True)
    if arg.strip() == "high":
        rs = [r for r in rs if rating(r) == "High"]
    rows = [[r["id"], r["title"], f"{r['prob']} x {r['impact']} = {score(r)}", rating(r), r["owner"], r["mitigation"]]
            for r in rs]
    return table(["ID", "Risk", "P x I", "Rating", "Owner", "Mitigation"], rows, "1.1,3.4,1.6,1.4,3,6.5",
                 "Programme risks by score" if arg.strip() != "high" else "High risks", size="7.5")


def render_raci(arg: str) -> list[str]:
    areas = [a.strip() for a in arg.split(",") if a.strip()]
    cols = ["BPO", "BO", "KU", "BIT", "ISEC", "APS", "OPS", "IPM", "ISA", "IDEV", "IQA", "ICT"]
    acts = [x for x in RACI["activities"] if not areas or x[0] in areas]
    rows = [[a[1]] + [raci_map(a[3]).get(c, "") for c in cols] for a in acts]
    return table(["Activity"] + cols, rows, "6.2" + ",0.95" * len(cols),
                 "RACI extract (all roles in the RACI workbook)", size="7.5")


def render_roles(_: str) -> list[str]:
    rows = [[c, RACI["roles"][c][0], RACI["roles"][c][1], RACI["roles"][c][2]] for c in RACI["order"]]
    return table(["Code", "Role", "Organisation", "Responsibility"], rows, "1.2,4,2.8,8.6", "Roles of the programme",
                 size="8")


def render_personas(_: str) -> list[str]:
    rows = [[p["name"], p["brds"], p["home"], str(p["users"]), len(p["tasks"])] for p in GUIDES["personas"]]
    return table(["Persona", "BRDs", "Home page", "Users", "Main tasks"], rows, "3,4.4,4.2,3.6,1.4",
                 "Main personas of BIBS", size="8")


def render_qrg_index(_: str) -> list[str]:
    rows = [[p["name"], f"BIBS_Adoption_BRD-00_QRG_{p['file']}_v{VERSION}", "; ".join(t["title"] for t in p["tasks"])]
            for p in GUIDES["personas"]]
    return table(["Persona", "Quick-reference guide", "Main tasks covered"], rows, "3.4,5.6,7.6",
                 "Quick-reference guides", size="8")


RENDER = {k[len("render_"):]: v for k, v in globals().items() if k.startswith("render_")}


def expand(lines: list[str]) -> list[str]:
    out: list[str] = []
    for line in lines:
        m = PLACEHOLDER.match(line.strip())
        if m:
            out += RENDER[m.group(1)](m.group(2))
        else:
            out.append(line)
    return out


# ------------------------------------------------------------------------------------------------ Word documents

DOCS = {
    "PROJECT_PLAN_SUMMARY.md": OUT_DELIVERY / f"BIBS_Delivery_BRD-00_Project_Plan_Summary_v{VERSION}.docx",
    "DELIVERY_METHODOLOGY.md": OUT_DELIVERY / f"BIBS_Delivery_BRD-00_Delivery_Methodology_v{VERSION}.docx",
    "KNOWLEDGE_TRANSFER_PLAN.md": OUT_ADOPTION / f"BIBS_Adoption_BRD-00_Knowledge_Transfer_Plan_v{VERSION}.docx",
    "OCM_TRAINING_FRAMEWORK.md":
        OUT_ADOPTION / f"BIBS_Adoption_BRD-00_Change_Management_and_Training_Framework_v{VERSION}.docx",
    "USER_MANUAL_FRAMEWORK.md": OUT_ADOPTION / f"BIBS_Adoption_BRD-00_User_Manual_Framework_v{VERSION}.docx",
}


def build_doc(src: Path, target: Path, pdf: bool, lines: list[str] | None = None,
              front: dict[str, Any] | None = None) -> Path:
    if lines is None:
        front, lines = load_source(src)
    doc = BdoiDocument(meta_from(front), h1_page_break=bool(front.get("h1_page_break", True)), base_dir=HERE)
    doc.cover()
    doc.front_matter()
    render_body(doc, expand(lines))
    target.parent.mkdir(parents=True, exist_ok=True)
    return doc.publish(target, pdf=pdf)[0]


def qrg_source(p: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    front = {
        "title": f"Quick-Reference Guide: {p['name']}",
        "subtitle": "What you do in BIBS, your main tasks step by step, tips and where to get help",
        "doc_type": "Quick-Reference Guide", "doc_code": "QRG", "brd": "BRD-00", "name": p["name"],
        "doc_id": f"BIBS-QRG-{p['key']}", "version": VERSION, "date": DATE, "status": "Issued for BDOI review",
        "header_title": f"QRG {p['name']}", "h1_page_break": False,
        "control": [{"version": VERSION, "date": "08 Oct 2026", "author": "iorta TechNXT change and training lead",
                     "reviewer": "iorta TechNXT Business Analyst lead", "approver": "BIBS Product Owner (pending)",
                     "change": "First issue, from the FRS walkthroughs and functional requirements"}],
        "distribution": [
            {"name": "BIBS Product Owner", "role": "Approver", "organisation": "BDOI", "purpose": "Approval"},
            {"name": f"Key users of the {p['name']} persona", "role": "Reviewer", "organisation": "BDOI",
             "purpose": "Check against the screens in UAT"},
            {"name": "BDOI trainers", "role": "User", "organisation": "BDOI", "purpose": "End-user training"}],
    }
    L: list[str] = []
    L += ["# Your role in BIBS", "", p["summary"].strip(), ""]
    L += ["```keyvalues", yaml.safe_dump({"Persona": p["name"], "BRDs": p["brds"], "Home page": p["home"],
                                           "Sign-in": "BDO single sign-on (EIAM); MFA as set by Information Security"},
                                          sort_keys=False, allow_unicode=True).strip(), "```", ""]
    L += ["## Your menus", "", "The sidebar shows only the screens your roles allow. Your main menus:", ""]
    L += [f"- {m}" for m in p["menus"]] + [""]
    L += ["# Your main tasks", "",
          "Each task lists the steps in order, the screen (ID and name of the screen specification) and the step of "
          "the FRS walkthrough or the functional requirement it follows. The FRS walkthroughs show every step with "
          "its screenshot.", ""]
    for t in p["tasks"]:
        L += [f"## {t['title']}", "", f"Reference: {t['ref']}.", ""]
        rows = [[str(i), s[0], s[1], s[2]] for i, s in enumerate(t["steps"], start=1)]
        L += table(["No.", "What you do", "Screen", "Ref."], rows, "0.8,9.4,4.2,2.2", t["title"], size="8.5")
    L += ["# Tips", ""] + [f"- {x}" for x in p["tips"]] + [""]
    L += ["# Where to get help", ""] + [f"- {x}" for x in GUIDES["help_common"]] + [""]
    L += ["Report a problem with: the record number, the screen ID or name, what you did, the message shown word "
          "for word and the time.", ""]
    return front, L


def build_qrgs(pdf: bool) -> list[Path]:
    out = []
    for p in GUIDES["personas"]:
        front, lines = qrg_source(p)
        target = OUT_ADOPTION / f"BIBS_Adoption_BRD-00_QRG_{p['file']}_v{VERSION}.docx"
        out.append(build_doc(HERE / "guides.yaml", target, pdf, lines=lines, front=front))
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--only", help="xlsx, docs, qrg or a Word source name")
    args = ap.parse_args()
    problems = check()
    for p in problems:
        print("problem:", p)
    if args.check or problems:
        print(f"check: {len(problems)} problem(s)")
        return 1 if problems else 0
    pdf = not args.no_pdf
    only = args.only
    built: list[Path] = []
    if only in (None, "xlsx"):
        gantt_png(HERE / "figures" / "pp_gantt.png")
        built += [project_plan_xlsx(), raci_xlsx(), raid_xlsx()]
    for src, target in DOCS.items():
        if only in (None, "docs", src):
            built.append(build_doc(HERE / src, target, pdf))
    if only in (None, "qrg"):
        built += build_qrgs(pdf)
    for b in built:
        print(f"{b.relative_to(REPO)}  {b.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
