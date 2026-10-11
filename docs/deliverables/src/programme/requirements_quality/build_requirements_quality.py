"""Builds the requirements and quality documents of the BIBS programme (BRD-00).

Outputs (docs/deliverables/out/Programme/Requirements/ and .../Quality/):
  * BIBS_SRS_BRD-00_System_Requirements_Specification_v<v>.docx          from SRS_BIBS.md
  * BIBS_NFR_BRD-00_Non_Functional_Requirements_v<v>.docx                from NFR_CATALOGUE.md
  * BIBS_NFR_BRD-00_Non_Functional_Requirements_Register_v<v>.xlsx       from nfr_catalogue.yaml
  * BIBS_RTM_BRD-00_Requirements_Traceability_Matrix_v<v>.xlsx           from the FRS and the test plans
  * BIBS_TestStrategy_BRD-00_Test_Strategy_v<v>.docx                     from TEST_STRATEGY.md

Usage
  python docs/deliverables/src/programme/requirements_quality/build_requirements_quality.py --check
  python docs/deliverables/src/programme/requirements_quality/build_requirements_quality.py [--previews] [--no-pdf]

Inputs
  * the FRS of every BRD (brand.src_dir(brd)/FRS_*.md; BRD-13: HANDBOOK_*.md): front matter (version, date), the
    ```fr blocks (FR, title, BRD IDs, priority, rules) and the tables of the chapter "Traceability" (BRD ID, page,
    requirement, FR, screen, test conditions); the appendices "User-story view" and "Storyboard index" (row counts);
  * the test plans brdNN_cases.yaml, loaded with the test-plan builder (src/testplans/build_test_plan.py, read only):
    test conditions, test cases with their type, polarity, screen path and screen ID;
  * nfr_catalogue.yaml (this folder): the NFR rows and the reconciliation with the FRS, the register and the IER;
  * src/alignment/alignment_data.yaml (read only): the integration inventory for the SRS interface chapter.

The Word sources use two kinds of placeholders, expanded before the build: a line <!-- rq:<name> --> becomes a
table (placeholders()), and a token {{name}} or {{nfr:<NFR ID>}} becomes a figure or an NFR target (figures()).
Every row of the matrix states the FRS version it comes from. --check runs the wording rules of the client pack
(restricted words, development-status wording) on the expanded sources and on the workbook cells.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
sys.path.insert(0, str(REPO / "docs" / "deliverables" / "src" / "testplans"))
import brand  # noqa: E402
import check_pack  # noqa: E402
from bdoi_docx import BdoiDocument, lint_source, load_source, meta_from, output_path, render_body  # noqa: E402
from bdoi_xlsx import BdoiWorkbook, Column  # noqa: E402

VERSION = "1.0"
DATE = "8 October 2026"
OUT_REQ = brand.OUT_DIR / brand.DROPS["Programme"]["folder"] / "Requirements"
OUT_QA = brand.OUT_DIR / brand.DROPS["Programme"]["folder"] / "Quality"
RTM_NAME = f"BIBS_RTM_BRD-00_Requirements_Traceability_Matrix_v{VERSION}.xlsx"
NFR_XLSX = f"BIBS_NFR_BRD-00_Non_Functional_Requirements_Register_v{VERSION}.xlsx"
DOCS = {  # source -> output folder and file
    "SRS_BIBS.md": (OUT_REQ, f"BIBS_SRS_BRD-00_System_Requirements_Specification_v{VERSION}.docx"),
    "NFR_CATALOGUE.md": (OUT_REQ, f"BIBS_NFR_BRD-00_Non_Functional_Requirements_v{VERSION}.docx"),
    "TEST_STRATEGY.md": (OUT_QA, f"BIBS_TestStrategy_BRD-00_Test_Strategy_v{VERSION}.docx"),
}
BRDS = [f"BRD-{n:02d}" for n in range(14)]
BRD_TITLES = {"BRD-00": "Core Replacement (umbrella)", **brand.BRD_NAMES}
PLACEHOLDER = re.compile(r"^<!--\s*rq:(\w+)\s*-->\s*$")
TOKEN = re.compile(r"\{\{([\w:-]+)\}\}")
COVERAGE = ["COVERED", "NFR", "FR ONLY", "CLARIFICATION", "OUT", "REPLACED", "REMOVED", "GAP"]
COVERAGE_TEXT = {
    "COVERED": "Traced to at least one FR and at least one test case",
    "NFR": "Non-functional requirement traced to the NFR chapter of the FRS and to the NFR catalogue",
    "FR ONLY": "Traced to an FR that has no test case yet (umbrella FRs; cases come from the function test plans)",
    "CLARIFICATION": "The FRS proposes a rule that waits for a BDOI decision (CLR item of the FRS)",
    "OUT": "Out of scope of phase 1 (by the BRD, by a BDOI decision or reinsurance phase 2), as the FRS records",
    "REPLACED": "Replaced by a BIBS function (legacy-system reference); the replacing FR is named",
    "REMOVED": "Removed from scope by the new BRD version; the FR keeps its ID",
    "GAP": "No FR named in the traceability chapter",
}
COVERAGE_FILL = {"COVERED": (brand.SUCCESS_BG, brand.SUCCESS), "NFR": (brand.SUCCESS_BG, brand.SUCCESS), "FR ONLY": (brand.AMBER_BG, brand.AMBER),
                 "CLARIFICATION": (brand.AMBER_BG, brand.AMBER),
                 "OUT": (brand.DIRTY_WHITE, brand.MUTED), "REPLACED": (brand.DIRTY_WHITE, brand.MUTED),
                 "REMOVED": (brand.DIRTY_WHITE, brand.MUTED), "GAP": (brand.DANGER_BG, brand.DANGER)}
NFR_FILL = {"ALIGNED": (brand.SUCCESS_BG, brand.SUCCESS), "PROPOSED": (brand.BG_BLUE, brand.HEADER_BLUE),
            "DECISION": (brand.AMBER_BG, brand.AMBER)}
# Production Reconciliation of BRD-02 is a Drop 2 item (brand.DROP_SHARED); other items keep the primary drop.
DROP_BY_PREFIX = {("BRD-02", "PRCID"): "Drop 2"}


# --------------------------------------------------------------------------------------------- FRS reading

@dataclass
class Frs:
    brd: str
    sources: list[Path]
    version: str
    date: str
    doc_name: str
    status: str
    frs: "OrderedDict[str, dict[str, Any]]" = field(default_factory=OrderedDict)
    trace: list[dict[str, Any]] = field(default_factory=list)
    user_stories: int = 0
    storyboard: int = 0
    rules: Counter = field(default_factory=Counter)

    @property
    def label(self) -> str:
        return f"FRS {self.brd} v{self.version}" if self.brd != "BRD-13" else f"Handbook BRD-13 v{self.version}"


def front_matter(text: str) -> dict[str, Any]:
    m = re.match(r"---\n(.*?)\n---\n", text, re.S)
    return yaml.safe_load(m.group(1)) if m else {}


def chapter(lines: list[str], title_re: str) -> list[str]:
    """The lines of the first H1 chapter whose title matches title_re, up to the next H1."""
    out: list[str] = []
    inside = False
    for line in lines:
        if line.startswith("# "):
            if inside:
                break
            inside = bool(re.match(title_re, line[2:].strip()))
            continue
        if inside:
            out.append(line)
    return out


def tables(lines: list[str]) -> list[tuple[str, list[str], list[list[str]]]]:
    """Pipe tables of a chapter: (caption, headers, rows)."""
    out = []
    caption = ""
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        m = re.search(r'caption="([^"]*)"', line) if line.startswith("<!-- table:") else None
        if m:
            caption = m.group(1)
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|(\s*:?-+:?\s*\|)+\s*$", lines[i + 1].strip()):
            headers = [c.strip() for c in line.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append((caption, headers, rows))
            caption = ""
            continue
        i += 1
    return out


def table_rows(lines: list[str]) -> int:
    return sum(len(rows) for _, _, rows in tables(lines))


def expand_ids(cell: str, kind: str) -> list[str]:
    """'FR-PM-040, 041, 042' -> FR-PM-040, FR-PM-041, FR-PM-042; 'FR-UA-060 to 063' -> 060..063.
    kind 'FR' or 'TC' (test conditions TC-PM-001.1, 001.2)."""
    num = r"\d{3}[a-z]?" if kind == "FR" else r"\d{3}\.\d+"
    full = re.compile(rf"\b({kind}-[A-Z]{{2,4}})-({num})\b")
    short = re.compile(rf"(?<![\w.-])({num})\b")
    out: list[str] = []
    prefix = ""
    for part in re.split(r"[,;]", cell):
        part = part.strip()
        rng = re.match(rf"^(?:({kind}-[A-Z]{{2,4}})-)?({num})\s+to\s+(?:{kind}-[A-Z]{{2,4}}-)?({num})\b", part)
        if rng:
            prefix = rng.group(1) or prefix
            a, b = rng.group(2), rng.group(3)
            if kind == "FR" and prefix:
                for n in range(int(a[:3]), int(b[:3]) + 1):
                    out.append(f"{prefix}-{n:03d}")
            elif prefix:
                base, lo, hi = a.split(".")[0], int(a.split(".")[1]), int(b.split(".")[1])
                if b.split(".")[0] == base:
                    out += [f"{prefix}-{base}.{n}" for n in range(lo, hi + 1)]
                else:
                    out += [f"{prefix}-{a}", f"{prefix}-{b}"]
            continue
        found = full.findall(part)
        if found:
            for p, n in found:
                prefix = p
                out.append(f"{p}-{n}")
            continue
        m = short.match(part)
        if m and prefix:
            out.append(f"{prefix}-{m.group(1)}")
    seen: list[str] = []
    for x in out:
        if x not in seen:
            seen.append(x)
    return seen


def brd_ids(ref: str) -> list[str]:
    text = re.sub(r"\([^)]*\)", "", str(ref))
    return [p.strip() for p in re.split(r"[;,]", text) if p.strip()]


def read_frs(brd: str) -> Frs | None:
    folder = brand.src_dir(brd)
    files = sorted(folder.glob("FRS_*.md")) or sorted(folder.glob("HANDBOOK_*.md"))
    if not files:
        return None
    text0 = files[0].read_text(encoding="utf-8")
    front = front_matter(text0)
    doc = Frs(brd, files, str(front.get("version", "")), str(front.get("date", "")),
              output_path(front, files[0]).name, str(front.get("status", "")))
    for path in files:
        text = path.read_text(encoding="utf-8")
        for block in re.findall(r"```(?:fr|requirement)\n(.*?)```", text, re.S):
            fr = yaml.safe_load(block)
            refs = fr.get("brd") or []
            refs = refs if isinstance(refs, list) else [refs]
            ids: list[str] = []
            for r in refs:
                ids += [b for b in brd_ids(r) if b not in ids]
            blob = " ".join(str(x) for x in [fr.get("title", "")] + list(fr.get("description") or []))
            for rule in fr.get("rules") or []:
                if isinstance(rule, list) and len(rule) > 2:
                    doc.rules[str(rule[2]).strip()] += 1
            doc.frs[fr["id"]] = {
                "id": fr["id"], "title": str(fr.get("title", "")).strip(), "brd": ids,
                "priority": str(fr.get("priority", "")), "screens": str(fr.get("screens", "") or ""),
                "removed": bool(re.search(r"removed from (?:the )?scope", blob, re.I)), "brd_code": brd}
        for m in re.finditer(r"^#+\s+(FR-[A-Z]{2,4}-\d{3}[a-z]?)\s+(.*?removed from (?:the )?scope.*?)\s*(?:\{-\})?$",
                             text, re.M | re.I):
            if m.group(1) not in doc.frs:
                doc.frs[m.group(1)] = {"id": m.group(1), "title": m.group(2).strip(), "brd": [], "priority": "",
                                       "screens": "", "removed": True, "brd_code": brd}
        lines = text.splitlines()
        doc.trace += parse_trace(brd, chapter(lines, r"Traceability\b"))
        doc.user_stories += table_rows(chapter(lines, r"(?:Appendix: )?User-story view"))
        doc.storyboard += table_rows(chapter(lines, r"(?:Appendix: )?Storyboard index"))
    return doc


def _col(headers: list[str], *names: str) -> int | None:
    low = [h.lower().strip("* ") for h in headers]
    for n in names:
        if n in low:
            return low.index(n)
    return None


def parse_trace(brd: str, lines: list[str]) -> list[dict[str, Any]]:
    """Rows of the traceability tables: one per BRD requirement ID (line IDs of BRD-6 exploded)."""
    out: list[dict[str, Any]] = []
    for caption, headers, rows in tables(lines):
        c_id = _col(headers, "brd id", "id")
        c_lines = _col(headers, "line ids")
        c_fr = _col(headers, "fr")
        c_repl = _col(headers, "replaced by")
        if (c_id is None and c_lines is None) or (c_fr is None and c_repl is None):
            continue
        c_page = _col(headers, "page", "pages")
        c_req = _col(headers, "activity and requirement", "capability (core brd page)", "requirement", "brd function")
        c_scr = _col(headers, "screen", "main screen")
        c_tc = _col(headers, "test cases")
        c_core = _col(headers, "core br")
        c_cov = _col(headers, "covered by")
        req_carry = ""
        for row in rows:
            cell = (lambda i: row[i].replace("**", "").strip() if i is not None and i < len(row) else "")
            req = cell(c_req)
            if c_lines is not None:
                req = req or req_carry
                req_carry = req
                ids = [x.strip() for x in cell(c_lines).split(",") if x.strip()]
            else:
                ids = [cell(c_id)] if cell(c_id) else []
            if c_core is not None and cell(c_core) not in ("", "-"):
                req = f"{req} (umbrella {cell(c_core)})"
            if c_cov is not None and cell(c_cov) not in ("", "-"):
                req = f"{req}; covered by {cell(c_cov)}"
            fr_cell = cell(c_fr) if c_fr is not None else cell(c_repl)
            for raw in ids:
                if raw in ("", "-") or raw.lower().startswith("total"):
                    continue
                page = cell(c_page)
                m = re.search(r"\((p\.[^)]*)\)", raw) or re.search(r"\((p\.[^)]*)\)", req)
                if m and not page:
                    page = m.group(1)
                bid = re.sub(r"\s*\(p\.[^)]*\)", "", raw).strip()
                status = ""
                if re.search(r"\(removed\)", bid, re.I):
                    bid = re.sub(r"\s*\(removed\)", "", bid, flags=re.I).strip()
                    status = "REMOVED"
                elif c_repl is not None:
                    status = "REPLACED"
                elif re.search(r"\bOUT\b|out of scope|not in scope|phase 2", fr_cell, re.I):
                    status = "OUT"
                elif not expand_ids(fr_cell, "FR") and re.search(r"\bCLR-", fr_cell):
                    status = "CLARIFICATION"
                elif not expand_ids(fr_cell, "FR") and re.search(r"\b(?:section|chapter) \d", fr_cell, re.I):
                    status = "NFR"
                out.append({
                    "brd": brd, "id": bid, "page": page, "req": req, "frs": expand_ids(fr_cell, "FR"),
                    "fr_text": fr_cell, "screen": cell(c_scr), "conds": expand_ids(cell(c_tc), "TC"),
                    "status": status, "table": caption})
    return out


# --------------------------------------------------------------------------------------------- test plans

@dataclass
class Plans:
    by_brd: dict[str, Any]
    cases_by_fr: dict[str, list[Any]]
    cases_by_cond: dict[str, list[Any]]


def load_plans() -> Plans:
    import build_test_plan as tp  # noqa: PLC0415  (read-only use of the shared test-plan builder)

    by_brd: dict[str, Any] = {}
    by_fr: dict[str, list[Any]] = defaultdict(list)
    by_cond: dict[str, list[Any]] = defaultdict(list)
    for brd in BRDS:
        files = sorted(brand.src_dir(brd).glob("brd*_cases.yaml"))
        if not files:
            continue
        plan = tp.load(files[0])
        by_brd[brd] = plan
        for c in plan.cases:
            if c.fr:
                by_fr[c.fr].append(c)
            if c.cond and c.cond != "-":
                by_cond[c.cond].append(c)
    return Plans(by_brd, by_fr, by_cond)


def _brd_match(bid: str, case_brd: list[str]) -> bool:
    b = bid.replace("BRD ", "").strip()
    for x in case_brd:
        x = x.replace("BRD ", "").strip()
        if x == b or b.startswith(x + ".") or x.startswith(b + "."):
            return True
    return False


def compact_cases(ids: list[str]) -> str:
    """TC-PM-040.1-01, TC-PM-040.1-02, TC-PM-041.1-01 -> 'TC-PM-040.1-01/02, 041.1-01'."""
    groups: "OrderedDict[str, list[str]]" = OrderedDict()
    others: list[str] = []
    for i in ids:
        m = re.match(r"^(TC-[A-Z]+)-(\d{3}\.\d+)-(\d+)$", i)
        if m:
            groups.setdefault(f"{m.group(1)}|{m.group(2)}", []).append(m.group(3))
        else:
            others.append(i)
    parts: list[str] = []
    last_prefix = ""
    for key, seqs in groups.items():
        prefix, cond = key.split("|")
        head = f"{prefix}-{cond}" if prefix != last_prefix else cond
        last_prefix = prefix
        parts.append(f"{head}-{'/'.join(seqs)}")
    return ", ".join(parts + others)


# --------------------------------------------------------------------------------------------- the matrix

def drop_of_row(brd: str, bid: str) -> str:
    for (b, prefix), drop in DROP_BY_PREFIX.items():
        if b == brd and bid.startswith(prefix):
            return drop
    return brand.drop_of(brd)


def build_matrix(docs: dict[str, Frs], plans: Plans) -> list[dict[str, Any]]:
    all_frs = {k: v for d in docs.values() for k, v in d.frs.items()}
    rows = []
    for brd in BRDS:
        doc = docs.get(brd)
        if not doc:
            continue
        plan = plans.by_brd.get(brd)
        for t in doc.trace:
            frs = t["frs"]
            known = [f for f in frs if f in all_frs]
            cases: list[Any] = []
            basis = ""
            if t["conds"]:
                for cond in t["conds"]:
                    cases += plans.cases_by_cond.get(cond, [])
                basis = "Test conditions of the FRS"
            if not cases:
                for f in frs:
                    cases += [c for c in plans.cases_by_fr.get(f, []) if c.cond != "-" and _brd_match(t["id"], c.brd)]
                basis = "BRD ID of the test cases" if cases else basis
            if not cases:
                for f in frs:
                    cases += [c for c in plans.cases_by_fr.get(f, []) if c.cond != "-"]
                basis = "FR of the test cases" if cases else ""
            # screen and message cases of the sign-off set for the same FRs
            extra = [c for f in frs for c in plans.cases_by_fr.get(f, []) if c.cond == "-" and _brd_match(t["id"], c.brd)]
            seen: set[str] = set()
            cases = [c for c in cases + extra if not (c.id in seen or seen.add(c.id))]
            status = t["status"]
            if not status:
                if not frs:
                    status = "GAP"
                elif known and all(all_frs[f]["removed"] for f in known):
                    status = "REMOVED"
                elif cases:
                    status = "COVERED"
                else:
                    status = "FR ONLY"
            titles = "; ".join(f"{f} {all_frs[f]['title']}" for f in frs if f in all_frs)
            prio = sorted({all_frs[f]["priority"].split("(")[0].strip() for f in known if all_frs[f]["priority"]})
            screen_ids = sorted({c.screen_id for c in cases if getattr(c, "screen_id", "")})
            screens_tested = sorted({c.screen.split(" > ")[-1] for c in cases if c.screen})
            rows.append({
                "brd": brd, "module": BRD_TITLES[brd], "drop": drop_of_row(brd, t["id"]), "id": t["id"],
                "page": t["page"], "req": t["req"] or (titles.split("; ")[0].split(" ", 1)[-1] if titles else ""),
                "frs": ", ".join(frs) if frs else t["fr_text"], "titles": titles, "prio": ", ".join(prio),
                "screen": t["screen"], "conds": ", ".join(t["conds"]), "cases": compact_cases([c.id for c in cases]),
                "n": len(cases), "pos": sum(1 for c in cases if not c.negative),
                "neg": sum(1 for c in cases if c.negative), "screen_ids": ", ".join(screen_ids),
                "screens_tested": "; ".join(screens_tested[:8]) + (" ..." if len(screens_tested) > 8 else ""),
                "status": status, "basis": basis, "frs_doc": f"{doc.doc_name} (v{doc.version}, {doc.date})",
                "frs_version": doc.version,
                "tp_version": str(plan.meta.get("version")) if plan else "-", "table": t["table"]})
    return rows


def fr_register(docs: dict[str, Frs], plans: Plans) -> list[dict[str, Any]]:
    rows = []
    for brd in BRDS:
        doc = docs.get(brd)
        if not doc:
            continue
        for fr in doc.frs.values():
            cases = plans.cases_by_fr.get(fr["id"], [])
            conds = sorted({c.cond for c in cases if c.cond != "-"})
            rows.append({
                "brd": brd, "drop": brand.drop_of(brd), "fr": fr["id"], "title": fr["title"],
                "prio": fr["priority"], "brd_ids": ", ".join(fr["brd"]), "screens": fr["screens"],
                "conds": len(conds), "cases": len(cases), "pos": sum(1 for c in cases if not c.negative),
                "neg": sum(1 for c in cases if c.negative),
                "screen_ids": ", ".join(sorted({c.screen_id for c in cases if getattr(c, "screen_id", "")})),
                "state": "Removed from scope" if fr["removed"] else ("Tested" if cases else "No test case"),
                "frs_doc": f"{doc.doc_name} (v{doc.version})"})
    return rows


def summary(docs: dict[str, Frs], plans: Plans, matrix: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for brd in BRDS:
        doc = docs.get(brd)
        if not doc:
            continue
        rows = [r for r in matrix if r["brd"] == brd]
        st = Counter(r["status"] for r in rows)
        in_scope = len(rows) - st["OUT"] - st["REPLACED"] - st["REMOVED"] - st["CLARIFICATION"]
        plan = plans.by_brd.get(brd)
        traced = {f for r in rows for f in r["frs"].split(", ") if f.startswith("FR-")}
        own = [f for f in doc.frs if not doc.frs[f]["removed"]]
        out.append({
            "brd": brd, "module": BRD_TITLES[brd], "drop": brand.drop_of(brd), "frs_doc": doc.doc_name,
            "frs_version": doc.version, "frs_date": doc.date,
            "tp": f"v{plan.meta.get('version')}" if plan else "No test plan (umbrella; cases in the function test plans)",
            "ids": len(rows), "covered": st["COVERED"] + st["NFR"], "fr_only": st["FR ONLY"], "clr": st["CLARIFICATION"],
            "out": st["OUT"] + st["REPLACED"] + st["REMOVED"], "gap": st["GAP"],
            "pct": round(100.0 * (st["COVERED"] + st["NFR"]) / in_scope, 1) if in_scope else 0.0,
            "frs": len(own), "frs_traced": len([f for f in own if f in traced]),
            "frs_tested": len([f for f in own if plans.cases_by_fr.get(f)]),
            "cases": len(plan.cases) if plan else 0, "conds": len(plan.conditions) if plan else 0,
            "stories": doc.user_stories, "frames": doc.storyboard})
    return out


# --------------------------------------------------------------------------------------------- workbooks

def _fill_values(ws, col_letter: str, first: int, last: int, fills: dict[str, tuple[str, str]]) -> None:
    rng = f"{col_letter}{first}:{col_letter}{last + 500}"
    for value, (bg, fg) in fills.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{value}"'], stopIfTrue=True,
                                                      fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                                                      font=Font(color=fg, bold=True)))


def build_rtm(docs: dict[str, Frs], plans: Plans, matrix, frreg, summ) -> Path:
    wb = BdoiWorkbook("Requirements Traceability Matrix", doc_type="Requirements traceability matrix",
                      brd="BRD-00", version=VERSION, date=DATE,
                      subtitle="BRD requirement to FR, test cases, screens and drop, for every BRD of BIBS")
    wb.legend = [(k, v) for k, v in COVERAGE_TEXT.items()]
    wb.cover_notes = [
        f"Generated on {DATE} from the traceability chapter of each FRS and the test cases of each test plan. Each row "
        "names the FRS document and version it comes from; a re-issued FRS updates its rows when the matrix is next generated.",
        "Matrix: one row per BRD requirement ID. FR register: one row per FR (backward trace). Summary: coverage "
        "per BRD. Sources: the FRS and test plan version of each BRD."]
    tot = {k: sum(s[k] for s in summ) for k in ("ids", "covered", "fr_only", "clr", "out", "gap", "frs", "frs_traced",
                                                 "frs_tested", "cases", "conds", "stories", "frames")}
    in_scope = tot["ids"] - tot["out"] - tot["clr"]
    total_row = {"brd": "Total", "module": "All BRDs", "drop": "", "frs_doc": "", "frs_version": "", "frs_date": "",
                 "tp": "", **tot, "pct": round(100.0 * tot["covered"] / in_scope, 1) if in_scope else 0.0}
    ws = wb.sheet("Summary", [
        Column("brd", "BRD", 9, "BRD code"),
        Column("module", "Module", 24, "BRD name"),
        Column("drop", "Drop", 10, "Primary BDOI drop of the BRD (brand drop map)"),
        Column("frs_doc", "FRS document", 40, "FRS (or Handbook) whose traceability chapter gives the rows"),
        Column("frs_version", "FRS version", 9, "Version of that FRS when this matrix was generated"),
        Column("frs_date", "FRS date", 14, "Date of that FRS version"),
        Column("tp", "Test plan", 16, "Version of the test plan whose cases are linked"),
        Column("ids", "BRD IDs traced", 10, "Rows of the Matrix sheet for the BRD", kind="number"),
        Column("covered", "Covered", 10, "Rows with FR and at least one test case, or traced to the NFR catalogue", kind="number"),
        Column("fr_only", "FR only", 9, "Rows with FR but no test case", kind="number"),
        Column("clr", "Clarification", 11, "Rows waiting for a BDOI decision on a proposed rule", kind="number"),
        Column("out", "Out / replaced / removed", 12, "Rows out of scope, replaced or removed", kind="number"),
        Column("gap", "Gap", 7, "Rows without FR", kind="number"),
        Column("pct", "Coverage %", 10, "Covered / (BRD IDs traced - clarification - out / replaced / removed)",
               kind="percent"),
        Column("frs", "FRs in scope", 9, "FRs of the FRS not removed from scope", kind="number"),
        Column("frs_traced", "FRs traced from a BRD ID", 11, "FRs of the FRS named in the Matrix", kind="number"),
        Column("frs_tested", "FRs with test cases", 10, "FRs with at least one test case", kind="number"),
        Column("conds", "Test conditions", 10, "Test conditions of the test plan", kind="number"),
        Column("cases", "Test cases", 9, "Test cases of the test plan (incl. screen and message cases)",
               kind="number"),
        Column("stories", "User-story rows", 9, "Rows of the FRS appendix User-story view (0 = appendix follows "
               "with the next FRS issue)", kind="number"),
        Column("frames", "Storyboard frames", 9, "Rows of the FRS appendix Storyboard index", kind="number"),
    ], summ + [total_row], description="Coverage per BRD; the Matrix sheet gives every row", freeze_first_column=True)
    last = 4 + len(summ) + 1
    for c in ws[last]:
        c.font = Font(name=brand.FONT, bold=True, color=brand.HEADER_BLUE)
    ws = wb.sheet("Matrix", [
        Column("brd", "BRD", 8, "BRD code"),
        Column("drop", "Drop", 9, "BDOI drop of the requirement (Production Reconciliation of BRD-2 is Drop 2)"),
        Column("id", "BRD requirement ID", 16, "Requirement ID of the BRD (or the umbrella row, or the BRD line ID)"),
        Column("page", "Page", 9, "BRD page as the FRS cites it"),
        Column("req", "Requirement", 40, "BRD wording where the FRS table carries it, else the title of the first FR"),
        Column("frs", "FR", 18, "FRs that meet the requirement (traceability chapter of the FRS)"),
        Column("titles", "FR titles", 40, "Title of each FR"),
        Column("prio", "Priority", 10, "Priority of the FRs"),
        Column("screen", "Screen (FRS)", 24, "Main screen named in the traceability chapter"),
        Column("conds", "Test conditions (FRS)", 22, "Test conditions the FRS names for the requirement"),
        Column("cases", "Test cases", 34, "Test cases of the test plan: TC-<module>-<FR>.<condition>-<case>; "
               "SCR and MSG are the screen and message cases of the sign-off set"),
        Column("n", "Cases", 7, "Number of test cases", kind="number"),
        Column("pos", "Positive", 8, "Cases where the action succeeds", kind="number"),
        Column("neg", "Negative", 8, "Cases where BIBS refuses the action", kind="number"),
        Column("screen_ids", "Screen IDs", 18, "Screen IDs of the sign-off set (UX screen register) tested by the cases"),
        Column("screens_tested", "Screens tested", 30, "Last menu level of the screens the cases use"),
        Column("status", "Coverage", 11, "Coverage of the row", values=COVERAGE),
        Column("basis", "Link basis", 18, "How the cases were linked: test conditions named by the FRS, BRD ID of "
               "the cases, or FR of the cases"),
        Column("frs_doc", "FRS source (version)", 36, "FRS document, version and date the row comes from"),
        Column("tp_version", "Test plan version", 9, "Version of the test plan of the BRD"),
        Column("table", "FRS table", 26, "Caption of the traceability table in the FRS"),
    ], matrix, description="BRD requirement - FR - test cases - screens - drop; one row per BRD requirement ID")
    _fill_values(ws, "Q", 5, 4 + len(matrix), COVERAGE_FILL)  # column Q = Coverage
    wb.sheet("FR register", [
        Column("brd", "BRD", 8, "BRD code"),
        Column("drop", "Drop", 9, "Primary drop of the BRD"),
        Column("fr", "FR", 12, "Functional requirement ID"),
        Column("title", "Title", 40, "FR title"),
        Column("prio", "Priority", 14, "FR priority"),
        Column("brd_ids", "BRD IDs", 30, "BRD requirement IDs the FR traces (FR header)"),
        Column("screens", "Screens", 30, "Screens of the FR (FR header)"),
        Column("conds", "Conditions", 9, "Test conditions with cases", kind="number"),
        Column("cases", "Cases", 7, "Test cases of the FR", kind="number"),
        Column("pos", "Positive", 8, "Positive cases", kind="number"),
        Column("neg", "Negative", 8, "Negative cases", kind="number"),
        Column("screen_ids", "Screen IDs", 18, "Screen IDs of the sign-off set tested by the cases"),
        Column("state", "State", 14, "Tested, No test case or Removed from scope"),
        Column("frs_doc", "FRS source", 36, "FRS document and version"),
    ], frreg, description="Backward trace: every FR of every FRS with its BRD IDs and test cases")
    wb.sheet("Sources", [
        Column("brd", "BRD", 8, "BRD code"),
        Column("module", "Module", 24, "BRD name"),
        Column("frs_doc", "FRS document", 44, "Document whose traceability chapter and FRs feed the matrix"),
        Column("frs_version", "Version", 8, "Version"),
        Column("frs_date", "Date", 16, "Date"),
        Column("tp", "Test plan", 18, "Test plan version"),
    ], [{k: s[k] for k in ("brd", "module", "frs_doc", "frs_version", "frs_date", "tp")} for s in summ],
        description="The FRS and test plan version of each BRD when this matrix was generated")
    return wb.save(OUT_REQ / RTM_NAME)


def nfr_rows(nfr: dict[str, Any]) -> list[dict[str, Any]]:
    cats = nfr["categories"]
    return [{**n, "catname": cats[n["cat"]]} for n in nfr["nfrs"]]


def build_nfr_register(nfr: dict[str, Any]) -> Path:
    wb = BdoiWorkbook("Non-Functional Requirements Register", doc_type="NFR register", brd="BRD-00",
                      version=VERSION, date=DATE, subtitle="Measurable NFRs of BIBS with verification and owners")
    wb.legend = [(k, v) for k, v in nfr["statuses"].items()]
    wb.cover_notes = [
        "Companion of the Word document 'Non-Functional Requirements' of the same version. The targets here are the "
        "ones quoted in the System Requirements Specification and the Test Strategy.",
        "Precedence of sources: BDOI decision, IER workbook v20, BIBS-wide value of the discrepancy register v1.2, "
        "BRD values, iorta TechNXT proposal."]
    rows = nfr_rows(nfr)
    ws = wb.sheet("NFR register", [
        Column("id", "NFR ID", 12, "NFR-<category>-<nn>"),
        Column("catname", "Category", 18, "Category of the NFR"),
        Column("title", "Title", 24, "Short name"),
        Column("req", "Requirement", 40, "What BIBS must do"),
        Column("target", "Measurable target", 40, "The value accepted as met"),
        Column("measure", "How it is measured", 36, "Metric, conditions and evidence"),
        Column("method", "Verification method", 12, "Test, Analysis, Inspection or Walk-through",
               values=["Test", "Analysis", "Inspection", "Walk-through"]),
        Column("test", "Test type", 20, "Test type of the Test Strategy"),
        Column("stage", "Verified in", 26, "Test stage and dates"),
        Column("owner", "Owner (iorta TechNXT)", 22, "Role accountable for meeting the target"),
        Column("acceptor", "Accepted by (BDOI)", 22, "Role that accepts the evidence"),
        Column("source", "Source", 36, "BRD pages, IER, register items and design documents"),
        Column("frs", "FRS reconciliation", 36, "How the FRS non-functional chapters relate to the target"),
        Column("status", "Status", 11, "ALIGNED, DECISION or PROPOSED", values=list(nfr["statuses"])),
        Column("prio", "Priority", 9, "Must or Should", values=["Must", "Should"]),
    ], rows, description="One row per NFR; filter by category, status or test type")
    _fill_values(ws, "N", 5, 4 + len(rows), NFR_FILL)
    recon = [dict(zip(["theme", "brd", "frs", "ier", "cat", "action", "refs"], r)) for r in nfr["reconciliation"]]
    wb.sheet("Reconciliation", [
        Column("theme", "Theme", 16, "NFR theme"),
        Column("brd", "BRD values", 40, "Values of the BRDs with pages (register v1.2, NFR comparison sheet)"),
        Column("frs", "FRS statements", 32, "What the non-functional chapters of the FRS state"),
        Column("ier", "IER workbook v20", 28, "What the infrastructure estimate states"),
        Column("cat", "Catalogue target", 30, "NFR of this register"),
        Column("action", "Difference and action", 34, "What changes where"),
        Column("refs", "Register / questions", 18, "Register items and open questions"),
    ], recon, description="The catalogue against the FRS non-functional chapters, the register and the IER")
    cats = nfr["categories"]
    by = [{"cat": f"{k} {v}", "n": sum(1 for r in rows if r["cat"] == k),
           "must": sum(1 for r in rows if r["cat"] == k and r["prio"] == "Must"),
           **{s.lower(): sum(1 for r in rows if r["cat"] == k and r["status"] == s) for s in nfr["statuses"]}}
          for k, v in cats.items()]
    wb.sheet("Summary", [
        Column("cat", "Category", 34, "Category code and name"),
        Column("n", "NFRs", 8, "Number of NFRs", kind="number"),
        Column("must", "Must", 8, "Must-have NFRs", kind="number"),
        Column("aligned", "Aligned", 9, "Status ALIGNED", kind="number"),
        Column("decision", "BDOI decision", 10, "Status DECISION", kind="number"),
        Column("proposed", "Proposed", 9, "Status PROPOSED", kind="number"),
    ], by, description="NFRs per category and status")
    return wb.save(OUT_REQ / NFR_XLSX)


# --------------------------------------------------------------------------------------------- Word placeholders

def md_table(headers, rows, widths, caption, size=8.5, bold_first=True, status=None) -> list[str]:
    opts = f'widths={",".join(str(w) for w in widths)} caption="{caption}" size={size}'
    if bold_first:
        opts += " bold=first"
    if status:
        opts += f" status={status}"
    out = [f"<!-- table: {opts} -->", "| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    for r in rows:
        out.append("| " + " | ".join(str(c).replace("|", "/").replace("\n", " ") for c in r) + " |")
    out.append("")
    return out


def load_alignment() -> dict[str, Any]:
    p = brand.SRC_DIR / "alignment" / "alignment_data.yaml"
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else {}


@dataclass
class Data:
    docs: dict[str, Frs]
    plans: Plans
    matrix: list[dict[str, Any]]
    frreg: list[dict[str, Any]]
    summ: list[dict[str, Any]]
    nfr: dict[str, Any]
    align: dict[str, Any]


def load_all() -> Data:
    docs = {b: d for b in BRDS if (d := read_frs(b))}
    plans = load_plans()
    matrix = build_matrix(docs, plans)
    frreg = fr_register(docs, plans)
    summ = summary(docs, plans, matrix)
    nfr = yaml.safe_load((HERE / "nfr_catalogue.yaml").read_text(encoding="utf-8"))
    return Data(docs, plans, matrix, frreg, summ, nfr, load_alignment())


def nfr_by_id(d: Data) -> dict[str, dict[str, Any]]:
    return {n["id"]: n for n in d.nfr["nfrs"]}


def figures(d: Data) -> dict[str, Any]:
    s = d.summ
    tot = lambda k: sum(x[k] for x in s)  # noqa: E731
    in_scope = tot("ids") - tot("out") - tot("clr")
    cfg = sum(doc.rules["Configurable"] for doc in d.docs.values())
    fixed = sum(doc.rules["Fixed"] for doc in d.docs.values())
    plans = d.plans.by_brd
    cases = [c for p in plans.values() for c in p.cases]
    nfr = d.nfr["nfrs"]
    out = {
        "version": VERSION, "date": DATE,
        "frs_count": len(d.docs), "fr_total": sum(len(x.frs) for x in d.docs.values()),
        "brd_ids": tot("ids"), "covered": tot("covered"), "fr_only": tot("fr_only"), "out_ids": tot("out"),
        "clr_ids": tot("clr"),
        "gap": tot("gap"), "coverage_pct": f"{100.0 * tot('covered') / in_scope:.1f}" if in_scope else "0",
        "in_scope_ids": in_scope,
        "plans": len(plans), "cases": len(cases), "conds": sum(len(p.conditions) for p in plans.values()),
        "pos_cases": sum(1 for c in cases if not c.negative), "neg_cases": sum(1 for c in cases if c.negative),
        "screen_cases": sum(1 for c in cases if c.type == "Screen"),
        "msg_cases": sum(1 for c in cases if c.type == "Message"),
        "cfg_rules": cfg, "fixed_rules": fixed, "all_rules": cfg + fixed,
        "cfg_pct": f"{100.0 * cfg / (cfg + fixed):.0f}" if cfg + fixed else "0",
        "nfr_total": len(nfr), "nfr_must": sum(1 for n in nfr if n["prio"] == "Must"),
        "nfr_aligned": sum(1 for n in nfr if n["status"] == "ALIGNED"),
        "nfr_decision": sum(1 for n in nfr if n["status"] == "DECISION"),
        "nfr_proposed": sum(1 for n in nfr if n["status"] == "PROPOSED"),
        "stories_total": tot("stories"), "story_frs": sum(1 for x in s if x["stories"]),
        "integrations": len(d.align.get("integrations") or []),
        "rtm_rows": len(d.matrix), "fr_register_rows": len(d.frreg),
    }
    for n in nfr:
        out[f"nfr:{n['id']}"] = n["target"]
    return out


def placeholders(d: Data) -> dict[str, list[str]]:
    nf = nfr_by_id(d)
    cats = d.nfr["categories"]
    s = d.summ

    def frs_row(x):
        return [x["brd"], x["module"], x["drop"], f"v{x['frs_version']} ({x['frs_date']})", x["frs"], x["ids"],
                x["tp"].split(" (")[0], x["stories"] or "Next issue"]

    def nfr_table(cat):
        rows = [[n["id"], n["title"], n["target"], n["method"] + "; " + n["test"], n["owner"].replace("iorta ", ""),
                 n["status"]] for n in d.nfr["nfrs"] if n["cat"] == cat]
        return md_table(["ID", "NFR", "Measurable target", "Verification", "Owner", "Status"], rows,
                        [2.1, 2.5, 5.8, 2.8, 2.3, 2.1], f"{cats[cat]}", size=7.5, status="Status")

    def nfr_head(ids, caption):
        rows = [[i, nf[i]["title"], nf[i]["target"], nf[i]["status"]] for i in ids]
        return md_table(["NFR", "Topic", "Target", "Status"], rows, [2.0, 3.0, 9.4, 2.2], caption, size=8,
                        status="Status")

    integ = d.align.get("integrations") or []
    blocks: dict[str, list[str]] = {
        "frs_summary": md_table(
            ["BRD", "Module", "Drop", "FRS version", "FRs", "BRD IDs traced", "Test plan", "User stories"],
            [frs_row(x) for x in s], [1.3, 3.6, 1.6, 2.9, 1.0, 1.5, 1.6, 1.6],
            "FRS of each BRD: version, FRs, traced BRD requirement IDs, test plan and user-story view", size=8),
        "coverage": md_table(
            ["BRD", "Module", "BRD IDs", "Covered", "FR only", "Clarif.", "Out", "Gap", "Coverage %", "Cases"],
            [[x["brd"], x["module"], x["ids"], x["covered"], x["fr_only"], x["clr"], x["out"], x["gap"], x["pct"],
              x["cases"]] for x in s], [1.3, 3.8, 1.3, 1.4, 1.3, 1.3, 1.1, 1.0, 1.7, 1.3],
            "Coverage per BRD (Summary sheet of the traceability matrix)", size=8),
        "test_plans": md_table(
            ["BRD", "Module", "Drop", "Test plan", "FRs tested", "Conditions", "Test cases"],
            [[x["brd"], x["module"], x["drop"], x["tp"].split(" (")[0], x["frs_tested"], x["conds"], x["cases"]]
             for x in s if x["cases"]], [1.3, 4.6, 1.8, 1.8, 1.8, 1.8, 1.8],
            "Test plans per BRD (test conditions and cases as approved for SIT and UAT)", size=8),
        "interfaces": md_table(
            ["ID", "System", "Drop", "Direction", "Data exchanged"],
            [[i["id"], i["name"], i["drop"], i.get("direction", ""), i.get("data", "")] for i in integ],
            [1.4, 3.4, 2.2, 3.6, 6.0], "External interfaces of BIBS (integration inventory of the programme "
            "alignment pack)", size=7.5),
        "nfr_categories": md_table(
            ["Category", "NFRs", "Must", "Aligned", "BDOI decision", "Proposed"],
            [[f"{v}", sum(1 for n in d.nfr["nfrs"] if n["cat"] == k),
              sum(1 for n in d.nfr["nfrs"] if n["cat"] == k and n["prio"] == "Must"),
              *[sum(1 for n in d.nfr["nfrs"] if n["cat"] == k and n["status"] == st) for st in d.nfr["statuses"]]]
             for k, v in cats.items()], [6.2, 1.6, 1.6, 1.8, 2.2, 1.8], "NFRs per category and status", size=8.5),
        "nfr_reconciliation": md_table(
            ["Theme", "BRD values", "FRS statements", "IER workbook", "Catalogue", "Action"],
            [r[:6] for r in d.nfr["reconciliation"]], [1.9, 4.0, 3.0, 2.6, 2.6, 2.6],
            "Reconciliation with the FRS non-functional chapters and the IER workbook", size=7),
        "nfr_owners": md_table(
            ["Owner (iorta TechNXT)", "NFRs owned", "Accepted by (BDOI)"],
            [[o, ", ".join(n["id"] for n in d.nfr["nfrs"] if n["owner"] == o),
              "; ".join(sorted({n["acceptor"] for n in d.nfr["nfrs"] if n["owner"] == o}))]
             for o in sorted({n["owner"] for n in d.nfr["nfrs"]})], [3.4, 8.0, 5.2],
            "Ownership of the NFRs", size=7.5),
        "nfr_verification": md_table(
            ["Test type", "NFRs verified"],
            [[t, ", ".join(n["id"] for n in d.nfr["nfrs"] if n["test"].split(";")[0].strip() == t)]
             for t in sorted({n["test"].split(";")[0].strip() for n in d.nfr["nfrs"]})], [4.2, 12.4],
            "NFRs by test type (first test type of each NFR)", size=7.5),
        "nfr_headline": nfr_head(["NFR-PRF-01", "NFR-PRF-02", "NFR-PRF-03", "NFR-CAP-01", "NFR-CAP-02",
                                  "NFR-AVL-01", "NFR-AVL-02", "NFR-AVL-03", "NFR-REC-01", "NFR-REC-02",
                                  "NFR-REC-03", "NFR-SEC-01", "NFR-SEC-05", "NFR-ACC-01", "NFR-USE-01",
                                  "NFR-RSP-01", "NFR-RET-01", "NFR-MNT-02"],
                                 "Headline non-functional targets (full catalogue in the NFR document and register)"),
        "nft_targets": nfr_head(["NFR-PRF-01", "NFR-PRF-02", "NFR-PRF-03", "NFR-PRF-04", "NFR-PRF-05",
                                 "NFR-PRF-07", "NFR-CAP-01", "NFR-CAP-02", "NFR-CAP-05", "NFR-SCL-01",
                                 "NFR-SCL-04", "NFR-AVL-04", "NFR-AVL-05", "NFR-REC-01", "NFR-REC-02"],
                                "Targets of the non-functional tests (from the NFR register)"),
    }
    for cat in cats:
        blocks[f"nfr_{cat.lower()}"] = nfr_table(cat)
    return blocks


def expanded_source(d: Data, src: Path) -> Path:
    blocks = placeholders(d)
    figs = figures(d)
    lines: list[str] = []
    for line in src.read_text(encoding="utf-8").splitlines():
        m = PLACEHOLDER.match(line.strip())
        if m:
            if m.group(1) not in blocks:
                raise KeyError(f"{src.name}: unknown placeholder rq:{m.group(1)}")
            lines.extend(blocks[m.group(1)])
            continue
        lines.append(TOKEN.sub(lambda t: str(figs[t.group(1)]), line))
    tmp = src.with_name(f"_{src.stem}.expanded.md")
    tmp.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return tmp


def build_doc(d: Data, name: str, pdf: bool, keep_pdf: bool) -> tuple[Path, Path | None]:
    folder, file = DOCS[name]
    tmp = expanded_source(d, HERE / name)
    try:
        front, lines = load_source(tmp)
        doc = BdoiDocument(meta_from(front), h1_page_break=bool(front.get("h1_page_break", True)), base_dir=HERE)
        doc.cover()
        doc.front_matter()
        render_body(doc, lines)
        return doc.publish(folder / file, pdf=pdf, keep_pdf=keep_pdf)
    finally:
        tmp.unlink(missing_ok=True)


# --------------------------------------------------------------------------------------------- checks

def wording(text: str, where: str) -> list[str]:
    out = []
    for m in check_pack.RESTRICTED.finditer(text):
        out.append(f"{where}: restricted word '{m.group(0)}'")
    for line in text.splitlines():
        for label, pat in check_pack.BUILD_STATUS:
            for m in pat.finditer(line):
                out.append(f"{where}: {label} '{m.group(0).strip()}' in: {line.strip()[:110]}")
    return out


def check(d: Data) -> list[str]:
    problems: list[str] = []
    for name in DOCS:
        tmp = expanded_source(d, HERE / name)
        try:
            problems += [f"{name}: {p}" for p in lint_source(tmp)]
            body = tmp.read_text(encoding="utf-8")
            body = re.sub(r"<!--.*?-->", " ", body, flags=re.S)
            body = re.sub(r"(!\[[^\]]*\])\([^)]*\)(\{[^}]*\})?", r"\1", body)  # figure paths are not shown
            problems += wording(body.split("\n---", 2)[-1], name)
            for t in ("TBD", "lorem"):
                if re.search(rf"\b{t}\b", body, re.I):
                    problems.append(f"{name}: placeholder word {t}")
        finally:
            tmp.unlink(missing_ok=True)
    cells = []
    for r in d.matrix + d.frreg + d.summ + nfr_rows(d.nfr):
        cells += [str(v) for v in r.values()]
    cells += [str(c) for r in d.nfr["reconciliation"] for c in r]
    problems += wording("\n".join(cells), "workbook cells")
    ids = [n["id"] for n in d.nfr["nfrs"]]
    problems += [f"duplicate NFR ID {i}" for i, k in Counter(ids).items() if k > 1]
    for n in d.nfr["nfrs"]:
        miss = [k for k in ("id", "cat", "title", "req", "target", "measure", "method", "test", "stage", "owner",
                            "acceptor", "source", "frs", "status", "prio") if not n.get(k)]
        if miss:
            problems.append(f"{n.get('id')}: missing {', '.join(miss)}")
        if n["status"] not in d.nfr["statuses"]:
            problems.append(f"{n['id']}: status {n['status']}")
    return problems


def report(d: Data) -> str:
    f = figures(d)
    lines = [f"FRS read: {f['frs_count']} ({f['fr_total']} FRs); test plans: {f['plans']} ({f['cases']} cases)",
             f"matrix rows: {f['rtm_rows']}; covered {f['covered']}, FR only {f['fr_only']}, clarification "
             f"{f['clr_ids']}, out {f['out_ids']}, "
             f"gap {f['gap']}; coverage {f['coverage_pct']}%",
             f"NFRs: {f['nfr_total']} ({f['nfr_aligned']} aligned, {f['nfr_decision']} decision, "
             f"{f['nfr_proposed']} proposed)"]
    for x in d.summ:
        lines.append(f"  {x['brd']} v{x['frs_version']}: ids {x['ids']} cov {x['covered']} fr-only {x['fr_only']} "
                     f"out {x['out']} gap {x['gap']} ({x['pct']}%) cases {x['cases']} stories {x['stories']}")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="checks only")
    ap.add_argument("--previews", action="store_true", help="page previews of the outputs")
    ap.add_argument("--no-pdf", action="store_true", help="skip the PDF pass (no page numbers in the TOC)")
    ap.add_argument("--only", choices=["rtm", "nfr", "srs", "ts", "docs", "xlsx"], help="build one part only")
    args = ap.parse_args(argv)
    d = load_all()
    print(report(d))
    problems = check(d)
    for p in problems:
        print(p)
    if problems:
        return 1
    print("check: OK")
    if args.check:
        return 0
    import render  # noqa: PLC0415

    built: list[Path] = []
    if args.only in (None, "rtm", "xlsx"):
        built.append(build_rtm(d.docs, d.plans, d.matrix, d.frreg, d.summ))
    if args.only in (None, "nfr", "xlsx"):
        built.append(build_nfr_register(d.nfr))
    want = {"srs": ["SRS_BIBS.md"], "nfr": ["NFR_CATALOGUE.md"], "ts": ["TEST_STRATEGY.md"]}
    names = list(DOCS) if args.only in (None, "docs") else want.get(args.only or "", [])
    for name in names:
        path, pdf = build_doc(d, name, pdf=not args.no_pdf, keep_pdf=args.previews)
        built.append(path)
        if pdf and args.previews:
            print(f"{path.name}: {render.page_count(pdf)} pages")
            render.previews(pdf)
            pdf.unlink(missing_ok=True)
    bad = 0
    for p in built:
        print(f"built: {p.relative_to(REPO)} ({p.stat().st_size // 1024} KB)")
        for problem in wording("\n".join(check_pack.office_lines(p)), p.name):
            print(problem)
            bad += 1
    if args.previews:
        for p in built:
            if p.suffix == ".xlsx":
                pdf = render.to_pdf(p)
                print(f"{p.name}: {render.page_count(pdf)} pages")
                render.previews(pdf)
                pdf.unlink(missing_ok=True)
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
