"""Builds the BRD Discrepancy, Conflict, Impact and Clarification Register (client deliverable 2).

Inputs
  * discrepancy_register.yaml (this folder): register items, NFR comparison, owners, rules,
    open-question status and grouping;
  * the "Open questions" tables of the specs in docs/requirements/ and the XQ table of
    BDOI_CROSS_BRD_DECISIONS.md (read at build time, so the question list follows the specs).

Outputs (docs/deliverables/out/Registers/)
  * BIBS_Register_BRD-00_Discrepancies_and_Clarifications_v<version>.xlsx
    (Cover, README, Summary, Register, Open questions, NFR comparison);
  * BIBS_Register_BRD-00_Discrepancies_and_Clarifications_Summary_v<version>.pdf (the Summary sheet).

Usage
  python docs/deliverables/src/registers/build_discrepancy_register.py [--no-pdf] [--previews]

The Summary sheet of the workbook uses COUNTIFS formulas on the Register, so it follows edits made
in Excel (for example a status changed to Answered). The PDF is made from the same layout with the
values computed at build time.
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
import sys
import tempfile
from collections import Counter, OrderedDict
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402
from bdoi_xlsx import BORDER, HEADER_FILL, HEADER_ROW, BdoiWorkbook, Column, _SheetInfo, _font  # noqa: E402

DATA = HERE / "discrepancy_register.yaml"
REQ = REPO / "docs" / "requirements"
OUT = REPO / "docs" / "deliverables" / "out" / "Registers"

TYPES = ["Contradiction within BRD", "Conflict between BRDs", "Gap / missing detail", "Ambiguity",
         "BRD vs platform", "NFR inconsistency", "Data / numbering issue", "Document quality"]
SEVERITIES = ["High", "Medium", "Low"]
STATUSES = ["Open", "Answered", "Closed"]
Q_STATUSES = ["Open", "Partially answered", "Answered"]
SEV_COLOURS = {"High": (brand.DANGER_BG, brand.DANGER), "Medium": (brand.AMBER_BG, brand.AMBER),
               "Low": (brand.DIRTY_WHITE, brand.MUTED)}

# Open-question section of each spec. The prefix qualifies IDs that two specs share (SQ).
QUESTION_SOURCES = [
    ("BRD-00", "BDOI_CORE_BRD_SPEC.md", "## 13. New open questions (Core Replacement)", ""),
    ("BRD-01", "BDOI_NB_BRD_SPEC.md", "## 9. Open questions", ""),
    ("BRD-02", "BDOI_OPS_BRD_SPEC.md", "## 10. Open questions", ""),
    ("BRD-03", "BDOI_PM_BRD_SPEC.md", "## 7. Open questions", ""),
    ("BRD-04", "BDOI_CLXN_BRD_SPEC.md", "## 10. New open questions", ""),
    ("BRD-05", "BDOI_ACCT_BRD_SPEC.md", "## 9. Open questions", ""),
    ("BRD-06", "BDOI_RN_BRD_SPEC.md", "## 10. New open questions", ""),
    ("BRD-07", "BDOI_CLM_BRD_SPEC.md", "## 11. Open questions", ""),
    ("BRD-08", "BDOI_EB_BRD_SPEC.md", "## 10. Open questions", ""),
    ("BRD-09", "BDOI_CSF_BRD_SPEC.md", "## 10. Open questions", ""),
    ("BRD-10", "BDOI_SANC_BRD_SPEC.md", "## 10. Open questions", "SANC-"),
    ("BRD-11", "BDOI_UAM_BRD_SPEC.md", "## 9. Open questions", ""),
    ("BRD-12", "BDOI_SP_BRD_SPEC.md", "## 10. New open questions", "SP-"),
    ("BRD-13", "BDOI_DM_BRD_SPEC.md", "## 7. Open questions for BDOI (Data Migration)", ""),
]
CROSS_SOURCE = ("CROSS", "BDOI_CROSS_BRD_DECISIONS.md", "## 5. Remaining cross-BRD conflicts")


# --------------------------------------------------------------------------- data

def clean(text: str) -> str:
    """Plain text from a Markdown table cell."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return re.sub(r"\s+", " ", text).strip()


def table_rows(path: Path, heading: str) -> list[list[str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith(heading))
    rows = []
    for line in lines[start + 1:]:
        if line.startswith("## ") or line.startswith("### "):
            break
        if line.startswith("|") and not re.match(r"\|\s*-", line):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            rows.append(cells)
    return rows[1:]  # header row dropped


def load_questions(data: dict) -> list[dict]:
    questions = []
    for brd, fname, heading, prefix in QUESTION_SOURCES:
        for cells in table_rows(REQ / fname, heading):
            qid = clean(cells[0])
            if not re.fullmatch(r"[A-Z]+\d+", qid):
                continue
            questions.append({"id": prefix + qid, "brd": data["brds"][brd]["short"], "topic": clean(cells[1]),
                              "question": clean(cells[2]), "source": clean(cells[3]) if len(cells) > 3 else ""})
    _, fname, heading = CROSS_SOURCE
    for cells in table_rows(REQ / fname, heading):
        qid = clean(cells[0])
        if not qid.startswith("XQ"):
            continue
        topic = re.match(r"\*\*(.+?)\*\*", cells[1])
        conflict = clean(cells[1][topic.end():] if topic else cells[1])
        questions.append({"id": qid, "brd": "Cross-BRD", "topic": clean(topic.group(1) if topic else cells[1]).rstrip("."),
                          "question": f"{clean(cells[3])} Background: {conflict}",
                          "source": "BDOI_CROSS_BRD_DECISIONS.md section 5"})
    # Status, answers, grouping, register links
    status = data.get("question_status", {})
    primary_of: dict[str, str] = {}
    for primary, members in data.get("question_groups", {}).items():
        for m in members:
            primary_of[m] = primary
    item_links: dict[str, list[str]] = {}
    for item in data["items"]:
        for q in item.get("related", []):
            item_links.setdefault(q, []).append(item["ref"])
    ids = {q["id"] for q in questions}
    for q in questions:
        st = status.get(q["id"])
        q["status"] = st[0] if st else "Open"
        q["answer"] = f"{st[2]} ({st[1]})" if st else ""
        q["under"] = primary_of.get(q["id"], "")
        members = data.get("question_groups", {}).get(q["id"], [])
        q["same"] = ", ".join(members)
        q["items"] = ", ".join(item_links.get(q["id"], []))
    unknown = sorted(set(primary_of) | set(status) | set(item_links) - ids)
    unknown = [u for u in unknown if u not in ids]
    if unknown:
        print("warning: question IDs not found in the specs:", ", ".join(unknown), file=sys.stderr)
    return questions


def load() -> dict:
    data = yaml.safe_load(DATA.read_text(encoding="utf-8"))
    refs = [i["ref"] for i in data["items"]]
    expected = [f"DCR-{n:03d}" for n in range(1, len(refs) + 1)]
    if refs != expected:
        raise SystemExit("Register refs must run DCR-001 upwards without gaps")
    for item in data["items"]:
        if item["type"] not in TYPES or item["sev"] not in SEVERITIES or item["status"] not in STATUSES:
            raise SystemExit(f"{item['ref']}: invalid type, severity or status")
        for b in item["brds"]:
            if b not in data["brds"]:
                raise SystemExit(f"{item['ref']}: unknown BRD {b}")
    return data


def owner_of(item: dict, data: dict) -> str:
    if item.get("owner"):
        return item["owner"]
    brds = item["brds"]
    return data["owners"][brds[0]] if len(brds) == 1 else data["owners"]["CROSS"]


def brd_label(item: dict, data: dict) -> str:
    return "; ".join(data["brds"][b]["short"] for b in item["brds"])


# --------------------------------------------------------------------------- workbook

class RegisterWorkbook(BdoiWorkbook):
    """BDOI workbook with a README that also carries the rules of the register."""

    readme_sections: list[tuple[str, list[tuple[str, str]]]] = []

    def custom_sheet(self, name: str, description: str):
        ws = self.wb.create_sheet(name[:31])
        self._sheets.append(_SheetInfo(ws.title, description, [], ""))
        return ws

    def _build_readme(self) -> None:
        super()._build_readme()
        ws = self._readme
        row = ws.max_row + 2
        for title, pairs in self.readme_sections:
            ws.cell(row=row, column=1, value=title).font = _font(12, True, brand.HEADER_BLUE)
            row += 1
            for key, text in pairs:
                a = ws.cell(row=row, column=1, value=key)
                b = ws.cell(row=row, column=2, value=text)
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
                a.font = _font(10, True)
                b.font = _font(10)
                for c in (a, b):
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    c.border = BORDER
                ws.row_dimensions[row].height = max(15, 13.5 * (1 + len(text) // 120))
                row += 1
            row += 1


def fit_rows(ws, columns: list[Column], n: int, font_size: float = 10) -> None:
    """Sets data row heights from the wrapped text (Excel does not grow rows on opening)."""
    per_line = 12.6 if font_size >= 10 else 11.6
    for r in range(HEADER_ROW + 1, HEADER_ROW + n + 1):
        lines = 1
        for c, col in enumerate(columns, start=1):
            value = ws.cell(row=r, column=c).value
            if value is None or not isinstance(value, str):
                continue
            chars = max(4.0, col.width * (1.15 if font_size >= 10 else 1.3) - 1)
            cell_lines = sum(max(1, -(-len(part) // int(chars))) for part in value.split("\n"))
            lines = max(lines, cell_lines)
        ws.row_dimensions[r].height = round(lines * per_line + 4, 1)


def prioritise(ws, n_new: int) -> None:
    """Gives the last n_new conditional-format rules the highest priority (above the banding)."""
    rules = [r for cf in ws.conditional_formatting for r in cf.rules]
    new, old = rules[-n_new:], rules[:-n_new]
    for i, r in enumerate(new + old, start=1):
        r.priority = i


def colour_rule(ws, rng: str, value: str, fill: str, font: str) -> None:
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=[f'"{value}"'], stopIfTrue=True,
        fill=PatternFill("solid", fgColor=fill, bgColor=fill), font=Font(color=font, bold=True)))


REGISTER_COLUMNS = [
    Column("ref", "Ref", 10, "Register reference DCR-nnn. Stable: a new item gets the next number."),
    Column("brds", "BRD(s) involved", 15, "BRDs the item involves (BRD-n and short name; Report List where relevant). Several BRDs are separated by semicolons."),
    Column("loc", "BRD ID / section / page", 26, "Requirement ID, section and PDF page of the source file in docs/source-documents, for each BRD involved."),
    Column("type", "Type", 16, "Kind of issue (definitions below).", values=TYPES),
    Column("desc", "Description", 50, "What the BRD says, with the conflicting statements quoted briefly."),
    Column("impact", "Impact (business / build / test)", 38, "Effect on the business process, on what is built, and on testing."),
    Column("sev", "Severity", 10, "High, Medium or Low (rule below).", values=SEVERITIES),
    Column("modules", "Affected modules", 20, "BIBS modules affected."),
    Column("resolution", "Proposed resolution / recommendation", 36, "What the project proposes, and builds as the default until BDOI answers."),
    Column("question", "Clarification question to BDOI", 38, "The question BDOI is asked to answer."),
    Column("owner", "Owner (BDOI)", 22, "BDOI business unit that owns the answer, from the BRD approval sheets. Items involving several BRDs go to the BIBS Product Owner with the owners involved."),
    Column("raised", "Raised on", 12, "Date the item was raised.", kind="date"),
    Column("status", "Status", 12, "Open, Answered or Closed (rule below).", values=STATUSES, status=True),
    Column("response", "BDOI response", 32, "BDOI's answer, or the signed BDOI document that answers the item."),
    Column("resolved", "Resolution date", 12, "Date the item was Closed.", kind="date"),
    Column("related", "Related question IDs", 18, "Open-question IDs of the specs (CRQ, Q, OQ, PQ, CQ, AQ, RQ, CLQ, EBQ, CSQ, SP-SQ, SANC-SQ, UQ, DMQ, XQ); see the Open questions sheet."),
]

QUESTION_COLUMNS = [
    Column("id", "Question ID", 11, "Original ID of the question. SQ is used by two specs, so it is written SP-SQnn (Submitted Policies) and SANC-SQnn (Sanction Screening)."),
    Column("brd", "BRD", 12, "Spec that raised the question (Cross-BRD: BDOI_CROSS_BRD_DECISIONS.md)."),
    Column("topic", "Topic", 20, "Topic of the question."),
    Column("question", "Question", 62, "Question as asked in the spec."),
    Column("source", "Source reference", 22, "Requirement IDs and earlier questions the question refers to."),
    Column("status", "Status", 14, "Open, Partially answered (a later BRD answers part of it) or Answered.", values=Q_STATUSES, status=True),
    Column("answer", "Answer so far (source)", 42, "Short answer from a later BRD, with the answering document and ID."),
    Column("under", "Consolidated under", 13, "Primary question of the topic when several specs ask the same thing. Answer the primary question; the others follow."),
    Column("same", "Same topic as", 20, "For a primary question: the questions consolidated under it."),
    Column("items", "Register items", 16, "Register items (DCR) that cite the question."),
]


def build_register_rows(data: dict) -> list[dict]:
    raised = data["meta"]["raised_on"]
    rows = []
    for item in data["items"]:
        rows.append({
            "ref": item["ref"], "brds": brd_label(item, data), "loc": item["loc"], "type": item["type"],
            "desc": item["desc"], "impact": item["impact"], "sev": item["sev"], "modules": item["modules"],
            "resolution": item["resolution"], "question": item["question"], "owner": owner_of(item, data),
            "raised": item.get("raised", raised), "status": item["status"], "response": item.get("response", ""),
            "resolved": item.get("resolved"), "related": ", ".join(item.get("related", [])),
        })
    return rows


def style_register(ws, n: int) -> None:
    letters = {c.key: get_column_letter(i) for i, c in enumerate(REGISTER_COLUMNS, start=1)}
    last = HEADER_ROW + n + 500
    rng = f"{letters['sev']}{HEADER_ROW + 1}:{letters['sev']}{last}"
    for sev, (fill, font) in SEV_COLOURS.items():
        colour_rule(ws, rng, sev, fill, font)
    prioritise(ws, len(SEV_COLOURS))
    for r in range(HEADER_ROW + 1, HEADER_ROW + n + 1):
        ws.cell(row=r, column=list(letters).index("sev") + 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=1).font = _font(10, True, brand.HEADER_BLUE)


def style_questions(ws, n: int) -> None:
    col = get_column_letter([c.key for c in QUESTION_COLUMNS].index("status") + 1)
    rng = f"{col}{HEADER_ROW + 1}:{col}{HEADER_ROW + n + 500}"
    colour_rule(ws, rng, "Partially answered", brand.BG_BLUE, brand.HEADER_BLUE)
    prioritise(ws, 1)


def build_nfr_sheet(wb: RegisterWorkbook, data: dict):
    brd_keys = [k for k in data["brds"] if k.startswith("BRD-")]
    cols = [Column("theme", "NFR theme", 18, "Non-functional theme compared across the BRDs.")]
    for k in brd_keys:
        cols.append(Column(k, data["brds"][k]["short"], 17, f"Value in {data['brds'][k]['short']} {data['brds'][k]['name']}, with the page."))
    cols += [
        Column("conflict", "Aligned?", 11, "CONFLICT when the BRDs give different values; ALIGNED when they agree or are silent."),
        Column("proposed", "Proposed BIBS-wide value", 24, "Value the project proposes for all modules, for BDOI decision."),
        Column("refs", "Register items", 13, "Register items on the theme."),
    ]
    rows = []
    for t in data["nfr"]:
        row = {"theme": t["theme"], "conflict": "CONFLICT" if t.get("conflict") else "ALIGNED",
               "proposed": t["proposed"], "refs": ", ".join(t.get("refs", []))}
        for k in brd_keys:
            row[k] = t.get(k, "Not stated")
        rows.append(row)
    ws = wb.sheet("NFR comparison", cols, rows,
                  description="Every NFR theme against each BRD's value. Conflicting rows are highlighted; grey cells: not stated or 'follow QPS'.")
    # Direct fills instead of banding: highlight conflicts and silent values.
    ws.conditional_formatting = ConditionalFormattingList()
    amber = PatternFill("solid", fgColor=brand.AMBER_BG)
    grey = PatternFill("solid", fgColor=brand.DIRTY_WHITE)
    red = PatternFill("solid", fgColor=brand.DANGER_BG)
    green = PatternFill("solid", fgColor=brand.SUCCESS_BG)
    for r, t in enumerate(data["nfr"], start=HEADER_ROW + 1):
        conflict = bool(t.get("conflict"))
        ws.cell(row=r, column=1).font = _font(10, True, brand.HEADER_BLUE)
        for c, k in enumerate(brd_keys, start=2):
            cell = ws.cell(row=r, column=c)
            value = str(cell.value or "")
            silent = value.startswith("Not stated") or "QPS" in value or value.startswith("Blank")
            if silent:
                cell.fill = grey
                cell.font = _font(9, False, brand.MUTED)
            else:
                if conflict:
                    cell.fill = amber
                cell.font = _font(9)
        flag = ws.cell(row=r, column=len(brd_keys) + 2)
        flag.fill = red if conflict else green
        flag.font = Font(name=brand.FONT, size=9, bold=True, color=brand.DANGER if conflict else brand.SUCCESS)
        flag.alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2)
    fit_rows(ws, cols, len(rows), font_size=9)
    return ws


# --------------------------------------------------------------------------- summary

def write_summary(ws, data: dict, questions: list[dict], live: bool, register_title: str = "Register",
                  n_items: int = 0) -> None:
    """Writes the Summary layout. live=True: COUNTIFS formulas on the Register; else static values."""
    items = data["items"]
    letters = {c.key: get_column_letter(i) for i, c in enumerate(REGISTER_COLUMNS, start=1)}
    first, last = HEADER_ROW + 1, HEADER_ROW + max(n_items, len(items)) + 500

    def rng(key: str) -> str:
        return f"'{register_title}'!${letters[key]}${first}:${letters[key]}${last}"

    def count(conds: list[tuple[str, str]], static: int):
        if not live:
            return static
        parts = ",".join(f'{rng(k)},"{v}"' for k, v in conds)
        return f"=COUNTIFS({parts})"

    ws.sheet_view.showGridLines = False
    widths = {"A": 4, "B": 34, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 40}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws["B1"] = "Summary"
    ws["B1"].font = _font(14, True, brand.HEADER_BLUE)
    ws["B2"] = (f"{data['meta']['title']}, version {data['meta']['version']}, {data['meta']['date']}. "
                f"{len(items)} register items; {len(questions)} open questions from {len(QUESTION_SOURCES)} BRD specs and the cross-BRD decisions.")
    ws["B2"].font = _font(9, False, brand.MUTED)
    for c in "BCDEFGH":
        ws[f"{c}3"].border = Border(bottom=Side(style="thin", color=brand.YELLOW))
    row = 5

    def header(r: int, labels: list[str]) -> None:
        for i, h in enumerate(labels):
            cell = ws.cell(row=r, column=2 + i, value=h)
            cell.font = _font(10, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center", wrap_text=True)

    def line(r: int, values: list, bold: bool = False, band: bool = False) -> None:
        for i, v in enumerate(values):
            cell = ws.cell(row=r, column=2 + i, value=v)
            cell.font = _font(10, bold or i == 0)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="top", wrap_text=True)
            if band:
                cell.fill = PatternFill("solid", fgColor=brand.BG_BLUE)

    def title(r: int, text: str, note: str = "") -> int:
        ws.cell(row=r, column=2, value=text).font = _font(12, True, brand.HEADER_BLUE)
        if note:
            ws.cell(row=r + 1, column=2, value=note).font = _font(9, False, brand.MUTED)
            return r + 2
        return r + 1

    def total_formula(r: int, c1: str, c2: str):
        return f"=SUM({c1}{r}:{c2}{r})"

    # 1. Severity x status (status pivot)
    row = title(row, "Items by severity and status", "Status pivot. Counts follow the Status column of the Register.")
    header(row, ["Severity", "Open", "Answered", "Closed", "Total"])
    row += 1
    top = row
    for i, sev in enumerate(SEVERITIES):
        vals = [count([("sev", sev), ("status", s)],
                      sum(1 for x in items if x["sev"] == sev and x["status"] == s)) for s in STATUSES]
        tot = total_formula(row, "C", "E") if live else sum(vals)
        line(row, [sev] + vals + [tot], band=i % 2 == 1)
        ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=SEV_COLOURS[sev][0])
        ws.cell(row=row, column=2).font = _font(10, True, SEV_COLOURS[sev][1])
        row += 1
    tots = [f"=SUM({c}{top}:{c}{row - 1})" if live else sum(1 for x in items if x["status"] == s)
            for c, s in zip("CDE", STATUSES)]
    line(row, ["Total"] + tots + [f"=SUM(F{top}:F{row - 1})" if live else len(items)], bold=True)
    row += 2

    # 2. By BRD
    row = title(row, "Items by BRD", "An item that involves several BRDs is counted once for each BRD involved.")
    header(row, ["BRD", "High", "Medium", "Low", "Total", "Open"])
    row += 1
    for i, (key, b) in enumerate(data["brds"].items()):
        label = f"{b['short']} - {b['name']}" if key != "RL" else b["name"]
        pat = f"*{b['short']}*"
        mine = [x for x in items if key in x["brds"]]
        vals = [count([("brds", pat), ("sev", s)], sum(1 for x in mine if x["sev"] == s)) for s in SEVERITIES]
        tot = total_formula(row, "C", "E") if live else len(mine)
        opn = count([("brds", pat), ("status", "Open")], sum(1 for x in mine if x["status"] == "Open"))
        line(row, [label] + vals + [tot, opn], band=i % 2 == 1)
        row += 1
    multi = [x for x in items if len(x["brds"]) > 1]
    vals = [count([("brds", "*;*"), ("sev", s)], sum(1 for x in multi if x["sev"] == s)) for s in SEVERITIES]
    line(row, ["Items involving two or more BRDs"] + vals +
         [total_formula(row, "C", "E") if live else len(multi),
          count([("brds", "*;*"), ("status", "Open")], sum(1 for x in multi if x["status"] == "Open"))], bold=True)
    row += 2

    # 3. By type
    row = title(row, "Items by type")
    header(row, ["Type", "High", "Medium", "Low", "Total", "Open"])
    row += 1
    top = row
    for i, t in enumerate(TYPES):
        mine = [x for x in items if x["type"] == t]
        vals = [count([("type", t), ("sev", s)], sum(1 for x in mine if x["sev"] == s)) for s in SEVERITIES]
        line(row, [t] + vals + [total_formula(row, "C", "E") if live else len(mine),
                                count([("type", t), ("status", "Open")], sum(1 for x in mine if x["status"] == "Open"))],
             band=i % 2 == 1)
        row += 1
    tots = [f"=SUM({c}{top}:{c}{row - 1})" if live else v for c, v in zip(
        "CDEFG", [sum(1 for x in items if x["sev"] == s) for s in SEVERITIES] + [len(items),
                                                                               sum(1 for x in items if x["status"] == "Open")])]
    line(row, ["Total"] + tots, bold=True)
    row += 2

    # 4. Open questions
    row = title(row, "Open questions", "From the Open questions sheet (status as recorded in the data file).")
    header(row, ["Status", "Questions", "", "", "", ""])
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
    row += 1
    qc = Counter(q["status"] for q in questions)
    for i, s in enumerate(Q_STATUSES):
        line(row, [s, qc.get(s, 0)], band=i % 2 == 1)
        row += 1
    grouped = sum(1 for q in questions if q["under"])
    line(row, ["Consolidated under another question", grouped])
    row += 1
    line(row, ["Total", len(questions)], bold=True)
    row += 2

    # 5. Resolve first
    row = title(row, "Resolve first", "The 15 items that most affect the build, money or compliance, in the order proposed for resolution.")
    header(row, ["Item", "Ref", "Severity", "Status", "BRD(s)", "", "Why first"])
    ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
    row += 1
    by_ref = {x["ref"]: x for x in items}
    for i, (ref, short, why) in enumerate(data["priority"], start=1):
        x = by_ref[ref]
        brds = ", ".join(data["brds"][b]["short"] for b in x["brds"]) if len(x["brds"]) <= 3 else "Cross-BRD"
        line(row, [f"{i}. {short}", ref, x["sev"], x["status"], brds, "", why], band=i % 2 == 0)
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 27
        row += 1


def summary_pdf(wb_template: RegisterWorkbook, data: dict, questions: list[dict], out_pdf: Path) -> Path:
    import render  # noqa: E402 (tools/deliverables)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_summary(ws, data, questions, live=False)
    wb_template._print_setup(ws, None)
    ws.page_setup.orientation = "portrait"
    ws.page_margins.top = 0.8
    tmp = Path(tempfile.mkdtemp())
    src = tmp / (out_pdf.stem + ".xlsx")
    wb.save(src)
    pdf = render.to_pdf(src, tmp)
    shutil.copyfile(pdf, out_pdf)
    shutil.rmtree(tmp, ignore_errors=True)
    return out_pdf


# --------------------------------------------------------------------------- main

def build(make_pdf: bool = True, previews: bool = False) -> Path:
    data = load()
    meta = data["meta"]
    questions = load_questions(data)
    wb = RegisterWorkbook(meta["title"], doc_type=meta["doc_type"], brd=meta["brd"], version=meta["version"],
                          date=meta["date"], subtitle=meta["subtitle"])
    wb.legend = [("Open", "Waiting for a BDOI answer"),
                 ("Answered", "A signed BDOI document answers the item; BDOI confirms to close"),
                 ("Closed", "BDOI confirmed the resolution")]
    wb.cover_notes = [
        "Severity rule, types and status values: see the README sheet.",
        "Page numbers are PDF pages of the source files in docs/source-documents.",
        "The CSF pack holds meeting credentials in an attached e-mail (DCR-098); they are not copied here.",
    ]
    wb.readme_sections = [
        ("Severity rule", [(k, v) for k, v in data["severity_rule"]]),
        ("Types", [(k, v) for k, v in data["type_rule"]]),
        ("Status", [(k, v) for k, v in data["status_rule"]]),
        ("Sources", [(b["short"], f"{b['name']}: docs/source-documents/{b['file']}; baseline docs/requirements/{b['spec']}")
                     for b in data["brds"].values()] + [
            ("Cross-BRD", "docs/requirements/BDOI_CROSS_BRD_DECISIONS.md (decisions D1-D8, questions XQ01-XQ13)"),
            ("Method", "Each item was checked against the source page (text layer, or the rendered page for scanned pages) and cites it. "
                       "Items come from the observations, NFR and open-question sections of the specs, the cross-BRD decisions, "
                       "the Report List, and a comparison with BIBS as built or designed."),
            ("Refresh", "Edit docs/deliverables/src/registers/discrepancy_register.yaml and run "
                        "python docs/deliverables/src/registers/build_discrepancy_register.py. "
                        "Open questions are re-read from the specs at every build.")]),
    ]
    summary_ws = wb.custom_sheet("Summary", "Counts by BRD, type and severity; status pivot; items to resolve first")
    rows = build_register_rows(data)
    reg = wb.sheet("Register", REGISTER_COLUMNS, rows,
                   description="One row per discrepancy, conflict, gap or clarification. Filter on BRD, type, severity or status.")
    style_register(reg, len(rows))
    fit_rows(reg, REGISTER_COLUMNS, len(rows))
    reg.page_setup.paperSize = reg.PAPERSIZE_A3
    qws = wb.sheet("Open questions", QUESTION_COLUMNS, questions,
                   description="Every open question of the specs, with its status, the consolidated question and the register items.")
    style_questions(qws, len(questions))
    fit_rows(qws, QUESTION_COLUMNS, len(questions))
    qws.page_setup.paperSize = qws.PAPERSIZE_A3
    nfr = build_nfr_sheet(wb, data)
    nfr.page_setup.paperSize = nfr.PAPERSIZE_A3
    write_summary(summary_ws, data, questions, live=True, n_items=len(rows))
    wb._print_setup(summary_ws, None)
    summary_ws.page_setup.orientation = "portrait"
    summary_ws.page_margins.top = 0.8

    name = brand.output_name(meta["doc_type"], meta["brd"], meta["name"], meta["version"], "xlsx")
    out = wb.save(OUT / name)
    print(out.relative_to(REPO))
    if make_pdf:
        pdf_name = brand.output_name(meta["doc_type"], meta["brd"], meta["name"] + " Summary", meta["version"], "pdf")
        pdf = summary_pdf(wb, data, questions, OUT / pdf_name)
        print(pdf.relative_to(REPO))
    if previews:
        import render  # noqa: E402

        full = render.to_pdf(out, Path(tempfile.mkdtemp()))
        for p in render.previews(full, dpi=60):
            print(p)
    counts = OrderedDict((("items", len(rows)), ("questions", len(questions))))
    print(", ".join(f"{k}: {v}" for k, v in counts.items()))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--no-pdf", action="store_true", help="skip the Summary PDF")
    ap.add_argument("--previews", action="store_true", help="render the whole workbook to PNG previews (not committed)")
    args = ap.parse_args()
    build(make_pdf=not args.no_pdf, previews=args.previews)
    return 0


if __name__ == "__main__":
    sys.exit(main())
