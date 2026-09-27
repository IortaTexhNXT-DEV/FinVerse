"""Builds the Sign-off Pack Guide deck of a BRD business sign-off release set (PowerPoint, BDO theme).

The same deck is produced for every BRD set: the content comes from the pack of the BRD
(docs/deliverables/src/signoff/<brd>/pack.yaml and guide.yaml) and from the system as built through the pack
(personas, menus, cross-BRD contract, screenshots). Every slide carries speaker notes.

    python docs/deliverables/src/signoff/build_guide_deck.py brd01
"""

from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Sequence

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
sys.path.insert(0, str(ROOT / "tools" / "deliverables"))
sys.path.insert(0, str(HERE))

import brand  # noqa: E402
from bdoi_docx import render_dot  # noqa: E402
from bdoi_pptx import MARGIN, H, W, BdoiDeck, _rgb  # noqa: E402
import signoff_pack  # noqa: E402

CONTENT_W = W - 2 * MARGIN
TOP = Inches(1.45)


class GuideDeck(BdoiDeck):
    """The BDOI deck with speaker notes, a step flow, image and gallery slides."""

    @property
    def slide(self):
        return self.prs.slides[-1]

    def notes(self, text: str | Sequence[str]) -> None:
        body = text if isinstance(text, str) else "\n".join(text)
        self.slide.notes_slide.notes_text_frame.text = " ".join(body.split()) if "\n" not in body else body.strip()

    def page(self, title: str):
        s = self._slide()
        self._chrome(s, title)
        return s

    def image(self, title: str, png: Path, caption: str = "", side: Sequence[str] | None = None) -> None:
        s = self.page(title)
        area_w = CONTENT_W if not side else Inches(8.2)
        area_h = Inches(4.9)
        self._picture(s, png, MARGIN, TOP, area_w, area_h)
        if caption:
            self._text(s, MARGIN, TOP + area_h + Inches(0.05), area_w, Inches(0.35), caption, size=11, italic=True,
                       colour=brand.MUTED, align=PP_ALIGN.CENTER)
        if side:
            x = MARGIN + area_w + Inches(0.3)
            w = W - MARGIN - x
            self._rect(s, x, TOP, w, area_h + Inches(0.4), brand.BG_BLUE)
            self._rect(s, x, TOP, Pt(4), area_h + Inches(0.4), brand.HEADER_BLUE)
            self._bullets(s, x + Inches(0.15), TOP + Inches(0.1), w - Inches(0.25), area_h, side, size=13)

    def _picture(self, s, png: Path, x, y, w, h) -> None:
        from PIL import Image

        with Image.open(png) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = int(iw * scale), int(ih * scale)
        pic = s.shapes.add_picture(str(png), x + (w - pw) // 2, y, width=pw, height=ph)
        pic.line.color.rgb = _rgb(brand.BORDER)
        pic.line.width = Pt(0.75)

    def gallery(self, title: str, items: Sequence[tuple[Path, str]]) -> None:
        s = self.page(title)
        cols = 2
        cell_w = (CONTENT_W - Inches(0.3)) / cols
        cell_h = Inches(2.35)
        for i, (png, caption) in enumerate(items):
            x = MARGIN + (i % cols) * (cell_w + Inches(0.3))
            y = TOP + (i // cols) * (cell_h + Inches(0.1))
            self._picture(s, png, x, y, cell_w, cell_h - Inches(0.3))
            self._text(s, x, y + cell_h - Inches(0.3), cell_w, Inches(0.3), caption, size=11, italic=True,
                       colour=brand.MUTED, align=PP_ALIGN.CENTER)

    def flow(self, title: str, steps: Sequence[dict[str, Any]]) -> None:
        """Numbered step boxes in two rows of four, joined by arrows."""
        s = self.page(title)
        per_row = 4
        gap = Inches(0.45)
        bw = (CONTENT_W - gap * (per_row - 1)) / per_row
        bh = Inches(2.2)
        for i, st in enumerate(steps):
            row, col = divmod(i, per_row)
            x = MARGIN + col * (bw + gap)
            y = TOP + Inches(0.1) + row * (bh + Inches(0.45))
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh)
            box.fill.solid()
            box.fill.fore_color.rgb = _rgb(brand.BG_BLUE)
            box.line.color.rgb = _rgb(brand.FIELD_BLUE)
            box.shadow.inherit = False
            circle = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.12), y + Inches(0.12), Inches(0.42), Inches(0.42))
            circle.fill.solid()
            circle.fill.fore_color.rgb = _rgb(brand.YELLOW)
            circle.line.fill.background()
            tf = circle.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(st["step"])
            r.font.size, r.font.bold, r.font.name = Pt(14), True, brand.FONT
            r.font.color.rgb = _rgb(brand.NEAR_BLACK)
            self._text(s, x + Inches(0.62), y + Inches(0.08), bw - Inches(0.7), Inches(0.55), st["name"], size=13,
                       bold=True, colour=brand.HEADER_BLUE, anchor=MSO_ANCHOR.MIDDLE)
            self._text(s, x + Inches(0.12), y + Inches(0.7), bw - Inches(0.24), Inches(0.9), st["outputs"], size=10,
                       colour=brand.TEXT)
            self._text(s, x + Inches(0.12), y + bh - Inches(0.5), bw - Inches(0.24), Inches(0.42), st["duration"],
                       size=10, bold=True, colour=brand.NEAR_BLACK)
            if col < per_row - 1 and i < len(steps) - 1:
                arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.07), y + bh / 2 - Inches(0.14),
                                           gap - Inches(0.14), Inches(0.28))
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = _rgb(brand.HEADER_BLUE)
                arrow.line.fill.background()

    def columns(self, title: str, blocks: Sequence[tuple[str, Sequence[str]]], size: int = 13) -> None:
        s = self.page(title)
        n = len(blocks)
        gap = Inches(0.3)
        cw = (CONTENT_W - gap * (n - 1)) / n
        for i, (head, items) in enumerate(blocks):
            x = MARGIN + i * (cw + gap)
            self._rect(s, x, TOP, cw, Inches(0.5), brand.HEADER_BLUE)
            self._text(s, x + Inches(0.1), TOP, cw, Inches(0.5), head, size=15, bold=True, colour=brand.WHITE,
                       anchor=MSO_ANCHOR.MIDDLE)
            self._rect(s, x, TOP + Inches(0.5), cw, H - TOP - Inches(1.3), brand.BG_BLUE)
            self._bullets(s, x + Inches(0.1), TOP + Inches(0.65), cw - Inches(0.2), H - TOP - Inches(1.5), items,
                          size=size)


# ------------------------------------------------------------------ figures


def workbook_image(xlsx: Path, sheet: str, out: Path, rows: int = 12) -> Path:
    """A picture of the first rows of one sheet of the sign-off workbook (LibreOffice, then pdftoppm)."""
    import openpyxl
    from PIL import Image, ImageOps

    with tempfile.TemporaryDirectory() as tmp:
        wb = openpyxl.load_workbook(xlsx)
        for ws in list(wb.worksheets):
            if ws.title != sheet:
                wb.remove(ws)
        ws = wb[sheet]
        # Only the columns the reviewer reads and fills in, so the picture stays legible.
        keep = {"Screen ID", "No.", "Field", "Type", "Mandatory", "Message when it fails", "BU review", "BU comment",
                "Reviewer", "Review date"}
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(4, c).value or "") not in keep:
                ws.column_dimensions[ws.cell(4, c).column_letter].hidden = True
        # Two rows filled in as an illustration of the review (only in this picture).
        heads = {str(ws.cell(4, c).value or ""): c for c in range(1, ws.max_column + 1)}
        for r, (decision, comment) in ((5, ("Accept", "")), (7, ("Change requested", "Show the code of a confirmed "
                                                                 "client and of a prospect in one column"))):
            if "BU review" in heads:
                ws.cell(r, heads["BU review"], decision)
                ws.cell(r, heads["BU comment"], comment)
                ws.cell(r, heads["Reviewer"], "Marketing reviewer")
                ws.cell(r, heads["Review date"], "12-Oct-2026")
        ws.print_area = f"A1:{ws.cell(4, ws.max_column).column_letter}{4 + rows}"
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        src = Path(tmp) / "sheet.xlsx"
        wb.save(src)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        subprocess.run(["pdftoppm", "-png", "-r", "150", "-f", "1", "-l", "1", "-singlefile",
                        str(Path(tmp) / "sheet.pdf"), str(Path(tmp) / "page")], check=True)
        im = Image.open(Path(tmp) / "page.png").convert("RGB")
        # Keep the table: from the top to the first long white gap (the page footer is left out).
        gray = ImageOps.invert(im.convert("L"))
        used = [y for y in range(im.height) if gray.crop((0, y, im.width, y + 1)).getbbox()]
        end = used[0]
        for a, b in zip(used, used[1:]):
            if b - a > 80:
                break
            end = b
        box = ImageOps.invert(im.crop((0, 0, im.width, end + 1))).getbbox()
        if box:
            im = im.crop((max(0, box[0] - 20), max(0, box[1] - 20), min(im.width, box[2] + 20), end + 20))
        out.parent.mkdir(parents=True, exist_ok=True)
        im.save(out)
    return out


def contract_map(pack: signoff_pack.Pack, out_dir: Path) -> Path:
    """Graphviz map of what the module takes from and hands to the other BRDs and systems."""
    module = pack.meta["name"]
    takes: dict[str, list[str]] = {}
    gives: dict[str, list[str]] = {}
    for c in pack.contract:
        party = c["party"]
        what = " ".join(str(c["what"]).split())
        short = what.split(";")[0].split(",")[0][:48]
        if "In" in c["direction"]:
            takes.setdefault(party, []).append(short)
        if "Out" in c["direction"]:
            gives.setdefault(party, []).append(short)

    def label(party: str, items: list[str]) -> str:
        return party.replace('"', "'") + "\\n" + "\\n".join(f"- {i}" for i in items[:2]).replace('"', "'")

    lines = ["digraph contract {",
             '  graph [rankdir=LR, fontname="Arial", nodesep=0.12, ranksep=0.9, bgcolor="white"];',
             '  node [shape=box, style="rounded,filled", fillcolor="@BG_BLUE", color="@FIELD_BLUE", fontname="Arial",'
             ' fontsize=9, fontcolor="@NEAR_BLACK", margin="0.1,0.05"];',
             '  edge [color="@CTA_BLUE", arrowsize=0.6];',
             f'  core [label="{module}", fillcolor="@HEADER_BLUE", fontcolor="white", fontsize=16, color="@HEADER_BLUE",'
             ' margin="0.3,0.25"];']
    for i, (party, items) in enumerate(sorted(takes.items())):
        lines.append(f'  in{i} [label="{label(party, items)}"];')
        lines.append(f"  in{i} -> core;")
    for i, (party, items) in enumerate(sorted(gives.items())):
        lines.append(f'  out{i} [label="{label(party, items)}"];')
        lines.append(f"  core -> out{i};")
    lines.append("}")
    src = out_dir / f"{pack.meta['brd'].lower().replace('-', '')}_contract_map.dot"
    src.parent.mkdir(parents=True, exist_ok=True)
    src.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return render_dot(src)


# ------------------------------------------------------------------ deck


def build(brd_dir: str) -> Path:
    pack = signoff_pack.Pack(HERE / brd_dir / "pack.yaml")
    g = pack.guide
    m = pack.meta
    gm = g["meta"]
    module = m["name"]
    label = m["brd_label"]
    shots = (pack.dir / m.get("screenshot_dir", "screenshots")).resolve()
    figs = HERE / brd_dir / "figures"
    deck = GuideDeck(f"{gm['deck_title']} {label} {module}", version=str(m["version"]), date=str(m["date"]),
                     subtitle=f"{label} {module} business sign-off pack, release set v{m['version']}")
    counts = (f"{len(pack.screens)} screens, {sum(len(s.fields) for s in pack.screens)} fields, "
              f"{len(pack.rules())} business rules, {len(pack.messages)} messages")

    # 1. Title
    deck.title(f"{gm['deck_title']}: {label} {module}")
    deck.notes(f"This deck guides the business units through the {label} {module} sign-off pack, release set "
               f"v{m['version']}. It explains what the pack is, how to read it, the steps up to sign-off and what "
               "happens after. It is presented at the kick-off and stays in the pack as file 01.")

    # 2. Purpose
    deck.content("Purpose of the pack and what sign-off means", [
        f"The pack shows {module} as built in BIBS, screen by screen, so each business unit confirms what it will get",
        (1, f"Content: {counts}, from the system as built"),
        (1, "Every screen with numbered fields, the navigation per persona, three end-to-end walkthroughs"),
        "Signing freezes the content, the screens and the navigation of the module",
        (1, "After sign-off a change goes through the Change Management Register with its mandays"),
        (1, "Configuration values marked default (SLA hours, thresholds, list entries, templates) stay open"),
        "The screenshots use fictitious seed data only, in the BDO theme",
    ])
    deck.notes(["Stress that the pack describes the system as built on the status date, not a design.",
                "Signing means: this is what we get; the look, feel, fields, rules, messages and navigation are frozen.",
                "Data on the screenshots is seed data of the SIT environment; names and numbers are fictitious."])

    # 3. Pack at a glance
    rows = [[signoff_pack.doc_file(pack, d), d["what"], d["readers"], d["when"]] for d in g["documents"]]
    deck.table("The pack at a glance", ["File", "What it is for", "Who reads it", "When"], rows,
               widths=[3.6, 5.2, 2.6, 1.8], size=10)
    deck.notes("The files are numbered 00 to 05 so they sort in reading order in the BRD folder. 00 Start Here is "
               "the short version of this deck; 02 FRS and 03 workbook carry the same rows; 04 and 05 are the test "
               "plan traced to the screens.")

    # 4. Where to start
    deck.table("Where to start: reading order per role", ["Role", "Read, in this order"],
               [[r["role"], r["order"]] for r in g["reading"]], widths=[2.6, 9.5], size=11)
    deck.notes("Nobody needs to read the whole FRS. Each role starts with its chapters and its sheets of the workbook; "
               "the walkthroughs in chapter 14 are the fastest way to see the whole process.")

    # 5-7. Approach and steps
    deck.flow("Approach: from issue to closure", g["steps"])
    deck.notes([f"{x['step']}. {x['name']}: {x['duration']}." for x in g["steps"]])
    half = (len(g["steps"]) + 1) // 2
    for part, chunk in enumerate((g["steps"][:half], g["steps"][half:]), start=1):
        deck.table(f"Steps in detail ({part} of 2): who, inputs, outputs, duration",
                   ["Step", "Who (R, A, C, I)", "Inputs", "Outputs", "Duration"],
                   [[f"{x['step']}. {x['name']}", x["who"], x["inputs"], x["outputs"], x["duration"]] for x in chunk],
                   widths=[2.2, 3.4, 2.8, 3.2, 1.9], size=10)
        deck.notes("R responsible, A accountable, C consulted, I informed. The durations are indicative; the dates "
                   "are those of the Start Here guide.")

    # 8. Workbook
    wb_png = figs / "signoff-workbook-field-register.png"
    xlsx = brand.out_dir(m["brd"], "Signoff") / brand.output_name("Signoff", m["brd"], module, str(m["version"]), "xlsx")
    if xlsx.exists():
        workbook_image(xlsx, "Field register", wb_png)
    deck.image("How to fill in the sign-off workbook", wb_png, "Field register sheet of the sign-off workbook", side=[
        "Review the rows of your area on Screen standards, Screen catalogue, Field register, Business rules, Messages",
        "BU review: Accept, Change requested or Comment",
        "BU comment: the change asked for",
        "Reviewer and Review date",
        "Questions go to the Comments log",
        "Leave the other columns unchanged; one workbook per unit",
    ])
    deck.notes("The No. column is the numbered marker on the screenshot of the FRS. The drop-down columns accept only "
               "the listed values. The Comments log, Meeting minutes and Version history sheets are kept by the "
               "project team; the Sign-off certificate is signed at the end.")

    # 9-11. Module at a glance
    process = ROOT / "docs" / "deliverables" / "src" / "frs" / "figures" / f"{m['brd'].lower().replace('-', '')}_process_flow.png"
    if process.exists():
        deck.image(f"{module} at a glance: process flow", process, f"{label} process from the client to the booked "
                   "invoice")
        deck.notes("The process runs from client onboarding and KYC, through the quotation or the PRF, the account, "
                   "payment and placement, issuance and booking, to the hand-off to Operations and Accounting.")
    personas = []
    for role, rows_ in pack.menus().items():
        own = [r for r in rows_ if r["own"]]
        sections = []
        for r in own:
            if r["section"] not in sections:
                sections.append(r["section"])
        personas.append([pack.persona_label(role), pack.personas[role].get("user", ""), str(len(own)),
                         ", ".join(sections)])
    deck.table(f"{module} at a glance: personas and menus", ["Persona", "SIT user", "Screens", "Menu sections"],
               personas, widths=[3.0, 1.3, 1.1, 7.2], size=10)
    deck.notes("Each persona sees only the menu entries of its role; the full menu of each persona is in FRS chapter "
               "12 and on the Menu by persona sheet of the workbook.")
    deck.gallery(f"{module} at a glance: key screens",
                 [(shots / f"{k['image']}.png", k["caption"]) for k in g["key_screens"]])
    deck.notes("Screenshots of the SIT environment with seed data; every screen of the module is in FRS chapter 13.")

    # 12. Rules
    deck.content("Key business rules and validations", list(g["rules"]) + [
        (1, f"Full list: FRS chapter 13 per screen, {len(pack.rules())} rules on the Business rules sheet and "
            f"{len(pack.messages)} messages in chapter 15 and on the Messages sheet")], size=14)
    deck.notes("These are the rules the business units ask about most. Each rule is stated with its FR on the screen "
               "where it applies.")

    # 13. Cross-module map
    deck.image("Cross-module dependencies and impacts", contract_map(pack, figs),
               f"What {module} takes from and hands to the other BRDs and systems (FRS chapter 19)")
    deck.notes(["The contract is part of the signed set. " + " ".join(pack.change_rule.split())] +
               [f"{c['id']} {c['party']} ({c['direction']}): {' '.join(str(c['what']).split())}" for c in pack.contract])

    # 14-15. Caveats
    deck.table("Caveats and possibilities", ["Topic", "What can happen", "How it is handled"],
               [[c["topic"], c["text"], c["handling"]] for c in g["caveats"]], widths=[2.2, 5.2, 5.2], size=10)
    deck.notes("Three ways to handle a case: an in-scope clarification (no change to the signed set, answered in the "
               "comments log); a change request with mandays in the Change Management Register; a delta sign-off of "
               "the affected pages only.")
    deck.table("Caveats: worked examples", ["Topic", "Example and how it is handled"],
               [[c["topic"], c["example"]] for c in g["caveats"]], widths=[2.4, 10.2], size=11)
    deck.notes("Use these examples to agree the handling of a case during the review.")

    # 16. Criteria
    deck.columns("Entry and exit criteria, definition of done",
                 [("Entry", g["entry"]), ("Exit", g["exit"]), ("Done", g["done"])], size=12)
    deck.notes("The sign-off meeting starts when the exit criteria are met; the set is done when the certificate is "
               "signed and the build is frozen.")

    # 17. Handover
    deck.content("Handover to development: build readiness checklist",
                 [f"☐  {x}" for x in g["handover"]] + [(1, g["handover_note"])], size=16)
    deck.notes(g["handover_note"])

    # 18. Change control
    deck.content("Change control after sign-off", [
        "A change to anything frozen is raised in the Change Management Register",
        (1, "The screen, field, rule or message concerned, the reason and the priority"),
        "The project team assesses the mandays and the effect on the other BRDs through the cross-BRD contract",
        "The owners of every BRD concerned approve it",
        "Delivered as a new version of the set (v2.1, v2.2...) with the changed pages signed again (delta sign-off)",
        "Configuration values marked default change in BIBS without a change request",
    ])
    deck.notes("Nothing in a signed set changes without a request in the register. The as-built refresh at the end of "
               "the build re-issues the set with the system as delivered.")

    # 19. Governance
    deck.table("Governance, contacts and next steps", ["Role", "Who", "Does"],
               [[x["role"], x["who"], x["does"]] for x in g["governance"]], widths=[3.0, 4.2, 5.4], size=10)
    deck.notes(["Next steps:"] + list(g["next_steps"]))
    deck.content("Next steps and the release order of the other BRD sets", list(g["next_steps"]), size=16)
    deck.notes("The other BRD sets follow the same pack, guide and steps.")

    out = brand.out_dir(m["brd"], "GuideDeck") / brand.output_name("GuideDeck", m["brd"], module, str(m["version"]),
                                                                    "pptx")
    return deck.save(out)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("brd", help="folder of the pack, e.g. brd01")
    args = ap.parse_args(argv)
    print(f"pptx: {build(args.brd)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
