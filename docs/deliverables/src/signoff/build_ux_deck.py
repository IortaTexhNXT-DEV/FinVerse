"""Builds the UX screen documents of a business sign-off set for the BDOI UX Design team (UXD).

The UXD team works on the screens in parallel with the review of the FRS. When the business changes the FRS, the
change reference names the screens it affects, and the UXD team revises those screens only. The three files of a set:

* 07 UX Screen Deck (PowerPoint): cover, purpose and use, how an FRS change reaches the UXD team, contacts; the
  design foundation (colours, type, spacing, icons, the shared components with a real screenshot each, the screen
  standards); the navigation (landing page and full menu of each persona); one section per persona; every
  end-to-end flow as a swimlane overview and one slide per step; the screen catalogue (every screen with its numbered
  callouts, fields, actions and every state captured); the outputs (documents, notifications, exports); the UX
  points for confirmation (presentation items of the clarifications chapter, screen observations, accessibility,
  desktop assumptions); the map of screen IDs, FRS sections and requirements. Screens are shown at full slide width,
  split over several slides when they are tall, never shrunk to thumbnails;
* 08 UX screen register (Excel, guided workbook): one row per screen image with the screen ID and name, the
  personas, the flow and step, the state, the FRS section and requirements, the image file, the UXD status and
  comments, and the "Changed in this issue" flag with the change reference;
* 09 image package (ZIP): every screen image at twice the screen resolution (PNG) named <Screen ID>_<state>.png, the
  crops of the shared components, and the register as CSV.

Sources: the pack of the BRD (pack.yaml, screens, walkthroughs, documents, notifications, messages, ownership.yaml),
its ux.yaml (persona texts, images taken for the deck only, flows of a pack without walkthroughs, the changes of the
issue, the clarifications and observations for the UXD team), ux_foundation.yaml (the shared components), the design
tokens of the screens, the screen standards appendix of the FRS and the FRS as issued (section numbers).

    python docs/deliverables/src/signoff/build_ux_deck.py BRD-03            # the three files
    python docs/deliverables/src/signoff/build_ux_deck.py BRD-03 --check    # checks the images and the files written
"""

from __future__ import annotations

import argparse
import csv
import io
import math
import re
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Sequence

import yaml
from lxml import etree
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
sys.path.insert(0, str(ROOT / "tools" / "deliverables"))
sys.path.insert(0, str(HERE))

import brand  # noqa: E402
import guided_xlsx as g  # noqa: E402
import signoff_pack  # noqa: E402
from bdoi_pptx import MARGIN, H, W, _rgb  # noqa: E402
from build_guide_deck import GuideDeck  # noqa: E402

FOUNDATION = HERE / "ux_foundation.yaml"
TOKENS = ROOT / "frontend" / "src" / "styles" / "tokens.css"

# The UX states of a screen image, in the order of the catalogue.
STATES: dict[str, str] = {
    "default": "Default",
    "filled": "Filled",
    "validation-error": "Validation error",
    "confirmation": "Confirmation with reason",
    "success": "Success message",
    "empty": "Empty list",
    "read-only": "Read-only or locked",
    "pending-approval": "Pending approval",
    "rejected-returned": "Rejected or returned",
    "document": "Document as received",
    "landing": "Landing page",
    "menu": "Menu",
}
# The states every screen is checked for in the state matrix of the catalogue.
SCREEN_STATES = ["default", "filled", "validation-error", "confirmation", "success", "empty", "read-only",
                 "pending-approval", "rejected-returned"]
STATE_TONES = {  # fill, text of the state pill on the slides
    "validation-error": (brand.DANGER_BG, brand.DANGER),
    "rejected-returned": (brand.DANGER_BG, brand.DANGER),
    "success": (brand.SUCCESS_BG, brand.SUCCESS),
    "pending-approval": (brand.AMBER_BG, brand.AMBER),
    "confirmation": (brand.AMBER_BG, brand.AMBER),
    "read-only": (brand.DIRTY_WHITE, brand.MUTED),
}
UXD_STATUSES = ["Not started", "In design", "Ready for review", "Approved"]

# Slide geometry: content area under the title rule and above the footer.
X0 = MARGIN
CW = W - 2 * MARGIN
Y0 = Inches(1.38)
Y1 = H - Inches(0.62)
CAPTION_H = Inches(0.32)
# Largest scale of a screen image: 1.3 times its size on a 96 dpi screen (the images hold two pixels per screen
# pixel), so a small dialog is shown larger but stays sharp.
MAX_SCALE = Inches(1) * 1.3 / 192


# ============================================================================ model


@dataclass
class Img:
    """One screen image of the set (a row of the register, a file of the package)."""

    slug: str
    file: Path
    screen: str
    title: str
    state: str
    name: str
    caption: str
    kind: str  # screen, step, ux, document, nav
    personas: list[str] = field(default_factory=list)
    flows: list[tuple[str, int]] = field(default_factory=list)  # (flow id, step)
    frs: list[str] = field(default_factory=list)
    section: str = "-"
    changed: list[str] = field(default_factory=list)
    in_frs: bool = True

    @property
    def state_label(self) -> str:
        return STATES.get(self.state, self.state)


def _words(text: Any) -> str:
    return " ".join(str(text or "").split())


class Ux:
    """Everything the three files need, read once from the pack."""

    def __init__(self, brd: str):
        self.pack = signoff_pack.Pack(brand.src_dir(brand.brd_of_code(brd)) / "pack" / "pack.yaml")
        p = self.pack
        self.meta = p.meta
        self.ux: dict[str, Any] = p.ux or {}
        self.brd = p.meta["brd"]
        self.code = p.meta["code"]
        self.version = str(p.meta["version"])
        self.module = p.module
        self.label = p.meta["brd_label"]
        self.shots_dir = (p.dir / p.meta.get("screenshot_dir", "screenshots")).resolve()
        self.problems: list[str] = []
        self.user_role = {v.get("user"): r for r, v in p.menu_roles.items() if v.get("user")}
        self.sections = self._frs_sections()
        self.flows = self._flows()
        self.images = self._images()
        self.by_slug = {i.slug: i for i in self.images}
        self.foundation = yaml.safe_load(FOUNDATION.read_text(encoding="utf-8"))
        self.components = [c for c in self.foundation["components"]]

    # ------------------------------------------------------------------ FRS as issued

    def frs_docx(self) -> Path | None:
        folder = brand.out_dir(self.brd, "FRS")
        hits = sorted(folder.glob(f"02_BIBS_*_{self.brd}_*_v{self.version}.docx"))
        return hits[0] if hits else None

    def frs_name(self) -> str:
        return "Handbook" if (self.pack.brd_dir / self.meta["frs"]).name.startswith("HANDBOOK") else "FRS"

    def _frs_sections(self) -> dict[str, str]:
        """Section number of each screen, walkthrough, document and of the menus, from the headings of the issued
        FRS (or handbook)."""
        path = self.frs_docx()
        out: dict[str, str] = {}
        if path is None:
            self.problems.append(f"{self.brd}: the issued FRS v{self.version} is not in the release folder")
            return out
        import docx  # noqa: PLC0415

        for para in docx.Document(str(path)).paragraphs:
            if not para.style.name.startswith("Heading"):
                continue
            text = para.text.replace(" ", " ").strip()
            m = re.match(r"^(\d+(?:\.\d+)*)\s+(SCR-[A-Z]+-\d+|WT-[A-Z]|DO-\d+)\b", text)
            if m:
                out.setdefault(m.group(2), m.group(1))
            m = re.match(r"^(\d+(?:\.\d+)*)\s+(Menu by persona|Personas|Common screen elements|Messages(?: catalogue)?|"
                         r"Notifications(?: catalogue| and alerts)?|Document outputs|Proposed business rules and "
                         r"clarifications for confirmation|Traceability|Appendix: Screen standards)$", text)
            if m:
                out.setdefault(m.group(2), m.group(1))
        return out

    def section_of(self, key: str) -> str:
        return self.sections.get(key, "-")

    # ------------------------------------------------------------------ flows

    def _flows(self) -> list[dict[str, Any]]:
        """The end-to-end flows: the walkthroughs of the pack, then the flows of ux.yaml (a pack without
        walkthroughs shows its flows with the screen images of the FRS). Each step: role, screen, image slug, does,
        system, sees."""
        out = []
        for w in self.pack.walkthroughs:
            steps = []
            for role, scr, action, sees, result, slug in w["steps"]:
                steps.append({"role": role, "screen": scr, "slug": slug, "does": _words(action),
                              "sees": _words(sees), "system": _words(result)})
            out.append({"id": w["id"], "title": w["title"], "summary": _words(w.get("summary")),
                        "data": _words(w.get("data")), "steps": steps, "downstream": w.get("downstream") or [],
                        "captured": True})
        for f in self.ux.get("flows") or []:
            steps = [{"role": s["role"], "screen": s["screen"], "slug": s["image"], "does": _words(s["does"]),
                      "sees": _words(s.get("sees")), "system": _words(s.get("system"))} for s in f["steps"]]
            out.append({"id": f["id"], "title": f["title"], "summary": _words(f.get("summary")),
                        "data": _words(f.get("data")), "steps": steps, "downstream": f.get("downstream") or [],
                        "captured": False})
        return out

    # ------------------------------------------------------------------ images

    def _classify(self, slug: str, state: str, text: str) -> str:
        """The UX state of an image: ux.yaml classify, else from the state name and the caption or step texts."""
        fixed = (self.ux.get("classify") or {}).get(slug)
        if fixed:
            return fixed
        s, t = state.lower(), text.lower()
        if s in STATES:
            return s
        if re.search(r"error|invalid", s) or re.search(r"field messages|under the fields|\brefus|cannot|may not|"
                                                       r"is not allowed|must be|already waiting|still needed|"
                                                       r"you used this password|needs a|\bno decision button", t):
            return "validation-error"
        if re.search(r"return|reject|deactivate|waive|explain", s):
            return "confirmation"
        if re.search(r"\breturned\b|\brejected\b", t):
            return "rejected-returned"
        if re.search(r"^toast|\btoast\b|confirmation on|was created|implemented|released\b|approved\b|\bsaved\b", t):
            return "success"
        if re.search(r"pending|waiting|for approval|second approval|for validation|to authorise|to authorize", t):
            return "pending-approval"
        if re.search(r"filled|entered", s + " " + t) and re.search(r"new|form|dialog|request", t):
            return "filled"
        if re.search(r"history|log|released|read-only|locked", s):
            return "read-only"
        return "default"

    def _images(self) -> list[Img]:
        p = self.pack
        out: list[Img] = []
        names: set[str] = set()

        def unique(name: str) -> str:
            base, n = name, 2
            while name in names:
                name = base.replace(".png", f"-{n}.png")
                n += 1
            names.add(name)
            return name

        changes = self.ux.get("issue", {}).get("changes") or []

        def changed(slug: str, screen: str) -> list[str]:
            refs = []
            for c in changes:
                if slug in (c.get("images") or []) or (screen in (c.get("screens") or []) and not c.get("images")):
                    refs.append(c["ref"])
            return refs

        for s in p.screens:
            for i, shot in enumerate(s.get("shots") or [], start=1):
                slug = f"{s.id.lower()}-{i:02d}-{shot.get('state', 'view')}"
                role = self.user_role.get(shot.get("user"))
                state = self._classify(slug, str(shot.get("state", "view")), str(shot.get("caption", "")))
                out.append(Img(slug, p.shot_file(slug), s.id, s.title, state,
                               unique(f"{s.id}_{str(shot.get('state', 'view')).replace('_', '-')}.png"),
                               _words(shot.get("caption", s.title)), "screen",
                               [p.persona_label(role)] if role else self._screen_personas(s),
                               frs=list(s.get("frs") or []), section=self.section_of(s.id)))
        for w in self.flows:
            for n, st in enumerate(w["steps"], start=1):
                if not w["captured"]:
                    continue
                scr = p.step_screen(st["screen"])
                text = f"{st['sees']} {st['system']}"
                state = self._classify(st["slug"], "", text)
                out.append(Img(st["slug"], p.shot_file(st["slug"]), st["screen"], scr.title if scr else st["screen"],
                               state, unique(f"{st['screen']}_{w['id'].lower()}-step-{n:02d}.png"),
                               f"{w['id']} step {n}: {st['does']}", "step", [p.persona_label(st["role"])],
                               flows=[(w["id"], n)], frs=list(scr.get("frs") or []) if scr else [],
                               section=self.section_of(w["id"])))
        for d in p.documents:
            out.append(Img(d["shot"], p.shot_file(d["shot"]), d["id"], d["name"], "document",
                           unique(f"{d['id']}_document.png"), f"{d['name']}: first page as generated", "document",
                           [], frs=list(d.get("frs") or []), section=self.section_of(d["id"])))
        for shot in p.ux_shots():
            if shot.get("component"):
                continue
            if shot["state"] in ("landing", "menu"):
                n = int(shot["slug"].split("-")[2])
                out.append(Img(shot["slug"], p.shot_file(shot["slug"]), shot["screen"],
                               f"{p.persona_label(shot['role'])}: " + ("landing page" if shot["state"] == "landing"
                                                                       else "menu"),
                               shot["state"], unique(f"{p.nav_id(n)}_{shot['state']}.png"), shot["caption"], "nav",
                               [p.persona_label(shot["role"])], section=self.section_of("Menu by persona"),
                               in_frs=False))
                continue
            s = p.by_id[shot["screen"]]
            role = self.user_role.get(shot.get("user"))
            state = shot.get("ux") if isinstance(shot.get("ux"), str) else self._classify(
                shot["slug"], str(shot["state"]), str(shot.get("caption", "")))
            out.append(Img(shot["slug"], p.shot_file(shot["slug"]), s.id, s.title, state,
                           unique(f"{s.id}_{str(shot['state']).replace('_', '-')}.png"), _words(shot["caption"]), "ux",
                           [p.persona_label(role)] if role else self._screen_personas(s),
                           frs=list(s.get("frs") or []), section=self.section_of(s.id), in_frs=False))
        # The flows of ux.yaml show images of the screens: the image records the flow and the step.
        by_slug = {i.slug: i for i in out}
        for w in (x for x in self.flows if not x["captured"]):
            for n, st in enumerate(w["steps"], start=1):
                img = by_slug.get(st["slug"])
                if img is None:
                    self.problems.append(f"{w['id']} step {n}: no image {st['slug']}")
                    continue
                img.flows.append((w["id"], n))
        for img in out:
            img.changed = changed(img.slug, img.screen)
            if not img.file.exists():
                self.problems.append(f"image missing: {img.file.relative_to(ROOT)}")
        return out

    def _screen_personas(self, s: signoff_pack.Screen) -> list[str]:
        return [self.pack.persona_label(r) for r in self.pack.personas_of(s)]

    def component_file(self, c: dict[str, Any]) -> Path:
        return brand.src_dir(c["brd"]) / "screenshots" / f"{c['image']}.png"

    def screen_images(self, screen: str) -> list[Img]:
        return [i for i in self.images if i.screen == screen]

    # ------------------------------------------------------------------ texts of the FRS

    def frs_table(self, heading: str, source: Path | None = None) -> list[dict[str, str]]:
        """The first markdown table under a heading of the FRS source (or of another source)."""
        text = (source or (self.pack.brd_dir / self.meta["frs"])).read_text(encoding="utf-8")
        m = re.search(r"^#+ " + re.escape(heading) + r"\s*$", text, re.M)
        if not m:
            return []
        lines = text[m.end():].splitlines()
        rows: list[list[str]] = []
        started = False
        for line in lines:
            if line.startswith("|"):
                started = True
                rows.append([c.strip() for c in line.strip().strip("|").split("|")])
            elif started:
                break
            elif line.startswith("#"):
                break
        if len(rows) < 2:
            return []
        head = rows[0]
        return [dict(zip(head, r)) for r in rows[2:]]

    def screen_standards(self) -> list[dict[str, str]]:
        rows = self.frs_table("Appendix: Screen standards")
        if rows:
            return rows
        # Every FRS carries the same standards; a set without the appendix shows those of Product Maintenance.
        return self.frs_table("Appendix: Screen standards",
                              brand.src_dir("BRD-03") / "FRS_BRD03_PRODUCT_MAINTENANCE.md")

    def persona_texts(self) -> dict[str, dict[str, str]]:
        """Who each persona is and what the persona does in this BRD: ux.yaml personas, and the responsibilities of
        the persona table of the FRS."""
        table = self.frs_table("Personas")
        out = {}
        for role in self.pack.nav_roles():
            label = self.pack.persona_label(role)
            row = None
            for r in table:
                roles = r.get("BIBS role", "")
                if roles and re.search(rf"\b{re.escape(role)}\b", roles):
                    row = r
                    break
            if row is None:
                key = " ".join(label.lower().split()[:2])
                row = next((r for r in table if r.get("Persona", "").lower().startswith(key)), None)
            mine = (self.ux.get("personas") or {}).get(role) or {}
            does = mine.get("does") or (next((v for k, v in (row or {}).items() if k.startswith("Responsibilities")),
                                             "") if row else "")
            if not does:
                self.problems.append(f"persona {role}: no text of what the persona does (ux.yaml personas)")
            out[role] = {"who": mine.get("who", ""), "does": does}
        return out

    def clarifications(self) -> list[dict[str, str]]:
        refs = (self.ux.get("confirm") or {}).get("clarifications") or []
        rows = self.frs_table("Proposed business rules and clarifications for confirmation")
        by_ref = {r.get("Ref", ""): r for r in rows}
        catalogue = self.pack.dir / "catalogue.yaml"
        if catalogue.exists():
            # The handbook of a migration set renders its proposed rules from the catalogue.
            data = yaml.safe_load(catalogue.read_text(encoding="utf-8"))
            for x in data.get("proposals") or []:
                by_ref.setdefault(x[0], {"Ref": x[0], "Topic": x[1], "Proposed rule or screen behaviour": x[2],
                                         "Decision requested from BDOI": x[4]})
        out = []
        for ref in refs:
            if ref not in by_ref:
                self.problems.append(f"ux.yaml: clarification {ref} is not in the clarifications chapter")
                continue
            out.append(by_ref[ref])
        return out


# ============================================================================ images on slides


def _png_size(path: Path) -> tuple[int, int]:
    from PIL import Image  # noqa: PLC0415

    with Image.open(path) as im:
        return im.size


def _cut_rows(path: Path, scale: float, height: int) -> list[tuple[int, int]]:
    """The vertical slices (top, bottom in pixels) of an image shown at `scale` (EMU per pixel) in areas `height`
    EMU high: each cut is moved up to a quiet row (a line of one colour) so that no text line is cut."""
    import numpy as np  # noqa: PLC0415
    from PIL import Image  # noqa: PLC0415

    with Image.open(path) as im:
        gray = np.asarray(im.convert("L"), dtype=np.int16)
    total = gray.shape[0]
    span = max(1, int(height / scale))
    if total <= span:
        return [(0, total)]
    spread = gray.max(axis=1) - gray.min(axis=1)
    out = []
    top = 0
    while top < total:
        end = min(total, top + span)
        if end < total:
            window = spread[max(top + span // 2, end - span // 4):end]
            if len(window):
                quiet = np.where(window <= window.min() + 2)[0]
                end = max(top + span // 2, end - span // 4) + int(quiet[-1]) + 1
        out.append((top, end))
        top = end
    return out


class UxDeck(GuideDeck):
    """The BDOI deck with the UX slides: images at full width split over slides, tables that flow over slides,
    swimlanes, state pills and the Changed badge."""

    def __init__(self, *args: Any, **kwargs: Any):
        super().__init__(*args, **kwargs)
        self.tmp = Path(tempfile.mkdtemp(prefix="uxdeck-"))
        self.counts: dict[str, int] = {}
        self.part = "Cover"
        self._crops = 0

    def _slide(self):
        self.counts[self.part] = self.counts.get(self.part, 0) + 1
        return super()._slide()

    def section_slide(self, title: str, subtitle: str = "") -> None:
        self.part = title
        self.section(title, subtitle)

    # ------------------------------------------------------------------ primitives

    def pill(self, s, x, y, text: str, fill: str, colour: str, w=None, size: int = 10):
        w = w or Inches(max(1.2, 0.09 * len(text) + 0.35))
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        shape.line.color.rgb = _rgb(colour)
        shape.line.width = Pt(1)
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(size), True, brand.FONT
        r.font.color.rgb = _rgb(colour)
        return shape

    def badges(self, s, state: str | None, changed: Sequence[str]) -> None:
        """The state pill and the Changed badge at the top right, under the title rule."""
        x = W - MARGIN
        if changed:
            text = "Changed in this issue: " + ", ".join(changed)
            w = Inches(0.085 * len(text) + 0.4)
            x -= w
            self.pill(s, x, Inches(0.98), text, brand.AMBER_BG, brand.AMBER, w=w)
            x -= Inches(0.1)
        if state:
            label = STATES.get(state, state)
            fill, colour = STATE_TONES.get(state, (brand.BG_BLUE, brand.HEADER_BLUE))
            w = Inches(max(1.3, 0.095 * len(label) + 0.4))
            self.pill(s, x - w, Inches(0.98), label, fill, colour, w=w)

    def picture(self, s, png: Path, x, y, w, h, border: bool = True):
        pw, ph = _png_size(png)
        scale = min(w / pw, h / ph)
        pic = s.shapes.add_picture(str(png), x + int((w - pw * scale) / 2), y, width=int(pw * scale),
                                   height=int(ph * scale))
        if border:
            pic.line.color.rgb = _rgb(brand.BORDER)
            pic.line.width = Pt(0.75)
        return pic

    def missing(self, s, x, y, w, h, name: str) -> None:
        self._rect(s, x, y, w, h, brand.DIRTY_WHITE, brand.BORDER)
        self._text(s, x, y, w, h, f"Image to come: {name}", size=14, colour=brand.MUTED, align=PP_ALIGN.CENTER,
                   anchor=MSO_ANCHOR.MIDDLE)

    def _trim(self, png: Path) -> Path:
        """The image without the empty background below its content (a short menu in a tall window)."""
        import numpy as np  # noqa: PLC0415
        from PIL import Image  # noqa: PLC0415

        with Image.open(png) as im:
            gray = np.asarray(im.convert("L"), dtype=np.int16)
            spread = gray.max(axis=1) - gray.min(axis=1)
            used = np.where(spread >= 12)[0]
            if not len(used):
                return png
            bottom = min(gray.shape[0], int(used[-1]) + 24)
            if bottom >= gray.shape[0] * 0.95:
                return png
            self._crops += 1
            out = self.tmp / f"trim-{self._crops:05d}.png"
            im.crop((0, 0, im.width, bottom)).save(out, optimize=True)
        return out

    def _slice(self, png: Path, top: int, bottom: int) -> Path:
        from PIL import Image  # noqa: PLC0415

        self._crops += 1
        out = self.tmp / f"slice-{self._crops:05d}.png"
        with Image.open(png) as im:
            im.crop((0, top, im.width, bottom)).save(out, optimize=True)
        return out

    def image_slides(self, title: str, png: Path, caption: str = "", state: str | None = None,
                     changed: Sequence[str] = (), notes: str | Sequence[str] = "",
                     first: Any = None, side_w=Inches(0), fit_whole: bool = False) -> int:
        """A screen image at full slide width (never larger than MAX_SCALE), split into slices over as many slides
        as needed; narrow slices sit side by side. With `side_w`, the image leaves a panel of that width on the right,
        which `first(slide)` fills on the first slide. Returns the number of slides."""
        area_w = CW - (side_w + Inches(0.2) if side_w else 0)
        area_h = Y1 - Y0 - CAPTION_H
        if not png.exists():
            s = self.page(title)
            self.badges(s, state, changed)
            self.missing(s, X0, Y0, area_w, area_h, png.name)
            if first:
                first(s)
            if notes:
                self.notes(notes)
            return 1
        png = self._trim(png)
        pw, ph = _png_size(png)
        scale = min(area_w / pw, MAX_SCALE)
        # A screen that fits on one slide at 72 % or more of the full width is shown whole (no split); a landing
        # page is always shown whole.
        fit = min(scale, area_h / ph)
        if fit >= 0.72 * scale or fit_whole:
            scale = fit
        cols = 1
        if pw * scale < area_w / 2:
            cols = max(1, int((area_w + Inches(0.25)) // (pw * scale + Inches(0.25))))
        cuts = _cut_rows(png, scale, area_h)
        cuts = [c for i, c in enumerate(cuts) if i == 0 or not _blank(png, *c)]
        groups = [cuts[i:i + cols] for i in range(0, len(cuts), cols)]
        n = len(groups)
        for k, group in enumerate(groups, start=1):
            s = self.page(title if n == 1 else f"{title} ({k} of {n})")
            self.badges(s, state, changed)
            gap = Inches(0.25)
            total_w = sum(int(pw * scale) for _ in group) + gap * (len(group) - 1)
            x = X0 + max(0, (area_w - total_w) // 2)
            for top, bottom in group:
                part = png if (top, bottom) == (0, ph) else self._slice(png, top, bottom)
                w = int(pw * scale)
                h = int((bottom - top) * scale)
                if h > area_h:  # a slice may end a few pixels lower than planned
                    w, h = int(w * area_h / h), int(area_h)
                pic = s.shapes.add_picture(str(part), x + (int(pw * scale) - w) // 2, Y0, width=w, height=h)
                pic.line.color.rgb = _rgb(brand.BORDER)
                pic.line.width = Pt(0.75)
                x += int(pw * scale) + gap
            if caption:
                self._text(s, X0, Y1 - CAPTION_H, area_w, CAPTION_H,
                           caption + ("" if n == 1 else f" (part {k} of {n})"), size=11, italic=True,
                           colour=brand.MUTED, align=PP_ALIGN.CENTER)
            if first and k == 1:
                first(s)
            if notes:
                self.notes(notes)
        return n

    # ------------------------------------------------------------------ tables that flow over slides

    def flow_table(self, title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]],
                   widths: Sequence[float], size: float = 10, notes: str | Sequence[str] = "",
                   top=None, bottom=None, first_col_bold: bool = False) -> int:
        """A table over as many slides as its rows need (row heights estimated from the text), with the header on
        every slide. Returns the number of slides."""
        top = top or Y0 + Inches(0.07)
        limit = (bottom or Y1) - top
        total = sum(widths)
        col_w = [CW * w / total for w in widths]
        chars_per_in = 15.2 * 10 / size

        def height(cells: Sequence[Any], bold: bool = False) -> int:
            lines = 1
            for text, w in zip(cells, col_w):
                width_in = max(0.3, w / Inches(1) - 0.16)
                parts = str(text if text is not None else "").split("\n")
                n = sum(max(1, math.ceil(len(p) * (1.08 if bold else 1) / (width_in * chars_per_in))) for p in parts)
                lines = max(lines, n)
            return int(Pt(size * 1.2) * lines + Inches(0.1))

        head_h = height(headers, True)
        chunks: list[list[Sequence[Any]]] = []
        cur: list[Sequence[Any]] = []
        used = head_h
        for r in rows:
            h = height(r)
            if cur and used + h > limit:
                chunks.append(cur)
                cur, used = [], head_h
            cur.append(r)
            used += h
        if cur or not chunks:
            chunks.append(cur)
        for k, chunk in enumerate(chunks, start=1):
            s = self.page(title if len(chunks) == 1 else f"{title} ({k} of {len(chunks)})")
            self.table_at(s, top, headers, chunk, col_w, size, height, first_col_bold)
            if notes:
                self.notes(notes)
        return len(chunks)

    def table_at(self, s, top, headers: Sequence[str], rows: Sequence[Sequence[Any]], col_w: Sequence[int],
                 size: float, height: Any = None, first_col_bold: bool = False, x=None) -> Any:
        n_rows = len(rows) + 1
        x = X0 if x is None else x
        shape = s.shapes.add_table(n_rows, len(headers), x, top, sum(col_w), Inches(0.3) * n_rows)
        tbl = shape.table
        for i, w in enumerate(col_w):
            tbl.columns[i].width = Emu(int(w))
        for r in range(n_rows):
            data = headers if r == 0 else rows[r - 1]
            if height:
                tbl.rows[r].height = Emu(height(data, r == 0))
            for c in range(len(headers)):
                cell = tbl.cell(r, c)
                value = data[c] if c < len(data) else ""
                cell.text = str(value if value is not None else "")
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.name = brand.FONT
                        run.font.size = Pt(size)
                        run.font.bold = r == 0 or (first_col_bold and c == 0)
                        run.font.color.rgb = _rgb(brand.WHITE if r == 0 else brand.TEXT)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(brand.HEADER_BLUE if r == 0 else
                                                (brand.BG_BLUE if r % 2 == 0 else brand.WHITE))
        return shape

    def key_values(self, s, x, y, w, rows: Sequence[tuple[str, str]], size: float = 11, label_w=Inches(1.9)) -> None:
        def height(cells: Sequence[Any], bold: bool = False) -> int:
            text = str(cells[1])
            width_in = (w - label_w) / Inches(1) - 0.16
            lines = sum(max(1, math.ceil(len(p) / (width_in * 15.2 * 10 / size))) for p in text.split("\n"))
            return int(Pt(size * 1.2) * lines + Inches(0.1))

        shape = s.shapes.add_table(len(rows), 2, x, y, w, Inches(0.3) * len(rows))
        tbl = shape.table
        tbl.columns[0].width = Emu(int(label_w))
        tbl.columns[1].width = Emu(int(w - label_w))
        tbl.first_row = False
        for r, (k, v) in enumerate(rows):
            tbl.rows[r].height = Emu(height((k, v)))
            for c, text in enumerate((k, v)):
                cell = tbl.cell(r, c)
                cell.text = str(text)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.name, run.font.size = brand.FONT, Pt(size)
                        run.font.bold = c == 0
                        run.font.color.rgb = _rgb(brand.HEADER_BLUE if c == 0 else brand.TEXT)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(brand.BG_BLUE if c == 0 else brand.WHITE)

    # ------------------------------------------------------------------ swimlanes

    def swimlane(self, title: str, lanes: Sequence[str], steps: Sequence[dict[str, Any]], first_no: int,
                 notes: Sequence[str]) -> None:
        """Persona lanes, one box per step in its lane and column, arrows between the steps (a hand-off to another
        persona in CTA Blue); approval steps have a yellow outline and the tag Approval."""
        s = self.page(title)
        lane_w = Inches(1.95)
        top = Y0 + Inches(0.05)
        lane_h = min(Inches(0.95), int((Y1 - top - Inches(0.35)) / max(1, len(lanes))))
        cols = len(steps)
        col_w = int((CW - lane_w) / max(cols, 1))
        for i, lane in enumerate(lanes):
            y = top + i * lane_h
            self._rect(s, X0, y, lane_w, lane_h - Inches(0.04), brand.HEADER_BLUE)
            self._text(s, X0 + Inches(0.05), y, lane_w - Inches(0.1), lane_h - Inches(0.04), lane, size=10,
                       bold=True, colour=brand.WHITE, anchor=MSO_ANCHOR.MIDDLE)
            self._rect(s, X0 + lane_w, y, CW - lane_w, lane_h - Inches(0.04),
                       brand.BG_BLUE if i % 2 == 0 else brand.DIRTY_WHITE)
        boxes = []
        for c, st in enumerate(steps):
            lane = lanes.index(st["lane"])
            bw, bh = col_w - Inches(0.22), lane_h - Inches(0.16)
            x = X0 + lane_w + c * col_w + Inches(0.11)
            y = top + lane * lane_h + Inches(0.06)
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh)
            box.fill.solid()
            box.fill.fore_color.rgb = _rgb(brand.WHITE)
            box.line.color.rgb = _rgb(brand.YELLOW if st["approval"] else brand.FIELD_BLUE)
            box.line.width = Pt(2.25 if st["approval"] else 1)
            box.shadow.inherit = False
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.04)
            tf.margin_top = tf.margin_bottom = Inches(0.02)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = f"{first_no + c}. "
            r.font.bold, r.font.size, r.font.name = True, Pt(9), brand.FONT
            r.font.color.rgb = _rgb(brand.HEADER_BLUE)
            r = p.add_run()
            r.text = st["label"]
            r.font.size, r.font.name = Pt(8), brand.FONT
            r.font.color.rgb = _rgb(brand.NEAR_BLACK)
            if st["approval"]:
                tag = self.pill(s, x + bw - Inches(0.78), y - Inches(0.05), "Approval", brand.AMBER_BG, brand.AMBER,
                                w=Inches(0.8), size=7)
                tag.height = Inches(0.2)
            boxes.append((x, y, bw, bh, lane))
        for (x1, y1, w1, h1, l1), (x2, y2, w2, h2, l2) in zip(boxes, boxes[1:]):
            handoff = l1 != l2
            if handoff:
                line = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, x1 + w1, y1 + h1 // 2, x2, y2 + h2 // 2)
            else:
                line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1 + w1, y1 + h1 // 2, x2, y2 + h2 // 2)
            line.line.color.rgb = _rgb(brand.CTA_BLUE if handoff else brand.MUTED)
            line.line.width = Pt(1.5 if handoff else 1)
            ln = line.line._get_or_add_ln()
            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
            tail.set("w", "med")
            tail.set("len", "med")
        legend_y = Y1 - Inches(0.3)
        self._text(s, X0, legend_y, CW, Inches(0.3), "Boxes: step number and screen. Blue arrow: hand-off to another "
                   "persona. Yellow outline: an approval, authorisation, sign-off or validation step.", size=9,
                   italic=True, colour=brand.MUTED)
        self.notes(notes)


def _blank(path: Path, top: int, bottom: int) -> bool:
    """True when the rows top..bottom of an image hold (almost) nothing but the background and card borders."""
    import numpy as np  # noqa: PLC0415
    from PIL import Image  # noqa: PLC0415

    with Image.open(path) as im:
        gray = np.asarray(im.convert("L").crop((0, top, im.width, bottom)), dtype=np.int16)
    if gray.size == 0:
        return True
    rows = (gray.max(axis=1) - gray.min(axis=1)) >= 12
    return float(rows.mean()) < 0.04


APPROVAL = re.compile(r"\b(approv\w*|authori[sz]\w*|sign(?:s|ed)? off|validate\w*|recommend\w*|accept\w*|"
                      r"sign(?:s)? (?:the )?(?:validation|reconciliation|mapping)|\bGO\b)", re.I)


# ============================================================================ deck


def tokens() -> list[tuple[str, str, str]]:
    """(token, value, comment) of the design tokens of the screens."""
    out = []
    for m in re.finditer(r"--([\w-]+):\s*([^;]+);(?:\s*/\*\s*(.*?)\s*\*/)?", TOKENS.read_text(encoding="utf-8")):
        out.append((m.group(1), m.group(2).strip(), (m.group(3) or "").strip()))
    return out


def build_deck(ux: Ux) -> tuple[Path, dict[str, int]]:
    p = ux.pack
    m = ux.meta
    deck = UxDeck(f"UX Screen Deck {ux.label} {ux.module}", version=ux.version, date=str(m["date"]),
                  subtitle=f"{ux.label} {ux.module}: every screen, persona by persona and flow by flow, for the "
                           "BDOI UX Design team")
    issue = ux.ux.get("issue") or {}
    doc02 = "Data Migration Handbook" if ux.frs_name() == "Handbook" else "FRS"
    screens = p.screens
    shown = [i for i in ux.images]
    n_states = len({(i.screen, i.state) for i in shown if i.kind in ("screen", "step", "ux")})

    # ------------------------------------------------------------------ 1. cover and use
    deck.part = "Cover and use"
    deck.title(f"UX Screen Deck: {ux.label} {ux.module}")
    deck.notes(f"File 07 of the {ux.label} {ux.module} sign-off set, release set v{ux.version}, for the BDOI UX Design "
               "team. The screens are those of the proposed system on SIT with seed data; the names and numbers on "
               "them are fictitious.")
    deck.content("Purpose of this deck for the UX Design team", [
        f"Every screen of {ux.label} {ux.module} as the {doc02} v{ux.version} specifies it, so that the UX Design "
        "team designs in parallel with the business review",
        (1, f"{len(screens)} screens, {len(shown)} screen images, {n_states} screen states, "
            f"{len(p.nav_roles())} personas, {len(ux.flows)} end-to-end flows"),
        "Persona by persona: who they are, their landing page and menu, the screens they open, the flows they take "
        "part in",
        "Flow by flow: every end-to-end flow as a swimlane of the personas, then one slide per step with the screen, "
        "what the user does, what the system does, the next step and the messages shown",
        "Screen by screen: the numbered callouts, fields, actions and every state of each screen",
        "Business and design content only: the screens, the words on them, the rules the user meets and the "
        "points to confirm",
    ], size=16)
    deck.notes("The deck does not replace the FRS: the FRS is signed by the business; the deck is the working "
               "material of the UX Design team and follows each issue of the FRS.")
    deck.columns("How to use the deck, the register and the image package", [
        (f"07 UX Screen Deck", [
            "Read the design foundation and the navigation first",
            "Then a persona, a flow or a screen: the screen ID (SCR-...) links the three parts",
            "Each image slide names its state; tall screens continue on the next slide",
            "Speaker notes give the image file and the FRS section",
        ]),
        (f"08 UX screen register", [
            "One row per screen image, with the screen ID, personas, flow and step, state, FRS section and "
            "requirements",
            "The UX Design team sets UXD status and UXD comments on each row",
            "Changed in this issue marks the rows to revisit, with the change reference",
            "Questions go to the Questions and comments sheet",
        ]),
        (f"09 Image package", [
            "Every screen image at twice the screen resolution (PNG), named <Screen ID>_<state>.png as in the "
            "register",
            "The crops of the shared components in the folder components",
            "The register as CSV, to import into the design tool",
            "Replaced as a whole with each issue of the set",
        ]),
    ], size=12)
    deck.notes("Start with the register when a new issue arrives: filter Changed in this issue = Y.")
    deck.flow("How an FRS change reaches the UX Design team", [
        {"step": 1, "name": "Change raised", "outputs": "A comment of the sign-off workbook or a change request of the "
         "Change Management Register gets its change reference", "duration": "Business and project team"},
        {"step": 2, "name": "Screens named", "outputs": "The change names the FRS section; the map at the end of this "
         "deck turns the section into the screen IDs", "duration": "Project team"},
        {"step": 3, "name": "Delta re-issue", "outputs": "The affected screens are captured again; the register rows "
         "get Changed in this issue = Y with the change reference; the slides carry the Changed badge",
         "duration": "Next version of the set"},
        {"step": 4, "name": "UXD revises", "outputs": "The UX Design team revises only the changed screens and sets "
         "their UXD status; the other rows keep theirs", "duration": "UX Design team"},
    ])
    deck.notes("A change never waits for the whole set: the delta issue carries the new images and the register "
               "with the changed rows flagged; the unchanged rows keep their UXD status and comments.")
    changes = issue.get("changes") or []
    rows = []
    for c in changes:
        imgs = [i for i in ux.images if c["ref"] in i.changed]
        scr = sorted({i.screen for i in imgs})
        rows.append([c["ref"], _words(c["what"]), ", ".join(scr) or "-", str(len(imgs))])
    if rows:
        deck.flow_table(f"Changes in this issue (v{ux.version}, {issue.get('date', m['date'])})",
                        ["Change reference", "What changed", "Screens", "Images"], rows, [2.2, 7.4, 2.6, 1.0],
                        size=11, notes="These images carry the Changed badge in this deck and Changed in this issue "
                                       "= Y in the register.")
    else:
        deck.content(f"Changes in this issue (v{ux.version})", ["First issue of the UX screen documents of the set: "
                                                                 "every screen is new to the UX Design team"])
    own = p.ownership or {}
    roles = own.get("roles") or []
    rows = [[r["capacity"], r["role"], r["org"], _words(r["confirms"])] for r in roles]
    rows.append(["Receives 07 to 09", "BDOI UX Design team (UXD)", "BDOI",
                 "Designs the screens from this deck; keeps the UXD status and comments of the register; raises "
                 "questions on the Questions and comments sheet"])
    deck.flow_table("Contacts and roles", ["Capacity", "Role", "Organisation", "Part in the set"], rows,
                    [1.7, 4.3, 1.6, 5.4], size=9,
                    notes=f"Roles as on the BRD approval sheet ({own.get('source', '-')}); the names are in the "
                          "sign-off certificate of the workbook.")

    # ------------------------------------------------------------------ 2. design foundation
    deck.section_slide("Design foundation", "Colours, type, spacing, icons, the shared components and the screen "
                                            "standards as the screens use them")
    toks = tokens()
    colours = [(n, v, c) for n, v, c in toks if v.startswith("#")]
    s = deck.page("Colour tokens")
    per_row = 7
    sw_w = (CW - Inches(0.2) * (per_row - 1)) / per_row
    for i, (name, value, comment) in enumerate(colours):
        r, c = divmod(i, per_row)
        x = X0 + c * (sw_w + Inches(0.2))
        y = Y0 + r * Inches(1.08)
        deck._rect(s, x, y, sw_w, Inches(0.42), value.lstrip("#").upper(), brand.BORDER)
        deck._text(s, x, y + Inches(0.43), sw_w, Inches(0.22), name, size=8, bold=True, colour=brand.NEAR_BLACK)
        deck._text(s, x, y + Inches(0.62), sw_w, Inches(0.4), f"{value.upper()} {comment.split(':')[0][:40]}",
                   size=7, colour=brand.MUTED)
    deck.notes("The colour tokens of the screens (BDO Style Guide names in the comments). Blue is the dominant "
               "colour; yellow is an accent only, never a fill behind text; never pure black.")
    deck.flow_table("Type scale", ["Use", "Size / line height", "Weight and colour"], [
        ["Page title (h1)", "28 / 36", "Bold, Near Black #2E2E2E"],
        ["Section title", "18 / 28", "Bold, Near Black"],
        ["Group title", "16 / 24", "Bold, Near Black"],
        ["Lead paragraph", "18 / 28", "Regular, Dark Grey #4B4B4B"],
        ["Content", "16 / 24", "Regular, Dark Grey"],
        ["Application default (tables, forms)", "14 / 20", "Regular, Dark Grey"],
        ["Fine print, format hints", "12 / 16", "Regular, Base Grey #656565"],
        ["Buttons", "14, Title Case", "Bold; white on CTA Blue (primary), CTA Blue on white (secondary)"],
        ["Status pill", "12, bold, centred", "24 px high, 112 to 184 px wide, one colour per state group"],
    ], [4.2, 3.0, 5.8], size=12, notes="Nunito everywhere; Arial is the approved fallback. Display sizes (72 to 24) "
                                        "only for dashboards and landing panels.")
    sizes = [(n, v, c) for n, v, c in toks if not v.startswith("#") and re.match(r"^[\d.]+px$|^\d+vh$", v)]
    deck.flow_table("Spacing, sizes and radius", ["Token", "Value", "Use"],
                    [[n, v, c or "-"] for n, v, c in sizes], [3.4, 1.6, 8.0], size=11,
                    notes="An 8 px baseline grid; desktop web uses a 12-column grid; cards have radius 12 on the "
                          "Dirty White page.")
    deck.content("Iconography", [
        "Outline icons in CTA Blue, one stroke weight, at 16, 24, 32 and 48 px (the Lucide outline set as the "
        "equivalent of the BDO icon set)",
        "Menu entries: a 20 px icon before the label",
        "Messages: the icon of the kind in its colour (error, warning, information, success)",
        "Actions: an icon only with its label or its tooltip; destructive actions in red, last",
        "Flags and statuses are text chips and pills, never an icon alone",
    ], size=16)
    deck.notes("The BDO outline icon master files replace the equivalents when BDO Marketing Communications supplies "
               "them.")
    for comp in ux.components:
        png = ux.component_file(comp)
        where = [sid for sid in comp.get("screens", {}).get(ux.brd, [])]
        side = [_words(x) for x in comp["rules"]]
        title = f"Component: {comp['name']}"
        if where:
            side.append("In this set: " + ", ".join(where))
        s = deck.page(title)
        img_w = Inches(8.4)
        if png.exists():
            deck.picture(s, png, X0, Y0, img_w, Y1 - Y0 - CAPTION_H)
        else:
            deck.missing(s, X0, Y0, img_w, Y1 - Y0 - CAPTION_H, png.name)
        deck._text(s, X0, Y1 - CAPTION_H, img_w, CAPTION_H, _words(comp["caption"]), size=10, italic=True,
                   colour=brand.MUTED, align=PP_ALIGN.CENTER)
        x = X0 + img_w + Inches(0.25)
        w = W - MARGIN - x
        deck._rect(s, x, Y0, w, Y1 - Y0, brand.BG_BLUE)
        deck._rect(s, x, Y0, Pt(4), Y1 - Y0, brand.HEADER_BLUE)
        deck._bullets(s, x + Inches(0.15), Y0 + Inches(0.1), w - Inches(0.25), Y1 - Y0 - Inches(0.2), side, size=11)
        deck.notes(f"{_words(comp['what'])} Image: components/{png.name} of the image package.")
    std = ux.screen_standards()
    deck.flow_table("Screen standards", ["Area", "Standard"], [[r["Area"], r["Standard"]] for r in std],
                    [2.4, 10.6], size=10, first_col_bold=True,
                    notes="The screen standards appendix of every FRS, in the same words: BDOI agrees the look of the "
                          "screens once and checks every screenshot against it.")

    # ------------------------------------------------------------------ 3. navigation
    deck.section_slide("Navigation", "The landing page and the full menu of each persona")
    menus = p.menus()
    for n, role in enumerate(p.nav_roles(), start=1):
        label = p.persona_label(role)
        land = ux.by_slug.get(f"ux-nav-{n:02d}-landing")
        menu = ux.by_slug.get(f"ux-nav-{n:02d}-menu")
        if land:
            deck.image_slides(f"{label}: landing page", land.file, f"{p.nav_id(n)} The page {label} sees after "
                              "sign-in", "landing", land.changed, notes=f"Image {land.name}.", fit_whole=True)
        if menu:
            deck.image_slides(f"{label}: menu", menu.file, f"{p.nav_id(n)} Every group of the menu open",
                              "menu", menu.changed, notes=f"Image {menu.name}.")
        rows_ = menus.get(role) or []
        own_rows = [[r["group"], r["section"], r["screen"]] for r in rows_ if r["own"]]
        others: dict[str, int] = {}
        for r in rows_:
            if not r["own"]:
                others[r["brd"]] = others.get(r["brd"], 0) + 1
        extra = "; ".join(f"{b}: {c}" for b, c in sorted(others.items()))
        if own_rows:
            deck.flow_table(f"{label}: menu entries of {ux.module}", ["Group", "Section", "Screen"], own_rows,
                            [3.0, 4.0, 6.0], size=11,
                            notes=f"{len(rows_)} menu entries in all; entries of the other BRDs: {extra or 'none'}. "
                                  f"The full menu is in {doc02} section {ux.section_of('Menu by persona')}.")

    # ------------------------------------------------------------------ 4. personas
    deck.section_slide("Personas", "Who they are, what they do in this BRD, their screens and flows")
    texts = ux.persona_texts()
    for n, role in enumerate(p.nav_roles(), start=1):
        label = p.persona_label(role)
        mine = [s for s in screens if role in p.personas_of(s)]
        steps = [(w["id"], k) for w in ux.flows for k, st in enumerate(w["steps"], start=1) if st["role"] == role]
        by_flow: dict[str, list[int]] = {}
        for fid, k in steps:
            by_flow.setdefault(fid, []).append(k)
        s = deck.page(f"Persona: {label}")
        half = (CW - Inches(0.3)) / 2
        kv = [("Who", texts[role]["who"] or label), ("What they do here", _words(texts[role]["does"])),
              ("Landing page and menu", p.nav_id(n)),
              ("Flows", "; ".join(f"{fid} steps {', '.join(map(str, ks))}" for fid, ks in by_flow.items()) or
               "No step of the flows of this set")]
        deck.key_values(s, X0, Y0 + Inches(0.05), half, kv, size=11)
        x2 = X0 + half + Inches(0.3)
        deck._text(s, x2, Y0, half, Inches(0.35), f"Screens this persona opens ({len(mine)})", size=13, bold=True,
                   colour=brand.HEADER_BLUE)
        items = [f"{sc.id} {sc.title}" for sc in mine] or ["None of the screens of this set"]
        size = 11 if len(items) <= 14 else (9 if len(items) <= 22 else 8)
        deck._bullets(s, x2, Y0 + Inches(0.4), half, Y1 - Y0 - Inches(0.4), items, size=size)
        deck.notes(f"{label}. Menu and landing page: {p.nav_id(n)} in the navigation section.")

    # ------------------------------------------------------------------ 5. flows
    deck.section_slide("End-to-end flows", "Every flow as a swimlane of the personas, then one slide per step")
    for w in ux.flows:
        steps = w["steps"]
        lanes_all: list[str] = []
        for st in steps:
            lab = p.persona_label(st["role"])
            if lab not in lanes_all:
                lanes_all.append(lab)
        kv = [("Flow", f"{w['id']} {w['title']}"), ("Summary", w["summary"] or "-"),
              ("Seed data", w["data"] or "-"), ("Personas", ", ".join(lanes_all)),
              ("Steps", f"{len(steps)} steps, "
                        f"{sum(1 for a, b in zip(steps, steps[1:]) if a['role'] != b['role'])} hand-offs, "
                        f"{sum(1 for st in steps if APPROVAL.search(st['does']))} approval steps"),
              (f"{doc02} section", ux.section_of(w["id"]) if w["captured"] else
               "The screens of the flow in the screen specifications (images of the screen catalogue)")]
        s = deck.page(f"{w['id']} {w['title']}")
        deck.key_values(s, X0, Y0 + Inches(0.05), CW, kv, size=12, label_w=Inches(2.2))
        deck.notes("Downstream: " + "; ".join(_words(x) for x in w["downstream"]) if w["downstream"] else w["summary"])
        per = 8
        chunks = [steps[i:i + per] for i in range(0, len(steps), per)]
        for k, chunk in enumerate(chunks):
            lanes = []
            for st in chunk:
                lab = p.persona_label(st["role"])
                if lab not in lanes:
                    lanes.append(lab)
            data = []
            for st in chunk:
                scr = p.step_screen(st["screen"])
                data.append({"lane": p.persona_label(st["role"]), "label": scr.title if scr else st["screen"],
                             "approval": bool(APPROVAL.search(st["does"]))})
            part = f" ({k + 1} of {len(chunks)})" if len(chunks) > 1 else ""
            deck.swimlane(f"{w['id']} swimlane{part}", lanes, data, k * per + 1,
                          [f"{k * per + i + 1}. {p.persona_label(st['role'])}: {st['does']}"
                           for i, st in enumerate(chunk)])
        for n, st in enumerate(steps, start=1):
            scr = p.step_screen(st["screen"])
            img = ux.by_slug.get(st["slug"])
            png = img.file if img else p.shot_file(st["slug"])
            nxt = steps[n] if n < len(steps) else None
            nxt_text = (f"Step {n + 1}: {p.persona_label(nxt['role'])} on "
                        f"{(p.step_screen(nxt['screen']).title if p.step_screen(nxt['screen']) else nxt['screen'])}"
                        if nxt else "End of the flow")
            blocks = [("Persona does", f"{p.persona_label(st['role'])}: {st['does']}"),
                      ("System does", st["system"] or "-"),
                      ("Notices and messages shown", st["sees"] or "-"),
                      ("Next", nxt_text)]

            def strip(s, blocks=blocks):
                pw_ = Inches(3.55)
                x = W - MARGIN - pw_
                bh = (Y1 - Y0 - Inches(0.1) * 3) / 4
                for i, (head, text) in enumerate(blocks):
                    y = Y0 + i * (bh + Inches(0.1))
                    deck._rect(s, x, y, pw_, bh, brand.BG_BLUE)
                    deck._rect(s, x, y, Pt(4), bh, brand.HEADER_BLUE)
                    deck._text(s, x + Inches(0.12), y + Inches(0.04), pw_ - Inches(0.2), Inches(0.25), head.upper(),
                               size=9, bold=True, colour=brand.HEADER_BLUE)
                    size = 10 if len(text) <= 170 else (9 if len(text) <= 240 else 8)
                    deck._text(s, x + Inches(0.12), y + Inches(0.28), pw_ - Inches(0.2), bh - Inches(0.3), text,
                               size=size, colour=brand.NEAR_BLACK)

            title = f"{w['id']} step {n} of {len(steps)}: {scr.title if scr else st['screen']}"
            deck.image_slides(title, png, f"{st['screen']} {img.caption if img and not w['captured'] else ''}".strip()
                              if img else st["screen"], img.state if img else None, img.changed if img else (),
                              notes=[f"Image {img.name if img else png.name}.", f"Persona does: {st['does']}",
                                     f"System does: {st['system']}", f"Shown: {st['sees']}"],
                              first=strip, side_w=Inches(3.55))

    # ------------------------------------------------------------------ 6. screen catalogue
    deck.section_slide("Screen catalogue", "Every screen with its numbered callouts, fields, actions and every state")
    for area in p.areas:
        for sc in [x for x in screens if x.area == area]:
            imgs = ux.screen_images(sc.id)
            personas = ", ".join(p.persona_label(r) for r in p.personas_of(sc)) or "Every user"
            entries = [f"Menu: {p.menu_path(sc)}"] + [str(e) for e in sc.get("entry") or []]
            s = deck.page(f"{sc.id} {sc.title}")
            deck.badges(s, None, sorted({c for i in imgs for c in i.changed}))
            half = CW * 0.58
            deck.key_values(s, X0, Y0 + Inches(0.05), half, [
                ("Area", area), ("Purpose", _words(sc.get("purpose"))), ("Personas", personas),
                ("Ways in", "\n".join(entries)), (f"{doc02} section", ux.section_of(sc.id)),
                ("Requirements", ", ".join(sc.get("frs") or []) or "-"),
                ("Outcome", _words(sc.get("outcome")) or "-")], size=9.5, label_w=Inches(1.5))
            x2 = X0 + half + Inches(0.25)
            present: dict[str, list[str]] = {}
            for i in imgs:
                present.setdefault(i.state, []).append(i.name)
            rows_ = []
            for st in SCREEN_STATES:
                names = present.get(st, [])
                rows_.append([STATES[st], f"{len(names)} image{'s' if len(names) != 1 else ''}" if names else "-"])
            deck._text(s, x2, Y0, CW - half - Inches(0.25), Inches(0.3), f"States shown ({len(imgs)} images)",
                       size=12, bold=True, colour=brand.HEADER_BLUE)
            col_w = [int((CW - half - Inches(0.25)) * 0.62), int((CW - half - Inches(0.25)) * 0.38)]
            deck.table_at(s, Y0 + Inches(0.35), ["State", "Images"], rows_, col_w, 9.5, x=x2)
            deck.notes([f"Images of {sc.id}: " + ", ".join(i.name for i in imgs),
                        "A dash: the state does not apply to this screen or is shown on another screen of the flow."])
            if sc.fields:
                deck.flow_table(f"{sc.id} fields and columns (No. = callout on the first image)",
                                ["No.", "Section", "Field", "Type", "Length / format", "Mand.", "Validation",
                                 "Message when it fails"],
                                [[f["no"], f["section"], f["label"], f["type"], f["format"], f["mandatory"],
                                  f["validation"], f["message"]] for f in sc.fields],
                                [0.5, 1.6, 2.2, 1.2, 2.4, 0.7, 2.4, 2.4], size=8.5)
            if sc.actions:
                deck.flow_table(f"{sc.id} actions", ["Button", "Who", "Enabled when", "What happens",
                                                     "Resulting status", "Notification"],
                                [[a["button"], a["who"], a["when"], a["what"], a["status"], a["notification"]]
                                 for a in sc.actions], [1.8, 2.0, 2.0, 4.2, 1.5, 1.8], size=8.5)
            order = {k: i for i, k in enumerate(STATES)}
            rules = [_words(r) for r in sc.get("rules") or []]
            if rules:
                deck.flow_table(f"{sc.id} rules the user meets", ["Rule"], [[r] for r in rules], [1], size=10)
            order = {k: i for i, k in enumerate(STATES)}
            # The images of the captured flow steps are shown in their flow; the catalogue names them.
            steps = [i for i in imgs if i.kind == "step"]
            if steps:
                deck.flow_table(f"{sc.id} states shown in the flows", ["State", "Flow and step", "What the image shows",
                                                                     "Image"],
                                [[i.state_label, ", ".join(f"{a} step {b}" for a, b in i.flows),
                                  i.caption.split(": ", 1)[-1], i.name] for i in
                                 sorted(steps, key=lambda i: (order.get(i.state, 99), i.slug))],
                                [1.8, 1.6, 6.6, 3.0], size=9,
                                notes="Each of these images has its slide in the end-to-end flows section.")
            for img in sorted((i for i in imgs if i.kind != "step"),
                              key=lambda i: (order.get(i.state, 99), i.kind != "screen", i.slug)):
                where = f"{img.flows and ', '.join(f'{a} step {b}' for a, b in img.flows) or ''}"
                cap = img.caption if img.kind != "step" else img.caption
                deck.image_slides(f"{sc.id} {sc.title}: {img.state_label}", img.file, cap, img.state, img.changed,
                                  notes=[f"Image {img.name}; {doc02} section {img.section}.",
                                         f"Personas: {', '.join(img.personas) or '-'}." + (f" Flows: {where}."
                                                                                          if where else ""),
                                         "Shown in the FRS." if img.in_frs else "Taken for this deck (not in the FRS)."])

    # ------------------------------------------------------------------ 7. outputs
    deck.section_slide("Outputs", "Documents, e-mails and notifications, and the exports as the user receives them")
    for d in p.documents:
        img = ux.by_slug.get(d["shot"])
        s = deck.page(f"{d['id']} {d['name']}")
        img_w = Inches(6.2)
        if img and img.file.exists():
            deck.picture(s, img.file, X0, Y0, img_w, Y1 - Y0)
        else:
            deck.missing(s, X0, Y0, img_w, Y1 - Y0, f"{d['shot']}.png")
        x = X0 + img_w + Inches(0.3)
        deck.key_values(s, x, Y0 + Inches(0.05), W - MARGIN - x, [
            ("Template", _words(d["template"])), ("Format", _words(d["format"])),
            ("Produced by", _words(d["produced"])), ("Password protected", _words(d["protected"])),
            ("Fields", "\n".join(_words(f) for f in d["fields"][:12]) + ("\n…" if len(d["fields"]) > 12 else ""))],
            size=9, label_w=Inches(1.5))
        deck.notes([f"Image {img.name if img else d['shot']}.", "Fields: " + "; ".join(_words(f) for f in d["fields"])])
    if p.notifications:
        deck.flow_table("Notifications and e-mails", ["ID", "Channel", "Trigger", "Recipient", "Content"],
                        [[n_["id"], n_["channel"], _words(n_["trigger"]), _words(n_["recipient"]),
                          _words(n_["template"])] for n_ in p.notifications], [0.8, 1.3, 3.2, 2.6, 5.1], size=8.5,
                        notes="In-app notices appear under the bell and on the Notifications page; e-mails follow "
                              "the protected e-mail rule of the set.")
    exports = []
    for sc in screens:
        for a in sc.actions:
            if re.search(r"\b(download|export|print|runbook|template|workbook)\w*", a["what"], re.I) and \
                    re.search(r"^(Downloads|Exports|Prints|Opens the (?:PDF|print))", a["what"]):
                exports.append([sc.id, a["button"], a["who"], a["what"]])
    if exports:
        deck.flow_table("Exports and downloads", ["Screen", "Button", "Who", "What the user receives"], exports,
                        [1.4, 2.4, 3.0, 6.2], size=9)
    if not p.documents and not p.notifications and not exports:
        deck.content("Outputs", ["No generated document in this set"])

    # ------------------------------------------------------------------ 8. UX points for confirmation
    deck.section_slide("UX points for confirmation", "Presentation items to confirm, screen observations, "
                                                     "accessibility and the desktop assumptions")
    clar = ux.clarifications()
    if clar:
        deck.flow_table("Presentation items of the clarifications chapter", ["Ref", "Topic", "Proposed", "To confirm"],
                        [[r.get("Ref", ""), r.get("Topic", ""), r.get("Proposed rule or screen behaviour", ""),
                          r.get("Decision requested from BDOI", "")] for r in clar], [1.2, 2.4, 6.0, 3.4], size=9,
                        notes=f"From the chapter Proposed business rules and clarifications for confirmation of the "
                              f"{doc02} (section {ux.section_of('Proposed business rules and clarifications for confirmation')}).")
    findings = (ux.ux.get("confirm") or {}).get("findings") or []
    if findings:
        deck.flow_table("Screen observations: proposals for the UX Design team",
                        ["No.", "Screens", "Observation", "Proposal"],
                        [[f"UX-{ux.code}-{i:02d}", ", ".join(f["screens"]), _words(f["observation"]),
                          _words(f["proposal"])] for i, f in enumerate(findings, start=1)],
                        [1.1, 1.8, 5.0, 5.1], size=9,
                        notes="Neutral proposals for the design; the business decides them with the review of the "
                              "set.")
    deck.flow_table("Accessibility notes", ["Topic", "What the screens do", "For the UX Design team"], [
        ["Keyboard", "Tabs move with the arrow keys, Home and End; table rows open with Enter; tree tables expand and "
                     "collapse with the arrow keys; dialogs close with Escape", "Keep every action reachable without "
                                                                              "a mouse; keep the order of the tabs "
                                                                              "and fields"],
        ["Focus", "A visible focus ring (3 px CTA Blue halo) on every control; a dialog takes the focus and returns "
                  "it to the button that opened it", "Never hide the focus ring; show it in the designs"],
        ["Contrast", "Text and pills meet AA contrast; no dark text on a dark fill or light text on a light fill; "
                     "yellow only as an accent", "Check new colours of pills, tags and charts for AA"],
        ["Error identification", "Errors are named in words under the field with the field outlined; a long form "
                                 "lists its errors at the top with a link to the first field; colour is never the "
                                 "only sign", "Keep the text of the message with its icon and title"],
        ["Labels and names", "Every field has a visible label above it; icons have a text label or a tooltip; users "
                             "are named by their names", "No placeholder as the only label"],
    ] + [[a["topic"], _words(a["screens"]), _words(a["design"])] for a in (ux.ux.get("confirm") or {}).get(
        "accessibility") or []], [2.0, 6.0, 5.0], size=10)
    deck.content("Desktop and responsive assumptions", [
        "BIBS is a desktop web application used in the office; the designs are for 1440 x 900 at 100 % zoom and "
        "work from 1280 px wide",
        "The images of this deck are taken at 1440 x 900 at twice the screen resolution",
        "Narrow widths: the menu collapses; the workflow step bar keeps the first, the last and the current stage "
        "and its neighbours; tables scroll inside their card with the header row in view",
        "No phone or tablet layout is in scope; printing uses the BDO print layout (A4 landscape, no menu or "
        "buttons)",
        "Browser zoom up to 200 % keeps every function usable",
    ] + [_words(x) for x in (ux.ux.get("confirm") or {}).get("desktop") or []], size=15)
    deck.notes("To confirm with BDOI with the review of the set.")

    # ------------------------------------------------------------------ 9. appendix
    deck.section_slide("Appendix", "Screen ID, FRS section and requirements")
    rows = []
    for sc in screens:
        imgs = ux.screen_images(sc.id)
        states = sorted({i.state for i in imgs}, key=lambda x: list(STATES).index(x) if x in STATES else 99)
        rows.append([sc.id, sc.title, ux.section_of(sc.id), ", ".join(sc.get("frs") or []) or "-", str(len(imgs)),
                     ", ".join(STATES.get(x, x) for x in states)])
    for n, role in enumerate(p.nav_roles(), start=1):
        rows.append([p.nav_id(n), f"Landing page and menu: {p.persona_label(role)}", ux.section_of("Menu by persona"),
                     "-", str(len([i for i in ux.images if i.screen == p.nav_id(n)])), "Landing page, Menu"])
    for d in p.documents:
        rows.append([d["id"], d["name"], ux.section_of(d["id"]), ", ".join(d.get("frs") or []) or "-", "1",
                     "Document as received"])
    deck.flow_table("Screen ID, FRS section and requirements", ["Screen ID", "Screen", f"{doc02} section",
                                                                "Requirements", "Images", "States"],
                    rows, [1.3, 3.3, 1.1, 3.0, 0.8, 3.5], size=8.5,
                    notes="The FRS section is the section of the screen specification in the issued FRS; a change "
                          "reference that names a section names these screens.")
    out = brand.out_dir(ux.brd, "UXDeck") / brand.output_name("UXDeck", ux.brd, ux.module, ux.version, "pptx")
    path = deck.save(out)
    return path, dict(deck.counts)


# ============================================================================ register and package


REGISTER_COLUMNS = [
    ("Screen ID", "Y", "SCR-, NAV- or DO- and a number", "The screen of the image, as in the FRS and the deck",
     "SCR-PM-03", ""),
    ("Screen name", "Y", "Text", "The title of the screen", "New Package Request and Edit Package Request", ""),
    ("Persona(s)", "Y", "Text", "The personas who see the screen in this state", "Marketing Account Officer", ""),
    ("Flow", "N", "Flow ID and title", "The end-to-end flow that shows the image, if any", "-", ""),
    ("Step", "N", "Whole number", "The step of the flow", "-", "whole:1-99"),
    ("State", "Y", "From the list", "The state of the screen on the image", "Filled", "list:UX_STATE"),
    ("FRS section", "Y", "Section number", "The section of the FRS (or handbook) that specifies the screen", "13.1.3",
     ""),
    ("Requirement IDs", "N", "FR-... separated by commas", "The functional requirements of the screen", "FR-PM-020",
     ""),
    ("Image file", "Y", "<Screen ID>_<state>.png", "The file of the image package (09)", "SCR-PM-03_filled.png", ""),
    ("What the image shows", "N", "Text", "The caption of the image", "The request form with the terms entered", ""),
    ("UXD status", "Y", "From the list", "The status of the design of this screen state, kept by the UX Design team",
     "Not started", "list:UXD_STATUS"),
    ("UXD comments", "N", "Text", "Comments of the UX Design team on the screen state", "-", ""),
    ("Changed in this issue", "Y", "Y or N", "Y when the image is new or changed in this issue of the set", "N", "yn"),
    ("Change reference", "Cond.: when Changed in this issue is Y", "Reference",
     "The change reference (comment, change request or issue change) that changed the image", "-", ""),
]


def register_rows(ux: Ux) -> list[list[Any]]:
    flows = {w["id"]: w for w in ux.flows}
    rows = []
    for img in ux.images:
        flow = "; ".join(f"{fid} {flows[fid]['title']}" for fid, _ in img.flows) if img.flows else ""
        step = img.flows[0][1] if len(img.flows) == 1 else ("; ".join(str(k) for _, k in img.flows) or None)
        rows.append([img.screen, img.title, ", ".join(img.personas) or "-", flow or "-", step if step else None,
                     img.state_label, img.section, ", ".join(img.frs) or "-", img.name, img.caption, "Not started",
                     None, "Y" if img.changed else "N", ", ".join(img.changed) or None])
    return rows


def build_register(ux: Ux) -> Path:
    title = f"UX screen register {ux.label} {ux.module}"
    book = g.GuidedBook(title, ux.version)
    book.add_list("UX_STATE", "State of a screen image", [(v, "") for v in STATES.values()])
    book.add_list("UXD_STATUS", "UXD status", [(v, "") for v in UXD_STATUSES])
    rows = register_rows(ux)
    # The example row shows the first image of the set.
    first = rows[0] if rows else [None] * len(REGISTER_COLUMNS)
    cols = [g.GuideColumn(h, mand, fmt, what, "-" if v in (None, "", "-") else str(v), check)
            for (h, mand, fmt, what, _, check), v in zip(REGISTER_COLUMNS, first)]
    doc02 = "Data Migration Handbook" if ux.frs_name() == "Handbook" else "FRS"
    t = g.Template("UXR", "UX screen register", "UX screen register",
                   f"One row per screen image of {ux.label} {ux.module} (the 07 deck and the 09 image package): the "
                   f"screen, personas, flow and step, state, {doc02} section and requirements, the image file, and the "
                   "UXD status and comments of the UX Design team. Changed in this issue marks the images that are "
                   "new or changed in this issue, with the change reference.",
                   "BDOI UX Design team (UXD status and comments); iorta TechNXT project team (the other columns)",
                   "With each design review; the rows are re-issued with each version of the set",
                   "Kept by the UX Design team; at a new issue the UXD status and comments of the unchanged rows are "
                   "carried over and the changed rows are flagged",
                   cols, source=f"{doc02} v{ux.version} and the screens of the SIT environment (seed data)",
                   capacity=len(rows) + 60,
                   confirm=("Reviewed by (UX Design team lead)", "Date"))
    book.plan_all([t])
    ws = book.template_sheet(t, 1, 1)
    pl = book.plans["UXR"]
    for r, row in enumerate(rows, start=pl.example_row + 1):
        for c, v in enumerate(row, start=2):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = g.font(10)
            cell.alignment = g.Alignment(wrap_text=True, vertical="top")
    changed = sum(1 for r in rows if r[12] == "Y")
    idx = [{"Step": 1, "ID": "UXR", "Template": "UX screen register", "What it holds": f"{len(rows)} screen images "
            f"of {len({r[0] for r in rows})} screens; {changed} changed in this issue",
            "Kept by": "BDOI UX Design team", "Rows entered": pl.rows_formula(),
            "Mandatory cells missing": pl.missing_formula(), "Status": "Not started", "Open": ("Open →", pl.sheet)}]
    steps = [
        "Open the sheet UX screen register: one row per screen image of the 07 deck and the 09 image package, "
        "with the screen, the personas, the flow and step, the state and the FRS section.",
        "Filter Changed in this issue = Y to see the images that are new or changed in this issue; the change "
        "reference names the comment or change request behind the change.",
        "Set the UXD status of each row (Not started, In design, Ready for review, Approved) and write the UXD "
        "comments; leave the other columns as issued.",
        "Import the CSV of the image package into the design tool to link each design to its screen image.",
        "Raise questions on the sheet Questions and comments; the project team answers there. At the next issue "
        "the UXD status and comments of the unchanged rows are carried over.",
    ]
    identity = [("Client", brand.CLIENT), ("Release set", f"{ux.label} {ux.module} ({brand.drop_of(ux.brd)})"),
                ("Version and date", f"Version {ux.version}, {ux.meta['date']}"),
                ("For", "BDOI UX Design team (UXD)"), ("Prepared by", brand.VENDOR),
                ("Classification", brand.CLASSIFICATION)]
    columns = [("Step", 6), ("ID", 8), ("Template", 24), ("What it holds", 44), ("Kept by", 26), ("Rows entered", 10),
               ("Mandatory cells missing", 11), ("Status", 14), ("Open", 8)]
    issue = ux.ux.get("issue") or {}

    def more(ws_, row):
        row = book.small_table(ws_, row, "Changes in this issue", ["Change reference", "What changed", "Screens"],
                               [[c["ref"], _words(c["what"]), ", ".join(sorted({i.screen for i in ux.images
                                                                               if c["ref"] in i.changed}))]
                                for c in issue.get("changes") or []] or [["-", "First issue", "-"]],
                               spans=[2, 5, 2])
        return book.small_table(ws_, row, "Other sheets", ["Sheet", "What it holds"], [
            ((g.LISTS, g.LISTS), "The lists of the drop-downs: states and UXD statuses"),
            ((g.QUESTIONS, g.QUESTIONS), "Questions and comments of the UX Design team, with the answer"),
        ], spans=[3, 6])

    book.questions_sheet([t])
    book.start_sheet(title, "The screen images of the set for the BDOI UX Design team", identity, steps,
                     "The register", columns, idx, after=more)
    book.lists_sheet()
    path = brand.out_dir(ux.brd, "UXScreens") / brand.output_name("UXScreens", ux.brd, ux.module, ux.version, "xlsx")
    return book.save(path, [g.START, t.sheet, g.LISTS, g.QUESTIONS],
                     {"title": title, "keywords": f"UX screen register, {ux.brd}, BIBS, BDOI"})


def build_package(ux: Ux) -> Path:
    path = brand.out_dir(ux.brd, "UXScreens") / brand.output_name("UXScreens", ux.brd, ux.module, ux.version, "zip")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in REGISTER_COLUMNS])
    for row in register_rows(ux):
        w.writerow(["" if v is None else v for v in row])
    folder = f"{ux.brd}_{ux.module.replace(' ', '_')}_UX_screens"
    path.parent.mkdir(parents=True, exist_ok=True)
    fixed = (2026, 9, 28, 0, 0, 0)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as z:
        info = zipfile.ZipInfo(f"{folder}/UX_screen_register_{ux.brd}.csv", fixed)
        info.compress_type = zipfile.ZIP_DEFLATED
        z.writestr(info, "﻿" + buf.getvalue())
        for img in ux.images:
            if img.file.exists():
                info = zipfile.ZipInfo(f"{folder}/images/{img.name}", fixed)
                z.writestr(info, img.file.read_bytes())
        for comp in ux.components:
            f = ux.component_file(comp)
            if f.exists():
                info = zipfile.ZipInfo(f"{folder}/components/{f.name}", fixed)
                z.writestr(info, f.read_bytes())
    return path


# ============================================================================ checks


def check(ux: Ux) -> list[str]:
    problems = list(ux.problems)
    for comp in ux.components:
        if not ux.component_file(comp).exists():
            problems.append(f"component {comp['name']}: image missing {ux.component_file(comp).relative_to(ROOT)}")
    for img in ux.images:
        if img.file.exists():
            w, h = _png_size(img.file)
            if w < 400:
                problems.append(f"{img.name}: only {w} pixels wide")
    names = [i.name for i in ux.images]
    if len(names) != len(set(names)):
        problems.append("duplicate image names in the package")
    for sc in ux.pack.screens:
        if not ux.screen_images(sc.id):
            problems.append(f"{sc.id}: no image")
        if ux.section_of(sc.id) == "-":
            problems.append(f"{sc.id}: no section in the issued FRS")
    xlsx = brand.out_dir(ux.brd, "UXScreens") / brand.output_name("UXScreens", ux.brd, ux.module, ux.version, "xlsx")
    if xlsx.exists():
        problems += g.verify(xlsx, ["UX screen register"])
    return problems


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build the UX screen deck, register and image package of a set")
    ap.add_argument("brd", help="BRD of the set, e.g. BRD-03")
    ap.add_argument("--check", action="store_true", help="check the images and the files written only")
    args = ap.parse_args(argv)
    ux = Ux(args.brd)
    states = {}
    for i in ux.images:
        states[i.state] = states.get(i.state, 0) + 1
    print(f"{ux.brd}: {len(ux.pack.screens)} screens, {len(ux.images)} images "
          f"({', '.join(f'{v} {STATES.get(k, k).lower()}' for k, v in states.items())}), {len(ux.flows)} flows, "
          f"{len(ux.pack.nav_roles())} personas, {sum(1 for i in ux.images if i.changed)} changed")
    if args.check:
        problems = check(ux)
        for p in problems:
            print(p)
        return 1 if problems else 0
    deck, counts = build_deck(ux)
    print(f"pptx: {deck}")
    print("slides: " + ", ".join(f"{k} {v}" for k, v in counts.items()) + f"; total {sum(counts.values())}")
    print(f"xlsx: {build_register(ux)}")
    print(f"zip: {build_package(ux)}")
    problems = check(ux)
    for p in problems:
        print(p)
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
