"""Business sign-off pack of a BRD: screen specifications, field register, messages, notifications, menus by
persona, cross-BRD contract, walkthroughs and the sign-off workbook (client deliverable: one release set per BRD).

Usage
  python docs/deliverables/src/signoff/signoff_pack.py docs/deliverables/src/signoff/brd01/pack.yaml
  python docs/deliverables/src/signoff/signoff_pack.py brd01/pack.yaml --check        # checks only
  python docs/deliverables/src/signoff/signoff_pack.py brd01/pack.yaml --manifest     # screenshot manifest (JSON)

What it reads
  * <brd>/pack.yaml             metadata, personas, sections of the BRD, screen-flow links, common screen elements;
  * <brd>/screens/*.yaml        one entry per screen: route, purpose, entry points, screenshots, field rows,
                                action rows, rules, expected outcome, FRs and test-plan screen aliases;
  * <brd>/messages.yaml         where each message appears and the fix the user makes (texts come from the code);
  * <brd>/notifications.yaml, contract.yaml, documents.yaml, walkthroughs.yaml;
  * the code itself through tools/deliverables/code_facts.py: sidebar menus, role grants, messages, upload templates,
    so the personas, menu paths and message texts are never typed by hand;
  * the test plan YAML of the BRD (docs/deliverables/src/testplans), for the test cases of each screen.

What it writes
  * BIBS_Signoff_BRD-nn_<Name>_v<version>.xlsx in the BRD's release folder (brand.out_dir): How to review, Screen
    catalogue, Field register, Actions, Business rules, Messages, Notifications, Menu by persona, Upload templates,
    Cross-BRD contract, Sign-off, with the BU review columns (Accept / Change requested / Comment, Comment,
    Reviewer, Date) on the screen, field, rule and message sheets;
  * the chapters of the FRS through ```pack blocks: render(doc, render=<part>, source=<pack.yaml>).

Row formats (screens/*.yaml; cells separated by " | ")
  fields:  section | label | type | length / format | mandatory | source / master | default | editable in |
           validation rule | message shown when it fails
  actions: button | who | enabled when | what happens | resulting status | notification sent
"""

from __future__ import annotations

import argparse
import codecs
import json
import re
import sys
from dataclasses import dataclass, field
from functools import cached_property
from pathlib import Path
from typing import Any

import yaml

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "tools" / "deliverables"))
import brand  # noqa: E402
import code_facts  # noqa: E402

FIELD_COLS = ["section", "label", "type", "format", "mandatory", "source", "default", "editable", "validation",
              "message"]
ACTION_COLS = ["button", "who", "when", "what", "status", "notification"]
REVIEW_VALUES = ["Accept", "Change requested", "Comment"]
SIGNOFF_VALUES = ["Approved", "Approved with comments", "Not approved"]
IMPERATIVE = re.compile(r"^(Enter|Select|Choose|Give|Attach|Upload|Add|Complete|Describe|Record|Compute|Use|Explain|"
                        r"Name|Keep|Confirm|Close|Correct|Check|Save|Submit|Write|Reduce|Leave|Ask)\b")
BANNED = re.compile(r"\b(" + codecs.decode("qrzb|qhzzl|snxr|fnzcyr qngn|cebgbglcr|cbp|fnaqobk|yberz vcfhz|gbqb|svkzr", "rot13") + r")\b", re.I)


def split_row(text: str, cols: list[str], where: str, problems: list[str]) -> dict[str, str]:
    parts = [p.strip() for p in str(text).split(" | ")]
    if len(parts) != len(cols):
        problems.append(f"{where}: {len(parts)} cells, expected {len(cols)}: {str(text)[:80]}")
        parts = (parts + ["-"] * len(cols))[: len(cols)]
    return dict(zip(cols, parts))


@dataclass
class Screen:
    id: str
    title: str
    area: str
    route: str
    raw: dict[str, Any]
    fields: list[dict[str, str]] = field(default_factory=list)
    actions: list[dict[str, str]] = field(default_factory=list)

    def get(self, key: str, default: Any = None) -> Any:
        return self.raw.get(key, default)


class Pack:
    def __init__(self, path: str | Path):
        self.path = Path(path).resolve()
        self.dir = self.path.parent
        raw = yaml.safe_load(self.path.read_text(encoding="utf-8"))
        self.meta: dict[str, Any] = raw["meta"]
        self.personas: dict[str, dict[str, Any]] = raw.get("personas") or {}
        self.own_sections: list[str] = raw.get("own_sections") or []
        self.section_brd: dict[str, str] = raw.get("section_brd") or {}
        self.flow: list[list[str]] = raw.get("flow") or []
        self.common: list[dict[str, str]] = raw.get("common") or []
        self.problems: list[str] = []
        self.areas: list[str] = []
        self.screens: list[Screen] = []
        for f in sorted((self.dir / self.meta.get("screens_dir", "screens")).glob("*.yaml")):
            data = yaml.safe_load(f.read_text(encoding="utf-8"))
            self.areas.append(data["area"])
            for s in data["screens"]:
                scr = Screen(s["id"], s["title"], data["area"], s["route"], s)
                for i, row in enumerate(s.get("fields") or [], start=1):
                    d = split_row(row, FIELD_COLS, f"{scr.id} field {i}", self.problems)
                    d["no"] = str(i)
                    scr.fields.append(_described(d, "label"))
                for i, row in enumerate(s.get("actions") or [], start=1):
                    scr.actions.append(_described(split_row(row, ACTION_COLS, f"{scr.id} action {i}", self.problems),
                                                  "button"))
                self.screens.append(scr)
        self.by_id = {s.id: s for s in self.screens}
        self.notifications = self._load("notifications.yaml").get("notifications", [])
        contract = self._load("contract.yaml")
        self.contract = contract.get("contract", [])
        self.change_rule = contract.get("change_rule", "")
        self.documents = self._load("documents.yaml").get("documents", [])
        self.walkthroughs = self._load("walkthroughs.yaml").get("walkthroughs", [])
        self.msg_cfg = self._load("messages.yaml")

    def _load(self, name: str) -> dict[str, Any]:
        p = self.dir / name
        return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else {}

    # ------------------------------------------------------------------ code facts

    @cached_property
    def groups(self) -> list[code_facts.Group]:
        return code_facts.frontend_menu()

    @cached_property
    def grants(self) -> dict[str, set[str]]:
        return code_facts.role_grants()

    @cached_property
    def routes(self) -> dict[str, tuple[code_facts.Group, code_facts.Section, code_facts.Screen]]:
        return {sc.path: (g, s, sc) for g, s, sc in code_facts.all_screens(self.groups)}

    def screen_def(self, route: str) -> code_facts.Screen | None:
        hit = self.routes.get(route)
        return hit[2] if hit else None

    def personas_of(self, scr: Screen) -> list[str]:
        sd = self.screen_def(scr.route)
        if sd is None:
            return []
        return [r for r in self.personas if sd.may_open(self.grants.get(r, set()))]

    def permission_text(self, scr: Screen) -> str:
        sd = self.screen_def(scr.route)
        if scr.get("public"):
            return "none, before sign-in"
        if sd is None:
            return "-"
        if not sd.opened_by:
            return "Every signed-in user"
        return " or ".join(sd.opened_by)

    def menu_path(self, scr: Screen) -> str:
        hit = self.routes.get(scr.route)
        if hit is None:
            return scr.route
        g, s, sd = hit
        if not sd.hidden:
            return code_facts.menu_path(g, s, sd)
        parent = self.routes.get(scr.get("parent", ""))
        base = code_facts.menu_path(*parent) if parent else ""
        return f"{base} › {sd.label}" if base else sd.label

    def persona_label(self, role: str) -> str:
        p = self.personas.get(role, {})
        return p.get("name", role)

    # ------------------------------------------------------------------ menus

    def menus(self) -> dict[str, list[dict[str, str]]]:
        out: dict[str, list[dict[str, str]]] = {}
        for role in self.personas:
            rows = []
            for g in code_facts.menu_for(self.grants.get(role, set()), self.groups):
                for s in g.sections:
                    for sc in s.screens:
                        rows.append({"role": role, "persona": self.persona_label(role),
                                     "group": g.title or "Home", "section": s.title, "screen": sc.label,
                                     "path": sc.path, "brd": self.section_brd.get(s.id, "-"),
                                     "own": "Yes" if s.id in self.own_sections else ""})
            out[role] = rows
        return out

    # ------------------------------------------------------------------ tests

    @cached_property
    def test_plan(self) -> Any | None:
        tp = self.meta.get("test_plan")
        if not tp:
            return None
        sys.path.insert(0, str(REPO / "docs" / "deliverables" / "src" / "testplans"))
        import build_test_plan  # noqa: E402

        return build_test_plan.load(REPO / "docs" / "deliverables" / "src" / "testplans" / tp)

    def tests_of(self, scr: Screen) -> list[str]:
        plan = self.test_plan
        if plan is None:
            return []
        paths = {plan.screens.get(a, a) for a in scr.get("tests") or []}
        return [c.id for c in plan.cases if getattr(c, "screen_id", "") == scr.id or c.screen in paths]

    def screen_of_alias(self) -> dict[str, str]:
        return {a: s.id for s in self.screens for a in s.get("tests") or []}

    # ------------------------------------------------------------------ rules

    def rules(self) -> list[dict[str, str]]:
        out = []
        n = 0
        for s in self.screens:
            for r in s.get("rules") or []:
                n += 1
                frs = ", ".join(re.findall(r"FR-[A-Z]+-\d+", r))
                text = re.sub(r"\s*\((FR-[A-Z]+-\d+(, )?)+\)\.?$", ".", r).strip()
                out.append({"id": f"BR-{self.meta['code']}-{n:03d}", "screen": s.id, "title": s.title,
                            "rule": text, "frs": frs})
        return out

    # ------------------------------------------------------------------ messages

    def _place(self, name: str) -> str:
        for pattern, where in self.msg_cfg.get("places") or []:
            if re.search(pattern, name):
                return where
        return "Several screens"

    def _fix(self, code: str, text: str, kind: str) -> str:
        fixes = self.msg_cfg.get("fixes") or {}
        if code in fixes:
            return fixes[code]
        if text in fixes:
            return fixes[text]
        if IMPERATIVE.match(text):
            return text.rstrip(".") + "."
        if kind in ("Confirmation", "Information"):
            return "No action."
        return "-"

    @cached_property
    def messages(self) -> list[dict[str, str]]:
        cfg = self.msg_cfg.get("sources") or {}
        exclude = set(cfg.get("exclude_classes") or [])
        rows: list[dict[str, str]] = []
        seen: set[tuple[str, str, str]] = set()

        def add(layer: str, code: str, text: str, kind: str, where: str, source: str) -> None:
            key = (code, text, where)
            if key in seen or not re.sub(r"<[^>]*>", "", text).strip(" .:,"):
                return
            seen.add(key)
            rows.append({"layer": layer, "code": code, "text": text, "kind": kind, "where": where,
                         "fix": self._fix(code, text, kind), "source": source})
        kinds = {"Business rule": "Error", "Field check": "Validation", "Not found": "Error", "Duplicate": "Error",
                 "Warning": "Warning", "Request check": "Validation"}
        backend = code_facts.backend_messages(cfg.get("backend_packages") or [])
        extra = set(cfg.get("backend_classes_extra") or [])
        if extra:
            pkgs = {p.parent.name for p in code_facts.JAVA_ROOT.rglob("*.java") if p.stem in extra}
            backend += [m for m in code_facts.backend_messages(sorted(pkgs)) if m.cls in extra]
        for m in backend:
            if m.cls in exclude or not re.fullmatch(r"[A-Z][A-Z0-9_]*", m.code):
                continue
            add("Server", m.code, m.text, kinds.get(m.kind, "Error"), self._place(m.cls), m.where)
        for t in code_facts.bulk_templates(cfg.get("bulk_handlers") or []):
            for text in t.row_messages:
                add("Upload row", "-", text, "Validation", f"Bulk upload wizard (SCR-NB-36): {t.title}",
                    Path(t.file).stem)
        for text in ("<Column> is mandatory", "<Column> '<value>' is not a valid <number, date (yyyy-mm-dd), Y/N value>"):
            add("Upload row", "-", text, "Validation", "Bulk upload wizard (SCR-NB-36): every upload type",
                "BulkRowValidator")
        for m in code_facts.platform_messages():
            add("Server", m.code, m.text, "Error", "Every screen (Common screen elements)", m.where)
        root = REPO / "frontend" / "src"
        for u in code_facts.frontend_messages(cfg.get("frontend_dirs") or []):
            add("Screen", "-", u.text, u.kind, self._place(Path(u.file).stem.split(".")[0]),
                str(Path(u.file).relative_to("frontend/src")) if u.file.startswith("frontend/src") else u.file)
        order = {w: i for i, (_, w) in enumerate(self.msg_cfg.get("places") or [])}
        rows.sort(key=lambda r: (order.get(r["where"], 999), r["layer"] != "Screen", r["code"], r["text"]))
        for i, r in enumerate(rows, start=1):
            r["id"] = f"MSG-{self.meta['code']}-{i:03d}"
        del root
        return rows

    # ------------------------------------------------------------------ uploads

    @cached_property
    def uploads(self) -> list[code_facts.BulkTemplate]:
        return code_facts.bulk_templates((self.msg_cfg.get("sources") or {}).get("bulk_handlers") or [])

    # ------------------------------------------------------------------ screenshots

    def shots(self) -> list[dict[str, Any]]:
        """Every screenshot of the pack (screens, walkthroughs, documents), with its file name."""
        out = []
        for s in self.screens:
            for i, shot in enumerate(s.get("shots") or [], start=1):
                slug = f"{s.id.lower()}-{i:02d}-{shot.get('state', 'view')}"
                out.append({"slug": slug, "screen": s.id, "route": s.route, "caption": shot.get("caption", s.title),
                            "callouts": [[f["no"], f["label"], f["type"]] for f in s.fields] if i == 1 else [],
                            **{k: v for k, v in shot.items() if k != "caption"}})
        for w in self.walkthroughs:
            for step in w["steps"]:
                role, scr, action, sees, result, slug = step
                out.append({"slug": slug, "screen": scr, "route": self.by_id[scr].route if scr in self.by_id else "",
                            "caption": f"{w['id']}: {action}", "user": self.personas.get(role, {}).get("user"),
                            "state": "walkthrough", "walkthrough": w["id"]})
        for d in self.documents:
            out.append({"slug": d["shot"], "screen": "-", "route": "", "caption": d["name"], "state": "document"})
        return out

    def shot_file(self, slug: str) -> Path:
        return self.dir / self.meta.get("screenshot_dir", "screenshots") / f"{slug}.png"

    # ------------------------------------------------------------------ checks

    @cached_property
    def _frontend_text(self) -> str:
        parts = []
        for p in (REPO / "frontend" / "src").rglob("*.ts*"):
            if not re.search(r"\.(test|spec)\.tsx?$", p.name):
                parts.append(p.read_text(encoding="utf-8", errors="ignore"))
        for wf in code_facts.workflows().values():
            parts += [str(x.get("name") or "") for x in wf["stages"]]
            parts += [str(x.get("label") or "") for x in wf["transitions"]]
            parts += [str(x.get("action") or "").replace("_", " ") for x in wf["transitions"]]
        return re.sub(r"\s+", " ", "\n".join(parts)).lower()

    def on_screen(self, label: str) -> bool:
        """True when the label is a text of the frontend (or a workflow stage or action of the database).
        A trailing descriptive noun (tag, banner, tiles...) is ignored; n and numbers match any value."""
        label = re.sub(r"\s+(tags?|banner|notice|tiles?|chips?|check boxes|reason|list|names)$", "", label.lower())
        pattern = r"\s*".join(".{0,30}?" if w == "n" or w.isdigit() else re.escape(w) for w in label.split())
        return re.search(pattern, self._frontend_text) is not None

    def check(self) -> list[str]:
        problems = list(self.problems)
        frs_text = (REPO / "docs" / "deliverables" / "src" / "frs" / self.meta["frs"]).read_text(encoding="utf-8")
        fr_ids = set(re.findall(r"^id: (FR-[A-Z]+-\d+)", frs_text, re.M))
        ids = [s.id for s in self.screens]
        if len(ids) != len(set(ids)):
            problems.append("duplicate screen ids")
        aliases = set(self.test_plan.screens) if self.test_plan else set()
        for s in self.screens:
            for r in [s.route] + list(s.get("also") or []):
                if r not in self.routes and not s.get("public"):
                    problems.append(f"{s.id}: route {r} is not a screen of navigation/modules.ts")
            if not self.personas_of(s) and not s.get("public"):
                problems.append(f"{s.id}: no persona of the pack may open {s.route}")
            for fr in s.get("frs") or []:
                if fr not in fr_ids:
                    problems.append(f"{s.id}: {fr} is not an FR of {self.meta['frs']}")
            for a in s.get("tests") or []:
                if aliases and a not in aliases:
                    problems.append(f"{s.id}: test screen alias {a} is not in the test plan")
            for f in s.fields:
                for label in [] if f.get("_described") else _labels(f["label"]):
                    if not self.on_screen(label):
                        problems.append(f"{s.id} field {f['no']}: label '{label}' not found in frontend/src")
            for a in s.actions:
                for label in [] if a.get("_described") else _labels(a["button"]):
                    if not self.on_screen(label):
                        problems.append(f"{s.id} action: button '{label}' not found in frontend/src")
            for text in self._texts(s):
                if BANNED.search(text):
                    problems.append(f"{s.id}: banned word in '{text[:60]}'")
        for m in self.messages:
            if m["fix"] == "-" and m["kind"] in ("Error", "Validation", "Warning"):
                problems.append(f"message without a fix: {m['code']} {m['text'][:70]} ({m['source']})")
        for a, b, _ in self.flow:
            for x in (a, b):
                if x not in self.by_id:
                    problems.append(f"flow: unknown screen {x}")
        for w in self.walkthroughs:
            for st in w["steps"]:
                if st[0] not in self.personas:
                    problems.append(f"{w['id']}: persona {st[0]}")
                if st[1] not in self.by_id:
                    problems.append(f"{w['id']}: screen {st[1]}")
        for u in getattr(code_facts.role_grants, "unread", []):
            problems.append(f"grant statement not read: {u}")
        return problems

    @staticmethod
    def _texts(s: Screen) -> list[str]:
        texts = [s.title, str(s.get("purpose", "")), str(s.get("outcome", ""))]
        texts += [str(r) for r in s.get("rules") or []]
        texts += [" ".join(v for k, v in f.items() if not k.startswith("_")) for f in s.fields] + [" ".join(a.values()) for a in s.actions]
        return texts


def _described(row: dict[str, str], key: str) -> dict[str, str]:
    """A label in [square brackets] describes screen content (tiles, tags...) rather than quoting a screen text."""
    if row[key].startswith("[") and row[key].endswith("]"):
        row[key] = row[key][1:-1]
        row["_described"] = "yes"
    return row


def _labels(label: str) -> list[str]:
    """The screen texts named in a field label or button cell: 'A / B', 'A, B', 'Vehicle: A, B' and (row) notes."""
    label = re.sub(r"\([^)]*\)|\[[^\]]*\]", "", label)
    if ":" in label:
        label = label.split(":", 1)[1]
    parts = re.split(r"\s+/\s+|,\s+", label)
    out = []
    for p in parts:
        p = p.strip().rstrip(".")
        p = re.sub(r"\s+n$|\s+\(n\)$", "", p)
        if not p or p in ("-",) or re.search(r"<|\.\.\.|…", p) or len(p) < 3:
            continue
        out.append(p)
    return out


# ============================================================================ Word rendering (```pack blocks)


def render(doc: Any, render: str, source: str, **opts: Any) -> None:  # noqa: A002 - block key
    """Entry point of the ```pack blocks of the FRS source (tools/deliverables/bdoi_docx.py render_plugin)."""
    pack = _pack((doc.base_dir / source).resolve())
    RENDERERS[render](doc, pack, **opts)


_PACKS: dict[Path, Pack] = {}


def _pack(path: Path) -> Pack:
    if path not in _PACKS:
        _PACKS[path] = Pack(path)
    return _PACKS[path]


def r_screen_index(doc: Any, pack: Pack, **_: Any) -> None:
    rows = []
    for s in pack.screens:
        rows.append([s.id, s.title, pack.menu_path(s), ", ".join(pack.personas_of(s))])
    doc.table(["ID", "Screen", "Menu path", "Personas (role codes)"], rows, widths=[1.9, 3.6, 6.4, 5.7],
              caption="New Business screens", size=7.5)


def r_menus(doc: Any, pack: Pack, **_: Any) -> None:
    menus = pack.menus()
    for role, rows in menus.items():
        p = pack.personas[role]
        doc.heading(f"{p['name']} ({role})", level=3)
        own = sum(1 for r in rows if r["own"])
        doc.paragraph(f"SIT/UAT user {p.get('user', '-')}. {len(rows)} menu entries, {own} of them New Business "
                      "screens; the other entries belong to the BRD shown.")
        grouped: dict[tuple[str, str], list[dict[str, str]]] = {}
        for r in rows:
            grouped.setdefault((r["group"], r["section"]), []).append(r)
        table = [[g, s, ", ".join(x["screen"] for x in items), items[0]["brd"]] for (g, s), items in grouped.items()]
        doc.table(["Group", "Section", "Screens", "BRD"], table, widths=[3.0, 3.8, 8.6, 2.2],
                  caption=f"Sidebar of the {p['name']}", size=7.5, first_col_bold=True)


def flow_dot(pack: Pack) -> Path:
    """Writes the screen-flow figure source (Graphviz) next to the FRS figures and returns its path."""
    lines = ["digraph flow {", '  graph [rankdir=TB, fontname="Arial", nodesep=0.2, ranksep=0.35, bgcolor="white", '
             'newrank=true, compound=true];',
             '  node [shape=box, style="rounded,filled", fillcolor="@BG_BLUE", color="@HEADER_BLUE", fontname="Arial",'
             ' fontsize=10, fontcolor="@NEAR_BLACK", margin="0.12,0.06"];',
             '  edge [color="@CTA_BLUE", fontname="Arial", fontsize=8, fontcolor="@MUTED", arrowsize=0.6];']
    used = {x for a, b, _ in pack.flow for x in (a, b)}
    for area in pack.areas:
        members = [s for s in pack.screens if s.area == area and s.id in used]
        if not members:
            continue
        lines.append(f'  subgraph "cluster_{area}" {{ label="{area}"; fontname="Arial"; fontsize=10; '
                     'fontcolor="@HEADER_BLUE"; color="@BORDER"; style="rounded";')
        for s in members:
            lines.append(f'    "{s.id}" [label="{s.title}\\n{s.id}"];')
        lines.append("  }")
    for a, b, label in pack.flow:
        lines.append(f'  "{a}" -> "{b}" [label="{label}"];')
    lines.append("}")
    target = REPO / "docs" / "deliverables" / "src" / "frs" / "figures" / f"{pack.meta['brd'].lower().replace('-', '')}_screen_flow.dot"
    text = "\n".join(lines) + "\n"
    if not target.exists() or target.read_text(encoding="utf-8") != text:
        target.write_text(text, encoding="utf-8")
    return target


def r_flow(doc: Any, pack: Pack, **_: Any) -> None:
    doc.figure(flow_dot(pack), "Screen flow of New Business: how the screens link (list, record, action, next screen)")


def r_common(doc: Any, pack: Pack, **_: Any) -> None:
    doc.table(["Element", "What the user sees and does"], [[c["name"], c["text"]] for c in pack.common],
              widths=[4.0, 13.6], caption="Common screen elements", size=8.5, first_col_bold=True)


def r_screens(doc: Any, pack: Pack, **_: Any) -> None:
    for area in pack.areas:
        doc.heading(area, level=2)
        for s in [x for x in pack.screens if x.area == area]:
            _screen(doc, pack, s)


def _screen(doc: Any, pack: Pack, s: Screen) -> None:
    doc.heading(f"{s.id} {s.title}", level=3)
    personas = ", ".join(f"{pack.persona_label(r)}" for r in pack.personas_of(s)) or "Every user"
    entries = [f"Menu: {pack.menu_path(s)}"] + list(s.get("entry") or [])
    doc.key_values([("Purpose", " ".join(str(s.get("purpose", "")).split())),
                    ("Who can open it", f"{personas} (permission {pack.permission_text(s)})"),
                    ("Navigation", entries),
                    ("FRs", ", ".join(s.get("frs") or []))], columns=1, label_width=3.2, keep_together=False)
    shots = s.get("shots") or []
    for i, shot in enumerate(shots, start=1):
        slug = f"{s.id.lower()}-{i:02d}-{shot.get('state', 'view')}"
        legend = [(f["no"], f["label"]) for f in s.fields] if i == 1 else None
        doc.screenshot(pack.shot_file(slug), f"{s.id} {shot.get('caption', s.title)}", legend=legend)
    if s.fields:
        doc.label("Fields (No. = callout number on the first screenshot)")
        rows = [[f["no"], f["section"], f["label"], f["type"], f["format"], f["mandatory"], f["source"], f["default"],
                 f["editable"], f["validation"], f["message"]] for f in s.fields]
        doc.table(["No.", "Section", "Field", "Type", "Length / format", "Mand.", "Source / list", "Default",
                   "Editable in", "Validation", "Message when it fails"], rows,
                  widths=[0.7, 1.8, 2.6, 1.5, 2.4, 1.9, 2.2, 1.6, 2.0, 3.0, 3.3], size=7, keep_rows=False)
    if s.actions:
        doc.label("Actions")
        rows = [[a["button"], a["who"], a["when"], a["what"], a["status"], a["notification"]] for a in s.actions]
        doc.table(["Button", "Who", "Enabled when", "What happens", "Resulting status", "Notification"], rows,
                  widths=[2.8, 3.0, 3.6, 7.2, 2.6, 3.4], size=7.5, keep_rows=False)
    rules = s.get("rules") or []
    if rules:
        doc.label("Business rules")
        doc.bullets(rules, size=9)
    doc.label("Expected outcome")
    doc.paragraph(str(s.get("outcome", "")))
    tests = pack.tests_of(s)
    doc.paragraph(f"**Test cases:** {', '.join(tests) if tests else 'see the test plan screen cases'}.", size=8.5)


def r_messages(doc: Any, pack: Pack, **_: Any) -> None:
    by_where: dict[str, list[dict[str, str]]] = {}
    for m in pack.messages:
        by_where.setdefault(m["where"], []).append(m)
    counts = {}
    for m in pack.messages:
        counts[m["kind"]] = counts.get(m["kind"], 0) + 1
    doc.paragraph(f"{len(pack.messages)} messages: " + ", ".join(f"{v} {k.lower()}" for k, v in sorted(counts.items()))
                  + ". The text is quoted exactly as the system shows it; <name> marks a value filled in by the "
                    "system. Server messages also show their code as the Reference under the message.")
    for where, rows in by_where.items():
        doc.heading(where, level=2, toc=False)
        table = [[m["id"], m["code"], m["text"], m["kind"], m["fix"]] for m in rows]
        doc.table(["ID", "Code (Reference)", "Message", "Type", "What the user does"], table,
                  widths=[1.9, 3.8, 7.2, 1.7, 5.0], size=7, keep_rows=False)


def r_notifications(doc: Any, pack: Pack, **_: Any) -> None:
    rows = [[n["id"], n["channel"], n["trigger"], n["recipient"], n["template"], ", ".join(n.get("frs") or [])]
            for n in pack.notifications]
    doc.table(["ID", "Channel", "Trigger", "Recipient", "Template and content", "FR"], rows,
              widths=[1.2, 1.6, 5.2, 4.2, 9.0, 2.0], caption="Notifications and e-mails of New Business", size=7,
              keep_rows=False)


def r_documents(doc: Any, pack: Pack, **_: Any) -> None:
    for d in pack.documents:
        doc.heading(f"{d['id']} {d['name']}", level=2)
        doc.key_values([("Template", d["template"]), ("Format", d["format"]), ("Produced by", d["produced"]),
                        ("Password protected", d["protected"]), ("FRs", ", ".join(d.get("frs") or []))],
                       columns=1, label_width=3.2)
        doc.label("Fields and their source")
        doc.bullets(d["fields"], size=9)
        doc.screenshot(pack.shot_file(d["shot"]), f"{d['id']} {d['name']}: first page as generated")


def r_uploads(doc: Any, pack: Pack, **_: Any) -> None:
    for t in pack.uploads:
        doc.heading(f"{t.title} ({t.code})", level=2)
        doc.key_values([("Permission", f"BULK_PROCESS and {t.permission}"), ("Instructions", t.instructions or "-")],
                       columns=1, label_width=3.2)
        rows = [[i, c["header"], "Y" if c["required"] else "N", c["type"].replace("_", "/").title(), c["description"],
                 c["example"]] for i, c in enumerate(t.columns, start=1)]
        doc.table(["No.", "Column", "Mand.", "Type", "Content", "Example"], rows, widths=[0.8, 3.4, 1.1, 1.4, 7.6, 3.3],
                  caption=f"Template columns of {t.title}", size=7.5, keep_rows=False)
        if t.row_messages:
            doc.label("Row checks of this upload (besides the column checks)")
            doc.bullets(t.row_messages, size=8.5)


def r_contract(doc: Any, pack: Pack, **_: Any) -> None:
    rows = [[c["id"], c["party"], c["direction"], c["what"], c["how"], c["owner"], ", ".join(c.get("frs") or [])]
            for c in pack.contract]
    doc.table(["ID", "BRD or system", "Direction", "What", "When and how", "Owner of the data", "FR"], rows,
              widths=[1.1, 2.6, 1.4, 5.6, 6.2, 3.2, 2.6], caption="Interface contract of New Business", size=7,
              keep_rows=False)
    doc.callout(" ".join(pack.change_rule.split()), kind="decision", title="Change against this signed set")


def r_walkthrough(doc: Any, pack: Pack, id: str, **_: Any) -> None:  # noqa: A002
    w = next(x for x in pack.walkthroughs if x["id"] == id)
    doc.paragraph(" ".join(w["summary"].split()))
    doc.paragraph(f"*{w['data']}*", size=9)
    for n, (role, scr, action, sees, result, slug) in enumerate(w["steps"], start=1):
        screen = pack.by_id[scr]
        doc.label(f"Step {n}. {pack.persona_label(role)} – {screen.title} ({scr})")
        doc.key_values([("Does", action), ("Sees", sees), ("Result", result)], columns=1, label_width=2.2)
        doc.screenshot(pack.shot_file(slug), f"{w['id']} step {n}: {screen.title}", max_height_ratio=0.42)
    if w.get("downstream"):
        doc.label("What happens next, downstream")
        doc.bullets(w["downstream"], size=9)


def r_counts(doc: Any, pack: Pack, **_: Any) -> None:
    fields = sum(len(s.fields) for s in pack.screens)
    actions = sum(len(s.actions) for s in pack.screens)
    rows = [["Screens specified", len(pack.screens)], ["Fields and columns described", fields],
            ["Actions described", actions], ["Business rules on screens", len(pack.rules())],
            ["Messages (server, screen and upload row)", len(pack.messages)],
            ["Notifications and e-mails", len(pack.notifications)], ["Upload types", len(pack.uploads)],
            ["Generated documents", len(pack.documents)], ["Walkthroughs", len(pack.walkthroughs)],
            ["Interface contract lines", len(pack.contract)], ["Personas with their menus", len(pack.personas)]]
    doc.table(["Content of the set", "Count"], rows, widths=[10, 3], caption="The set in numbers", size=9)


RENDERERS = {
    "screen-index": r_screen_index,
    "menus": r_menus,
    "flow": r_flow,
    "common": r_common,
    "screens": r_screens,
    "messages": r_messages,
    "notifications": r_notifications,
    "documents": r_documents,
    "uploads": r_uploads,
    "contract": r_contract,
    "walkthrough": r_walkthrough,
    "counts": r_counts,
}


# ============================================================================ workbook


def review_columns() -> list[Any]:
    from bdoi_xlsx import Column

    return [
        Column("bu_decision", "BU review", 16, "Business unit decision on the row", values=REVIEW_VALUES),
        Column("bu_comment", "BU comment", 40, "What should change, or the comment"),
        Column("bu_reviewer", "Reviewer", 18, "Name of the BU reviewer"),
        Column("bu_date", "Review date", 13, "Date of the review (dd-mmm-yyyy)", kind="date"),
    ]


def _date_validation(ws: Any, column_letter: str, first: int = 5, last: int = 5000) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True,
                        showErrorMessage=True, errorTitle="Review date", error="Enter a date (dd-mmm-yyyy).")
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}{first}:{column_letter}{last}")


def build_workbook(pack: Pack) -> Path:
    from bdoi_xlsx import BdoiWorkbook, Column
    from openpyxl.utils import get_column_letter

    m = pack.meta
    wb = BdoiWorkbook(f"Sign-off Workbook {m['brd_label']} {m['name']}", doc_type="Business sign-off workbook",
                      brd=m["brd"], version=str(m["version"]), date=str(m["date"]),
                      subtitle=f"{m['release_set']} – screens, fields, rules, messages and menus for BU review")
    wb.legend = [("Accept", "The row is correct as specified"),
                 ("Change requested", "The row must change: describe the change in BU comment"),
                 ("Comment", "A remark that does not change the row")]
    wb.cover_notes = [
        f"Status as of {m['status_as_of']}. The rows are generated from the FRS v{m['version']} sources and the "
        "system as built; they match the FRS screen specifications row for row.",
        "Fill in the BU review columns only. After sign-off the screens, fields, navigation and messages are "
        "frozen; later changes go through the Change Management Register.",
    ]
    date_sheets = []

    howto = [
        ("1", "Read the FRS chapter Screen specifications with the screenshots, or walk through the screens on the SIT "
              "environment during the review sessions of the release note."),
        ("2", "On the sheets Screen catalogue, Field register, Business rules and Messages, set BU review to Accept, "
              "Change requested or Comment for each row you review, and write the change in BU comment."),
        ("3", "Put your name in Reviewer and the date in Review date. Leave the other columns unchanged."),
        ("4", "Menu by persona shows what each role sees in the sidebar; Cross-BRD contract lists what New Business "
              "takes from and hands to the other BRDs."),
        ("5", "Return the workbook to the iorta TechNXT project team by the date of the release note. Every "
              "Change requested row is answered in the sign-off tracker before sign-off."),
        ("6", "On the Sign-off sheet, each signatory records the decision and the date. Signing freezes the content, "
              "screens and navigation of this set; the screenshots use fictitious seed data."),
    ]
    wb.sheet("How to review", [Column("step", "Step", 8, "Step number"), Column("text", "How to review", 110, "Instruction")],
             [{"step": a, "text": b} for a, b in howto], description="How the business unit reviews this workbook",
             freeze_first_column=False)

    screen_rows = []
    for s in pack.screens:
        screen_rows.append({
            "id": s.id, "area": s.area, "title": s.title, "menu": pack.menu_path(s),
            "personas": ", ".join(pack.persona_label(r) for r in pack.personas_of(s)),
            "permission": pack.permission_text(s), "purpose": " ".join(str(s.get("purpose", "")).split()),
            "entry": "\n".join(s.get("entry") or []), "outcome": s.get("outcome", ""),
            "frs": ", ".join(s.get("frs") or []), "tests": ", ".join(pack.tests_of(s)),
            "shots": "\n".join(x.get("caption", "") for x in s.get("shots") or []),
            "fields": len(s.fields), "actions": len(s.actions)})
    ws = wb.sheet("Screen catalogue", [
        Column("id", "Screen ID", 12, "Screen identifier used in the FRS"),
        Column("area", "Area", 22, "Process area"),
        Column("title", "Screen", 26, "Screen name as in the menu or page title"),
        Column("menu", "Menu path", 38, "Where the screen is in the sidebar (› separates the levels)"),
        Column("personas", "Personas", 36, "Roles that can open the screen (from the role grants)"),
        Column("permission", "Permission", 24, "Permission(s) that open the screen"),
        Column("purpose", "Purpose", 50, "What the screen is for"),
        Column("entry", "Other entry points", 36, "How the user reaches it from other screens"),
        Column("outcome", "Expected outcome", 40, "What is achieved on the screen"),
        Column("frs", "FRs", 20, "Functional requirements of the FRS"),
        Column("tests", "Test cases", 30, "Test cases of the test plan on this screen"),
        Column("shots", "Screenshots", 36, "Screenshots of the screen in the FRS"),
        Column("fields", "Fields", 8, "Number of fields and columns", kind="number"),
        Column("actions", "Actions", 8, "Number of actions", kind="number"),
    ] + review_columns(), screen_rows, description="One row per New Business screen")
    date_sheets.append(ws)

    field_rows = []
    for s in pack.screens:
        for f in s.fields:
            field_rows.append({"id": s.id, "screen": s.title, **f})
    ws = wb.sheet("Field register", [
        Column("id", "Screen ID", 12, "Screen identifier"),
        Column("screen", "Screen", 22, "Screen name"),
        Column("no", "No.", 6, "Field number = callout number on the screenshot"),
        Column("section", "Section", 20, "Card, tab, step or dialog of the field"),
        Column("label", "Field", 28, "Label as shown on screen"),
        Column("type", "Type", 14, "Text, Date, Drop-down, Check box, Column, ..."),
        Column("format", "Length / format", 28, "Length, format, allowed values"),
        Column("mandatory", "Mandatory", 16, "Y, N or the condition that makes it mandatory"),
        Column("source", "Source / list", 22, "List of values, master or record the value comes from"),
        Column("default", "Default", 16, "Value proposed by the system"),
        Column("editable", "Editable in", 22, "Statuses in which the field can be changed"),
        Column("validation", "Validation rule", 32, "Check applied"),
        Column("message", "Message when it fails", 36, "Message shown, exactly as on screen"),
    ] + review_columns(), field_rows, description="Every field and list column of every screen")
    date_sheets.append(ws)

    action_rows = [{"id": s.id, "screen": s.title, **a} for s in pack.screens for a in s.actions]
    wb.sheet("Actions", [
        Column("id", "Screen ID", 12, "Screen identifier"),
        Column("screen", "Screen", 22, "Screen name"),
        Column("button", "Button", 26, "Button or link as shown"),
        Column("who", "Who", 28, "Permission or role that sees it"),
        Column("when", "Enabled when", 30, "Condition for the button to be offered"),
        Column("what", "What happens", 50, "Effect of the action"),
        Column("status", "Resulting status", 20, "Status after the action"),
        Column("notification", "Notification", 28, "Notice or e-mail sent"),
    ], action_rows, description="Every action (button) of every screen")

    rule_rows = pack.rules()
    ws = wb.sheet("Business rules", [
        Column("id", "Rule ID", 12, "Rule identifier in this workbook"),
        Column("screen", "Screen ID", 12, "Screen where the rule applies"),
        Column("title", "Screen", 22, "Screen name"),
        Column("rule", "Business rule", 70, "The rule as built"),
        Column("frs", "FRs", 20, "Functional requirements that state the rule"),
    ] + review_columns(), rule_rows, description="The business rules stated on the screens")
    date_sheets.append(ws)

    ws = wb.sheet("Messages", [
        Column("id", "Message ID", 13, "Message identifier in this set"),
        Column("where", "Where it appears", 34, "Screen or dialog that shows it"),
        Column("layer", "Shown by", 11, "Screen (checked as you type), Server (checked on save or action) or Upload row"),
        Column("code", "Code (Reference)", 26, "Code shown as the Reference under a server message"),
        Column("text", "Message", 60, "Exact text; <name> is filled in by the system"),
        Column("kind", "Type", 13, "Validation, Error, Warning, Confirmation or Information",
               values=["Validation", "Error", "Warning", "Confirmation", "Information"]),
        Column("fix", "What the user does", 44, "The correction or next step"),
        Column("source", "Source", 30, "Where the message is defined (for the project team)"),
    ] + review_columns(), pack.messages, description="Every validation, error, warning and confirmation message")
    date_sheets.append(ws)

    wb.sheet("Notifications", [
        Column("id", "ID", 8, "Notification identifier"),
        Column("channel", "Channel", 12, "In-app, E-mail or Alert inbox"),
        Column("trigger", "Trigger", 40, "What causes it"),
        Column("recipient", "Recipient", 34, "Who receives it"),
        Column("template", "Template and text", 60, "Subject / title and body"),
        Column("content", "Content", 40, "Attachments and data"),
        Column("screen", "Screen", 26, "Where it is raised or opened"),
        Column("frs", "FRs", 16, "Functional requirements"),
    ], [{**n, "frs": ", ".join(n.get("frs") or [])} for n in pack.notifications],
        description="In-app notifications and e-mails")

    menu_rows = [r for rows in pack.menus().values() for r in rows]
    wb.sheet("Menu by persona", [
        Column("persona", "Persona", 26, "Persona name"),
        Column("role", "Role", 16, "BIBS role code"),
        Column("group", "Group", 18, "Sidebar group"),
        Column("section", "Section", 24, "Sidebar section"),
        Column("screen", "Screen", 26, "Menu entry"),
        Column("brd", "BRD", 12, "BRD that owns the screen"),
        Column("own", "New Business", 12, "Yes for a screen of this set"),
        Column("path", "Route", 30, "Address of the screen"),
    ], menu_rows, description="What each New Business persona sees in the sidebar (from the role grants)")

    upload_rows = []
    for t in pack.uploads:
        for i, c in enumerate(t.columns, start=1):
            upload_rows.append({"type": t.title, "code": t.code, "permission": t.permission, "no": i,
                                "header": c["header"], "required": "Y" if c["required"] else "N",
                                "ctype": c["type"].replace("_", "/").title(), "description": c["description"],
                                "example": c["example"]})
    wb.sheet("Upload templates", [
        Column("type", "Upload type", 28, "Upload type as shown on Bulk Uploads"),
        Column("code", "Code", 20, "Upload type code"),
        Column("permission", "Permission", 18, "Permission of the type (with BULK_PROCESS)"),
        Column("no", "No.", 6, "Column order in the template", kind="number"),
        Column("header", "Column", 24, "Template header, unchanged"),
        Column("required", "Mandatory", 10, "Y or N"),
        Column("ctype", "Type", 10, "Text, Number, Date, Yes/No"),
        Column("description", "Content", 50, "What to enter"),
        Column("example", "Example", 22, "Example value (fictitious)"),
    ], upload_rows, description="Columns of every New Business upload template")

    wb.sheet("Cross-BRD contract", [
        Column("id", "ID", 8, "Contract line"),
        Column("party", "BRD or system", 26, "Other BRD or external system"),
        Column("direction", "Direction", 10, "In: taken by New Business; Out: handed over"),
        Column("what", "What", 50, "Data or function exchanged"),
        Column("how", "When and how", 50, "Trigger and mechanism"),
        Column("owner", "Owner of the data", 26, "Who maintains it"),
        Column("screens", "Screens", 30, "New Business screens concerned"),
        Column("frs", "FRs", 20, "Functional requirements"),
    ], [{**c, "frs": ", ".join(c.get("frs") or [])} for c in pack.contract],
        description="What New Business takes from and hands to the other BRDs and systems; a change is a change request")

    signatories = [
        ("Product Owner", "BDOI"), ("Head, Marketing Business Services and System Support", "BDOI"),
        ("Unit Head, Processing", "BDOI"), ("Head, Retail Marketing", "BDOI"),
        ("Unit Head, Combank and Corbank", "BDOI"), ("Head, Technical Support Unit", "BDOI"),
        ("Head, Comptrollership", "BDOI"), ("Program Manager, Business Project Services", "BDO Unibank ESG"),
        ("Project Manager", brand.VENDOR)]
    ws = wb.sheet("Sign-off", [
        Column("role", "Role", 40, "Signatory role"),
        Column("org", "Organisation", 20, "Organisation"),
        Column("name", "Name", 28, "Name of the signatory"),
        Column("decision", "Decision", 22, "Approved, Approved with comments or Not approved", values=SIGNOFF_VALUES),
        Column("comments", "Comments", 40, "Conditions of the approval"),
        Column("date", "Date", 14, "Date of the decision", kind="date"),
    ], [{"role": r, "org": o} for r, o in signatories],
        description=f"Sign-off of the {m['release_set']} v{m['version']}", freeze_first_column=False)
    date_sheets.append(ws)
    for sheet in date_sheets:
        for cell in sheet[4]:
            if cell.value in ("Review date", "Date"):
                _date_validation(sheet, get_column_letter(cell.column))
    name = brand.output_name("Signoff", m["brd"], m["name"], str(m["version"]), "xlsx")
    return wb.save(brand.out_dir(m["brd"], "Signoff") / name)


# ============================================================================ CLI


def report(pack: Pack) -> str:
    return (f"{pack.meta['brd']} {pack.meta['name']}: {len(pack.screens)} screens, "
            f"{sum(len(s.fields) for s in pack.screens)} fields, {sum(len(s.actions) for s in pack.screens)} actions, "
            f"{len(pack.rules())} rules, {len(pack.messages)} messages, {len(pack.notifications)} notifications, "
            f"{len(pack.uploads)} upload types, {len(pack.documents)} documents, {len(pack.walkthroughs)} walkthroughs, "
            f"{len(pack.shots())} screenshots")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build the business sign-off workbook of a BRD from its pack")
    ap.add_argument("pack", help="pack.yaml of the BRD (e.g. brd01/pack.yaml)")
    ap.add_argument("--check", action="store_true", help="check the pack against the code and the FRS only")
    ap.add_argument("--manifest", metavar="JSON", help="write the screenshot manifest for capture_pack.cjs")
    args = ap.parse_args(argv)
    path = Path(args.pack)
    if not path.exists():
        path = HERE / args.pack
    pack = Pack(path)
    problems = pack.check()
    for p in problems:
        print(f"{path.parent.name}: {p}")
    print(report(pack))
    if args.manifest:
        manifest = {"brd": pack.meta["brd"], "out": str(pack.dir / pack.meta.get("screenshot_dir", "screenshots")),
                    "personas": pack.personas, "shots": pack.shots()}
        Path(args.manifest).write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        print(f"manifest: {args.manifest}")
    if args.check or args.manifest:
        return 1 if problems else 0
    print(f"xlsx: {build_workbook(pack)}")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
