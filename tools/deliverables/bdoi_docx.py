"""Word builder for the BIBS client pack (BDO Insure template).

Two ways to use it:

* **Python API.** ``BdoiDocument`` builds a document piece by piece: ``cover()``, ``front_matter()``,
  ``heading()``, ``paragraph()``, ``bullets()``, ``table()``, ``requirement()``, ``callout()``,
  ``figure()``, ``glossary()``, ``signoff()``. ``save()`` writes the .docx; ``publish()`` also
  writes the PDF and fills the table of contents with the real page numbers.
* **Markdown-like sources.** ``build_markdown(src)`` reads a source under ``docs/deliverables/src``
  (YAML front matter plus the syntax described in ``tools/deliverables/README.md``) and publishes
  it to ``docs/deliverables/out``.

Layout: US Letter portrait (the BDOI BRDs are Letter), Arial 10.5 pt body, headings in Header Blue
numbered 1 / 1.1 / 1.1.1, tables with a Header Blue header row, banded #E5F5FF rows and 0.5 pt
#C2C2C1 borders, Yellow only as a thin accent rule under the running header and on the cover.
"""

from __future__ import annotations

import copy
import datetime as _dt
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent))
import brand  # noqa: E402

BODY_PT = 10.5
TABLE_PT = 9.0
SMALL_PT = 8.0
PAGE_W_CM = 21.59
PAGE_H_CM = 27.94
MARGIN_LR_CM = 2.0
MARGIN_TB_CM = 2.2

CALLOUT_KINDS = {
    # kind: (label, bar colour, background)
    "note": ("Note", brand.HEADER_BLUE, brand.BG_BLUE),
    "info": ("Note", brand.HEADER_BLUE, brand.BG_BLUE),
    "warning": ("Warning", brand.DANGER, brand.DANGER_BG),
    "parked": ("Parked", brand.MUTED, brand.DIRTY_WHITE),
    "decision": ("Decision", brand.SUCCESS, brand.SUCCESS_BG),
    "question": ("Open question", brand.AMBER, brand.AMBER_BG),
}


# --------------------------------------------------------------------------------------------
# Low-level XML helpers
# --------------------------------------------------------------------------------------------


def _set_run_font(rpr_owner, name: str = brand.FONT) -> None:
    """Sets every font slot (ascii, hAnsi, cs, eastAsia) and removes theme font references."""
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        if "Theme" in attr or "theme" in attr:
            del rfonts.attrib[attr]
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(slot), name)


def _shade(cell, fill_hex: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcpr.append(shd)


def _cell_margins(cell, top=50, bottom=50, left=90, right=90) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def _cell_borders(cell, **sides: tuple[int, str] | None) -> None:
    """Sets borders of one cell: side=(size in eighths of a point, colour) or None for 'nil'."""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for side, spec in sides.items():
        el = OxmlElement(f"w:{side}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec[1])
        borders.append(el)


def _table_borders(table: Table, size: int = 4, colour: str = brand.BORDER) -> None:
    """0.5 pt borders (size 4 = 4/8 pt) on every edge and inner line."""
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tblpr.append(borders)


def _table_no_borders(table: Table) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def _table_fixed(table: Table, widths_cm: Sequence[float]) -> None:
    """Fixed layout with explicit column widths (Word and LibreOffice both honour tblGrid)."""
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    grid = table._tbl.tblGrid
    for i, col in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_cm):
            col.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _cant_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _row_height(row, cm: float) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(cm * 567)))
    el.set(qn("w:hRule"), "atLeast")
    trpr.append(el)


def _para_border(paragraph: Paragraph, side: str, size: int, colour: str, space: int = 1) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), colour)
    pbdr.append(el)


def _add_field(paragraph: Paragraph, instr: str, placeholder: str = "1", size: float | None = None,
               colour: str | None = None) -> None:
    """Adds a simple field (PAGE, NUMPAGES, ...) as begin / instr / separate / result / end runs."""

    def run_with(el):
        r = paragraph.add_run()
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = RGBColor.from_string(colour)
        r._r.append(el)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    r = paragraph.add_run(placeholder)
    if size:
        r.font.size = Pt(size)
    if colour:
        r.font.color.rgb = RGBColor.from_string(colour)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def _bookmark(paragraph: Paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _keep_with_next(paragraph: Paragraph, on: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = on


def keep_table_together(table: Table) -> None:
    """Keeps a short table on one page and with the paragraph after it.

    Word follows keep-with-next on the cell paragraphs; LibreOffice follows the same flags, so a
    requirement header table never splits from its heading or across two pages.
    """
    for row in table.rows:
        _cant_split(row)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------------------------
# Inline Markdown
# --------------------------------------------------------------------------------------------

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\)|<br\s*/?>)")


def add_inline(paragraph: Paragraph, text: str, size: float | None = None, colour: str | None = None,
               bold: bool = False) -> None:
    """Adds text with **bold**, *italic*, `code`, [link](url) (text only) and <br> line breaks."""
    for part in _INLINE.split(text):
        if not part:
            continue
        run_bold, italic, code = bold, False, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, run_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            part = part[1:part.index("](")]
        elif re.fullmatch(r"<br\s*/?>", part):
            paragraph.add_run().add_break()
            continue
        run = paragraph.add_run(part)
        run.bold = run_bold or None
        run.italic = italic or None
        if code:
            _set_run_font(run._r, "Courier New")
            run.font.size = Pt((size or BODY_PT) - 1)
            run.font.color.rgb = RGBColor.from_string(brand.NEAR_BLACK)
        else:
            if size:
                run.font.size = Pt(size)
            if colour:
                run.font.color.rgb = RGBColor.from_string(colour)


# --------------------------------------------------------------------------------------------
# The document
# --------------------------------------------------------------------------------------------


@dataclass
class HeadingEntry:
    level: int
    text: str
    bookmark: str
    in_toc: bool = True


@dataclass
class DocMeta:
    """Document identity shown on the cover, in the header and in the footer."""

    title: str
    doc_type: str = ""  # e.g. "Functional Requirements Specification"
    doc_code: str = ""  # e.g. "FRS"
    subtitle: str = ""
    brd: str = ""  # e.g. "BRD-03"
    doc_id: str = ""  # e.g. "BIBS-FRS-BRD-03"
    version: str = "1.0"
    date: str = ""
    status: str = "Issued for BDOI review"
    classification: str = brand.CLASSIFICATION
    header_title: str = ""
    control: list[dict[str, str]] = field(default_factory=list)
    distribution: list[dict[str, str]] = field(default_factory=list)


class BdoiDocument:
    """A Word document in the BDO Insure template.

    Typical use::

        doc = BdoiDocument(DocMeta(title="Product Maintenance", doc_type="Functional Requirements
                           Specification", doc_code="FRS", brd="BRD-03", version="1.0"))
        doc.cover(); doc.front_matter()
        doc.heading("Introduction"); doc.paragraph("...")
        doc.publish(Path("out/FRS/BIBS_FRS_BRD-03_Product_Maintenance_v1.0.docx"))
    """

    def __init__(self, meta: DocMeta, *, h1_page_break: bool = True, base_dir: Path | None = None):
        self.meta = meta
        if not meta.date:
            meta.date = _dt.date.today().strftime("%d %B %Y").lstrip("0")
        self.h1_page_break = h1_page_break
        self.base_dir = base_dir or Path.cwd()
        self.doc = Document()
        self.headings: list[HeadingEntry] = []
        self._numbers = [0, 0, 0, 0]
        self._bookmark_id = 100
        self._figure_no = 0
        self._table_no = 0
        self._toc_anchor: Paragraph | None = None
        self._toc_paragraphs: list[Any] = []
        self._first_body_heading = True
        self._setup_page(self.doc.sections[0])
        self._setup_styles()
        self._header_footer(self.doc.sections[0])

    # ---------------------------------------------------------------- page and styles

    @property
    def text_width_cm(self) -> float:
        s = self.doc.sections[-1]
        return (s.page_width - s.left_margin - s.right_margin) / 360000

    def _setup_page(self, section, landscape: bool = False) -> None:
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Cm(PAGE_H_CM), Cm(PAGE_W_CM)
        else:
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
        section.left_margin = section.right_margin = Cm(MARGIN_LR_CM)
        section.top_margin = Cm(MARGIN_TB_CM + 0.3)
        section.bottom_margin = Cm(MARGIN_TB_CM)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(0.9)

    def _style_font(self, style, size: float, colour: str, bold: bool = False, italic: bool = False):
        style.font.name = brand.FONT
        _set_run_font(style.element.get_or_add_rPr())
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(colour)
        style.font.bold = bold
        style.font.italic = italic

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        self._style_font(normal, BODY_PT, brand.TEXT)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        # Document defaults, so table cells and headers inherit Arial too.
        rpr_default = self.doc.styles.element.find(qn("w:docDefaults"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
            if rpr is not None:
                _set_run_font(rpr)

        spec = {
            1: (16, 18, 8, brand.HEADER_BLUE),
            2: (13, 14, 6, brand.HEADER_BLUE),
            3: (11.5, 12, 4, brand.HEADER_BLUE),
            4: (10.5, 8, 3, brand.NEAR_BLACK),
        }
        for level, (size, before, after, colour) in spec.items():
            st = styles[f"Heading {level}"]
            self._style_font(st, size, colour, bold=True)
            pf = st.paragraph_format
            pf.space_before = Pt(before)
            pf.space_after = Pt(after)
            pf.keep_with_next = True
            pf.line_spacing = 1.0
        # Heading 4 is a run-in label inside requirements: not in the outline or the TOC.
        h4 = styles["Heading 4"]
        ppr = h4.element.get_or_add_pPr()
        for old in ppr.findall(qn("w:outlineLvl")):
            ppr.remove(old)
        ol = OxmlElement("w:outlineLvl")
        ol.set(qn("w:val"), "9")
        ppr.append(ol)

        for name, size, colour, bold, italic in (
            ("Caption", 9, brand.MUTED, False, True),
            ("BDOI Table", TABLE_PT, brand.TEXT, False, False),
            ("BDOI Table Header", TABLE_PT, brand.WHITE, True, False),
            ("BDOI Small", SMALL_PT, brand.MUTED, False, False),
            ("BDOI Front Heading", 16, brand.HEADER_BLUE, True, False),
            ("BDOI Label", 10, brand.HEADER_BLUE, True, False),
            ("BDOI Spacer", 3, brand.TEXT, False, False),
        ):
            st = styles[name] if name in [s.name for s in styles] else styles.add_style(
                name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            self._style_font(st, size, colour, bold, italic)
            pf = st.paragraph_format
            if name.startswith("BDOI Table"):
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.05
            elif name == "BDOI Front Heading":
                pf.space_before = Pt(0)
                pf.space_after = Pt(10)
                pf.keep_with_next = True
            elif name == "BDOI Label":
                pf.space_before = Pt(8)
                pf.space_after = Pt(3)
                pf.keep_with_next = True
            elif name == "BDOI Spacer":
                pf.space_before = Pt(0)
                pf.space_after = Pt(5)
                pf.line_spacing = 1.0
            elif name == "Caption":
                pf.space_before = Pt(3)
                pf.space_after = Pt(10)

        for level in (1, 2, 3):
            name = f"toc {level}"
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            self._style_font(st, 10 if level > 1 else 10.5, brand.NEAR_BLACK if level == 1 else brand.TEXT,
                             bold=(level == 1))
            pf = st.paragraph_format
            pf.left_indent = Cm(0.6 * (level - 1))
            pf.space_before = Pt(6 if level == 1 else 0)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.0
            pf.tab_stops.add_tab_stop(Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    def _header_footer(self, section) -> None:
        section.different_first_page_header_footer = True
        width = (section.page_width - section.left_margin - section.right_margin) / 360000
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.text = ""
        p.style = self.doc.styles["Normal"]  # the built-in Header style has its own tab stops
        p.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_after = Pt(0)
        title = self.meta.header_title or self.meta.title
        r = p.add_run(title)
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(brand.HEADER_BLUE)
        r.bold = True
        if self.meta.doc_type:
            r2 = p.add_run(f"   {self.meta.doc_type}")
            r2.font.size = Pt(8)
            r2.font.color.rgb = RGBColor.from_string(brand.MUTED)
        p.add_run("\t")
        if brand.BDO_LOGO.exists():
            p.add_run().add_picture(str(brand.BDO_LOGO), height=Cm(0.62))
        _para_border(p, "bottom", 6, brand.YELLOW, space=4)

        footer = section.footer
        footer.is_linked_to_previous = False
        f = footer.paragraphs[0]
        f.text = ""
        f.style = self.doc.styles["Normal"]
        f.paragraph_format.tab_stops.add_tab_stop(Cm(width / 2), WD_TAB_ALIGNMENT.CENTER)
        f.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT)
        _para_border(f, "top", 4, brand.BORDER, space=4)
        for text in (brand.FOOTER_TEXT, "\tPage "):
            r = f.add_run(text)
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(brand.MUTED)
        _add_field(f, "PAGE", "1", 8, brand.MUTED)
        r = f.add_run(" of ")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(brand.MUTED)
        _add_field(f, "NUMPAGES", "1", 8, brand.MUTED)
        r = f.add_run(f"\tVersion {self.meta.version}")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(brand.MUTED)
        # First page (cover): empty header and footer.
        section.first_page_header.is_linked_to_previous = False
        section.first_page_footer.is_linked_to_previous = False

    def new_section(self, landscape: bool = False) -> None:
        """Starts a new page section (portrait or landscape) with its own header and footer."""
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(section, landscape)
        self._header_footer(section)
        section.different_first_page_header_footer = False

    # ---------------------------------------------------------------- cover and front matter

    def cover(self) -> None:
        """Cover page: BDO Insure logo, title block, document facts and the iorta TechNXT line."""
        m = self.meta
        doc = self.doc
        p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        if brand.BDO_LOGO.exists():
            p.add_run().add_picture(str(brand.BDO_LOGO), width=Cm(6.2))
        rule = doc.add_paragraph()
        rule.paragraph_format.space_before = Pt(10)
        rule.paragraph_format.space_after = Pt(0)
        _para_border(rule, "bottom", 12, brand.YELLOW, space=1)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(80)
        spacer.paragraph_format.space_after = Pt(0)

        if m.doc_type:
            kicker = doc.add_paragraph()
            r = kicker.add_run(m.doc_type.upper())
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor.from_string(brand.CTA_BLUE)
            kicker.paragraph_format.space_after = Pt(6)
        t = doc.add_paragraph()
        t.paragraph_format.space_after = Pt(4)
        t.paragraph_format.line_spacing = 1.0
        r = t.add_run(m.title)
        r.bold = True
        r.font.size = Pt(28)
        r.font.color.rgb = RGBColor.from_string(brand.HEADER_BLUE)
        if m.subtitle:
            s = doc.add_paragraph()
            r = s.add_run(m.subtitle)
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor.from_string(brand.NEAR_BLACK)
            s.paragraph_format.space_after = Pt(2)
        sysline = doc.add_paragraph()
        r = sysline.add_run(brand.SYSTEM)
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor.from_string(brand.MUTED)
        sysline.paragraph_format.space_after = Pt(28)

        facts = [
            ("Client", brand.CLIENT),
            ("Document", m.doc_id or m.doc_code),
            ("Version", m.version),
            ("Date", m.date),
            ("Status", m.status),
            ("Classification", m.classification),
        ]
        facts = [(k, v) for k, v in facts if v]
        table = doc.add_table(rows=len(facts), cols=2)
        _table_no_borders(table)
        _table_fixed(table, [3.6, 11.0])
        for row, (k, v) in zip(table.rows, facts):
            for cell in row.cells:
                _cell_margins(cell, 40, 40, 0, 60)
            _cell_borders(row.cells[0], bottom=(4, brand.BORDER_LIGHT))
            _cell_borders(row.cells[1], bottom=(4, brand.BORDER_LIGHT))
            a = row.cells[0].paragraphs[0]
            a.style = self.doc.styles["BDOI Table"]
            add_inline(a, k, size=9.5, colour=brand.MUTED)
            b = row.cells[1].paragraphs[0]
            b.style = self.doc.styles["BDOI Table"]
            add_inline(b, v, size=10, colour=brand.NEAR_BLACK, bold=(k == "Classification"))

        gap = doc.add_paragraph()
        gap.paragraph_format.space_before = Pt(150 if len(facts) >= 6 else 180)
        gap.paragraph_format.space_after = Pt(0)
        foot = doc.add_paragraph()
        _para_border(foot, "top", 4, brand.BORDER, space=6)
        foot.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_width_cm), WD_TAB_ALIGNMENT.RIGHT)
        r = foot.add_run("Prepared by iorta TechNXT for " + brand.CLIENT_SHORT)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(brand.MUTED)
        foot.add_run("\t")
        if brand.IORTA_LOGO.exists():
            foot.add_run().add_picture(str(brand.IORTA_LOGO), height=Cm(0.8))
        note = doc.add_paragraph()
        r = note.add_run(f"{brand.PLATFORM}. This document is confidential to {brand.CLIENT} and iorta TechNXT.")
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(brand.PLACEHOLDER)
        self.page_break()

    def front_heading(self, text: str) -> Paragraph:
        """A heading of the front matter (not numbered, not in the table of contents)."""
        return self.doc.add_paragraph(text, style="BDOI Front Heading")

    def document_control(self, rows: list[dict[str, str]] | None = None) -> None:
        rows = rows or self.meta.control or [{
            "version": self.meta.version, "date": self.meta.date, "author": "iorta TechNXT",
            "reviewer": "", "approver": "", "change": "First issue"}]
        self.front_heading("Document Control")
        self.table(
            ["Version", "Date", "Author", "Reviewer", "Approver", "Change"],
            [[r.get("version", ""), r.get("date", ""), r.get("author", ""), r.get("reviewer", ""),
              r.get("approver", ""), r.get("change", "")] for r in rows],
            widths=[1.6, 2.4, 3.0, 3.0, 3.0, 4.6], caption=None)

    def distribution(self, rows: list[dict[str, str]] | None = None) -> None:
        rows = rows or self.meta.distribution
        if not rows:
            return
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        self.front_heading("Distribution List")
        self.table(
            ["Name / Unit", "Role", "Organisation", "Purpose"],
            [[r.get("name", ""), r.get("role", ""), r.get("organisation", ""), r.get("purpose", "")]
             for r in rows],
            widths=[4.6, 4.0, 4.0, 5.0])

    def toc(self) -> None:
        """Table of contents: a real Word TOC field (levels 1-3).

        The field result is pre-filled with the headings and, after ``publish()``, with the page
        numbers of the generated PDF, so the PDF and a Word reader both see a complete table.
        Word refreshes it with Update Field.
        """
        self.front_heading("Contents").paragraph_format.page_break_before = True
        self._toc_anchor = self.doc.add_paragraph()
        self._toc_anchor.paragraph_format.space_after = Pt(0)

    def front_matter(self) -> None:
        """Document control, distribution list and table of contents, then a page break."""
        self.document_control()
        self.distribution()
        self.toc()

    # ---------------------------------------------------------------- body blocks

    def page_break(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def heading(self, text: str, level: int = 1, numbered: bool = True, toc: bool = True) -> Paragraph:
        """Heading 1-3 (numbered 1 / 1.1 / 1.1.1 unless numbered=False) or 4 (a label)."""
        # A new chapter starts on a new page. "Page break before" on the heading itself never leaves
        # an empty page when the previous page is exactly full (a break paragraph would).
        break_before = level == 1 and (self.h1_page_break or self._first_body_heading)
        if level == 1:
            self._first_body_heading = False
        label = text
        if numbered and level <= 3:
            self._numbers[level - 1] += 1
            for i in range(level, len(self._numbers)):
                self._numbers[i] = 0
            label = ".".join(str(n) for n in self._numbers[:level]) + " " + text
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        if break_before:
            p.paragraph_format.page_break_before = True
        add_inline(p, label)
        if level <= 3:
            self._bookmark_id += 1
            name = f"_Toc{self._bookmark_id:06d}"
            _bookmark(p, name, self._bookmark_id)
            self.headings.append(HeadingEntry(level, label, name, toc))
        if level == 1:
            _para_border(p, "bottom", 4, brand.FIELD_BLUE, space=4)
        return p

    def label(self, text: str) -> Paragraph:
        """A small blue run-in heading used inside requirements ("Main flow", ...)."""
        p = self.doc.add_paragraph(style="BDOI Label")
        add_inline(p, text)
        return p

    def paragraph(self, text: str = "", style: str | None = None, keep_with_next: bool = False,
                  align: str | None = None, size: float | None = None) -> Paragraph:
        p = self.doc.add_paragraph(style=style)
        add_inline(p, text, size=size)
        if keep_with_next:
            _keep_with_next(p)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def bullets(self, items: Iterable[str | tuple[int, str]], numbered: bool = False, size: float | None = None,
                space_after: float = 2) -> None:
        """Bulleted or numbered list. An item may be (level, text) for a second level."""
        counters = [0, 0, 0]
        items = list(items)
        for idx, item in enumerate(items):
            level, text = item if isinstance(item, tuple) else (0, item)
            counters[level] += 1
            for i in range(level + 1, 3):
                counters[i] = 0
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            indent = 0.55 + 0.6 * level
            pf.left_indent = Cm(indent)
            pf.first_line_indent = Cm(-0.5)
            pf.space_after = Pt(space_after if idx < len(items) - 1 else 6)
            pf.tab_stops.add_tab_stop(Cm(indent))
            if numbered and level == 0:
                marker = f"{counters[0]}."
            elif numbered:
                marker = f"{chr(96 + counters[level])})"
            else:
                marker = "•" if level == 0 else "–"
            r = p.add_run(marker + "\t")
            r.font.color.rgb = RGBColor.from_string(brand.HEADER_BLUE)
            if size:
                r.font.size = Pt(size)
            add_inline(p, text, size=size)

    def caption(self, kind: str, text: str) -> Paragraph:
        if kind == "Figure":
            self._figure_no += 1
            n = self._figure_no
        else:
            self._table_no += 1
            n = self._table_no
        p = self.doc.add_paragraph(style="Caption")
        r = p.add_run(f"{kind} {n}: ")
        r.bold = True
        r.italic = False
        add_inline(p, text)
        if kind == "Figure":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            _keep_with_next(p)
            p.paragraph_format.space_after = Pt(3)
        return p

    def table(self, headers: Sequence[str] | None, rows: Sequence[Sequence[Any]],
              widths: Sequence[float] | None = None, caption: str | None = None,
              size: float = TABLE_PT, status_cols: Sequence[int] = (), banded: bool = True,
              first_col_bold: bool = False, keep_rows: bool = True, align: Sequence[str] | None = None) -> Table:
        """A BDOI table: Header Blue header row (white bold), banded rows, 0.5 pt grey borders.

        widths are in cm (relative values are scaled to the text width); status_cols are column
        indexes whose values get the status colours of brand.STATUS_COLOURS.
        """
        ncols = len(headers) if headers else max(len(r) for r in rows)
        total = self.text_width_cm
        if widths is None:
            widths = [total / ncols] * ncols
        scale = total / sum(widths)
        widths = [w * scale for w in widths]
        if caption:
            self.caption("Table", caption)
        table = self.doc.add_table(rows=(1 if headers else 0) + len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(table)
        _table_fixed(table, widths)
        r0 = 0
        if headers:
            hdr = table.rows[0]
            _repeat_header(hdr)
            _cant_split(hdr)
            for i, h in enumerate(headers):
                cell = hdr.cells[i]
                _shade(cell, brand.HEADER_BLUE)
                _cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.style = self.doc.styles["BDOI Table Header"]
                add_inline(p, str(h), size=size, colour=brand.WHITE, bold=True)
            r0 = 1
        for ri, values in enumerate(rows):
            row = table.rows[r0 + ri]
            if keep_rows:
                _cant_split(row)
            for ci in range(ncols):
                value = values[ci] if ci < len(values) else ""
                cell = row.cells[ci]
                _cell_margins(cell)
                if banded and ri % 2 == 1:
                    _shade(cell, brand.BG_BLUE)
                self._fill_cell(cell, value, size, bold=(first_col_bold and ci == 0))
                if align and ci < len(align) and align[ci] in ("center", "right"):
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align[ci] == "center" else WD_ALIGN_PARAGRAPH.RIGHT
                if ci in status_cols:
                    key = re.sub(r"[*`]", "", str(value)).strip().upper()
                    if key in brand.STATUS_COLOURS:
                        fill, font = brand.STATUS_COLOURS[key]
                        _shade(cell, fill)
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.color.rgb = RGBColor.from_string(font)
                                run.bold = True
        self._after_table()
        return table

    def _fill_cell(self, cell, value: Any, size: float, bold: bool = False) -> None:
        """Cell text; a list value becomes bullet lines, '<br>' a line break."""
        p = cell.paragraphs[0]
        p.style = self.doc.styles["BDOI Table"]
        if isinstance(value, (list, tuple)):
            for i, item in enumerate(value):
                q = p if i == 0 else cell.add_paragraph(style="BDOI Table")
                q.paragraph_format.left_indent = Cm(0.3)
                q.paragraph_format.first_line_indent = Cm(-0.3)
                q.add_run("• ").font.size = Pt(size)
                add_inline(q, str(item), size=size)
        else:
            add_inline(p, "" if value is None else str(value), size=size, bold=bold)

    def _after_table(self) -> None:
        self.doc.add_paragraph(style="BDOI Spacer")

    def key_values(self, pairs: Sequence[tuple[str, Any]], columns: int = 2, size: float = TABLE_PT,
                   label_width: float = 2.6, keep_together: bool = True) -> Table:
        """A compact label / value grid (1 or 2 pairs per row), labels shaded blue."""
        total = self.text_width_cm
        rows = [pairs[i:i + columns] for i in range(0, len(pairs), columns)]
        value_w = (total - label_width * columns) / columns
        widths = [label_width, value_w] * columns
        table = self.doc.add_table(rows=len(rows), cols=2 * columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(table)
        _table_fixed(table, widths)
        for ri, row_pairs in enumerate(rows):
            row = table.rows[ri]
            _cant_split(row)
            for pi in range(columns):
                lc, vc = row.cells[2 * pi], row.cells[2 * pi + 1]
                _cell_margins(lc)
                _cell_margins(vc)
                _shade(lc, brand.BG_BLUE)
                if pi < len(row_pairs):
                    k, v = row_pairs[pi]
                    p = lc.paragraphs[0]
                    p.style = self.doc.styles["BDOI Table"]
                    add_inline(p, k, size=size, colour=brand.HEADER_BLUE, bold=True)
                    self._fill_cell(vc, v, size)
            if len(row_pairs) < columns:  # merge the empty trailing pair into the last value
                a = row.cells[2 * len(row_pairs) - 1]
                b = row.cells[2 * columns - 1]
                a.merge(b)
        if keep_together:
            keep_table_together(table)
        self._after_table()
        return table

    def callout(self, text: str | Sequence[str], kind: str = "note", title: str | None = None) -> None:
        """A shaded box with a coloured left bar: note, warning, parked, decision or question."""
        label, bar, bg = CALLOUT_KINDS.get(kind, CALLOUT_KINDS["note"])
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_no_borders(table)
        _table_fixed(table, [self.text_width_cm])
        cell = table.rows[0].cells[0]
        _cant_split(table.rows[0])
        _shade(cell, bg)
        _cell_margins(cell, 90, 90, 160, 120)
        _cell_borders(cell, left=(24, bar), top=None, bottom=None, right=None)
        lines = [text] if isinstance(text, str) else list(text)
        p = cell.paragraphs[0]
        p.style = self.doc.styles["BDOI Table"]
        head = p.add_run((title or label) + ". ")
        head.bold = True
        head.font.size = Pt(9.5)
        head.font.color.rgb = RGBColor.from_string(bar)
        add_inline(p, lines[0], size=9.5)
        for extra in lines[1:]:
            q = cell.add_paragraph(style="BDOI Table")
            q.paragraph_format.space_before = Pt(3)
            if extra.startswith("- "):
                q.paragraph_format.left_indent = Cm(0.35)
                q.paragraph_format.first_line_indent = Cm(-0.3)
                q.add_run("• ").font.size = Pt(9.5)
                extra = extra[2:]
            add_inline(q, extra, size=9.5)
        self._after_table()

    def figure(self, path: str | Path, caption: str, width_cm: float | None = None) -> None:
        """A picture (PNG, JPG or a Graphviz .dot source rendered in BDO colours) with a caption."""
        path = Path(path)
        if not path.is_absolute():
            path = (self.base_dir / path).resolve()
        if path.suffix == ".dot":
            path = render_dot(path)
        width = width_cm or self.text_width_cm
        # Keep tall figures on one page: cap the height at 80 % of the page body.
        try:
            from PIL import Image

            with Image.open(path) as im:
                w, h = im.size
            max_h = 0.8 * (PAGE_H_CM - 2 * MARGIN_TB_CM)
            if h / w * width > max_h:
                width = max_h * w / h
        except Exception:  # pragma: no cover - PIL missing
            pass
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        _keep_with_next(p)
        p.add_run().add_picture(str(path), width=Cm(width))
        self.caption("Figure", caption)

    def glossary(self, entries: dict[str, str] | Sequence[tuple[str, str]], caption: str | None = None) -> None:
        items = list(entries.items()) if isinstance(entries, dict) else list(entries)
        items.sort(key=lambda kv: kv[0].lower())
        self.table(["Term", "Meaning"], items, widths=[3.6, 14.0], caption=caption, first_col_bold=True)

    def signoff(self, rows: Sequence[dict[str, str]], intro: str | None = None) -> None:
        """Sign-off block: name, role, organisation, signature and date, with tall signing rows."""
        if intro:
            self.paragraph(intro)
        headers = ["Name", "Role", "Organisation", "Signature", "Date"]
        table = self.table(headers, [[r.get("name", ""), r.get("role", ""), r.get("organisation", ""), "", ""]
                                     for r in rows], widths=[4.0, 4.2, 4.2, 3.2, 2.0], banded=False)
        for row in table.rows[1:]:
            _row_height(row, 1.3)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def requirement(self, fr: dict[str, Any]) -> None:
        """A functional requirement in the house pattern (see README "Functional requirement block").

        Keys: id, title, brd (list), actor, priority, fit, screens, api, description, preconditions,
        main_flow, alternate_flows, rules, validations, fields, notifications, audit, acceptance.
        """
        self.heading(f"{fr['id']} {fr['title']}", level=3, numbered=False)
        brd_refs = fr.get("brd", [])
        if isinstance(brd_refs, str):
            brd_refs = [brd_refs]
        pairs = [
            ("BRD trace", ", ".join(brd_refs)),
            ("Actor", fr.get("actor", "")),
            ("Priority", fr.get("priority", "Must have")),
            ("Fit", fr.get("fit", "")),
        ]
        if fr.get("screens"):
            pairs.append(("Screens", fr["screens"]))
        if fr.get("api"):
            pairs.append(("API", fr["api"]))
        self.key_values(pairs, columns=2, label_width=2.2)
        if fr.get("description"):
            self.label("Description")
            for para in _as_list(fr["description"]):
                self.paragraph(para)
        if fr.get("preconditions"):
            self.label("Preconditions")
            self.bullets(_as_list(fr["preconditions"]))
        if fr.get("main_flow"):
            self.label("Main flow")
            self.bullets(_levels(fr["main_flow"]), numbered=True)
        if fr.get("alternate_flows"):
            self.label("Alternate and exception flows")
            self.bullets(_levels(fr["alternate_flows"]))
        if fr.get("rules"):
            self.label("Business rules")
            self.table(["Rule", "Statement", "Kind", "Maintained in"],
                       [_pad(r, 4) for r in fr["rules"]], widths=[1.2, 9.0, 2.2, 5.2])
        if fr.get("validations"):
            self.label("Validations and messages")
            self.table(["Check", "Message shown", "Code"],
                       [_pad(r, 3) for r in fr["validations"]], widths=[4.6, 7.2, 5.8])
        if fr.get("fields"):
            self.label("Screens and fields" + (f": {fr['fields_screen']}" if fr.get("fields_screen") else ""))
            mand = {"conditional": "Cond.", "yes": "Yes", "no": "No", "true": "Yes", "false": "No"}
            rows = []
            for r in fr["fields"]:
                r = _pad(r, 5)
                r[2] = mand.get(str(r[2]).strip().lower(), r[2])
                rows.append(r)
            self.table(["Field", "Type", "Mand.", "Source / LOV", "Validation"],
                       rows, widths=[3.4, 2.1, 1.3, 5.0, 5.8],
                       align=[None, None, "center"])
        if fr.get("notifications"):
            self.label("Notifications")
            self.bullets(_as_list(fr["notifications"]))
        if fr.get("audit"):
            self.label("Audit")
            self.bullets(_as_list(fr["audit"]))
        if fr.get("acceptance"):
            self.label("Acceptance criteria")
            self.bullets(_as_list(fr["acceptance"]), numbered=True)

    # ---------------------------------------------------------------- TOC and output

    def _write_toc(self, pages: dict[str, int] | None) -> None:
        """(Re)writes the TOC field after the anchor paragraph; pages maps bookmark -> page."""
        if self._toc_anchor is None:
            return
        for el in self._toc_paragraphs:
            el.getparent().remove(el)
        self._toc_paragraphs = []
        entries = [h for h in self.headings if h.in_toc and h.level <= 3]
        body = self._toc_anchor._p.getparent()
        anchor = self._toc_anchor._p
        new_paras = []
        for i, h in enumerate(entries):
            p = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), self.doc.styles[f"toc {h.level}"].style_id)
            ppr.append(ps)
            p.append(ppr)
            para = Paragraph(p, self._toc_anchor._parent)
            if i == 0:
                self._fld(para, "begin")
                instr = OxmlElement("w:r")
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-3" \\h \\z \\u '
                instr.append(it)
                p.append(instr)
                self._fld(para, "separate")
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), h.bookmark)
            link.set(qn("w:history"), "1")
            page = str(pages.get(h.bookmark, "")) if pages else "0"
            for text in (h.text, "\t", page):
                r = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = text
                r.append(t)
                link.append(r)
            p.append(link)
            if i == len(entries) - 1:
                self._fld(para, "end")
            new_paras.append(p)
        idx = list(body).index(anchor)
        for off, p in enumerate(new_paras, start=1):
            body.insert(idx + off, p)
        self._toc_paragraphs = new_paras

    @staticmethod
    def _fld(para: Paragraph, kind: str) -> None:
        r = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            fc.set(qn("w:dirty"), "false")
        r.append(fc)
        para._p.append(r)

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        if self._toc_anchor is not None and not self._toc_paragraphs:
            self._write_toc(None)
        props = self.doc.core_properties
        props.title = f"{self.meta.doc_type} {self.meta.title}".strip()
        props.subject = brand.SYSTEM
        props.author = brand.VENDOR
        props.company = brand.CLIENT if hasattr(props, "company") else None
        props.keywords = ", ".join(x for x in (self.meta.doc_code, self.meta.brd, "BIBS", "BDOI") if x)
        props.comments = f"Version {self.meta.version}. {self.meta.classification}."
        props.last_modified_by = brand.VENDOR
        self.doc.save(str(path))
        return path

    def publish(self, path: str | Path, pdf: bool = True, keep_pdf: bool = False) -> tuple[Path, Path | None]:
        """Saves the .docx. With pdf=True a PDF is rendered to fill the TOC page numbers; it is kept only
        with keep_pdf=True (the Word file is the editable master; PDFs are produced at issue)."""
        path = self.save(path)
        if not pdf:
            return path, None
        import render  # local module

        pdf_path = render.to_pdf(path)
        if self._toc_anchor is not None:
            pages = heading_pages(pdf_path, self.headings)
            self._write_toc(pages)
            self.save(path)
            if keep_pdf:
                pdf_path = render.to_pdf(path)
        if keep_pdf:
            return path, pdf_path
        pdf_path.unlink(missing_ok=True)
        return path, None


# --------------------------------------------------------------------------------------------
# Helpers for requirements and figures
# --------------------------------------------------------------------------------------------


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    return [str(v) for v in value]


def _levels(items: Any) -> list[str | tuple[int, str]]:
    """Flow steps: a string is level 0; a nested list gives level-1 sub-steps of the previous one."""
    out: list[str | tuple[int, str]] = []
    for item in _as_list(items) if isinstance(items, str) else items:
        if isinstance(item, (list, tuple)):
            out.extend((1, str(sub)) for sub in item)
        elif isinstance(item, dict):  # {"step": [sub, sub]}
            for k, subs in item.items():
                out.append(str(k))
                out.extend((1, str(s)) for s in subs)
        else:
            out.append(str(item))
    return out


def _pad(row: Any, n: int) -> list[Any]:
    row = list(row) if isinstance(row, (list, tuple)) else [row]
    return (row + [""] * n)[:n]


def render_dot(src: Path, dpi: int = 200) -> Path:
    """Renders a Graphviz source to PNG with the BDO tokens and defaults; returns the PNG path.

    The PNG is written next to the source (same name, .png) so it can also be reviewed on its own.
    """
    text = src.read_text(encoding="utf-8")
    for token in sorted(brand.DOT_TOKENS, key=len, reverse=True):
        text = text.replace(token, brand.DOT_TOKENS[token])
    out = src.with_suffix(".png")
    dot = shutil.which("dot")
    if not dot:
        if out.exists():
            return out
        raise RuntimeError("Graphviz 'dot' is not installed and no rendered PNG exists for " + str(src))
    subprocess.run([dot, "-Tpng", f"-Gdpi={dpi}", "-o", str(out)], input=text.encode("utf-8"),
                   check=True)
    return out


def heading_pages(pdf_path: Path, headings: list[HeadingEntry]) -> dict[str, int]:
    """Page number (1-based) of each heading, read from the PDF outline written by LibreOffice."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(pdf_path))
    outline = []
    for item in pdf.get_toc(max_depth=6):
        title = (item.get_title() if hasattr(item, "get_title") else item.title) or ""
        page = item.page_index if hasattr(item, "page_index") else None
        if page is None and hasattr(item, "get_dest"):
            dest = item.get_dest()
            page = dest.get_index() if dest else None
        outline.append((_norm(title), page))
    pdf.close()
    pages: dict[str, int] = {}
    pos = 0
    for h in headings:
        target = _norm(h.text)
        for j in range(pos, len(outline)):
            if outline[j][0] == target and outline[j][1] is not None:
                pages[h.bookmark] = outline[j][1] + 1
                pos = j + 1
                break
    return pages


def _norm(text: str) -> str:
    text = re.sub(r"[*`]", "", text)
    return re.sub(r"\s+", " ", text.replace(" ", " ")).strip().lower()


# --------------------------------------------------------------------------------------------
# Markdown-like sources
# --------------------------------------------------------------------------------------------

_DIRECTIVE = re.compile(r"<!--\s*(\w+)\s*:?\s*(.*?)\s*-->")
_TABLE_OPTS = re.compile(r'(\w+)=("[^"]*"|\S+)')


def _parse_opts(text: str) -> dict[str, str]:
    return {k: v.strip('"') for k, v in _TABLE_OPTS.findall(text)}


def load_source(src: Path) -> tuple[dict[str, Any], list[str]]:
    text = src.read_text(encoding="utf-8")
    meta: dict[str, Any] = {}
    if text.startswith("---"):
        _, fm, body = text.split("---", 2)
        meta = yaml.safe_load(fm) or {}
    else:
        body = text
    return meta, body.splitlines()


def meta_from(front: dict[str, Any]) -> DocMeta:
    return DocMeta(
        title=str(front.get("title", "Untitled")),
        doc_type=str(front.get("doc_type", "")),
        doc_code=str(front.get("doc_code", "")),
        subtitle=str(front.get("subtitle", "")),
        brd=str(front.get("brd", "")),
        doc_id=str(front.get("doc_id", "")),
        version=str(front.get("version", "1.0")),
        date=str(front.get("date", "")),
        status=str(front.get("status", "Issued for BDOI review")),
        classification=str(front.get("classification", brand.CLASSIFICATION)),
        header_title=str(front.get("header_title", "")),
        control=list(front.get("control", []) or []),
        distribution=list(front.get("distribution", []) or []),
    )


def output_path(front: dict[str, Any], src: Path) -> Path:
    if front.get("output"):
        return brand.OUT_DIR / str(front["output"])
    folder = str(front.get("doc_code", "misc"))
    name = brand.output_name(folder, str(front.get("brd", "BRD-00")), str(front.get("name", src.stem)),
                             str(front.get("version", "1.0")), "docx")
    return brand.OUT_DIR / folder / name


def render_body(doc: BdoiDocument, lines: list[str]) -> None:
    """Renders the body syntax (see README) into the document."""
    i = 0
    para: list[str] = []
    pending_opts: dict[str, str] = {}

    def flush_para():
        nonlocal para
        if para:
            doc.paragraph(" ".join(s.strip() for s in para))
            para = []

    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_para()
            i += 1
            continue
        m = _DIRECTIVE.fullmatch(stripped)
        if m:
            flush_para()
            name, rest = m.group(1).lower(), m.group(2)
            if name == "pagebreak":
                doc.page_break()
            elif name == "landscape":
                doc.new_section(landscape=True)
            elif name == "portrait":
                doc.new_section(landscape=False)
            elif name == "table":
                pending_opts = _parse_opts(rest)
            i += 1
            continue
        if stripped.startswith("```"):
            flush_para()
            kind = stripped[3:].strip().lower()
            j = i + 1
            block = []
            while j < n and not lines[j].strip().startswith("```"):
                block.append(lines[j])
                j += 1
            _render_block(doc, kind, "\n".join(block))
            i = j + 1
            continue
        hm = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if hm:
            flush_para()
            level = len(hm.group(1))
            text = hm.group(2)
            numbered = True
            if text.endswith("{-}"):
                text, numbered = text[:-3].rstrip(), False
            doc.heading(text, level=level, numbered=numbered)
            i += 1
            continue
        if stripped.startswith("|"):
            flush_para()
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            header, body_rows = rows[0], [r for r in rows[1:] if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
            opts, pending_opts = pending_opts, {}
            widths = [float(w) for w in opts["widths"].split(",")] if "widths" in opts else None
            status = []
            if "status" in opts:
                names = [s.strip().lower() for s in opts["status"].split(",")]
                status = [k for k, h in enumerate(header) if h.lower() in names]
            doc.table(header, body_rows, widths=widths, caption=opts.get("caption"),
                      size=float(opts.get("size", TABLE_PT)), status_cols=status,
                      first_col_bold=opts.get("bold") == "first")
            continue
        if stripped.startswith(">"):
            flush_para()
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            kind, title = "note", None
            km = re.match(r"^\[!(\w+)\]\s*(.*)$", quote[0])
            if km:
                kind = km.group(1).lower()
                title = km.group(2) or None
                quote = quote[1:]
            text_lines: list[str] = []
            for q in quote:
                if q.startswith("- ") or not text_lines:
                    text_lines.append(q)
                elif q == "":
                    continue
                else:
                    text_lines[-1] += " " + q
            doc.callout(text_lines or [""], kind=kind, title=title)
            continue
        fm = re.match(r"^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)\})?$", stripped)
        if fm:
            flush_para()
            doc.figure(fm.group(2), fm.group(1), float(fm.group(4)) if fm.group(4) else None)
            i += 1
            continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if lm:
            flush_para()
            numbered = lm.group(2)[0].isdigit()
            items: list[str | tuple[int, str]] = []
            while i < n:
                lm2 = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if lm2 and not lm2.group(1) and lm2.group(2)[0].isdigit() != numbered:
                    break  # a top-level item of the other list kind starts a new list
                if lm2:
                    level = 1 if len(lm2.group(1)) >= 2 else 0
                    items.append((level, lm2.group(3)) if level else lm2.group(3))
                    i += 1
                elif lines[i].startswith("   ") and lines[i].strip() and items:
                    last = items[-1]
                    if isinstance(last, tuple):
                        items[-1] = (last[0], last[1] + " " + lines[i].strip())
                    else:
                        items[-1] = last + " " + lines[i].strip()
                    i += 1
                else:
                    break
            doc.bullets(items, numbered=numbered)
            continue
        para.append(line)
        i += 1
    flush_para()


def _render_block(doc: BdoiDocument, kind: str, text: str) -> None:
    if kind in ("fr", "requirement"):
        doc.requirement(yaml.safe_load(text))
    elif kind == "glossary":
        data = yaml.safe_load(text)
        doc.glossary(data if isinstance(data, dict) else [tuple(x) for x in data])
    elif kind == "signoff":
        data = yaml.safe_load(text)
        doc.signoff(data.get("rows", []), intro=data.get("intro")) if isinstance(data, dict) else doc.signoff(data)
    elif kind == "keyvalues":
        data = yaml.safe_load(text)
        doc.key_values([(str(k), v) for k, v in (data.items() if isinstance(data, dict) else data)], columns=1,
                       label_width=4.0)
    elif kind == "table":  # YAML table: {headers, rows, widths, caption, status}
        data = yaml.safe_load(text)
        headers = data.get("headers")
        status = [headers.index(s) for s in data.get("status", []) if headers and s in headers]
        doc.table(headers, data.get("rows", []), widths=data.get("widths"), caption=data.get("caption"),
                  size=float(data.get("size", TABLE_PT)), status_cols=status,
                  first_col_bold=bool(data.get("bold_first", False)))
    else:  # plain preformatted text
        for line in text.splitlines():
            p = doc.paragraph(style="BDOI Small")
            add_inline(p, "`" + line + "`" if line.strip() else "")


FR_ROW_WIDTHS = {"rules": 4, "validations": 3, "fields": 5}


YAML_BLOCK_KINDS = ("fr", "requirement", "glossary", "table", "signoff")
_QUOTED_OR_STRUCTURED = ("'", '"', "|", ">", "[", "{", "&", "*", "!")


def _plain_scalar(line: str) -> str | None:
    """The unquoted text value of a YAML line (``key: text`` or ``- text``), or None."""
    body = line.strip()
    if body.startswith("- "):
        body = body[2:].lstrip()
        key = re.match(r"^[\w-]+:(?: |$)", body)
        if not key:
            return body if body and not body.startswith(_QUOTED_OR_STRUCTURED) else None
    key = re.match(r"^[\w-]+: (.*)$", body)
    if not key:
        return None
    value = key.group(1).strip()
    return value if value and not value.startswith(_QUOTED_OR_STRUCTURED) else None


def _lint_plain_scalars(kind: str, block: str) -> list[str]:
    """Finds unquoted values that YAML silently changes: ' #' starts a comment (the rest of the
    text is dropped) and ': ' inside a list item turns it into a mapping."""
    problems: list[str] = []
    for number, line in enumerate(block.splitlines(), start=1):
        value = _plain_scalar(line)
        if value is None:
            continue
        if " #" in value:
            problems.append(f"{kind} block line {number}: ' #' starts a YAML comment and cuts the text; "
                            f"quote the value: {line.strip()[:90]}")
        if line.strip().startswith("- ") and ": " in value:
            problems.append(f"{kind} block line {number}: ': ' in a list item makes it a mapping; "
                            f"quote the value: {line.strip()[:90]}")
    return problems


def lint_source(src: str | Path) -> list[str]:
    """Checks a source for the usual YAML slips in ```fr blocks.

    * a rules / validations / fields row with the wrong number of cells (an unquoted comma split a
      cell, or a cell is missing);
    * a one-line flow list ``key: [a, b]`` whose items contain a comma inside parentheses or text
      (YAML splits it into several items);
    * missing mandatory keys (id, title, brd, actor, description, main_flow, acceptance).
    """
    src = Path(src)
    text = src.read_text(encoding="utf-8")
    problems: list[str] = []
    for kind, block in re.findall(r"```(\w+)\n(.*?)```", text, re.S):
        if kind in YAML_BLOCK_KINDS:
            problems.extend(_lint_plain_scalars(kind, block))
        if kind not in ("fr", "requirement"):
            continue
        try:
            fr = yaml.safe_load(block)
        except yaml.YAMLError as exc:
            problems.append(f"YAML error: {exc}".replace("\n", " "))
            continue
        rid = fr.get("id", "?")
        for key in ("id", "title", "brd", "actor", "description", "main_flow", "acceptance"):
            if not fr.get(key):
                problems.append(f"{rid}: missing '{key}'")
        for key, width in FR_ROW_WIDTHS.items():
            for row in fr.get(key) or []:
                if not isinstance(row, list) or len(row) != width:
                    problems.append(f"{rid}: {key} row has {len(row) if isinstance(row, list) else 1} cells, "
                                    f"expected {width}: {row}")
                elif any(isinstance(cell, (dict, list)) for cell in row):
                    problems.append(f"{rid}: {key} row has a cell read as a mapping or list (quote text "
                                    f"that contains ': ' or brackets): {row}")
        for key in ("description", "preconditions", "alternate_flows", "notifications", "audit", "acceptance"):
            value = fr.get(key)
            for item in value if isinstance(value, list) else []:
                if not isinstance(item, str):
                    problems.append(f"{rid}: {key} item read as {type(item).__name__} (quote text with ': '): {item}")
        for line in block.splitlines():
            m = re.match(r"^(\w+): \[(.*)\]$", line.strip())
            if m and m.group(1) != "brd" and "," in m.group(2):
                items = yaml.safe_load(line)[m.group(1)]
                if any(not str(i).rstrip().endswith((".", ")")) for i in items) or "(" in m.group(2):
                    problems.append(f"{rid}: flow list may be split by commas, use a block list: {line.strip()[:90]}")
    return problems


def build_markdown(src: str | Path, out: str | Path | None = None, pdf: bool = True,
                   keep_pdf: bool = False) -> tuple[Path, Path | None]:
    """Builds one Markdown-like source into .docx (and .pdf). Returns the output paths."""
    src = Path(src).resolve()
    front, lines = load_source(src)
    doc = BdoiDocument(meta_from(front), h1_page_break=bool(front.get("h1_page_break", True)),
                       base_dir=src.parent)
    doc.cover()
    doc.front_matter()
    render_body(doc, lines)
    target = Path(out) if out else output_path(front, src)
    return doc.publish(target, pdf=pdf, keep_pdf=keep_pdf)


def main(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="Build BIBS Word documents from Markdown-like sources")
    ap.add_argument("sources", nargs="+", help="source .md files under docs/deliverables/src")
    ap.add_argument("--out", help="output .docx (only with one source)")
    ap.add_argument("--no-pdf", action="store_true", help="skip the PDF pass (and the TOC page numbers)")
    ap.add_argument("--keep-pdf", action="store_true", help="keep the PDF next to the .docx (for issue)")
    ap.add_argument("--previews", action="store_true", help="also render page PNG previews and contact sheets")
    ap.add_argument("--check", action="store_true", help="only lint the YAML blocks (fr, glossary, table, signoff) and exit")
    args = ap.parse_args(argv)
    failed = False
    for s in args.sources:
        for problem in lint_source(s):
            failed = True
            print(f"{s}: {problem}")
    if args.check:
        return 1 if failed else 0
    for s in args.sources:
        keep = args.keep_pdf or args.previews
        docx_path, pdf_path = build_markdown(s, args.out, pdf=not args.no_pdf, keep_pdf=keep)
        print(f"docx: {docx_path}")
        if pdf_path and args.previews:
            import render

            sheets = render.previews(pdf_path)
            print(f"previews: {sheets[0].parent} ({len(sheets)} files)")
        if pdf_path and not args.keep_pdf:
            pdf_path.unlink(missing_ok=True)
        elif pdf_path:
            print(f"pdf:  {pdf_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
