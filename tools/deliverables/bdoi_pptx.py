"""PowerPoint builder for the BIBS client pack (BDO Insure template, 16:9).

Layouts (all drawn on the blank layout, so no template file is needed):

* ``title(title, subtitle)``: BDO Insure logo, title, yellow rule, BIBS line, iorta TechNXT line;
* ``section(title, subtitle)``: Header Blue divider slide;
* ``content(title, bullets)``: title, yellow rule, bullets (a (level, text) tuple for level 2);
* ``two_column(title, left_title, left, right_title, right)``;
* ``screenshot(title, image, caption, persona, action, expected, observation)``: the screen on the
  left, a side panel with persona / action / expected outcome / observation on the right;
* ``table(title, headers, rows, widths)``: Header Blue header row, banded rows.

Every slide except the title and section slides has the small logo top right and a footer with
"Confidential - BDOI", the deck title and the slide number.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import brand  # noqa: E402

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
FONT = brand.FONT


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


class BdoiDeck:
    def __init__(self, title: str, version: str = "1.0", date: str = "", subtitle: str = ""):
        self.title_text = title
        self.version = version
        self.date = date
        self.subtitle = subtitle
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self._blank = self.prs.slide_layouts[6]
        self._n = 0

    # ------------------------------------------------------------------ primitives

    def _slide(self):
        self._n += 1
        return self.prs.slides.add_slide(self._blank)

    def _text(self, slide, x, y, w, h, text, size=18, bold=False, colour=brand.NEAR_BLACK, align=None,
              anchor=MSO_ANCHOR.TOP, italic=False):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text
        f = run.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = _rgb(colour)
        if align:
            p.alignment = align
        return box

    def _rect(self, slide, x, y, w, h, fill, line=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        if line:
            shape.line.color.rgb = _rgb(line)
            shape.line.width = Pt(0.75)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def _bullets(self, slide, x, y, w, h, items, size=16):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for item in items:
            level, text = item if isinstance(item, tuple) else (0, item)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(6)
            marker = p.add_run()
            marker.text = ("•  " if level == 0 else "–  ")
            marker.font.color.rgb = _rgb(brand.HEADER_BLUE if level == 0 else brand.MUTED)
            marker.font.size = Pt(size - 2 * level)
            marker.font.name = FONT
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size - 2 * level)
            r.font.name = FONT
            r.font.color.rgb = _rgb(brand.TEXT)
            if level:
                p.level = 1
        return box

    def _chrome(self, slide, title: str) -> None:
        """Title, yellow rule, small logo and footer of a content slide."""
        self._text(slide, MARGIN, Inches(0.35), Inches(10.4), Inches(0.8), title, size=26, bold=True,
                   colour=brand.HEADER_BLUE, anchor=MSO_ANCHOR.MIDDLE)
        self._rect(slide, MARGIN, Inches(1.2), W - 2 * MARGIN, Pt(2), brand.YELLOW)
        if brand.BDO_LOGO.exists():
            slide.shapes.add_picture(str(brand.BDO_LOGO), W - MARGIN - Inches(1.8), Inches(0.5), width=Inches(1.8))
        self._rect(slide, MARGIN, H - Inches(0.55), W - 2 * MARGIN, Pt(0.75), brand.BORDER)
        self._text(slide, MARGIN, H - Inches(0.5), Inches(4), Inches(0.35), brand.FOOTER_TEXT, size=9,
                   colour=brand.MUTED)
        self._text(slide, Inches(4.7), H - Inches(0.5), Inches(4), Inches(0.35),
                   f"{self.title_text} | v{self.version}", size=9, colour=brand.MUTED, align=PP_ALIGN.CENTER)
        self._text(slide, W - MARGIN - Inches(1.5), H - Inches(0.5), Inches(1.5), Inches(0.35), str(self._n),
                   size=9, colour=brand.MUTED, align=PP_ALIGN.RIGHT)

    # ------------------------------------------------------------------ layouts

    def title(self, title: str | None = None, subtitle: str | None = None) -> None:
        s = self._slide()
        if brand.BDO_LOGO.exists():
            s.shapes.add_picture(str(brand.BDO_LOGO), MARGIN, Inches(0.6), width=Inches(3.2))
        self._text(s, MARGIN, Inches(2.4), Inches(11), Inches(1.4), title or self.title_text, size=40, bold=True,
                   colour=brand.HEADER_BLUE, anchor=MSO_ANCHOR.BOTTOM)
        self._rect(s, MARGIN, Inches(3.95), Inches(2.2), Pt(4), brand.YELLOW)
        self._text(s, MARGIN, Inches(4.15), Inches(11), Inches(0.6), subtitle or self.subtitle or "", size=20,
                   colour=brand.NEAR_BLACK)
        self._text(s, MARGIN, Inches(4.75), Inches(11), Inches(0.5), brand.SYSTEM, size=16, colour=brand.MUTED)
        meta = " | ".join(x for x in (brand.CLIENT, f"Version {self.version}", self.date, brand.CLASSIFICATION) if x)
        self._text(s, MARGIN, H - Inches(1.0), Inches(9), Inches(0.4), meta, size=11, colour=brand.MUTED)
        self._text(s, W - MARGIN - Inches(3.6), H - Inches(1.05), Inches(2.0), Inches(0.4), "Prepared by",
                   size=10, colour=brand.MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        if brand.IORTA_LOGO.exists():
            s.shapes.add_picture(str(brand.IORTA_LOGO), W - MARGIN - Inches(1.5), H - Inches(1.15), width=Inches(1.5))

    def section(self, title: str, subtitle: str = "") -> None:
        s = self._slide()
        self._rect(s, 0, 0, W, H, brand.HEADER_BLUE)
        self._text(s, MARGIN, Inches(2.6), Inches(11.5), Inches(1.3), title, size=38, bold=True,
                   colour=brand.WHITE, anchor=MSO_ANCHOR.BOTTOM)
        self._rect(s, MARGIN, Inches(4.05), Inches(1.6), Pt(4), brand.YELLOW)
        if subtitle:
            self._text(s, MARGIN, Inches(4.25), Inches(11.5), Inches(0.8), subtitle, size=18, colour=brand.BG_BLUE)
        self._text(s, MARGIN, H - Inches(0.6), Inches(6), Inches(0.35), f"{brand.FOOTER_TEXT} | {self.title_text}",
                   size=9, colour=brand.FIELD_BLUE)

    def content(self, title: str, bullets: Sequence[str | tuple[int, str]], size: int = 18) -> None:
        s = self._slide()
        self._chrome(s, title)
        self._bullets(s, MARGIN, Inches(1.5), W - 2 * MARGIN, H - Inches(2.3), bullets, size=size)

    def two_column(self, title: str, left_title: str, left: Sequence, right_title: str, right: Sequence) -> None:
        s = self._slide()
        self._chrome(s, title)
        col_w = (W - 2 * MARGIN - Inches(0.5)) / 2
        for i, (head, items) in enumerate(((left_title, left), (right_title, right))):
            x = MARGIN + i * (col_w + Inches(0.5))
            self._rect(s, x, Inches(1.5), col_w, Inches(0.5), brand.BG_BLUE)
            self._text(s, x + Inches(0.15), Inches(1.5), col_w, Inches(0.5), head, size=16, bold=True,
                       colour=brand.HEADER_BLUE, anchor=MSO_ANCHOR.MIDDLE)
            self._bullets(s, x, Inches(2.15), col_w, H - Inches(3.0), items, size=15)

    def screenshot(self, title: str, image: str | Path, caption: str = "", persona: str = "", action: str = "",
                   expected: str = "", observation: str = "") -> None:
        s = self._slide()
        self._chrome(s, title)
        area_w, area_h = Inches(8.3), Inches(4.9)
        x0, y0 = MARGIN, Inches(1.45)
        image = Path(image)
        if image.exists():
            from PIL import Image

            with Image.open(image) as im:
                iw, ih = im.size
            scale = min(area_w / iw, area_h / ih)
            w, h = int(iw * scale), int(ih * scale)
            pic = s.shapes.add_picture(str(image), x0 + (area_w - w) // 2, y0, width=w, height=h)
            pic.line.color.rgb = _rgb(brand.BORDER)
            pic.line.width = Pt(0.75)
        else:
            self._rect(s, x0, y0, area_w, area_h, brand.DIRTY_WHITE, brand.BORDER)
            self._text(s, x0, y0, area_w, area_h, f"Screenshot missing: {image.name}", size=14, colour=brand.MUTED,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if caption:
            self._text(s, x0, y0 + area_h + Inches(0.05), area_w, Inches(0.4), caption, size=11, italic=True,
                       colour=brand.MUTED, align=PP_ALIGN.CENTER)
        px = x0 + area_w + Inches(0.3)
        pw = W - MARGIN - px
        self._rect(s, px, y0, pw, area_h + Inches(0.45), brand.BG_BLUE)
        self._rect(s, px, y0, Pt(4), area_h + Inches(0.45), brand.HEADER_BLUE)
        y = y0 + Inches(0.1)
        block_h = (area_h + Inches(0.3)) / 4
        for label, text in (("Persona", persona), ("Action", action), ("Expected outcome", expected),
                            ("Observation", observation)):
            self._text(s, px + Inches(0.2), y, pw - Inches(0.3), Inches(0.35), label.upper(), size=10, bold=True,
                       colour=brand.HEADER_BLUE)
            self._text(s, px + Inches(0.2), y + Inches(0.3), pw - Inches(0.3), block_h - Inches(0.35), text or "-",
                       size=12, colour=brand.NEAR_BLACK)
            y += block_h

    def table(self, title: str, headers: Sequence[str], rows: Sequence[Sequence], widths: Sequence[float] | None = None,
              size: int = 12) -> None:
        s = self._slide()
        self._chrome(s, title)
        n_rows, n_cols = len(rows) + 1, len(headers)
        total_w = W - 2 * MARGIN
        row_h = Inches(0.4)
        shape = s.shapes.add_table(n_rows, n_cols, MARGIN, Inches(1.5), total_w, row_h * n_rows)
        tbl = shape.table
        widths = widths or [1] * n_cols
        scale = total_w / sum(widths)
        for i, w in enumerate(widths):
            tbl.columns[i].width = Emu(int(w * scale))
        for r in range(n_rows):
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                value = headers[c] if r == 0 else (rows[r - 1][c] if c < len(rows[r - 1]) else "")
                cell.text = str(value)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                para = cell.text_frame.paragraphs[0]
                for run in para.runs:
                    run.font.name = FONT
                    run.font.size = Pt(size)
                    run.font.bold = r == 0
                    run.font.color.rgb = _rgb(brand.WHITE if r == 0 else brand.TEXT)
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = _rgb(brand.HEADER_BLUE)
                else:
                    key = str(value).upper()
                    if key in brand.STATUS_COLOURS:
                        fill, font = brand.STATUS_COLOURS[key]
                        cell.fill.fore_color.rgb = _rgb(fill)
                        for run in para.runs:
                            run.font.color.rgb = _rgb(font)
                            run.font.bold = True
                    else:
                        cell.fill.fore_color.rgb = _rgb(brand.BG_BLUE if r % 2 == 0 else brand.WHITE)

    def save(self, path: str | Path) -> Path:
        props = self.prs.core_properties
        props.title = self.title_text
        props.author = brand.VENDOR
        props.subject = brand.SYSTEM
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def self_test(path: str | Path, screenshot: str | Path | None = None) -> Path:
    """A sample deck with every layout (used by the README and the smoke test)."""
    deck = BdoiDeck("Product Maintenance walkthrough", version="1.0", date="25 September 2026",
                    subtitle="BRD-3 Product Maintenance: end-to-end by persona")
    deck.title()
    deck.section("1. Package request", "Marketing AO, Marketing TL, TSU TL, TSU Head")
    deck.content("What this walkthrough covers", [
        "A new package request from the Marketing AO to release",
        (1, "Two negotiation rounds with two insurers"),
        "ManCom sign-off, MBS set-up and validation by the TSU Head",
        "Release advisory to Marketing, TSU, MBS and Operations"])
    deck.two_column("Roles and screens", "Personas", ["Marketing AO", "TSU Officer", "MBS"], "Screens",
                    ["Package Requests", "Package Request page", "Version editor"])
    deck.screenshot("Submit the Package Request Form", screenshot or "missing.png",
                    caption="New Package Request, step 1", persona="Marketing AO (ao)",
                    action="Enters the programme name, line, cover type, requested terms and target insurers; "
                           "clicks Submit for Approval.",
                    expected="The request moves to For Marketing approval; the Marketing TL is notified.",
                    observation="As expected.")
    deck.table("Stage SLA (defaults)", ["Stage", "Owner", "SLA hours", "Status"],
               [["For Marketing approval", "Marketing TL / TH / UH", "24", "PASS"],
                ["For TSU review", "TSU TL", "24", "PASS"],
                ["Negotiation", "TSU Officer", "120", "OPEN"]], widths=[3, 3, 1.4, 1.4])
    return deck.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "deck-sample.pptx")
    print(self_test(out, sys.argv[2] if len(sys.argv) > 2 else None))
