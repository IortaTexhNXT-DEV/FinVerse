"""Excel builder for the BIBS client pack (BDO Insure template).

A workbook has:

* a **Cover** sheet: BDO Insure logo, title, BIBS line, client, version, date, classification, a
  sheet index and a colour legend for the status values;
* a **README** sheet that describes every column of every data sheet (generated from the column
  definitions, including the allowed values);
* **data sheets**: a title band, a Header Blue header row, frozen panes, auto-filter, set column
  widths with wrapped text, drop-down lists for status / fit / priority columns, conditional fills
  for status values, banded rows, and print set-up (landscape, fit to width, header row repeated,
  "Confidential - BDOI" footer with page x of y and the version).

Example::

    wb = BdoiWorkbook("BRD Discrepancy Register", doc_type="Register", brd="BRD-03", version="1.0")
    wb.sheet("Register", [
        Column("id", "ID", 10, "Register row identifier"),
        Column("status", "Status", 14, "Resolution status", values=["OPEN", "ANSWERED", "CLOSED"], status=True),
    ], rows=[{"id": "PQ01", "status": "OPEN"}])
    wb.save("docs/deliverables/out/Registers/BIBS_Register_BRD-03_Discrepancies_v1.0.xlsx")
"""

from __future__ import annotations

import datetime as _dt
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

sys.path.insert(0, str(Path(__file__).resolve().parent))
import brand  # noqa: E402

FONT = brand.FONT
HEADER_ROW = 4  # title band in rows 1-3, header in row 4, data from row 5
THIN = Side(style="thin", color=brand.BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=brand.HEADER_BLUE)
BAND_FILL = PatternFill("solid", fgColor=brand.BG_BLUE)
LABEL_FILL = PatternFill("solid", fgColor=brand.BG_BLUE)


def _font(size: float = 10, bold: bool = False, colour: str = brand.TEXT, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=colour, italic=italic)


@dataclass
class Column:
    """A data column: key in the row dicts, header text, width (characters) and README text.

    values: allowed values (drop-down list); status: colour the values with brand.STATUS_COLOURS;
    kind: text | number | percent | date | amount.
    """

    key: str
    header: str
    width: float = 18
    description: str = ""
    values: Sequence[str] | None = None
    status: bool = False
    kind: str = "text"
    wrap: bool = True


@dataclass
class _SheetInfo:
    name: str
    description: str
    columns: list[Column]
    rows: int


class BdoiWorkbook:
    def __init__(self, title: str, doc_type: str = "", brd: str = "", version: str = "1.0",
                 date: str | None = None, subtitle: str = "", classification: str = brand.CLASSIFICATION):
        self.title = title
        self.doc_type = doc_type
        self.brd = brd
        self.version = version
        self.date = date or _dt.date.today().strftime("%d %B %Y").lstrip("0")
        self.subtitle = subtitle
        self.classification = classification
        self.wb = Workbook()
        self._cover = self.wb.active
        self._cover.title = "Cover"
        self._readme = self.wb.create_sheet("README")
        self._lists = self.wb.create_sheet("_lists")
        self._lists.sheet_state = "hidden"
        self._list_col = 0
        self._sheets: list[_SheetInfo] = []
        self.legend: list[tuple[str, str]] = []
        self.cover_notes: list[str] = []

    # ------------------------------------------------------------------ data sheets

    def sheet(self, name: str, columns: Sequence[Column], rows: Sequence[dict[str, Any] | Sequence[Any]],
              description: str = "", freeze_first_column: bool = True):
        """Adds a data sheet. rows are dicts keyed by Column.key, or sequences in column order."""
        ws = self.wb.create_sheet(name[:31])
        cols = list(columns)
        last_col = get_column_letter(len(cols))
        # Title band
        ws.row_dimensions[1].height = 24
        ws["A1"] = name
        ws["A1"].font = _font(14, True, brand.HEADER_BLUE)
        ws["A2"] = description or f"{self.title} – {brand.SYSTEM}"
        ws["A2"].font = _font(9, False, brand.MUTED)
        ws.merge_cells(f"A2:{last_col}2")
        ws["A2"].alignment = Alignment(wrap_text=False, vertical="top")
        for c in range(1, len(cols) + 1):
            ws.cell(row=3, column=c).border = Border(bottom=Side(style="thin", color=brand.YELLOW))
        ws.row_dimensions[3].height = 6
        # Header
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=HEADER_ROW, column=c, value=col.header)
            cell.font = _font(10, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(c)].width = col.width
        ws.row_dimensions[HEADER_ROW].height = 30
        # Rows
        for r, row in enumerate(rows, start=HEADER_ROW + 1):
            for c, col in enumerate(cols, start=1):
                value = row.get(col.key) if isinstance(row, dict) else (row[c - 1] if c - 1 < len(row) else None)
                if isinstance(value, (list, tuple)):
                    value = "\n".join(f"• {v}" for v in value)
                cell = ws.cell(row=r, column=c, value=value)
                cell.font = _font(10, col.status and value is not None, brand.TEXT)
                cell.alignment = Alignment(wrap_text=col.wrap, vertical="top",
                                           horizontal="center" if col.status else None)
                cell.border = BORDER
                if col.kind == "percent":
                    cell.number_format = "0.00"
                elif col.kind == "amount":
                    cell.number_format = "#,##0.00"
                elif col.kind == "date":
                    cell.number_format = "dd-mmm-yyyy"
        n = len(rows)
        first, last = HEADER_ROW + 1, HEADER_ROW + max(n, 1)
        # Room for new rows: validations and formats extend 500 rows past the data.
        ext_last = last + 500
        data_ref = f"A{first}:{last_col}{ext_last}"
        # Status fills first (higher priority), then banding.
        for c, col in enumerate(cols, start=1):
            letter = get_column_letter(c)
            rng = f"{letter}{first}:{letter}{ext_last}"
            if col.values:
                dv = DataValidation(type="list", formula1=self._list_ref(col.values), allow_blank=True,
                                    showErrorMessage=True, errorTitle=col.header,
                                    error=f"Choose a value from the list ({col.header}).")
                ws.add_data_validation(dv)
                dv.add(rng)
            if col.status:
                values = list(col.values or []) or list(brand.STATUS_COLOURS)
                for v in values:
                    colours = brand.STATUS_COLOURS.get(str(v).upper())
                    if not colours:
                        continue
                    fill, font = colours
                    ws.conditional_formatting.add(rng, CellIsRule(
                        operator="equal", formula=[f'"{v}"'], stopIfTrue=True,
                        fill=PatternFill("solid", fgColor=fill, bgColor=fill),
                        font=Font(color=font, bold=True)))
        ws.conditional_formatting.add(data_ref, FormulaRule(
            formula=[f"AND(MOD(ROW(),2)=0,$A{first}<>\"\")"], fill=PatternFill("solid", fgColor=brand.BG_BLUE,
                                                                                 bgColor=brand.BG_BLUE)))
        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2 if freeze_first_column else 1)
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last}"
        self._print_setup(ws, f"{HEADER_ROW}:{HEADER_ROW}")
        ws.sheet_view.zoomScale = 90
        self._sheets.append(_SheetInfo(ws.title, description, cols, n))
        return ws

    def _list_ref(self, values: Sequence[str]) -> str:
        self._list_col += 1
        letter = get_column_letter(self._list_col)
        for i, v in enumerate(values, start=1):
            self._lists.cell(row=i, column=self._list_col, value=v)
        return f"'_lists'!${letter}$1:${letter}${len(values)}"

    def _print_setup(self, ws, title_rows: str | None) -> None:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = 0.6
        ws.page_margins.bottom = 0.6
        if title_rows:
            ws.print_title_rows = title_rows
        ws.oddHeader.left.text = f"{self.title}"
        ws.oddHeader.left.size = 8
        ws.oddHeader.left.color = brand.HEADER_BLUE
        ws.oddHeader.right.text = "BDO Insure | " + brand.SYSTEM
        ws.oddHeader.right.size = 8
        ws.oddHeader.right.color = brand.MUTED
        ws.oddFooter.left.text = brand.FOOTER_TEXT
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.right.text = f"Version {self.version}"
        for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
            part.size = 8
            part.color = brand.MUTED

    # ------------------------------------------------------------------ cover and README

    def _build_cover(self) -> None:
        ws = self._cover
        ws.sheet_view.showGridLines = False
        widths = {"A": 3, "B": 24, "C": 60, "D": 16, "E": 16}
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        if brand.BDO_LOGO.exists():
            img = XlImage(str(brand.BDO_LOGO))
            img.width, img.height = 227, 46
            ws.add_image(img, "B2")
        ws.row_dimensions[2].height = 40
        for c in "BCDE":
            ws[f"{c}4"].border = Border(bottom=Side(style="medium", color=brand.YELLOW))
        ws.row_dimensions[4].height = 6
        row = 6
        if self.doc_type:
            ws[f"B{row}"] = self.doc_type.upper()
            ws[f"B{row}"].font = _font(10, True, brand.CTA_BLUE)
            row += 1
        ws[f"B{row}"] = self.title
        ws[f"B{row}"].font = _font(22, True, brand.HEADER_BLUE)
        ws.row_dimensions[row].height = 32
        row += 1
        if self.subtitle:
            ws[f"B{row}"] = self.subtitle
            ws[f"B{row}"].font = _font(12, False, brand.NEAR_BLACK)
            row += 1
        ws[f"B{row}"] = brand.SYSTEM
        ws[f"B{row}"].font = _font(11, False, brand.MUTED)
        row += 2
        facts = [("Client", brand.CLIENT), ("BRD", self.brd), ("Version", self.version), ("Date", self.date),
                 ("Classification", self.classification), ("Prepared by", brand.VENDOR)]
        for k, v in facts:
            if not v:
                continue
            ws[f"B{row}"] = k
            ws[f"B{row}"].font = _font(10, False, brand.MUTED)
            ws[f"C{row}"] = v
            ws[f"C{row}"].font = _font(10, k == "Classification", brand.NEAR_BLACK)
            for c in "BC":
                ws[f"{c}{row}"].border = Border(bottom=Side(style="thin", color=brand.BORDER_LIGHT))
            row += 1
        row += 1
        ws[f"B{row}"] = "Sheets"
        ws[f"B{row}"].font = _font(12, True, brand.HEADER_BLUE)
        row += 1
        for c, h in zip("BCD", ("Sheet", "Content", "Rows")):
            cell = ws[f"{c}{row}"]
            cell.value = h
            cell.font = _font(10, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.border = BORDER
        row += 1
        index = [("README", "What each column of each sheet means, with the allowed values", "")]
        index += [(s.name, s.description or s.name, s.rows) for s in self._sheets]
        for i, (name, desc, n) in enumerate(index):
            for c, v in zip("BCD", (name, desc, n)):
                cell = ws[f"{c}{row}"]
                cell.value = v
                cell.font = _font(10, c == "B", brand.TEXT)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2 == 1:
                    cell.fill = BAND_FILL
            ws[f"B{row}"].hyperlink = f"#'{name}'!A1"
            ws[f"B{row}"].font = _font(10, True, brand.CTA_BLUE)
            row += 1
        legend = self.legend or self._default_legend()
        if legend:
            row += 1
            ws[f"B{row}"] = "Legend"
            ws[f"B{row}"].font = _font(12, True, brand.HEADER_BLUE)
            row += 1
            for value, meaning in legend:
                fill, font = brand.STATUS_COLOURS.get(value.upper(), (brand.WHITE, brand.TEXT))
                ws[f"B{row}"] = value
                ws[f"B{row}"].fill = PatternFill("solid", fgColor=fill)
                ws[f"B{row}"].font = Font(name=FONT, size=10, bold=True, color=font)
                ws[f"B{row}"].alignment = Alignment(horizontal="center")
                ws[f"C{row}"] = meaning
                ws[f"C{row}"].font = _font(10)
                ws[f"C{row}"].alignment = Alignment(wrap_text=True)
                for c in "BC":
                    ws[f"{c}{row}"].border = BORDER
                row += 1
        for note in self.cover_notes:
            row += 1
            ws[f"B{row}"] = note
            ws[f"B{row}"].font = _font(9, False, brand.MUTED, italic=True)
        row += 2
        ws[f"B{row}"] = f"Prepared by {brand.VENDOR} for {brand.CLIENT_SHORT}. {brand.FOOTER_TEXT}."
        ws[f"B{row}"].font = _font(9, False, brand.MUTED)
        if brand.IORTA_LOGO.exists():
            img = XlImage(str(brand.IORTA_LOGO))
            img.width, img.height = 102, 30
            ws.add_image(img, f"D{row}")
        ws.row_dimensions[row].height = 26
        self._print_setup(ws, None)
        ws.page_setup.orientation = "portrait"

    def _default_legend(self) -> list[tuple[str, str]]:
        seen: list[str] = []
        for s in self._sheets:
            for col in s.columns:
                if col.status:
                    for v in col.values or []:
                        if v.upper() in brand.STATUS_COLOURS and v not in seen:
                            seen.append(v)
        return [(v, "") for v in seen]

    def _build_readme(self) -> None:
        ws = self._readme
        ws.sheet_view.showGridLines = False
        ws["A1"] = "README – columns of each sheet"
        ws["A1"].font = _font(14, True, brand.HEADER_BLUE)
        ws["A2"] = f"{self.title}, version {self.version}. Drop-down columns accept only the listed values."
        ws["A2"].font = _font(9, False, brand.MUTED)
        headers = ["Sheet", "Column", "Meaning", "Allowed values"]
        widths = [22, 26, 70, 40]
        for c, (h, w) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=HEADER_ROW, column=c, value=h)
            cell.font = _font(10, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(c)].width = w
        r = HEADER_ROW + 1
        for s in self._sheets:
            for col in s.columns:
                values = ", ".join(col.values) if col.values else ""
                for c, v in enumerate((s.name, col.header, col.description, values), start=1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.font = _font(10, c == 2)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.border = BORDER
                    if (r - HEADER_ROW) % 2 == 0:
                        cell.fill = BAND_FILL
                r += 1
        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
        ws.auto_filter.ref = f"A{HEADER_ROW}:D{max(r - 1, HEADER_ROW)}"
        self._print_setup(ws, f"{HEADER_ROW}:{HEADER_ROW}")

    def save(self, path: str | Path) -> Path:
        self._build_cover()
        self._build_readme()
        # Order: Cover, README, data sheets; the hidden list sheet last.
        order = [self._cover, self._readme] + [self.wb[s.name] for s in self._sheets] + [self._lists]
        self.wb._sheets = order
        self.wb.active = 0
        props = self.wb.properties
        props.title = self.title
        props.creator = brand.VENDOR
        props.subject = brand.SYSTEM
        props.keywords = ", ".join(x for x in (self.doc_type, self.brd, "BIBS", "BDOI") if x)
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        return path


def demo(path: str | Path) -> Path:
    """A small sample workbook that exercises every feature (used by the README and smoke test)."""
    wb = BdoiWorkbook("Fit-Gap Sample", doc_type="Fit-gap workbook", brd="BRD-03", version="1.0",
                      subtitle="BRD-3 Product Maintenance")
    wb.legend = [("FIT", "Works today"), ("CONFIGURE", "Set-up only"), ("CHANGE", "Extends a capability"),
                 ("NEW", "New capability"), ("OUT", "Out of scope")]
    cols = [
        Column("id", "BRD ID", 12, "Requirement ID in the BRD, with the page"),
        Column("req", "Requirement", 48, "Requirement text, shortened"),
        Column("fit", "Fit", 12, "Fit class", values=["FIT", "CONFIGURE", "CHANGE", "NEW", "OUT"], status=True),
        Column("priority", "Priority", 12, "BRD priority", values=["Must have", "Should have", "Could have"]),
        Column("module", "Module", 16, "Target module"),
    ]
    rows = [
        {"id": "BRPM.001", "req": "Access and log in to the system", "fit": "FIT", "priority": "Must have",
         "module": "platform"},
        {"id": "BRPM.007", "req": "New business uses the latest approved rate scheme", "fit": "CHANGE",
         "priority": "Must have", "module": "catalog"},
        {"id": "BRPM.008", "req": "Submit an approved Package Request Form", "fit": "NEW", "priority": "Must have",
         "module": "productmaint"},
    ]
    wb.sheet("Fit-Gap", cols, rows, description="One row per BRD requirement")
    return wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "fitgap-sample.xlsx")
    print(demo(out))
