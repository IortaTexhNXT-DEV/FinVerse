"""Drop closure set of Drop 0 (Setup and Data Migration): the configuration inputs workbook and the closure summary.

    python tools/deliverables/drop_closure.py            # 01 workbook and 02 summary (Word)
    python tools/deliverables/drop_closure.py --check    # checks only
    python tools/deliverables/drop_closure.py --xlsx     # the workbook only

Sources (docs/deliverables/src/Drop-0_Closure/):
  * drop0.yaml: the register items, the Drop 0 templates D0-nn, the owners, due milestones, decisions, criteria,
    risks and dependencies;
  * CLOSURE_SUMMARY_DROP0.md: the Word source of the closure summary; its tables are ```pack blocks rendered by
    render() of this module.
Read, not typed:
  * the configuration input templates of BRD-03 and BRD-11 (pack/config_inputs.yaml of each BRD, the source of their
    06 workbooks), through src/signoff/config_inputs.py;
  * the migration objects, milestones, proposals and open decisions of BRD-13 (pack/catalogue.yaml);
  * the clarification chapters and open questions of the FRS of BRD-03 and BRD-11 and of the BRD-13 handbook;
  * from the platform: system parameters, lists of values, notification events, document templates, accounting
    events and code map sets of the migration scripts (sql_facts.py), the master tables the SIT seed data fills,
    the screens of the menu and the chart of accounts upload template (code_facts.py).

Outputs, in the drop folder Drop-0_Closure/ (brand.drop_set_dir):
  01_BIBS_Drop-0_Configuration_Inputs_v2.0.xlsx and 02_BIBS_Drop-0_Closure_Summary_v2.0.docx.
"""

from __future__ import annotations

import argparse
import datetime as dt
import importlib.util
import re
import sys
from functools import lru_cache
from pathlib import Path
from typing import Any

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import brand  # noqa: E402
import code_facts  # noqa: E402
import sql_facts  # noqa: E402

DROP = "Drop 0"
SRC = brand.SRC_DIR / "Drop-0_Closure"
SOURCE = SRC / "drop0.yaml"
SUMMARY = SRC / "CLOSURE_SUMMARY_DROP0.md"
ROUTES = {"screen": "Screen", "template": "Template", "migration": "Migration object"}
BRD_TEMPLATE_PREFIX = {"BRD-03": "PM", "BRD-11": "UA"}
LISTS = {
    "LIST-LOV": "Ref Lists of values",
    "LIST-PARAMETERS": "Ref System parameters",
    "LIST-NUMBERING": "Ref Document numbering",
    "LIST-NOTIFICATIONS": "Ref Notifications",
    "LIST-DOCUMENTS": "Ref Document templates",
    "LIST-EVENTS": "Ref Accounting events",
}
MIGRATION_SHEET = "Ref Migration cross-reference"


# ============================================================================ loading


def _module(path: Path, name: str):
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


@lru_cache(maxsize=1)
def data() -> dict[str, Any]:
    return yaml.safe_load(SOURCE.read_text(encoding="utf-8"))


def config_inputs():
    return _module(brand.SRC_DIR / "signoff" / "config_inputs.py", "signoff_config_inputs")


@lru_cache(maxsize=1)
def brd_templates() -> dict[str, dict[str, Any]]:
    """The configuration input templates of the BRD-03 and BRD-11 sets, keyed PM-nn / UA-nn; the checks and depends
    name the templates by these IDs."""
    out: dict[str, dict[str, Any]] = {}
    for brd, prefix in BRD_TEMPLATE_PREFIX.items():
        src = brand.src_dir(brd) / "pack" / "config_inputs.yaml"
        data = config_inputs().load(src)
        sid_of = lambda tid, prefix=prefix: f"{prefix}-{tid.split('-')[-1]}"  # noqa: E731
        for t in data["templates"]:
            sid = sid_of(t["id"])
            rows = []
            for r in t["rows"]:
                kind, _, arg = r["check"].partition(":")
                if kind in ("code", "ref"):
                    ref, _, col = arg.partition("/")
                    r = {**r, "check": f"{kind}:{sid_of(ref)}/{col}"}
                rows.append(r)
            out[sid] = {**t, "id": sid, "sid": sid, "brd": brd, "ref": f"{brd} template {t['id']}", "rows": rows,
                        "depends": [sid_of(d) for d in t.get("depends") or []], "lists_source": data["lists"]}
    return out


@lru_cache(maxsize=1)
def drop_templates() -> dict[str, dict[str, Any]]:
    """The Drop 0 templates D0-nn, with their rows (the upload columns for an upload template)."""
    out: dict[str, dict[str, Any]] = {}
    for t in data()["templates"]:
        t = dict(t)
        checks = t.get("checks") or {}
        if t.get("upload"):
            tpl = code_facts.bulk_templates([t["upload"]])[0]
            t["rows"] = [{"no": i, "header": c["header"], "mandatory": "Y" if c["required"] else "N",
                          "format": {"YES_NO": "Y or N", "DATE": "yyyy-MM-dd", "NUMBER": "Number"}.get(c["type"], "Text"),
                          "what": c["description"], "example": c.get("example") or "-"}
                         for i, c in enumerate(tpl.columns, start=1)]
            t["content"] = f"{t['content']}. {tpl.instructions}".strip()
        else:
            t["rows"] = []
            for i, raw in enumerate(t["columns"], start=1):
                parts = [p.strip() for p in str(raw).split(" | ")]
                if len(parts) != 5:
                    raise ValueError(f"{t['id']} column {i}: {len(parts)} cells, expected 5: {raw}")
                t["rows"].append(dict(zip(["header", "mandatory", "format", "what", "example"], parts), no=i))
        for r in t["rows"]:
            r["check"] = str(checks.get(r["header"], ""))
        t["sid"] = t["id"]
        t["ref"] = f"Drop 0 template {t['id']}"
        out[t["id"]] = t
    return out


def combined() -> dict[str, Any]:
    """Every template of the workbook with the lists of the three sources, in the form of a templates file."""
    lists: dict[str, Any] = {}
    for t in brd_templates().values():
        prefix = BRD_TEMPLATE_PREFIX[t["brd"]]
        for code, spec in t["lists_source"].items():
            if spec.get("parameters_from_column"):
                tid, n = spec["parameters_from_column"]
                spec = {**spec, "parameters_from_column": [f"{prefix}-{tid.split('-')[-1]}", n]}
            lists[code] = spec
    lists.update(data().get("lists") or {})
    return {"templates": list(all_templates().values()), "lists": lists}


def all_templates() -> dict[str, dict[str, Any]]:
    return {**brd_templates(), **drop_templates()}


@lru_cache(maxsize=1)
def catalogue() -> dict[str, Any]:
    return yaml.safe_load((brand.src_dir("BRD-13") / "pack" / "catalogue.yaml").read_text(encoding="utf-8"))


@lru_cache(maxsize=1)
def menu() -> dict[str, str]:
    """Menu path text of every screen route of the sidebar."""
    out = {}
    for g, s, sc in code_facts.all_screens():
        out.setdefault(sc.path, code_facts.menu_path(g, s, sc))
    return out


def sheet_name(t: dict[str, Any]) -> str:
    return f"{t['sid']} {t.get('sheet', t['name'])}"[:31]


def go_live() -> dt.date:
    v = data()["meta"]["go_live"]
    return v if isinstance(v, dt.date) else dt.date.fromisoformat(str(v))


@lru_cache(maxsize=1)
def dues() -> dict[str, dict[str, Any]]:
    out = {}
    for d in data()["due"]:
        date = d["date"] if isinstance(d["date"], dt.date) else dt.date.fromisoformat(str(d["date"]))
        weeks = (go_live() - date).days // 7
        out[d["code"]] = {**d, "date": date, "weeks": weeks, "text": f"{fmt(date)} (T-{weeks} weeks)"}
    return out


def fmt(date: dt.date) -> str:
    return date.strftime("%d-%b-%Y").lstrip("0")


def module(code: str) -> dict[str, str]:
    return data()["modules"][code]


def _clean(text: str | None) -> str:
    """Description text for the client: without the references to requirement, question and design sections."""
    text = (text or "").strip()
    ref = (r"\b(?:BRD|NFR|[A-Z]{1,3}Q\d|EBQ|CLQ|DMQ|DCR|p\.\s?\d|FR-|BR[A-Z]+\.|UX-\d|design|spec|section|seam|stub|"
           r"R\d\b|Q\d)|\b\d+\.\d+(?:\.\d+|\.x)?\b|\b[A-Z]{2,}[- .]?\d|@[A-Z]+")
    text = re.sub(rf"\s*\((?:[^()]|\([^()]*\))*?(?:{ref})(?:[^()]|\([^()]*\))*\)", "", text)
    text = re.sub(r",?\s*\bBR[A-Z]+\.\d+\b", "", text)
    text = text.replace("BrokerVerse", "BIBS")
    text = re.sub(r"\b[a-z]+(?:_[a-z0-9]+)+\b", lambda m: m.group(0).replace("_", " "), text)
    return re.sub(r"\s{2,}", " ", text).strip()


# ============================================================================ generated lists


def _module_of_lov(code: str) -> str:
    d = data()
    if code in d["lov_codes"]:
        return d["lov_codes"][code]
    for prefix in sorted(d["lov_prefixes"], key=len, reverse=True):
        if code.startswith(prefix):
            return d["lov_prefixes"][prefix]
    raise KeyError(f"list of values {code}: no module in lov_prefixes or lov_codes")


def _module_of_doc(code: str) -> str:
    d = data()["document_prefixes"]
    for prefix in sorted(d, key=len, reverse=True):
        if code.startswith(prefix):
            return d[prefix]
    raise KeyError(f"document template {code}: no module in document_prefixes")


@lru_cache(maxsize=1)
def migrated_lovs() -> dict[str, str]:
    """List code -> code map set, for the lists whose values the data migration loads (object R01)."""
    out = {}
    for m in sql_facts.rows("mig_code_map_set"):
        used = [u.strip() for u in (m.get("used_by") or "").split(",")]
        if m.get("target_kind") == "LOV" and "R01" in used and m.get("target_ref"):
            out[m["target_ref"]] = m["code"]
    return out


def _template_lists() -> dict[str, dict[str, str]]:
    """parameter / list / document code -> template id from list_templates."""
    lt = data()["list_templates"]
    out: dict[str, dict[str, str]] = {"parameters": {}, "lovs": {}, "documents": {}}
    params = sql_facts.rows("sys_parameter", "param_key")
    for spec in lt["parameters"]:
        if "keys_from_column" in spec:
            t = all_templates()[spec["template"]]
            first = t["rows"][spec["keys_from_column"] - 1]["format"]
            keys = [k.strip() for k in first.split(",")]
        else:
            keys = [p["param_key"] for p in params if p["category"] in spec["groups"]]
        for k in keys:
            if k in out["parameters"]:
                raise ValueError(f"parameter {k} in two templates")
            out["parameters"][k] = spec["template"]
    for spec in lt["lovs"]:
        for c in spec["codes"]:
            out["lovs"][c] = spec["template"]
    for spec in lt["documents"]:
        for r in sql_facts.rows("doc_template"):
            if r["code"].startswith(spec["prefix"]):
                out["documents"][r["code"]] = spec["template"]
    return out


def _prepared(tid: str) -> str:
    if tid in LISTS:
        return f"This workbook, sheet {LISTS[tid]}"
    t = all_templates()[tid]
    return f"{t['ref']}, sheet {sheet_name(t)}"


def lov_rows() -> list[dict[str, Any]]:
    tl = _template_lists()["lovs"]
    counts: dict[str, int] = {}
    for v in sql_facts.rows("lov_value", ("type_code", "code")):
        counts[v["type_code"]] = counts.get(v["type_code"], 0) + 1
    rows = []
    for r in sorted(sql_facts.rows("lov_type"), key=lambda x: x["code"]):
        mod = _module_of_lov(r["code"])
        m = module(mod)
        if r["code"] in migrated_lovs():
            route, prepared = ROUTES["migration"], f"Migration Workbook, object R01, code map {migrated_lovs()[r['code']]}"
        else:
            tid = tl.get(r["code"], "LIST-LOV")
            route, prepared = ROUTES["template"], _prepared(tid)
        rows.append({"code": r["code"], "name": r["name"], "module": m["name"], "drop": m["drop"], "brd": m["brd"],
                     "owner": m["owner"], "values": counts.get(r["code"], 0), "route": route, "prepared": prepared,
                     "due": dues()["D3"]["text"] if route == ROUTES["template"] else dues()["D4"]["text"],
                     "status": "Not started"})
    return rows


def parameter_rows() -> list[dict[str, Any]]:
    tl = _template_lists()["parameters"]
    d_text = data().get("parameter_text") or {}
    rows = []
    for p in sorted(sql_facts.rows("sys_parameter", "param_key"), key=lambda x: (x["category"], x["param_key"])):
        mod = data()["parameter_groups"][p["category"]]
        m = module(mod)
        key = p["param_key"]
        tid = tl.get(key) or ("LIST-NUMBERING" if key.endswith("_PREFIX") else "LIST-PARAMETERS")
        rng = ""
        if p.get("min_value") or p.get("max_value"):
            rng = f"{p.get('min_value') or '-'} to {p.get('max_value') or '-'}"
        rows.append({"key": key, "group": p["category"].replace("_", " ").title(), "module": m["name"],
                     "drop": m["drop"], "brd": m["brd"], "owner": m["owner"],
                     "description": d_text.get(key) or _clean(p["description"]),
                     "type": p["value_type"].replace("_", " ").lower(), "value": p.get("param_value") or "(blank)",
                     "range": rng, "route": ROUTES["template"], "template": tid, "prepared": _prepared(tid),
                     "due": dues()["D5"]["text"], "status": "Not started"})
    return rows


def numbering_rows() -> list[dict[str, Any]]:
    return [r for r in parameter_rows() if r["key"].endswith("_PREFIX")]


def event_rows() -> list[dict[str, Any]]:
    rows = []
    for e in sorted(sql_facts.rows("msg_notification_event"), key=lambda x: (x["module"], int(x.get("sort_order") or 0))):
        m = module(data()["event_modules"][e["module"]])
        rows.append({"code": e["code"], "name": e["name"], "module": m["name"], "drop": m["drop"], "owner": m["owner"],
                     "description": _clean(e.get("description")), "in_app": "Y" if e["default_in_app"] == "true" else "N",
                     "email": "Y" if e["default_email"] == "true" else "N", "route": ROUTES["template"],
                     "prepared": _prepared("LIST-NOTIFICATIONS"), "due": dues()["D5"]["text"], "status": "Not started"})
    return rows


def document_rows() -> list[dict[str, Any]]:
    tl = _template_lists()["documents"]
    rows = []
    for r in sorted(sql_facts.rows("doc_template"), key=lambda x: x["code"]):
        m = module(_module_of_doc(r["code"]))
        tid = tl.get(r["code"], "LIST-DOCUMENTS")
        fields = sorted(set(re.findall(r"\{\{(\w+)\}\}", r.get("body") or "")))
        rows.append({"code": r["code"], "title": r["title"], "module": m["name"], "drop": m["drop"], "owner": m["owner"],
                     "fields": ", ".join(fields), "route": ROUTES["template"], "prepared": _prepared(tid),
                     "due": dues()["D5"]["text"], "status": "Not started"})
    return rows


def accounting_event_rows() -> list[dict[str, Any]]:
    rows = []
    for r in sorted(sql_facts.rows("acc_event_type"), key=lambda x: (x["category"], x["code"])):
        rows.append({"code": r["code"], "name": r["name"], "category": r["category"].title(),
                     "journal": (r.get("journal_type") or "").replace("_", " ").title(),
                     "components": (r.get("amount_components") or "").replace(",", ", "),
                     "prepared": _prepared("D0-08"), "status": "Not started"})
    return rows


# ============================================================================ register


def register_rows() -> list[dict[str, Any]]:
    cat = {o["code"]: o for o in catalogue()["objects"]}
    rows = []
    for it in data()["items"]:
        m = module(it["module"])
        owner = it.get("owner") or m["owner"]
        where, prepared = "", ""
        if it.get("screen"):
            where = menu()[it["screen"]]
        if it["route"] == "migration":
            o = cat[it["object"]]
            where = f"Data migration object {o['code']} {o['name']}" + (f"; after go-live: {where}" if where else "")
            prepared = f"Migration Workbook, layout {', '.join(o.get('layouts') or [o['code']])}"
            if it.get("template"):
                prepared = f"{_prepared(it['template'])}; delivered in the layout of {o['code']}"
        elif it["route"] == "template":
            prepared = _prepared(it["template"])
            if it["template"] in brd_templates():
                owner = brd_templates()[it["template"]]["owner"]
            elif it["template"] in drop_templates():
                owner = drop_templates()[it["template"]]["owner"]
        else:
            prepared = "This index row (decision and values in BDOI comments)"
        d = dues()[it["due"]]
        rows.append({"id": it["id"], "area": it["area"], "name": it["name"], "provides": it["provides"],
                     "drop": m["drop"], "brd": m["brd"], "owner": owner, "due": d["label"], "due_date": fmt(d["date"]),
                     "weeks": f"T-{d['weeks']} weeks", "route": ROUTES[it["route"]], "where": where,
                     "prepared": prepared, "status": "Not started", "template": it.get("template"),
                     "object": it.get("object"), "due_code": it["due"]})
    return rows


def migration_rows() -> list[dict[str, Any]]:
    items = data()["items"]
    notes = {"R10": "Not migrated (excluded): users are created by access requests from BRD-11 template CI-01 "
                    "(CFG-110); legacy user IDs are only mapped (code map USER) so the officer, handler and collector "
                    "fields of migrated records resolve"}
    rows = []
    for o in catalogue()["objects"]:
        linked = [f"{i['id']} {i['name']}" for i in items if i["route"] == "migration" and i.get("object") == o["code"]]
        if linked:
            route = "Loaded by the migration: " + "; ".join(linked)
        elif o["code"] in notes:
            route = notes[o["code"]]
        elif o["category"] != "Reference":
            route = "Business data, not configuration: prepared and signed in the Migration Workbook only"
        else:
            route = ""
        rows.append({"code": o["code"], "name": o["name"], "category": o["category"],
                     "decision": o["decision"].replace("_", " ").title(), "owner": o["owner"],
                     "due": o.get("due") or "-", "route": route,
                     "rule": "One route: not typed in this workbook" if linked else "-"})
    return rows


# ============================================================================ checks


def seeded_config_tables() -> set[str]:
    records = {t for group in data()["records_not_configuration"] for t in group}
    return set(sql_facts.seed_tables()) - records


def check() -> list[str]:
    problems: list[str] = []
    d = data()
    cat = {o["code"] for o in catalogue()["objects"]}
    ids = [i["id"] for i in d["items"]]
    if len(ids) != len(set(ids)):
        problems.append("duplicate item ids")
    templates = all_templates()
    claimed: dict[str, list[str]] = {}
    item_keys = {"id", "area", "module", "name", "due", "route", "template", "screen", "provides", "tables", "object",
                 "owner"}
    for it in d["items"]:
        where = it["id"]
        if set(it) - item_keys:
            problems.append(f"{where}: unknown keys {sorted(set(it) - item_keys)} (quote a value with a comma)")
        if it["area"] not in d["areas"]:
            problems.append(f"{where}: unknown area {it['area']}")
        if it["module"] not in d["modules"]:
            problems.append(f"{where}: unknown module {it['module']}")
        if it["due"] not in dues():
            problems.append(f"{where}: unknown due {it['due']}")
        route = it.get("route")
        if route not in ROUTES:
            problems.append(f"{where}: route must be one of {', '.join(ROUTES)}")
        if route == "screen" and not it.get("screen"):
            problems.append(f"{where}: a screen route names its screen")
        if route == "template" and not it.get("template"):
            problems.append(f"{where}: a template route names its template")
        if route == "migration" and it.get("object") not in cat:
            problems.append(f"{where}: migration object {it.get('object')} is not in the BRD-13 catalogue")
        if route != "migration" and it.get("object"):
            problems.append(f"{where}: an object is named only with the migration route (one route per item)")
        if it.get("template") and it["template"] not in templates and it["template"] not in LISTS:
            problems.append(f"{where}: unknown template {it['template']}")
        if it.get("screen") and it["screen"] not in menu():
            problems.append(f"{where}: screen {it['screen']} is not in the menu")
        for t in it.get("tables") or []:
            claimed.setdefault(t, []).append(where)
    seeded = seeded_config_tables()
    for t in sorted(seeded):
        if t not in claimed:
            problems.append(f"seeded master table {t} is not claimed by an item (or listed in records_not_configuration)")
    for t, who in sorted(claimed.items()):
        if len(who) > 1:
            problems.append(f"table {t} claimed by {', '.join(who)} (one route per item)")
    # templates of the register exist once
    used = {i.get("template") for i in d["items"]}
    for tid in drop_templates():
        if tid not in used:
            problems.append(f"template {tid} is not used by any item")
    for tid, t in drop_templates().items():
        if t.get("table"):
            cols = set(sql_facts.columns(t["table"]))
            if not cols:
                problems.append(f"{tid}: table {t['table']} not found in the migrations")
            for f in t.get("fields") or []:
                if f and f not in cols:
                    problems.append(f"{tid}: {f} is not a column of {t['table']}")
            if t.get("fields") and len(t["fields"]) != len(t["rows"]):
                problems.append(f"{tid}: {len(t['fields'])} fields for {len(t['rows'])} columns")
    # every generated list has a module and exactly one route
    try:
        lovs, params, events, docs = lov_rows(), parameter_rows(), event_rows(), document_rows()
        accounting_event_rows()
    except KeyError as e:
        problems.append(str(e))
        lovs, params, events, docs = [], [], [], []
    tl = _template_lists()
    for code in tl["lovs"]:
        if code in migrated_lovs():
            problems.append(f"list {code} is in a template and loaded by R01 (one route per list)")
    known_lovs = {r["code"] for r in lovs}
    for code in tl["lovs"]:
        if code not in known_lovs:
            problems.append(f"list {code} of list_templates is not a list of the platform")
    known_params = {r["key"] for r in params}
    for key in tl["parameters"]:
        if key not in known_params:
            problems.append(f"parameter {key} of a template is not a parameter of the platform")
    # every reference object that is migrated or carried has a register item or a note
    for r in migration_rows():
        if r["category"] == "Reference" and not r["route"]:
            problems.append(f"reference object {r['code']} has no register item and no note")
    # decisions refer to items of the clarification chapters and open decisions
    texts = " ".join(Path(brand.SRC_DIR / s["source"]).read_text(encoding="utf-8") for s in d["sets"])
    texts += " " + (brand.src_dir("BRD-13") / "pack" / "catalogue.yaml").read_text(encoding="utf-8")
    for dec in d["decisions"]:
        for ref in dec["refs"]:
            if not re.search(rf"\b{re.escape(ref)}\b", texts):
                problems.append(f"{dec['id']}: {ref} is not in the clarification chapters or open decisions")
    # the templates: checks, lists, examples of mandatory columns, dependencies without a cycle
    problems += [f"templates: {p}" for p in config_inputs().check(combined())]
    # banned words
    banned = config_inputs().BANNED
    for text in [str(i) for i in d["items"]] + [str(t) for t in d["templates"]] + [str(x) for x in d["decisions"]]:
        m = banned.search(text)
        if m:
            problems.append(f"banned word '{m.group(0)}' in {text[:60]}")
    return problems


# ============================================================================ clarification counts (summary)


def clarification_counts() -> list[dict[str, Any]]:
    out = []
    cat = catalogue()
    for s in data()["sets"]:
        text = (brand.SRC_DIR / s["source"]).read_text(encoding="utf-8")
        if s["brd"] == "BRD-13":
            rules = len(cat["proposals"])
            statuses = [r[6] for r in cat["decisions"]]
        else:
            rules = len(re.findall(rf"^\| {s['clarifications']}-\d+ \|", text, re.M))
            oq = text.split("## Open questions", 1)[1].split("\n# ", 1)[0]
            statuses = [line.rsplit("|", 2)[-2].strip() for line in oq.splitlines()
                        if re.match(r"^\| [A-Z]{1,3}Q\d+ \|", line)]
        open_n = sum(1 for x in statuses if x in ("OPEN", "PARTIAL"))
        rec = sum(1 for x in statuses if x == "RECOMMENDED")
        answered = sum(1 for x in statuses if x == "ANSWERED")
        out.append({"brd": s["brd"], "name": brand.BRD_NAMES[s["brd"]], "rules": rules, "questions": len(statuses),
                    "open": open_n, "recommended": rec, "answered": answered,
                    "chapter": "Proposed business rules and clarifications for confirmation"})
    return out


# ============================================================================ workbook


def template_order() -> list[Any]:
    """The guided templates of the workbook in the order they are filled in (dependencies first)."""
    import guided_xlsx as g  # noqa: PLC0415

    return g.fill_in_order(guided())


@lru_cache(maxsize=1)
def guided() -> list[Any]:
    ci = config_inputs()
    comb = combined()
    deps = {t["id"]: ci.depends(t) for t in comb["templates"]}
    dues_ = ci.template_dues([t["id"] for t in comb["templates"]], deps, ci.register_due,
                             "With the Drop 0 sign-off")
    return ci.guided_templates(comb, dues_, sheet_of=sheet_name)


def due_conflicts() -> list[str]:
    """Templates due after a template that uses their codes (for information: the index keeps the due dates)."""
    ci = config_inputs()
    out = []
    for t in combined()["templates"]:
        mine = ci.register_due(t["id"])
        for d in ci.depends(t):
            theirs = ci.register_due(d)
            if mine and theirs and theirs[0] > mine[0]:
                out.append(f"{t['id']} (due {fmt(mine[0])}) uses the codes of {d} (due {fmt(theirs[0])})")
    return out


def build_workbook() -> Path:
    import guided_xlsx as g  # noqa: PLC0415
    from bdoi_xlsx import Column  # noqa: PLC0415

    meta = data()["meta"]
    statuses = list(meta["statuses"])
    book = g.GuidedBook(meta["title"], str(meta["version"]))
    ci = config_inputs()
    for code, x in ci.resolve_lists(combined()).items():
        book.add_list(code, x["name"], x["values"], x["note"], x["strict"])
    templates = guided()
    book.plan_all(templates)
    ci.finish_checks(book, templates)
    order = template_order()
    for i, t in enumerate(order, start=1):
        book.template_sheet(t, i, len(order))
    status_col = lambda: Column("status", "Status", 14, "Status of the row", values=statuses)  # noqa: E731
    comment_col = lambda: Column("comments", "BDOI comments", 36, "Decision, value or comment of the owner")  # noqa: E731

    lov_cols = [
        Column("code", "List", 26, "Code of the list on Lists of Values"),
        Column("name", "Name", 30, "Name of the list"),
        Column("module", "Module", 22, "Module that uses the list"),
        Column("drop", "Drop", 8, "Drop"),
        Column("owner", "Owner (BDOI unit)", 30, "Unit that owns the values"),
        Column("values", "Values proposed", 10, "Number of values proposed with the platform", kind="number"),
        Column("route", "Route", 14, "Template or Migration object", values=list(ROUTES.values())),
        Column("prepared", "Prepared in", 40, "Template or code map where the values are given"),
        Column("due", "Due", 22, "Due date"),
        Column("keep", "Values to keep, add or retire", 40, "The owner's decision on the values"),
        status_col(), comment_col(),
    ]
    ref_sheets = [
        ("LIST-LOV", "Every list of values of BIBS with its owner and its one route", lov_cols, lov_rows()),
        ("LIST-PARAMETERS", "Every system parameter with the value proposed and the value BDOI wants", [
            Column("key", "Parameter", 30, "Parameter on System Parameters"),
            Column("group", "Group", 16, "Group of the parameter"),
            Column("module", "Module", 20, "Module"),
            Column("drop", "Drop", 8, "Drop"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that decides the value"),
            Column("description", "What it controls", 60, "Meaning of the parameter"),
            Column("type", "Type", 12, "Kind of value"),
            Column("value", "Value proposed", 22, "Value proposed with the platform"),
            Column("range", "Allowed range", 12, "Minimum and maximum"),
            Column("route", "Route", 12, "Template: the value is given here or in the named template",
                   values=list(ROUTES.values())),
            Column("prepared", "Prepared in", 34, "Where the value wanted is given"),
            Column("due", "Due", 22, "Due date"),
            Column("wanted", "Value wanted", 22, "Value BDOI wants"),
            Column("reason", "Reason", 30, "Policy reference or reason of a change"),
            Column("approver", "Approved by", 20, "BDOI approver of the value"),
            status_col(), comment_col(),
        ], parameter_rows()),
        ("LIST-NUMBERING", "Prefixes of the numbered documents; receipt series come from object R11, check numbers "
                           "from D0-12", [
            Column("key", "Parameter", 30, "Prefix parameter"),
            Column("module", "Module", 20, "Module of the document"),
            Column("description", "Numbered document", 60, "Document numbered with the prefix"),
            Column("value", "Prefix proposed", 16, "Prefix proposed with the platform"),
            Column("prepared", "Prepared in", 34, "Where the prefix wanted is given"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that decides the prefix"),
            Column("wanted", "Prefix wanted", 16, "Prefix BDOI wants; the year and a running number follow it"),
            status_col(), comment_col(),
        ], numbering_rows()),
        ("LIST-NOTIFICATIONS", "Every notification event with its default channels; users choose their own channels "
                               "in Notification Settings", [
            Column("code", "Event", 28, "Notification event"),
            Column("name", "Name", 32, "Name shown in Notification Settings"),
            Column("module", "Module", 20, "Module"),
            Column("drop", "Drop", 8, "Drop"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that confirms the event"),
            Column("description", "Who is told and when", 50, "What the event tells the user"),
            Column("in_app", "In-app proposed", 10, "Y or N"),
            Column("email", "E-mail proposed", 10, "Y or N"),
            Column("in_app_wanted", "In-app wanted", 10, "Y or N", values=["Y", "N"]),
            Column("email_wanted", "E-mail wanted", 10, "Y or N", values=["Y", "N"]),
            Column("mailbox", "Unit mailbox", 30, "Mailbox that also receives the e-mail (optional)"),
            status_col(), comment_col(),
        ], event_rows()),
        ("LIST-DOCUMENTS", "Every generated document with the unit that gives its BDOI wording and layout", [
            Column("code", "Template", 26, "Template code on Document Templates"),
            Column("title", "Title", 34, "Title of the template"),
            Column("module", "Module", 20, "Module that generates the document"),
            Column("drop", "Drop", 8, "Drop"),
            Column("owner", "Owner (BDOI unit)", 28, "Unit that gives the wording and layout"),
            Column("fields", "Merge fields available", 50, "Fields BIBS fills in"),
            Column("prepared", "Prepared in", 34, "Where the text is given"),
            Column("due", "Due", 22, "Due date"),
            Column("layout", "BDOI layout attached", 12, "Y when the layout is sent", values=["Y", "N"]),
            status_col(), comment_col(),
        ], document_rows()),
        ("LIST-EVENTS", "The accounting events that need a rule of template D0-08; owner Head, Comptrollership", [
            Column("code", "Event type", 30, "Accounting event of the platform"),
            Column("name", "Name", 36, "Name of the event"),
            Column("category", "Category", 14, "Category"),
            Column("journal", "Journal type", 14, "Journal type of the entries"),
            Column("components", "Amount components", 50, "Amounts the rule lines can post"),
            Column("prepared", "Prepared in", 30, "Template of the rules"),
            status_col(), comment_col(),
        ], accounting_event_rows()),
    ]
    for tid, desc, cols, rows in ref_sheets:
        book.table_sheet(LISTS[tid], f"Reference: {LISTS[tid][4:]}", desc, cols, rows, statuses)
    book.table_sheet(MIGRATION_SHEET, "Reference: migration cross-reference", "Reference data loaded through the "
                     "migration and items entered on screens: one route per item", [
        Column("code", "Object", 8, "Data migration object of the Migration Workbook"),
        Column("name", "Name", 40, "Object name"),
        Column("category", "Category", 12, "Object category"),
        Column("decision", "Proposed decision", 16, "Decision proposed in the Data Migration set"),
        Column("owner", "Business owner", 30, "Owner of the object"),
        Column("due", "Extracts due", 10, "Milestone of the extracts"),
        Column("route", "Configuration route", 70, "Register items loaded by the object, or why none is"),
        Column("rule", "Duplication rule", 22, "The item is not also typed in this workbook"),
    ], migration_rows(), statuses)

    # Start here: the register items by due date, and within a due date the templates in their fill-in order.
    plans = book.plans
    position = {t.id: i for i, t in enumerate(order)}
    reg = register_rows()
    route_rank = {ROUTES["template"]: 0, ROUTES["migration"]: 1, ROUTES["screen"]: 2}
    reg.sort(key=lambda r: (dues()[r["due_code"]]["date"], route_rank[r["route"]],
                            position.get(r["template"] or "", 999), r["id"]))
    idx = []
    for i, r in enumerate(reg, start=1):
        tid = r["template"]
        opens, depends_on, entered, missing = "", "-", "", ""
        if tid in plans:
            p = plans[tid]
            opens = (f"{tid} →", p.sheet)
            depends_on = ", ".join(p.template.depends) or "-"
            entered, missing = p.rows_formula(), p.missing_formula()
        elif tid in LISTS:
            opens = ("Open →", LISTS[tid])
        elif r["object"]:
            opens = (f"{r['object']} →", MIGRATION_SHEET)
        idx.append({"Step": i, "ID": r["id"], "Configuration item": r["name"], "What BDOI provides": r["provides"],
                    "Owner (BDOI unit)": r["owner"], "Due": f"{r['due_date']} – {r['due']}", "Route": r["route"],
                    "Entered or loaded on": r["where"], "Prepared in": r["prepared"], "Depends on": depends_on,
                    "Rows entered": entered, "Mandatory cells missing": missing, "Status": "Not started",
                    "BDOI comments": None, "Provided on": None, "Open": opens})
    columns = [("Step", 6), ("ID", 9), ("Configuration item", 24), ("What BDOI provides", 40),
               ("Owner (BDOI unit)", 22), ("Due", 20), ("Route", 11), ("Entered or loaded on", 28),
               ("Prepared in", 28), ("Depends on", 12), ("Rows entered", 9), ("Mandatory cells missing", 10),
               ("Status", 14), ("BDOI comments", 28), ("Provided on", 11), ("Open", 9)]
    steps = [
        "The index below lists every configuration input of Drop 0 and of the platform set-up that Drop 1 depends "
        "on, one row per item, by due date; within a due date the templates come in the order they are filled in, "
        "a template whose codes others use first.",
        "Each item has one route. Template: fill in the template sheet (Open →); its rows are then entered or "
        "uploaded on the named screen and authorised by a second user. Screen: enter the item on the named BIBS "
        "screen and record the decision in BDOI comments. Migration object: the item comes from the legacy systems "
        "through the named object of the Migration Workbook (BRD-13); it is not typed here.",
        "On a template sheet, read the header block and the guide above each column (Mandatory, Format, Allowed "
        "values, What to enter; * marks a mandatory column), overwrite or delete the grey example row and enter one "
        "row per record. Drop-downs offer the allowed values (sheet Reference lists) and the codes of the templates "
        "filled in before; a mandatory cell left empty turns red.",
        "The sheets starting with Ref list every list of values, system parameter, numbering prefix, notification "
        "event, document template and accounting event of BIBS, and the migration cross-reference; set the value "
        "wanted, the decision or the comment on each row you own.",
        f"Set the Status of each item, follow Rows entered and Mandatory cells missing, raise questions on the sheet "
        f"Questions and comments, and return the workbook to the {brand.VENDOR} project team by each due date. The "
        "status is reviewed at the weekly programme meeting. Codes must be the codes of the other sheets: branch "
        "codes of D0-02, accounts of D0-06, cost centres of D0-04, bank accounts of D0-11, group profiles of UA-02 "
        "and users of UA-01.",
    ]
    identity = [("Client", brand.CLIENT), ("Scope", "Drop 0 (BRD-03, BRD-11, BRD-13) and the platform set-up "
                                                     "Drop 1 depends on"),
                ("Version and date", f"Version {meta['version']}, {meta['date']}"),
                ("Go-live", meta["go_live_text"]),
                ("Owner", "Program Manager, Business Project Services, with the BDOI owner of each item"),
                ("Prepared by", brand.VENDOR), ("Classification", brand.CLASSIFICATION)]
    per_due: dict[str, int] = {}
    for it in data()["items"]:
        per_due[it["due"]] = per_due.get(it["due"], 0) + 1

    def more(ws, row):
        row = book.small_table(ws, row, "Due milestones", ["Due", "Milestone", "Due date", "Before go-live",
                                                            "Items", "Why this date"],
                               [[c, x["label"], fmt(x["date"]), f"T-{x['weeks']} weeks", per_due.get(c, 0), x["note"]]
                                for c, x in dues().items()], spans=[1, 3, 1, 1, 1, 6])
        return book.small_table(ws, row, "Reference sheets and other sheets", ["Sheet", "What it holds"], [
            ((name, name), desc) for name, desc in
            [(LISTS[tid], desc) for tid, desc, _, _ in ref_sheets] +
            [(MIGRATION_SHEET, "Reference data loaded by the migration and the register items it carries"),
             (g.LISTS, "Every list of allowed values of the drop-downs, with code and label"),
             (g.QUESTIONS, "Questions and comments per template and column, with the answer")]
        ], spans=[3, 8])

    book.questions_sheet(order)
    ws = book.start_sheet(meta["title"], "Every configuration input BDOI provides before go-live for Drop 0 and the "
                                         "platform set-up Drop 1 depends on", identity, steps,
                          "Configuration inputs by due date, templates in the order they are filled in", columns, idx,
                          statuses=statuses, after=more)
    # Provided on: a date
    from openpyxl.worksheet.datavalidation import DataValidation  # noqa: PLC0415

    heads = {ws.cell(r, 1).value: r for r in range(1, 40)}
    header_row = next(r for r in range(1, 60) if ws.cell(r, 1).value == "Step" and ws.cell(r, 2).value == "ID")
    col = [h for h, _ in columns].index("Provided on") + 1
    letter = ws.cell(header_row, col).column_letter
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True,
                        showErrorMessage=True, errorTitle="Provided on", error="Enter a date.")
    ws.add_data_validation(dv)
    dv.add(f"{letter}{header_row + 1}:{letter}{header_row + len(idx)}")
    for r in range(header_row + 1, header_row + len(idx) + 1):
        ws.cell(r, col).number_format = g.DATE_FORMAT
    del heads
    book.lists_sheet()
    file = drop_output_name("Configuration_Inputs", "xlsx")
    sheets = ([g.START] + [t.sheet for t in order] + [LISTS[tid] for tid, *_ in ref_sheets] + [MIGRATION_SHEET]
              + [g.LISTS, g.QUESTIONS])
    return book.save(brand.drop_set_dir(DROP) / file, sheets,
                     {"title": meta["title"], "keywords": f"{meta['doc_type']}, Drop 0, BIBS, BDOI"})


def drop_output_name(name: str, ext: str) -> str:
    return brand.drop_output_name(DROP, name, str(data()["meta"]["version"]), ext)


# ============================================================================ Word (```pack blocks of the summary)


def render(doc: Any, render: str, **_: Any) -> None:  # noqa: A002 - block key
    d = data()
    counts = {c["brd"]: c for c in clarification_counts()}
    if render == "sets":
        rows = []
        for s in d["sets"]:
            folder = brand.out_dir(s["brd"], "FRS")
            files = sorted(p.name for p in folder.glob("*_v2.0.*"))
            rows.append([f"{s['brd']} {brand.BRD_NAMES[s['brd']]}", files, s["confirms"], s["owner"], s["signoff"]])
        doc.table(["Sign-off set", "Files (v2.0)", "What the business confirms", "Signed by", "Sign-off"], rows,
                  widths=[2.4, 5.4, 5.8, 2.8, 1.6], caption="The three sign-off sets of Drop 0", size=7.5, keep_rows=False)
        return
    if render == "clarifications":
        rows = [[f"{c['brd']} {c['name']}", c["rules"], c["questions"], c["open"], c["recommended"], c["answered"]]
                for c in counts.values()]
        tot = [sum(c[k] for c in counts.values()) for k in ("rules", "questions", "open", "recommended", "answered")]
        rows.append(["Total", *tot])
        doc.table(["Sign-off set", "Proposed rules and clarifications for confirmation", "Open questions and decisions",
                   "Open or partly answered", "Recommended, to confirm", "Answered"], rows,
                  widths=[4.2, 3.2, 3.0, 2.4, 2.4, 2.0], caption="Items for confirmation per sign-off set",
                  align=["left", "center", "center", "center", "center", "center"], size=8.5)
        return
    if render == "decisions":
        rows = [[x["id"], x["topic"], x["decision"], ", ".join(x["refs"]), x["owner"], x["by"]] for x in d["decisions"]]
        doc.table(["ID", "Topic", "Decision BDOI takes", "From", "BDOI owner", "Needed by"], rows,
                  widths=[1.4, 2.8, 6.4, 2.6, 2.6, 2.4], caption="Consolidated decisions BDOI takes to close Drop 0",
                  size=7.5, keep_rows=False)
        return
    if render in ("close_entry", "close_exit", "drop1_entry"):
        doc.bullets(d["entry_exit"][render])
        return
    if render == "timeline":
        by_due: dict[str, list[str]] = {}
        for it in d["items"]:
            by_due.setdefault(it["due"], []).append(it["name"])
        rows = [[c, x["label"], fmt(x["date"]), f"T-{x['weeks']} weeks", by_due.get(c, [])] for c, x in dues().items()]
        doc.table(["Due", "Milestone", "Date", "Before go-live", "Configuration inputs due"], rows,
                  widths=[1.0, 3.6, 2.2, 2.2, 9.0], caption="Configuration inputs timeline (go-live T = 3-Jan-2028)",
                  size=7.5, keep_rows=False)
        return
    if render == "inputs":
        areas: dict[str, dict[str, int]] = {}
        for it in d["items"]:
            a = areas.setdefault(it["area"], {"screen": 0, "template": 0, "migration": 0})
            a[it["route"]] += 1
        rows = [[a, v["template"], v["screen"], v["migration"], sum(v.values())] for a, v in areas.items()]
        rows.append(["Total", *[sum(v[k] for v in areas.values()) for k in ("template", "screen", "migration")],
                     len(d["items"])])
        doc.table(["Area", "Template", "Screen", "Migration object", "Items"], rows, widths=[7.0, 2.4, 2.4, 3.0, 2.0],
                  caption="Register items per area and route", align=["left", "center", "center", "center", "center"],
                  size=8.5)
        lists = [["Lists of values", len(lov_rows()), sum(1 for r in lov_rows() if r["route"] == ROUTES["migration"])],
                 ["System parameters", len(parameter_rows()), 0],
                 ["Numbered document prefixes", len(numbering_rows()), 0],
                 ["Notification events", len(event_rows()), 0],
                 ["Document templates", len(document_rows()), 0],
                 ["Accounting events (rules of D0-08)", len(accounting_event_rows()), 0]]
        doc.table(["List", "Rows", "Of which loaded by the migration (R01)"], lists, widths=[8.0, 3.0, 5.0],
                  caption="Lists of the workbook with one row per value", align=["left", "center", "center"], size=8.5)
        return
    if render == "risks":
        doc.table(["ID", "Risk", "Impact", "Mitigation", "Owner"],
                  [[r["id"], r["risk"], r["impact"], r["mitigation"], r["owner"]] for r in d["risks"]],
                  widths=[1.4, 4.4, 4.2, 5.2, 2.6], caption="Risks", size=7.5, keep_rows=False)
        return
    if render == "dependencies":
        doc.table(["ID", "Dependency", "From", "Needed by"],
                  [[r["id"], r["what"], r["from"], r["needed"]] for r in d["dependencies"]],
                  widths=[1.4, 9.0, 4.0, 3.4], caption="Dependencies", size=8)
        return
    if render == "signoff":
        doc.signoff([{"name": "", "role": s["role"], "organisation": s["organisation"]} for s in d["signatories"]])
        return
    if render == "count":  # inline numbers are not supported; kept for completeness
        return
    raise ValueError(f"unknown render {render}")


def build_summary() -> Path:
    import bdoi_docx

    results = bdoi_docx.build_markdown(SUMMARY, brand.drop_set_dir(DROP) / drop_output_name("Closure_Summary", "docx"))
    docx, pdf = results[0]
    if pdf:
        pdf.unlink(missing_ok=True)
    return docx


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build the Drop 0 closure set")
    ap.add_argument("--check", action="store_true", help="checks only")
    ap.add_argument("--xlsx", action="store_true", help="the workbook only")
    args = ap.parse_args(argv)
    problems = check()
    for p in problems:
        print(p)
    d = data()
    print(f"{len(d['items'])} register items, {len(all_templates())} templates, {len(lov_rows())} lists, "
          f"{len(parameter_rows())} parameters, {len(event_rows())} events, {len(document_rows())} document templates, "
          f"{len(accounting_event_rows())} accounting events")
    for c in clarification_counts():
        print(f"{c['brd']}: {c['rules']} proposed rules, {c['questions']} open questions / decisions "
              f"({c['open']} open, {c['recommended']} recommended, {c['answered']} answered)")
    for c in due_conflicts():
        print(f"note: {c}")
    if args.check:
        import guided_xlsx  # noqa: PLC0415

        path = brand.drop_set_dir(DROP) / drop_output_name("Configuration_Inputs", "xlsx")
        if path.exists():
            found = guided_xlsx.verify(path, [sheet_name(t) for t in all_templates().values()])
            for p in found:
                print(p)
            problems += found
    if args.check or problems:
        return 1 if problems else 0
    print(f"xlsx: {build_workbook()}")
    if not args.xlsx:
        print(f"docx: {build_summary()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
