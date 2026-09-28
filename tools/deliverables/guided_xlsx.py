"""Guided input workbooks: templates the business fills in, one self-contained sheet per template.

Used by the configuration input templates of a Drop 0 set (06, src/signoff/config_inputs.py), the Drop 0
configuration inputs workbook (drop_closure.py) and the load templates of the Migration Workbook (BRD-13).

A guided workbook has:

* a **Start here** sheet: who the workbook is for, a few plain steps and the index of the templates in the order
  they are filled in (a template whose codes others use comes first), with the rows entered, the mandatory cells
  still missing (formulas), a status drop-down and a link to each sheet;
* one **template sheet** per template, self-contained: a header block (purpose, provided by, due, how it is loaded,
  depends on with links, progress, confirmed by and date, back to Start here), a guide band aligned column by column
  directly above the input header (Mandatory, Format, Allowed values, What to enter), the header row with the
  mandatory columns marked * and a note with the full description, one example row in grey italics labelled
  "Example – overwrite or delete", then the input rows as an Excel table with drop-downs, date and number checks,
  input messages and the missing mandatory cells highlighted;
* a visible **Reference lists** sheet with every list of allowed values (code and label), one named range per list;
* a **Questions and comments** sheet to raise a question per template and column and record the answer.

The same design is used by the upload templates and the load templates exported by the platform: title block,
guide band, mandatory header marked *, example row, visible lists.

Validation of a column (``GuideColumn.check``):
  yn                      Y or N (list YES_NO)
  date                    a date (dd-mmm-yyyy)
  number | number:0-100   a number, optionally within a range
  whole | whole:1-12      a whole number, optionally within a range
  text:40                 at most 40 characters
  list:CODE               a value of the list CODE of the Reference lists sheet (add_list)
  code:ID/Column          a code entered in the column Column of the template ID (a template filled in before)
"""

from __future__ import annotations

import datetime as _dt
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
import brand  # noqa: E402

FONT = brand.FONT
START = "Start here"
LISTS = "Reference lists"
QUESTIONS = "Questions and comments"
EXAMPLE_LABEL = "Example – overwrite or delete"
BACK = "← Back to Start here"
STATUSES = ["Not started", "In progress", "Ready for review", "Confirmed"]
QUESTION_STATUSES = ["Open", "Answered", "Closed"]
STATUS_STYLE = {
    "Not started": (brand.DIRTY_WHITE, brand.MUTED),
    "In progress": (brand.AMBER_BG, brand.AMBER),
    "Ready for review": (brand.BG_BLUE, brand.HEADER_BLUE),
    "Confirmed": (brand.SUCCESS_BG, brand.SUCCESS),
    "Not applicable": (brand.BORDER_LIGHT, brand.MUTED),
    "Open": (brand.AMBER_BG, brand.AMBER),
    "Answered": (brand.BG_BLUE, brand.HEADER_BLUE),
    "Closed": (brand.SUCCESS_BG, brand.SUCCESS),
}
GUIDE_ROWS = ["Mandatory", "Format", "Allowed values", "What to enter"]
LABEL_WIDTH = 17

THIN = Side(style="thin", color=brand.BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=brand.HEADER_BLUE)
LABEL_FILL = PatternFill("solid", fgColor=brand.BG_BLUE)
GUIDE_FILL = PatternFill("solid", fgColor=brand.DIRTY_WHITE)
EXAMPLE_FILL = PatternFill("solid", fgColor=brand.DIRTY_WHITE)
INPUT_FILL = PatternFill("solid", fgColor=brand.WHITE)
MANDATORY_FILL = PatternFill("solid", fgColor=brand.AMBER_BG)
MISSING_FILL = PatternFill("solid", fgColor=brand.DANGER_BG, bgColor=brand.DANGER_BG)
DATE_FORMAT = "dd-mmm-yyyy"


def font(size: float = 10, bold: bool = False, colour: str = brand.TEXT, italic: bool = False,
         underline: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=colour, italic=italic,
                underline="single" if underline else None)


def set_link(cell, target: str) -> None:
    """An internal hyperlink to 'Sheet!A1' (or '#Sheet!A1', or a sheet name alone)."""
    target = target.lstrip("#")
    sheet, _, ref = target.rpartition("!") if "!" in target else (target, "", "A1")
    sheet = sheet.strip("'").replace("''", "'")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{quote_sheetname(sheet)}!{ref or 'A1'}")


# ============================================================================ model


@dataclass
class GuideColumn:
    """One column of a template with its guide: header, mandatory (Y, N, C, 'Cond.: condition', 'Y (part)' or
    'Y for ...'), format, what to enter, example ('-' for none), check (see the module doc) and extra guide rows."""

    header: str
    mandatory: str
    format: str
    what: str
    example: str = "-"
    check: str = ""
    extra: dict[str, str] = field(default_factory=dict)
    allowed: str = ""  # replaces the allowed values text derived from the check
    allowed_link: str = ""  # hyperlink of the allowed values cell (another sheet)
    text_cell: bool = True  # keep what is typed as text (codes with leading zeros, file dates)


@dataclass
class Template:
    id: str
    name: str
    sheet: str
    purpose: str
    owner: str
    due: str
    load: str
    columns: list[GuideColumn]
    source: str = ""
    depends: list[str] = field(default_factory=list)
    capacity: int = 500
    extra_rows: list[str] = field(default_factory=list)  # labels of guide rows above Mandatory (e.g. review rows)
    input_rows: list[str] = field(default_factory=list)  # which of extra_rows are filled in by the reader
    extra_lists: dict[str, str] = field(default_factory=dict)  # extra row label -> list code (drop-down)
    confirm: tuple[str, str] = ("Confirmed by (BDOI owner)", "Date")
    notes: list[str] = field(default_factory=list)  # more lines of the header block (label, text) as "Label: text"


def mandatory_kind(raw: str) -> tuple[str, str]:
    """('yes' | 'no' | 'cond', guide text) of a mandatory cell of the sources."""
    r = " ".join(str(raw).split())
    if r in ("Y", "Yes"):
        return "yes", "Yes"
    if r in ("N", "No", "-", ""):
        return "no", "No"
    if r == "C":
        return "cond", "Conditional (see What to enter)"
    m = re.match(r"Cond\.?:?\s*(.*)$", r)
    if m:
        return "cond", f"Conditional: {m.group(1)}"
    m = re.match(r"Y \((.*)\)$", r)
    if m:
        return "yes", f"Yes ({m.group(1)})"
    m = re.match(r"Y (for .*)$", r)
    if m:
        return "cond", f"Conditional: {m.group(1)}"
    return "cond", f"Conditional: {r}"


def fill_in_order(templates: Sequence[Template]) -> list[Template]:
    """The templates in the order they are filled in: by the length of their dependency chain (a template whose
    codes others use comes first); at the same level the templates others depend on first, then by ID; templates
    that neither use nor give codes are filled in at any time and come last. A cycle raises ValueError."""
    by_id = {t.id: t for t in templates}
    deps = {t.id: [d for d in t.depends if d in by_id] for t in templates}
    users = {t.id: [u for u in by_id if t.id in deps[u]] for t in templates}
    level: dict[str, int] = {}

    def depth(tid: str, seen: tuple[str, ...] = ()) -> int:
        if tid in seen:
            raise ValueError(f"dependency cycle: {' > '.join(seen + (tid,))}")
        if tid not in level:
            level[tid] = 1 + max((depth(d, seen + (tid,)) for d in deps[tid]), default=-1)
        return level[tid]

    for tid in by_id:
        depth(tid)
    standalone = {tid for tid in by_id if not deps[tid] and not users[tid]}
    return sorted(templates, key=lambda t: (t.id in standalone, level[t.id], not users[t.id], t.id))


# ============================================================================ planning


@dataclass
class Plan:
    """Where everything of a template sheet is: rows of the guide, header and inputs, letters of the columns."""

    template: Template
    sheet: str
    first_col: int  # column of the first template column (B)
    guide_top: int
    header_row: int
    example_row: int
    last_row: int
    progress_row: int
    confirm_row: int
    letters: dict[str, str]
    key: str  # header of the column that identifies a used row of the example
    key_example: Any

    @property
    def last_letter(self) -> str:
        return get_column_letter(self.first_col + len(self.template.columns) - 1)

    def rows_formula(self) -> str:
        return f"={quote_sheetname(self.sheet)}!$C${self.progress_row}"

    def missing_formula(self) -> str:
        return f"={quote_sheetname(self.sheet)}!$E${self.progress_row}"

    def column_range(self, header: str) -> str:
        col = self.letters[header]
        return f"{quote_sheetname(self.sheet)}!${col}${self.example_row}:${col}${self.last_row}"


HEADER_BLOCK = ["Purpose", "Provided by", "Due", "How it is loaded", "Source", "Depends on"]


def plan(t: Template) -> Plan:
    block = len(HEADER_BLOCK) + len(t.notes)
    progress_row = 3 + block
    confirm_row = progress_row + 1
    guide_top = confirm_row + 2
    header_row = guide_top + len(t.extra_rows) + len(GUIDE_ROWS)
    example_row = header_row + 1
    letters = {c.header: get_column_letter(2 + i) for i, c in enumerate(t.columns)}
    key_col = next((c for c in t.columns if example_value(c) not in (None, "") and not c.check.startswith(("date", "number", "whole"))),
                   None) or next((c for c in t.columns if example_value(c) not in (None, "")), t.columns[0])
    return Plan(t, t.sheet, 2, guide_top, header_row, example_row, example_row + t.capacity, progress_row,
                confirm_row, letters, key_col.header, example_value(key_col))


def example_value(c: GuideColumn) -> Any:
    v = str(c.example).strip()
    if v in ("-", ""):
        return None
    kind = c.check.split(":", 1)[0]
    if kind == "date":
        for f in ("%d-%b-%Y", "%Y-%m-%d"):
            try:
                return _dt.datetime.strptime(v, f).date()
            except ValueError:
                pass
    if kind in ("number", "whole"):
        try:
            n = float(v.replace(",", ""))
            return int(n) if kind == "whole" or n.is_integer() and "." not in v else n
        except ValueError:
            pass
    return v


# ============================================================================ workbook


@dataclass
class RefList:
    code: str
    name: str
    values: list[tuple[str, str]]
    note: str = ""
    strict: bool = True  # False: a value outside the list is allowed after a warning (lists BDOI may extend)
    used_in: list[str] = field(default_factory=list)
    row: int = 0  # first row on the Reference lists sheet

    @property
    def range_name(self) -> str:
        return "LIST_" + re.sub(r"[^A-Z0-9_]", "_", self.code.upper())


class GuidedBook:
    """Writes the guided sheets into an openpyxl workbook (a new one, or the workbook of a BdoiWorkbook).

    register(ws, description, rows) is called for each sheet written (BdoiWorkbook keeps its cover index)."""

    def __init__(self, title: str, version: str, wb: Workbook | None = None,
                 register: Callable[[Any, str, int], None] | None = None,
                 start: str = START, questions: str | None = QUESTIONS) -> None:
        self.title = title
        self.version = version
        self.own = wb is None
        self.wb = wb or Workbook()
        if self.own:
            self.wb.remove(self.wb.active)
        self.register = register or (lambda ws, description, rows: None)
        self.start = start
        self.questions = questions
        self.lists: dict[str, RefList] = {}
        self.plans: dict[str, Plan] = {}
        self._next_list_row = 5  # first value row of the Reference lists sheet; lists keep their rows once added
        self.add_list("YES_NO", "Yes or no", [("Y", "Yes"), ("N", "No")])

    # ------------------------------------------------------------------ lists

    def add_list(self, code: str, name: str, values: Sequence[tuple[str, str] | str], note: str = "",
                 strict: bool = True) -> RefList:
        vals = [(v, "") if isinstance(v, str) else (str(v[0]), str(v[1] or "")) for v in values]
        if code in self.lists:
            known = self.lists[code]
            if [v for v, _ in known.values] != [v for v, _ in vals]:
                raise ValueError(f"list {code} defined twice with other values")
            return known
        if not vals:
            raise ValueError(f"list {code} has no values")
        ref = RefList(code, name, vals, note, strict, row=self._next_list_row)
        self._next_list_row += len(vals)
        self.lists[code] = ref
        return ref

    def list_for_values(self, name: str, values: Sequence[str]) -> RefList:
        """A list for a drop-down of a table column (status, route, ...), named after the column."""
        code = re.sub(r"[^A-Z0-9]+", "_", name.upper()).strip("_")
        for ref in self.lists.values():
            if [v for v, _ in ref.values] == list(values):
                return ref
        while code in self.lists:
            code += "_2"
        return self.add_list(code, name, list(values))

    # ------------------------------------------------------------------ helpers

    def _print_setup(self, ws, title_rows: str | None = None, portrait: bool = False) -> None:
        ws.page_setup.orientation = "portrait" if portrait else "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        if title_rows:
            ws.print_title_rows = title_rows
        ws.oddHeader.left.text = self.title
        ws.oddHeader.left.size = 8
        ws.oddHeader.left.color = brand.HEADER_BLUE
        ws.oddHeader.right.text = "BDO Insure | " + brand.SYSTEM
        ws.oddHeader.right.size = 8
        ws.oddHeader.right.color = brand.MUTED
        ws.oddFooter.left.text = brand.FOOTER_TEXT
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.right.text = f"Version {self.version}"
        for part in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
            part.size = 8
            part.color = brand.MUTED

    def _title(self, ws, title: str, subtitle: str = "", back: bool = True, width_cols: int = 6) -> None:
        ws.sheet_view.showGridLines = False
        ws["A1"] = title
        ws["A1"].font = font(15, True, brand.HEADER_BLUE)
        ws.row_dimensions[1].height = 26
        if back:
            ws["A2"] = BACK if self.start == START else f"← Back to {self.start}"
            set_link(ws["A2"], self.start)
            ws["A2"].font = font(9, False, brand.CTA_BLUE, underline=True)
        if subtitle:
            cell = ws.cell(row=2, column=2 if back else 1, value=subtitle)
            cell.font = font(9, False, brand.MUTED)

    def _status_colours(self, ws, rng: str, values: Sequence[str]) -> None:
        for v in values:
            if v in STATUS_STYLE:
                fill, colour = STATUS_STYLE[v]
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=[f'"{v}"'], stopIfTrue=True,
                    fill=PatternFill("solid", fgColor=fill, bgColor=fill), font=Font(color=colour, bold=True)))

    def _list_validation(self, ws, rng: str, ref: RefList, title: str, prompt: str = "") -> None:
        ref.used_in.append(f"{ws.title}: {title}") if f"{ws.title}: {title}" not in ref.used_in else None
        dv = DataValidation(type="list", formula1=f"={ref.range_name}", allow_blank=True, showErrorMessage=True,
                            errorStyle="stop" if ref.strict else "warning", errorTitle=title[:32],
                            error=(f"Choose a value of the list {ref.name} (sheet {LISTS})." if ref.strict else
                                   f"The value is not in the list {ref.name} (sheet {LISTS}). Keep it only for a new "
                                   "value you also add to its template.")[:225],
                            showInputMessage=bool(prompt), promptTitle=title[:32], prompt=prompt[:250])
        ws.add_data_validation(dv)
        dv.add(rng)

    # ------------------------------------------------------------------ template sheets

    def plan_all(self, templates: Sequence[Template]) -> None:
        for t in templates:
            self.plans[t.id] = plan(t)

    def guide_texts(self, t: Template, c: GuideColumn) -> tuple[str, str, str, str]:
        """(allowed values text, its hyperlink, format text, list code of the drop-down) of a column."""
        kind, _, arg = c.check.partition(":")
        fmt, allowed, target, list_code = c.format, "Free entry", "", ""
        if kind == "yn":
            allowed, list_code = "Y or N", "YES_NO"
        elif kind == "list":
            ref = self.lists[arg]
            list_code = arg
            if len(ref.values) <= 6 and sum(len(v) for v, _ in ref.values) <= 70:
                allowed = f"{', '.join(v for v, _ in ref.values)} (list {ref.name})"
            else:
                allowed = f"List {ref.name}: {len(ref.values)} values, see {LISTS}"
            target = f"{LISTS}!A{ref.row}" if ref.row else ""
        elif kind == "code":
            tid, _, col = arg.partition("/")
            p = self.plans[tid]
            allowed = f"A code of {p.template.id} {p.template.name}, column {col}"
            target = f"{p.sheet}!{p.letters[col]}{p.header_row}"
        elif kind == "date":
            allowed = "A date, for example 15-Jan-2028"
        elif kind in ("number", "whole"):
            lo, _, hi = arg.partition("-")
            what = "A whole number" if kind == "whole" else "A number"
            allowed = f"{what} from {lo} to {hi}" if arg else what
        elif kind == "text":
            allowed = f"Free entry, up to {arg} characters"
        if c.allowed:
            allowed = c.allowed
        if c.allowed_link:
            target = c.allowed_link
        return allowed, target, fmt, list_code

    def template_sheet(self, t: Template, step: int, steps: int, review_list: str = "") -> Any:
        p = self.plans[t.id]
        ws = self.wb.create_sheet(p.sheet[:31])
        n = len(t.columns)
        last_col = 1 + n
        span = max(last_col, 7)
        # Column widths first: the header block merges across them.
        ws.column_dimensions["A"].width = LABEL_WIDTH
        widths = {}
        for i, c in enumerate(t.columns, start=2):
            w = max(len(c.header) + 5, len(str(c.example)) + 2, 12, min(len(c.format), 60) / 2.2)
            widths[i] = min(40, w)
            ws.column_dimensions[get_column_letter(i)].width = widths[i]
        for i in range(last_col + 1, span + 1):
            ws.column_dimensions[get_column_letter(i)].width = 12
        merged_chars = sum(widths.get(i, 12) for i in range(2, span + 1))

        self._title(ws, f"{t.id} {t.name}", f"Step {step} of {steps}. Fill in one row per record below the "
                                            "column headers; the guide above each column says what to enter.")
        ws.cell(row=1, column=span, value=f"Step {step} of {steps}").font = font(10, True, brand.CTA_BLUE)
        ws.cell(row=1, column=span).alignment = Alignment(horizontal="right")
        if self.questions:
            q = ws.cell(row=2, column=span, value="Ask a question →")
            set_link(q, self.questions)
            q.font = font(9, False, brand.CTA_BLUE, underline=True)
            q.alignment = Alignment(horizontal="right")
        ws.cell(row=2, column=2).alignment = Alignment(wrap_text=False)

        # Header block
        values = {"Purpose": t.purpose, "Provided by": t.owner, "Due": t.due, "How it is loaded": t.load,
                  "Source": t.source or "-"}
        row = 3
        for label in HEADER_BLOCK + [x.split(":", 1)[0] for x in t.notes]:
            ws.cell(row=row, column=1, value=label).font = font(9, True, brand.HEADER_BLUE)
            ws.cell(row=row, column=1).fill = LABEL_FILL
            ws.cell(row=row, column=1).alignment = Alignment(vertical="top", wrap_text=True)
            if label == "Depends on":
                deps = [self.plans[d] for d in t.depends if d in self.plans]
                if not deps:
                    ws.cell(row=row, column=2, value="Nothing: this template can be filled in first").font = font(9)
                    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=span)
                else:
                    # One linked cell per template, the last one spans the rest of the block.
                    col = 2
                    for i, d in enumerate(deps):
                        end = span if i == len(deps) - 1 else min(span, col + max(1, math.ceil(24 / widths.get(col, 12))) - 1)
                        cell = ws.cell(row=row, column=col, value=f"{d.template.id} {d.template.name} →")
                        set_link(cell, d.sheet)
                        cell.font = font(9, False, brand.CTA_BLUE, underline=True)
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        if end > col:
                            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end)
                        col = end + 1
                        if col > span:
                            break
                    ws.row_dimensions[row].height = 26 if any(len(f"{d.template.id} {d.template.name}") > 24
                                                              for d in deps) else 15
            else:
                text = values.get(label) or next((x.split(":", 1)[1].strip() for x in t.notes
                                                  if x.split(":", 1)[0] == label), "")
                cell = ws.cell(row=row, column=2, value=text)
                cell.font = font(9, label == "Due", brand.NEAR_BLACK if label == "Due" else brand.TEXT)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=span)
                lines = max(1, math.ceil(len(str(text)) / max(merged_chars * 1.2, 20)))
                ws.row_dimensions[row].height = max(15, 12.5 * lines + 3)
            ws.row_dimensions[row].outlineLevel = 1
            row += 1
        # Progress and confirmation
        assert row == p.progress_row, (row, p.progress_row)
        used, missing = self._progress_formulas(p)
        cells = [(1, "Progress", True), (2, "Rows entered", False), (3, used, False), (4, "Mandatory cells missing",
                                                                                         False), (5, missing, False)]
        for col, v, bold in cells:
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = font(9, bold or col in (3, 5), brand.HEADER_BLUE if col == 1 else brand.NEAR_BLACK)
            cell.alignment = Alignment(vertical="center", horizontal="center" if col in (3, 5) else None,
                                       wrap_text=col in (2, 4))
            if col == 1:
                cell.fill = LABEL_FILL
            if col in (3, 5):
                cell.border = BORDER
        ws.conditional_formatting.add(f"E{row}", CellIsRule(operator="greaterThan", formula=["0"],
                                                             fill=MISSING_FILL, font=Font(color=brand.DANGER, bold=True)))
        ws.row_dimensions[row].height = 26
        row += 1
        who, when = t.confirm
        for col, v in ((1, who), (4, when)):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = font(9, col == 1, brand.HEADER_BLUE if col == 1 else brand.NEAR_BLACK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col == 1:
                cell.fill = LABEL_FILL
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        for col in (2, 3, 5):
            ws.cell(row=row, column=col).border = BORDER
            ws.cell(row=row, column=col).fill = INPUT_FILL
        ws.cell(row=row, column=5).number_format = DATE_FORMAT
        dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True,
                            showErrorMessage=True, errorTitle=when[:32], error="Enter a date (dd-mmm-yyyy).",
                            showInputMessage=True, promptTitle=when[:32], prompt="Date the owner confirmed the content")
        ws.add_data_validation(dv)
        dv.add(f"E{row}")
        ws.row_dimensions[row].height = 26
        row += 1
        for c in range(1, span + 1):
            ws.cell(row=row, column=c).border = Border(bottom=Side(style="thin", color=brand.YELLOW))
        ws.row_dimensions[row].height = 6
        row += 1

        # Guide band
        assert row == p.guide_top
        labels = t.extra_rows + GUIDE_ROWS
        texts: dict[str, list[tuple[str, str]]] = {label: [] for label in labels}
        list_codes = []
        for c in t.columns:
            allowed, target, fmt, list_code = self.guide_texts(t, c)
            list_codes.append(list_code)
            texts["Mandatory"].append((mandatory_kind(c.mandatory)[1], ""))
            texts["Format"].append((fmt or "-", ""))
            texts["Allowed values"].append((allowed, target))
            texts["What to enter"].append((c.what, ""))
            for label in t.extra_rows:
                texts[label].append((c.extra.get(label, ""), ""))
        for label in labels:
            is_input = label in t.input_rows
            head = ws.cell(row=row, column=1, value=label)
            head.font = font(9, True, brand.HEADER_BLUE)
            head.fill = LABEL_FILL
            head.alignment = Alignment(vertical="top", wrap_text=True)
            head.border = BORDER
            max_lines = 1
            for i, (text, target) in enumerate(texts[label]):
                col = 2 + i
                cell = ws.cell(row=row, column=col, value=text or None)
                kind = mandatory_kind(t.columns[i].mandatory)[0] if label == "Mandatory" else ""
                cell.font = font(9, kind == "yes", brand.DANGER if kind == "yes" else brand.TEXT)
                if target:
                    set_link(cell, target)
                    cell.font = font(9, False, brand.CTA_BLUE, underline=True)
                cell.fill = INPUT_FILL if is_input else (MANDATORY_FILL if kind == "yes" else GUIDE_FILL)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = BORDER
                max_lines = max(max_lines, math.ceil(len(str(text)) / max(widths[col] * 1.2 - 1, 6)))
            if is_input:
                code = t.extra_lists.get(label)
                if code:
                    self._list_validation(ws, f"B{row}:{get_column_letter(last_col)}{row}", self.lists[code], label,
                                          f"Your review of the column: {', '.join(v for v, _ in self.lists[code].values)}")
                max_lines = max(max_lines, 2)
            ws.row_dimensions[row].height = min(13 * max_lines + 4, 150)
            row += 1

        # Header row
        assert row == p.header_row
        head = ws.cell(row=row, column=1, value="Row")
        head.font = font(10, True, brand.WHITE)
        head.fill = HEADER_FILL
        head.border = BORDER
        head.alignment = Alignment(vertical="center")
        for i, c in enumerate(t.columns, start=2):
            kind, mtext = mandatory_kind(c.mandatory)
            allowed = texts["Allowed values"][i - 2][0]
            cell = ws.cell(row=row, column=i, value=c.header + (" *" if kind == "yes" else ""))
            cell.font = font(10, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
            note = f"{c.header}\n\n{c.what}\n\nMandatory: {mtext}\nFormat: {c.format or '-'}\nAllowed values: {allowed}"
            if str(c.example).strip() not in ("-", ""):
                note += f"\nExample: {c.example}"
            comment = Comment(note, brand.VENDOR)
            comment.width, comment.height = 320, min(60 + 14 * math.ceil(len(note) / 45), 400)
            cell.comment = comment
        ws.row_dimensions[row].height = max(30, 13 * max(math.ceil(len(c.header) * 1.1 / widths[i])
                                                          for i, c in enumerate(t.columns, start=2)) + 6)
        row += 1

        # Example row
        assert row == p.example_row
        cell = ws.cell(row=row, column=1, value=EXAMPLE_LABEL)
        cell.font = font(8, False, brand.PLACEHOLDER, italic=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = EXAMPLE_FILL
        for i, c in enumerate(t.columns, start=2):
            v = example_value(c)
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = font(10, False, brand.PLACEHOLDER, italic=True)
            cell.fill = EXAMPLE_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, _dt.date):
                cell.number_format = DATE_FORMAT
            elif c.text_cell and not c.check.startswith(("number", "whole", "date")):
                cell.number_format = "@"
        ws.row_dimensions[row].height = max(18, 13 * max(math.ceil(len(str(c.example)) / widths[i])
                                                          for i, c in enumerate(t.columns, start=2)) + 4)

        # Input rows: number formats (text cells keep leading zeros), then the table, validations and highlights.
        first, last = p.example_row, p.last_row
        for i, c in enumerate(t.columns, start=2):
            kind = c.check.split(":", 1)[0]
            fmt = DATE_FORMAT if kind == "date" else ("General" if kind in ("number", "whole") or not c.text_cell
                                                      else "@")
            if fmt != "General":
                for r in range(first + 1, last + 1):
                    ws.cell(row=r, column=i).number_format = fmt
        table = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", f"T_{t.id}"),
                      ref=f"B{p.header_row}:{get_column_letter(last_col)}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
        for i, c in enumerate(t.columns, start=2):
            letter = get_column_letter(i)
            rng = f"{letter}{first}:{letter}{last}"
            kind, _, arg = c.check.partition(":")
            mkind, mtext = mandatory_kind(c.mandatory)
            prompt = f"{mtext}. {c.what}"
            allowed = texts["Allowed values"][i - 2][0]
            if allowed != "Free entry":
                prompt += f" Allowed: {allowed}"
            if list_codes[i - 2]:
                self._list_validation(ws, rng, self.lists[list_codes[i - 2]], c.header, prompt)
            elif kind == "code":
                tid, _, col = arg.partition("/")
                name = self.code_range_name(tid, col)
                dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True, showErrorMessage=True,
                                    errorStyle="warning", errorTitle=c.header[:32],
                                    error=f"The code is not in {tid}, column {col}. Enter it there first, or keep it "
                                          "if it is a code BIBS already holds."[:225],
                                    showInputMessage=True, promptTitle=c.header[:32], prompt=prompt[:250])
                ws.add_data_validation(dv)
                dv.add(rng)
            elif kind == "date":
                dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(1900,1,1)", allow_blank=True,
                                    showErrorMessage=True, errorTitle=c.header[:32],
                                    error="Enter a date, for example 15-Jan-2028.", showInputMessage=True,
                                    promptTitle=c.header[:32], prompt=prompt[:250])
                ws.add_data_validation(dv)
                dv.add(rng)
            elif kind in ("number", "whole"):
                lo, _, hi = arg.partition("-")
                kw = ({"operator": "between", "formula1": lo, "formula2": hi} if arg else
                      {"operator": "greaterThanOrEqual", "formula1": "-1E+15"})
                dv = DataValidation(type="decimal" if kind == "number" else "whole", allow_blank=True,
                                    showErrorMessage=True, errorTitle=c.header[:32],
                                    error=f"Enter {allowed.lower()}."[:225], showInputMessage=True,
                                    promptTitle=c.header[:32], prompt=prompt[:250], **kw)
                ws.add_data_validation(dv)
                dv.add(rng)
            elif kind == "text":
                dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=arg, allow_blank=True,
                                    showErrorMessage=True, errorTitle=c.header[:32],
                                    error=f"At most {arg} characters.", showInputMessage=True,
                                    promptTitle=c.header[:32], prompt=prompt[:250])
                ws.add_data_validation(dv)
                dv.add(rng)
            else:
                dv = DataValidation(allow_blank=True, showInputMessage=True, promptTitle=c.header[:32],
                                    prompt=prompt[:250])
                ws.add_data_validation(dv)
                dv.add(rng)
            if mkind == "yes":
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f"AND(LEN({self._concat(p, first)})>0,{letter}{first}=\"\")"], fill=MISSING_FILL))

        ws.freeze_panes = ws.cell(row=p.example_row, column=3)
        ws.print_area = f"A1:{get_column_letter(span)}{min(last, p.example_row + 40)}"
        self._print_setup(ws, f"{p.guide_top}:{p.header_row}")
        ws.sheet_properties.outlinePr.summaryBelow = False
        ws.sheet_view.zoomScale = 90
        ws.sheet_properties.tabColor = brand.HEADER_BLUE
        self.register(ws, f"{t.id} {t.name}: {t.purpose}"[:250], 1)
        return ws

    def code_range_name(self, tid: str, col: str) -> str:
        name = "CODES_" + re.sub(r"[^A-Z0-9]+", "_", f"{tid} {col}".upper()).strip("_")
        if name not in self.wb.defined_names:
            p = self.plans[tid]
            self.wb.defined_names[name] = DefinedName(name, attr_text=p.column_range(col))
        return name

    def _concat(self, p: Plan, row: int | None = None) -> str:
        """Text of a whole row of the template (relative row), or of the input range (row None)."""
        if row is not None:
            return "&".join(f"${p.letters[c.header]}{row}" for c in p.template.columns)
        return "&".join(f"${p.letters[c.header]}${p.example_row}:${p.letters[c.header]}${p.last_row}"
                        for c in p.template.columns)

    def _progress_formulas(self, p: Plan) -> tuple[str, str]:
        """Rows entered and mandatory cells missing; the example row counts only once it is overwritten."""
        key = f"${p.letters[p.key]}${p.example_row}"
        ex = p.key_example
        if isinstance(ex, _dt.date):
            ex_lit = f"DATE({ex.year},{ex.month},{ex.day})"
        elif isinstance(ex, (int, float)):
            ex_lit = repr(ex)
        else:
            ex_lit = '"' + str(ex).replace('"', '""') + '"'
        rows = f"${p.letters[p.key]}${p.example_row}:${p.letters[p.key]}${p.last_row}"
        counted = f"((ROW({rows})>{p.example_row})+({key}<>{ex_lit})>0)"
        used = f"(LEN({self._concat(p)})>0)"
        entered = f"=SUMPRODUCT({counted}*{used})"
        mand = [c for c in p.template.columns if mandatory_kind(c.mandatory)[0] == "yes"]
        if not mand:
            return entered, 0
        blanks = "+".join(f"(${p.letters[c.header]}${p.example_row}:${p.letters[c.header]}${p.last_row}=\"\")"
                          for c in mand)
        return entered, f"=SUMPRODUCT({counted}*{used}*({blanks}))"

    # ------------------------------------------------------------------ Start here

    def start_sheet(self, title: str, subtitle: str, identity: Sequence[tuple[str, str]], steps: Sequence[str],
                    index_title: str, columns: Sequence[tuple[str, float]], rows: Sequence[dict[str, Any]],
                    status_col: str = "Status", statuses: Sequence[str] = STATUSES,
                    after: Callable[[Any, int], int] | None = None, sheet: str | None = None) -> Any:
        """The Start here sheet. rows: {header: value}; a value may be a (text, link) tuple for a hyperlink and a
        string starting with = for a formula. after(ws, row) may add more blocks and returns the next row."""
        ws = self.wb.create_sheet(sheet or self.start, 0)
        ws.sheet_view.showGridLines = False
        ncol = len(columns)
        for i, (_, w) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        total = sum(w for _, w in columns)
        if brand.BDO_LOGO.exists():
            img = XlImage(str(brand.BDO_LOGO))
            img.width, img.height = 190, 38
            ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 34
        if brand.IORTA_LOGO.exists():
            img = XlImage(str(brand.IORTA_LOGO))
            img.width, img.height = 92, 27
            ws.add_image(img, f"{get_column_letter(ncol)}1")
        for c in range(1, ncol + 1):
            ws.cell(row=2, column=c).border = Border(bottom=Side(style="medium", color=brand.YELLOW))
        ws.row_dimensions[2].height = 6
        ws["A3"] = title
        ws["A3"].font = font(18, True, brand.HEADER_BLUE)
        ws.row_dimensions[3].height = 28
        ws["A4"] = subtitle
        ws["A4"].font = font(11, False, brand.NEAR_BLACK)
        row = 6
        for label, value in identity:
            ws.cell(row=row, column=1, value=label).font = font(9, True, brand.HEADER_BLUE)
            ws.cell(row=row, column=1).fill = LABEL_FILL
            ws.cell(row=row, column=1).border = BORDER
            if ncol > 2:
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            cell = ws.cell(row=row, column=3 if ncol > 2 else 2, value=value)
            cell.font = font(10, label in ("Version",), brand.NEAR_BLACK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            end = min(ncol, 8)
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=end)
            for c in range(3, end + 1):
                ws.cell(row=row, column=c).border = BORDER
            row += 1
        row += 1
        ws.cell(row=row, column=1, value="How to fill in this workbook").font = font(12, True, brand.HEADER_BLUE)
        row += 1
        for i, s in enumerate(steps, start=1):
            ws.cell(row=row, column=1, value=i).font = font(11, True, brand.CTA_BLUE)
            ws.cell(row=row, column=1).alignment = Alignment(horizontal="right", vertical="top")
            cell = ws.cell(row=row, column=2, value=s)
            cell.font = font(10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=ncol)
            ws.row_dimensions[row].height = max(15, 13 * math.ceil(len(s) / max(total - columns[0][1], 30)) + 3)
            row += 1
        row += 1
        ws.cell(row=row, column=1, value=index_title).font = font(12, True, brand.HEADER_BLUE)
        row += 1
        header_row = row
        for c, (h, _) in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = font(9, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
        ws.row_dimensions[row].height = 30
        row += 1
        first = row
        for i, r in enumerate(rows):
            lines = 1
            for c, (h, w) in enumerate(columns, start=1):
                v = r.get(h)
                target = None
                if isinstance(v, tuple):
                    v, target = v
                cell = ws.cell(row=row, column=c, value=v)
                cell.font = font(9, h in ("Step", "ID"), brand.TEXT)
                cell.alignment = Alignment(wrap_text=True, vertical="top",
                                           horizontal="center" if isinstance(v, (int, float)) or
                                           (isinstance(v, str) and v.startswith("=")) else None)
                cell.border = BORDER
                if i % 2 == 1:
                    cell.fill = LABEL_FILL
                if target:
                    set_link(cell, target)
                    cell.font = font(9, True, brand.CTA_BLUE, underline=True)
                if isinstance(v, str) and not v.startswith("="):
                    lines = max(lines, math.ceil(len(v) / max(w * 1.2 - 1, 4)))
            ws.row_dimensions[row].height = min(13 * lines + 4, 120)
            row += 1
        last = max(row - 1, first)
        if status_col in [h for h, _ in columns]:
            letter = get_column_letter([h for h, _ in columns].index(status_col) + 1)
            rng = f"{letter}{first}:{letter}{last}"
            self._list_validation(ws, rng, self.list_for_values("Status", statuses), status_col,
                                  "Status of the template: " + ", ".join(statuses))
            self._status_colours(ws, rng, statuses)
        for h in ("Mandatory cells missing",):
            if h in [x for x, _ in columns]:
                letter = get_column_letter([x for x, _ in columns].index(h) + 1)
                ws.conditional_formatting.add(f"{letter}{first}:{letter}{last}", CellIsRule(
                    operator="greaterThan", formula=["0"], fill=MISSING_FILL, font=Font(color=brand.DANGER, bold=True)))
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncol)}{last}"
        row = last + 2
        if after:
            row = after(ws, row)
        ws.cell(row=row, column=1, value=f"Prepared by {brand.VENDOR} for {brand.CLIENT_SHORT}. {brand.FOOTER_TEXT}.") \
            .font = font(8, False, brand.MUTED)
        ws.freeze_panes = None
        self._print_setup(ws, f"{header_row}:{header_row}")
        ws.sheet_properties.tabColor = brand.YELLOW
        self.register(ws, "Who provides what, the steps and the templates in the order they are filled in", len(rows))
        return ws

    def small_table(self, ws, row: int, title: str, headers: Sequence[str], rows: Sequence[Sequence[Any]],
                    spans: Sequence[int] | None = None) -> int:
        """A small table under the index (milestones, other sheets); returns the next free row."""
        ws.cell(row=row, column=1, value=title).font = font(12, True, brand.HEADER_BLUE)
        row += 1
        spans = list(spans or [1] * len(headers))
        for data, is_head in [(headers, True)] + [(r, False) for r in rows]:
            col = 1
            for v, span in zip(data, spans):
                target = None
                if isinstance(v, tuple):
                    v, target = v
                cell = ws.cell(row=row, column=col, value=v)
                cell.font = font(9, is_head, brand.WHITE if is_head else brand.TEXT)
                if is_head:
                    cell.fill = HEADER_FILL
                if target:
                    set_link(cell, target) if not target.startswith("#") else target
                    cell.font = font(9, True, brand.CTA_BLUE, underline=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                for c in range(col, col + span):
                    ws.cell(row=row, column=c).border = BORDER
                if span > 1:
                    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
                col += span
            row += 1
        return row + 1

    # ------------------------------------------------------------------ reference sheets (tables)

    def table_sheet(self, name: str, title: str, description: str, columns: Sequence[Any],
                    rows: Sequence[dict[str, Any]], statuses: Sequence[str] = STATUSES) -> Any:
        """A reference sheet of the workbook (bdoi_xlsx.Column definitions), with a back link, drop-downs from the
        Reference lists sheet and the status colours."""
        ws = self.wb.create_sheet(name[:31])
        cols = list(columns)
        self._title(ws, title, description)
        header = 4
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=header, column=c, value=col.header)
            cell.font = font(9, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
            if col.description:
                cell.comment = Comment(f"{col.header}: {col.description}", brand.VENDOR)
                cell.comment.width, cell.comment.height = 260, 90
            ws.column_dimensions[get_column_letter(c)].width = col.width
        ws.row_dimensions[header].height = 30
        for r, row in enumerate(rows, start=header + 1):
            for c, col in enumerate(cols, start=1):
                v = row.get(col.key)
                target = None
                if isinstance(v, tuple):
                    v, target = v
                if isinstance(v, (list, tuple)):
                    v = "\n".join(f"• {x}" for x in v)
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = font(9, False, brand.TEXT)
                cell.alignment = Alignment(wrap_text=col.wrap, vertical="top")
                cell.border = BORDER
                if target:
                    set_link(cell, target)
                    cell.font = font(9, False, brand.CTA_BLUE, underline=True)
                if col.kind == "date":
                    cell.number_format = DATE_FORMAT
        first, last = header + 1, header + max(len(rows), 1)
        for c, col in enumerate(cols, start=1):
            letter = get_column_letter(c)
            rng = f"{letter}{first}:{letter}{last}"
            if col.values:
                self._list_validation(ws, rng, self.list_for_values(col.header, list(col.values)), col.header,
                                      col.description)
                self._status_colours(ws, rng, list(col.values))
            elif col.kind == "date":
                dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True,
                                    showErrorMessage=True, errorTitle=col.header[:32], error="Enter a date.")
                ws.add_data_validation(dv)
                dv.add(rng)
        ws.conditional_formatting.add(f"A{first}:{get_column_letter(len(cols))}{last}", FormulaRule(
            formula=[f"AND(MOD(ROW(),2)=0,$A{first}<>\"\")"], fill=PatternFill("solid", fgColor=brand.BG_BLUE,
                                                                                 bgColor=brand.BG_BLUE)))
        ws.freeze_panes = ws.cell(row=first, column=2)
        ws.auto_filter.ref = f"A{header}:{get_column_letter(len(cols))}{last}"
        self._print_setup(ws, f"{header}:{header}")
        ws.sheet_view.zoomScale = 90
        self.register(ws, description, len(rows))
        return ws

    # ------------------------------------------------------------------ Reference lists and questions

    def lists_sheet(self) -> Any:
        ws = self.wb.create_sheet(LISTS)
        self._title(ws, "Reference lists", "Every list of allowed values used by the drop-downs, with its code and "
                                            "label. A drop-down accepts the codes of its list.")
        widths = [30, 34, 26, 44, 60]
        heads = ["List", "Used in", "Code", "Label", "Note"]
        for c, (h, w) in enumerate(zip(heads, widths), start=1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font = font(9, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(c)].width = w
        row = 5
        for i, ref in enumerate(self.lists.values()):
            assert ref.row == row, (ref.code, ref.row, row)
            fill = LABEL_FILL if i % 2 == 0 else PatternFill(fill_type=None)
            for j, (code, label) in enumerate(ref.values):
                values = [f"{ref.name} ({ref.code})" if j == 0 else None,
                          "\n".join(ref.used_in) if j == 0 else None, code, label or None,
                          (ref.note or None) if j == 0 else None]
                lines = 1
                for c, v in enumerate(values, start=1):
                    cell = ws.cell(row=row + j, column=c, value=v)
                    cell.font = font(9, c in (1, 3), brand.HEADER_BLUE if c == 1 else brand.TEXT)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.fill = fill
                    cell.border = BORDER
                    if c == 3:
                        cell.number_format = "@"
                    if v:
                        lines = max(lines, sum(math.ceil(max(len(x), 1) / (widths[c - 1] * 1.2))
                                               for x in str(v).split("\n")))
                if lines > 1:
                    ws.row_dimensions[row + j].height = 12.5 * lines + 3
            end = row + len(ref.values) - 1
            self.wb.defined_names[ref.range_name] = DefinedName(
                ref.range_name, attr_text=f"{quote_sheetname(LISTS)}!$C${row}:$C${end}")
            row = end + 1
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:E{row - 1}"
        self._print_setup(ws, "4:4")
        self.register(ws, "Every list of allowed values used by the drop-downs (code and label)", row - 5)
        return ws

    def questions_sheet(self, templates: Sequence[Template], rows: int = 100) -> Any:
        ws = self.wb.create_sheet(self.questions)
        self._title(ws, "Questions and comments", "Raise a question or a comment on a template or a column; the "
                                                   f"{brand.VENDOR} project team answers here.")
        cols = [("No.", 6), ("Template", 30), ("Column", 26), ("Question / comment", 50), ("Raised by", 20),
                ("Date", 13), ("Answer", 50), ("Answered by", 20), ("Status", 12)]
        for c, (h, w) in enumerate(cols, start=1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font = font(9, True, brand.WHITE)
            cell.fill = HEADER_FILL
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = w
        for r in range(5, 5 + rows):
            ws.cell(row=r, column=1, value=r - 4).font = font(9, False, brand.MUTED)
            for c in range(1, len(cols) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if c > 1:
                    cell.font = font(9)
            ws.cell(row=r, column=6).number_format = DATE_FORMAT
        last = 4 + rows
        names = self.add_list("TEMPLATES", "Templates of this workbook",
                              [(f"{t.id} {t.name}", "") for t in templates] + [("General", "")])
        self._list_validation(ws, f"B5:B{last}", names, "Template", "The template the question is about")
        dv = DataValidation(allow_blank=True, showInputMessage=True, promptTitle="Column",
                            prompt="The column header as on the template sheet; blank for the whole template")
        ws.add_data_validation(dv)
        dv.add(f"C5:C{last}")
        dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True,
                            showErrorMessage=True, errorTitle="Date", error="Enter a date.")
        ws.add_data_validation(dv)
        dv.add(f"F5:F{last}")
        status = self.list_for_values("Question status", QUESTION_STATUSES)
        self._list_validation(ws, f"I5:I{last}", status, "Status", "Open, Answered or Closed")
        self._status_colours(ws, f"I5:I{last}", QUESTION_STATUSES)
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:I{last}"
        self._print_setup(ws, "4:4")
        self.register(ws, "Questions and comments per template and column, with the answer", 0)
        return ws

    # ------------------------------------------------------------------ save

    def save(self, path: Path, order: Sequence[str], properties: dict[str, str]) -> Path:
        self.wb._sheets = [self.wb[name] for name in order]
        self.wb.active = 0
        for ws in self.wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == order[0]
        props = self.wb.properties
        props.title = properties.get("title", self.title)
        props.creator = brand.VENDOR
        props.subject = brand.SYSTEM
        props.keywords = properties.get("keywords", "")
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        return path


# ============================================================================ verification (used by the --check of the builders)


def verify(path: Path, expect_templates: Sequence[str] = ()) -> list[str]:
    """Problems of a written guided workbook: a Columns sheet, a hidden sheet, a template column without its guide
    cells, a drop-down whose list does not resolve, a broken internal hyperlink, a restricted word."""
    import codecs

    from openpyxl import load_workbook

    problems: list[str] = []
    wb = load_workbook(path)
    names = set(wb.sheetnames)
    restricted = re.compile(r"\b(" + codecs.decode(
        "qrzb|qhzzl|snxr|fnzcyr qngn|cebgbglcr|fnaqobk|yberz vcfhz|gbqb|svkzr|pynhqr|naguebcvp|pungtcg|bcranv|"
        "pbcvybg|tvguho", "rot13") + r"|json|sql|api)\b", re.I)
    for ws in wb.worksheets:
        if ws.title.endswith("Columns") or ws.title == "Template columns":
            problems.append(f"{path.name}: sheet {ws.title} (column descriptions belong on the template sheet)")
        if ws.sheet_state != "visible" and ws.title != "_lists":
            problems.append(f"{path.name}: hidden sheet {ws.title}")
    for dn_name, dn in wb.defined_names.items():
        for sheet, rng in dn.destinations:
            if sheet not in names:
                problems.append(f"{path.name}: named range {dn_name} points to the missing sheet {sheet}")
                continue
            cells = wb[sheet][rng.replace("$", "")]
            if dn_name.startswith("LIST_") and not any(c.value not in (None, "") for row in cells for c in row):
                problems.append(f"{path.name}: named range {dn_name} is empty")
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            if dv.type == "list":
                f = (dv.formula1 or "").lstrip("=")
                if f.startswith('"'):
                    continue
                if f not in wb.defined_names and "!" not in f:
                    problems.append(f"{path.name}: {ws.title} {dv.sqref}: list {f} does not resolve")
                elif "!" in f and f.split("!")[0].strip("'") not in names:
                    problems.append(f"{path.name}: {ws.title} {dv.sqref}: list {f} does not resolve")
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink and c.hyperlink.location is None and str(c.hyperlink.target or "").startswith("#"):
                    pass
                target = c.hyperlink.location if c.hyperlink else None
                if c.hyperlink and not target and c.hyperlink.target and str(c.hyperlink.target).startswith("#"):
                    target = str(c.hyperlink.target)[1:]
                if target:
                    sheet = target.split("!")[0].strip("'").replace("''", "'")
                    if sheet not in names:
                        problems.append(f"{path.name}: {ws.title}!{c.coordinate}: link to missing sheet {sheet}")
                if isinstance(c.value, str) and restricted.search(c.value):
                    problems.append(f"{path.name}: {ws.title}!{c.coordinate}: restricted word "
                                    f"'{restricted.search(c.value).group(0)}'")
    for sheet in expect_templates:
        if sheet not in names:
            problems.append(f"{path.name}: template sheet {sheet} missing")
            continue
        ws = wb[sheet]
        labels = {str(ws.cell(r, 1).value): r for r in range(1, 60) if ws.cell(r, 1).value}
        if not all(g in labels for g in GUIDE_ROWS) or "Row" not in labels:
            problems.append(f"{path.name}: {sheet}: guide band or header missing")
            continue
        header = labels["Row"]
        if [labels[g] for g in GUIDE_ROWS] != list(range(header - 4, header)):
            problems.append(f"{path.name}: {sheet}: the guide rows are not directly above the header")
        col = 2
        while ws.cell(header, col).value:
            for g in ("Mandatory", "Format", "Allowed values", "What to enter"):
                if ws.cell(labels[g], col).value in (None, ""):
                    problems.append(f"{path.name}: {sheet}: column {ws.cell(header, col).value}: no {g}")
            col += 1
        if ws.cell(header + 1, 1).value != EXAMPLE_LABEL:
            problems.append(f"{path.name}: {sheet}: example row not labelled")
    return problems
