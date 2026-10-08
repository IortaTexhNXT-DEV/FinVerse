"""Prepares generated documents for their screenshot (used by packs/brd01_documents.cjs `render`).

    python3 doc_render.py xlsx <file.xlsx> <options.json>   prepare a workbook for printing its first page
    python3 doc_render.py crop <file.png>                    crop a rendered sheet to its used part
    python3 doc_render.py text <file.txt> <out.png> [dpi]    draw a text file (no wrapping) as an image

Workbook options (JSON):
    keep   names of the columns to show (others are hidden), in any order; empty = all columns
    above  "keep" (default): the rows above the heading row (logo, title, report parameters) stay, with their text
           and images moved to the first column shown; "drop": they are deleted together with their images (a
           template whose guide band would make the rows print too small)
    rows   {"column": name, "equals": value, "max": n}: only the data rows whose column has this value, at most n
           (an extract); the first data row is kept when it is not a data row of that column (e.g. a totals row)
    first  number of data rows to keep at most (an extract of a long list)

Every column shown is as wide as its longest value as printed plus a margin, so that a right-aligned date or amount
never touches the text of the next column. The sheet prints landscape, fitted to the page width.
"""
from __future__ import annotations

import datetime as dt
import json
import sys

MARGIN = 3
MIN_CHARS = 6
MAX_CHARS = 50
ONE_LINE_HEADING = 20


def _clean(value) -> str:
    return str(value or "").replace("*", "").strip()


def _shown(cell) -> str:
    """The value of a cell roughly as Excel prints it."""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return "00-Mon-0000"
    if isinstance(v, (int, float)):
        fmt = cell.number_format or ""
        if "0.00" in fmt:
            return f"{v:,.2f}"
        if "#,##0" in fmt:
            return f"{v:,.0f}"
        return f"{v:g}"
    return str(v)


def _heading_chars(text: str) -> int:
    text = text.strip()
    if len(text) <= ONE_LINE_HEADING:
        return len(text)
    return max((len(w) for w in text.split()), default=0)


def _header_row(ws, keep: list[str]) -> int | None:
    for r in range(1, ws.max_row + 1):
        names = {_clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        if any(k in names for k in keep):
            return r
    return None


def prepare_xlsx(path: str, options: dict) -> None:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import RowBreak

    keep = [k for k in options.get("keep") or [] if k]
    above = options.get("above", "keep")
    rows = options.get("rows") or {}
    first = options.get("first")
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        hr = _header_row(ws, keep) if keep else None
        if keep and hr is None:
            continue
        for m in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(m))
        if hr and hr > 1 and above == "drop":
            ws.delete_rows(1, hr - 1)
            ws._images = []
            hr = 1
        heading = {c: _clean(ws.cell(hr, c).value) for c in range(1, ws.max_column + 1)} if hr else {}
        shown = [c for c in range(1, ws.max_column + 1) if not keep or heading.get(c) in keep]
        first_shown = shown[0] if shown else 1
        if hr and hr > 1 and first_shown > 1:
            # The header block (company, title, parameters) and the logo move to the first column shown.
            for r in range(1, hr):
                for c in range(1, first_shown):
                    cell = ws.cell(r, c)
                    if cell.value not in (None, "") and ws.cell(r, first_shown).value in (None, ""):
                        target = ws.cell(r, first_shown)
                        target.value = cell.value
                        target._style = cell._style
                        cell.value = None
            for image in getattr(ws, "_images", []):
                anchor = getattr(image.anchor, "_from", None)
                if anchor is not None and anchor.col < first_shown - 1:
                    to = getattr(image.anchor, "to", None)
                    if to is not None:
                        to.col += first_shown - 1 - anchor.col
                    anchor.col = first_shown - 1
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(c)].hidden = c not in shown
        start = (hr or 0) + 1
        # A group or total label written in a column left out moves to the first column shown.
        for r in range(start, ws.max_row + 1):
            for c in range(1, first_shown):
                cell = ws.cell(r, c)
                if isinstance(cell.value, str) and cell.value.strip() and ws.cell(r, first_shown).value in (None, ""):
                    target = ws.cell(r, first_shown)
                    target.value = cell.value
                    target.font = cell.font.copy()
                    cell.value = None
        if rows or first:
            column = next((c for c, n in heading.items() if n == rows.get("column")), None) if rows else None
            limit = rows.get("max") if rows else first
            taken = 0
            for r in range(start, ws.max_row + 1):
                if column is not None:
                    match = _clean(ws.cell(r, column).value) == rows.get("equals")
                    # A leading row of another kind (the enabled users of each profile) stays.
                    wanted = match or r == start
                else:
                    match = wanted = True
                if wanted and match:
                    taken += 1
                    wanted = not limit or taken <= limit
                ws.row_dimensions[r].hidden = not wanted
        for c in shown:
            chars = _heading_chars(heading.get(c, "")) if hr else 0
            for r in range(start, ws.max_row + 1):
                if not ws.row_dimensions[r].hidden:
                    chars = max(chars, max((len(x) for x in _shown(ws.cell(r, c)).split("\n")), default=0))
            dim = ws.column_dimensions[get_column_letter(c)]
            if keep or (dim.width or 0) < min(chars + MARGIN, MAX_CHARS):
                dim.width = max(MIN_CHARS, min(chars + MARGIN, MAX_CHARS))
        ws.print_title_rows = None
        ws.print_title_cols = None
        ws.print_area = None
        ws.row_breaks = RowBreak()
        ws.freeze_panes = None
        ws.auto_filter.ref = None
        for part in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter, ws.firstHeader, ws.firstFooter):
            part.left.text = part.center.text = part.right.text = None
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(path)


def crop(path: str) -> None:
    """A spreadsheet fills only the top of the page: keep the used part with a margin."""
    from PIL import Image, ImageOps

    im = Image.open(path).convert("RGB")
    box = ImageOps.invert(im).getbbox()
    if box:
        left, top, right, bottom = box
        im = im.crop((max(0, left - 40), max(0, top - 40), min(im.width, right + 40),
                      min(im.height, max(bottom + 40, top + 120))))
    im.save(path)


def text_image(path: str, out: str, dpi: int = 200) -> None:
    """A text file as printed in a monospace font at 10 pt, one file line per image line, never wrapped.

    A file of delimited records (every line with the same number of | separators) is shown field by field: the field
    names of the first line down the left, each record in a column, so that a long record stays readable.
    """
    from PIL import Image, ImageDraw, ImageFont

    raw = open(path, "rb").read().decode("utf-8", errors="replace").replace("\r\n", "\n").rstrip("\n")
    lines = raw.split("\n")
    size = round(10 * dpi / 72)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", size)
    bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf", size)
    pad = round(size * 1.5)
    line_h = round(size * 1.35)
    char_w = font.getlength("M")
    delimited = len(lines) > 1 and "|" in lines[0] and len({ln.count("|") for ln in lines}) == 1
    if delimited:
        names = lines[0].split("|")
        records = [ln.split("|") for ln in lines[1:]]
        name_w = max(len(n) for n in names)
        value_w = [max(len(rec[i]) for i in range(len(names))) for rec in records]
        width = round(pad * 2 + char_w * (name_w + 3 + sum(w + 3 for w in value_w)))
        height = pad * 2 + line_h * (len(names) + 1)
        im = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(im)
        y = pad
        draw.text((pad, y), "Field", font=bold, fill="black")
        x = pad + char_w * (name_w + 3)
        for i, w in enumerate(value_w):
            draw.text((x, y), f"Record {i + 1}", font=bold, fill="black")
            x += char_w * (w + 3)
        y += line_h
        draw.line((pad, y - round(line_h * 0.15), width - pad, y - round(line_h * 0.15)), fill="#999999", width=2)
        for f, name in enumerate(names):
            draw.text((pad, y), name, font=font, fill="#444444")
            x = pad + char_w * (name_w + 3)
            for i, rec in enumerate(records):
                draw.text((x, y), rec[f], font=font, fill="black")
                x += char_w * (value_w[i] + 3)
            y += line_h
    else:
        width = round(pad * 2 + char_w * max(len(ln) for ln in lines))
        height = pad * 2 + line_h * len(lines)
        im = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(im)
        for i, ln in enumerate(lines):
            draw.text((pad, pad + i * line_h), ln, font=font, fill="black")
    im.save(out, dpi=(dpi, dpi))


if __name__ == "__main__":
    command = sys.argv[1]
    if command == "xlsx":
        prepare_xlsx(sys.argv[2], json.loads(sys.argv[3]))
    elif command == "crop":
        crop(sys.argv[2])
    elif command == "text":
        text_image(sys.argv[2], sys.argv[3], int(sys.argv[4]) if len(sys.argv) > 4 else 200)
    else:
        raise SystemExit(f"unknown command {command}")
