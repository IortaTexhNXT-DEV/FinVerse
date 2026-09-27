"""Build the iNXT BrokerVerse User Manual (Word) on the iorta TechNXT template.

    pip install python-docx pillow pypdfium2      # LibreOffice (soffice) is needed for the PDF pass
    python tools/user-manual/build_user_manual.py [--edition 2.0] [--pdf] [--previews DIR]

The text lives in content.py (edition 1.0) and content_v2.py (edition 2.0, built from the screen inventory in
screens/inventory.json and the screen captures next to it). The builder renders the document twice: the first PDF gives the page of every
heading, which is written into the table of contents (a real Word TOC field, so Update Field in Word refreshes
it). The .docx goes to docs/user-manual/; --pdf also keeps the PDF next to it (not committed).
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "frontend" / "src" / "assets" / "brand" / "iorta-technxt.png"
SCREENS = Path(__file__).resolve().parent / "screens"
OUT_DIR = ROOT / "docs" / "user-manual"
TITLE = "iNXT BrokerVerse"
SUBTITLE = "User Manual"

# Each edition keeps its own text module, so an older edition can still be rebuilt as issued.
EDITIONS = {
    "1.0": {"content": "content", "date": date(2026, 9, 27),
            "history": [("1.0", date(2026, 9, 27), "First issue")]},
    "2.0": {"content": "content_v2", "date": date(2026, 9, 27),
            "history": [("1.0", date(2026, 9, 27), "First issue"),
                        ("2.0", date(2026, 9, 27), "Screen-by-screen edition: every menu screen with its screenshot, "
                                                   "fields, validation messages, actions, rules and statuses")]},
}
VERSION = "2.0"
ISSUE_DATE = EDITIONS[VERSION]["date"]
CHAPTERS: list = []


def use_edition(version: str) -> None:
    global VERSION, ISSUE_DATE, CHAPTERS
    import importlib
    VERSION = version
    ISSUE_DATE = EDITIONS[version]["date"]
    CHAPTERS = importlib.import_module(EDITIONS[version]["content"]).CHAPTERS


def file_name() -> str:
    return f"iNXT_BrokerVerse_User_Manual_v{VERSION}.docx"

# iorta TechNXT palette, taken from the logo: shield blue, cube teal, dark text.
NAVY = "0B4F9C"
BLUE = "0090D8"
SKY = "00A8F0"
TEAL = "00A098"
TEXT = "333333"
MUTED = "6B7280"
BAND = "EEF6FC"
TIP_BG = "E6F5F4"
NOTE_BG = "FFF6E0"
NOTE_BAR = "E0A100"
BORDER = "C9D6E3"
WHITE = "FFFFFF"
FONT = "Arial"

CONTENT_WIDTH_CM = 16.6


# ---------------------------------------------------------------- low-level XML helpers

def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def set_fonts(rpr_parent, name: str = FONT) -> None:
    """Pin every script of a run or style to one font (drops the theme font attributes)."""
    rpr = rpr_parent.get_or_add_rPr() if hasattr(rpr_parent, "get_or_add_rPr") else rpr_parent
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_borders(cell, **edges) -> None:
    """edges: top/left/bottom/right = (size in eighths of a point, hex) or None for no border."""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        borders.append(el)
    tcpr.append(borders)


def cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def table_borders(table, size=4, color=BORDER) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def fixed_layout(table, widths_cm) -> None:
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths_cm):
        col.set(qn("w:w"), str(int(Cm(width).twips)))
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(Cm(sum(widths_cm)).twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def no_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def para_border(paragraph, edge: str, size: int, color: str, space: int = 4) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def add_field(paragraph, instr: str, placeholder: str = "1", size=None, color=None, bold=False):
    """A simple field (PAGE, NUMPAGES) with a cached result."""
    def run_with(el):
        run = paragraph.add_run()
        run._r.append(el)
        return run

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run_with(instr_el)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    result = paragraph.add_run(placeholder)
    style_run(result, size=size, color=color, bold=bold)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def style_run(run, size=None, color=None, bold=None, italic=None) -> None:
    set_fonts(run._r)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_rich(paragraph, text: str, size=None, color=None, bold_color=None) -> None:
    """Text with **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        is_bold = i % 2 == 1
        style_run(run, size=size, color=(bold_color if is_bold and bold_color else color), bold=is_bold or None)


def spacing(paragraph, before=0, after=6, line=None, keep_next=False) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True


# ---------------------------------------------------------------- document styles

def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_fonts(normal.element)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    heading_specs = {
        "Heading 1": (20, NAVY, 0, 10),
        "Heading 2": (13.5, BLUE, 16, 6),
        "Heading 3": (11, NAVY, 10, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        st = styles[name]
        set_fonts(st.element)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.1

    for name, indent in (("TOC 1", 0), ("TOC 2", 0.8)):
        st = styles[name] if name in [s.name for s in styles] else styles.add_style(name, 1)
        st.base_style = styles["Normal"]
        st.paragraph_format.left_indent = Cm(indent)
        st.paragraph_format.space_after = Pt(3 if name == "TOC 2" else 4)
        st.paragraph_format.space_before = Pt(6 if name == "TOC 1" else 0)
        st.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        st.font.bold = name == "TOC 1"
        st.font.color.rgb = rgb(NAVY if name == "TOC 1" else TEXT)

    for name in ("Header", "Footer"):
        st = styles[name]
        set_fonts(st.element)
        st.font.size = Pt(8)
        st.font.color.rgb = rgb(MUTED)

    bullet = styles["List Bullet"]
    set_fonts(bullet.element)
    bullet.font.size = Pt(10.5)


# ---------------------------------------------------------------- page frame

def page_setup(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def clear_default_tabs(paragraph) -> None:
    for position in (Cm(8.255), Cm(16.51)):  # the Header/Footer style tabs of the default template
        paragraph.paragraph_format.tab_stops.add_tab_stop(position, WD_TAB_ALIGNMENT.CLEAR)


def header_footer(section) -> None:
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    clear_default_tabs(p)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run().add_picture(str(LOGO), height=Cm(0.95))
    run = p.add_run(f"\t{TITLE}  |  {SUBTITLE}")
    style_run(run, size=8.5, color=MUTED)
    spacing(p, after=0)
    para_border(p, "bottom", 8, SKY, space=6)

    footer = section.footer
    p = footer.paragraphs[0]
    clear_default_tabs(p)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    para_border(p, "top", 4, BORDER, space=6)
    run = p.add_run(f"© {ISSUE_DATE.year} iorta TechNXT. Confidential.   Version {VERSION}")
    style_run(run, size=8, color=MUTED)
    run = p.add_run("\tPage ")
    style_run(run, size=8, color=MUTED)
    add_field(p, "PAGE", size=8, color=MUTED)
    run = p.add_run(" of ")
    style_run(run, size=8, color=MUTED)
    add_field(p, "NUMPAGES", size=8, color=MUTED)


def cover(doc: Document) -> None:
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.add_run().add_picture(str(LOGO), width=Cm(6.2))
    spacing(p, after=0)

    for _ in range(6):
        spacing(doc.add_paragraph(), after=10)

    p = doc.add_paragraph()
    run = p.add_run(TITLE)
    style_run(run, size=34, color=NAVY, bold=True)
    spacing(p, after=2, line=1.0)

    p = doc.add_paragraph()
    run = p.add_run(SUBTITLE)
    style_run(run, size=22, color=SKY)
    spacing(p, after=14, line=1.0)
    para_border(p, "bottom", 24, TEAL, space=12)

    p = doc.add_paragraph()
    run = p.add_run("A guide for business users of the iNXT BrokerVerse insurance broking platform")
    style_run(run, size=12, color=TEXT)
    spacing(p, before=6, after=4)

    for _ in range(9):
        spacing(doc.add_paragraph(), after=10)

    rows = [
        ("Version", VERSION),
        ("Date", ISSUE_DATE.strftime("%d %B %Y")),
        ("Prepared by", "iorta TechNXT"),
        ("Classification", "Confidential"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    fixed_layout(table, [3.6, 7.0])
    for row, (label, value) in zip(table.rows, rows):
        for cell in row.cells:
            cell_borders(cell)
            cell_margins(cell, top=30, bottom=30, left=0, right=80)
        lp = row.cells[0].paragraphs[0]
        style_run(lp.add_run(label), size=10, color=MUTED)
        spacing(lp, after=0)
        vp = row.cells[1].paragraphs[0]
        style_run(vp.add_run(value), size=10, color=NAVY, bold=True)
        spacing(vp, after=0)

    p = doc.add_paragraph()
    spacing(p, before=28, after=0)
    run = p.add_run("www.iortatechnxt.com")
    style_run(run, size=9.5, color=BLUE)


# ---------------------------------------------------------------- content blocks

class Builder:
    def __init__(self, doc: Document, pages: dict[str, int] | None):
        self.doc = doc
        self.pages = pages or {}
        self.h1 = 0
        self.h2 = 0
        self.headings: list[tuple[int, str]] = []
        self.figures = 0

    # -- front matter

    def document_control(self) -> None:
        self.doc.add_page_break()
        self.front_title("Document control")
        self.table(["Version", "Date", "Description", "Author"],
                   [[v, d.strftime("%d-%b-%Y"), text, "iorta TechNXT"] for v, d, text in EDITIONS[VERSION]["history"]],
                   [2.2, 3.0, 7.4, 4.0])
        self.paragraph("**Distribution**: customer project team and key users.")
        self.paragraph("**Copyright**: this document and its contents are the property of iorta TechNXT and are "
                       "provided to the customer for the use of the iNXT BrokerVerse platform. The screens you see "
                       "may differ slightly from the descriptions here, depending on your organisation's set-up "
                       "and your role.")

    def front_title(self, text: str) -> None:
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        style_run(run, size=20, color=NAVY, bold=True)
        spacing(p, after=12)
        para_border(p, "bottom", 12, TEAL, space=6)

    def contents(self, entries: list[tuple[int, str]]) -> None:
        self.doc.add_page_break()
        self.front_title("Contents")
        paragraphs = []
        for level, text in entries:
            p = self.doc.add_paragraph(style=f"TOC {level}")
            paragraphs.append((p, text))
        first, last = paragraphs[0][0], paragraphs[-1][0]
        # Open the TOC field in the first entry and close it after the last one.
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = ' TOC \\o "1-2" \\h \\z \\u '
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        for el in (begin, instr, sep):
            first.add_run()._r.append(el)
        for p, text in paragraphs:
            page = self.pages.get(text)
            style_run(p.add_run(text))
            style_run(p.add_run(f"\t{page if page else ''}"))
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        last.add_run()._r.append(end)

    # -- body

    def heading(self, level: int, text: str) -> None:
        if level == 1:
            self.h1 += 1
            self.h2 = 0
            number = f"{self.h1}"
            label = f"{number}  {text}"
            self.drop_trailing_spacer()
            p = self.doc.add_paragraph(style="Heading 1")
            p.paragraph_format.page_break_before = True
            style_run(p.add_run(label))
            para_border(p, "bottom", 12, TEAL, space=6)
            self.headings.append((1, label))
        elif level == 2:
            self.h2 += 1
            label = f"{self.h1}.{self.h2}  {text}"
            p = self.doc.add_paragraph(style="Heading 2")
            style_run(p.add_run(label))
            self.headings.append((2, label))
        else:
            p = self.doc.add_paragraph(style="Heading 3")
            style_run(p.add_run(text))

    def drop_trailing_spacer(self) -> None:
        """Remove an empty spacer paragraph left at the end of a chapter (it can spill onto a blank page)."""
        body = self.doc.element.body
        items = [el for el in body if el.tag != qn("w:sectPr")]
        if not items:
            return
        last = items[-1]
        if last.tag == qn("w:p") and not "".join(last.itertext()).strip() and not last.findall(".//" + qn("w:drawing")):
            body.remove(last)

    def paragraph(self, text: str) -> None:
        p = self.doc.add_paragraph()
        add_rich(p, text, bold_color=None)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            add_rich(p, item)
            spacing(p, after=4)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.5)

    def steps(self, items: list[str]) -> None:
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            num = p.add_run(f"{i}.\t")
            style_run(num, color=BLUE, bold=True)
            add_rich(p, item)
            spacing(p, after=4, keep_next=i < len(items))
        spacing(self.doc.paragraphs[-1], after=8)

    def table(self, headers, rows, widths) -> None:
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table_borders(table)
        fixed_layout(table, widths)
        head = table.rows[0]
        repeat_header(head)
        no_split(head)
        for cell, text in zip(head.cells, headers):
            shade(cell, NAVY)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_run(p.add_run(text), size=9.5, color=WHITE, bold=True)
            spacing(p, after=0, line=1.1)
        for r, values in enumerate(rows):
            row = table.rows[r + 1]
            no_split(row)
            for c, (cell, text) in enumerate(zip(row.cells, values)):
                if r % 2 == 1:
                    shade(cell, BAND)
                cell_margins(cell)
                p = cell.paragraphs[0]
                add_rich(p, text, size=9.5, color=NAVY if c == 0 else TEXT)
                if c == 0:
                    for run in p.runs:
                        run.bold = True
                spacing(p, after=0, line=1.15)
        after = self.doc.add_paragraph()
        spacing(after, after=4, line=1.0)

    def info(self, rows) -> None:
        """A small two-column panel without a header row (for "Used by" and "Where to find it")."""
        table = self.doc.add_table(rows=len(rows), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table_borders(table, color=BORDER)
        fixed_layout(table, [4.2, 12.4])
        for row, (label, value) in zip(table.rows, rows):
            no_split(row)
            shade(row.cells[0], BAND)
            for cell in row.cells:
                cell_margins(cell, top=60, bottom=60)
            p = row.cells[0].paragraphs[0]
            style_run(p.add_run(label), size=9.5, color=NAVY, bold=True)
            spacing(p, after=0, line=1.1)
            p = row.cells[1].paragraphs[0]
            add_rich(p, value, size=9.5, color=TEXT)
            spacing(p, after=0, line=1.1)
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True  # keep the panel with the screenshot
        spacing(self.doc.add_paragraph(), after=4, line=1.0, keep_next=True)

    def callout(self, label: str, text: str, bar: str, fill: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        fixed_layout(table, [CONTENT_WIDTH_CM])
        no_split(table.rows[0])
        if self.doc.paragraphs:
            self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        cell = table.rows[0].cells[0]
        shade(cell, fill)
        cell_borders(cell, left=(36, bar))
        cell_margins(cell, top=110, bottom=110, left=200, right=160)
        p = cell.paragraphs[0]
        style_run(p.add_run(f"{label}  "), size=10, color=bar if bar != TEAL else "00857E", bold=True)
        add_rich(p, text, size=10)
        spacing(p, after=0, line=1.2)
        spacing(self.doc.add_paragraph(), after=4, line=1.0)

    def figure(self, key: str, caption: str, figures_dir: Path) -> None:
        path = figures_dir / f"{key}.png"
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(CONTENT_WIDTH_CM))
        spacing(p, before=6, after=2, keep_next=True)
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(c.add_run(caption), size=9, color=MUTED, italic=True)
        spacing(c, after=10)

    def shot(self, name: str, caption: str, figures_dir: Path, crop: str = "full") -> None:
        """A screen capture from tools/user-manual/screens, cropped and framed for the page."""
        path = prepare_shot(name, crop, figures_dir)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(CONTENT_WIDTH_CM if crop == "full" else 15.2))
        spacing(p, before=4, after=2, keep_next=True)
        self.figures += 1
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(c.add_run(f"Figure {self.figures}. {caption}"), size=9, color=MUTED, italic=True)
        spacing(c, after=10)

    def chapter(self, blocks, figures_dir: Path) -> None:
        for block in blocks:
            kind = block[0]
            if kind in ("h1", "h2", "h3"):
                self.heading(int(kind[1]), block[1])
            elif kind == "p":
                self.paragraph(block[1])
            elif kind == "bullets":
                self.bullets(block[1])
            elif kind == "steps":
                self.steps(block[1])
            elif kind == "table":
                self.table(block[1], block[2], block[3])
            elif kind == "tip":
                self.callout("Tip", block[1], TEAL, TIP_BG)
            elif kind == "note":
                self.callout("Note", block[1], NOTE_BAR, NOTE_BG)
            elif kind == "figure":
                self.figure(block[1], block[2], figures_dir)
            elif kind == "shot":
                self.shot(block[1], block[2], figures_dir, *block[3:])
            elif kind == "info":
                self.info(block[1])
            elif kind == "keep":
                self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
            else:
                raise ValueError(f"unknown block {kind}")


# ---------------------------------------------------------------- figure

def draw_cycle(path: Path) -> None:
    """The broking cycle as two rows of numbered stage boxes joined by arrows."""
    from PIL import Image, ImageDraw, ImageFont

    font_dir = Path("/usr/share/fonts/truetype/liberation")
    bold = ImageFont.truetype(str(font_dir / "LiberationSans-Bold.ttf"), 44)
    small = ImageFont.truetype(str(font_dir / "LiberationSans-Regular.ttf"), 31)
    num_font = ImageFont.truetype(str(font_dir / "LiberationSans-Bold.ttf"), 28)

    stages = [
        ("Client", "Onboarding and KYC"),
        ("Quotation", "or TSU proposal"),
        ("Account", "ARN, submitted"),
        ("Placement", "Slip to insurer"),
        ("Issuance", "E-policy to client"),
        ("Booking", "Invoice and entries"),
        ("Collection", "Receipts applied"),
        ("Remittance", "Premium to insurer"),
    ]
    width, height = 2000, 560
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    cols, box_w, box_h, gap = 4, 410, 190, 76
    left = (width - (cols * box_w + (cols - 1) * gap)) // 2
    tops = [20, 340]
    colours = ["#" + NAVY, "#" + BLUE]
    centres = []
    for i, (title, sub) in enumerate(stages):
        row, col = divmod(i, cols)
        if row == 1:
            col = cols - 1 - col  # snake back along the second row
        x, y = left + col * (box_w + gap), tops[row]
        d.rounded_rectangle([x, y, x + box_w, y + box_h], radius=22, fill=colours[row])
        d.rectangle([x + 22, y + box_h - 10, x + box_w - 22, y + box_h], fill="#" + TEAL)
        r = 22
        d.ellipse([x + 20, y + 20, x + 20 + 2 * r, y + 20 + 2 * r], fill="white")
        d.text((x + 20 + r, y + 20 + r), str(i + 1), font=num_font, fill=colours[row], anchor="mm")
        d.text((x + box_w / 2, y + 92), title, font=bold, fill="white", anchor="mm")
        d.text((x + box_w / 2, y + 140), sub, font=small, fill="#E6F4FD", anchor="mm")
        centres.append((x, y, col, row))

    arrow = "#" + TEAL

    def arrow_head(x, y, direction):
        s = 16
        if direction == "right":
            d.polygon([(x, y), (x - s, y - s), (x - s, y + s)], fill=arrow)
        elif direction == "left":
            d.polygon([(x, y), (x + s, y - s), (x + s, y + s)], fill=arrow)
        else:
            d.polygon([(x, y), (x - s, y - s), (x + s, y - s)], fill=arrow)

    for i in range(len(stages) - 1):
        x, y, col, row = centres[i]
        nx, ny, ncol, nrow = centres[i + 1]
        if row == nrow:
            cy = y + box_h // 2
            if ncol > col:
                d.line([x + box_w + 12, cy, nx - 14, cy], fill=arrow, width=6)
                arrow_head(nx - 8, cy, "right")
            else:
                d.line([x - 12, cy, nx + box_w + 14, cy], fill=arrow, width=6)
                arrow_head(nx + box_w + 8, cy, "left")
        else:
            cx = x + box_w // 2
            d.line([cx, y + box_h + 12, cx, ny - 14], fill=arrow, width=6)
            arrow_head(cx, ny - 8, "down")
    img.save(path, dpi=(300, 300))


def prepare_shot(name: str, crop: str, out_dir: Path) -> Path:
    """Crop a 1600x1000 capture ("full" keeps the menu, "content" drops it), trim the empty bottom band and
    save a JPEG so the Word file stays small."""
    from PIL import Image, ImageChops

    out = out_dir / f"{Path(name).stem}-{crop}.jpg"
    if out.exists():
        return out
    img = Image.open(SCREENS / name).convert("RGB")
    if crop == "content":
        img = img.crop((300, 0, img.width, img.height))
    # Trim rows at the bottom that are one flat colour (the page background under short screens).
    bg = Image.new("RGB", img.size, img.getpixel((img.width - 5, img.height - 5)))
    box = ImageChops.difference(img.crop((300 if crop == "full" else 0, 0, img.width, img.height)),
                                bg.crop((300 if crop == "full" else 0, 0, img.width, img.height))).getbbox()
    if box and box[3] < img.height - 40:
        img = img.crop((0, 0, img.width, max(box[3] + 24, 420)))
    img.save(out, "JPEG", quality=84, optimize=True)
    return out


# ---------------------------------------------------------------- build

def build(pages: dict[str, int] | None, figures_dir: Path, out: Path) -> list[tuple[int, str]]:
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    page_setup(section)
    header_footer(section)
    core = doc.core_properties
    core.title = f"{TITLE} {SUBTITLE}"
    core.subject = "User manual"
    core.author = "iorta TechNXT"
    core.keywords = "iNXT BrokerVerse, user manual, iorta TechNXT"

    cover(doc)
    builder = Builder(doc, pages)
    builder.document_control()

    # First collect the headings (a dry run of the numbering), then write the contents and the chapters.
    probe = Builder(Document(), None)
    for blocks in CHAPTERS:
        for block in blocks:
            if block[0] in ("h1", "h2"):
                probe.heading(int(block[0][1]), block[1])
    builder.contents(probe.headings)

    for blocks in CHAPTERS:
        builder.chapter(blocks, figures_dir)
    doc.save(out)
    return builder.headings


def to_pdf(docx: Path, out_dir: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise SystemExit("LibreOffice (soffice) is needed to paginate the table of contents")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx)],
                   check=True, capture_output=True)
    return out_dir / (docx.stem + ".pdf")


def heading_pages(pdf: Path, headings: list[tuple[int, str]]) -> dict[str, int]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(pdf))
    texts = []
    for page in document:
        textpage = page.get_textpage()
        lines = textpage.get_text_range().splitlines()
        texts.append([re.sub(r"\s+", " ", ln).strip() for ln in lines])
        textpage.close()
        page.close()
    document.close()
    pages: dict[str, int] = {}
    # The contents can run over several pages and lists every heading, so start at the page that carries the
    # first chapter heading for the last time (its contents line comes first).
    first = re.sub(r"\s+", " ", headings[0][1]).strip()
    start = max(i for i, lines in enumerate(texts) if first in lines)
    for _, label in headings:
        target = re.sub(r"\s+", " ", label).strip()
        for index in range(start, len(texts)):
            if target in texts[index]:
                pages[label] = index + 1
                start = index
                break
        else:
            raise SystemExit(f"heading not found in the PDF: {label}")
    return pages


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--pdf", action="store_true", help="keep the PDF next to the .docx")
    parser.add_argument("--previews", type=Path, help="write page PNGs to this folder")
    parser.add_argument("--edition", choices=sorted(EDITIONS), default="2.0", help="edition to build (default 2.0)")
    args = parser.parse_args()
    use_edition(args.edition)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / file_name()
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        draw_cycle(tmp_dir / "cycle.png")
        headings = build(None, tmp_dir, out)
        pages = heading_pages(to_pdf(out, tmp_dir), headings)
        build(pages, tmp_dir, out)
        pdf = to_pdf(out, tmp_dir)
        check = heading_pages(pdf, headings)
        if check != pages:
            raise SystemExit("page numbers moved between passes; run again")
        if args.pdf:
            shutil.copy(pdf, out.with_suffix(".pdf"))
        if args.previews:
            import pypdfium2 as pdfium

            args.previews.mkdir(parents=True, exist_ok=True)
            for i, page in enumerate(pdfium.PdfDocument(str(pdf)), 1):
                page.render(scale=1.3).to_pil().save(args.previews / f"page-{i:03d}.png")
    print(f"{out.relative_to(ROOT)}: {len(check)} headings, {max(check.values())} + pages")


if __name__ == "__main__":
    main()
